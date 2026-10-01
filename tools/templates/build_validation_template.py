"""Build the generic Word layout; never use this to rewrite a regional report.

Use the bundled workspace Python runtime. An optional output path supports
temporary QA copies; the default is the adjacent maintained template.
"""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def build(output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.03
    for name, size in [('Title', 23), ('Heading 1', 17), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)

    def paragraph(value, style=None):
        return doc.add_paragraph(value, style)

    def table(headers, rows, widths, label_column=False):
        result = doc.add_table(rows=1, cols=len(headers))
        result.alignment = WD_TABLE_ALIGNMENT.CENTER
        result.autofit = False
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
        for index, value in enumerate(headers):
            result.rows[0].cells[index].text = value
        for values in rows:
            row = result.add_row()
            for cell, value in zip(row.cells, values):
                lines = value if isinstance(value, list) else [value]
                cell.text = lines[0]
                for line in lines[1:]:
                    cell.add_paragraph(line)
        for row_index, row in enumerate(result.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if row_index == 0:
                properties.append(OxmlElement('w:tblHeader'))
            for column_index, cell in enumerate(row.cells):
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E5ECF0' if row_index == 0 else 'FFFFFF')
                properties.append(shade)
                margins = OxmlElement('w:tcMar')
                vertical_margin = '60' if label_column else '90'
                for side, value in [('top', vertical_margin), ('bottom', vertical_margin), ('left', '110'), ('right', '110')]:
                    element = OxmlElement('w:' + side)
                    element.set(qn('w:w'), value)
                    element.set(qn('w:type'), 'dxa')
                    margins.append(element)
                properties.append(margins)
                borders = OxmlElement('w:tcBorders')
                for side in ['top', 'left', 'bottom', 'right']:
                    element = OxmlElement('w:' + side)
                    for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
                        element.set(qn('w:' + key), value)
                    borders.append(element)
                properties.append(borders)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    for run in para.runs:
                        run.font.name, run.font.size = 'Calibri', Pt(10.5)
                        if row_index == 0 or (label_column and column_index == 0):
                            run.bold = True
        return result

    doc.add_heading('[Region name] model validation', 0)
    table(['Field', 'Entry'], [
        ('Region', '[Region name | identifier | Regional workbook]'),
        ('Catch source', ['[Provider; catch years; catch basis; date/version]', '[Region | Catch files]']),
        ('Selected article', ['[Authors (year). Title.]', '[DOI | Paper PDF | Paper folder]']),
        ('Other known articles', '[Author/year; short title; useful DOI or local record, one entry per article]'),
        ('Selected model', ['[Exact model ID/name; period; domain; source groups]', '[Folder | JSON | Saved runtime, if relevant]']),
        ('Other models from the same article', '[Other model: period; area; distinguishing features, one entry per model]'),
        ('Selection rationale', ['Article - [documented reason, or ?]', 'Model - [documented reason, or ?]']),
        ('Model extraction', ['[Material source conflicts and treatment | Useful reconstruction/conflict links]', '[Diet discrepancy and treatment; missing/default GS; derived BA; other material changes, where relevant]']),
        ('GE diagnostics', ['[Evidence-based verdict]', '[Negative SPPR finding; each affected source and group by name/ID when present]', 'rho_living = [value]', 'b = [value]', 'Detritus [pool name/ID]: SPPR = [value]', '[Material balance finding, if relevant]']),
        ('TE diagnostics', ['[Evidence-based verdict]', '[Negative SPPR finding; each affected source and group by name/ID when present]', 'rho_living = [value]', 'Detritus [pool name/ID]: SPPR = [value]', '[Material balance finding, if relevant]']),
        ('SPPR calculation', '[Researcher calculation choices and notes]'),
        ('Geographic fit', ['A. Region covered by study: [approximate value or range]%.', 'B. Study area covered by region: [approximate value or range]%.']),
        ('Temporal fit', '[Model period; catch years; material implications of transferring fixed coefficients or weights]'),
        ('Other', ''),
        ('Open issues and next action', ''),
        ('Review and reproducibility', 'Researcher name: [name] | review date: [dd/mm/yyyy]'),
    ], [1.53, 5.44], label_column=True)

    doc.add_page_break()
    doc.add_heading('Taxon mapping and coverage', 1)
    for value in [
        'Reference: [year and catch basis]; [source scope and taxon/group filters].',
        '[n] source groups [+ computational imports separately]; [n] catch taxa.',
        'Simple-chain PPR totals [value] t C in [year].',
        '[Missing TL/coefficient note only when needed; distinguish zero catch from unknown annual PPR.]',
        'Method: [simple-chain method name].',
        'Excel taxon appendix and descriptive Sources sheet: [relative hyperlink].',
    ]:
        paragraph(value)
    table(['Overall confidence', 'Taxa (n)', 'Catch (%)', 'Simple-chain PPR (%)'],
          [(level, '[n]', '[%]', '[%]') for level in ['High', 'Medium', 'Low', 'Very low', 'Unresolved']],
          [2.7, 1, 1.6, 1.65])
    paragraph('Overall confidence uses the weaker required membership or allocation component.')
    doc.add_heading('Group assignment rules', 2)
    table(['Plain-language rule', 'Confidence', 'PPR percentage'],
          [('[Actual rules, highest PPR share first]', '[Level]', '[%]')], [4.6, 1.05, 1.3])
    paragraph('Membership evidence: [concise descriptive sources and useful links].')
    doc.add_heading('Allocation weight rules', 2)
    table(['Plain-language rule', 'Confidence', 'PPR percentage'],
          [('[Actual allocation rule in full words]', '[Level]', '[%]')], [4.6, 1.05, 1.3])
    paragraph('Weights and assumptions: [actual basis and material transfer assumptions; source links].')
    paragraph('Each percentage is the independent simple-chain PPR associated with taxa using that rule.')
    paragraph('Unresolved taxa: [names and short reasons, or none]. [n] Very low decisions cover [%] of catch and [%] of simple-chain PPR.')
    doc.add_heading('Very low decisions', 2)
    table(['Affected taxa', 'Why confidence is very low'],
          [('[Exact taxa sharing one evidenced reason]', '[Region/model-specific composition, analogue or allocation uncertainty]')], [2.29, 4.66])

    doc.add_page_break()
    doc.add_heading('Geographic evidence screenshots', 1)
    paragraph('Target region R and selected model study area S. [Overlap method and approximation, if applicable.]')
    for heading, header, caption in [
        ('Region area', 'Target region boundary', '[Region name/ID; boundary source/version | Source | Image]'),
        ('Study area in the article', 'Study boundary for the selected model', '[Article; figure/page; represented model area | Source PDF | Image]'),
    ]:
        doc.add_heading(heading, 2)
        table([header], [('[Insert readable boundary image]\n\n\n\n',)], [6.97])
        paragraph(caption)
    paragraph('[If a figure is unavailable, state the missing source or reason.]')
    doc.core_properties.title = 'Regional model validation template'
    doc.core_properties.subject = 'Concise evidence summary for researcher review'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', nargs='?', type=Path, default=Path(__file__).with_name('Model_validation_template.docx'))
    build(parser.parse_args().output)
