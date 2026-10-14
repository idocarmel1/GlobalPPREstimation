"""Complete, content-checked SPPR inventories for registered validated models.

Coverage is an execution record, not researcher approval. Each method must have
fresh saved coefficients or a specific unsuccessful outcome on the same inputs.
"""
from __future__ import annotations

import copy
import json
import os
import tempfile
import zipfile
from datetime import datetime, timezone
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path


def validated_identities(metadata):
    identities = [(r['unit_id'], r['model_id']) for r in metadata
                  if r.get('researcher_review_status') == 'Validated by researcher']
    if len(identities) != len(set(identities)):
        raise ValueError('Duplicate registered validated model identity')
    return sorted(identities)


def plan_methods(expected, context, saved, row_hashes):
    if saved.get('context') != context:
        return list(expected)
    pending = []
    for key in expected:
        outcome = saved.get('methods', {}).get(key, {})
        terminal = outcome.get('status') in {'ok', 'reused', 'failed', 'timeout', 'crashed', 'unavailable', 'NOT_RUN'}
        reason_ok = outcome.get('status') in {'ok', 'reused'} or bool(outcome.get('reason'))
        if (not terminal or not reason_ok or key not in row_hashes
                or outcome.get('group_rows_sha256') != row_hashes[key]):
            pending.append(key)
    return pending


def _same_context(saved, current):
    """Adapter metadata alone cannot invalidate compatible numerical evidence."""
    if not isinstance(saved,dict) or not isinstance(current,dict):return saved==current
    return {key:value for key,value in saved.items() if key!='adapter_sha256'} == {
        key:value for key,value in current.items() if key!='adapter_sha256'}


def _planning_coverage(coverage,context):
    return {**coverage,'context':context} if _same_context(coverage.get('context'),context) else coverage


def _json(value):
    if isinstance(value, dict):
        return value
    try:
        return json.loads(value) if value else {}
    except (TypeError, ValueError):
        return {}


def _metadata(root):
    from tools.project_core.validation.researcher_review import table_rows
    with zipfile.ZipFile(root / 'Project.xlsx') as archive:
        return [r for _, r in table_rows(archive, 'Models & coverage', 'Models')[3]]


def _settings_book(path):
    from tools.project_core.validation.researcher_review import table_rows
    with zipfile.ZipFile(path) as archive:
        entries = [r for _, r in table_rows(archive, 'Overview', 'Settings')[3]]
    return {'Overview': {'Settings': (['field', 'value'], [[r['field'], r.get('value')] for r in entries])}}


def _context(root, model, flags):
    from tools.project_core.workbooks.workbooks import sha
    from tools.scientific_helpers.sppr.method_coverage import engine_identity
    packages = {}
    for name in ('numpy', 'pandas', 'scipy', 'sympy', 'igraph'):
        try:
            packages[name] = version(name)
        except PackageNotFoundError:
            packages[name] = None
    return {'model_sha256': sha(model), 'flags': flags,
            'code_identity': engine_identity(flags, model),
            'adapter_sha256': sha(root / 'tools/scientific_helpers/sppr/method_coverage.py'),
            'dependencies': packages}


def _arithmetic_identity():
    from tools.project_core.workbooks.workbooks import sha
    core = Path(__file__).resolve().parents[1]
    return {name:sha(core/name) for name in
            ('calculations/coverage.py','calculations/coverage_candidates.py',
             'calculations/ee_floor_migration.py','calculations/ee_floor_execution.py',
             'calculations/regional.py','workbooks/workbooks.py')}


def _method_rows(book):
    from tools.project_core.workbooks.workbooks import rows, digest_tables
    result = {}
    for row in rows(book, 'Selected model groups', 'Group SPPR'):
        result.setdefault(row[3], []).append(row)
    return result, {key: digest_tables(value) for key, value in result.items()}


_UNSET = object()


def _write_bytes(path, data, *, expected_sha256=_UNSET):
    from tools.project_core.workbooks.workbooks import sha
    path.parent.mkdir(parents=True, exist_ok=True)
    before=sha(path) if path.is_file() else None
    if expected_sha256 is not _UNSET and before!=expected_sha256:
        raise ValueError('Output changed before replacement: '+str(path))
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=path.suffix); os.close(fd)
    try:
        Path(temporary).write_bytes(data)
        if (sha(path) if path.is_file() else None)!=before:
            raise ValueError('Output changed during replacement: '+str(path))
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def _write_json(path, value, *, expected_sha256=_UNSET):
    _write_bytes(path,json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False).encode('utf-8'),expected_sha256=expected_sha256)


def _cached_methods(model_dir, coverage, context, book, expected):
    from tools.project_core.workbooks.workbooks import sha
    _, row_hashes = _method_rows(book)
    pending = plan_methods(expected, context, _planning_coverage(coverage,context), row_hashes)
    for evidence in coverage.get('execution_evidence',{}).values():
        path = (model_dir/'results/diagnostics'/str(evidence.get('path',''))).resolve()
        if not path.is_relative_to(model_dir.resolve()) or not path.is_file() or sha(path)!=evidence.get('sha256'):
            return list(expected)
    for key in expected:
        if key in pending:
            continue
        evidence = coverage['methods'][key].get('evidence') or {}
        path = (model_dir / 'results/diagnostics' / str(evidence.get('path', ''))).resolve()
        if (not path.is_relative_to(model_dir.resolve()) or not path.is_file()
                or sha(path) != evidence.get('sha256')):
            pending.append(key)
    return [key for key in expected if key in pending]


def _package_cache_valid(model_dir, path, manifest, context, expected):
    """Unchanged exact workbook bytes reuse the previously verified inventory."""
    from tools.project_core.workbooks.workbooks import sha
    coverage = manifest.get('method_coverage', {})
    if coverage.get('arithmetic_code_identity') != _arithmetic_identity():
        return False
    if coverage.get('expected_methods') != list(expected):
        return False
    if sha(path) != manifest.get('snapshot_sha256'):
        return False
    snapshot = model_dir/'results/regional_snapshot.xlsx'
    if not snapshot.is_file() or sha(snapshot) != manifest.get('snapshot_sha256'):
        return False
    source = model_dir/'sppr_source.xlsx'
    if not source.is_file() or sha(source) != manifest.get('coefficient_source_sha256'):
        return False
    hashes = {key:value.get('group_rows_sha256') for key,value in coverage.get('methods', {}).items()}
    if any(not value for value in hashes.values()) or plan_methods(expected, context, _planning_coverage(coverage,context), hashes):
        return False
    references = [coverage['methods'][key].get('evidence') or {} for key in expected]
    references.extend(coverage.get('execution_evidence',{}).values())
    for evidence in references:
        artifact = (model_dir/'results/diagnostics'/str(evidence.get('path', ''))).resolve()
        if not artifact.is_relative_to(model_dir.resolve()) or not artifact.is_file() or sha(artifact) != evidence.get('sha256'):
            return False
    return True


def _source_matches(model_dir, book, expected):
    """A diagnostic-only exporter may overwrite only the source workbook."""
    from tools.project_core.workbooks.workbooks import records, digest_tables
    from tools.scientific_helpers.sppr.ppr_scopes import read_scopes
    from tools.scientific_helpers.sppr.retained_methods import _saved_cell_equal
    source = model_dir / 'sppr_source.xlsx'
    if not source.is_file():
        return False
    scopes = read_scopes(source)
    for scope in ('all', 'inner', 'PP'):
        methods, groups = scopes.get(scope, ([], {}))
        if not set(expected).issubset(methods):
            return False
        retained = {(r['group'], r['method']): r['sppr'] for r in
                    records(book, 'Selected model groups', 'Group SPPR') if r['scope'] == scope}
        for method in expected:
            index = methods.index(method)
            actual = {(group, method): values[index] for group, values in groups.items()}
            wanted = {key: value for key, value in retained.items() if key[1] == method}
            if actual.keys() != wanted.keys() or any(not _saved_cell_equal(actual[key], value) for key, value in wanted.items()):
                return False
    return True


def _update_source(path, book, outcomes, *, replace_methods=(), effective_groups=None, expected_sha256=_UNSET):
    """Append missing scoped columns without rewriting retained source cells."""
    import openpyxl
    from xml.etree import ElementTree as etree
    from tools.project_core.validation.researcher_review import sheet_xml, S
    from tools.project_core.workbooks.workbooks import records, clean, sha
    baseline=sha(path) if path.is_file() else None
    if expected_sha256 is not _UNSET and baseline!=expected_sha256:
        raise ValueError('Source workbook changed before replacement')
    workbook = openpyxl.load_workbook(path) if path.is_file() else openpyxl.Workbook()
    def formula(cell):
        value = cell.find(f'{{{S}}}f')
        return None if value is None else (value.text, sorted(value.attrib.items()))
    original_numbers = {}; changed_cells = {}; replace_methods = set(replace_methods)
    if path.is_file():
        with zipfile.ZipFile(path) as archive:
            for name in workbook.sheetnames:
                if name == 'method_coverage': continue
                _, sheet = sheet_xml(archive,name)
                original_numbers[name] = {cell.get('r'):(cell.find(f'{{{S}}}v').text,formula(cell))
                    for cell in sheet.iter(f'{{{S}}}c') if cell.get('t') in (None,'n')
                    and cell.find(f'{{{S}}}v') is not None}
    if workbook.sheetnames == ['Sheet']:
        workbook.remove(workbook['Sheet'])
    groups = records(book, 'Selected model groups', 'Groups')
    sequences = {r.get('group_name'): r.get('seq') for r in groups}
    methods = list(dict.fromkeys(r['method'] for r in records(book, 'Selected model groups', 'Group SPPR')))
    added_methods = set()
    for scope in ('all', 'inner', 'PP'):
        name = 'sppr_' + scope
        if name in workbook:
            sheet = workbook[name]
        else:
            sheet = workbook.create_sheet(name)
            sheet.append(['seq', 'group_name'])
        header = [cell.value for cell in sheet[1]]
        if header.count('group_name') != 1:
            raise ValueError('Saved source has no unique group-name column: ' + name)
        group_column = header.index('group_name') + 1
        positions = {sheet.cell(row,group_column).value:row for row in range(2,sheet.max_row+1)
                     if sheet.cell(row,group_column).value is not None}
        values = {}
        for row in records(book, 'Selected model groups', 'Group SPPR'):
            if row['scope'] == scope:
                values.setdefault(row['group'], {})[row['method']] = row['sppr']
        for method in methods:
            if method in header and method not in replace_methods:
                continue
            added_methods.add(method)
            if method in header:
                column = header.index(method) + 1
            else:
                column = len(header) + 1; header.append(method)
                sheet.cell(1,column,method)
            for group, coefficients in values.items():
                if group not in positions:
                    positions[group] = sheet.max_row + 1
                    sheet.cell(positions[group],group_column,group)
                    if 'seq' in header:
                        sheet.cell(positions[group],header.index('seq')+1,sequences.get(group))
                sheet.cell(positions[group],column).value=coefficients.get(method)
                changed_cells.setdefault(name,set()).add(sheet.cell(positions[group],column).coordinate)
    if effective_groups is not None:
        if 'groups_df' in workbook: workbook.remove(workbook['groups_df'])
        sheet = workbook.create_sheet('groups_df')
        frame = effective_groups.reset_index()
        sheet.append(list(frame.columns))
        for record in frame.itertuples(index=False,name=None):
            sheet.append([clean(value) for value in record])
        original_numbers.pop('groups_df',None)
    name = 'method_coverage'
    if name in workbook:
        workbook.remove(workbook[name])
    sheet = workbook.create_sheet(name)
    sheet.append(['method', 'status', 'reason', 'all_status', 'inner_status', 'PP_status', 'evidence'])
    for key, outcome in outcomes.items():
        sheet.append([key, outcome['status'], outcome.get('reason'),
                      *[outcome.get('annual_status', {}).get(scope) for scope in ('all', 'inner', 'PP')],
                      outcome.get('evidence', {}).get('path')])
    for name in ('model_health', 'mc_diagnostics'):
        header, values = book.get('Diagnostics', {}).get(name, ([], []))
        if name in workbook:
            options = {'new_GE':'GE','new_TE_EEfix':'TE','new_WithEgestion':'With Egestion'}
            relevant = ({options[key] for key in added_methods if key in options} if name=='model_health'
                        else {key for key in added_methods if key.startswith('MC_')})
            if not relevant:
                continue
            if replace_methods:
                workbook.remove(workbook[name]); sheet=workbook.create_sheet(name)
                sheet.append(header)
                for row in values:sheet.append([clean(value) for value in row])
                original_numbers.pop(name,None)
                continue
            selected = [row for row in values if any(dict(zip(header,row)).get(key) in relevant
                        for key in ('TE_option','config_TE_option','method'))]
            sheet = workbook[name]
            existing_header = [cell.value for cell in sheet[1]]
            for key in header:
                if key not in existing_header:
                    existing_header.append(key); sheet.cell(1,len(existing_header),key)
            for row in selected:
                record = dict(zip(header,row))
                sheet.append([clean(record.get(key)) for key in existing_header])
        else:
            sheet = workbook.create_sheet(name)
            sheet.append(header)
            for row in values:
                sheet.append([clean(value) for value in row])
    if 'run_notes' not in workbook:
        workbook.create_sheet('run_notes').append(['topic','note','status'])
    note = workbook['run_notes']
    note.append(['method coverage', 'Per-method outcomes and retained or executed evidence are in method_coverage and results/diagnostics/method_coverage.', 'recorded'])
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix='.xlsx'); os.close(fd)
    try:
        workbook.save(temporary); workbook.close()
        # openpyxl serializes even untouched floats at 16 digits. Restore the
        # original numeric text, including unchanged formula caches, by address.
        with zipfile.ZipFile(temporary) as archive:
            replacements = {}
            for name, values in original_numbers.items():
                member, sheet = sheet_xml(archive,name); changed = False
                for cell in sheet.iter(f'{{{S}}}c'):
                    if cell.get('r') in changed_cells.get(name,set()): continue
                    previous = values.get(cell.get('r'))
                    if previous is None or cell.get('t') not in (None,'n') or formula(cell) != previous[1]: continue
                    value = cell.find(f'{{{S}}}v')
                    if value is None: value = etree.SubElement(cell,f'{{{S}}}v')
                    if value.text != previous[0]: value.text = previous[0]; changed = True
                if changed: replacements[member] = etree.tostring(sheet,encoding='utf-8')
            parts = [(info,replacements.get(info.filename,archive.read(info.filename))) for info in archive.infolist()]
        if replacements:
            with zipfile.ZipFile(temporary,'w') as archive:
                for info,data in parts: archive.writestr(info,data)
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip() is not None: raise ValueError('Updated source workbook ZIP failed verification')
        if (sha(path) if path.is_file() else None)!=baseline:
            raise ValueError('Source workbook changed during replacement')
        os.replace(temporary, path)
    finally:
        workbook.close(); Path(temporary).unlink(missing_ok=True)


def _apply_outcomes(book, model_id, outcomes, pending, *, migration=False):
    from tools.project_core.workbooks.workbooks import rows, records, ANNUAL_HEADER, finite, digest_tables
    # Adopting evidence for a saved method is not a new scientific execution.
    # Its coefficients, annual restrictions and diagnostics stay exactly saved.
    pending = [key for key in pending if not outcomes[key].get('preserve_saved')]
    current = rows(book, 'Selected model groups', 'Group SPPR')
    old_status = {(row[1],row[2]):str(row[6]) for row in rows(book,'PPR','Annual')
                  if row[3]=='landings' and row[4]=='method' and row[5]=='ppr'}
    unchanged_scopes = {}
    for row in current:
        if row[3] not in pending:
            continue
        key = (row[2],row[3])
        value = outcomes[row[3]].get('scopes',{}).get(row[2],{}).get(row[1])
        value = value if finite(value) else None
        unchanged_scopes[key] = unchanged_scopes.get(key,True) and digest_tables(value)==digest_tables(row[4])
    group_names = [r['group_name'] for r in records(book, 'Selected model groups', 'Groups')]
    # Preserve row order and exact unchanged reviewed values when extending.
    retained = []
    present = set()
    for row in current:
        updated = list(row)
        group, scope, method = row[1:4]
        present.add((group, scope, method))
        if method in pending:
            value = outcomes[method].get('scopes', {}).get(scope, {}).get(group)
            value = value if finite(value) else None
            if digest_tables(value) != digest_tables(row[4]):
                updated[4] = value
        retained.append(updated)
    for method in pending:
        outcome = outcomes[method]
        for scope in ('all', 'inner', 'PP'):
            values = outcome.get('scopes', {}).get(scope, {})
            for group in group_names:
                if (group, scope, method) in present:
                    continue
                value = values.get(group)
                retained.append([model_id, group, scope, method, value if finite(value) else None])
    book.setdefault('Selected model groups', {})['Group SPPR'] = (
        ['model_id', 'group', 'scope', 'method', 'sppr'], retained)
    annual = [row for row in rows(book, 'PPR', 'Annual') if row[2] not in pending]
    for method in pending:
        for scope in ('all', 'inner', 'PP'):
            status = outcomes[method].get('annual_status', {}).get(scope, 'unavailable: no scoped outcome')
            executed_status = status
            prior = old_status.get((scope,method),'')
            if migration and prior and prior != 'ok':
                values = outcomes[method].get('scopes',{}).get(scope,{}).values()
                if prior.startswith('provisional:') and any(finite(value) for value in values):
                    status = prior if status=='ok' or status in prior else prior+'; refreshed diagnosis: '+status
                elif not prior.startswith('provisional:'):
                    status = prior if status=='ok' or status in prior else prior+'; refreshed diagnosis: '+status
            elif not method.startswith('MC_'):
                if status == 'ok' and prior.startswith(('FAIL','DIVERGED','NOT_RUN')):
                    status = prior
                elif (prior.startswith('provisional:') and unchanged_scopes.get((scope,method))
                      and any(finite(value) for value in outcomes[method].get('scopes',{}).get(scope,{}).values())
                      and not status.startswith('unavailable:')):
                    status = prior if status=='ok' or status in prior else prior+'; refreshed diagnosis: '+status
            if status in {'unavailable: NOT_RUN','unavailable: failed','unavailable: timeout','unavailable: crashed'}:
                status += '; ' + outcomes[method].get('reason','No further execution detail')
            if status != executed_status:
                outcomes[method].setdefault('execution_annual_status',{})[scope] = executed_status
                outcomes[method].setdefault('annual_status',{})[scope] = status
            annual.append([model_id, scope, method, 'landings', 'method', 'ppr', status, *[None]*70])
    book.setdefault('PPR', {})['Annual'] = (ANNUAL_HEADER, annual)
    book.setdefault('Diagnostics', {})['Method coverage'] = (
        ['method', 'status', 'reason', 'evidence'],
        [[key, value['status'], value.get('reason'), value.get('evidence', {}).get('path')]
        for key, value in outcomes.items()])
    from tools.project_core.workbooks.workbooks import table_dict
    health = {r.get('TE_option') or r.get('config_TE_option'): r
              for r in records(book, 'Diagnostics', 'model_health')}
    mc = {r['method']: r for r in records(book, 'Diagnostics', 'mc_diagnostics')}
    exact = {'new_GE':'GE', 'new_TE_EEfix':'TE', 'new_WithEgestion':'With Egestion'}
    for method in pending:
        outcome = outcomes[method]
        if method in exact:
            diagnostic = outcome.get('diagnostics')
            if outcome.get('retained_diagnostics_flat'):
                health[exact[method]] = copy.deepcopy(outcome['retained_diagnostics_flat'])
            elif diagnostic:
                from tools.scientific_helpers.sppr.method_coverage import _engine, diagnostics_for_table
                flat = _engine()[2]._flatten_diagnostics(diagnostics_for_table(diagnostic))
                health[exact[method]] = {'TE_option':exact[method], **flat}
            else:
                health[exact[method]] = {'TE_option':exact[method], 'status':'NOT_RUN', 'error':outcome.get('reason')}
        if method.startswith('MC_'):
            diagnostic = outcome.get('mc_diagnostics') or {}
            mc[method] = {'method':method, **diagnostic}
    if health and any(method in exact for method in pending):
        book['Diagnostics']['model_health'] = table_dict(health.values())
    if mc and any(method.startswith('MC_') for method in pending):
        book['Diagnostics']['mc_diagnostics'] = table_dict(mc.values())


def _preserved_inputs(book):
    from tools.project_core.workbooks.workbooks import digest_tables
    return digest_tables([(sheet, name, data) for sheet, tables in book.items()
                          for name, data in tables.items()
                          if (sheet in {'Catch', 'NPP'} or (sheet, name) in {
                              ('Classic PPR', 'Taxa'), ('Selected model groups', 'Groups'),
                              ('PPR', 'Matching'), ('PPR', 'Allocation assumptions'), ('PPR', 'Mapping review')})])


def _arithmetic(book, path, model, flags, original_book=None, *, migration=False, invalidated_methods=()):
    from tools.project_core.workbooks.workbooks import overview, sha, rows, digest_tables, input_hash
    from tools.project_core.calculations.regional import recalculate, set_setting, set_result_hash, result_hash
    old_overview = copy.deepcopy(overview(book))
    classic = copy.deepcopy(book.get('Classic PPR', {}))
    # A candidate snapshot can honestly retain a different historical selection;
    # operate on its actual results identity without changing the regional choice.
    set_setting(book, 'selected_model_id', model.parent.name)
    set_setting(book, 'model_path', os.path.relpath(model, path.parent).replace('\\', '/'))
    set_setting(book, 'results_model_id', model.parent.name)
    set_setting(book, 'results_model_sha256', sha(model))
    if flags is not None:
        set_setting(book, 'effective_loader_flags', json.dumps(flags, ensure_ascii=False))
    recalculate(book, path)
    book['Classic PPR'] = classic
    # Recalculate restored independent-series ratios from those actual saved
    # numerators, never from an intermediate Classic table discarded above.
    from tools.project_core.workbooks.workbooks import records, YEARS, finite
    denominator = {r['method']:[r.get(year) for year in YEARS] for r in records(book,'NPP','NPP')}
    ratio_header, ratio_rows = book['PPR–NPP']['Ratios']
    ratio_rows[:] = [r for r in ratio_rows if r[0]]
    for row in rows(book,'Classic PPR','Annual'):
        if row[5] != 'ppr': continue
        for method, values in denominator.items():
            ratio_rows.append([*row[:5],method,*[100*p/9/n if finite(p) and finite(n) and n>0 else None
                                               for p,n in zip(row[7:],values)]])
    if original_book is not None and _preserved_inputs(original_book) == _preserved_inputs(book):
        old_groups, new_groups = {}, {}
        for target, value in ((old_groups,original_book), (new_groups,book)):
            for row in rows(value,'Selected model groups','Group SPPR'):
                target.setdefault((row[2],row[3]),[]).append(row)
        unchanged = {key for key in old_groups.keys() & new_groups.keys()
                     if key[1] not in invalidated_methods and digest_tables(old_groups[key]) == digest_tables(new_groups[key])}
        previous = overview(original_book)
        input_fresh = not previous.get('calculation_input_sha256') or previous['calculation_input_sha256'] == input_hash(original_book)
        result_fresh = not previous.get('calculation_result_sha256') or previous['calculation_result_sha256'] == result_hash(original_book)
        if input_fresh and result_fresh:
            # A method extension is not permission to revise saved arithmetic,
            # catch-specific diagnostic restrictions or sensitivity intervals.
            # Keep their exact rows and append only newly calculated scopes.
            for sheet, table in [('PPR','Annual'), ('PPR','Taxon SPPR'),
                                 ('PPR','Taxon PPR inspected year'), ('PPR–NPP','Ratios')]:
                if table not in original_book.get(sheet, {}): continue
                if table == 'Taxon PPR inspected year' and any(
                        previous.get(key, default) != old_overview.get(key, default)
                        for key, default in [('taxon_detail_year',2019), ('catch_basis','landings')]):
                    continue
                old_header, old_rows = original_book[sheet][table]
                new_header, generated = book[sheet][table]
                if old_header != new_header:
                    raise ValueError('Cannot preserve saved arithmetic with a changed table schema: ' + sheet + '/' + table)
                scope_column, method_column = old_header.index('scope'), old_header.index('method')
                model_column = old_header.index('model_id')
                def eligible(row):
                    return ((row[scope_column], row[method_column]) in unchanged or
                            (sheet == 'PPR–NPP' and not row[model_column]))
                retained = [copy.deepcopy(row) for row in old_rows if eligible(row)]
                represented = {(row[model_column], row[scope_column], row[method_column]) for row in retained}
                remaining = [row for row in generated
                             if (row[model_column], row[scope_column], row[method_column]) not in represented]
                book[sheet][table] = (old_header, retained + remaining)
        elif input_fresh:
            # Historical callers can append sensitivity bounds after computing
            # their result hash. Retain only those bounds when removing them
            # exactly recovers the otherwise fresh recorded output hash.
            without_bounds = copy.deepcopy(original_book)
            header, annual = without_bounds.get('PPR', {}).get('Annual', ([], []))
            without_bounds.setdefault('PPR', {})['Annual'] = (header, [r for r in annual if r[5] not in {'min_tC','max_tC'}])
            if previous.get('calculation_result_sha256') == result_hash(without_bounds):
                old_status = {(r[1],r[2],r[3],r[4]):r[6] for r in rows(original_book,'PPR','Annual') if r[5]=='ppr'}
                new_status = {(r[1],r[2],r[3],r[4]):r[6] for r in rows(book,'PPR','Annual') if r[5]=='ppr'}
                for row in rows(original_book,'PPR','Annual'):
                    key = (row[1],row[2]); status_key = (*key,row[3],row[4])
                    if row[5] in {'min_tC','max_tC'} and key in unchanged and old_status.get(status_key) == new_status.get(status_key):
                        book['PPR']['Annual'][1].append(copy.deepcopy(row))
    for key in ('selected_model_id', 'selection_rationale', 'model_path'):
        set_setting(book, key, old_overview.get(key))
    if migration and original_book is not None:
        from tools.project_core.calculations.ee_floor_migration import preserve_restrictions
        preserve_restrictions(book,original_book)
    set_result_hash(book)


def _restore_retained_numbers(path, requested, original):
    """Undo writer rounding only in exact retained rows, wherever they moved.

    Python's numeric repr round-trips the desired original value even when the
    existing XLSX already lost a digit. New or changed rows are never restored.
    """
    from xml.etree import ElementTree as etree
    from tools.project_core.validation.researcher_review import sheet_xml, column_index, S
    from tools.project_core.workbooks.workbooks import finite, sha
    if original is None: return False
    def identity(row): return tuple((type(value).__name__, repr(value)) for value in row)
    wanted = {}
    for name, tables in requested.items():
        for table, (header, values) in tables.items():
            prior_header, prior_rows = original.get(name, {}).get(table, ([], []))
            if prior_header != header: continue
            prior = {identity(row) for row in prior_rows}
            retained = {index: {column: repr(value) for column, value in enumerate(row, 1) if finite(value)}
                        for index, row in enumerate(values) if identity(row) in prior}
            retained = {index: cells for index, cells in retained.items() if cells}
            if retained: wanted.setdefault(name, {})[table] = (header, len(values), retained)
    if not wanted: return False
    baseline = sha(path)
    replacements = {}
    with zipfile.ZipFile(path) as archive:
        shared = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            shared = [''.join(item.itertext()) for item in etree.fromstring(archive.read('xl/sharedStrings.xml'))]
        def text(cell):
            if cell is None: return None
            if cell.get('t') == 'inlineStr': return ''.join(node.text or '' for node in cell.findall(f'.//{{{S}}}t'))
            value = cell.find(f'{{{S}}}v')
            value = value.text if value is not None else None
            return shared[int(value)] if cell.get('t') == 's' and value is not None else value
        for name, tables in wanted.items():
            member, sheet = sheet_xml(archive, name)
            current = None; header_pending = False; index = 0; changed = False; seen = set()
            for row in sheet.findall(f'{{{S}}}sheetData/{{{S}}}row'):
                cells = {column_index(cell.get('r')): cell for cell in row}
                values = {column: text(cell) for column, cell in cells.items()}
                if not any(value not in (None, '') for value in values.values()): continue
                if values.get(1) == '@table':
                    current = values.get(2); header_pending = True; index = 0
                    continue
                if current not in tables: continue
                header, count, retained = tables[current]
                if header_pending:
                    if [values.get(column) for column in range(1, len(header)+1)] != [str(value) if value is not None else None for value in header]:
                        raise ValueError('Saved table header changed during numeric preservation: ' + name + '/' + current)
                    header_pending = False
                    continue
                for column, number in retained.get(index, {}).items():
                    cell = cells.get(column)
                    value = cell.find(f'{{{S}}}v') if cell is not None else None
                    if (cell is None or cell.get('t') not in (None, 'n') or cell.find(f'{{{S}}}f') is not None or
                            value is None or float(value.text) not in (float(number), float(format(float(number), '.16g')))):
                        raise ValueError('Saved numeric cell does not match the requested retained row')
                    if value.text != number: value.text = number; changed = True
                    seen.add((current, index, column))
                index += 1
            expected = {(table, index, column) for table, (_, _, retained) in tables.items()
                        for index, cells in retained.items() for column in cells}
            if seen != expected: raise ValueError('Saved retained numeric rows were not found: ' + name)
            if changed: replacements[member] = etree.tostring(sheet, encoding='utf-8')
        if not replacements: return False
        parts = [(info, replacements.get(info.filename, archive.read(info.filename))) for info in archive.infolist()]
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix='.xlsx'); os.close(fd)
    try:
        with zipfile.ZipFile(temporary, 'w') as archive:
            for info, data in parts: archive.writestr(info, data)
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip() is not None: raise ValueError('Retained numeric workbook ZIP failed verification')
        if sha(path) != baseline: raise ValueError('Workbook changed during retained numeric preservation')
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return True


def _save_calculation(path, book, model, flags, original_book, *, pending_reason=None, on_write=None, before_write=None, migration=False, invalidated_methods=()):
    """Certify persisted tables, including Excel's signed-zero normalization."""
    from tools.project_core.workbooks.workbooks import update_book, read_book, overview, input_hash
    from tools.project_core.calculations.regional import set_result_hash, result_hash
    def save(value):
        if before_write is not None:before_write()
        update_book(path,value)
        if on_write is not None:on_write()
        if _restore_retained_numbers(path,value,original_book) and on_write is not None:on_write()
        return read_book(path)
    saved = save(book)
    if overview(saved).get('calculation_input_sha256') != input_hash(saved):
        if pending_reason:
            from tools.project_core.calculations.coverage_candidates import suppress_candidate_arithmetic
            suppress_candidate_arithmetic(saved,pending_reason)
        else:
            _arithmetic(saved,path,model,flags,original_book,migration=migration,invalidated_methods=invalidated_methods)
        saved = save(saved)
    # Rounded tiny negative taxon coefficients can serialize as -0 and reload
    # as integer 0. Hash these actual saved outputs without altering historical
    # input/review hash rules or repeating the scientific calculation.
    if overview(saved).get('calculation_result_sha256') != result_hash(saved):
        set_result_hash(saved)
        saved = save(saved)
    settings = overview(saved)
    if (settings.get('calculation_input_sha256') != input_hash(saved)
            or settings.get('calculation_result_sha256') != result_hash(saved)):
        raise ValueError('Persisted calculation fingerprints are not stable')
    return saved


def _complete_model(root, metadata, expected, *, progress, compute, ee_floor=None):
    from tools.project_core.registry.discovery import region_directory, resolve_model
    from tools.project_core.workbooks.workbooks import read_book, overview, sha, input_hash, update_book
    from tools.project_core.calculations.snapshots import dependencies
    region = region_directory(root, metadata['unit_id'])
    model = resolve_model(region, metadata['model_id']); model_dir = model.parent
    regional_path = region / (region.name + '.xlsx')
    local = _settings_book(regional_path)
    local_settings = overview(local)
    is_current = (local_settings.get('selected_model_id') == model_dir.name
                  and local_settings.get('results_model_id') == model_dir.name)
    path = regional_path if is_current else model_dir / 'results/regional_snapshot.xlsx'
    if not path.is_file():
        raise ValueError(f"{metadata['unit_id']}/{model_dir.name}: missing model-local calculation inputs")
    settings_book = local if is_current else _settings_book(path)
    settings = overview(settings_book)
    if settings.get('results_model_id') != model_dir.name:
        raise ValueError('Saved calculation workbook belongs to another model: ' + str(path))
    current_model_hash = sha(model)
    recorded_model_hashes = [settings.get('results_model_sha256'), metadata.get('reviewed_model_sha256')]
    if any(value and value != current_model_hash for value in recorded_model_hashes):
        raise ValueError('Model extraction changed; an explicit scientific refresh is required before method completion')
    manifest_path = model_dir / 'results/result_manifest.json'
    manifest = _json(manifest_path.read_text(encoding='utf-8')) if manifest_path.is_file() else {}
    from tools.scientific_helpers.sppr.configuration import resolve_audited_configuration
    configuration_error = None
    try:
        flags = resolve_audited_configuration(model_dir, settings_book, manifest)
    except ValueError:
        # Legacy settings/sample counts may require the full saved tables.
        flags = None
    requested_change = False
    if ee_floor and flags is not None:
        from tools.project_core.calculations.ee_floor_migration import target_flags
        from tools.scientific_helpers.sppr.method_coverage import ENGINE
        requested_change = target_flags(flags,ee_floor['use_ee_floor'],ee_floor['ee_floor'],sha(ENGINE/'PPRCalculator.py')) != flags
    if is_current and manifest.get('method_coverage') and flags is not None and not requested_change:
        context = _context(root, model, flags)
        if _package_cache_valid(model_dir, path, manifest, context, expected):
            progress(f"{metadata['unit_id']}/{model_dir.name}: reusing complete fresh {len(expected)}-method package")
            return False, {'unit_id': metadata['unit_id'], 'model_id': model_dir.name,
                           'reused': True, 'methods': manifest['method_coverage']['methods']}
    consumed_hashes = {path:sha(path), model:sha(model), root/'Project.xlsx':sha(root/'Project.xlsx')}
    book = read_book(path)
    if sha(path)!=consumed_hashes[path]:
        raise ValueError('Calculation workbook changed while loading method inputs')
    original_book = copy.deepcopy(book)
    candidate_changed = False; candidate_pending = None
    if not is_current:
        from tools.project_core.calculations.coverage_candidates import prepare_candidate_book
        consumed_hashes[regional_path] = sha(regional_path)
        current_book = read_book(regional_path)
        if sha(regional_path)!=consumed_hashes[regional_path]:
            raise ValueError('Regional independent inputs changed while loading candidate inputs')
        prepared = prepare_candidate_book(book,current_book)
        book = prepared['book']; candidate_changed = prepared['changed']; candidate_pending = prepared['pending_reason']
    try:
        flags = resolve_audited_configuration(model_dir, book, manifest)
    except ValueError as error:
        flags = None; configuration_error = str(error)
    migration = None
    original_context = _context(root, model, flags)
    if ee_floor:
        if configuration_error:raise ValueError('EE-floor migration requires audited settings: '+configuration_error)
        from tools.project_core.calculations.ee_floor_migration import prepare
        migration = prepare(model,flags,expected,ee_floor)
        if migration:flags=migration['flags']
    context = _context(root, model, flags)
    if context.get('code_identity'):
        from tools.scientific_helpers.sppr.method_coverage import _engine_paths, _calculation_engine_path
        code_paths=list(_engine_paths(flags,model).values())
        calculation_path=_calculation_engine_path(flags,model)
        if calculation_path:code_paths.append(calculation_path)
        from tools.scientific_helpers.sppr import method_coverage as adapter
        code_paths.append(Path(adapter.__file__))
        consumed_hashes.update({code:sha(code) for code in code_paths})
    if configuration_error:
        context['configuration_error'] = configuration_error
    coverage = manifest.get('method_coverage', {})
    pending = _cached_methods(model_dir, coverage, context, book, expected)
    if migration:pending=list(expected)
    source = model_dir / 'sppr_source.xlsx'
    arithmetic_fresh = settings.get('calculation_input_sha256') == input_hash(book)
    from tools.project_core.calculations.regional import result_hash
    arithmetic_fresh = arithmetic_fresh and settings.get('calculation_result_sha256') == result_hash(book)
    arithmetic_fresh = arithmetic_fresh and coverage.get('arithmetic_code_identity') == _arithmetic_identity()
    arithmetic_fresh = arithmetic_fresh and not candidate_changed
    snapshot = model_dir/'results/regional_snapshot.xlsx'
    package_fresh = (snapshot.is_file() and manifest.get('snapshot_sha256')==sha(snapshot)
                     and manifest.get('snapshot_sha256')==sha(path) and source.is_file()
                     and manifest.get('coefficient_source_sha256')==sha(source))
    if not migration and not pending and arithmetic_fresh and package_fresh and _source_matches(model_dir, book, expected):
        progress(f"{metadata['unit_id']}/{model_dir.name}: reusing complete fresh {len(expected)}-method package")
        return False, {'unit_id': metadata['unit_id'], 'model_id': model_dir.name, 'reused': True,
                       'methods': coverage['methods']}
    progress(f"{metadata['unit_id']}/{model_dir.name}: completing {len(pending)} of {len(expected)} methods")
    before = consumed_hashes
    protected = _preserved_inputs(book)
    original_review_hash = settings.get('calculation_input_sha256')
    original_actual_hash = input_hash(original_book)
    reviewed_methods = list(_method_rows(book)[0])
    saved_method_rows = copy.deepcopy(_method_rows(book)[0])
    # Stage every model-owned output. Publication errors restore the package and
    # regional workbook together; signed reports and model inputs are never edited.
    destinations = [source, path, manifest_path, model_dir/'results/regional_snapshot.xlsx']
    evidence_dir = model_dir / 'results/diagnostics/method_coverage'
    floor_dir = model_dir/'results/diagnostics/ee_floor'
    existing_evidence = {p:p.read_bytes() for directory in (evidence_dir,floor_dir) if directory.exists()
                         for p in directory.rglob('*') if p.is_file()}
    backups = {p: p.read_bytes() if p.is_file() else None for p in set(destinations)}
    backups.update(existing_evidence)
    owned = {p: sha(p) if p.is_file() else None for p in backups}
    written = set()
    before.update({p:value for p,value in owned.items() if value is not None})
    def check_inputs():
        for consumed,fingerprint in before.items():
            expected_hash=owned.get(consumed,fingerprint)
            if (sha(consumed) if consumed.is_file() else None)!=expected_hash:
                raise ValueError('Inputs changed before publication replacement: '+str(consumed))
    evidence_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = tempfile.TemporaryDirectory(dir=evidence_dir.parent, prefix='.method-coverage-')
    staged_dir = Path(staging.name).resolve()
    if not staged_dir.is_relative_to(evidence_dir.parent.resolve()):
        raise ValueError('Evidence staging escaped the model result directory')
    try:
        outcomes = copy.deepcopy(coverage.get('methods', {})) if _same_context(coverage.get('context'),context) else {}
        execution_evidence = copy.deepcopy(coverage.get('execution_evidence',{})) if _same_context(coverage.get('context'),context) else {}
        promoted_references = []
        floor_promotions = []
        baseline = None
        execution_cache = None
        if pending:
            if migration:
                from tools.project_core.calculations.ee_floor_migration import save_baseline, retained_outcomes, _write_gzip
                floor_staging=staged_dir/'ee_floor'; floor_staging.mkdir()
                baseline=save_baseline(model_dir,{p:backups[p] for p in destinations},manifest,floor_staging)
                baseline['path']='ee_floor/'+baseline['path']
                floor_promotions.append(baseline)
                reused=retained_outcomes(manifest,migration['compatible'],book,staged_dir,context,baseline,
                                        original_context=original_context,model_dir=model_dir,
                                        reviewed_model_sha256=metadata.get('reviewed_model_sha256'))
                record={key:migration[key] for key in ('before_runtime','after_runtime','old_warnings','new_warnings')}
                from tools.scientific_helpers.sppr.method_coverage import encode_lossless
                record.update(info=encode_lossless(migration['info']),ledger=encode_lossless(migration['ledger']),
                              baseline=baseline,original_flags=copy.deepcopy(coverage.get('context',{}).get('flags',manifest.get('effective_flags'))),requested_flags=flags)
                runtime=_write_gzip(floor_staging/'runtime.json.gz',record);runtime['path']='ee_floor/'+runtime['path']
                floor_promotions.append(runtime);execution_evidence['ee_floor_migration']=runtime
                code_hash=__import__('hashlib').sha256(migration['calculation_code']).hexdigest()
                relative='ee_floor/executed_code/'+code_hash+'/PPRCalculator.py'
                code=staged_dir/relative;code.parent.mkdir(parents=True);code.write_bytes(migration['calculation_code'])
                floor_promotions.append({'path':relative,'sha256':code_hash})
            else:
                from tools.scientific_helpers.sppr.retained_methods import adopt_saved_methods
                reused = adopt_saved_methods(model_dir, manifest, context, book, staged_dir, pending,
                                             reviewed_model_sha256=metadata.get('reviewed_model_sha256'))
            unretained = [key for key in pending if key in saved_method_rows and key not in reused] if not migration else []
            if unretained:
                raise ValueError('Saved method evidence needs an explicit scientific refresh; existing results were preserved: '
                                 + ', '.join(unretained))
            if reused:
                progress(f"{metadata['unit_id']}: retained saved methods without rerunning: {', '.join(reused)}")
            execute = [key for key in pending if key not in reused]
            if migration and execute:
                from tools.project_core.calculations.ee_floor_execution import load_execution, save_execution
                from tools.scientific_helpers.sppr.method_coverage import method_settings
                execution_settings={key:method_settings(key,flags,model) for key in execute}
                cached=load_execution(model_dir,context,execute,execution_settings,staged_dir)
                if cached:
                    calculated=cached['calculated'];execution_cache=cached['reference']
                    progress(f"{metadata['unit_id']}: reusing {len(execute)} completed executions from a prior publication attempt")
                else:
                    calculated=compute(model,flags,execute,staged_dir,progress=progress)
                    check_inputs()
                    execution_cache=save_execution(model_dir,context,execute,execution_settings,calculated,staged_dir)
            else:
                calculated=compute(model, flags, execute, staged_dir, progress=progress) if execute else {'methods':{}}
            calculated['methods'].update(reused)
            for key, outcome in calculated['methods'].items():
                evidence = outcome.get('evidence') or {}
                if 'path' in evidence and not str(evidence['path']).startswith('method_coverage/'):
                    evidence['path'] = 'method_coverage/' + evidence['path']
                promoted_references.append((key,evidence))
                outcomes[key] = outcome
                if configuration_error:
                    outcomes[key]['reason'] = configuration_error + '; ' + outcomes[key].get('reason', '')
            missing = set(pending) - set(calculated['methods'])
            if missing:
                raise ValueError('Executor omitted requested method outcomes: ' + ', '.join(sorted(missing)))
            for key in ('constructor_evidence','evidence_index'):
                if key in calculated:
                    evidence = copy.deepcopy(calculated[key])
                    if not str(evidence.get('path','')).startswith('method_coverage/'):
                        evidence['path'] = 'method_coverage/'+evidence['path']
                    execution_evidence[key] = evidence
                    promoted_references.append((key,evidence))
        if (any(sha(p) != fingerprint for p, fingerprint in before.items())
                or any(p.exists() for p,value in owned.items() if value is None)):
            raise ValueError('Inputs changed during method completion; no mixed input versions published')
        # Only promote exact returned evidence. A concurrent unrelated file is
        # neither owned by this run nor removed by its rollback.
        for key,evidence in promoted_references:
            relative = Path(str(evidence.get('path', ''))).relative_to('method_coverage')
            staged = (staged_dir / relative).resolve()
            target = (evidence_dir / relative).resolve()
            if (not staged.is_relative_to(staged_dir) or not target.is_relative_to(evidence_dir.resolve())
                    or not staged.is_file() or sha(staged) != evidence.get('sha256')):
                raise ValueError('Missing or changed returned method evidence: ' + key)
            if target not in backups:
                if target.exists():
                    raise ValueError('Method evidence appeared during calculation: ' + str(target))
                backups[target] = None; owned[target] = None
            target.parent.mkdir(parents=True, exist_ok=True)
            os.replace(staged, target); owned[target] = sha(target); written.add(target)
        for evidence in floor_promotions:
            relative=Path(evidence['path']); staged=(staged_dir/relative).resolve()
            target=(model_dir/'results/diagnostics'/relative).resolve()
            if not staged.is_relative_to(staged_dir) or not target.is_relative_to(floor_dir.resolve()) or sha(staged)!=evidence['sha256']:
                raise ValueError('Invalid EE-floor migration evidence')
            if target not in backups:
                if target.exists():raise ValueError('EE-floor evidence appeared during calculation')
                backups[target]=None;owned[target]=None
            target.parent.mkdir(parents=True,exist_ok=True);os.replace(staged,target);owned[target]=sha(target);written.add(target)
        _apply_outcomes(book, model_dir.name, outcomes, pending, migration=bool(migration))
        changed_methods=set(pending)-{key for key,value in outcomes.items() if value.get('preserve_saved')} if migration else set()
        if any(_method_rows(book)[0].get(key) != values for key, values in saved_method_rows.items() if key not in changed_methods):
            raise ValueError('Method completion attempted to change existing saved coefficients')
        if candidate_pending:
            from tools.project_core.calculations.coverage_candidates import suppress_candidate_arithmetic
            from tools.project_core.calculations.regional import set_setting
            set_setting(book,'results_model_sha256',sha(model))
            if flags is not None:set_setting(book,'effective_loader_flags',json.dumps(flags,ensure_ascii=False))
            suppress_candidate_arithmetic(book,candidate_pending)
        elif pending or not arithmetic_fresh:
            _arithmetic(book, path, model, flags, original_book,migration=bool(migration),invalidated_methods=changed_methods)
        if migration:
            from tools.project_core.calculations.ee_floor_migration import add_diagnostics
            from tools.project_core.calculations.regional import set_setting
            add_diagnostics(book,migration)
            set_setting(book,'sppr_engine_hashes',json.dumps({k:v for k,v in context['code_identity'].items() if k!='create_PPRS_excel.py'},ensure_ascii=False))
        if _preserved_inputs(book) != protected:
            raise ValueError('Method completion changed a protected scientific input')
        def record_saved():
            owned[path] = sha(path); written.add(path)
        saved = _save_calculation(path,book,model,flags,original_book,
                                  pending_reason=candidate_pending,on_write=record_saved,before_write=check_inputs,
                                  migration=bool(migration),invalidated_methods=changed_methods)
        if _preserved_inputs(saved) != protected:
            raise ValueError('Workbook serialization changed a protected scientific input')
        if any(_method_rows(saved)[0].get(key) != values for key, values in saved_method_rows.items() if key not in changed_methods):
            raise ValueError('Workbook serialization changed existing saved coefficients')
        # A complete source is already authoritative. Re-saving it merely to
        # attach coverage metadata could round its original numeric cells.
        if migration or not _source_matches(model_dir, saved, expected):
            check_inputs()
            _update_source(source, saved, outcomes,replace_methods=changed_methods,
                           effective_groups=migration['effective_groups'] if migration else None,expected_sha256=owned[source])
            owned[source] = sha(source); written.add(source)
            if not _source_matches(model_dir, saved, expected):
                raise ValueError('Saved source coefficients conflict with the completed workbook; retained source cells were not replaced')
        snapshot = model_dir / 'results/regional_snapshot.xlsx'
        if path != snapshot:
            check_inputs();_write_bytes(snapshot,path.read_bytes(),expected_sha256=owned[snapshot])
        owned[snapshot] = sha(snapshot); written.add(snapshot)
        _, row_hashes = _method_rows(saved)
        for key in expected:
            outcomes[key]['group_rows_sha256'] = row_hashes[key]
        current_input = input_hash(saved)
        extension = manifest.get('review_extension')
        if (original_review_hash == metadata.get('reviewed_calculation_input_sha256')
                and original_review_hash == original_actual_hash):
            extension = {'schema_version': 1, 'original_reviewed_input_sha256': original_review_hash,
                         'reviewed_methods': reviewed_methods}
        if extension:
            extension['current_input_sha256'] = current_input
            manifest['review_extension'] = extension
        if migration:
            manifest['ee_floor_migration']={'schema_version':1,'baseline':baseline,
                'requested_policy':ee_floor,'runtime_evidence':execution_evidence['ee_floor_migration'],
                'retained_methods':sorted(set(expected)-changed_methods),'executed_methods':sorted(changed_methods),
                'researcher_approval':'Preserved historical review; migration does not approve changed method outcomes'}
        manifest.update({'schema_version': 1, 'unit_id': metadata['unit_id'], 'model_id': model_dir.name,
                         'snapshot_selected_model_id': overview(saved).get('selected_model_id'),
                         'timestamp': datetime.now(timezone.utc).isoformat(), 'snapshot_sha256': sha(snapshot),
                         'canonical_model_sha256': sha(model), 'present_canonical_model_sha256': sha(model),
                         'coefficient_source_sha256': sha(source), 'effective_flags': flags,
                         'code_identity': {k:v for k,v in context['code_identity'].items() if k != 'create_PPRS_excel.py'},
                         'provenance_status': 'Per-method inventory distinguishes retained saved results from executions; unknown historical provenance and unsuccessful outcomes remain explicit',
                         'dependencies': dependencies(saved),
                         'artifacts': {'snapshot':'regional_snapshot.xlsx','model':'../model.json','coefficients':'../sppr_source.xlsx'},
                         'method_coverage': {'schema_version': 1, 'expected_methods': list(expected),
                                             'context': context, 'methods': outcomes,
                                             'execution_evidence': execution_evidence,
                                             'arithmetic_code_identity': _arithmetic_identity(),
                                             'protected_inputs_sha256': protected}})
        check_inputs();_write_json(manifest_path, manifest,expected_sha256=owned[manifest_path])
        owned[manifest_path] = sha(manifest_path); written.add(manifest_path)
        if plan_methods(expected, context, manifest['method_coverage'], row_hashes):
            raise ValueError('Incomplete final method inventory')
        return True, {'unit_id':metadata['unit_id'], 'model_id':model_dir.name, 'reused':False,
                      'written_region_path': str(path) if is_current else None,
                      'written_paths': [str(p) for p in written],
                      'written_hashes': {str(p):owned[p] for p in written}, 'methods':outcomes,
                      'execution_cache':execution_cache}
    except Exception:
        conflicts = []
        for destination, data in backups.items():
            actual = sha(destination) if destination.is_file() else None
            if actual != owned[destination]:
                conflicts.append(str(destination)); continue
            if data is None:
                destination.unlink(missing_ok=True)
            else:
                destination.parent.mkdir(parents=True, exist_ok=True); destination.write_bytes(data)
        if conflicts:
            raise RuntimeError('Coverage failed; concurrent edits preserved: ' + ', '.join(conflicts))
        raise
    finally:
        staging.cleanup()


def ensure_validated_coverage(root, *, progress=print, lock_held=False, ee_floor=None, migration_units=None):
    """Complete all exact validated identities without changing model selection.

    Central/map callers publish changed_units under their existing writer lock.
    This function writes only model packages and their current regional workbook.
    """
    from tools.project_core.registry.writes import central_lock
    root = Path(root).resolve()
    if not lock_held:
        with central_lock(root):
            return ensure_validated_coverage(root, progress=progress, lock_held=True,ee_floor=ee_floor,migration_units=migration_units)
    from tools.project_core.registry.discovery import region_directory, resolve_model
    from tools.project_core.workbooks.workbooks import sha
    metadata_hash = sha(root/'Project.xlsx')
    metadata = _metadata(root)
    if sha(root/'Project.xlsx')!=metadata_hash:
        raise ValueError('Project model registration changed while loading coverage identities')
    indexed = {(r['unit_id'], r['model_id']): r for r in metadata}
    identities = validated_identities(metadata)
    if not identities:
        return {'changed_units': [], 'models': []}
    from tools.scientific_helpers.sppr.method_coverage import method_inventory, compute_methods
    expected = method_inventory(); changed = set(); models = []
    protected = set()
    directories = []
    for unit, model_id in identities:
        region = region_directory(root, unit); model = resolve_model(region, model_id).parent
        protected.update([region/(unit+'.xlsx'), model/'sppr_source.xlsx',
                          model/'results/regional_snapshot.xlsx', model/'results/result_manifest.json'])
        for kind in ('method_coverage','ee_floor'):
            directory = model/'results/diagnostics'/kind; directories.append(directory)
            if directory.is_dir(): protected.update(p for p in directory.rglob('*') if p.is_file())
    backups = {p:p.read_bytes() if p.is_file() else None for p in protected}
    owned = {p:sha(p) if p.is_file() else None for p in protected}
    try:
        for identity in identities:
            if sha(root/'Project.xlsx')!=metadata_hash:
                raise ValueError('Project model registration changed during coverage completion')
            unit, model_id = identity
            region = region_directory(root,unit); model_dir = resolve_model(region,model_id).parent
            current_paths = {region/(unit+'.xlsx'), model_dir/'sppr_source.xlsx',
                             model_dir/'results/regional_snapshot.xlsx', model_dir/'results/result_manifest.json'}
            directory = model_dir/'results/diagnostics/method_coverage'
            floor_directory = model_dir/'results/diagnostics/ee_floor'
            current_paths.update(p for p in protected if p.is_relative_to(directory) or p.is_relative_to(floor_directory))
            if any((sha(p) if p.is_file() else None) != owned[p] for p in current_paths):
                raise ValueError('Model package changed during coverage batch: ' + model_id)
            request=ee_floor if migration_units is None or unit in migration_units else None
            modified, report = _complete_model(root, indexed[identity], expected, progress=progress, compute=compute_methods,ee_floor=request)
            if modified:
                changed.add(identity[0])
            models.append(report)
            if modified:
                for value in report.get('written_paths', []):
                    path = Path(value)
                    if path not in backups:
                        backups[path] = None
                    owned[path] = report['written_hashes'][str(path)]
        if sha(root/'Project.xlsx')!=metadata_hash:
            raise ValueError('Project model registration changed during coverage completion')
    except Exception:
        conflicts = []
        for path, value in backups.items():
            if (sha(path) if path.is_file() else None) != owned[path]:
                conflicts.append(str(path)); continue
            if value is None: path.unlink(missing_ok=True)
            else: path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(value)
        if conflicts:
            raise RuntimeError('Coverage batch failed; concurrent edits preserved: ' + ', '.join(conflicts))
        raise
    # A prior interrupted publication must consolidate completed regional state,
    # even when every scientific method is already cached and unchanged.
    from tools.project_core.validation.researcher_review import table_rows
    with zipfile.ZipFile(root/'Project.xlsx') as archive:
        registered = {r['unit_id']:r.get('sha256') for _,r in table_rows(archive,'Regions & status','Regions')[3]}
    for unit, _ in identities:
        path = region_directory(root,unit)/(unit+'.xlsx')
        if registered.get(unit) != sha(path): changed.add(unit)
    return {'changed_units': sorted(changed), 'models': models,
            'written_region_paths': [report['written_region_path'] for report in models if report.get('written_region_path')],
            'written_paths': sorted({value for report in models for value in report.get('written_paths', [])}),
            'written_hashes': {path:value for report in models for path,value in report.get('written_hashes',{}).items()}}
