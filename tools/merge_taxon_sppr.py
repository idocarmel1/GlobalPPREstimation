"""Join taxon-to-group mappings onto Ecopath SPPR, and turn the result into PPR.

This is the step that closes the pipeline. Sea Around Us gives catch per taxon per year.
`PPREstimation` gives SPPR per *group* under twenty different methods. Ecopath has no taxon
level, so the two cannot meet until every catch taxon is assigned to a model group -- the
job of `skills/ewe-species-to-group-mapper`. This script consumes that mapping and produces:

    data/<unit_id>/<unit_id>.xlsx, sheet "Taxon SPPR"      taxon x 20 methods
    data/<unit_id>/<unit_id>.xlsx, sheet "PPR by method"   year x 20 methods, tonnes
    data/<unit_id>/taxon_group_map.csv                     the mapping, as plain data

Mappings are read from a `Taxon-Group Map` sheet in the ecosystem workbook when the skill
has been run, and otherwise from the completed examples shipped with the skill.

Two rules the arithmetic depends on:

* A taxon mapped to `Unresolved` contributes no PPR. It is reported in the coverage line
  rather than silently dropped, because unresolved catch is the honest measure of how far
  the mapping got.
* A group name in the mapping that does not appear in the model's own `groups_df` is a
  hard error, not a warning. It means the mapping and the model disagree, and every PPR
  derived from that row would be fiction.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
SPPR_DIR = ROOT / "PPREstimation" / "output" / "top10"
EXAMPLES = ROOT / "skills" / "ewe-species-to-group-mapper" / "examples"
CATCH_DIR = ROOT / "SeaAroundUsExtraction" / "data" / "catch_by_taxon_year"

HEAD = Font(bold=True)


def unit_from_model_filename(name: str) -> str | None:
    m = re.match(r"^(\d+)HS_", name)
    if m:
        return f"HS_{int(m.group(1)):03d}"
    m = re.match(r"^(\d+)_", name)
    if m:
        return f"LME_{int(m.group(1)):03d}"
    return None


def read_mapping(unit: str) -> tuple[dict[str, dict[str, str]], str]:
    """taxon -> {model_column: group}. Prefers the workbook sheet, falls back to examples."""
    book = DATA / unit / f"{unit}.xlsx"
    if book.exists():
        wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
        if "Taxon-Group Map" in wb.sheetnames:
            rows = list(wb["Taxon-Group Map"].iter_rows(values_only=True))
            wb.close()
            return _parse_map_rows(rows), "workbook sheet 'Taxon-Group Map'"
        wb.close()
    for p in EXAMPLES.glob("*.xlsx"):
        if not p.name.startswith(unit):
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        sheet = "Species" if "Species" in wb.sheetnames else wb.sheetnames[0]
        rows = list(wb[sheet].iter_rows(values_only=True))
        wb.close()
        return _parse_map_rows(rows), f"skill example {p.name}"
    return {}, "none found"


def _parse_map_rows(rows) -> dict[str, dict[str, str]]:
    if not rows:
        return {}
    hdr = [str(h) if h is not None else "" for h in rows[0]]
    if "taxon" not in hdr:
        return {}
    ti = hdr.index("taxon")
    # model columns are the EwE_* ones that are not explanation columns
    cols = {i: h for i, h in enumerate(hdr)
            if h.startswith("EwE_") and not h.endswith("_explanation")}
    # a duplicated header (group + unnamed explanation) appears in some examples; keep the first
    seen, keep = set(), {}
    for i, h in cols.items():
        if h in seen:
            continue
        seen.add(h)
        keep[i] = h
    out: dict[str, dict[str, str]] = {}
    for r in rows[1:]:
        if ti >= len(r) or not r[ti]:
            continue
        out.setdefault(str(r[ti]), {}).update(
            {h: (str(r[i]).strip() if i < len(r) and r[i] else "") for i, h in keep.items()}
        )
    return out


def read_model_sppr(unit: str) -> dict[str, tuple[list[str], dict[str, list]]]:
    """model stem -> (method names, group_name -> [values])."""
    out = {}
    for p in sorted(SPPR_DIR.glob("*.xlsx")):
        if unit_from_model_filename(p.name) != unit:
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        if "sppr_all" not in wb.sheetnames:
            wb.close()
            continue
        rows = list(wb["sppr_all"].iter_rows(values_only=True))
        wb.close()
        if not rows:
            continue
        hdr = [str(h) for h in rows[0]]
        methods = hdr[2:]
        by_group = {}
        for r in rows[1:]:
            if len(r) < 2 or r[1] is None:
                continue
            by_group[str(r[1])] = list(r[2:])
        out[p.stem] = (methods, by_group)
    return out


def read_catch(unit: str):
    p = CATCH_DIR / f"{unit}.csv.gz"
    if not p.exists():
        return {}, []
    by_taxon = defaultdict(lambda: defaultdict(float))
    years = set()
    with gzip.open(p, "rt", encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            try:
                y = int(r["year"]); t = float(r["catch_tonnes"] or 0.0)
            except ValueError:
                continue
            years.add(y)
            by_taxon[r["taxon"]][y] += t
    return by_taxon, sorted(years)


def pick_model(mapping_cols: list[str], models: dict) -> dict[str, str]:
    """Associate each mapping column with a model stem, by best token overlap."""
    out = {}
    for col in mapping_cols:
        toks = {t.lower() for t in re.split(r"[_\W]+", col) if len(t) > 2}
        best, score = None, 0
        for stem in models:
            s = len(toks & {t.lower() for t in re.split(r"[_\W]+", stem) if len(t) > 2})
            if s > score:
                best, score = stem, s
        out[col] = best or (next(iter(models)) if models else None)
    return out


def build(unit: str) -> dict:
    mapping, src = read_mapping(unit)
    models = read_model_sppr(unit)
    catch, years = read_catch(unit)
    book = DATA / unit / f"{unit}.xlsx"
    result = {"unit": unit, "source": src, "taxa_mapped": 0, "unresolved": 0,
              "models": list(models), "written": False, "errors": []}
    if not mapping or not models or not book.exists():
        result["errors"].append(
            f"missing input (mapping={bool(mapping)} models={bool(models)} workbook={book.exists()})")
        return result

    cols = sorted({c for v in mapping.values() for c in v})
    col_model = pick_model(cols, models)

    # every mapped group must exist in its model, or the arithmetic is fiction
    for col, stem in col_model.items():
        if stem is None:
            result["errors"].append(f"no model matched mapping column {col}")
            continue
        known = set(models[stem][1])
        used = {v[col] for v in mapping.values()
                if v.get(col) and v[col].lower() != "unresolved"}
        unknown = sorted(used - known)
        if unknown:
            result["errors"].append(
                f"{col} -> {stem}: {len(unknown)} group(s) absent from the model: {unknown[:5]}")
    if result["errors"]:
        return result

    wb = openpyxl.load_workbook(book)
    for name in ("Taxon SPPR", "PPR by method"):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("Taxon SPPR")
    ws.append([f"Ecopath SPPR per taxon for {unit}, via the taxon-to-group mapping."])
    ws["A1"].font = HEAD
    ws.append([f"Mapping source: {src}. A taxon marked Unresolved has no model group and no PPR."])
    ws.append([])
    stems = [col_model[c] for c in cols]
    header = ["taxon"] + [f"{c} group" for c in cols]
    for c, stem in zip(cols, stems):
        header += [f"{c} :: {m}" for m in models[stem][0]]
    ws.append(header)
    for cell in ws[4]:
        cell.font = HEAD

    resolved = 0
    for taxon in sorted(mapping):
        row = [taxon]
        vals = []
        any_res = False
        for c, stem in zip(cols, stems):
            g = mapping[taxon].get(c, "")
            row.append(g)
            methods, by_group = models[stem]
            series = by_group.get(g)
            if g and g.lower() != "unresolved" and series:
                any_res = True
                vals += [v if isinstance(v, (int, float)) else None for v in series]
            else:
                vals += [None] * len(methods)
        resolved += 1 if any_res else 0
        ws.append(row + vals)
    result["taxa_mapped"] = len(mapping)
    result["unresolved"] = len(mapping) - resolved
    finish(ws, len(header))

    # PPR by method: year x method, tonnes. Aggregated, because per-taxon-per-year-per-method
    # would be 20 sheets of noise and the map only ever needs the total.
    ws2 = wb.create_sheet("PPR by method")
    ws2.append([f"PPR for {unit} by Ecopath SPPR method: sum over taxa of catch x SPPR(group)."])
    ws2["A1"].font = HEAD
    ws2.append(["Taxa without a resolved group contribute nothing; see the coverage line below."])
    ws2.append([])
    cols2 = ["year"]
    for c, stem in zip(cols, stems):
        cols2 += [f"{c} :: {m}" for m in models[stem][0]]
    ws2.append(cols2)
    for cell in ws2[4]:
        cell.font = HEAD

    for y in years:
        row = [y]
        for c, stem in zip(cols, stems):
            methods, by_group = models[stem]
            totals = [0.0] * len(methods)
            for taxon, tmap in mapping.items():
                g = tmap.get(c, "")
                if not g or g.lower() == "unresolved":
                    continue
                series = by_group.get(g)
                if not series:
                    continue
                tonnes = catch.get(taxon, {}).get(y, 0.0)
                if not tonnes:
                    continue
                for i, v in enumerate(series):
                    if isinstance(v, (int, float)):
                        totals[i] += tonnes * v
            row += [round(t, 3) if t else None for t in totals]
        ws2.append(row)

    ws2.append([])
    tot = sum(sum(v.values()) for v in catch.values())
    res_catch = sum(sum(catch.get(t, {}).values()) for t, m in mapping.items()
                    if any(g and g.lower() != "unresolved" for g in m.values()))
    ws2.append([f"Catch coverage: {res_catch:,.0f} of {tot:,.0f} tonnes "
                f"({100*res_catch/tot if tot else 0:.1f}%) sits on a resolved group."])
    finish(ws2, len(cols2))

    wb.save(book)
    wb.close()

    with (DATA / unit / "taxon_group_map.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["taxon"] + cols)
        for t in sorted(mapping):
            w.writerow([t] + [mapping[t].get(c, "") for c in cols])

    result["written"] = True
    result["catch_coverage_pct"] = round(100 * res_catch / tot, 1) if tot else 0.0
    return result


def finish(ws, ncols):
    ws.column_dimensions["A"].width = 34
    for i in range(2, min(ncols, 60) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 17
    ws.freeze_panes = "B5"
    ws.sheet_view.showGridLines = False


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--units", nargs="*", help="unit ids; default is every unit with a mapping")
    args = ap.parse_args()

    units = args.units
    if not units:
        units = sorted({p.name.split("_")[0] + "_" + p.name.split("_")[1]
                        for p in EXAMPLES.glob("*.xlsx")})
        units += [d.name for d in DATA.iterdir()
                  if d.is_dir() and (d / f"{d.name}.xlsx").exists() and d.name not in units]
        units = [u for u in dict.fromkeys(units) if (DATA / u).exists()]

    ok = 0
    for u in units:
        r = build(u)
        if r["written"]:
            ok += 1
            print(f"  {u}: {r['taxa_mapped']} taxa, {r['unresolved']} unresolved, "
                  f"{r['catch_coverage_pct']}% of catch on a resolved group  [{r['source']}]")
        elif r["errors"] and "missing input" not in r["errors"][0]:
            print(f"  {u}: NOT WRITTEN — {r['errors'][0]}", file=sys.stderr)
    print(f"\nmerged Ecopath SPPR into {ok} ecosystem workbook(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
