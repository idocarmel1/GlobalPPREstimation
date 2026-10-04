"""Isolated independent regressions; no project scientific data is written."""
import copy,json,sys,unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'Project.xlsx').is_file())
sys.path.insert(0,str(ROOT))
from tools.workflow_checks.selection import test_selection_refresh as selection_fixture
from tools.project_core.calculations import selection
from tools.project_core.calculations.snapshots import save_snapshot
from tools.project_core.calculations.regional import recalculate,set_setting,set_result_hash
from tools.project_core.workbooks.workbooks import read_book,write_book,update_book,overview,records,sha
from tools.project_core.registry.update_project import update

class ReviewRegressions(unittest.TestCase):
    def setUp(self):self.f=selection_fixture.SelectionTests();self.f.setUp()
    def tearDown(self):self.f.tearDown()
    def snapshot(self,book=None,mid='m',flags=None):
        f=self.f;book=book or read_book(f.path)
        return save_snapshot(f.path,f.models[mid],result_identity=selection.execution_identity(f.root,f.models[mid],book,flags=flags or {'normalize_DC':True}))
    def refresh(self):return selection.refresh_region(self.f.root,self.f.path,publisher=self.f.publish)
    def test_incompatible_matching_cannot_restore_annual_values(self):
        f=self.f;b=read_book(f.path);b['PPR']['Matching'][1][:]=[['m','a','A',.5,'high','explicit_member','source'],['m','a','B',.5,'high','explicit_member','source'],['m','b','B',1.,'high','explicit_member','source']]
        recalculate(b,f.path);update_book(f.path,b);self.snapshot(b)
        f.select('B');self.refresh();b=read_book(f.path);b['Catch']['Catch'][1][0][-1]=200;update_book(f.path,b)
        f.select('m');result=self.refresh();self.assertEqual(result['calculation_status'],'pending');self.assertFalse(records(read_book(f.path),'PPR','Annual'));self.assertFalse(records(read_book(f.path),'PPR','Taxon SPPR'))
    def test_unchanged_path_does_not_retag_execution_flags(self):
        f=self.f;self.snapshot();b=read_book(f.path);set_setting(b,'effective_loader_flags',json.dumps({'normalize_DC':False}));update_book(f.path,b)
        result=self.refresh();self.assertEqual(result['calculation_status'],'pending')
        manifest=json.loads((f.models['m']/'results/result_manifest.json').read_text(encoding='utf-8'));self.assertEqual(manifest['effective_flags'],{'normalize_DC':True})
    def test_target_model_does_not_inherit_outgoing_flags(self):
        f=self.f;original=f.path.read_bytes();b=read_book(f.path)
        for k,v in [('selected_model_id','B'),('results_model_id','B'),('results_model_sha256',None),('model_path','papers/P/models/B/model.json')]:set_setting(b,k,v)
        for sheet,table in [('Selected model groups','Group SPPR'),('PPR','Matching'),('PPR','Annual')]:
            for row in b[sheet][table][1]:row[0]='B'
        recalculate(b,f.path);write_book(f.path,b);self.snapshot(b,'B',{'normalize_DC':False});f.path.write_bytes(original)
        b=read_book(f.path);set_setting(b,'effective_loader_flags',json.dumps({'normalize_DC':True}));update_book(f.path,b);f.select('B')
        result=self.refresh();self.assertEqual(result['calculation_status'],'ready');self.assertEqual(json.loads(overview(read_book(f.path))['effective_loader_flags']),{'normalize_DC':False})
    def test_failed_publication_rolls_back_snapshot_package(self):
        f=self.f;self.snapshot();package=f.models['m']/'results';before={p:p.read_bytes() for p in package.iterdir()};f.select('B');active=f.path.read_bytes();central=(f.root/'Project.xlsx').read_bytes()
        def fail(*args):raise RuntimeError('injected failure')
        with self.assertRaisesRegex(RuntimeError,'injected'):selection.refresh_region(f.root,f.path,publisher=fail)
        self.assertTrue(all(p.read_bytes()==value for p,value in before.items()));self.assertEqual(f.path.read_bytes(),active);self.assertEqual((f.root/'Project.xlsx').read_bytes(),central)
    def test_central_updater_obeys_shared_lock(self):
        with selection.central_lock(self.f.root):
            with self.assertRaisesRegex(ValueError,'central writer'):update(self.f.root,[self.f.path])
    def test_outside_table_note_and_native_columns_preserved(self):
        f=self.f;w=openpyxl.load_workbook(f.path);w['Overview']['D3']='human note';w.save(f.path);w.close();b=read_book(f.path);set_setting(b,'selection_rationale','updated');update_book(f.path,b)
        w=openpyxl.load_workbook(f.path);self.assertEqual(w['Overview']['D3'].value,'human note');w.close();update(f.root,[f.path]);w=openpyxl.load_workbook(f.root/'Project.xlsx')
        try:
            for s in w:
                for table in s.tables.values():self.assertEqual(openpyxl.utils.cell.range_boundaries(table.ref)[2],len(table.tableColumns))
        finally:w.close()
    def test_human_edit_during_read_is_preserved(self):
        f=self.f;real_read=read_book;edited=False
        def edit_after_read(path):
            nonlocal edited
            book=real_read(path)
            if Path(path)==f.path and not edited:
                edited=True;latest=real_read(path);set_setting(latest,'researcher_note','new human edit');update_book(path,latest)
            return book
        with patch.object(selection,'read_book',edit_after_read):
            with self.assertRaisesRegex(ValueError,'preflight'):self.refresh()
        self.assertEqual(overview(real_read(f.path))['researcher_note'],'new human edit')
    def test_uncalculated_NOT_RUN_snapshot_stays_pending(self):
        f=self.f;b=read_book(f.path);b['Selected model groups']['Groups']=(['seq','group_name'],[[1,'A'],[2,'B']]);b['Selected model groups']['Group SPPR'][1].clear();b['PPR']['Annual'][1].clear();b['PPR']['Taxon SPPR'][1].clear();b['Diagnostics']['model_health']=(['method','status'],[['GE','NOT_RUN']]);update_book(f.path,b);self.snapshot(b)
        f.select('B');self.refresh();f.select('m');result=self.refresh();self.assertEqual(result['calculation_status'],'pending');self.assertFalse(records(read_book(f.path),'Selected model groups','Group SPPR'))
        self.assertEqual(records(read_book(f.path),'Diagnostics','model_health')[0]['status'],'NOT_RUN')
    def test_compatible_snapshot_preserves_true_eligibility(self):
        f=self.f;b=read_book(f.path);set_setting(b,'production_eligible',True);update_book(f.path,b);self.snapshot(b)
        f.select('B');self.refresh();f.select('m');result=self.refresh();self.assertEqual(result['calculation_status'],'ready');self.assertTrue(overview(read_book(f.path))['production_eligible'])
    def test_unknown_group_membership_stays_pending(self):
        f=self.f;b=read_book(f.path);b['PPR']['Matching'][1][0][2]='Ghost group';update_book(f.path,b);self.snapshot(b)
        f.select('B');self.refresh();f.select('m');result=self.refresh();self.assertEqual(result['calculation_status'],'pending');self.assertFalse(records(read_book(f.path),'PPR','Annual'))
    def test_stale_arithmetic_in_snapshot_is_recalculated(self):
        f=self.f;b=read_book(f.path);b['PPR']['Matching'][1][0][2]='B';update_book(f.path,b);self.snapshot(b)
        f.select('B');self.refresh();f.select('m');result=self.refresh();self.assertEqual(result['calculation_status'],'ready')
        a=next(r for r in records(read_book(f.path),'PPR','Taxon SPPR') if r['taxon']=='a' and r['method']=='one');self.assertIsNone(a['sppr'])
        annual=next(r for r in records(read_book(f.path),'PPR','Annual') if r['metric']=='ppr' and r['method']=='one' and r['catch_basis']=='landings' and r['unidentified']=='method');self.assertIsNone(annual[2019])
    def test_uncertainty_retained_when_fresh_and_invalidated_after_NPP_change(self):
        f=self.f;b=read_book(f.path);template=copy.deepcopy(b['PPR']['Annual'][1][0])
        for metric,value in [('min_tC',1),('max_tC',99)]:
            row=copy.deepcopy(template);row[5]=metric;row[6]='assessed';row[7:]=[value]*70;b['PPR']['Annual'][1].append(row)
        set_result_hash(b);update_book(f.path,b);self.snapshot(b);expected=copy.deepcopy(b['PPR']['Annual'])
        f.select('B');self.refresh();f.select('m');self.refresh();self.assertEqual(read_book(f.path)['PPR']['Annual'],expected)
        f.select('B');self.refresh();b=read_book(f.path);b['NPP']['NPP'][1][0][2]=None;update_book(f.path,b);f.select('m');self.refresh();active=read_book(f.path)
        self.assertIsNone(active['NPP']['NPP'][1][0][2]);self.assertFalse(any(r['metric'] in ['min_tC','max_tC'] for r in records(active,'PPR','Annual')))
        ratio=next(r for r in records(active,'PPR–NPP','Ratios') if r['model_id']=='m' and r['method']=='one' and r['catch_basis']=='landings' and r['unidentified']=='method');self.assertIsNone(ratio[1950])

if __name__=='__main__':
    files=['tools/project_core/calculations/selection.py','tools/project_core/calculations/snapshots.py','tools/project_core/workbooks/workbooks.py','tools/project_core/registry/update_project.py','tools/project_core/registry/writes.py']
    initial={p:sha(ROOT/p) for p in files}
    suite=unittest.defaultTestLoader.loadTestsFromNames(['ReviewRegressions.'+name for name in sys.argv[1:]],module=sys.modules[__name__]) if sys.argv[1:] else unittest.defaultTestLoader.loadTestsFromTestCase(ReviewRegressions)
    result=unittest.TextTestRunner(verbosity=2).run(suite);final={p:sha(ROOT/p) for p in files}
    Path(__file__).with_name('additional_regression_results.json' if sys.argv[1:] else 'regression_review_results.json').write_text(json.dumps({'tests':result.testsRun,'failures':[(t.id(),msg) for t,msg in result.failures],'errors':[(t.id(),msg) for t,msg in result.errors],'source_hashes':initial,'sources_unchanged_during_run':initial==final},indent=2),encoding='utf-8')
    sys.exit(not result.wasSuccessful())
