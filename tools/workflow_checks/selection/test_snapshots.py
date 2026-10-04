import copy, json, tempfile, unittest
from pathlib import Path
import openpyxl
from tools.project_core.workbooks.workbooks import read_book, write_book, overview, sha, YEARS, REGION_SHEETS
from tools.project_core.calculations.snapshots import save_snapshot, inspect_snapshot, restore_model_tables

def fixture():
    b={s:{} for s in REGION_SHEETS}
    b['Overview']['Settings']=(['field','value'],[['unit_id','LME_001'],['selected_model_id','B'],['results_model_id','A'],['selection_rationale','incoming B']])
    b['Catch']['Catch']=(['taxon','catch_basis',*YEARS],[['fish','landings',*([2]*70)]])
    b['Classic PPR']['Taxa']=(['taxon','sppr'],[['fish',10]])
    b['Selected model groups']['Groups']=(['seq','group_name'],[[1,'Fish']])
    b['Selected model groups']['Group SPPR']=(['model_id','group','scope','method','sppr'],[['A','Fish','all','GE',5]])
    b['PPR']['Matching']=(['model_id','taxon','group','weight'],[['A','fish','Fish',1]])
    b['PPR']['Annual']=(['model_id','scope','method','catch_basis','unidentified','metric','status',*YEARS],[['A','all','GE','landings','method','ppr','FAIL',*([None]*70)]])
    b['NPP']['NPP']=(['method','units',*YEARS],[['npp','tC/year',*([100]*70)]])
    b['Diagnostics']['manual']=(['note'],[['human note']])
    b['Diagnostics']['model_health']=(['method','status'],[['GE','FAIL']])
    return b

class SnapshotTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.r=Path(self.temp.name);self.m=self.r/'papers/P/models/A';self.m.mkdir(parents=True)
        (self.m/'model.json').write_text('{"group":[{"group_seq":"1","group_name":"Fish"}]}')
        self.path=self.r/'LME_001.xlsx';self.b=fixture();write_book(self.path,self.b)
        w=openpyxl.load_workbook(self.path);s=w.create_sheet('Researcher');s['A1']='manual';s['A1'].hyperlink='https://example.org';s['A1'].font=openpyxl.styles.Font(bold=True);w.save(self.path)
        self.identity={'model_id':'A','canonical_model_sha256':sha(self.m/'model.json'),'effective_flags':{'normalize_DC':True},'code_identity':{'calculator':'known'},'provenance_status':'verified'}
    def tearDown(self):self.temp.cleanup()
    def test_outgoing_identity_and_exact_full_copy(self):
        before=self.path.read_bytes();p=save_snapshot(self.path,self.m,result_identity=self.identity)
        self.assertEqual(p.read_bytes(),before)
        manifest=json.loads((self.m/'results/result_manifest.json').read_text())
        self.assertEqual(manifest['model_id'],'A');self.assertEqual(manifest['snapshot_selected_model_id'],'B')
        w=openpyxl.load_workbook(p);self.assertEqual(w['Researcher']['A1'].value,'manual');self.assertTrue(w['Researcher']['A1'].font.bold);w.close()
    def test_pending_does_not_overwrite_useful_snapshot(self):
        p=save_snapshot(self.path,self.m,result_identity=self.identity);before=p.read_bytes()
        self.b['Selected model groups']['Group SPPR'][1].clear();self.b['PPR']['Annual'][1].clear();write_book(self.path,self.b)
        self.assertIsNone(save_snapshot(self.path,self.m,result_identity=self.identity));self.assertEqual(p.read_bytes(),before)
    def test_corruption_and_model_or_flags_mismatch_reject_numerical_reuse(self):
        p=save_snapshot(self.path,self.m,result_identity=self.identity)
        expected={'effective_flags':{'normalize_DC':False},'code_identity':{'calculator':'known'}}
        result=inspect_snapshot(self.m,self.b,expected_identity=expected);self.assertFalse(result['coefficients']);self.assertIn('flags',' '.join(result['reasons']))
        p.write_bytes(b'corrupt');self.assertFalse(inspect_snapshot(self.m,self.b,expected_identity=self.identity)['coefficients'])
    def test_restore_keeps_independent_data_and_manual_diagnostics(self):
        target=copy.deepcopy(self.b);target['Catch']['Catch'][1][0][2]=9;target['NPP']['NPP'][1][0][2]=200
        target['Diagnostics']['manual'][1][0][0]='new human note';before=copy.deepcopy(target)
        restored=restore_model_tables(target,self.b,{'coefficients':True,'matching':True,'diagnostics':True})
        for s in ['Overview','Catch','Classic PPR','NPP']:self.assertEqual(restored[s],before[s])
        self.assertEqual(restored['Diagnostics']['manual'],before['Diagnostics']['manual'])
        self.assertEqual(restored['PPR']['Annual'][1][0][6],'FAIL')

    def test_authored_matching_records_restore_with_unknown_coefficient_provenance(self):
        saved=copy.deepcopy(self.b)
        saved['Selected model groups']['Group SPPR'][1].clear()
        saved['PPR']['Allocation assumptions']=(['taxon','model_id','assumption','evidence'],
            [['fish','A','Researcher retained equal allocation; not an inferred proportion','Source table 2; uncertainty retained']])
        saved['PPR']['Mapping review']=(['model_id','researcher','decision','limitations'],
            [['A','Researcher','Membership accepted; allocation remains uncertain','No size composition available']])
        write_book(self.path,saved)
        unknown={**self.identity,'effective_flags':None,'code_identity':None,'provenance_status':'unknown historical execution'}
        save_snapshot(self.path,self.m,result_identity=unknown)
        target=copy.deepcopy(self.b)
        target['Catch']['Catch'][1][0][2]=9
        target['NPP']['NPP'][1][0][2]=None
        target['Diagnostics']['manual'][1][0][0]='new unrelated researcher note'
        target['PPR']['Annual'][1].clear()
        target['PPR']['Unrelated researcher annotations']=(['text'],[['Preserve current independent note']])
        before=copy.deepcopy(target)
        compatibility=inspect_snapshot(self.m,target,expected_identity=unknown)
        self.assertTrue(compatibility['matching'])
        self.assertFalse(compatibility['coefficients'])
        restored=restore_model_tables(target,compatibility['book'],compatibility)
        for name in ['Allocation assumptions','Mapping review']:
            self.assertEqual(restored['PPR'][name],saved['PPR'][name])
        for sheet in ['Overview','Catch','Classic PPR','NPP']:
            self.assertEqual(restored[sheet],before[sheet])
        self.assertEqual(restored['Diagnostics']['manual'],before['Diagnostics']['manual'])
        self.assertEqual(restored['PPR']['Unrelated researcher annotations'],before['PPR']['Unrelated researcher annotations'])
        self.assertEqual(restored['PPR']['Annual'],before['PPR']['Annual'])
        self.assertEqual(target,before,'Selective restoration does not mutate caller-owned current data')

    def test_incompatible_matching_never_restores_authored_matching_records(self):
        saved=copy.deepcopy(self.b)
        saved['PPR']['Matching'][1][0][3]=.5
        saved['PPR']['Allocation assumptions']=(['model_id','assumption'],[['A','Historical allocation review']])
        saved['PPR']['Mapping review']=(['model_id','decision'],[['A','Historical membership review']])
        write_book(self.path,saved);save_snapshot(self.path,self.m,result_identity=self.identity)
        target=copy.deepcopy(self.b)
        target['PPR']['Allocation assumptions']=(['model_id','assumption'],[['B','Current authored record']])
        target['PPR']['Mapping review']=(['model_id','decision'],[['B','Current researcher decision']])
        target['Catch']['Catch'][1][0][2]=9
        target['Diagnostics']['manual'][1][0][0]='current unrelated review'
        before=copy.deepcopy(target)
        compatibility=inspect_snapshot(self.m,target,expected_identity=self.identity)
        self.assertFalse(compatibility['matching'])
        self.assertIn('saved matching weights do not sum to one',compatibility['reasons'])
        restored=restore_model_tables(target,compatibility['book'],compatibility)
        for name in ['Allocation assumptions','Mapping review','Matching','Annual']:
            self.assertEqual(restored['PPR'][name],before['PPR'][name])
        for sheet in ['Overview','Catch','Classic PPR','NPP']:
            self.assertEqual(restored[sheet],before[sheet])
        self.assertEqual(restored['Diagnostics']['manual'],before['Diagnostics']['manual'])
        self.assertEqual(target,before)
