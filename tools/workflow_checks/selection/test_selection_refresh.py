import copy,json,tempfile,unittest,gzip
from pathlib import Path
from unittest.mock import patch
import openpyxl
from tools.project_core.workbooks.workbooks import write_book,read_book,sha,overview,records,input_hash
from tools.project_core.calculations.regional import recalculate,set_setting,set_result_hash
from tools.project_core.calculations.selection import refresh_region,execution_identity
from tools.project_core.calculations.snapshots import save_snapshot
from tools.workflow_checks.calculations.test_workflow import fixture

class SelectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.region=self.root/'regions/LME/LME_001';self.region.mkdir(parents=True)
        self.path=self.region/'LME_001.xlsx';self.models={}
        engine=self.root/'tools/scientific_code/PPREstimation';engine.mkdir(parents=True)
        for name in ['ModelData.py','PPRCalculator.py','utils.py']:(engine/name).write_text('synthetic fixture code identity')
        context=self.root/'common_reference_data/atlas_source_context';context.mkdir(parents=True)
        with gzip.open(context/'catalog.json.gz','wt',encoding='utf-8') as f:json.dump({'curated_region_ids':[]},f)
        for mid in ['m','B']:
            d=self.region/'papers/P/models'/mid;d.mkdir(parents=True);(d/'model.json').write_text('{"group":[{"group_seq":"1","group_name":"A"},{"group_seq":"2","group_name":"B"}]}');self.models[mid]=d
        self.book=fixture();set_setting(self.book,'region_name','Fixture');set_setting(self.book,'region_type','LME');set_setting(self.book,'model_path','papers/P/models/m/model.json');set_setting(self.book,'selection_rationale','choice m')
        recalculate(self.book,self.path);write_book(self.path,self.book)
        self.project={'Models & coverage':{'Models':(['unit_id','model_id','model_path','human_note'],[['LME_001',mid,(d/'model.json').relative_to(self.root).as_posix(),'retain me'] for mid,d in self.models.items()])},'Papers':{'Papers':(['unit_id','article_id'],[])},'Definitions & build':{}}
        write_book(self.root/'Project.xlsx',self.project)
        w=openpyxl.load_workbook(self.path);s=w.create_sheet('Human');s['A1']='preserve manual sheet';s['A1'].font=openpyxl.styles.Font(bold=True);w.save(self.path)
    def tearDown(self):self.tmp.cleanup()
    def publish(self,root,units):
        p=root/'interactive_map';p.mkdir(exist_ok=True);(p/'index.html').write_text(overview(read_book(self.path))['selected_model_id'])
    def select(self,mid):
        from tools.project_core.workbooks.workbooks import update_book
        b=read_book(self.path);set_setting(b,'selected_model_id',mid);set_setting(b,'selection_rationale','choice '+mid);update_book(self.path,b)
    def test_pending_B_and_return_to_A_full_snapshot_without_input_rollback(self):
        save_snapshot(self.path,self.models['m'],result_identity=execution_identity(self.root,self.models['m'],self.book,flags={'normalize_DC':True}))
        self.select('B');result=refresh_region(self.root,self.path,publisher=self.publish)
        self.assertEqual(result['calculation_status'],'pending');self.assertFalse(records(read_book(self.path),'PPR','Annual'));self.assertEqual((self.root/'interactive_map/index.html').read_text(),'B')
        self.select('m');result=refresh_region(self.root,self.path,publisher=self.publish)
        self.assertEqual(result['calculation_status'],'ready');self.assertEqual(overview(read_book(self.path))['results_model_id'],'m')
        w=openpyxl.load_workbook(self.path);self.assertEqual(w['Human']['A1'].value,'preserve manual sheet');self.assertTrue(w['Human']['A1'].font.bold);w.close()
        self.assertTrue(all(r['human_note']=='retain me' for r in records(read_book(self.root/'Project.xlsx'),'Models & coverage','Models')))
    def test_unchanged_refresh_and_invalid_choice_preflight(self):
        before=copy.deepcopy(read_book(self.path)['PPR']);refresh_region(self.root,self.path,publisher=self.publish);self.assertEqual(read_book(self.path)['PPR'],before)
        self.select('missing');before=self.path.read_bytes();central=(self.root/'Project.xlsx').read_bytes()
        with self.assertRaisesRegex(ValueError,'Unknown'):refresh_region(self.root,self.path,publisher=self.publish)
        self.assertEqual(self.path.read_bytes(),before);self.assertEqual((self.root/'Project.xlsx').read_bytes(),central)
    def test_publication_failure_rolls_back_authoritative_files(self):
        self.select('B');before=self.path.read_bytes();central=(self.root/'Project.xlsx').read_bytes()
        def fail(root,units):raise RuntimeError('injected publication failure')
        with self.assertRaisesRegex(RuntimeError,'injected'):refresh_region(self.root,self.path,publisher=fail)
        self.assertEqual(self.path.read_bytes(),before);self.assertEqual((self.root/'Project.xlsx').read_bytes(),central)
