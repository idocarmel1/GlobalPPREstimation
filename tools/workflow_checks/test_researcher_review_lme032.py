"""Review-source regression checks for the approved Karnataka report."""
import json, sys, unittest, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from researcher_review import read_report

ROOT = Path(__file__).resolve().parents[2]
MODEL = '32_1_Arabian_Sea_off_Karnataka_(2000)'
REPORT = ROOT / 'regions/LME_032' / ('Model_validation_' + MODEL + '.docx')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class KarnatakaReviewTests(unittest.TestCase):
    def test_only_new_table_changes_word_content_and_covers_all_very_low_taxa(self):
        qa = REPORT.parent/'validation_reports'/MODEL/'researcher_review_20261001'
        proof = json.loads((qa/'word_preservation.json').read_text(encoding='utf-8'))
        with zipfile.ZipFile(qa/('baseline_'+REPORT.name)) as old, zipfile.ZipFile(REPORT) as new:
            self.assertEqual(old.namelist(),new.namelist())
            changed = [name for name in old.namelist() if old.read(name)!=new.read(name)]
            self.assertEqual(changed,['word/document.xml'])
            raw = new.read('word/document.xml'); offset = proof['inserted_xml_offset']; length = proof['inserted_xml_bytes']
            self.assertEqual(raw[:offset]+raw[offset+length:],old.read('word/document.xml'))
            nodes = list(ET.fromstring(raw).find(W+'body'))
        text = lambda e: ''.join(t.text or '' for t in e.iter(W+'t'))
        index = next(i for i,n in enumerate(nodes) if n.tag==W+'p' and text(n)=='Very low decisions')
        self.assertEqual(text(nodes[index+2]),'Geographic evidence')
        rows = nodes[index+1].findall(W+'tr')
        self.assertEqual([text(c) for c in rows[0].findall(W+'tc')],['Affected taxa','Why confidence is very low'])
        names = [name for row in rows[1:] for name in text(row.find(W+'tc')).split('; ')]
        book = load_workbook(ROOT/'regions/LME_032/LME032_taxon_mapping_appendix.xlsx',read_only=True,data_only=True)
        expected = [r[0] for r in list(book['Taxon appendix'].values)[5:] if r[5]=='Very low']; book.close()
        self.assertEqual(len(names),40); self.assertEqual(len(set(names)),40)
        self.assertEqual(set(names),set(expected))

    def test_exact_review_identity_text_and_region_specific_appendix(self):
        name, date, summary = read_report(ROOT, REPORT, MODEL)
        self.assertEqual((name, date), ('Ido Carmel', '2026-10-01'))
        self.assertIn('rho_living = 0.58', summary['sections'][1]['rows'][0]['text'])
        self.assertIn('(34) marine mammals', summary['sections'][1]['rows'][1]['text'])
        self.assertIn('marine mammals (GE=1.5%', summary['sections'][2]['rows'][0]['text'])
        self.assertEqual(summary['sections'][4]['table'][2], ['Medium','205','48.5275%','64.2253%'])
        self.assertEqual(summary['sections'][4]['appendix_path'], 'regions/LME_032/LME032_taxon_mapping_appendix.xlsx')
        self.assertTrue(any('2019 landings' in p for p in summary['sections'][4]['reference']))
        self.assertIn('marine mammals are excluded from displayed PPR', summary['note'])
        self.assertNotIn('pinnipeds', summary['note'])

    def test_existing_assignment_table_matches_current_regional_appendix(self):
        with zipfile.ZipFile(REPORT) as z:
            body = ET.fromstring(z.read('word/document.xml')).find(W+'body')
        text = lambda e: ''.join(t.text or '' for t in e.iter(W+'t'))
        nodes = list(body)
        locations = [i for i,n in enumerate(nodes) if n.tag == W+'p' and text(n) == 'Group assignment rules']
        self.assertEqual(len(locations), 1)
        table = nodes[locations[0]+1]
        self.assertEqual(table.tag, W+'tbl')
        actual = [[text(c) for c in row.findall(W+'tc')] for row in table.findall(W+'tr')]
        self.assertEqual(actual[0], ['Plain-language rule','Confidence','PPR percentage'])
        book = load_workbook(ROOT/'regions/LME_032/LME032_taxon_mapping_appendix.xlsx',read_only=True,data_only=True)
        rows = list(book['Coverage'].values); book.close()
        start = next(i for i,r in enumerate(rows) if r[0] == 'Group assignment rules')+2
        expected = []
        for row in rows[start:]:
            if row[0] is None: break
            expected.append([row[0],row[1],f'{row[2]:.4f}%'])
        self.assertEqual(actual[1:], expected)
        self.assertEqual(len(expected), 9)

if __name__ == '__main__': unittest.main()
