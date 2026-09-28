import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from workbooks import ANNUAL_HEADER, SIMPLE, read_book, records, write_book


class AtlasRankTests(unittest.TestCase):
    def test_ensemble_current_values_ties_missing_and_metadata(self):
        from atlas_ranks import apply_atlas_ranks
        header = ['unit_id', *ANNUAL_HEADER]
        def annual(unit, value, basis='catch', status='ok', model=None):
            return [unit, model, 'all', SIMPLE, basis, 'method', 'ppr', status,
                    *([None] * 69), value]
        book = {
            'Regional PPR': {'Annual': (header, [annual('LME_001', 10), annual('EEZ_001', 20),
                annual('HS_001', 20), annual('EEZ_999', 1000), annual('LME_002', None),
                annual('LME_003', -1), annual('LME_004', 99, status='failed'),
                annual('LME_001', 900, basis='landings'), annual('LME_001', 999, model='alternative')])},
            'Regions & status': {'Regions': (['unit_id', 'name'], [[u, u] for u in
                ['LME_001', 'EEZ_001', 'HS_001', 'EEZ_999', 'LME_002', 'LME_003', 'LME_004']])},
            'Papers': {'Papers': (['unit_id', 'notes'], [['EEZ_001', 'User approval'], ['EEZ_001', 'Alternative']])},
            'Models & coverage': {'Models': (['unit_id', 'model_id'], [['LME_001', 'a'], ['LME_001', 'b']])},
        }
        apply_atlas_ranks(book, {'LME_001', 'EEZ_001', 'HS_001', 'LME_002', 'LME_003', 'LME_004'})
        ranks = {r['unit_id']: r['atlas_region_rank'] for r in records(book, 'Regions & status', 'Regions')}
        self.assertEqual(ranks, {'LME_001': 3, 'EEZ_001': 1, 'HS_001': 1, 'EEZ_999': None,
                                'LME_002': None, 'LME_003': None, 'LME_004': None})
        self.assertEqual([r['atlas_region_rank'] for r in records(book, 'Papers', 'Papers')], [1, 1])
        self.assertEqual(records(book, 'Papers', 'Papers')[0]['notes'], 'User approval')
        self.assertEqual([r['atlas_region_rank'] for r in records(book, 'Models & coverage', 'Models')], [3, 3])
        apply_atlas_ranks(book, {'LME_001'})
        self.assertEqual(book['Papers']['Papers'][0].count('atlas_region_rank'), 1)
        self.assertIsNone(records(book, 'Papers', 'Papers')[0]['atlas_region_rank'])

    def test_duplicate_reference_rows_rejected(self):
        from atlas_ranks import apply_atlas_ranks
        row = ['LME_001', None, 'all', SIMPLE, 'catch', 'method', 'ppr', 'ok', *([1]*70)]
        with self.assertRaisesRegex(ValueError, 'Duplicate'):
            apply_atlas_ranks({'Regional PPR': {'Annual': (['unit_id', *ANNUAL_HEADER], [row, row])}}, {'LME_001'})

    def test_project_native_tables_preserve_year_keys_and_multiple_blocks(self):
        import openpyxl
        book = {'Regional PPR': {'Annual': (['unit_id', 2019], [['LME_001', 25]])},
                'Definitions & build': {'First': (['field', 'value'], [['a', 'b']]),
                                        'Second': (['field', 'value'], [['c', 'd']])}}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'Project.xlsx'
            write_book(path, book)
            actual = read_book(path)
            self.assertEqual(actual, book)
            wb = openpyxl.load_workbook(path)
            self.assertEqual(sum(len(s.tables) for s in wb), 3)
            self.assertEqual(wb['Regional PPR']['B2'].value, '2019')
            self.assertEqual(wb['Regional PPR'].tables.values().__iter__().__next__().autoFilter.ref, 'A2:B3')
            wb.close()

if __name__ == '__main__':
    unittest.main()
