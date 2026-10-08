import copy
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest

from tools.project_core.maps.build_html import linked_layout
from tools.project_core.maps.original_atlas_data import detail_from_book
from tools.project_core.maps import original_atlas_data


class GroupCatchTests(unittest.TestCase):
    @unittest.skipUnless(shutil.which('node'), 'Node required for generated dialog behavior')
    def test_generated_group_dialog_and_catch_weights(self):
        root = Path(__file__).resolve().parents[3]
        with tempfile.TemporaryDirectory() as directory:
            pages = []
            for name in ['index.html', 'trends.html']:
                page = Path(directory) / name
                page.write_text(linked_layout((root / 'tools/project_core/maps/original_html_layout' / name).read_text('utf-8')), encoding='utf-8')
                pages.append(str(page))
            result = subprocess.run([shutil.which('node'), str(Path(__file__).with_name('check_group_catch.js')), *pages], capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_detail_exposes_only_saved_model_catch(self):
        book = {'Selected model groups': {'Groups': (['group_name', 'catch'], [['Fish', 2.5], ['Zero', 0], ['Missing', None]])}}
        _, model = detail_from_book(book, 'U', 'm')
        self.assertEqual([g.get('model_catch') for g in model['group_data']['groups']], [2.5, 0, None])

    def test_enrichment_uses_current_catch_and_canonical_alternative_without_changing_inputs(self):
        from tools.project_core.workbooks.workbooks import write_book, sha
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            region = root / 'regions/LME/LME_001'
            region.mkdir(parents=True)
            write_book(region / 'LME_001.xlsx', {
                'Overview': {'Settings': (['field', 'value'], [['selected_model_id', 'current'], ['results_model_id', 'current']])},
                'Selected model groups': {'Groups': (['group_name', 'catch'], [['Fish', 2.5], ['Zero', 0], ['Missing', None]])}})
            for model_id in ['current', 'alternative', 'absent']:
                model_path = region / 'papers/P/models' / model_id / 'model.json'
                model_path.parent.mkdir(parents=True)
                metadata = {'_reconstruction': {'currency': 'tonnes carbon per km2; annual rates'}} if model_id == 'alternative' else {}
                model_path.write_text(json.dumps(metadata))
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'groups_df'
            sheet.append(['group_name', 'catch', 'flow_units'])
            sheet.append(['Fish', 4.5, 'tC/km2/year'])
            sheet.append(['Zero', 0, 'tC/km2/year'])
            sheet.append(['Missing', None, 'tC/km2/year'])
            provenance = workbook.create_sheet('Provenance')
            provenance.append(['field', 'value'])
            provenance.append(['model_id', 'alternative'])
            provenance.append(['model_sha256', sha(region / 'papers/P/models/alternative/model.json')])
            provenance.append(['native_model_currency', 'tC/km2, peryear forflows'])
            native = workbook.create_sheet('native_C_all')
            native.append(['group_name', 'new_GE'])
            native.append(['Fish', 18])
            native.append(['Zero', 9])
            native.append(['Missing', None])
            workbook.save(region / 'papers/P/models/alternative/sppr_source.xlsx')
            workbook.close()
            units = {'LME_001': {'models': [
                {'id': mid, 'group_data': {'groups': [{'id': name, 'sppr': 9} for name in ['Fish', 'Zero', 'Missing']], 'methods': ['new_GE']}}
                for mid in ['current', 'alternative', 'absent']]}}
            units['LME_001']['models'][0].update(workbook='regions/LME/LME_001/LME_001.xlsx', workbook_sha256=sha(region / 'LME_001.xlsx'))
            source_path = region / 'papers/P/models/alternative/sppr_source.xlsx'
            units['LME_001']['models'][1].update(source=source_path.relative_to(root).as_posix(), source_sha256=sha(source_path))
            before = copy.deepcopy(units)
            input_bytes = {path: path.read_bytes() for path in region.rglob('*.xlsx')}
            original_atlas_data.add_group_catch(root, units)
            values = [[group['model_catch'] for group in model['group_data']['groups']] for model in units['LME_001']['models']]
            self.assertEqual(values, [[2.5, 0, None], [4.5, 0, None], [None, None, None]])
            self.assertEqual(units['LME_001']['models'][1]['model_catch_carbon_factor'], 1)
            self.assertIsNone(units['LME_001']['models'][0]['model_catch_carbon_factor'])
            self.assertEqual(units['LME_001']['models'][1]['group_data']['groups'][0]['model_catch_sppr']['all']['new_GE'], 18)
            for model in units['LME_001']['models']:
                for field in ['model_catch_units', 'model_catch_carbon_factor', 'model_catch_unit_basis']:
                    model.pop(field)
                for group in model['group_data']['groups']:
                    group.pop('model_catch')
                    group.pop('model_catch_sppr', None)
            self.assertEqual(units, before)
            for path, data in input_bytes.items():
                self.assertEqual(path.read_bytes(), data)
            # A changed canonical source cannot reuse the native-carbon table.
            (region / 'papers/P/models/alternative/model.json').write_text('{"changed": true}')
            original_atlas_data.add_group_catch(root, units)
            alternative = units['LME_001']['models'][1]
            self.assertIsNone(alternative['model_catch_carbon_factor'])
            self.assertTrue(all('model_catch_sppr' not in group for group in alternative['group_data']['groups']))
            source_path.write_bytes(source_path.read_bytes() + b'changed since embedded coefficient source')
            original_atlas_data.add_group_catch(root, units)
            self.assertTrue(all(group['model_catch'] is None for group in alternative['group_data']['groups']))

    def test_wet_density_requires_documented_currency_and_annual_catch(self):
        from tools.project_core.workbooks.workbooks import write_book, sha
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            region = root / 'regions/LME/LME_001'
            model = region / 'papers/P/models/current/model.json'
            model.parent.mkdir(parents=True)
            model.write_text('{}')
            write_book(region / 'LME_001.xlsx', {
                'Overview': {'Settings': (['field', 'value'], [['selected_model_id', 'current'], ['results_model_id', 'current']])},
                'Selected model groups': {'Groups': (['group_name', 'catch'], [['Fish', 1.5]])}})
            source = model.parent / 'extracted_tables/evidence/source'
            source.mkdir(parents=True)
            (source / 'REPORT.md').write_text('Canonical B is interpreted as t wet weight/km²; catches are t/km²/year.', encoding='utf-8')
            units = {'LME_001': {'models': [{'id': 'current', 'workbook': 'regions/LME/LME_001/LME_001.xlsx', 'workbook_sha256': sha(region / 'LME_001.xlsx'), 'group_data': {'groups': [{'id': 'Fish'}]}}]}}
            original_atlas_data.add_group_catch(root, units)
            result = units['LME_001']['models'][0]
            self.assertIsNone(result['model_catch_carbon_factor'])
            (source / 'MODEL_PROFILE.md').write_text('Biomass uses wet-weight densities.', encoding='utf-8')
            original_atlas_data.add_group_catch(root, units)
            self.assertEqual(result['model_catch_carbon_factor'], 1/9)
            self.assertEqual(result['model_catch_units'], 't wet weight/km²/year')
            self.assertEqual(result['group_data']['groups'][0]['model_catch'], 1.5)
            self.assertEqual(len(result['model_catch_unit_basis']['sources']), 3)
            self.assertTrue(all(len(item['sha256']) == 64 for item in result['model_catch_unit_basis']['sources']))

    def test_regional_catch_requires_actual_result_identity_and_embedded_workbook_hash(self):
        from tools.project_core.workbooks.workbooks import write_book, sha
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            region = root / 'regions/LME/LME_001'
            model = region / 'papers/P/models/current/model.json'
            model.parent.mkdir(parents=True)
            model.write_text('{}')
            workbook = region / 'LME_001.xlsx'
            def save(result_id, catch):
                write_book(workbook, {
                    'Overview': {'Settings': (['field', 'value'], [['selected_model_id', 'current'], ['results_model_id', result_id]])},
                    'Selected model groups': {'Groups': (['group_name', 'catch'], [['Fish', catch]])}})
            save('current', 2)
            item = {'id': 'current', 'workbook': workbook.relative_to(root).as_posix(), 'workbook_sha256': sha(workbook), 'group_data': {'groups': [{'id': 'Fish'}]}}
            units = {'LME_001': {'models': [item]}}
            original_atlas_data.add_group_catch(root, units)
            self.assertEqual(item['group_data']['groups'][0]['model_catch'], 2)
            save('current', 7)
            original_atlas_data.add_group_catch(root, units)
            self.assertIsNone(item['group_data']['groups'][0]['model_catch'], 'Edited catch cannot be paired with older embedded coefficients')
            import openpyxl
            source = model.with_name('sppr_source.xlsx')
            source_book = openpyxl.Workbook()
            source_book.active.title = 'groups_df'
            source_book.active.append(['group_name', 'catch'])
            source_book.active.append(['Fish', 100])
            source_book.save(source)
            source_book.close()
            item.update(source=source.relative_to(root).as_posix(), source_sha256=sha(source))
            original_atlas_data.add_group_catch(root, units)
            self.assertIsNone(item['group_data']['groups'][0]['model_catch'], 'Stale regional coefficients cannot silently fall back to source-workbook catch')
            save('outgoing', 11)
            item['workbook_sha256'] = sha(workbook)
            original_atlas_data.add_group_catch(root, units)
            self.assertIsNone(item['group_data']['groups'][0]['model_catch'], 'Pending selection cannot relabel outgoing group catches')
            save('current', 2)
            item.pop('workbook_sha256')
            original_atlas_data.add_group_catch(root, units)
            self.assertIsNone(item['group_data']['groups'][0]['model_catch'], 'Unknown embedded source identity is not fresh')
            # Simulate a concurrent writer after the first fingerprint was read.
            # The real hash function and real workbook reader still run.
            from unittest.mock import patch
            item['workbook_sha256'] = sha(workbook)
            before = copy.deepcopy(units)
            changed = False
            def change_after_first_fingerprint(path):
                nonlocal changed
                digest = sha(path)
                if Path(path).resolve() == workbook.resolve() and not changed:
                    changed = True
                    workbook.write_bytes(workbook.read_bytes() + b'concurrent edit')
                return digest
            with patch.object(original_atlas_data, 'sha', side_effect=change_after_first_fingerprint):
                with self.assertRaisesRegex(ValueError, 'Model-catch source changed'):
                    original_atlas_data.add_group_catch(root, units)
            self.assertEqual(units, before, 'Source mutation aborts the complete enrichment before any metadata is applied')


if __name__ == '__main__':
    unittest.main()
