"""Review-source regression checks for the approved Karnataka report."""
import json, sys, unittest, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
from tools.project_core.validation.researcher_review import read_report
from tools.project_core.workbooks.workbooks import sha
from tools.project_core.validation.validation_percentage_format import format_percent

ROOT = Path(__file__).resolve().parents[3]
MODEL = '32_1_Arabian_Sea_off_Karnataka_(2000)'
REPORT = ROOT / 'regions/LME/LME_032/papers/ARAB-2005/models/32_1_Arabian_Sea_off_Karnataka_(2000)' / 'model_validation/validation.docx'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class KarnatakaReviewTests(unittest.TestCase):
    def test_current_manual_table_covers_all_very_low_taxa_without_mutating_report(self):
        protected=sha(REPORT)
        with zipfile.ZipFile(REPORT) as new:
            self.assertIn('word/styles.xml',new.namelist())
            raw = new.read('word/document.xml')
            nodes = list(ET.fromstring(raw).find(W+'body'))
        text = lambda e: ''.join(t.text or '' for t in e.iter(W+'t'))
        index = next(i for i,n in enumerate(nodes) if n.tag==W+'p' and text(n)=='Very low decisions')
        self.assertEqual(text(nodes[index+2]),'Geographic evidence')
        rows = nodes[index+1].findall(W+'tr')
        self.assertEqual([text(c) for c in rows[0].findall(W+'tc')],['Affected taxa','Why confidence is very low'])
        names = [name for row in rows[1:] for name in text(row.find(W+'tc')).split('; ')]
        book = load_workbook(ROOT/'regions/LME/LME_032/papers/ARAB-2005/models/32_1_Arabian_Sea_off_Karnataka_(2000)/model_validation/taxon_mapping.xlsx',read_only=True,data_only=True)
        expected = [r[0] for r in list(book['Taxon appendix'].values)[5:] if r[5]=='Very low']; book.close()
        self.assertEqual(len(names),40); self.assertEqual(len(set(names)),40)
        self.assertEqual(set(names),set(expected))
        self.assertEqual(sha(REPORT),protected)

    def test_exact_review_identity_text_and_region_specific_appendix(self):
        name, date, summary = read_report(ROOT, REPORT, MODEL)
        self.assertEqual((name, date), ('Ido Carmel', '2026-10-01'))
        self.assertIn('rho_living = 0.58', summary['sections'][1]['rows'][0]['text'])
        self.assertIn('(34) marine mammals', summary['sections'][1]['rows'][1]['text'])
        self.assertIn('marine mammals (GE=1.5%', summary['sections'][2]['rows'][0]['text'])
        self.assertEqual(summary['sections'][4]['table'][2], ['Medium','205','49%','64%'])
        book=load_workbook(REPORT.with_name('taxon_mapping.xlsx'),read_only=True,data_only=True)
        medium=next(r for r in book['Coverage'].values if r[0]=='Medium');book.close()
        self.assertEqual(medium[1:],(205,2534929.25610556,48.52750233216555,64.2252681737161))
        self.assertEqual(summary['sections'][4]['appendix_path'], 'regions/LME/LME_032/papers/ARAB-2005/models/32_1_Arabian_Sea_off_Karnataka_(2000)/model_validation/taxon_mapping.xlsx')
        self.assertTrue(any('2019 landings' in p for p in summary['sections'][4]['reference']))
        self.assertIn('marine mammals are excluded from displayed PPR', summary['note'])
        self.assertNotIn('pinnipeds', summary['note'])

    def test_existing_assignment_table_matches_current_regional_appendix(self):
        with zipfile.ZipFile(REPORT) as z:
            body = ET.fromstring(z.read('word/document.xml')).find(W+'body')
        text = lambda e: ''.join(t.text or '' for t in e.iter(W+'t'))
        nodes = list(body)
        locations = [i for i,n in enumerate(nodes) if n.tag == W+'p' and text(n) == 'Group assignment rules']
        self.assertEqual(len(locations), 1)
        table = nodes[locations[0]+1]
        self.assertEqual(table.tag, W+'tbl')
        actual = [[text(c) for c in row.findall(W+'tc')] for row in table.findall(W+'tr')]
        self.assertEqual(actual[0], ['Plain-language rule','Confidence','PPR percentage'])
        book = load_workbook(ROOT/'regions/LME/LME_032/papers/ARAB-2005/models/32_1_Arabian_Sea_off_Karnataka_(2000)/model_validation/taxon_mapping.xlsx',read_only=True,data_only=True)
        rows = list(book['Coverage'].values); book.close()
        start = next(i for i,r in enumerate(rows) if r[0] == 'Group assignment rules')+2
        expected = []
        for row in rows[start:]:
            if row[0] is None: break
            expected.append([row[0],row[1],format_percent(str(row[2]))])
        self.assertEqual(actual[1:], expected)
        self.assertEqual(len(expected), 9)

if __name__ == '__main__': unittest.main()
