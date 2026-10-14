"""Certifying saved returns must fail closed on changed sources or evidence."""
import json
from pathlib import Path
import tempfile
import unittest
import copy
import gzip
from unittest.mock import patch

import numpy as np
import openpyxl
import pandas as pd

from tools.scientific_helpers.sppr import retained_methods as retained
from tools.scientific_helpers.sppr.method_coverage import _hash_bytes, method_settings, method_inventory, engine_identity
from tools.workflow_checks.calculations.test_sppr_method_execution import FLAGS


def typed(value):
    if isinstance(value, pd.DataFrame):
        return {'type':'DataFrame','index':typed(value.index.tolist()),
                'columns':typed(value.columns.tolist()),'index_name':typed(None),
                'columns_name':typed(None),'dtypes':[str(v) for v in value.dtypes],
                'data':[[typed(v) for v in row] for row in value.to_numpy().tolist()]}
    if isinstance(value, dict):
        return {'type':'dict','items':[[typed(k),typed(v)] for k,v in value.items()]}
    if isinstance(value, (list,tuple)):
        return {'type':type(value).__name__,'items':[typed(v) for v in value]}
    return {'type':'null'} if value is None else {'type':type(value).__name__,'value':value}


def fixture(root):
    model=root/'model'; model.mkdir(); (model/'model.json').write_text('{}',encoding='utf-8')
    results=model/'results';results.mkdir();(results/'regional_snapshot.xlsx').write_bytes(b'fixture snapshot')
    run=model/'model_validation/evidence/source_review/diagnostics';run.mkdir(parents=True)
    option=run/'GE';option.mkdir()
    axes=[{'id':1,'name':'Producer','trophic_type':'PP','synthetic':False},
          {'id':2,'name':'Consumer','trophic_type':'Regular','synthetic':False}]
    frame=pd.DataFrame([[1.],[2.]],index=[1,2],columns=[1])
    report={'status':'WARN','config':{'TE_option':'GE','det_collapse_mode':'never',
        'det_open_mode':'none','det_theta':1.,'det_external_sppr':0.,'explicit_TE':False}}
    (option/'full_return_lossless.json').write_text(json.dumps(typed((report,frame,frame,frame))),encoding='utf-8')
    for name in ('SPPR','A','L'):
        np.save(option/f'{name}.npy',frame.to_numpy())
        np.savez(option/f'{name}_masks.npz',**{k:np.zeros(frame.shape,dtype=bool) for k in
                 ('null','nan','positive_infinity','negative_infinity')})
        (option/f'{name}_axes.json').write_text(json.dumps({'rows':axes,'columns':[axes[0]],'shape':[2,1]}),encoding='utf-8')
    workbook=openpyxl.Workbook();workbook.remove(workbook.active)
    rows=[]
    for scope in ('all','inner','PP'):
        ws=workbook.create_sheet('sppr_'+scope);ws.append(['seq','group_name','new_GE'])
        for seq,name,value in [(1,'Producer',1.),(2,'Consumer',2.)]:
            ws.append([seq,name,value]);rows.append(['model',name,scope,'new_GE',value])
    workbook.save(model/'sppr_source.xlsx');workbook.close()
    source_hash=_hash_bytes((model/'model.json').read_bytes())
    flags = copy.deepcopy(FLAGS)
    flags['constructor'] = retained._constructor(flags)
    flags['constructor']['use_ee_floor'] = False
    context={'model_sha256':source_hash,'flags':flags,
             'code_identity':{'PPRCalculator.py':'engine','ModelData.py':'loader','utils.py':'utils'},
             'dependencies':{'numpy':'fixture'}}
    manifest={'canonical_model_sha256':source_hash,'effective_flags':flags,
              'coefficient_source_sha256':_hash_bytes((model/'sppr_source.xlsx').read_bytes()),
              'snapshot_sha256':_hash_bytes((results/'regional_snapshot.xlsx').read_bytes()),
              'code_identity':context['code_identity']}
    run_record={'input_sha256':source_hash,'model_id':'model','engine':context['code_identity'],
                'constructor':flags['constructor'],'method_call':{'short':False,'flat':False,'return_sppr':True},
                'thresholds':method_settings('new_GE',FLAGS)['kwargs']['thresholds'],
                'packages':{'numpy':'fixture'},'constructor_status':'CONSTRUCTED',
                'source_input_unchanged':True,'persisted_runtime_reload_exact':True,
                'reconciliation':{'scopes':True}}
    (run/'run_manifest.json').write_text(json.dumps(run_record),encoding='utf-8')
    book={'Selected model groups':{'Group SPPR':(['model_id','group','scope','method','sppr'],rows),
            'Groups':(['seq','group_name','trophic_info'],[[1,'Producer','PP'],[2,'Consumer','Regular']])}}
    return model,manifest,context,book,option


class RetainedTests(unittest.TestCase):
    def test_audited_aggregate_export_reuses_all_methods_and_blocks_low_mc_acceptance(self):
        with tempfile.TemporaryDirectory() as directory:
            model,manifest,context,book,option=fixture(Path(directory))
            keys=method_inventory(); context['code_identity']=engine_identity()
            workbook=openpyxl.load_workbook(model/'sppr_source.xlsx')
            coefficients=[]
            for scope in ('all','inner','PP'):
                ws=workbook['sppr_'+scope];ws.delete_rows(1,ws.max_row)
                ws.append(['seq','group_name',*keys])
                for seq,name,value in [(1,'Producer',1.),(2,'Consumer',2.)]:
                    values=[None if scope=='PP' and key.startswith(('SPPR_1986','SPPR_1995')) else value for key in keys]
                    ws.append([seq,name,*values])
                    coefficients.extend([['model',name,scope,key,v] for key,v in zip(keys,values)])
            mc=[{'method':key,'n_samples':100,'n_accepted':70 if key=='MC_new_GE' else 100}
                for key in keys if key.startswith('MC_')]
            ws=workbook.create_sheet('mc_diagnostics');ws.append(['method','n_samples','n_accepted'])
            for row in mc:ws.append(list(row.values()))
            ws=workbook.create_sheet('model_health');ws.append(['TE_option','status'])
            for name in ('GE','TE','With Egestion'):ws.append([name,'WARN'])
            workbook.save(model/'sppr_source.xlsx');workbook.close()
            manifest['coefficient_source_sha256']=retained._sha(model/'sppr_source.xlsx')
            manifest['code_identity']=None
            book['Selected model groups']['Group SPPR']=(book['Selected model groups']['Group SPPR'][0],coefficients)
            book['PPR']={'Annual':(['model_id','scope','method','catch_basis','unidentified','metric','status'],
                [['model',scope,key,'landings','method','ppr','ok'] for key in keys for scope in ('all','inner','PP')])}
            record={'canonical_model_sha256':context['model_sha256'],'fresh_sppr_sha256':manifest['coefficient_source_sha256'],
                'engine_hashes':context['code_identity'],'constructor':copy.deepcopy(context['flags']['constructor']),
                'methods':list(keys),'method_status':dict.fromkeys(keys,'ok'),'mc_samples':100,
                'final_mc_diagnostics':mc,'monte_carlo_worker_seed':'Unspecified native worker RNG'}
            export=option.parent/'refresh_configuration.json';export.write_text(json.dumps(record))
            output=model/'results/diagnostics'
            result=retained.reuse_methods(model,manifest,context,book,output,keys)
            self.assertEqual(set(result),set(keys))
            self.assertEqual(result['new_GE']['evidence_kind'],'retained_aggregate_export')
            self.assertEqual(result['new_GE']['retained_diagnostics_flat']['status'],'WARN')
            self.assertIn('below 80%',result['MC_new_GE']['annual_status']['all'])
            self.assertEqual(result['MC_new_GE']['scopes']['all']['Consumer'],2.)
            self.assertIn('unsupported',result['SPPR_1986']['annual_status']['PP'])
            record['fresh_sppr_sha256']='0'*64;export.write_text(json.dumps(record))
            self.assertEqual(retained.reuse_methods(model,manifest,context,book,output,keys),{})

    def test_exact_evidence_is_reused_without_calling_engine(self):
        with tempfile.TemporaryDirectory() as directory:
            model,manifest,context,book,_=fixture(Path(directory))
            result=retained.reuse_methods(model,manifest,context,book,model/'results/diagnostics',['new_GE','new_TE_EEfix'])
            self.assertEqual(list(result),['new_GE'])
            self.assertEqual(result['new_GE']['status'],'reused')
            self.assertEqual(result['new_GE']['scopes']['all']['Consumer'],2.)
            self.assertTrue((model/'results/diagnostics'/result['new_GE']['evidence']['path']).is_file())

    def test_changed_matrix_or_model_is_not_reused(self):
        with tempfile.TemporaryDirectory() as directory:
            model,manifest,context,book,option=fixture(Path(directory))
            np.save(option/'SPPR.npy',np.zeros((2,1)))
            self.assertEqual(retained.reuse_methods(model,manifest,context,book,model/'results/diagnostics',['new_GE']),{})
            (model/'model.json').write_text('{"changed":true}',encoding='utf-8')
            self.assertEqual(retained.reuse_methods(model,manifest,context,book,model/'results/diagnostics',['new_GE']),{})

    def test_decoder_does_not_evaluate_arbitrary_objects(self):
        with self.assertRaises(ValueError):
            retained.decode_typed({'type':'pickle','value':'anything'})

    def test_unknown_code_identity_refuses_reuse_without_crashing(self):
        with tempfile.TemporaryDirectory() as directory:
            model,manifest,context,book,_=fixture(Path(directory))
            for unknown in (None, '', [], 'unknown'):
                manifest['code_identity']=unknown
                self.assertEqual(retained.reuse_methods(model,manifest,context,book,
                    model/'results/diagnostics',['new_GE']),{})

    def test_explicit_active_run_certifies_legacy_manifest_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            model,manifest,context,book,option=fixture(Path(directory))
            manifest['input_sha256']=manifest.pop('canonical_model_sha256')
            manifest.pop('coefficient_source_sha256')
            manifest.pop('code_identity')
            manifest['engine_hashes']={'PPRCalculator.py':'obsolete descriptive field'}
            manifest['evidence']=str((option.parent/'run_manifest.json').relative_to(model.parent))
            result=retained.reuse_methods(model,manifest,context,book,
                model/'results/diagnostics',['new_GE'])
            self.assertEqual(result['new_GE']['status'],'reused')
            book['Selected model groups']['Group SPPR'][1][0][4]=12.
            self.assertEqual(retained.reuse_methods(model,manifest,context,book,
                model/'results/diagnostics',['new_GE']),{})

    def test_explicit_active_run_does_not_fallback_to_other_run(self):
        with tempfile.TemporaryDirectory() as directory:
            model,manifest,context,book,option=fixture(Path(directory))
            manifest.pop('coefficient_source_sha256');manifest.pop('code_identity')
            manifest['evidence']='missing/run_manifest.json'
            self.assertEqual(retained.reuse_methods(model,manifest,context,book,
                model/'results/diagnostics',['new_GE']),{})


class SavedMethodAdoptionTests(unittest.TestCase):
    def saved_fixture(self, directory, *, key='MC_new_GE', value=2., status='provisional: retained warning'):
        model, manifest, context, book, _ = fixture(Path(directory))
        manifest['code_identity'] = None
        manifest['effective_flags'] = None
        context = copy.deepcopy(context)
        # Administrative adoption with unknown historical settings must not
        # claim an EE-floor configuration or a numerical engine upgrade.
        context['flags']['constructor'].pop('use_ee_floor')
        context['flags']['constructor'].pop('ee_floor')
        context['adapter_sha256'] = 'new adapter must not trigger existing MC'
        context['dependencies'] = {'numpy': 'changed package metadata'}
        for row in book['Selected model groups']['Group SPPR'][1]:
            row[3] = key
            row[4] = value
        wb = openpyxl.load_workbook(model/'sppr_source.xlsx')
        for scope in ('all', 'inner', 'PP'):
            wb['sppr_'+scope].cell(1, 3, key)
            for row in (2, 3):
                wb['sppr_'+scope].cell(row, 3).value = value
        wb.save(model/'sppr_source.xlsx'); wb.close()
        manifest['coefficient_source_sha256'] = retained._sha(model/'sppr_source.xlsx')
        book['Overview'] = {'Settings': (['field', 'value'], [
            ['results_model_id', model.name], ['results_model_sha256', context['model_sha256']]])}
        book['PPR'] = {'Annual': (['model_id', 'scope', 'method', 'catch_basis', 'unidentified', 'metric', 'status', 2000],
            [[model.name, scope, key, 'landings', 'method', 'ppr', status, None if value is None else value*3]
             for scope in ('all', 'inner', 'PP')])}
        book['Diagnostics'] = {
            'mc_diagnostics': (['method', 'n_samples', 'n_accepted', 'warning'], [[key, 100, 70, 'original warning']]),
            'model_health': (['TE_option', 'status'], [['GE', 'FAIL']]),
            'Sensitivity': (['method', 'note'], [[key, 'keep this exact note']])}
        return model, manifest, context, book

    def test_unknown_provenance_adopts_exact_saved_rows_without_engine_or_mutation(self):
        with tempfile.TemporaryDirectory() as directory:
            model, manifest, context, book = self.saved_fixture(directory)
            original = copy.deepcopy(book)
            with patch.object(retained, '_engine', side_effect=AssertionError('must not load science')), \
                 patch.object(retained, 'method_settings', side_effect=AssertionError('must not invent settings')):
                result = retained.adopt_saved_methods(model, manifest, context, book,
                    model/'results/diagnostics', ['MC_new_GE', 'new_GE'])
            self.assertEqual(set(result), {'MC_new_GE'})
            outcome = result['MC_new_GE']
            self.assertTrue(outcome['preserve_saved'])
            self.assertEqual(outcome['annual_status']['all'], 'provisional: retained warning')
            self.assertEqual(outcome['mc_diagnostics']['n_accepted'], 70)
            self.assertEqual(outcome['scopes']['all']['Consumer'], 2.)
            self.assertNotIn('settings', outcome)
            self.assertEqual(book, original)
            evidence_path = model/'results/diagnostics'/outcome['evidence']['path']
            certificate = json.loads(gzip.decompress(evidence_path.read_bytes()))
            self.assertEqual(certificate['evidence_kind'], 'retained_saved_method')
            self.assertEqual(certificate['historical_execution']['code_identity'], None)
            self.assertEqual(certificate['historical_execution']['effective_flags'], None)
            self.assertFalse(certificate['historical_execution']['verified_by_adoption'])
            self.assertEqual(certificate['availability']['individual_mc_samples'], 'not_recovered_by_adoption')
            self.assertEqual(certificate['saved_diagnostic_tables']['Sensitivity'],
                             [original['Diagnostics']['Sensitivity'][0], original['Diagnostics']['Sensitivity'][1]])

    def test_complete_empty_failed_and_not_run_rows_are_existing_results(self):
        for status in ('FAIL: solver diverged', 'NOT_RUN: missing prerequisite', 'unavailable: MC acceptance 70/100 below 80%'):
            with self.subTest(status=status), tempfile.TemporaryDirectory() as directory:
                model, manifest, context, book = self.saved_fixture(directory, value=None, status=status)
                result = retained.adopt_saved_methods(model, manifest, context, book,
                    model/'results/diagnostics', ['MC_new_GE'])
                self.assertEqual(result['MC_new_GE']['annual_status'], dict.fromkeys(('all', 'inner', 'PP'), status))
                self.assertTrue(all(value is None for value in result['MC_new_GE']['scopes']['all'].values()))

    def test_changed_model_or_known_scientific_flags_requires_explicit_refresh(self):
        with tempfile.TemporaryDirectory() as directory:
            model, manifest, context, book = self.saved_fixture(directory)
            manifest['effective_flags'] = {'constructor': copy.deepcopy(FLAGS['constructor']),
                                           'diagnostics': {'det_collapse_mode': 'auto'}}
            context['flags']['diagnostics'] = {'det_collapse_mode': 'never'}
            with self.assertRaisesRegex(ValueError, 'configuration|flags'):
                retained.adopt_saved_methods(model, manifest, context, book, model/'out', ['MC_new_GE'])
            manifest['effective_flags'] = None
            (model/'model.json').write_text('{"changed":true}')
            with self.assertRaisesRegex(ValueError, 'model'):
                retained.adopt_saved_methods(model, manifest, context, book, model/'out', ['MC_new_GE'])
            self.assertFalse((model/'out').exists())

    def test_unknown_historical_floor_cannot_be_adopted_as_enabled(self):
        with tempfile.TemporaryDirectory() as directory:
            model, manifest, context, book = self.saved_fixture(directory)
            context['flags']['constructor'].update(use_ee_floor=True, ee_floor=.1)
            with self.assertRaisesRegex(ValueError, 'EE-floor calculation identity'):
                retained.adopt_saved_methods(model, manifest, context, book, model/'out', ['MC_new_GE'])
            self.assertFalse((model/'out').exists())

    def test_missing_rows_or_status_are_not_certified_as_complete(self):
        for missing in ('row', 'status', 'duplicate'):
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as directory:
                model, manifest, context, book = self.saved_fixture(directory)
                if missing == 'row': book['Selected model groups']['Group SPPR'][1].pop()
                elif missing == 'status': book['PPR']['Annual'][1].pop()
                else: book['Selected model groups']['Group SPPR'][1].append(book['Selected model groups']['Group SPPR'][1][0][:])
                self.assertEqual(retained.adopt_saved_methods(model, manifest, context, book,
                    model/'out', ['MC_new_GE']), {})

    def test_source_disagreement_rejects_adoption_without_replacing_values(self):
        with tempfile.TemporaryDirectory() as directory:
            model, manifest, context, book = self.saved_fixture(directory)
            book['Selected model groups']['Group SPPR'][1][0][4] += 1e-8
            with self.assertRaisesRegex(ValueError, 'source'):
                retained.adopt_saved_methods(model, manifest, context, book, model/'out', ['MC_new_GE'])
            self.assertFalse((model/'out').exists())

    def test_prior_full_return_evidence_is_hash_checked_and_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            model, manifest, context, book = self.saved_fixture(directory)
            output = model/'results/diagnostics/method_coverage'; output.mkdir(parents=True)
            full = output/'MC_new_GE.json.gz'; full.write_bytes(b'original full return evidence')
            original = full.read_bytes()
            manifest['method_coverage'] = {'methods': {'MC_new_GE': {'status': 'ok',
                'evidence': {'path': 'method_coverage/MC_new_GE.json.gz', 'sha256': retained._sha(full)}}}}
            result = retained.adopt_saved_methods(model, manifest, context, book, output, ['MC_new_GE'])
            self.assertEqual(result['MC_new_GE']['evidence']['path'], 'MC_new_GE.adopted.json.gz')
            self.assertEqual(full.read_bytes(), original)
            self.assertEqual(result['MC_new_GE']['prior_method_trace']['evidence_check']['status'], 'matching')
            full.write_bytes(b'changed evidence')
            result = retained.adopt_saved_methods(model, manifest, context, book, output, ['MC_new_GE'])
            self.assertEqual(result['MC_new_GE']['prior_method_trace']['evidence_check']['status'], 'hash_mismatch')


if __name__=='__main__':unittest.main()
