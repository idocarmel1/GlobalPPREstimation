"""Explicit engine migration retains unaffected results and historical evidence."""
import base64
import copy
import gzip
import io
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import pandas as pd

from tools.project_core.calculations import coverage


class EEFloorMigrationTests(unittest.TestCase):
    def helper(self):
        name = 'tools.project_core.calculations.ee_floor_migration'
        self.assertIsNotNone(importlib.util.find_spec(name), 'EE-floor migration helper is missing')
        return __import__(name, fromlist=['migration'])

    def calculator(self):
        series = pd.Series([2., 1.], index=[1, 2])
        result = SimpleNamespace(seq2name={1:'Fish', 2:'Algae'}, name2seq={'Fish':1, 'Algae':2})
        for name in ('p','q','catch','predation','growth','immigration','emigration',
                     'net_migration','respiration','egestion','M0','det_export','EE','GE','GS','TL'):
            setattr(result, name, series.copy())
        result._DC = pd.DataFrame([[0., 1.], [0., 0.]], index=[1,2], columns=[1,2])
        result._det_fate = pd.DataFrame(index=[1,2])
        result.get_PP_seq = lambda: [2]
        result.get_DET_seq = lambda: []
        result.get_Import_seq = lambda: []
        result.get_TL = lambda **kwargs: result.TL.copy()
        return result

    def test_target_flags_preserve_original_and_sampling_policy(self):
        helper = self.helper()
        original = {'constructor': {'normalize_DC':True}, 'diagnostics': {'det_collapse_mode':'auto'},
                    'mc': {'n_samples':75}, 'engine_source_hashes': {'PPRCalculator.py':'old'}}
        frozen = copy.deepcopy(original)
        requested = helper.target_flags(original, True, .1, 'a'*64)
        self.assertEqual(original, frozen)
        self.assertEqual(requested['mc'], original['mc'])
        self.assertEqual(requested['diagnostics'], original['diagnostics'])
        self.assertEqual(requested['engine_source_hashes'], original['engine_source_hashes'])
        self.assertEqual(requested['constructor']['use_ee_floor'], True)
        self.assertEqual(requested['constructor']['ee_floor'], .1)
        self.assertEqual(requested['calculation_engine_upgrade'], 'ee_floor_v1')
        self.assertEqual(requested['calculation_engine_sha256'], 'a'*64)
        self.assertEqual(helper.target_flags(original,True,0,'a'*64)['constructor']['ee_floor'],0.)
        for value in (True, -1, 1.01, float('nan')):
            with self.subTest(value=value), self.assertRaises(ValueError):
                helper.target_flags(original, True, value, 'a'*64)

    def test_unchanged_numerical_runtime_keeps_every_method(self):
        helper = self.helper(); old = self.calculator(); new = copy.deepcopy(old)
        new.source_EE = old.EE.copy(); new.ee_floor_info = {'affected_groups':[]}
        methods = ('SPPR_1986', 'SPPR_1995_TE0.1', 'new_GE', 'MC_new_GE')
        retained = helper.compatible_methods(old, new, methods)
        self.assertEqual(set(retained), set(methods))

    def test_ee_changes_keep_only_fixed_trophic_level_methods(self):
        helper = self.helper(); old = self.calculator(); new = copy.deepcopy(old)
        new.EE.loc[1] = .1; new.M0.loc[1] = .3
        methods = ('SPPR_1986', 'SPPR_1995_TE0.1', 'new_GE', 'MC_new_GE')
        self.assertEqual(set(helper.compatible_methods(old, new, methods)), set(methods[:2]))
        new._DC.loc[1,2] = .9
        self.assertEqual(helper.compatible_methods(old, new, methods), {})

    def test_baseline_package_preserves_exact_outgoing_bytes(self):
        helper = self.helper()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary); model = root/'model'; model.mkdir(); staged = root/'staged'; staged.mkdir()
            source = model/'sppr_source.xlsx'; source.write_bytes(b'17-digit source bytes')
            evidence = model/'results/diagnostics/method_coverage/MC_new_GE.json.gz'
            evidence.parent.mkdir(parents=True); evidence.write_bytes(b'original random draws')
            reference = {'path':'method_coverage/MC_new_GE.json.gz', 'sha256':coverage._context.__name__}
            from tools.project_core.workbooks.workbooks import sha
            reference['sha256'] = sha(evidence)
            manifest = {'effective_flags':None, 'method_coverage': {'methods': {'MC_new_GE': {'evidence':reference}}}}
            package = helper.save_baseline(model, {source:source.read_bytes()}, manifest, staged)
            with gzip.open(staged/package['path'], 'rt', encoding='utf-8') as handle:
                saved = json.load(handle)
            self.assertIsNone(saved['historical_manifest']['effective_flags'])
            self.assertEqual(base64.b64decode(saved['files']['sppr_source.xlsx']['base64']), source.read_bytes())
            self.assertEqual(base64.b64decode(saved['files']['results/diagnostics/method_coverage/MC_new_GE.json.gz']['base64']), evidence.read_bytes())

    def test_explicit_changed_method_preserves_provisional_and_mc_restrictions(self):
        self.assertIn('migration', __import__('inspect').signature(coverage._apply_outcomes).parameters,
                      'Explicit migration status preservation is missing')
        for method, prior in [('new_GE','provisional: original diagnostic FAIL'),
                              ('MC_new_GE','unavailable: Monte Carlo quality restriction')]:
            book = {'Selected model groups': {'Groups': (['group_name'], [['Fish']]),
                'Group SPPR': (['model_id','group','scope','method','sppr'], [['A','Fish','all',method,5.]])},
                'PPR': {'Annual': (['model_id','scope','method','catch_basis','unidentified','metric','status'],
                    [['A','all',method,'landings','method','ppr',prior]])}}
            outcome = {method: {'status':'ok', 'scopes': {'all': {'Fish':7.}}, 'annual_status': {'all':'ok'}}}
            coverage._apply_outcomes(book, 'A', outcome, [method], migration=True)
            self.assertEqual(book['PPR']['Annual'][1][0][6], prior)
            self.assertEqual(book['Selected model groups']['Group SPPR'][1][0][4], 7.)

    def test_source_overwrite_is_explicit_and_preserves_unaffected_cells(self):
        self.assertIn('replace_methods', __import__('inspect').signature(coverage._update_source).parameters,
                      'Explicit migration source replacement is missing')
        import openpyxl
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'sppr_source.xlsx'
            workbook = openpyxl.Workbook(); workbook.remove(workbook.active)
            for scope in ('all','inner','PP'):
                sheet = workbook.create_sheet('sppr_'+scope)
                sheet.append(['seq','group_name','old','changed']); sheet.append([1,'Fish',1.234567890123457,5.])
            workbook.save(path); workbook.close()
            book = {'Selected model groups': {'Groups': (['seq','group_name'], [[1,'Fish']]),
                'Group SPPR': (['model_id','group','scope','method','sppr'],
                    [['A','Fish',scope,method,value] for scope in ('all','inner','PP')
                     for method,value in [('old',1.234567890123457),('changed',7.)]])}}
            coverage._update_source(path, book, {}, replace_methods={'changed'},
                                    effective_groups=pd.DataFrame({'group_name':['Fish'],'ee':[.1],'source_EE':[.01]}, index=[1]))
            workbook = openpyxl.load_workbook(path, read_only=True)
            for scope in ('all','inner','PP'):
                self.assertEqual(workbook['sppr_'+scope]['C2'].value,1.234567890123457)
                self.assertEqual(workbook['sppr_'+scope]['D2'].value,7.)
            self.assertIn('source_EE', list(next(workbook['groups_df'].values)))
            workbook.close()

    def test_source_replacement_clears_finite_coefficients_for_unavailable_methods(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temporary:
            model=Path(temporary);source=model/'sppr_source.xlsx'
            workbook=openpyxl.Workbook();workbook.remove(workbook.active)
            for scope in ('all','inner','PP'):
                sheet=workbook.create_sheet('sppr_'+scope)
                sheet.append(['seq','group_name','SPPR_1986','MC_new_GE'])
                sheet.append([1,'Fish',1.234567890123457,5.])
            workbook.save(source);workbook.close()
            book={'Selected model groups':{'Groups':(['seq','group_name'],[[1,'Fish']]),
                'Group SPPR':(['model_id','group','scope','method','sppr'],
                    [['model','Fish',scope,method,value] for scope in ('all','inner','PP')
                     for method,value in [('SPPR_1986',1.234567890123457),('MC_new_GE',None)]])}}
            outcomes={'MC_new_GE':{'status':'failed','reason':'unavailable draw result',
                      'annual_status':{scope:'unavailable: failed' for scope in ('all','inner','PP')}}}
            coverage._update_source(source,book,outcomes,replace_methods={'MC_new_GE'})
            self.assertTrue(coverage._source_matches(model,book,('SPPR_1986','MC_new_GE')))
            workbook=openpyxl.load_workbook(source,read_only=True,data_only=True)
            try:
                for scope in ('all','inner','PP'):
                    self.assertIsNone(workbook['sppr_'+scope]['D2'].value)
                    self.assertEqual(workbook['sppr_'+scope]['C2'].value,1.234567890123457)
            finally:workbook.close()

    def test_command_defaults_to_validated_migration_without_all_switch(self):
        from tools.project_core.registry import update_project
        self.assertTrue(hasattr(update_project,'parse_arguments'), 'Migration CLI parsing is missing')
        parsed = update_project.parse_arguments(['--migrate-ee-floor'])
        self.assertTrue(parsed.migrate_ee_floor); self.assertEqual(parsed.ee_floor,.1)
        parsed = update_project.parse_arguments(['--migrate-ee-floor','--no-ee-floor'])
        self.assertTrue(parsed.no_ee_floor)
        from contextlib import redirect_stderr
        with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            update_project.parse_arguments(['--ee-floor','.2','--all'])

    def test_upgraded_selection_identity_uses_archived_loader_and_current_calculator(self):
        from tools.project_core.calculations.selection import execution_identity
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); model=root/'model'; model.mkdir(); (model/'model.json').write_text('{}')
            engine=root/'tools/scientific_code/PPREstimation'; engine.mkdir(parents=True)
            for name in ('PPRCalculator.py','ModelData.py','utils.py'):(engine/name).write_text('# current '+name)
            flags={'constructor':{'use_ee_floor':True,'ee_floor':.1},'calculation_engine_upgrade':'ee_floor_v1'}
            identity={'PPRCalculator.py':'archived_loader','calculation/PPRCalculator.py':'current_calculator',
                      'create_PPRS_excel.py':'registry'}
            with patch('tools.scientific_helpers.sppr.method_coverage.engine_identity',return_value=identity):
                actual=execution_identity(root,model,flags=flags)
            self.assertEqual(actual['code_identity'], {'PPRCalculator.py':'archived_loader',
                             'calculation/PPRCalculator.py':'current_calculator'})

    def test_publication_rolls_back_new_floor_evidence_and_keeps_author_race(self):
        from tools.project_core.registry.publication import publication_transaction
        from tools.project_core.workbooks.workbooks import write_book
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary); model=root/'model'; model.mkdir()
            write_book(root/'Project.xlsx', {'Models & coverage': {'Models': (['unit_id','model_id'],[])}})
            evidence=model/'results/diagnostics/ee_floor/baseline.json.gz'
            other=evidence.with_name('researcher-note.txt')
            with self.assertRaisesRegex(RuntimeError,'downstream failure'):
                with publication_transaction(root,extra_model_dirs=[model]) as transaction:
                    evidence.parent.mkdir(parents=True); evidence.write_bytes(b'exact old bytes')
                    transaction.checkpoint([evidence])
                    other.write_text('author evidence')
                    raise RuntimeError('downstream failure')
            self.assertFalse(evidence.exists()); self.assertEqual(other.read_text(),'author evidence')

    def test_explicit_migration_changes_only_affected_methods_with_lossless_baseline(self):
        self.assertIn('ee_floor', __import__('inspect').signature(coverage._complete_model).parameters,
                      'Explicit model migration is missing')
        from tools.workflow_checks.calculations.test_coverage_preservation import CoveragePreservationTests
        from tools.project_core.workbooks.workbooks import read_book, rows, sha
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); model,path,original,metadata=CoveragePreservationTests().package(root)
            prior_rows=copy.deepcopy(rows(original,'Selected model groups','Group SPPR'))
            flags={'constructor':{'normalize_DC':True},'mc':{'n_samples':100}}
            migration={'flags':self.helper().target_flags(flags,True,.1,'a'*64),
                       'compatible':{'MC_new_GE':{'reason':'no changed dependencies'}},
                       'effective_groups':pd.DataFrame({'group_name':['Fish'],'ee':[.1],'source_EE':[.01]},index=[1]),
                       'before_runtime':{},'after_runtime':{},'old_warnings':[],'new_warnings':[],
                       'info':{'use_ee_floor':True,'ee_floor':.1},'ledger':pd.DataFrame({'seq':[1],'before':[.01],'after':[.1]}),
                       'calculation_code':b'current calculator'}
            # The fixture lacks modern coverage evidence; establish its actual saved
            # identity using ordinary completion first, without rerunning it.
            adapter_version=['fixture']
            def context(root,model_path,resolved):
                return {'model_sha256':sha(model_path),'flags':resolved,'code_identity':{},'adapter_sha256':adapter_version[0]}
            def executor(model_path,resolved,keys,directory,progress):
                self.assertEqual(keys,['new_GE'])
                evidence=directory/'new_GE.json'; evidence.write_text('new numerical execution')
                return {'methods':{'new_GE':{'status':'ok','scopes':{scope:{'Fish':7.} for scope in ('all','inner','PP')},
                    'annual_status':{scope:'ok' for scope in ('all','inner','PP')},
                    'evidence':{'path':evidence.name,'sha256':sha(evidence)}}}}
            with patch('tools.scientific_helpers.sppr.configuration.resolve_audited_configuration',return_value=flags), \
                 patch.object(coverage,'_context',side_effect=context), \
                 patch('tools.scientific_helpers.sppr.method_coverage.method_settings',return_value={'method':'new_GE'}):
                coverage._complete_model(root,metadata,('new_GE','MC_new_GE'),
                    progress=lambda value:None,compute=lambda *args,**kwargs: self.fail('saved methods were rerun'))
                with patch('tools.project_core.calculations.ee_floor_migration.prepare',return_value=migration):
                    changed,report=coverage._complete_model(root,metadata,('new_GE','MC_new_GE'),
                        progress=lambda value:None,compute=executor,ee_floor={'use_ee_floor':True,'ee_floor':.1})
            saved=read_book(path)
            self.assertTrue(changed)
            self.assertEqual([row for row in rows(saved,'Selected model groups','Group SPPR') if row[3]=='MC_new_GE'],
                             [row for row in prior_rows if row[3]=='MC_new_GE'])
            self.assertTrue(all(row[4]==7. for row in rows(saved,'Selected model groups','Group SPPR') if row[3]=='new_GE'))
            self.assertEqual(saved['Selected model groups']['Groups'],original['Selected model groups']['Groups'])
            final=json.loads((model/'results/result_manifest.json').read_text())
            baseline=final['ee_floor_migration']['baseline']
            self.assertTrue((model/'results/diagnostics'/baseline['path']).is_file())
            self.assertEqual(final['effective_flags']['mc']['n_samples'],100)
            before={p:p.read_bytes() for p in model.rglob('*') if p.is_file()}
            regional_bytes=path.read_bytes()
            adapter_version[0]='metadata update'
            with patch('tools.scientific_helpers.sppr.configuration.resolve_audited_configuration',return_value=migration['flags']), \
                 patch.object(coverage,'_context',side_effect=context), \
                 patch('tools.project_core.calculations.ee_floor_migration.prepare',return_value=None):
                changed,_=coverage._complete_model(root,metadata,('new_GE','MC_new_GE'),progress=lambda value:None,
                    compute=lambda *args,**kwargs:self.fail('Fresh migrated package was rerun'),
                    ee_floor={'use_ee_floor':True,'ee_floor':.1})
            self.assertFalse(changed); self.assertEqual(path.read_bytes(),regional_bytes)
            self.assertEqual({p:p.read_bytes() for p in model.rglob('*') if p.is_file()},before)
            mc_evidence=model/'results/diagnostics'/final['method_coverage']['methods']['MC_new_GE']['evidence']['path']
            previous_mc=mc_evidence.read_bytes()
            with patch('tools.scientific_helpers.sppr.configuration.resolve_audited_configuration',return_value=migration['flags']), \
                 patch.object(coverage,'_context',side_effect=context), \
                 patch.object(coverage,'_arithmetic_identity',return_value={'metadata':'updated'}):
                changed,_=coverage._complete_model(root,metadata,('new_GE','MC_new_GE'),progress=lambda value:None,
                    compute=lambda *args,**kwargs:self.fail('Arithmetic metadata update redrew Monte Carlo'))
            self.assertTrue(changed);self.assertEqual(mc_evidence.read_bytes(),previous_mc)

    def test_changed_method_catch_restriction_clears_disallowed_numerators(self):
        helper=self.helper()
        original={'PPR': {'Annual': ([],[
            ['A','all','new_GE','catch','method','ppr','FAIL: original catch-specific restriction',None]])}}
        current={'PPR': {'Annual': ([],[
            ['A','all','new_GE','catch','method','ppr','ok',14.]])},
            'PPR–NPP': {'Ratios': ([],[['A','all','new_GE','catch','method','NPP',1.]])}}
        helper.preserve_restrictions(current,original)
        self.assertEqual(current['PPR']['Annual'][1][0][6],'FAIL: original catch-specific restriction')
        self.assertIsNone(current['PPR']['Annual'][1][0][7])
        self.assertIsNone(current['PPR–NPP']['Ratios'][1][0][6])

    def test_adapter_only_change_keeps_complete_method_evidence(self):
        self.assertTrue(hasattr(coverage,'_same_context'), 'Narrow coverage context compatibility is missing')
        saved={'model_sha256':'same','flags':{'constructor':{'ee_floor':.1}},
               'code_identity':{'calculation/PPRCalculator.py':'same'},'dependencies':{'numpy':'same'},
               'adapter_sha256':'historical'}
        current={**saved,'adapter_sha256':'metadata_update'}
        self.assertTrue(coverage._same_context(saved,current))
        for key,value in [('flags',{}),('code_identity',{}),('dependencies',{}),('model_sha256','changed')]:
            self.assertFalse(coverage._same_context(saved,{**current,key:value}))

    def test_source_serialization_preserves_concurrent_author_edit(self):
        import openpyxl
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temporary:
            source=Path(temporary)/'sppr_source.xlsx'; workbook=openpyxl.Workbook()
            workbook.active.title='sppr_all'; workbook.active.append(['seq','group_name']); workbook.save(source);workbook.close()
            original_save=openpyxl.Workbook.save
            def concurrent_save(workbook,filename):
                original_save(workbook,filename)
                if Path(filename)!=source:source.write_bytes(b'concurrent author edit')
            with patch.object(openpyxl.Workbook,'save',autospec=True,side_effect=concurrent_save), \
                 self.assertRaisesRegex(ValueError,'changed'):
                coverage._update_source(source,{}, {})
            self.assertEqual(source.read_bytes(),b'concurrent author edit')

    def test_catch_only_failure_does_not_override_landings_seed(self):
        book={'Selected model groups': {'Groups': (['group_name'],[['Fish']]),
            'Group SPPR': ([],[['A','Fish','all','new_GE',5.]])},
            'PPR': {'Annual': ([],[['A','all','new_GE','landings','method','ppr','ok'],
                                 ['A','all','new_GE','catch','method','ppr','FAIL: catch-only restriction']])}}
        outcomes={'new_GE':{'status':'ok','scopes':{'all':{'Fish':7.}},'annual_status':{'all':'ok'}}}
        coverage._apply_outcomes(book,'A',outcomes,['new_GE'],migration=True)
        self.assertEqual(book['PPR']['Annual'][1][0][6],'ok')

    def test_reexecuted_equal_point_coefficients_invalidate_old_bounds(self):
        self.assertIn('invalidated_methods',__import__('inspect').signature(coverage._arithmetic).parameters,
                      'Scientific reruns must invalidate old bounds despite equal point values')
        from tools.workflow_checks.calculations.test_coverage_preservation import CoveragePreservationTests
        from tools.project_core.calculations.regional import set_result_hash
        from tools.project_core.workbooks.workbooks import rows, YEARS
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary);model,path,original,_=CoveragePreservationTests().package(root)
            for method in ('new_GE','MC_new_GE'):
                original['PPR']['Annual'][1].append(['A','all',method,'landings','method','min_tC','ok',*([.4]*len(YEARS))])
            set_result_hash(original); current=copy.deepcopy(original)
            coverage._arithmetic(current,path,model/'model.json',{},original,migration=True,invalidated_methods={'new_GE'})
            retained=[row[2] for row in rows(current,'PPR','Annual') if row[5]=='min_tC']
            self.assertEqual(retained,['MC_new_GE'])

    def test_retention_rejects_changed_known_dependencies_and_edited_rows(self):
        helper=self.helper()
        self.assertIn('original_context',__import__('inspect').signature(helper.retained_outcomes).parameters,
                      'Retained floor methods must authenticate their original calculation context')
        from tools.project_core.workbooks.workbooks import sha, digest_tables
        with tempfile.TemporaryDirectory() as temporary:
            model=Path(temporary)/'model'; evidence_dir=model/'results/diagnostics/method_coverage';evidence_dir.mkdir(parents=True)
            (model/'model.json').write_text('{}'); evidence=evidence_dir/'MC_new_GE.json';evidence.write_text('old execution')
            staged=Path(temporary)/'staged';staged.mkdir()
            values=[['model','Fish',scope,'MC_new_GE',5.] for scope in ('all','inner','PP')]
            book={'Selected model groups':{'Groups':(['group_name'],[['Fish']]),
                'Group SPPR':(['model_id','group','scope','method','sppr'],values)},
                  'PPR':{'Annual':(['model_id','scope','method','catch_basis','unidentified','metric','status'],
                    [['model',scope,'MC_new_GE','landings','method','ppr','ok'] for scope in ('all','inner','PP')])}}
            original={'model_sha256':sha(model/'model.json'),'flags':{'constructor':{}},'code_identity':{},'dependencies':{'numpy':'old'}}
            manifest={'canonical_model_sha256':sha(model/'model.json'),'method_coverage':{'context':original,
                'methods':{'MC_new_GE':{'status':'ok','group_rows_sha256':digest_tables(values),
                    'evidence':{'path':'method_coverage/MC_new_GE.json','sha256':sha(evidence)}}}}}
            kwargs={'original_context':{**original,'dependencies':{'numpy':'new'}},'model_dir':model}
            self.assertEqual(helper.retained_outcomes(manifest,{'MC_new_GE':{}},book,staged,{}, {},**kwargs),{})
            kwargs['original_context']=original;book['Selected model groups']['Group SPPR'][1][0][4]=8.
            self.assertEqual(helper.retained_outcomes(manifest,{'MC_new_GE':{}},book,staged,{}, {},**kwargs),{})

    def test_registry_proof_preserves_nested_numeric_imports(self):
        helper=self.helper(); import ast,hashlib
        from unittest.mock import patch
        source='def numeric():\n    import json\n    return 3\n'
        protected=hashlib.sha256(ast.dump(ast.parse(source),include_attributes=False).encode()).hexdigest()
        with tempfile.TemporaryDirectory() as temporary:
            engine=Path(temporary); current=engine/'create_PPRS_excel.py'; current.write_text(source)
            with patch('tools.scientific_helpers.sppr.method_coverage.ENGINE',engine), \
                 patch.object(helper,'REGISTRY_PROTECTED_AST_SHA256',protected):
                self.assertIsNotNone(helper.registry_compatibility(helper.REGISTRY_PREDECESSOR_SHA256))
                current.write_text(source.replace('import json','import inspect'))
                self.assertIsNone(helper.registry_compatibility(helper.REGISTRY_PREDECESSOR_SHA256))


if __name__ == '__main__': unittest.main()
