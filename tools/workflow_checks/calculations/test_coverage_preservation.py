"""Method completion never replaces an already saved scientific result."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from tools.project_core.calculations import coverage
from tools.project_core.calculations.regional import recalculate, set_setting, set_result_hash
from tools.project_core.workbooks.workbooks import read_book, write_book, sha, rows, input_hash, YEARS, records
from tools.workflow_checks.selection.test_snapshots import fixture


class CoveragePreservationTests(unittest.TestCase):
    def test_publication_preserves_17_digit_retained_values_after_rows_and_tables_move(self):
        import zipfile
        from xml.etree import ElementTree as etree
        from tools.project_core.validation.researcher_review import table_rows, sheet_xml, S
        from tools.project_core.calculations.regional import result_hash
        from tools.project_core.workbooks.workbooks import overview
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'region.xlsx'
            original = {'Overview': {'Settings': (['field','value'], [])},
                'PPR': {
                    'Annual': (['model_id','scope','method','catch_basis','unidentified','metric','status',1950],
                        [['A','all','EwE_TE_noEE','catch','method','ppr','ok',216219124.64603513]]),
                    'Taxon SPPR': (['model_id','taxon','scope','method','sppr'],
                        [['A','fish','all','EwE_TE_noEE',1.2345678901234567]])}}
            write_book(path, original)
            # This is the observed old Excel value, not an openpyxl-rounded fixture.
            with zipfile.ZipFile(path) as archive:
                member, sheet = sheet_xml(archive,'PPR')
                for table, text in [('Annual','216219124.64603513'), ('Taxon SPPR','1.2345678901234567')]:
                    _, header, _, table_data = table_rows(archive,'PPR',table)
                    address = table_data[0][0][-1].get('r')
                    sheet.find(f'.//{{{S}}}c[@r="{address}"]/{{{S}}}v').text = text
                data = etree.tostring(sheet,encoding='utf-8')
                parts = [(info,data if info.filename == member else archive.read(info.filename)) for info in archive.infolist()]
            with zipfile.ZipFile(path,'w') as archive:
                for info, data in parts: archive.writestr(info,data)
            original = read_book(path)
            self.assertEqual(original['PPR']['Annual'][1][0][-1],216219124.64603513)
            extended = copy.deepcopy(original)
            extended['PPR']['Annual'][1].insert(0,['A','all','new_method','catch','method','ppr','ok',7.])
            extended['PPR']['Taxon SPPR'][1].insert(0,['A','fish','all','new_method',9.])
            set_setting(extended,'calculation_input_sha256',input_hash(extended)); set_result_hash(extended)
            saved = coverage._save_calculation(path,extended,Path(temporary)/'model.json',{},original)
            self.assertEqual(saved['PPR']['Annual'][1][1],original['PPR']['Annual'][1][0])
            self.assertEqual(saved['PPR']['Taxon SPPR'][1][1],original['PPR']['Taxon SPPR'][1][0])
            self.assertEqual(saved['PPR']['Annual'][1][0][-1],7.)
            self.assertEqual(overview(saved)['calculation_result_sha256'],result_hash(saved))
            with zipfile.ZipFile(path) as archive:
                _, _, _, data = table_rows(archive,'PPR','Annual')
                self.assertEqual(data[1][1]['1950'],'216219124.64603513')

    def test_publication_restores_desired_original_numbers_even_if_path_already_rounded(self):
        from tools.project_core.calculations.regional import result_hash
        from tools.project_core.workbooks.workbooks import overview
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'region.xlsx'
            desired = {'Overview': {'Settings': (['field','value'], [])},
                'PPR': {'Taxon SPPR': (['model_id','taxon','scope','method','sppr'],
                    [['A','fish','all','retained',216219124.64603513],
                     ['A','other','all','changed',1.2345678901234567]])}}
            write_book(path, desired)
            self.assertNotEqual(read_book(path)['PPR']['Taxon SPPR'][1][0][-1],desired['PPR']['Taxon SPPR'][1][0][-1])
            current = copy.deepcopy(desired)
            current['PPR']['Taxon SPPR'][1][1][-1] = 8.
            set_setting(current,'calculation_input_sha256',input_hash(current)); set_result_hash(current)
            saved = coverage._save_calculation(path,current,Path(temporary)/'model.json',{},desired)
            self.assertEqual(saved['PPR']['Taxon SPPR'][1][0][-1],216219124.64603513)
            self.assertEqual(saved['PPR']['Taxon SPPR'][1][1][-1],8.)
            self.assertEqual(overview(saved)['calculation_result_sha256'],result_hash(saved))

    def test_source_extension_preserves_original_cells_and_existing_diagnostics(self):
        import openpyxl
        import zipfile
        from xml.etree import ElementTree as etree
        from tools.project_core.validation.researcher_review import sheet_xml, S
        with tempfile.TemporaryDirectory() as temporary:
            model = Path(temporary)
            source = model/'sppr_source.xlsx'
            workbook = openpyxl.Workbook(); workbook.remove(workbook.active)
            original_value = 5.000000000000004
            for scope in ('all','inner','PP'):
                sheet = workbook.create_sheet('sppr_'+scope)
                sheet.append(['seq','group_name','GE'])
                sheet.append([1,'Fish',original_value])
                sheet['C2'].number_format = '0.000000000000000'
            workbook.create_sheet('model_health').append(['original diagnostic cells'])
            workbook.save(source); workbook.close()
            # A source produced by Excel can retain 17 significant digits;
            # creating the fixture with openpyxl alone already rounds them.
            raw_value = '1.2345678901234567'
            with zipfile.ZipFile(source) as archive:
                replacements = {}
                for scope in ('all','inner','PP'):
                    member, sheet = sheet_xml(archive, 'sppr_'+scope)
                    cell = sheet.find(f'.//{{{S}}}c[@r="C2"]')
                    cell.find(f'{{{S}}}v').text = raw_value
                    replacements[member] = etree.tostring(sheet,encoding='utf-8')
                parts = [(info,replacements.get(info.filename,archive.read(info.filename))) for info in archive.infolist()]
            with zipfile.ZipFile(source,'w') as archive:
                for info,data in parts: archive.writestr(info,data)
            original_value = float(raw_value)
            book = fixture()
            book['Selected model groups']['Group SPPR'][1][:] = [
                ['A','Fish',scope,method,value] for method,value in [('GE',1.234567890123457),('extra',7.)]
                for scope in ('all','inner','PP')]
            outcomes = {'GE': {'status':'reused','preserve_saved':True}, 'extra': {'status':'ok'}}
            coverage._update_source(source,book,outcomes)
            workbook = openpyxl.load_workbook(source)
            for scope in ('all','inner','PP'):
                sheet = workbook['sppr_'+scope]
                self.assertEqual(sheet['C2'].value,original_value)
                self.assertEqual(sheet['C2'].number_format,'0.000000000000000')
                self.assertEqual(sheet['D1'].value,'extra')
                self.assertEqual(sheet['D2'].value,7.)
            self.assertEqual(workbook['model_health']['A1'].value,'original diagnostic cells')
            workbook.close()
            with zipfile.ZipFile(source) as archive:
                for scope in ('all','inner','PP'):
                    _, sheet = sheet_xml(archive, 'sppr_'+scope)
                    self.assertEqual(sheet.find(f'.//{{{S}}}c[@r="C2"]/{{{S}}}v').text,raw_value)
            self.assertTrue(coverage._source_matches(model,book,('GE','extra')))

    def test_extension_preserves_exact_historical_outputs_and_catch_specific_restrictions(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model, path, original, _ = self.package(root)
            for row in rows(original, 'Selected model groups', 'Group SPPR'):
                if row[3] == 'new_GE': row[4] = -5.
            recalculate(original, path)
            for row in rows(original, 'PPR', 'Annual'):
                if row[2] == 'new_GE':
                    row[6] = 'FAIL: retained diagnostic for '+row[3]
            bound = ['A','all','MC_new_GE','landings','method','min_tC','ok',*([.4]*len(YEARS))]
            original['PPR']['Annual'][1].insert(0, bound)
            set_result_hash(original)
            extended = copy.deepcopy(original)
            coverage._apply_outcomes(extended, 'A', {'extra': {
                'status': 'ok', 'scopes': {s: {'Fish': 7.} for s in ('all','inner','PP')},
                'annual_status': {s: 'ok' for s in ('all','inner','PP')},
            }}, ['extra'])
            coverage._arithmetic(extended, path, model/'model.json', {}, original)
            for sheet, table in [('PPR','Annual'), ('PPR','Taxon SPPR'),
                                 ('PPR','Taxon PPR inspected year'), ('PPR–NPP','Ratios')]:
                old_header, old_rows = original[sheet][table]
                new_header, new_rows = extended[sheet][table]
                self.assertEqual(new_header, old_header)
                method_column = old_header.index('method')
                self.assertEqual([row for row in new_rows if row[method_column] in {'new_GE','MC_new_GE'}],
                                 [row for row in old_rows if row[method_column] in {'new_GE','MC_new_GE'}])
                self.assertTrue(any(row[method_column] == 'extra' for row in new_rows))
            extra = next(row for row in records(extended,'PPR','Annual') if row['method']=='extra' and
                         row['scope']=='all' and row['catch_basis']=='landings' and row['unidentified']=='method' and row['metric']=='ppr')
            self.assertEqual(extra[2019], 14.)

    def test_stale_input_or_result_hash_does_not_preserve_obsolete_annual_values(self):
        for stale in ('catch', 'result'):
            with self.subTest(stale=stale), tempfile.TemporaryDirectory() as temporary:
                model, path, original, _ = self.package(Path(temporary))
                if stale == 'catch': original['Catch']['Catch'][1][0][-1] = 30.
                else:
                    for row in rows(original,'PPR','Annual'):
                        if row[2]=='new_GE' and row[1]=='all' and row[3]=='landings' and row[4]=='method' and row[5]=='ppr':
                            row[-1] = 999.
                current = copy.deepcopy(original)
                coverage._arithmetic(current,path,model/'model.json',{},original)
                annual = next(row for row in records(current,'PPR','Annual') if row['method']=='new_GE' and
                              row['scope']=='all' and row['catch_basis']=='landings' and row['unidentified']=='method' and row['metric']=='ppr')
                self.assertEqual(annual[2019],150. if stale=='catch' else 10.)

    def test_retained_outcome_does_not_replace_coefficients_status_or_diagnostics(self):
        book = fixture()
        book['Diagnostics']['mc_diagnostics'] = (['method', 'n_valid'], [['MC_new_GE', 93]])
        original = copy.deepcopy(book)
        coverage._apply_outcomes(book, 'A', {'GE': {
            'status': 'reused', 'preserve_saved': True,
            'scopes': {'all': {'Fish': 999}}, 'annual_status': {'all': 'ok'},
        }}, ['GE'])
        self.assertEqual(book['Selected model groups'], original['Selected model groups'])
        self.assertEqual(book['PPR']['Annual'], original['PPR']['Annual'])
        self.assertEqual(book['Diagnostics']['model_health'], original['Diagnostics']['model_health'])
        self.assertEqual(book['Diagnostics']['mc_diagnostics'], original['Diagnostics']['mc_diagnostics'])

    def package(self, root):
        region = root/'regions/LME/LME_001'
        model = region/'papers/P/models/A'
        model.mkdir(parents=True)
        (root/'Project.xlsx').write_bytes(b'unchanged central metadata')
        source = model/'model.json'
        source.write_text('{"group":[{"group_seq":"1","group_name":"Fish"}]}')
        book = fixture()
        set_setting(book, 'selected_model_id', 'A')
        set_setting(book, 'results_model_sha256', sha(source))
        set_setting(book, 'model_path', source.relative_to(region).as_posix())
        book['Selected model groups']['Group SPPR'][1][:] = [
            ['A', 'Fish', scope, method, value]
            for method, value in [('new_GE', 5.), ('MC_new_GE', 5.123456789)]
            for scope in ('all', 'inner', 'PP')]
        book['PPR']['Annual'][1][:] = [
            ['A', scope, method, 'landings', 'method', 'ppr', 'ok', *([None]*70)]
            for method in ('new_GE', 'MC_new_GE') for scope in ('all', 'inner', 'PP')]
        book['Diagnostics']['mc_diagnostics'] = (
            ['method', 'n_samples', 'n_valid', 'acceptance_rate'], [['MC_new_GE', 100, 93, .93]])
        path = region/'LME_001.xlsx'
        recalculate(book, path)
        write_book(path, book)
        original = read_book(path)
        metadata = {'unit_id': 'LME_001', 'model_id': 'A',
                    'reviewed_model_sha256': sha(source),
                    'reviewed_calculation_input_sha256': input_hash(original)}
        return model, path, original, metadata

    def test_legacy_complete_results_are_adopted_without_any_engine_run(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model, path, original, metadata = self.package(root)
            def unexpected(*args, **kwargs):
                self.fail('Existing deterministic or Monte Carlo method was rerun')
            context = {'model_sha256': sha(model/'model.json'), 'flags': {}, 'code_identity': {}}
            with patch('tools.scientific_helpers.sppr.configuration.resolve_audited_configuration', return_value={}), \
                 patch.object(coverage, '_context', return_value=context):
                modified, _ = coverage._complete_model(root, metadata, ('new_GE', 'MC_new_GE'),
                                                       progress=lambda value: None, compute=unexpected)
                self.assertTrue(modified)
                saved = read_book(path)
                self.assertEqual(saved['Selected model groups']['Group SPPR'], original['Selected model groups']['Group SPPR'])
                self.assertEqual(saved['Diagnostics']['model_health'], original['Diagnostics']['model_health'])
                self.assertEqual(saved['Diagnostics']['mc_diagnostics'], original['Diagnostics']['mc_diagnostics'])
                before = path.read_bytes()
                modified, _ = coverage._complete_model(root, metadata, ('new_GE', 'MC_new_GE'),
                                                       progress=lambda value: None, compute=unexpected)
                self.assertFalse(modified)
                self.assertEqual(path.read_bytes(), before)

    def test_only_missing_method_executes_with_the_audited_collapse_flag(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model, path, original, metadata = self.package(root)
            calls = []
            flags = {'det_collapse_mode': 'auto'}
            def compute(model_path, actual_flags, methods, directory, progress):
                calls.append((methods, actual_flags))
                outcomes = {}
                for method in methods:
                    evidence = directory/(method+'.json')
                    evidence.write_text(json.dumps({'method': method}))
                    outcomes[method] = {'status': 'ok', 'scopes': {s: {'Fish': 7.} for s in ('all','inner','PP')},
                        'annual_status': {s: 'ok' for s in ('all','inner','PP')},
                        'evidence': {'path': evidence.name, 'sha256': sha(evidence)}}
                return {'methods': outcomes}
            context = {'model_sha256': sha(model/'model.json'), 'flags': flags, 'code_identity': {}}
            with patch('tools.scientific_helpers.sppr.configuration.resolve_audited_configuration', return_value=flags), \
                 patch.object(coverage, '_context', return_value=context):
                coverage._complete_model(root, metadata, ('new_GE','MC_new_GE','SPPR_1986'),
                                         progress=lambda value: None, compute=compute)
            self.assertEqual(calls, [(['SPPR_1986'], flags)])
            saved = read_book(path)
            original_rows = rows(original, 'Selected model groups', 'Group SPPR')
            self.assertEqual(rows(saved, 'Selected model groups', 'Group SPPR')[:len(original_rows)], original_rows)

    def test_changed_model_requires_explicit_scientific_refresh_without_overwriting(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            model, path, original, metadata = self.package(root)
            before = path.read_bytes()
            (model/'model.json').write_text('{"group":[{"group_seq":"1","group_name":"Fish","b":2}]}')
            def unexpected(*args, **kwargs):
                self.fail('Map method completion reran a changed reviewed model')
            context = {'model_sha256': sha(model/'model.json'), 'flags': {}, 'code_identity': {}}
            with patch('tools.scientific_helpers.sppr.configuration.resolve_audited_configuration', return_value={}), \
                 patch.object(coverage, '_context', return_value=context), \
                 self.assertRaisesRegex(ValueError, 'explicit scientific refresh'):
                coverage._complete_model(root, metadata, ('new_GE','MC_new_GE'),
                                         progress=lambda value: None, compute=unexpected)
            self.assertEqual(path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
