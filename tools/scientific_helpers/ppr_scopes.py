"""Read the exported SPPR scopes and diagnose_sppr results without recalculation."""
import math

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SCOPES = {'all': 'All basal sources, including imports',
          'inner': 'Internal sources: primary producers and detritus; excludes imports',
          'PP': 'Primary producers only; excludes detritus and imports'}


def finite(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else None


def read_scopes(path):
    result = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for scope in SCOPES:
        sheet = 'sppr_' + scope
        rows = list(wb[sheet].values) if sheet in wb.sheetnames else []
        if not rows:
            result[scope] = ([], {})
            continue
        header = list(rows[0])
        start = header.index('group_name') + 1
        methods = header[start:]
        result[scope] = (methods, {r[start - 1]: [finite(v) for v in r[start:]]
                                   for r in rows[1:] if r[start - 1]})
    wb.close()
    return result


def read_health(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb['model_health'].values) if 'model_health' in wb.sheetnames else []
    wb.close()
    result = {}
    for row in rows[1:]:
        d = dict(zip(rows[0], row))
        option = d.get('TE_option')
        if not option:
            continue
        result[option] = {
            'b': finite(d.get('divergence_b')),
            'rho_living': finite(d.get('divergence_rho_living')),
            'status': d.get('status') or 'unavailable',
            'b_converges': d.get('divergence_b_converges'),
            'living_converges': d.get('divergence_living_converges'),
            'balanced': d.get('model_input_is_model_balanced'),
            'config': {k[7:]: v for k, v in d.items() if k.startswith('config_')},
        }
    return result


def read_mc_diagnostics(path):
    """Read acceptance counts from the stored simulations, without rerunning them."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb['mc_diagnostics'].values) if 'mc_diagnostics' in wb.sheetnames else []
    finally:
        wb.close()
    result = {}
    for row in rows[1:]:
        entry = dict(zip(rows[0], row))
        method = entry.get('method')
        if isinstance(method, str) and method.startswith('MC_'):
            result[method] = {key: finite(entry.get(key)) for key in ('n_samples', 'n_accepted')}
    return result


def method_health_flags(methods, groups, health):
    """Only apply diagnose_sppr to configurations actually evaluated upstream.

    The exported TE diagnostic uses SPPR_new's default EE repair (True). Monte
    Carlo and other solvers have different configurations; do not inherit it.
    A negative result anywhere in the source model invalidates that method even
    when the negative compartment has no catch in this ecosystem.
    """
    matched = {'new_GE': 'GE', 'new_TE_EEfix': 'TE', 'new_WithEgestion': 'With Egestion'}
    flags = {}
    for i, method in enumerate(methods):
        if any(values[i] is not None and values[i] < 0 for values in groups.values()):
            flags[method] = 'DIVERGED - negative SPPR in a source-model group'
        h = health.get(matched.get(method), {})
        if h.get('status') == 'FAIL':
            flags[method] = 'FAILED diagnostic for this SPPR_new configuration'
    return flags


def format_sheet(ws, header=4):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'C{header + 1}'
    ws['A1'].font = Font(size=14, bold=True)
    ws.row_dimensions[2].height = 32
    ws['A2'].alignment = Alignment(wrap_text=False)
    for cell in ws[header]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E79')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[header].height = 32
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 46 if col == 1 else 32 if col == 2 else 20
    for row in ws.iter_rows(min_row=header + 1):
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                c.number_format = '#,##0.000000'


def add_recycling_sheet(wb, paths):
    ws = wb.create_sheet('Recycling')
    ws.append(['Recycling diagnostics from diagnose_sppr()'])
    ws.append(['b: detrital recycling; rho living: living food-web cycles. Each must be below 1 for convergence; these are necessary, not sufficient checks.'])
    ws.append([])
    ws.append(['model', 'TE_option', 'b', 'rho living', 'diagnostic status', 'b converges',
               'living converges', 'input balanced', 'configuration'])
    import json
    for path in paths:
        for te, d in read_health(path).items():
            ws.append([path.stem, te, d['b'], d['rho_living'], d['status'], d['b_converges'],
                       d['living_converges'], d['balanced'], json.dumps(d['config'], sort_keys=True)])
    format_sheet(ws)
    ws.column_dimensions['I'].width = 90


def add_group_scope_sheets(wb, paths):
    models = [(p.stem, read_scopes(p)) for p in paths]
    for scope, description in SCOPES.items():
        ws = wb.create_sheet('sppr_' + scope)
        ws.append([f'Group SPPR: {scope}'])
        ws.append([description + '. Models are separate. Blank means unavailable, not zero.'])
        ws.append([])
        methods = list(dict.fromkeys(m for _, s in models for m in s[scope][0]))
        ws.append(['model', 'group_name'] + methods)
        for model, scopes in models:
            names, groups = scopes[scope]
            for group, values in groups.items():
                row = dict(zip(names, values))
                ws.append([model, group] + [row.get(m) for m in methods])
        if not models:
            ws.append(['No extracted model is available for this ecosystem.'])
        format_sheet(ws)
    add_recycling_sheet(wb, paths)
