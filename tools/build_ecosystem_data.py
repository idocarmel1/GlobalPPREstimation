"""Build the per-ecosystem data spine and one lean workbook per ecosystem.

This is the integration layer. Everything upstream produces data keyed by different
things -- catch archives by unit, articles by region directory, Ecopath models by a
filename prefix, NPP by layer plus region id. This script joins all of it on `unit_id`
(`LME_003`, `EEZ_711`, `HS_018`) and writes, for each ecosystem:

    data/<unit_id>/metadata.json     identity, geography, coverage, pointers to bulk sources
    data/<unit_id>/npp.json          net primary production by satellite model, when known
    data/<unit_id>/<unit_id>.xlsx    the ecosystem workbook a human actually opens

This workbook holds catch, trophic-chain SPPR and PPR, annual NPP, and final taxon
mappings identified separately by model. Model-specific PPR lives in
`data/<unit_id>/models/<model_stem>.xlsx`, one workbook per model, built by
`tools/build_model_workbook.py`. Keeping them apart is not tidiness: an ecosystem with two
published models has two different answers for PPR, and a single sheet holding both invites
a reader to average them.

Bulk inputs are referenced by relative path rather than copied. The catch archives and
the article archive together run to gigabytes; duplicating them per ecosystem would add
nothing and cost a great deal.

The workbook keeps model identities distinct. Sheets that have no
data for an ecosystem say so in one line rather than appearing empty.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from ppr_scopes import add_group_scope_sheets
from npp_data import load_npp as load_annual_npp, npp_for_year, source_paths
from build_model_workbook import (sheet_npp as annual_npp_sheet, build_taxon_sppr,
                                  final_mapping_rows, sheet_final_mappings, mio, CATCH_DP,
                                  save_workbook_atomic)
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data"

CATCH_DIR = ROOT / "SeaAroundUsExtraction" / "data" / "catch_by_taxon_year"
REGION_TABLES = [
    ROOT / "SeaAroundUsExtraction" / "global_output" / "tables" / "regions",
    ROOT / "SeaAroundUsExtraction" / "eez_output" / "tables" / "regions",
]
ARCHIVE = ROOT / "PPRAtlas" / "archive" / "regions"
ATLAS_REGIONS = ROOT / "PPRAtlas" / "data" / "regions.csv"
SPPR_DIR = ROOT / "PPREstimation" / "output" / "top10"
MODEL_JSON_DIR = ROOT / "PPREstimation" / "real_models" / "global_cover_jsons"

TRANSFER_EFFICIENCY = 0.1

HEADER = Font(bold=True)


def unit_from_model_filename(name: str) -> str | None:
    """`077HS_1_Eastern...` -> HS_077;  `36_2_South_China_Sea...` -> LME_036."""
    m = re.match(r"^(\d+)HS_", name)
    if m:
        return f"HS_{int(m.group(1)):03d}"
    m = re.match(r"^(\d+)_", name)
    if m:
        return f"LME_{int(m.group(1)):03d}"
    return None


def load_atlas_regions() -> dict[str, dict]:
    if not ATLAS_REGIONS.exists():
        return {}
    with ATLAS_REGIONS.open(encoding="utf-8-sig") as fh:
        return {r["unit_id"]: r for r in csv.DictReader(fh)}


def load_npp() -> dict[str, dict]:
    return load_annual_npp(ROOT)


def npp_coverage(npp):
    years = sorted(int(y) for y in ((npp or {}).get('annual') or {'2019': npp})
                   if (npp_for_year(npp, y) or {}).get('ens_median_tC_yr') not in (None, ''))
    return {'has_npp': bool(years), 'npp_years_available': len(years),
            'npp_from': years[0] if years else None, 'npp_to': years[-1] if years else None}


def load_final_mappings(unit):
    """Resolve only this unit's available mapping sources, using model arithmetic."""
    books = [book for book in mio.model_workbooks(ROOT, unit)
             if (mio.mapping_dir(ROOT, unit) / f'{book.stem}.csv').exists()]
    if not books:
        return []
    taxa, _years = mio.read_catch(ROOT, unit)
    totals = {t: sum(round(v, CATCH_DP) for v in e['by_year'].values())
              for t, e in taxa.items()}
    out = []
    for book in books:
        mapping = mio.mapping_dir(ROOT, unit) / f'{book.stem}.csv'
        if not mapping.exists():
            continue
        rows = mio.read_mapping_csv(mapping)
        groups = {g['group_name']: g for g in mio.read_groups(book)}
        methods, sppr = mio.read_methods(book)
        for r in rows:
            for name in mio.parse_groups_cell(r.get('group')):
                if name.lower() != 'unresolved' and name not in groups:
                    raise ValueError(f'{book.stem}: mapped group {name!r} absent from model')
        seen = {r['taxon'] for r in rows}
        rows += [{'taxon': t, 'group': 'Unresolved', 'confidence': 'unresolved',
                  'evidence': 'none', 'explanation': 'No mapping row recorded for this catch taxon.'}
                 for t in taxa if t not in seen]
        resolved = build_taxon_sppr(rows, methods, sppr, groups, totals)
        out.extend(final_mapping_rows(unit, book.stem, taxa, resolved, totals))
    return out


def load_trophic_levels(unit: str) -> dict[str, float]:
    """Taxon -> trophic level, from whichever region table holds this unit."""
    for base in REGION_TABLES:
        p = base / unit / "species.csv"
        if not p.exists():
            continue
        tl = {}
        with p.open(encoding="utf-8-sig") as fh:
            for r in csv.DictReader(fh):
                v = (r.get("tl") or "").strip()
                if v:
                    try:
                        tl[r["taxon"]] = float(v)
                    except ValueError:
                        pass
        return tl
    return {}


def load_catch(unit: str):
    """Return (rows, years). Each row: taxon, common, fg, cg, {year: tonnes}."""
    p = CATCH_DIR / f"{unit}.csv.gz"
    if not p.exists():
        return [], []
    by_taxon: dict[str, dict] = {}
    years: set[int] = set()
    with gzip.open(p, "rt", encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            taxon = r["taxon"]
            try:
                year = int(r["year"])
                tonnes = float(r["catch_tonnes"] or 0.0)
            except ValueError:
                continue
            years.add(year)
            e = by_taxon.setdefault(
                taxon,
                {
                    "taxon": taxon,
                    "common_name": r.get("common_name", ""),
                    "functional_group": r.get("functional_group", ""),
                    "commercial_group": r.get("commercial_group", ""),
                    "by_year": defaultdict(float),
                },
            )
            e["by_year"][year] += tonnes
    return list(by_taxon.values()), sorted(years)


def load_sppr_methods(unit: str) -> list[dict]:
    """Per-group SPPR by method, from every Ecopath workbook belonging to this unit."""
    if not SPPR_DIR.exists():
        return []
    out = []
    for p in sorted(SPPR_DIR.glob("*.xlsx")):
        if unit_from_model_filename(p.name) != unit:
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception as exc:  # a corrupt workbook must not sink the whole run
            out.append({"model": p.stem, "error": str(exc), "header": [], "rows": []})
            continue
        if "sppr_all" not in wb.sheetnames:
            wb.close()
            continue
        ws = wb["sppr_all"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            continue
        out.append({"model": p.stem, "header": list(rows[0]), "rows": [list(r) for r in rows[1:]]})
    return out


def sheet_catch(wb, rows, years):
    ws = wb.create_sheet("Catch")
    if not rows:
        ws.append(["No catch data for this ecosystem."])
        return
    # Years are written as numbers, not text: the model workbooks match against them
    # with MATCH(), and a text header silently fails that lookup.
    head = ["taxon", "common_name", "functional_group", "commercial_group"] + list(years)
    ws.append(head)
    for r in sorted(rows, key=lambda x: -sum(x["by_year"].values())):
        ws.append(
            [r["taxon"], r["common_name"], r["functional_group"], r["commercial_group"]]
            + [round(r["by_year"].get(y, 0.0), 3) for y in years]
        )
    finish(ws, len(head), freeze="E2")


def sheet_sppr(wb, unit, models, tl):
    """The trophic-chain SPPR per taxon. Model SPPR lives in the model workbooks."""
    ws = wb.create_sheet("SPPR")
    ws.append([f"Specific PPR for {unit}"])
    ws["A1"].font = HEADER
    ws.append([])
    ws.append(["Simple method, applied per taxon: SPPR = (1/TE)^(TL-1), TE = 0.1"])
    ws.append(["taxon", "trophic_level", "sppr_simple"])
    for c in ws[4]:
        c.font = HEADER
    for taxon, level in sorted(tl.items(), key=lambda kv: -kv[1]):
        ws.append([taxon, level, round((1.0 / TRANSFER_EFFICIENCY) ** (level - 1.0), 6)])
    ws.append([])
    if models:
        ws.append([f"{len(models)} Ecopath model(s) also cover this ecosystem. Their "
                   "per-group SPPR under all 20 methods, and the PPR that follows from it,"])
        ws.append(["are in data/" + unit + "/models/ — one workbook per model, because two "
                   "models give two different answers and they are not averageable."])
        for m in models:
            ws.append([f"    {m['model']}.xlsx"])
    else:
        ws.append(["No Ecopath model has been extracted for this ecosystem yet,"])
        ws.append(["so the network-based SPPR methods are not available here."])
    finish(ws, 3)


def sheet_ppr(wb, rows, years, tl):
    """PPR per taxon per year: catch x SPPR, using the simple per-taxon method."""
    ws = wb.create_sheet("PPR")
    if not rows or not tl:
        ws.append(["PPR needs both catch and trophic levels; one is missing here."])
        return
    ws.append(["PPR = catch tonnes x (1/TE)^(TL-1), TE = 0.1, per taxon per year."])
    ws.append(["Taxa with no trophic level are omitted and their catch contributes no PPR."])
    ws.append([])
    head = ["taxon", "trophic_level"] + list(years)
    ws.append(head)
    for c in ws[4]:
        c.font = HEADER
    matched = [r for r in rows if r["taxon"] in tl]
    for r in sorted(matched, key=lambda x: -sum(x["by_year"].values())):
        sppr = (1.0 / TRANSFER_EFFICIENCY) ** (tl[r["taxon"]] - 1.0)
        ws.append([r["taxon"], tl[r["taxon"]]] + [round(r["by_year"].get(y, 0.0) * sppr, 3) for y in years])
    finish(ws, len(head), freeze="C5")


def sheet_npp(wb, npp, years=()):
    return annual_npp_sheet(wb, npp, years)


def sheet_summary(wb, unit, meta, rows, years, tl, npp, models):
    ws = wb.create_sheet("Summary", 0)
    ws.append([f"{unit} — {meta.get('region_name') or ''}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in [
        ("region type", meta.get("region_type") or ""),
        ("latitude", meta.get("marker_lat") or ""),
        ("longitude", meta.get("marker_lon") or ""),
        ("archived articles", meta.get("article_count", 0)),
        ("Ecopath models extracted", len(models)),
        ("model workbooks", f"data/{unit}/models/" if models else "none yet"),
        ("catch years", f"{years[0]}-{years[-1]}" if years else "none"),
        ("taxa in catch", len(rows)),
        ("taxa with a trophic level", sum(1 for r in rows if r["taxon"] in tl)),
        ("model source context", meta.get('model_context', {}).get('note', 'No pilot model selected')),
    ]:
        ws.append([k, v])
    ws.append([])

    ws.append(["year", "catch_tonnes", "ppr_tonnes", "ppr_over_npp_percent"])
    for c in ws[ws.max_row]:
        c.font = HEADER
    for y in years:
        annual = npp_for_year(npp, y) or {}
        npp_median = float(annual['ens_median_tC_yr']) if annual.get('ens_median_tC_yr') not in (None, '') else None
        catch = sum(r["by_year"].get(y, 0.0) for r in rows)
        ppr = sum(
            r["by_year"].get(y, 0.0) * (1.0 / TRANSFER_EFFICIENCY) ** (tl[r["taxon"]] - 1.0)
            for r in rows
            if r["taxon"] in tl
        )
        ratio = round(100.0 * ppr / 9.0 / npp_median, 4) if npp_median else None
        ws.append([y, round(catch, 3), round(ppr, 3), ratio])
    ws.append(['PPR/NPP converts wet-weight PP to carbon at 9:1 and uses the matching year\'s NPP ensemble median.'])
    ws.append(['PPR/NPP is blank for years without NPP; see NPP for availability and provenance.'])
    finish(ws, 4)
    for column, width in [('A', 30), ('C', 22), ('D', 30)]:
        ws.column_dimensions[column].width = width


def finish(ws, ncols, freeze=None):
    widths = [26, 22] + [14] * max(0, ncols - 2)
    for i, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze or "A2"
    ws.sheet_view.showGridLines = False
    for c in ws[1]:
        c.alignment = Alignment(vertical="center")


def build_unit(unit, atlas, npp_all, force=False):
    d = OUT / unit
    d.mkdir(parents=True, exist_ok=True)
    book = d / f"{unit}.xlsx"
    rows, years = load_catch(unit)
    tl = load_trophic_levels(unit)
    npp = npp_all.get(unit)
    models = load_sppr_methods(unit)
    meta_src = atlas.get(unit, {})

    article_dir = ARCHIVE / unit
    articles = sorted(p.name for p in article_dir.iterdir() if p.is_dir()) if article_dir.exists() else []

    # An extracted model is not the same thing as a usable answer. A model needs a
    # taxon-to-group mapping before any PPR can come out of it, and several extracted
    # models are marked unfit and will never get one. Counting both keeps the gap visible
    # rather than letting "10 ecosystems have models" stand in for "10 have results".
    mapped = sorted(q.stem for q in (d / "mapping").glob("*.csv")
                    if not q.name.endswith((".groups.csv", ".members.csv", ".resolved.csv", ".taxonomy.csv"))
                    ) if (d / "mapping").is_dir() else []
    built = sorted(q.stem for q in (d / "models").glob("*.xlsx")
                   if not q.name.startswith("~$")) if (d / "models").is_dir() else []

    meta = {
        "unit_id": unit,
        "region_name": meta_src.get("region_name") or "",
        "region_type": meta_src.get("region_type") or "",
        "marker_lat": meta_src.get("marker_lat") or None,
        "marker_lon": meta_src.get("marker_lon") or None,
        "ppr_rank_1995_method": meta_src.get("ppr_rank") or None,
        "article_count": len(articles),
        "articles": articles,
        "coverage": {
            "catch_years": [years[0], years[-1]] if years else None,
            "taxa": len(rows),
            "taxa_with_trophic_level": sum(1 for r in rows if r["taxon"] in tl),
            **npp_coverage(npp),
            "ecopath_models": len(models),
            "models_mapped": len(mapped),
            "model_workbooks": len(built),
        },
        "sources": {
            "npp": [f'../../{p}' for p in source_paths(ROOT)],
            "catch": f"../../SeaAroundUsExtraction/data/catch_by_taxon_year/{unit}.csv.gz",
            "articles": f"../../PPRAtlas/archive/regions/{unit}" if articles else None,
            "sppr_workbooks": [f"../../PPREstimation/output/top10/{m['model']}.xlsx" for m in models],
            "model_json": [
                f"../../PPREstimation/real_models/global_cover_jsons/{p.name}"
                for p in sorted(MODEL_JSON_DIR.glob("*.json"))
                if unit_from_model_filename(p.name) == unit
            ]
            if MODEL_JSON_DIR.exists()
            else [],
        },
        "note": "Bulk inputs are referenced, not copied. Paths are relative to this file.",
    }
    selection_path = ROOT / 'data/atlas_selection.json'
    if selection_path.exists():
        meta['model_context'] = json.loads(selection_path.read_text(encoding='utf-8')).get('units', {}).get(unit, {})
    (d / "metadata.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    if npp:
        (d / "npp.json").write_text(json.dumps(npp, indent=2, ensure_ascii=False), encoding="utf-8")
    elif (d / 'npp.json').exists():
        (d / 'npp.json').unlink()

    if book.exists() and not force:
        return meta

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_catch(wb, rows, years)
    sheet_final_mappings(wb, load_final_mappings(unit))
    sheet_sppr(wb, unit, models, tl)
    add_group_scope_sheets(wb, [p for p in sorted(SPPR_DIR.glob('*.xlsx'))
                               if unit_from_model_filename(p.name) == unit])
    sheet_ppr(wb, rows, years, tl)
    sheet_npp(wb, npp, years)
    sheet_summary(wb, unit, meta, rows, years, tl, npp, models)
    try:
        save_workbook_atomic(wb, book)
    finally:
        wb.close()
    return meta


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--units", nargs="*", help="limit to these unit ids")
    ap.add_argument("--limit", type=int, default=0, help="stop after N units")
    ap.add_argument("--force", action="store_true", help="rewrite workbooks that already exist")
    args = ap.parse_args()

    atlas = load_atlas_regions()
    npp_all = load_npp()
    units = sorted(p.name.replace(".csv.gz", "") for p in CATCH_DIR.glob("*.csv.gz"))
    if args.units:
        units = [u for u in units if u in set(args.units)]
    if args.limit:
        units = units[: args.limit]
    if not units:
        print("no units to build", file=sys.stderr)
        return 1

    OUT.mkdir(exist_ok=True)
    index = []
    for i, unit in enumerate(units, 1):
        meta = build_unit(unit, atlas, npp_all, force=args.force)
        c = meta["coverage"]
        index.append(
            {
                "unit_id": unit,
                "region_name": meta["region_name"],
                "region_type": meta["region_type"],
                "catch_from": (c["catch_years"] or ["", ""])[0],
                "catch_to": (c["catch_years"] or ["", ""])[1],
                "taxa": c["taxa"],
                "taxa_with_tl": c["taxa_with_trophic_level"],
                "articles": meta["article_count"],
                "ecopath_models": c["ecopath_models"],
                "models_mapped": c["models_mapped"],
                "model_workbooks": c["model_workbooks"],
                "has_npp": int(c["has_npp"]),
                "npp_years_available": c.get('npp_years_available', int(c['has_npp'])),
                "npp_from": c.get('npp_from', 2019 if c['has_npp'] else None),
                "npp_to": c.get('npp_to', 2019 if c['has_npp'] else None),
            }
        )
        if i % 25 == 0 or i == len(units):
            print(f"  [{i}/{len(units)}] {unit}")

    index_path = OUT / 'INDEX.csv'
    if (args.units or args.limit) and index_path.exists():
        with index_path.open(encoding='utf-8-sig', newline='') as fh:
            previous = {r['unit_id']: r for r in csv.DictReader(fh)}
        previous.update({r['unit_id']: r for r in index})
        index = [previous[u] for u in sorted(previous)]
        for r in index:
            for key in ('taxa', 'taxa_with_tl', 'articles', 'ecopath_models', 'models_mapped', 'model_workbooks', 'has_npp'):
                r[key] = int(r[key])
            # Older index rows describe legacy 2019 coverage until those units rebuild.
            r['npp_years_available'] = int(r.get('npp_years_available') or (1 if r['has_npp'] else 0))
            for key in ('npp_from', 'npp_to'):
                value = r.get(key)
                r[key] = int(value) if value not in (None, '') else (2019 if r['has_npp'] else None)
    with index_path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(index[0].keys()))
        w.writeheader()
        w.writerows(index)

    tot = len(index)
    print(f"\nbuilt {tot} ecosystems")
    print(f"  with an Ecopath model : {sum(1 for r in index if r['ecopath_models'])}")
    print(f"  with a mapping        : {sum(1 for r in index if r['models_mapped'])}")
    print(f"  with a model workbook : {sum(1 for r in index if r['model_workbooks'])}"
          f"  ({sum(r['model_workbooks'] for r in index)} workbooks)")
    print(f"  with NPP              : {sum(1 for r in index if r['has_npp'])}")
    print(f"  with articles         : {sum(1 for r in index if r['articles'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
