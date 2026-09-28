from pathlib import Path
import json, hashlib
import fitz
from docx import Document
ROOT=Path(__file__).resolve().parents[4]
OUT=Path(__file__).resolve().parent
SRC=ROOT/'regions/LME_024/papers/LME024-Hernvann-2020'
inventory=[]
for f in SRC.iterdir():
 if f.is_file(): inventory.append({'file':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(OUT/'source_inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf8')
d=Document(next(SRC.glob('*.docx')))
tables=[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
(OUT/'docx_tables.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2),encoding='utf8')
(OUT/'docx_paragraphs.txt').write_text('\n'.join(f'P{i+1}: {p.text}' for i,p in enumerate(d.paragraphs)),encoding='utf8')
for i,t in enumerate(tables):
 (OUT/f'table_{i+1:02}.txt').write_text('\n'.join(f'R{j+1}: '+' | '.join(row) for j,row in enumerate(t)),encoding='utf8')
 print('TABLE',i+1,'rows',len(t),'cols',len(t[0]),t[:2])
pdf=fitz.open(next(SRC.glob('*.pdf')))
(OUT/'article_text.txt').write_text('\n'.join(f'\nPAGE {i+1}\n'+p.get_text() for i,p in enumerate(pdf)),encoding='utf8')
from openpyxl import load_workbook
w=load_workbook(ROOT/'regions/LME_024/LME_024.xlsx',read_only=True,data_only=True)
print('OVERVIEW',list(w['Overview'].values))
