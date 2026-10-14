"""Fill the retained validation template for the two experimental candidates."""
from pathlib import Path
from copy import deepcopy
import json,csv,hashlib,zipfile,collections,sys,os
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT
from lxml import etree
ROOT=Path(__file__).resolve().parents[6];WORK=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from tools.project_core.validation.validation_percentage_format import format_percent,verify_report
PAPER=ROOT/'regions/EEZ/EEZ_939/papers/Corrales-2010';TEMPLATE=ROOT/'tools/templates/validation.docx';RUN=WORK.name
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def stable_xml(e):return etree.tostring(e,method='c14n',exclusive=True)
def hyperlink(p,label,target):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.extend([color,u]);run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);h.append(run);p._p.append(h)
def fill(cell,paragraphs):
 for p in list(cell.paragraphs)[1:]:cell._tc.remove(p._p)
 p=cell.paragraphs[0];p.clear()
 for i,value in enumerate(paragraphs):
  if i:p=cell.add_paragraph()
  if isinstance(value,str):p.add_run(value)
  else:
   for part in value:
    if isinstance(part,tuple):hyperlink(p,*part)
    else:p.add_run(part)
  p.paragraph_format.space_after=Pt(3)
 return cell
def replace(p,value):p.clear();p.add_run(value)
def table_rows(table,rows):
 exemplar=deepcopy(table.rows[-1]._tr)
 for r in list(table.rows)[1:]:table._tbl.remove(r._tr)
 for row in rows:
  tr=deepcopy(exemplar);table._tbl.append(tr)
  for c,v in zip(table.rows[-1].cells,row):fill(c,[str(v)])
 # A header repeats on each page and template body rows keep their padding.
 header=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');header.append(repeat)
def before(anchor,element):anchor._p.addprevious(element)
def para_before(doc,anchor,text,style='Normal'):
 p=doc.add_paragraph(text,style);before(anchor,p._p);return p
def build(data):
 model=PAPER/'models'/data['model_id'];v=model/'model_validation';ev=v/'evidence';wk=v/'work'/RUN
 for d in ['inputs','outputs','code','qa']:(wk/d).mkdir(parents=True,exist_ok=True)
 document=Document(TEMPLATE);pars=list(document.paragraphs);tables=list(document.tables)
 with zipfile.ZipFile(TEMPLATE) as z:parts={n:z.read(n) for n in z.namelist()}
 sec=document.sections[0]
 contract=f'''# Validation template contract
Reference: {TEMPLATE.as_posix()}\nSHA-256: {sha(TEMPLATE)}\nThree rendered reference pages inspected; one A4 portrait section.
Page geometry: width {sec.page_width.cm:.6f} cm, height {sec.page_height.cm:.6f} cm; margins left {sec.left_margin.cm:.6f}, right {sec.right_margin.cm:.6f}, top {sec.top_margin.cm:.6f}, bottom {sec.bottom_margin.cm:.6f} cm. Header/footer distances and all section XML preserved.
Typography: retained Normal Calibri 11, Title 23, Heading 1 17, Heading 2 12. Source styles, numbering, theme, custom XML, settings and opaque parts are preserve-only. Title and automatic slots are editable; manual SPPR calculation, Open issues and next action, and Review and reproducibility cell XML are unchanged. No fields, content controls, drawings or hyperlinks exist in the source template.
Flow: flowing main Field/Entry table; additional user-requested mass-balance section using cloned source table styling; separate coverage/confidence and two component-rule tables; Very low decisions; geographic figures. All automatic placeholders are replaced or removed. Exact-model identity and candidate draft state govern the content. Template margins, table grids, heading styles and page patterns remain authoritative.
Slots: word/document.xml paragraphs 0–24 and main table labels; automatic table body rows may expand. Existing manual cells are preserve-only. Media and hyperlinks may be added; source source/relationship parts stay unchanged except document.xml, its relationships and Content_Types for added media. All other ZIP parts are restored byte-for-byte after object-model editing. No manual cell is signed or populated.
''' 
 (wk/'inputs/artifact.md').write_text(contract,encoding='utf-8');dump(wk/'inputs/template_parts.json',{n:hashlib.sha256(b).hexdigest() for n,b in parts.items()})
 period=data['period'].replace('_','–');earlier=data['period']=='1990_1994'
 replace(pars[0],f'Israel Mediterranean model validation {period.replace("–"," to ")}')
 replace(pars[1],'Candidate review • Pending-alignment draft. The reconstructed model still fails strict mass balance.')
 main=tables[0];fields={r.cells[0].text:r.cells[1] for r in main.rows[1:]}
 main._tbl.addprevious(pars[1]._p)
 manual={label:deepcopy(fields[label]._tc) for label in ['SPPR calculation','Open issues and next action','Review and reproducibility']}
 fill(fields['Region'],[['Israel Mediterranean • EEZ_939 • ',('Regional workbook','../../../../../EEZ_939.xlsx')]])
 fill(fields['Catch source'],['Sea Around Us regional catch series, 1950–2019; reference year 2019, landings. Retained download/version date is unavailable.',[('Original catch bundle','../../../../../raw/catch/EEZ_939-catch.zip'),' | ',('Catch and classic PPR workbook','../../../../../EEZ_939.xlsx')]])
 fill(fields['Selected article'],['Corrales X, Ofir E, Coll M, Goren M, Edelist D, Heymans JJ and Gal G (2017). Modeling the role and impact of alien species and fisheries on the Israeli marine continental shelf ecosystem. Journal of Marine Systems 170:88–102.',[('DOI','https://doi.org/10.1016/j.jmarsys.2017.02.004'),' | ',('Article PDF','../../../sources/1-s2.0-S0924796316302329-main.pdf'),' | ',('Author supplement','../../../sources/supplements/1-s2.0-S0924796316302329-mmc1.docx')]])
 fill(fields['Other known articles'],['Corrales et al. (2017), Hindcasting the dynamics of an Eastern Mediterranean marine ecosystem under the impacts of multiple stressors. Related dynamic study; no additional numerical source adopted.',[('Related study DOI','https://doi.org/10.3354/meps12271')]])
 fill(fields['Selected model'],[data['model_id'],f'Reviewed candidate for {period}; {data["source_group_count"]} source groups (+1 synthetic diet import); shelf depths 0–200 m, area 3,725 km². Regional selection is currently blank.',[('Canonical model JSON','../model.json'),' | ',('Saved runtime and diagnostics','../results/diagnostics'),' | ',('Model folder','..')]])
 other='2008–2012 (41 groups)' if earlier else '1990–1994 (39 groups)'
 fill(fields['Other models from the same article'],[f'Published counterpart: {other}, same shelf domain. Both periods retain source-only, assumed-fate, gray-only fit and gray-fit/F-removal variants. These reconstructions are separate from the author’s models.',[('All retained models','../../')]])
 fill(fields['Selection rationale'],['Article: researcher supplied this paper for extraction. Model: both published periods assessed; no regional selection rationale recorded (?).'])
 fill(fields['Model extraction'],['Ordinary blank diet cells are assumed zero on runtime copies. Original gray cells (<0.01) remain −9999 in source models; this separate variant fits them within their published bound. Printed diet values are fixed before runtime normalization.', 'Missing total removals are reconstructed from published F × biomass where consistent with known components. GS defaults to 0.2; living accumulation defaults to zero. Detritus accumulation is calculated. Fate is assumed, and explicit fleet-discard returns are omitted by the calculator.',[('Departures and source evidence','../model_notes.md'),' | ',('Gray-cell fit','../extracted_tables/evidence/gray_diet_fit/fit_summary.json')]])
 for method,label in [('GE','GE diagnostics'),('TE','TE diagnostics')]:
  d=json.loads((model/f'results/diagnostics/{method}_full_return.json').read_text(encoding='utf-8'))['diagnostics'];div=d['divergence'];detids=tuple(div['sppr_det']);btext=f'; b = {div["b"]:.3g}' if method=='GE' else ''
  fill(fields[label],[f'{d["status"]}: input production residual up to {format_percent(100*d["model_input"]["p_max_rel_residual"])}; strict mass balance is false.','No negative basal-source contributions in any retained group/source pairing.',f'ρ living = {div["rho_living"]:.3g}{btext}; Detritus ({detids[0]}) SPPR = {div["sppr_det"][detids[0]]:.3g}; Discards ({detids[1]}) SPPR = {div["sppr_det"][detids[1]]:.3g}.',f'PP budget gap = {format_percent(100*d["balance"]["rel_gap"])}; below the grading threshold, but strict budget closure is false.'])
 fill(fields['Geographic fit'],['Approximate containment scenario: A (region covered by study) ≈ 13–14%; B (study covered by region) ≈ 90–100%.'])
 fill(fields['Temporal fit'],[f'Modeled period {period}; assessed catch is 2019, outside that period. Fixed composition and coefficients would assume no later food-web/fishery change. Earlier model groups 24 and 35 are absent; no corresponding catches are invented.'])
 fill(fields['Other'],['With Egestion also returned '+('WARN' if earlier else 'FAIL')+'. Finite SPPR and convergent recycling do not establish mass balance. Automatic detritus collapse and removal of Discards were deferred because the required mass-balanced prerequisite remains unmet.'])
 # Additional scientific section requested explicitly by the researcher.
 anchor=pars[2];para_before(document,anchor,'Mass balance findings','Heading 1')
 fit=json.loads((model/'extracted_tables/evidence/gray_diet_fit/fit_summary.json').read_text(encoding='utf-8'))['results']['derived_F_removals']
 para_before(document,anchor,'Filling only gray cells cannot balance the rounded translated model. Even after the authorized removal reconstruction, the smallest possible maximum production residual is '+format_percent(100*fit['minimax_relative_residual_lower_bound'])+'. Ordinary blanks stay zero, printed fractions remain fixed, and each fitted gray fraction is below 0.01. These are constraints on this reconstruction, not a finding that the original native Ecopath model was unbalanced.')
 budgets=list(csv.DictReader((model/'results/diagnostics/mass_balance_by_group.csv').open(encoding='utf-8')))
 raw=json.loads((model/'model.json').read_text(encoding='utf-8'));published={int(g['group_seq']):int(g['source_status']['published_group_id']) for g in raw['group']}
 worst=sorted((r for r in budgets if r['trophic_info'] in ['Regular','PP'] and abs(float(r['production_relative_residual']))>1e-8),key=lambda r:-abs(float(r['production_relative_residual'])))[:8]
 table=document.add_table(rows=1,cols=5);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 # Clone the template's table properties and header/cell properties.
 table._tbl.remove(table._tbl.tblPr);table._tbl.insert(0,deepcopy(tables[1]._tbl.tblPr))
 widths=[7.0,2.5,2.5,2.7,2.0];headers=['Group and published ID','Production','Catch','Predation','Residual']
 headerstyle=tables[1].rows[0].cells[0]._tc.get_or_add_tcPr()
 for i,(cell,title) in enumerate(zip(table.rows[0].cells,headers)):
  cell.width=Cm(widths[i]);cell._tc.remove(cell._tc.tcPr);cell._tc.insert(0,deepcopy(headerstyle));cell.width=Cm(widths[i]);fill(cell,[title]);
  for r in cell.paragraphs[0].runs:r.bold=True
 for r in worst:
  cells=table.add_row().cells;values=[f'{r["group_name"]} ({published[int(r["group_seq"])]})']+[f'{float(r[k]):.5g}' for k in ['production','catch','predation']]+[format_percent(100*float(r['production_relative_residual']))]
  for i,(cell,value) in enumerate(zip(cells,values)):
   cell._tc.remove(cell._tc.tcPr);cell._tc.insert(0,deepcopy(tables[1].rows[1].cells[0]._tc.get_or_add_tcPr()));cell.width=Cm(widths[i]);fill(cell,[value])
 # Grid widths must agree with cells for stable Word rendering.
 for grid,width in zip(table._tbl.tblGrid.gridCol_lst,widths):grid.set(qn('w:w'),str(round(width/2.54*1440)))
 h=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(h)
 before(anchor,table._tbl)
 p=para_before(document,anchor,'Flows are t km⁻² yr⁻¹. Residual = (catch + predation + accumulation + net migration + other mortality − production) / production. Negative means an unexplained production sink deficit; positive means excess demand. Published IDs are used in the table. ')
 hyperlink(p,'All groups and both balance equations','../results/diagnostics/mass_balance_by_group.csv')
 if earlier:
  para_before(document,anchor,'Demersal sharks require negative predation: catch 0.010 exceeds EE × production 0.008960 by 0.001040. Gray-cell filling cannot remove that contradiction. Mackerel needs predation 0.074430, below the smallest attainable 0.078645 with printed diets fixed. Rays/skates have no gray incoming diet cells; their required predation 0.004534 exceeds the attainable 0.002986. Large native demersal fishes also have no gray incoming diet cells, leaving a deficit that normalization can only worsen.')
 else:
  para_before(document,anchor,'Rays/skates have no gray incoming diet cells. Required predation is 0.011830, while the largest attainable value is 0.004958. Large pelagic fishes have known landings 0.060 exceeding EE × production 0.053460; their F × B estimate 0.051600 is smaller than that known component and was rejected. Alien lizardfish needs predation 0.002922, below its minimum attainable 0.006506. These restrictions prevent exact balance through gray cells alone.')
 para_before(document,anchor,'Turtles and dolphins have no incoming predation in the source diets but nonzero fishing mortality. Reconstructed removals reduce their original losses; residuals remain because rounded F, B, P/B and EE are not mutually exact. No published biomass accumulation, net migration or unrounded native values are available to resolve these terms. Seabird biomass printed as 0.00 prevents reliable reconstruction of its positive consumption.')
 para_before(document,anchor,'Discards retains zero modeled inflow because the calculator omits explicit fleet-discard returns, while animals still eat it. Its negative calculated accumulation cancels that predation arithmetically; this is not an observed stock trend or a validated fishery-return budget. Natural detritus routing does not repair this missing return. Catch rounding, missing removals and defaults remain separate from gray-cell uncertainty.')
 # Shared mapping handoff is consumed, never independently regenerated here.
 replace(pars[3],'Reference: 2019 landings; all recorded catch taxa, with no taxon or group exclusions. Proposed candidate mapping.')
 replace(pars[4],f'{data["source_group_count"]} source groups (+1 synthetic diet-import group); {data["taxon_count"]} catch taxa.')
 replace(pars[5],f'Simple-chain PPR totals {data["simple_ppr_total_tC"]:,.0f} t C in 2019.')
 replace(pars[6],f'{len(data["missing_coeff_zero_catch_taxa"])} taxa lack a saved TL/coefficient, all with zero reference-year catch; their annual contribution is zero. No positive catch has missing simple-chain PPR.')
 replace(pars[7],'Method: simple trophic chain using the retained taxon coefficients and transfer efficiency 0.1.')
 pars[8].clear();pars[8].add_run('Excel taxon appendix and descriptive Sources sheet: ');hyperlink(pars[8],'All 96 proposed taxon mappings','taxon_mapping.xlsx')
 table_rows(tables[1],[[s['confidence'],s['taxa'],format_percent(s['catch_percentage']),format_percent(s['ppr_percentage'])] for s in data['confidence_summary']])
 for i,component in [(2,'membership'),(3,'allocation')]:table_rows(tables[i],[[r['rule'],r['confidence'],format_percent(r['ppr_percentage'])] for r in data['rules'][component]])
 pars[11].clear();pars[11].add_run('Membership evidence: author supplement section 1 and verified historical-name crosswalks; ecological extensions remain explicit. ');hyperlink(pars[11],'Source group definitions','../extracted_tables/companions/Source_group_crosswalk.csv')
 pars[13].clear();pars[13].add_run('Weights and assumptions: complete source catch proportions first, then original biomass proportions; composition and catchability are assumed transferable to 2019. ');hyperlink(pars[13],'Allocation evidence','evidence/allocation_assumptions.json')
 vl=[r for r in data['rows'] if r['confidence']=='Very low'];summary=next(r for r in data['confidence_summary'] if r['confidence']=='Very low')
 replace(pars[15],f'Unresolved taxa: none under the stated proposals. {len(vl)} Very low decisions cover {format_percent(summary["catch_percentage"])} of catch and {format_percent(summary["ppr_percentage"])} of simple-chain PPR. These shares rely on uncertain composition or ecological analogues.')
 grouped=collections.defaultdict(list)
 for r in vl:
  if r['membership_rule'].startswith('Broad'):key='Unknown internal composition across several source guilds; model catch/biomass proportions are proxies.'
  elif r['taxon'] in ['Sarpa salpa']:key='Native herbivore approximated by an alien herbivore guild.'
  elif r['taxon'] in ['Acanthocybium solandri','Pomatomus saltatrix','Trachinotus ovatus']:key='Unlisted pelagic predator approximated by the large-pelagic guild; size or geographic representation uncertain.'
  else:key='Unlisted shark approximated by the demersal-shark guild, with feeding or shelf/deep-water habitat mismatch.'
  grouped[key].append(r['taxon'])
 table_rows(tables[4],[['; '.join(sorted(taxa)),reason] for reason,taxa in grouped.items()])
 replace(pars[18],'Geographic evidence')
 replace(pars[19],'The study covers the Israeli 0–200 m shelf. A ≈ 13–14% and B ≈ 90–100% use the reported 3,725 km² area and an assumed near-complete coastal-shelf containment within the 25,748 km² target EEZ. The study polygon and exact endpoints are unavailable; these are approximate containment scenarios, not a measured intersection.')
 for i,image in [(5,'region_boundary.png'),(6,'study_figure.jpg')]:
  fill(tables[i].rows[1].cells[0],['']);tables[i].rows[1].cells[0].paragraphs[0].add_run().add_picture(str(ev/image),width=Cm(16.8))
 pars[21].clear();pars[21].add_run('Israel Mediterranean EEZ_939; retained project boundary, spherical area approximation. ');hyperlink(pars[21],'Boundary evidence','evidence/region_boundary.geojson')
 pars[23].clear();pars[23].add_run('Corrales et al. (2017), Figure 1, PDF page 3 / printed page 90. The figure shows wider regional bathymetry; the study is restricted to the 0–200 m shelf. ');hyperlink(pars[23],'Original article','../../../sources/1-s2.0-S0924796316302329-main.pdf#page=3')
 pars[24]._p.getparent().remove(pars[24]._p)
 # Keep each figure and caption together and prevent a stranded heading.
 for index in [18,20,22]:pars[index].paragraph_format.keep_with_next=True
 for i in [5,6]:
  for row in tables[i].rows:
   trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  tables[i].rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next=True
 # Remove fixed source placeholder image heights while retaining table styling.
  for row in tables[i].rows:
   for height in row._tr.xpath('./w:trPr/w:trHeight'):height.getparent().remove(height)
 for label,original in manual.items():assert stable_xml(fields[label]._tc)==stable_xml(original),label
 output=v/'validation.docx';document.save(output)
 # Restore every preserve-only reference part byte for byte.
 with zipfile.ZipFile(output) as z:created={n:z.read(n) for n in z.namelist()}
 editable={'[Content_Types].xml','word/document.xml','word/_rels/document.xml.rels'}
 for name,value in parts.items():
  if name not in editable:created[name]=value
 with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
  for name,value in created.items():z.writestr(name,value)
 verify_report(output)
 reopened=Document(output);assert len(reopened.sections)==1
 finalfields={r.cells[0].text:r.cells[1] for r in reopened.tables[0].rows[1:]}
 assert all(stable_xml(finalfields[k]._tc)==stable_xml(old) for k,old in manual.items())
 dump(wk/'qa/template_fidelity.json',{'template_sha256':sha(TEMPLATE),'preserve_only_parts_identical':all(created[n]==parts[n] for n in parts if n not in editable),'manual_cells_unchanged':True,'section_geometry_unchanged':reopened.sections[0]._sectPr.xml==sec._sectPr.xml,'source_template_pages_inspected':3,'automatic_placeholders_removed':True,'added_component':'User-requested group-level mass-balance findings and source-style table'})
 print('Created',output,flush=True)
for data in json.loads((WORK/'outputs/review_data.json').read_text(encoding='utf-8')):build(data)
