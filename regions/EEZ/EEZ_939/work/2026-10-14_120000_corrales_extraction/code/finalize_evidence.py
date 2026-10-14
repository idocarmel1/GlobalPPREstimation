"""Verify saved scientific results and finalize portable evidence without solving again."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
import csv,json,hashlib,os,sys,collections,datetime,math,shutil,re
from urllib.parse import unquote,urlparse
from xml.etree import ElementTree as E
import numpy as np
import pandas as pd
import openpyxl
from docx import Document
ROOT=Path(__file__).resolve().parents[6];WORK=Path(__file__).resolve().parents[1]
PAPER=ROOT/'regions/EEZ/EEZ_939/papers/Corrales-2010';BASE=PAPER/'models'
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tools/skills/paper-to-ppr/resources/extraction/scripts'))
from tools.project_core.workbooks.workbooks import read_book,digest_tables
from tools.project_core.validation.validation_percentage_format import verify_report
from source_fidelity import reconstruct_source_workbook
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def load(p):return json.loads(p.read_text(encoding='utf-8'))
def frame(x):
 assert x['type']=='DataFrame'
 assert len(x['index'])==len(set(x['index'])) and len(x['columns'])==len(set(x['columns']))
 def value(v):
  if isinstance(v,dict) and 'nonfinite' in v:return {'NaN':np.nan,'+Infinity':np.inf,'-Infinity':-np.inf}[v['nonfinite']]
  return v
 return pd.DataFrame([[value(v) for v in row] for row in x['data']],index=x['index'],columns=x['columns'])
def numeric(v):return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v)
def relative(p,base):return Path(os.path.relpath(p,base)).as_posix()
def local_link(target,folder):
 if target.startswith(('https://','http://','#')):return True
 assert not re.match(r'^[A-Za-z]:|^[/\\]|^file:|^localhost',target),target
 p=folder/unquote(target.split('#')[0]);assert p.exists(),(target,str(p));assert p.resolve().is_relative_to(ROOT)
 return True
def add_excel_links(file,sources):
 # Artifact-tool has no documented hyperlink surface in the bounded help lookup.
 # Patch only OOXML relationships and cloned hyperlink font/style definitions.
 with ZipFile(file) as z:parts={n:z.read(n) for n in z.namelist()}
 X='http://schemas.openxmlformats.org/spreadsheetml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships'
 styles=E.fromstring(parts['xl/styles.xml']);fonts=styles.find('{'+X+'}fonts');xfs=styles.find('{'+X+'}cellXfs');sheet=E.fromstring(parts['xl/worksheets/sheet2.xml'])
 relpath='xl/worksheets/_rels/sheet2.xml.rels';rels=E.fromstring(parts[relpath]) if relpath in parts else E.Element('{'+P+'}Relationships')
 links=sheet.find('{'+X+'}hyperlinks')
 if links is not None:sheet.remove(links)
 links=E.SubElement(sheet,'{'+X+'}hyperlinks');newstyles={};checks=[]
 for i,row in enumerate(sources[1:],2):
  ref=f'C{i}';target=row[2];local_link(target,file.parent)
  cell=next(c for c in sheet.iter('{'+X+'}c') if c.get('r')==ref);old=int(cell.get('s','0'))
  if old not in newstyles:
   xf=deepcopy(list(xfs)[old]);font=deepcopy(list(fonts)[int(xf.get('fontId','0'))])
   for child in list(font):
    if child.tag in ['{'+X+'}u','{'+X+'}color']:font.remove(child)
   E.SubElement(font,'{'+X+'}u',{'val':'single'});E.SubElement(font,'{'+X+'}color',{'rgb':'FF0563C1'})
   xf.set('fontId',str(len(fonts)));fonts.append(font);xf.set('applyFont','1');newstyles[old]=len(xfs);xfs.append(xf)
  cell.set('s',str(newstyles[old]));rid='corralesLink'+str(i)
  E.SubElement(rels,'{'+P+'}Relationship',{'Id':rid,'Type':R+'/hyperlink','Target':target,'TargetMode':'External'})
  E.SubElement(links,'{'+X+'}hyperlink',{'ref':ref,'{'+R+'}id':rid});checks.append({'cell':ref,'target':target,'relative_or_public':True,'blue_underlined':True})
 fonts.set('count',str(len(fonts)));xfs.set('count',str(len(xfs)))
 parts['xl/styles.xml']=E.tostring(styles,encoding='utf-8',xml_declaration=True);parts['xl/worksheets/sheet2.xml']=E.tostring(sheet,encoding='utf-8',xml_declaration=True);parts[relpath]=E.tostring(rels,encoding='utf-8',xml_declaration=True)
 with ZipFile(file,'w',ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)
 return checks

snapshot=load(WORK/'inputs/region_tables.json');live=read_book(ROOT/'regions/EEZ/EEZ_939/EEZ_939.xlsx')
assert digest_tables(snapshot)==digest_tables(live),'Regional inputs changed; stop mixed-version publication'
source_manifest=load(PAPER/'source_manifest.json')
for record in source_manifest['files']:assert sha(PAPER/record['path'])==record['sha256']
allresults=[]
for model in sorted(BASE.iterdir()):
 if not (model/'model.json').exists():continue
 data=load(model/'model.json');period='1990_1994' if '1990_1994' in model.name else '2008_2012';source=BASE/f'Corrales_2017_ICS_{period}'
 ev=model/'extracted_tables/evidence';ev.mkdir(parents=True,exist_ok=True);diag=model/'results/diagnostics';diag.mkdir(parents=True,exist_ok=True)
 before_json_hash=sha(model/'model.json');group_hash=hashlib.sha256(json.dumps(data['group'],sort_keys=True,ensure_ascii=False).encode()).hexdigest()
 # Update only the embedded import-table evidence in earlier gray-only candidates.
 # Scientific group payload and loader-consumed metadata must remain unchanged.
 changed_evidence=False
 if model.name.endswith('_gray_fit'):
  rows=list(csv.reader((model/'extracted_tables/Diet_composition.csv').open(encoding='utf-8-sig',newline='')))
  if data['extraction_source_tables']['Diet_composition.csv']!=rows:data['extraction_source_tables']['Diet_composition.csv']=rows;dump(model/'model.json',data);changed_evidence=True
 if changed_evidence:
  assert group_hash==hashlib.sha256(json.dumps(data['group'],sort_keys=True,ensure_ascii=False).encode()).hexdigest()
  dump(diag/'source_identity_bridge.json',{'schema_version':1,'executed_source_sha256':before_json_hash,'current_source_sha256':sha(model/'model.json'),'scientific_group_payload_sha256':group_hash,'change':'Embedded diet import evidence updated to match the fitted CSV; group payload and all loader-consumed metadata unchanged. Existing diagnostics retained, no numerical rerun claimed.'})
 # Verify standard import cells against the exact embedded current tables.
 checks=[]
 for filename,embedded in data['extraction_source_tables'].items():
  path=model/'extracted_tables'/filename;assert path.exists()
  if filename.endswith('.csv'):
   actual=list(csv.reader(path.open(encoding='utf-8-sig',newline='')));assert actual==embedded,(model.name,filename)
  checks.append(filename)
 source_data=load(source/'model.json');source_groups={int(g['group_seq']):g for g in source_data['group']}
 gray=[]
 with (source/'extracted_tables/companions/Source_diet_cells.csv').open(encoding='utf-8',newline='') as f:
  gray=[r for r in csv.DictReader(f) if r['status']=='censored_less_than_0.01']
 graykeys={(int(r['consumer_seq']),int(r['prey_seq'])) for r in gray}
 fitted='gray_fit' in model.name;changedcells=[]
 for group in data['group']:
  seq=int(group['group_seq']);old=source_groups[seq]
  for field in ['group_name','biomass','pb','qb','ee','gs','biomass_accum','landings_by_fleet','discards_by_fleet','taxon_descr']:assert group[field]==old[field],(model.name,seq,field)
  od={int(d['prey_seq']):d['proportion'] for d in (old.get('diet_descr') or {}).get('diet',[])}
  for cell in (group.get('diet_descr') or {}).get('diet',[]):
   prey=int(cell['prey_seq']);v=cell['proportion'];original=od[prey]
   if fitted and (seq,prey) in graykeys:assert 0<=float(v)<.01;changedcells.append({'consumer_seq':seq,'prey_seq':prey,'source_value':original,'retained_value':v})
   else:assert v==original,(model.name,seq,prey)
 dump(ev/'CURRENT_VARIANT_CHECK.json',{'schema_version':1,'model_id':model.name,'canonical_sha256':sha(model/'model.json'),'standard_imports_verified':checks,'protected_biology_catch_taxonomy_unchanged':True,'all_non_gray_diet_cells_unchanged':True,'source_gray_cell_count':len(gray),'fitted_gray_cells':changedcells,'ordinary_diet_blanks':'Remain -9999 in canonical JSON; runtime zero assumption only','fate_change':'Authorized assumption only in derived variants','aggregate_catch_change':'Authorized F*B / known-component counterfactual only in F_removals variant'})
 # Current workbook reconstruction uses the maintained domain writer; source
 # templates/import schemas govern this artifact, including published companions.
 reconstruction=ev/'reconstructed.xlsx';reconstruct_source_workbook(data,str(reconstruction));roundtrip=load(ev/'SOURCE_FIDELITY_CHECK.json');assert roundtrip['numeric_value_and_missing_mask_match'],roundtrip
 if model==source:
  admission=load(WORK/f'qa/{period}_admission.json');assert admission['input_sha256']==sha(model/'model.json');dump(diag/'diagnostic_manifest.json',admission)
  for name in ['loader_receipt.json','check_outcomes.json','validate.log','massbalance.log','conversion.log']:
   p=WORK/f'qa/{period}_{name}'
   if p.exists():shutil.copy2(p,ev/name)
 else:
  manifest=load(diag/'diagnostic_manifest.json');assert manifest['constructor_outcome']=='passed'
  assert manifest['source_sha256'] in [sha(model/'model.json'),before_json_hash]
  for name,value in manifest['engine_hashes'].items():assert sha(ROOT/'tools/scientific_code/PPREstimation'/name)==value
  state=load(diag/'runtime_state.json');groups=frame(state['groups']);dc=frame(state['diet']);fate=frame(state['fate']);groups=groups.fillna({'catch':0,'p':0,'q':0,'M0':0,'biomass_accum':0,'net_migration':0,'detritus_import':0,'egestion':0,'respiration':0})
  production=groups['catch']+groups['predation']+groups['biomass_accum']+groups['net_migration']-groups['detritus_import']+groups['M0'];consumption=production+groups['egestion']+groups['respiration']
  budget=pd.DataFrame({'group_name':groups.group_name,'trophic_info':groups.trophic_info,'production':groups.p,'catch':groups['catch'],'predation':groups.predation,'biomass_accumulation':groups.biomass_accum,'net_migration':groups.net_migration-groups.detritus_import,'other_mortality':groups.M0,'production_sinks':production,'production_residual':production-groups.p,'consumption':groups.q,'egestion':groups.egestion,'respiration':groups.respiration,'consumption_recomputed':consumption,'consumption_residual':consumption-groups.q})
  budget['production_relative_residual']=budget.production_residual/np.maximum(abs(budget.production),1e-12);budget['consumption_relative_residual']=budget.consumption_residual/np.maximum(abs(budget.consumption),1e-12);budget.index.name='group_seq';budget.to_csv(diag/'mass_balance_by_group.csv')
  before=frame(load(diag/'loaded_before.json')['groups']);changes=[]
  for seq in groups.index:
   for field in groups.columns.intersection(before.columns):
    a=before.loc[seq,field];b=groups.loc[seq,field]
    if pd.isna(a) and pd.isna(b):continue
    same=a==b if not pd.isna(a) and not pd.isna(b) else False
    if not same:changes.append({'group_seq':int(seq),'group_name':groups.loc[seq,'group_name'],'field':field,'before':None if pd.isna(a) else (a.item() if isinstance(a,np.generic) else a),'after':None if pd.isna(b) else (b.item() if isinstance(b,np.generic) else b)})
  dump(diag/'runtime_transformations.json',changes)
  negative=[];coeff=[];matrixchecks=[]
  for method in ['GE','TE','With Egestion']:
   ret=load(diag/(method.replace(' ','_')+'_full_return.json'));s=frame(ret['SPPR']);assert len(s)==len(groups);assert set(s.index)==set(groups.index)
   imported=list(groups.index[groups.trophic_info.eq('Import')]);pp=list(groups.index[groups.trophic_info.eq('PP')]);dets=list(groups.index[groups.trophic_info.eq('DET')]);assert set(s.columns)==set(imported+pp+dets)
   for key in ['SPPR','A','L']:
    f=frame(ret[key]);a=f.to_numpy(dtype=float)
    assert len(f.index)==len(ret[key]['data']) and all(len(r)==len(f.columns) for r in ret[key]['data'])
    assert np.array_equal(np.isnan(a),ret[key]['nan_mask']) and np.array_equal(np.isposinf(a),ret[key]['positive_infinity_mask']) and np.array_equal(np.isneginf(a),ret[key]['negative_infinity_mask'])
    for ri,ci in zip(*np.where(a < -1e-12)):negative.append({'method':method,'matrix':key,'row_id':int(f.index[ri]),'row_name':ret['seq2name'].get(str(f.index[ri]),str(f.index[ri])),'column_id':int(f.columns[ci]),'column_name':ret['seq2name'].get(str(f.columns[ci]),str(f.columns[ci])),'value':float(a[ri,ci]),'tolerance':1e-12})
   scopes={'all':list(s.columns),'inner':[x for x in s.columns if x not in imported],'PP':pp}
   for scope,cols in scopes.items():
    totals=s[cols].sum(axis=1);actual=float(groups['catch'].reindex(s.index).dot(totals));key={'all':'ppr_all','inner':'ppr_inner','PP':'ppr_pp_only'}[scope];expected=ret['diagnostics']['footprint'][key]
    assert np.isclose(actual,expected,rtol=1e-12,atol=1e-12),(model.name,method,scope,actual,expected)
    for seq,value in totals.items():coeff.append({'model_id':model.name,'group_seq':int(seq),'group_name':groups.loc[seq,'group_name'],'scope':scope,'method':method,'sppr':float(value),'units':'wet-weight primary-production equivalent per wet-weight group production','diagnostic_status':ret['diagnostics']['status'],'strict_mass_balance':False,'production_eligible':False})
   assert np.isclose(max(abs(budget.production_relative_residual)),ret['diagnostics']['model_input']['p_max_rel_residual'],rtol=1e-12,atol=1e-12)
   assert np.isclose(max(abs(budget.consumption_relative_residual)),ret['diagnostics']['model_input']['q_max_rel_residual'],rtol=1e-12,atol=1e-12)
   matrixchecks.append({'method':method,'SPPR_shape':list(s.shape),'all_axes_unique':True,'all_source_groups_and_synthetic_import_retained':True,'null_nonfinite_masks_verified':True,'source_scopes_and_catch_footprints_reconciled':True,'max_p_relative_residual':float(max(abs(budget.production_relative_residual)))})
  pd.DataFrame(coeff).to_csv(diag/'group_sppr.csv',index=False);dump(diag/'negative_entries_all_matrices.json',negative);dump(diag/'matrix_verification.json',matrixchecks)
  assert not [r for r in negative if r['matrix']=='SPPR']
  shutil.copy2(WORK/'code/run_sppr.py',diag/'run_sppr.py')
  (diag/'README.md').write_text('# Direct SPPR diagnostics\n\nGE, TE and With Egestion were executed with the settings in diagnostic_manifest.json. Complete arrays, masks, axes and names remain in each full_return.json. The group_sppr.csv is a convenience scope-sum view, with all scientific restrictions retained.\n\nTo reproduce with the required scientific Python environment, run `python run_sppr.py` from this directory, or pass the model folder explicitly. The script finds the project root and restores the explicit fate matrix on runtime nonfeeding rows. It does not use the engine’s forced balanced_model copy. Fresh execution replaces diagnostics in this exact candidate, so preserve researcher work first.\n\nFitted gray values are assumptions, not author-recovered diets. Source fields and inferred removals are distinguished in model_notes.md. `inner` excludes the synthetic imported-diet column; `PP` also excludes both detritus columns. All saved methods retain det_collapse_mode=never.\n',encoding='utf-8')
 if fitted:
  shutil.copy2(WORK/'code/fit_gray_diet.py',ev/'gray_diet_fit/executed_fit_script.py')
  if 'F_removals' in model.name:shutil.copy2(ev/'gray_diet_fit/catch_reconstruction.csv',model/'extracted_tables/companions/Reconstructed_total_removals.csv')
 # Write a completed source/variant extraction account, superseding placeholders.
 status=load(diag/'diagnostic_manifest.json');grades={k:v.get('diagnostics',{}).get('status',v.get('status')) for k,v in status['methods'].items()}
 sourceonly=model==source
 report=f'''# Corrales Israeli continental shelf {period.replace('_','–')}

Corrales et al. (2017), Journal of Marine Systems 170:88–102, doi:10.1016/j.jmarsys.2017.02.004. The supplied Corrales-2010 alias is retained. The model describes 3,725 km² of shelf at 0–200 m depth, rather than the entire Israeli Mediterranean EEZ. There are {len(data['group'])} ecological groups; older source IDs 24 and 35 are absent and the import crosswalk retains the original IDs.

Article Table 1 (PDF 4 / printed 91) supplies B, P/B, Q/B, EE, P/Q, landings and discards for both periods. Table 2 (PDF 6 / printed 93) supplies TL, fishing mortality F, predation mortality M2, other mortality, F/Z and detritus flow. Table 3 (PDF 8 / printed 95) and all eight supplementary tables are retained as companions. Bold initial-input markers and author-estimated roles remain explicit. Supplement section 1 supplies group membership; section 5 supplies the two continued diet matrices and their totals. The gray status comes from the source cell shading, not inferred diet magnitudes.

Ordinary blank diet cells are assumed zero only at runtime. Gray cells mean less than 0.01; source models preserve their -9999 missing marker. There are {len(gray)} such cells in this period. Published known consumer diet subtotals are 0.98–1.03 in 1990–1994 and 0.97–1.01 in 2008–2012. Source values were not normalized; runtime normalization is recorded separately. Seabird biomass printed 0.00 loses biological precision. Biomass accumulation, GS, habitat fraction, native detritus fate and several catch components are unpublished, not observed zero. Biomass values are explicitly per whole model area.

Known rounded landings sum to 0.79 / 0.65, compared with published aggregate landings 0.80 / 0.64; known discards sum to 0.10 / 0.29 compared with aggregate 0.14 / 0.29 t km⁻² yr⁻¹. The source tables remain unchanged. These differences do not justify distributing an invented residual among groups. The original loader's incomplete aggregate-catch default can omit known partial components; this is distinguished from missing source catch.

The label Mullets means red mullets Mullus barbatus and M. surmuletus in the author definitions. Mugilidae grey mullets instead occur in Small native dem. fishes. Several author spellings are retained verbatim and crosswalks are separately evidenced.

This exact candidate has direct method outcomes {json.dumps(grades,ensure_ascii=False)}. A source checklist or rounded equation check does not establish balance of the author's original native model. Current import/JSON/reconstructed-workbook cells and missing masks were checked with the maintained domain routine; deliberate fitted-gray departures are tested separately in CURRENT_VARIANT_CHECK.json. These reconstruction checks establish value retention, not ecological validity.

''' 
 if sourceonly:report+='The source-only model loads into ModelData but cannot construct a multi-detritus calculator without per-pool fate. All three direct methods are NOT_RUN with the retained exception. No native EwE database was obtained.\n\n'
 else:report+='Natural mortality and egestion are routed to natural Detritus under the researcher-authorized fate assumption. Discards remains a separate pool; explicit fishery-return amounts are retained as a donor ledger, but the current calculator omits recipient fleet returns. Runtime GS is 0.2, living accumulation is zero by default, detritus EE becomes one, and detritus accumulation is calculated; these are not published observations. Full differences and all-group balances are in results/diagnostics.\n\n'
 if fitted:report+='Only gray raw fractions were fitted in [0,0.01), with all printed diet cells fixed before normalization and ordinary blanks zero on the runtime copy. Feasibility/minimax results and every fitted cell are in gray_diet_fit. Exact balance is infeasible under these constraints. This is one constrained reconstruction, not a uniquely recovered native diet.\n\n'
 if 'F_removals' in model.name:report+='Missing total removals were reconstructed from published F × B where this did not undercut a known component. Conflicting smaller estimates were rejected; known components were retained with unknown components zero only on the experimental copy. Complete known catches remain fixed. Canonical aggregate export records this experiment; original Landings/Discards imports preserve source components. The extra Reconstructed_total_removals.csv is an evidence companion, not a new native fleet table. The largest production residual remains 7.4286% for the older period and 14.6212% for the later period. Strict mass balance is false.\n\n'
 report+='The conditional detritus-collapse and discard-removal experiments were not run because the researcher required a mass-balanced prerequisite. No selection, signed approval or annual regional PPR adoption is inferred.\n\n[Canonical notes](../../model_notes.md) · [Current extraction check](CURRENT_VARIANT_CHECK.json) · [Full diagnostic returns](../../results/diagnostics)\n'
 (ev/'REPORT.md').write_text(report,encoding='utf-8')
 notes=f'''# {model.name}

Source identity: Corrales et al. (2017), Modeling the role and impact of alien species and fisheries on the Israeli marine continental shelf ecosystem, Journal of Marine Systems 170:88–102, doi:10.1016/j.jmarsys.2017.02.004. [Article](../../sources/1-s2.0-S0924796316302329-main.pdf), [author supplement](../../sources/supplements/1-s2.0-S0924796316302329-mmc1.docx). The researcher-supplied Corrales-2010 folder alias is retained; publication year is 2017.

Modeled period and area: {period.replace('_','–')}; Israeli Mediterranean continental shelf at 0–200 m depth, 3,725 km²; {len(data['group'])} source groups. Whole-model-area biomass basis is published. This is partial shelf coverage of EEZ_939, not a full-EEZ model.

Departures from publication: {'No fitted biological or diet values; source literals and unknowns retained.' if sourceonly else 'Researcher-authorized assumed natural detritus fate; source biology and fleet components retained.'} Ordinary blank diets become zero only at runtime; source gray diet cells are -9999, not measured zero. Unpublished BA, GS, habitat fraction and missing fisheries components remain unknown in source tables. The source prints seabird B=0.00 and rounded consumer diets; neither supplies unrounded native precision. See article Table 1 PDF 4, Table 2 PDF 6, Table 3 PDF 8 and supplement sections 1 and 5.

Actual direct outcomes: {json.dumps(grades,ensure_ascii=False)}. {'Calculator construction is blocked by absent per-pool detritus fate.' if sourceonly else 'The executed runtime uses recorded zero-catch/BA and GS defaults, restores explicit producer fate, and normalizes diets on a copy. Detritus EE/accumulation are software-derived. Explicit native fleet returns to Discards are not represented. Strict mass balance is false.'}

'''
 if fitted:notes+='The researcher explicitly authorized filling gray diet cells on 14 October 2026. Fitted raw values are nonnegative and strictly below 0.01. All printed diet fractions and non-diet biological parameters remain unchanged before runtime normalization. The constrained linear-program solution is not unique source recovery; exact balance is infeasible. [Fitted cells and feasibility](extracted_tables/evidence/gray_diet_fit/fit_summary.json).\n\n'
 if 'F_removals' in model.name:notes+='The researcher additionally authorized missing total-removal reconstruction using published F × biomass. Each group’s source operands, reconstruction, known catch subtotal, rejected contradictions and retained value are recorded in [catch reconstruction](extracted_tables/companions/Reconstructed_total_removals.csv). Known total catches and original fleet components remain unchanged. Lower F × B values conflicting with known components were rejected; the remaining missing component was zero only in the experimental known-subtotal counterfactual. The canonical export is an experimental aggregate; native Landings/Discards imports still preserve original missing components. This variant is runnable but is not an approved balanced model.\n\n'
 notes+='Evidence links: [extraction account](extracted_tables/evidence/REPORT.md), [current parameter/diet comparison](extracted_tables/evidence/CURRENT_VARIANT_CHECK.json), [round-trip check](extracted_tables/evidence/SOURCE_FIDELITY_CHECK.json), [diagnostics and settings](results/diagnostics).\n\n'
 if (model/'model_validation/validation.docx').exists():notes+='[Candidate validation report](model_validation/validation.docx) and [all-taxon mapping appendix](model_validation/taxon_mapping.xlsx) describe remaining balance issues and proposed mapping confidence. Mapping is pending alignment; no researcher verdict, selection or annual regional calculation has been adopted.\n\n'
 notes+='Automatic detritus collapse and removing Discards remain deferred: the researcher requested those only after a mass-balanced model still fails, and strict balance has not been reached.\n'
 (model/'model_notes.md').write_text(notes,encoding='utf-8')
 allresults.append({'model_id':model.name,'groups':len(data['group']),'grades':grades,'canonical_sha256':sha(model/'model.json'),'source_only':sourceonly})

# Validate the two finished reports and all-taxon appendices.
for period in ['1990_1994','2008_2012']:
 model=BASE/f'Corrales_2017_ICS_{period}_gray_fit_F_removals';v=model/'model_validation';ev=v/'evidence';data=load(ev/'mapping_handoff.json');spec=load(WORK/f'outputs/appendix_{period}.json')
 xlchecks=add_excel_links(v/'taxon_mapping.xlsx',spec['sources']);wb=openpyxl.load_workbook(v/'taxon_mapping.xlsx');s=wb['Taxon mapping'];assert len(list(s.iter_rows(min_row=6,max_col=7,values_only=True)))==96
 actual=[list(row) for row in s.iter_rows(min_row=5,max_row=101,max_col=7,values_only=True)]
 for row,expected in zip(actual,spec['matrix']):
  for a,b in zip(row,expected):assert np.isclose(a,b,rtol=1e-13,atol=1e-12) if numeric(a) and numeric(b) else a==b,(a,b)
 assert s.freeze_panes=='A6' and len(s.tables)==1
 assert all(abs(sum(g['weight'] for g in r['groups'])-1)<1e-12 for r in data['rows']);assert sum(r['taxa'] for r in data['confidence_summary'])==96
 assert abs(sum(r['catch_percentage'] for r in data['confidence_summary'])-100)<1e-8 and abs(sum(r['ppr_percentage'] for r in data['confidence_summary'])-100)<1e-8
 assert all(abs(sum(r['ppr_percentage'] for r in data['rules'][component])-100)<1e-8 for component in ['membership','allocation'])
 for r in data['rows']:
  if r['catch_t']==0:assert r['simple_ppr_tC']==0
  if r['catch_t']>0:assert r['simple_ppr_tC'] is not None
 for i in range(2,11):c=wb['Sources'].cell(i,3);assert c.hyperlink and c.font.underline=='single' and c.font.color.rgb=='FF0563C1'
 verify_report(v/'validation.docx');doc=Document(v/'validation.docx');linkchecks=[]
 for rel in doc.part.rels.values():
  if rel.reltype.endswith('/hyperlink'):local_link(rel.target_ref,v);linkchecks.append(rel.target_ref)
 xml=doc.part.element;W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
 for link in xml.iter(W+'hyperlink'):
  for run in link.findall(W+'r'):assert run.find(W+'rPr/'+W+'u').get(W+'val')=='single' and run.find(W+'rPr/'+W+'color').get(W+'val')=='0563C1'
 wk=v/'work'/WORK.name;fidelity=load(wk/'qa/template_fidelity.json');assert fidelity['preserve_only_parts_identical'] and fidelity['manual_cells_unchanged'] and fidelity['section_geometry_unchanged']
 dump(ev/'review_checks.json',{'schema_version':1,'model_id':model.name,'year':2019,'catch_basis':'landings','all_96_taxa_present_once':True,'source_group_inventory_matches':True,'numeric_types_and_unrounded_sort_verified':True,'frozen_filtered_appendix':True,'weight_sums_verified':True,'five_category_and_two_component_denominators_verified':True,'zero_catch_and_missing_coefficient_masks_verified':True,'simple_ppr_total_tC':data['simple_ppr_total_tC'],'catch_total_t':data['catch_total_t'],'mapping_adoption_state':'pending alignment','Word_percentage_format_verified':True,'manual_cells_preserved':True,'template_opaque_parts_and_section_geometry_preserved':True,'all_repository_links_relative_and_resolving':True,'Word_links':linkchecks,'Excel_links':xlchecks,'visible_links_blue_underlined':True,'regional_tables_unchanged':True})
 dump(ev/'render_verification.json',{'schema_version':1,'report_sha256':sha(v/'validation.docx'),'pages':6,'all_pages_visually_inspected':True,'renderer':'Hidden independent read-only Word export; Poppler 150 dpi rasterization','layout':'Readable tables and figures; repeated headers; no clipping, accidental blank pages or stranded headings','template_reference_pages_inspected':3,'packaged_renderer_limitation':'LibreOffice unavailable; no original Word document modified','manual_cells_unchanged':True})
 dump(ev/'source_recovery_search.json',{'date':'2026-10-14','source_supplied_by':'researcher','records':[{'url':'https://www.sciencedirect.com/science/article/pii/S0924796316302329','result':'Publisher access unavailable / 403; researcher local PDF and supplement used'},{'url':'https://ecobase.ecopath.org/','result':'Bounded catalogue search found no verified native export for these exact Israeli shelf periods; no accession invented'},{'url':'https://pureadmin.uhi.ac.uk/ws/portalfiles/portal/2266377/Corrales_et_al_2017_AAM.pdf','result':'Accepted author version consulted; rounded Table 1 and seabird 0.00 do not restore unrounded native parameters'},{'url':'https://hdl.handle.net/10261/154023','result':'Repository link unavailable'},{'url':'https://www.ncbi.nlm.nih.gov/Taxonomy/Browser/wwwtax.cgi?id=27405','result':'Full page read; names Penaeus japonicus and homotypic Marsupenaeus japonicus'},{'url':'https://fish-commercial-names.ec.europa.eu/fish-names/species/oblada-melanurus_en','result':'Primary search snippet supports Oblada melanura spelling crosswalk; full-page access limited'},{'url':'https://obis.org/taxon/127056','result':'Primary search snippet supports accepted Oblada melanurus; subsequent full-page open timed out'}],'limitations':'No author contact, native EwE database, unrounded biological inputs or published numerical fate matrix obtained'})
 print('Reviewed report and appendix',period,flush=True)
dump(WORK/'outputs/final_model_outcomes.json',allresults)
print('Saved evidence and verified results for',len(allresults),'models',flush=True)
