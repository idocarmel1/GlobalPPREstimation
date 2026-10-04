import json,sys,openpyxl
from pathlib import Path
sys.path.insert(0,'tools')
from workbooks import *
g=json.loads(Path('tools/knowledge_graph/graph.json').read_text(encoding='utf-8'))
print('GRAPH type',list(g)[:5])
for n in g.get('nodes',[]):
 if any(s in str(n).lower() for s in ['wcpo','wcp-2007','kiribati']):print('GRAPH',str(n)[:1200])
p=openpyxl.load_workbook('Project.xlsx',read_only=True,data_only=False)
print('SHEETS',p.sheetnames)
for name in ['Papers','Models','Models & coverage']:
 if name in p:
  for row in p[name].values:
   if 'EEZ_941' in row or any(v and str(v).startswith('941_') for v in row):print('META',name,json.dumps(row,default=str))
p.close()
from docx import Document
for file in ['tools/templates/Model_validation_template.docx']:
 d=Document(file)
 print('DOC',file)
 for i,t in enumerate(d.tables):
  print('TABLE',i,len(t.rows),len(t.columns))
  for r in t.rows: print([c.text for c in r.cells])
 print('PARAGRAPHS',[p.text for p in d.paragraphs])
