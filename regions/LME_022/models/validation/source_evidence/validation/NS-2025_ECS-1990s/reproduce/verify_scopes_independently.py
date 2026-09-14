"""Cross-check all taxon/scoped outputs against original exported group SPPR.

Uses only workbook cells and independent arithmetic; does not call the builder,
mapping weights resolver, or verifier implementation.
"""
from pathlib import Path
from collections import defaultdict
import openpyxl,csv,json,math
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'tools/run_sppr.py').is_file() and (p/'PPRAtlas/archive').is_dir());BASE=ROOT/'data/LME_022/validation/NS-2025_ECS-1990s'
STEM='22_20251990_East_Coast_of_Scotland_(1991-1995)'
up=openpyxl.load_workbook(BASE/'sppr'/f'{STEM}.xlsx',data_only=True)
out=openpyxl.load_workbook(BASE/'evaluation/data/LME_022/models'/f'{STEM}.xlsx',data_only=True)
mapping=list(csv.DictReader((BASE/'evaluation/data/LME_022/mapping'/f'{STEM}.csv').open(encoding='utf-8-sig')))
choices={r['taxon']:[x.strip() for x in r['group'].split('|')] for r in mapping}
cr=list(out['Catch'].values);years=list(cr[0][4:]);catch={r[0]:dict(zip(years,r[4:])) for r in cr[1:]}
solo=defaultdict(float)
for t,choices_t in choices.items():
    if len(choices_t)==1 and choices_t!=['Unresolved']:solo[choices_t[0]]+=sum(catch[t].values())
weights={}
for t,gs in choices.items():
    if len(gs)==1:weights[t]=[1.0]
    else:
        denom=sum(solo[g] for g in gs);assert denom>0
        weights[t]=[solo[g]/denom for g in gs]
checks=0;annualchecks=0;maxdiff=0;examples={}
for scope,ts,ps in [('all','SPPR','PPR by method'),('inner','sppr_inner','PPR inner'),('PP','sppr_PP','PPR PP')]:
    u=list(up['sppr_'+scope].values);methods=list(u[0][2:]);source={r[1]:dict(zip(methods,r[2:])) for r in u[1:]}
    s=list(out[ts].values);values={r[0]:dict(zip(methods,r[4:])) for r in s[4:] if r[0] in choices}
    for t,gs in choices.items():
        for m in methods:
            if gs==['Unresolved'] or any(source[g][m] is None and w>0 for g,w in zip(gs,weights[t])):expected=None
            else:expected=round(sum(source[g][m]*w for g,w in zip(gs,weights[t]) if w>0),6)
            actual=values[t][m]
            assert (actual is None)==(expected is None),(scope,t,m,expected,actual)
            if expected is not None:
                diff=abs(expected-actual);maxdiff=max(maxdiff,diff);assert diff<=1e-6,(scope,t,m,expected,actual)
            checks+=1
    p=list(out[ps].values);header=next(r for r in p if r[0]=='method');ppr={r[0]:dict(zip(header[2:],r[2:])) for r in p if r[0] in methods}
    for m in methods:
        for y in years:
            terms=[catch[t][y]*values[t][m] for t in values if values[t][m] is not None and catch[t][y]]
            expected=round(sum(terms),3) if terms else None;actual=ppr[m][y]
            assert (actual is None)==(expected is None),(scope,m,y)
            if expected is not None:assert abs(expected-actual)<=.001,(scope,m,y,expected,actual)
            annualchecks+=1
    examples[scope]={m:ppr[m][2019] for m in ['new_GE','new_WithEgestion','new_TE_EEfix','SPPR_1995_TE0.1']}
before=openpyxl.load_workbook(BASE/'baseline_before_health_fix/PPR_unflagged_TE.xlsx',data_only=True)
def data_rows(book):
    r=list(book['PPR by method'].values);h=next(x for x in r if x[0]=='method');return {x[0]:list(x[2:]) for x in r if x[0] in methods or str(x[0]).startswith('simple')}
before_rows=data_rows(before);after_rows=data_rows(out)
assert {k:v for k,v in before_rows.items() if not k.startswith('MC_')}=={k:v for k,v in after_rows.items() if not k.startswith('MC_')},'Deterministic all-source PPR numbers changed with health-only integration correction'
summary=dict((r[0],r[1]) for r in out['Summary'].values if r[0])
assert summary['headline method below']!='new_TE_EEfix'
status={r[0]:r[1] for r in out['PPR by method'].values if r[0] in methods}
assert status['new_TE_EEfix'].startswith('FAILED')
assert status['new_GE']==status['new_WithEgestion']=='ok'
result={'taxon_scope_cells_compared_to_upstream':checks,'annual_method_scope_totals_compared':annualchecks,'max_taxon_sppr_absolute_difference':maxdiff,'composite_weights':'independently recomputed from singly mapped catch across 70 years','all_source_deterministic_PPR_and_Jensen_values_unchanged_from_baseline':True,'all_source_Monte_Carlo_values_match_baseline_this_run':all(before_rows[k]==after_rows[k] for k in before_rows if k.startswith('MC_')),'Monte_Carlo_note':'MC equality is not required on reproduction because upstream draws are unseeded.','headline':summary['headline method below'],'method_status':status,'examples_2019_t_wet':examples}
(BASE/'SCOPES_VERIFICATION.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
