"""Reproduce the source extraction; never writes the selected pilot."""
from pathlib import Path
import csv, json, re, shutil, subprocess, sys, unicodedata
import openpyxl
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'tools/run_sppr.py').is_file() and (p/'PPRAtlas/archive').is_dir())
BASE=ROOT/'data/LME_022/validation/NS-2025_ECS-1990s'
SKILL=ROOT/'skills/codex/ecopath-paper-to-ppr'
SRC=ROOT/'PPRAtlas/archive/regions/LME_022/NS-2025'
STEM='22_20251990_East_Coast_of_Scotland_(1991-1995)'
MODEL_DIR=BASE/'extraction/20251990_East_Coast_of_Scotland_1991-1995'
BASE.mkdir(parents=True,exist_ok=True)
sys.path.insert(0,str(SKILL/'scripts'))
from pdfgrid import get_words
pdf=SRC/'pdf-aa2228c7.pdf'
lines=get_words(str(pdf),6,merge_gap=0.4)
header=next(l for l in lines if l.text.startswith('Both '))
anchors=[c.xc for c in header.cells]
assert len(anchors)==13,anchors
fields=['tl','B_1890','biomass','PB_1890','pb','QB_1890','qb','EE_1890','ee','PQ_1890','pq','E_1890','E_1990']
groups=[];source_rows=[]
for line in lines:
    if not line.cells or not re.fullmatch(r'\d+',line.cells[0].text):continue
    n=int(line.cells[0].text)
    if not 1<=n<=25 or line.cells[0].x0>110:continue
    vals={};name=[]
    for cell in line.cells[1:]:
        if cell.xc<anchors[0]-20:name.append(cell.text)
        elif cell.xc<anchors[-1]+20:
            idx=min(range(13),key=lambda i:abs(cell.xc-anchors[i]))
            if re.fullmatch(r'\d+(?:\.\d+)?',cell.text):
                assert fields[idx] not in vals,(n,cell.text,vals)
                vals[fields[idx]]=cell.text
    if not name:continue
    name=unicodedata.normalize('NFKC',' '.join(name))
    row={'n':n,'name':name,**vals};source_rows.append(row)
    groups.append({k:v for k,v in row.items() if k in ['n','name','tl','biomass','pb','qb','ee','pq']})
groups.sort(key=lambda x:x['n']);source_rows.sort(key=lambda x:x['n'])
assert [g['n'] for g in groups]==list(range(1,26)),groups
assert groups[9]['biomass']=='0.41' and groups[23]['qb']=='8.43'
book=next(SRC.glob('*.xlsx'));wb=openpyxl.load_workbook(book,data_only=False)
names={g['n']:g['name'] for g in groups}
tax=[];members=[]
for n,source_name,descr in list(wb['Table S2'].values)[2:27]:
    descr=(descr or '').strip()
    if descr=='-':descr='not documented - Table S2 provides no species list for this basal group.'
    tax.append({'seq':n,'group_name':names[n],'taxon_descr':descr})
    for part in re.findall(r'\(([^)]+)\)',descr):
        for member in part.split(','):
            member=member.strip()
            if not member:continue
            members.append({'printed_name':member,'accepted_name':member,'group_name':names[n],'source_page':f'Table S2 row {n+2}; spelling retained, no nomenclatural validation'})
diet={str(n):{} for n in range(2,25)}
ds=wb['Table S3'];consumer_ids=[ds.cell(2,c).value for c in range(3,26)]
assert consumer_ids==list(range(2,25))
for row in range(3,28):
    n=ds.cell(row,1).value;assert n==row-2
    for c,pred in enumerate(consumer_ids,3):
        v=ds.cell(row,c).value
        if v is not None:
            assert isinstance(v,(int,float)),(row,c,v)
            diet[str(pred)][str(n)]=v
fleets=['Herring fisheries','Whitefish fisheries','Creels'];landings={}
ls=wb['Table S1']
for row in range(4,29):
    n=ls.cell(row,1).value;assert n==row-3
    landings[str(n)]={f:ls.cell(row,c).value for f,c in zip(fleets,[7,8,9]) if ls.cell(row,c).value is not None}
model={'metadata':{'LME':'22','model_number':20251990,'model_name':'East Coast of Scotland','model_year':'1991-1995'},'groups':groups,'consumers':list(range(2,25)),'fleets':fleets,'landings':landings,'discards':{},'detritus_groups':['Detritus'],'detritus_fate':{},'diet':diet,'diet_rows':25,'landings_rows':25,'discards_rows':25}
MODEL_DIR.mkdir(parents=True,exist_ok=True)
(MODEL_DIR/'model.json').write_text(json.dumps(model,indent=2),encoding='utf-8')
(BASE/'source_table1.json').write_text(json.dumps(source_rows,indent=2),encoding='utf-8')
def write_csv(p,rows,fields):
    with p.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fields);w.writeheader();w.writerows(rows)
write_csv(MODEL_DIR/'taxonomy.csv',tax,['seq','group_name','taxon_descr'])
write_csv(BASE/(STEM+'.members.csv'),members,['printed_name','accepted_name','group_name','source_page'])
provenance={'paper':'Saygu et al. (2025), Historical ecosystem models can serve as a baseline for indicator-based assessment: the North Sea','doi':'10.3389/fmars.2025.1646031','source_pdf':str(pdf.relative_to(ROOT)),'source_workbook':str(book.relative_to(ROOT)),'local_model_number':'20251990 is an extraction-local identifier (publication 2025, 1990s model); no EcoBase ID asserted','basic_input':'PDF page 6 Table 1: Both TL and all right-hand 1990s subcolumns. Excludes E exploitation rate, which is not EE or migration. Names use Table 1 with Unicode ligatures expanded.','diet':'Table S3 C3:Y27, prey rows and predator columns; all explicit cells retained; no trace entries.','landings':'Table S1 G4:I28, three fleets; workbook J total is rounded independently and is not used to alter the fleet values.','taxonomy':'Table S2 A3:C27, Main Species column copied; supplementary labels aligned to Table 1 by checked group IDs','not_reported':['habitat_area','unassimilated_consumption','numeric_biomass_accumulation','detritus_fate','detritus_import','diet_import'],'discard':'Paper p.14 explicitly excludes discards for both models; extraction left blank rather than treating observed discards as known zero','rounding':'Table 1 has rounded P/Q 0.00 for seals/cetaceans and 0.01 for seabirds; preserved in Basic_input; database converter derives GE from PB/QB.'}
(BASE/'SOURCE_INVENTORY.json').write_text(json.dumps(provenance,indent=2),encoding='utf-8')
for command in [ ['write_outputs.py',str(MODEL_DIR/'model.json'),'--outdir',str(BASE/'extraction')], ['write_taxonomy.py',str(MODEL_DIR/'taxonomy.csv'),str(MODEL_DIR/'Taxonomy.xlsx')]]:
    subprocess.run([sys.executable,'-X','utf8',str(SKILL/'scripts'/command[0]),*command[1:]],check=True)
print('GROUPS',len(groups),'DIET_CELLS',sum(map(len,diet.values())),'DIET_SUMS',min(sum(d.values()) for d in diet.values()),max(sum(d.values()) for d in diet.values()))
print('MODEL_DIR',MODEL_DIR)
