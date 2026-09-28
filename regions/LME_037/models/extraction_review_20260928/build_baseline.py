from pathlib import Path
from decimal import Decimal
import json, csv, subprocess, sys, shutil, re, hashlib
from lxml import etree
from docx import Document
from openpyxl import Workbook
import pymupdf

ROOT=Path(__file__).resolve().parents[4]
R=Path(__file__).resolve().parent
P=ROOT/'regions/LME_037/papers/LME037-Bacalso-2026'
S=P/'recovered_sources_20260928'
MID='Bacalso2026_Visayan_Sea_1997_baseline'
MD=ROOT/'regions/LME_037/models'/MID
MD.mkdir(exist_ok=True)
ED=MD/'extracted_tables'
SK=Path('C:/Users/idoca/.agents/skills/ecopath-extraction/scripts')
def save(p,obj):p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf8')
def clean(s):return ' '.join(s.split())

cross=json.loads((R/'BASELINE_PARAMETER_CROSSWALK.json').read_text())
names={x['seq']:x['name'] for x in json.loads((R/'Bacalso2026_table1_coordinate_rows.json').read_text())}
groups=[]
for seq in range(1,34):
    g={'n':seq,'name':names[seq].replace(',',' -'),'hab_area':'1'}
    # B in the 2026 output table already represents the full model area.
    # Habitat fraction one encodes this density basis, not an ecological claim.
    for v in [x for x in cross if x['group']==seq]:g[v['field']]=v['2026_table1_1997'] or None
    groups.append(g)
assert len(groups)==33 and all(g['biomass'] for g in groups)
doc=Document(S/'Bacalso2023_DataSheet_1.docx')
diet={};evidence=[]
for ti in (2,3):
    t=doc.tables[ti]
    for ci in range(2,len(t.columns)):
        predator=t.rows[0].cells[ci].text
        diet[predator]={}
        for ri,row in enumerate(t.rows[1:],2):
            prey=row.cells[0].text;raw=row.cells[ci].text
            value='0' if raw=='--' else raw
            Decimal(value)
            diet[predator][prey]=value
            evidence.append({'source':'Bacalso2023_DataSheet_1.docx','table':'S2','docx_table':ti+1,'row':ri,'column':ci+1,'prey':prey,'predator':predator,'source_text':raw,'value':value,'interpretation':'-- denotes absent diet link; no normalization'})
landings={};fleet_cells=[]
for ri,row in enumerate(doc.tables[6].rows[1:-1],2):
    seq=row.cells[0].text;raw=row.cells[-1].text
    assert seq.isdigit()
    value='0' if raw=='--' else (None if '*' in raw else raw)
    landings[seq]={'Published total catch':value}
    evidence.append({'source':'Bacalso2023_DataSheet_1.docx','table':'S3','docx_table':7,'row':ri,'column':13,'group':seq,'source_text':raw,'value':value,'interpretation':'published group total; not sum of censored fleet cells'})
for ti in (4,5,6):
    t=doc.tables[ti]
    for ri,row in enumerate(t.rows[1:-1],2):
        for ci in range(2,len(t.columns)-(1 if ti==6 else 0)):
            fleet_cells.append({'group':int(row.cells[0].text),'fleet':t.rows[0].cells[ci].text,'source_text':row.cells[ci].text,'docx_table':ti+1,'row':ri,'column':ci+1})
assert len(fleet_cells)==33*34
model={'metadata':{'LME':'37','model_number':3702026,'model_name':'Visayan Sea Bacalso baseline','model_year':1997},'groups':groups,'consumers':list(range(1,31)),'fleets':['Published total catch'],'landings':landings,'discards':{},'detritus_groups':['Detritus'],'detritus_fate':{},'diet':diet,'diet_rows':33,'landings_rows':33,'discards_rows':33}
save(R/'extraction_input.json',model);save(R/'cell_evidence.json',evidence);save(R/'fleet_cells_source.json',fleet_cells)
proc=subprocess.run([sys.executable,str(SK/'write_outputs.py'),str(R/'extraction_input.json'),'--outdir',str(MD),'--dir-name','extracted_tables'],capture_output=True,text=True,encoding='utf8')
(R/'writer.log').write_text(proc.stdout+proc.stderr,encoding='utf8');assert proc.returncode==0,proc.stderr

# Expand publisher XML rowspan cells before associating representative taxa.
xml=etree.parse(str(S/'Bacalso2023_fmars_1099400.xml'))
grid=[];spans={}
for row in xml.xpath('//table-wrap[@id="T1"]//tr'):
    output=[];col=0
    for cell in row:
        while col in spans:
            text,left=spans[col];output.append(text)
            if left==1:del spans[col]
            else:spans[col]=(text,left-1)
            col+=1
        text=clean(' '.join(cell.itertext()));output.append(text)
        rowspan=int(cell.get('rowspan','1'))
        if rowspan>1:spans[col]=(text,rowspan-1)
        col+=1
    while col in spans:
        text,left=spans[col];output.append(text)
        if left==1:del spans[col]
        else:spans[col]=(text,left-1)
        col+=1
    grid.append(output)
assert len(grid)==34 and all(len(x)==5 for x in grid)
tax=[]
for g,row in zip(groups,grid[1:]):
    assert g['n']==int(row[0])
    desc='Representative taxa (Bacalso et al. 2023 Table 1; not an exhaustive list): '+row[3]
    if row[4]:desc+='; size range in catch: '+row[4]
    tax.append([g['n'],g['name'],desc])
wb=Workbook();ws=wb.active;ws.title='Taxonomy';ws.append(['seq','group_name','taxon_descr'])
for row in tax:ws.append(row)
wb.save(ED/'Taxonomy.xlsx')
with (ED/'taxonomy.csv').open('w',encoding='utf8',newline='') as f:
    w=csv.writer(f);w.writerow(['seq','group_name','taxon_descr']);w.writerows(tax)
save(ED/'taxonomy_evidence.json',{'source':'Bacalso2023_fmars_1099400.xml','table':'T1','expanded_rows':grid,'scope':'representative taxa, not exhaustive membership; 2026 S3 contains supplementary trait/catch records retained as source evidence, not automatically exhaustive memberships'})
for filename in ['cell_evidence.json','BASELINE_PARAMETER_CROSSWALK.json','fleet_cells_source.json']:
    shutil.copy2(R/filename,ED/filename)
for script,args,log in [('validate.py',[str(ED)],'validate.log'),('massbalance_check.py',[str(ED)],'massbalance.log'),('database_json.py',['-d',str(ED),'--update-report'],'database_json.log')]:
    p=subprocess.run([sys.executable,str(SK/script),*args],capture_output=True,text=True,encoding='utf8')
    (R/log).write_text(p.stdout+p.stderr,encoding='utf8');print(script,p.returncode, p.stdout[-1000:])
dbfiles=list(ED.glob('37_3702026_*.json'));assert len(dbfiles)==1
converted=json.loads(dbfiles[0].read_text(encoding='utf8'))
save(R/'converter_raw.json',converted)
source=json.loads(json.dumps(converted));changes=[]
for sg,g in zip(source['group'],groups):
    old=json.loads(json.dumps(sg));n=g['n']
    for input_field,target_field in [('biomass','biomass'),('biomass','biomass_habitat_area'),('pb','pb'),('qb','qb'),('ee','ee'),('pq','ge'),('tl','tl'),('unassim','gs'),('ba','biomass_accum'),('ba_rate','biomass_accum_rate')]:
        sg[target_field]=g.get(input_field) if g.get(input_field) is not None else '-9999'
    sg['ge_input']='true' if g.get('pq') is not None else 'false'
    sg['diet_imp']='-9999'
    entries=[]
    for prey in range(1,34):
        value=diet.get(str(n),{}).get(str(prey))
        if value is not None or prey==33:entries.append({'prey_seq':str(prey),'proportion':value if value is not None else '-9999','detritus_fate':'-9999'})
    sg['diet_descr']={'diet':entries}
    for key in sg:
        if old.get(key)!=sg[key]:changes.append({'group':n,'field':key,'converter_value':old.get(key),'source_faithful_value':sg[key]})
save(MD/'model.json',source);save(MD/dbfiles[0].name,source)
save(R/'converter_source_restoration.json',changes)
save(ED/'canonical_source_database.json',source)
save(dbfiles[0],source)
p=subprocess.run([sys.executable,str(SK/'database_json.py'),'-j',str(dbfiles[0])],capture_output=True,text=True,encoding='utf8')
(R/'canonical_roundtrip.log').write_text(p.stdout+p.stderr,encoding='utf8');assert p.returncode==0,p.stderr
save(R/'model_paths.json',{'model_id':MID,'model_dir':str(MD),'tables':str(ED),'canonical_filename':dbfiles[0].name})
# Prose sweep retained with page references; numeric assumptions remain unknown.
hits=[]
for pdf in [P/'Bacalso-c0d09ef6.pdf',P/'1-s2.0-S0964569125004909-mmc1-f97c887e.pdf',S/'Bacalso2023_fmars_1099400.pdf']:
    for i,page in enumerate(pymupdf.open(pdf),1):
        text=page.get_text()
        for key in ['accumul','assimil','steady','migration','discard','habitat','km2','km²']:
            for m in re.finditer(re.escape(key),text,re.I):hits.append({'file':pdf.name,'page':i,'term':key,'context':clean(text[max(0,m.start()-180):m.start()+400])})
save(R/'prose_sweep.json',hits)
print('MODEL',MD)
