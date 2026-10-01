from pathlib import Path
import copy,json,hashlib,zipfile,urllib.parse,re
from lxml import etree
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
Q=Path(__file__).parent;ROOT=Q.parents[4];R=ROOT/'regions/LME_038';MID='38_38003_Java_Sea_normalized_BA_completed_(mid1970s)';E=R/'validation_reports'/MID;REF=ROOT/'tools/templates/Model_validation_template.docx';OUT=R/f'Model_validation_{MID}.docx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
a=json.loads((E/'coverage_summary.json').read_text(encoding='utf-8'));diag=json.loads((E/'matrix_inspection.json').read_text(encoding='utf-8'));geo=json.loads((E/'geography/geographic_assessment.json').read_text(encoding='utf-8'));exposure=json.loads((E/'method_specific_mapping_exposure.json').read_text(encoding='utf-8'))
assert sha(REF)=='d3aca1a596bf9476789aa919fbbc4b0f578fed7a0eb625d3d290b8db66c5b1f0'
d=Document(REF);p=d.paragraphs;manual=[copy.deepcopy(d.tables[t].rows[r]._tr)for t,r in [(0,9),(1,6),(1,7)]]
def clear(par):
 for n in list(par._p):
  if n.tag!=qn('w:pPr'):par._p.remove(n)
 return par
def text(par,value):clear(par);par.add_run(value);return par
def cell(c,value):
 first=c.paragraphs[0]
 for pp in list(c._tc):
  if pp.tag==qn('w:p')and pp is not first._p:c._tc.remove(pp)
 return text(first,value)
def link(par,label,target):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),par.part.relate_to(urllib.parse.quote(target,safe='/:?=&%#'),RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'0563C1');rp.append(co);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def addlinks(par,links):
 for i,(label,target)in enumerate(links):par.add_run(' | 'if i else' ');link(par,label,target)
def put(t,r,v,links=[]):par=cell(d.tables[t].cell(r,1),v);addlinks(par,links)
def pct(v):return f'{v:.4f}%'if v>=.00005 or v==0 else'<0.0001%'
ev='validation_reports/'+MID+'/';pdf='papers/BUCHARY-1991/ubc_1999-0254.pdf.pdf'
text(p[0],'Java Sea regional model validation')
text(p[1],'Pending alignment draft')
put(0,1,'Indonesian Sea, LME_038. The source Java Sea domain is subregional.', [('Regional workbook','LME_038.xlsx')])
put(0,2,'Sea Around Us v50.1,1950–2019 original catch. Reference:2019 landings; all181 exact labels, including zero catch and residual categories.', [('Original catch','raw/LME_038-catch.zip')])
put(0,3,'Buchary,E.A.(1999). Evaluating the effect of the1980 trawl ban in the Java Sea, Indonesia: An ecosystem-based approach. MSc thesis, University of British Columbia. Local main PDF verified; inherited BUCHARY-1991 folder is an alias.', [('Original thesis',pdf),('Source review',ev+'source_review.md')])
put(0,4,'Zainuri & Endrawati(1999):115ha Awur Bay, a local12-component model. Nurhakim(2003):27-group1979 north Central Java survey model; printed diet defects prevent strict admission. Both are subregional alternatives.', [('Awur Bay source','papers/LME038-Zainuri-1999/4895-c62b530d.pdf'),('North Central Java source','papers/LME038-Nurhakim-2003/content-ce43630c.pdf')])
put(0,5,MID+'. Static mid1970s (1974–1976 catch context);28source groups plus one computational import. Reported Java Sea area471000km².', [('Accepted selected model',f'models/{MID}/model.json'),('Variant provenance',f'models/{MID}/SELECTED_MODEL_REPORT.md')])
put(0,6,'The thesis contains one documented mid1970s static baseline and Ecosim trawl-ban simulations. A second complete, independently verified static parameter/diet model was not identified; dynamic scenarios are not separate static baselines.')
put(0,7,'Article — ?\nModel — User-selected best available Java Sea derived variant. Full-LME and multi-decadal transfer remain provisional; source fidelity, warnings and applicability are assessed separately.')
put(0,8,'Source Tables3.1,3.9,3.10 and parameter prose reviewed. Accepted Macrozoobenthos diet0.660→1 proportional normalization and28 signed computational BA completions are retained. BA is not measured biomass change. Canonical bytes, loaded biological inputs and source catches remain unchanged.', [('Source/derived differences',ev+'source_review.md'),('Loaded-state proof',ev+'matrix_inspection.json')])
for row,opt in [(1,'GE'),(2,'TE')]:
 m=next(x for x in diag['methods']if x['method']==opt)
 value=f"No negative SPPR entries across the basal-source columns. {m['grade']} retained. rho_living={m['rho_living']:.8f}; b={m['b']:.8f}; Detritus(28) SPPR={m['detritus_sppr']['28']:.8f}."
 if opt=='TE':value+=' Marine mammals have near-zero transfer efficiency; finite results remain provisional.'
 put(1,row,value)
put(1,3,'A. Target fraction covered by study: approximately19–22%.\nB. Study fraction within target: approximately99–100%.\nLow geographic confidence: approximate prose-bound reconstruction with two eastern scenarios; no explicit source study polygon.')
put(1,4,'Static mid1970s coefficients and fixed source-catch/donor weights are applied across1950–2019 whole-LME landings, catch and discards. Reference2019 is about44years later. The source guild populations and historical trawl bycatch differ from modern recorded fisheries.')
put(1,5,'With Egestion is OK under accepted thresholds. All348 entries in the three direct29×4source matrices are finite/nonnegative, including unfished groups. GE/Egestion threshold grades do not establish ecological validity. Historical method-specific failures/warnings remain. Mixed bony/cartilage source pools and assumed composition limit coarse labels.', [('Full matrices and warnings',ev+'matrix_inspection.json')])
text(p[3],'Reference:2019 landings; all exact labels, including zero catch and residual categories; no group filter.')
text(p[4],'28source groups(+1computational import);181catch taxa.')
text(p[5],f"Simple-chain PPR totals {a['total_simple_chain_ppr_tC']:,.2f} t C in2019.")
text(p[6],'39labels lack a classic coefficient; all have zero2019 catch and zero annual contribution. Unavailable TL/coefficient stays ?, with no invented value.')
text(p[7],'Method: protected independent simple trophic-chain coefficients; wet coefficient×catch/9 once.')
clear(p[8]);link(p[8],'All 181 taxon mappings','LME038_taxon_mapping_appendix.xlsx');p[8].add_run(' and ');link(p[8],'descriptive Sources sheet','LME038_taxon_mapping_appendix.xlsx#Sources!A1')
for row,s in zip(d.tables[2].rows[1:],a['confidence_summary']):
 for c,v in zip(row.cells,[s['label'],str(s['taxa']),pct(s['catch_percentage']),pct(s['ppr_percentage'])]):cell(c,v)
desc={'M2':'Verified historical name or spelling bridge.','M3':'Documented source taxon/family criterion supports the particular taxon.','M5':'Supported ecological extension to represented source members.','M6':'Partial or competing source ecological/size fit.','M9':'Assumed complete source-containing family or stage set.','M10':'Broad reporting category approximated across compatible pools.','M11':'Closest represented analogue with material habitat/taxonomic mismatches.','W1':'One reviewed source group receives100%; no numerical split.','W4':'Complete printed source Harvest proportions, including genuine zeros.','W5':'Historical identified-catch donor proportions; population/basis mismatch.'}
for ti,key in [(3,'membership_summary'),(4,'allocation_summary')]:
 table=d.tables[ti];proto=copy.deepcopy(table.rows[1]._tr)
 for row in list(table.rows)[1:]:table._tbl.remove(row._tr)
 for s in a[key]:
  table._tbl.append(copy.deepcopy(proto))
  for c,v in zip(table.rows[-1].cells,[desc[s['rule']],s['confidence'],pct(s['ppr_percentage'])]):cell(c,v)
clear(p[11]);p[11].add_run('Membership evidence: ');link(p[11],'source Table3.1',pdf+'#page=58');p[11].add_run(', ');link(p[11],'habitat/size rule',pdf+'#page=55');p[11].add_run(', ');link(p[11],'native FAO support',ev+'supporting_sources_manifest.json');p[11].add_run(' and ');link(p[11],'all decisions',ev+'taxon_audit.json')
text(p[13],'Percentages are each rule’s share of simple-chain PPR. Complete Harvest includes genuine zero stages. Juvenile shrimp0.008 and juvenile large demersals0.066 are author assumptions. Living bottom structure catch is10×SAF assumed trawl bycatch and supplies71.43% of the generic-invertebrate allocation. Retained donor fractions use all eligible identified direct/stage donors, with no same-family restriction; mismatch remains Low. Compatible measured stage-mass partitions were not located. ');link(p[13],'Allocation evidence',ev+'identified_catch_proxy_review.json')
vl=next(x for x in a['confidence_summary']if x['label']=='Very low');names={'new_GE':'GE','new_TE_EEfix':'TE','new_WithEgestion':'With Egestion'};shares='; '.join(f"{names[x['method']]} {x['very_low_percentage']:.2f}%"for x in exposure)
text(p[14],f"Unresolved: none. Very low: {vl['taxa']}labels contribute {vl['catch_percentage']:.4f}% of landings and {vl['ppr_percentage']:.4f}% of simple-chain PPR. Broad bony fish/invertebrate pools and unlisted billfish, sharks, tarpon, catfish, flatfish and deep-water taxa use explicit approximations. Mixed-source pools retain non-target members. Very low all-source model PPR shares: {shares}. Broad marine fish alone supplies30.07%of TE; Lates calcarifer analogy18.58%. ");link(p[14],'Method-specific exposure',ev+'method_specific_mapping_exposure.json')
text(p[16],'R is the full Sea Around Us Indonesian Sea LME. S reconstructs the thesis’s approximate3°S–6°50′S,105°50′E–116°30′/117°30′E prose bounds, removing land consistently. Figure2.1 supplies bathymetry and provincial context; it has no explicit study polygon. A/B are approximate regional-extent estimates.')
for ti,title,img in [(5,'Indonesian Sea target and two source-bound scenarios',E/'geography/target_region.png'),(6,'Original source Figure2.1: Java Sea geography and bathymetry',E/'geography/source_figure2_1_landscape.png')]:
 cell(d.tables[ti].cell(0,0),title);cell(d.tables[ti].cell(1,0),'').add_run().add_picture(str(img),width=Inches(5.0))
text(p[18],f"Target R:{geo['target_area_km2']:,.0f}km². Reconstructed water areas:441114/488277km²; intersections:439052/486216km². A≈19–22%,B≈99–100%. Source-reported471000km² is separate context. All seven stated water/land controls pass. ");addlinks(p[18],[('Target polygon',ev+'geography/target_region.geojson'),('Approximate bounds',ev+'geography/reconstructed_study_bounds.geojson')])
text(p[20],'Original PDF30/printed21, rotated for readability. Its graticule and bordering provinces inform the reconstruction; source latitude bounds can truncate local coast-following water. Natural Earth1:50 million land removal is generalized. Two eastern bounds express source ambiguity, not statistical confidence intervals. No geographic multiplier is applied. ');addlinks(p[20],[('Original figure',pdf+'#page=30'),('Areas and limits',ev+'geography/geographic_assessment.json')])
text(p[21],'Original Table3.1, PDF58/printed49: source biological groups and representative taxa. The family summary does not exhaust modern taxonomic membership.');p[21].add_run().add_picture(str(E/'geography/original_pdf_page_58.png'),width=Inches(5.0))
for par in [*d.paragraphs,*[pp for table in d.tables for row in table.rows for c in row.cells for pp in c.paragraphs]]:
 for run in par.runs:
  if run._r.findall('.//'+qn('w:drawing')):continue
  value=run.text.replace(MID,'\ue000');value=re.sub(r'(?<=[A-Za-z])(?=\d)', ' ',value);value=re.sub(r'(?<=\d)(?=[A-Za-z])',' ',value);value=re.sub(r'([,:;])(?=[A-Za-z])',r'\1 ',value);value=re.sub(r'(\d{4}) s\b',r'\1s',value);run.text=value.replace('\ue000',MID)
for table in d.tables:
 table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in table.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
for idx in [2,12,15,19,21]:p[idx].paragraph_format.page_break_before=True
for idx in [17,19]:p[idx].paragraph_format.keep_with_next=True
for t,r,original in [(0,9,manual[0]),(1,6,manual[1]),(1,7,manual[2])]:d.tables[t].rows[r]._tr.getparent().replace(d.tables[t].rows[r]._tr,original)
temp=Q/'generated.docx';d.save(temp);changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with zipfile.ZipFile(REF)as ref,zipfile.ZipFile(temp)as gen,zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED)as final:
 for item in ref.infolist():final.writestr(item,gen.read(item.filename)if item.filename in changed else ref.read(item.filename))
 for item in gen.infolist():
  if item.filename not in ref.namelist():final.writestr(item,gen.read(item.filename))
check=Document(OUT);assert[etree.tostring(check.tables[t].rows[r]._tr)for t,r in [(0,9),(1,6),(1,7)]]==[etree.tostring(x)for x in manual];assert etree.tostring(check.sections[0]._sectPr)==etree.tostring(Document(REF).sections[0]._sectPr)
with zipfile.ZipFile(REF)as ref,zipfile.ZipFile(OUT)as final:
 preserved=[n for n in ref.namelist()if n not in changed];assert all(ref.read(n)==final.read(n)for n in preserved)
(E/'document_fidelity.json').write_text(json.dumps(dict(manual_rows_identical=True,section_geometry_identical=True,preserved_reference_package_parts=preserved,template_sha256=sha(REF)),indent=2),encoding='utf-8')
print('DOCX',OUT)
