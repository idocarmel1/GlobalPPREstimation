from pathlib import Path
import copy,json,hashlib,zipfile,urllib.parse,re
from lxml import etree
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
Q=Path(__file__).parent;ROOT=Q.parents[4];R=ROOT/'regions/LME_035';MID='35_412_Gulf_of_Thailande_(1963)';E=R/'validation_reports'/MID;REF=ROOT/'tools/templates/Model_validation_template.docx';OUT=R/f'Model_validation_{MID}.docx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
a=json.loads((E/'coverage_summary.json').read_text(encoding='utf-8'));diag=json.loads((E/'matrix_inspection.json').read_text(encoding='utf-8'));geo=json.loads((E/'geography/geographic_assessment.json').read_text(encoding='utf-8'))
assert sha(REF)=='d3aca1a596bf9476789aa919fbbc4b0f578fed7a0eb625d3d290b8db66c5b1f0'
d=Document(REF);p=d.paragraphs;manual=[copy.deepcopy(d.tables[t].rows[r]._tr)for t,r in [(0,9),(1,6),(1,7)]]
def clear(par):
 for n in list(par._p):
  if n.tag!=qn('w:pPr'):par._p.remove(n)
 return par
def text(par,value):clear(par);par.add_run(value);return par
def cell(c,value):
 first=c.paragraphs[0]
 for pp in list(c._tc):
  if pp.tag==qn('w:p')and pp is not first._p:c._tc.remove(pp)
 return text(first,value)
def link(par,label,target):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),par.part.relate_to(urllib.parse.quote(target,safe='/:?=&%#'),RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
 co=OxmlElement('w:color');co.set(qn('w:val'),'0563C1');rp.append(co);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def addlinks(par,links):
 for i,(label,target)in enumerate(links):par.add_run(' | 'if i else' ');link(par,label,target)
def put(t,r,v,links=[]):par=cell(d.tables[t].cell(r,1),v);addlinks(par,links)
def pct(v):return f'{v:.4f}%'if v>=.00005 or v==0 else'<0.0001%'

ev='validation_reports/'+MID+'/';pdf='papers/Christensen-1998/Pauly-Chuenpagdee-2003.pdf';pred=ev+'independent_review/Pauly_Christensen_1993.pdf';html='https://www.academia.edu/1483062/Fishery_induced_changes_in_a_marine_ecosystem_insight_from_models_of_the_Gulf_of_Thailand'
put(0,1,'Gulf of Thailand, LME_035; source10–50m shelf applied across the whole LME.', [('Regional workbook','LME_035.xlsx')])
put(0,2,'Sea Around Us v50.1, local1950–2019 catch; reference2019 landings. All247 exact taxon labels, including zero catch and unidentified categories.', [('Original catch','raw/LME_035-catch.zip')])
put(0,3,'Christensen,V.(1998). Fishery-induced changes in a marine ecosystem: insight from models of the Gulf of Thailand. Journal of Fish Biology53(Suppl.A):128–142. Author full HTML read; original PDF unavailable.', [('DOI','https://doi.org/10.1111/j.1095-8649.1998.tb01023.x'),('Author full text',html),('Retained source facts','papers/Christensen-1998/source-facts.json')])
put(0,4,'Pauly&Christensen(1993), Stratified models of large marine ecosystems: application to the South China Sea — acknowledged10–50m Gulf predecessor. Vibunpant et al.(2003), Trophic model of the coastal fisheries ecosystem in the Gulf of Thailand — separate40-group1973 model. Premcharoen(2012), Mae Klong estuary — subregional alternative.', [('1993 predecessor',pred),('1973 source','papers/GOT-2003'),('Estuary source','papers/LME035-Premcharoen-2012')])
put(0,5,MID+'; scientific payload1980 despite inherited1963 filename. EcoBase412 description and all11TableII biomass comparisons identify1980.29source ecological groups(+1computational import);10–50m shelf.', [('Canonical model',f'models/{MID}/model.json'),('Identity evidence',ev+'source_fidelity_review.json')])
put(0,6,'Article compares1963 and1980 baselines. Selected canonical file contains1980; a separate verified complete1963 parameter/diet set is not supplied.1993 predecessor and1973 candidate have different group structures.')
put(0,7,'Article — ?\nModel — Retained selected EcoBase412 payload; independent identity evidence establishes1980. Whole-LME/year application remains provisional; scientific approval is not established by finite coefficients.')
put(0,8,'Available source identity, habitat, illustrative membership and11biomasses were checked. Complete published B/PB/QB/EE/diet inputs are unavailable. Accepted canonical export, BA, diet, routing and runtime settings are retained; exact loaded groups agree. Zero juvenile exports are canonical values, not observed regional stage catches.', [('Source limitations',ev+'source_review.md'),('Loaded proof',ev+'matrix_inspection.json')])
for row,option in [(1,'GE'),(2,'TE')]:
 m=next(x for x in diag['methods']if x['method']==option)
 value=f"No negative source SPPR entries; WARN retained. rho_living={m['rho_living']:.8f}; "
 if option=='GE':value+=f"b={m['b']:.8f}; "
 value+=f"Detritus(29) SPPR={m['detritus_sppr']['29']:.8f}."
 put(1,row,value)
put(1,3,'A. Focal10–50m study fraction: unavailable from the accessible source geometry.\nB. Focal study fraction within LME: unavailable.\nContext only: bounded0–50m lineage trace A≈26.5%,B≈91%; reported larger shallowband/target≈38.7%. These do not establish focal coverage.')
put(1,4,'Static1980 coefficients and fixed composition/stage proxies apply across1950–2019 whole-LME catches; reference2019 is39years later. Source shelf and reporting-population differences are not reconstructed annual ecosystem dynamics.')
put(1,5,'GE,TE and With Egestion remain WARN. TE repairs consumed-production routing for twoEE0 groups (marine mammals,jellyfish) under accepted runtime settings. All270direct source entries are finite/nonnegative, including unfished groups; six symbolic methods remain failed. Source fidelity and spatial/temporal application limits remain.', [('Full diagnostics',ev+'matrix_inspection.json')])
text(p[3],'Reference:2019 landings; all exact taxon labels, including zero catch and unidentified categories; no group filter.')
text(p[4],'29source ecological groups(+1computational import);247catch taxa.')
text(p[5],f"Simple-chain PPR totals {a['total_simple_chain_ppr_tC']:,.2f} t C in2019.")
text(p[6],'33labels lack an assigned classic coefficient; all have zero2019 landings and zero annual PPR. Missing TL/coefficient stays ?, without an invented value.')
text(p[7],'Method: independent simple trophic chain using protected taxon coefficients; carbon conversion/9 once.')
clear(p[8]);link(p[8],'All taxon mappings','LME035_taxon_mapping_appendix.xlsx');p[8].add_run(' and ');link(p[8],'descriptive Sources sheet','LME035_taxon_mapping_appendix.xlsx#Sources!A1')
for row,s in zip(d.tables[2].rows[1:],a['confidence_summary']):
 for c,v in zip(row.cells,[s['label'],str(s['taxa']),pct(s['catch_percentage']),pct(s['ppr_percentage'])]):cell(c,v)
desc={'M1':'Explicit focal source taxon or genus assignment.','M2':'Verified source spelling or historical-name bridge.','M3':'Unambiguous documented generic taxonomic group fit.','M4':'Related-member extension from source representatives.','M5':'Supported coastal/guild or historical ecological extension.','M6':'Partial or competing ecological/taxonomic fit.','M9':'Assumed complete eligible family or suborder set.','M10':'Broad reporting category approximated over compatible pools.','M11':'Closest represented analogue with material source mismatches.','W1':'One reviewed group receives100%; no split.','W4':'Complete1980 canonical export proportions, including zero candidates.','W5':'Retained historical identified-catch proxy with population mismatch.'}
for ti,key in [(3,'membership_summary'),(4,'allocation_summary')]:
 table=d.tables[ti];proto=copy.deepcopy(table.rows[1]._tr)
 for row in list(table.rows)[1:]:table._tbl.remove(row._tr)
 for s in a[key]:
  table._tbl.append(copy.deepcopy(proto))
  for c,v in zip(table.rows[-1].cells,[desc[s['rule']],s['confidence'],pct(s['ppr_percentage'])]):cell(c,v)
clear(p[11]);p[11].add_run('Membership evidence: ');link(p[11],'focalp130 and identity','papers/Christensen-1998/source-facts.json');p[11].add_run(', ');link(p[11],'predecessor Table2',pred+'#page=19');p[11].add_run(' and ');link(p[11],'all taxon decisions',ev+'taxon_audit.json')
text(p[13],'Percentages are simple-chain PPR shares of taxa using each rule. Accepted48stage fractions remain unchanged. Revised broad candidate sets use complete model export proportions; historical identified-catch fractions with unchanged candidates remain at Low because donor country/gear/composition differ. No compatible observed stage-mass partitions were located. ');link(p[13],'Weights and donor evidence',ev+'candidate_and_allocation_evidence.json')
vl=next(x for x in a['confidence_summary']if x['label']=='Very low')
text(p[14],f"Unresolved: none. Very low: {vl['taxa']}labels contribute {vl['catch_percentage']:.4f}% of landings and {vl['ppr_percentage']:.4f}% of simple-chain PPR. Broad fish/invertebrate residual composition, historical serranid and carangid concepts, coastal herbivore/trashfish analogues and nine billfish rely on explicit assumptions. Billfish use the acknowledged predecessor large-pelagic connection to focal Tuna, with narrowed coastal taxonomy and depth limits. All affected taxa appear in the linked appendix. Model-method shares differ: Very low contributes 14.12% of GE, 12.01% of TE and 13.86% of With Egestion all-source PPR; Carangidae alone contributes 7.93% of GE. ");link(p[14],'Method-specific exposure',ev+'method_specific_mapping_exposure.json')
text(p[16],'R is the Sea Around Us target. The accessible focal article establishes a10–50m shelf domain but provides no usable boundary here. Pauly&Chuenpagdee2003 Figure14-1 gives geographic lineage context with a50m contour; its bounded trace is illustrative and does not substitute for focal studyS.')
for ti,title,img in [(5,'Gulf of Thailand target and shelf context',E/'geography/target_region.png'),(6,'Illustrative depth context from lineage map',E/'geography/study_domain.png')]:
 cell(d.tables[ti].cell(0,0),title);cell(d.tables[ti].cell(1,0),'').add_run().add_picture(str(img),width=Inches(5.0))
text(p[18],f"Target R:{geo['target_area_km2']:,.0f} km². Blue shows the target; orange is a bounded contextual shelf trace within the visible angular lineage boundary. Coastline/islands are simplified; all eight stated water/land controls pass. ");addlinks(p[18],[('Target polygon',ev+'geography/target_region.geojson'),('Context trace',ev+'geography/contextual_0_50m_shelf_trace.geojson')])
text(p[20],f"Contextual trace:{geo['contextual_trace_area_km2']:,.0f} km²; intersection:{geo['intersection_km2']:,.0f} km²; contextual A≈26.47%,B≈90.97%. The text's larger0–50m band≈150,000 km²/target≈38.74% is a separate area context, not a second measured polygon. Neither supplies the absent10m coastal boundary or full focal study extent. No spatial multiplier is applied. ");addlinks(p[20],[('Lineage Figure14-1',pdf+'#page=2'),('Calibration and limits',ev+'geography/geographic_assessment.json')])
text(p[21],'Original lineage Figure14-1 (printed338/PDF2). Its50m contour and angular jurisdiction/location line are geography context. Figure14-3 is an ecosystem diagram; no focal original-study figure is falsely claimed.');p[21].add_run().add_picture(str(E/'geography/source_figure14_1_cropped.png'),width=Inches(5.0))
# Expand compressed historical evidence labels into ordinary readable prose without changing template styling.
for par in [*d.paragraphs,*[pp for table in d.tables for row in table.rows for c in row.cells for pp in c.paragraphs]]:
 for run in par.runs:
  if run._r.findall('.//'+qn('w:drawing')):continue
  run.text=re.sub(r'(?<=[A-Za-z])(?=\d)', ' ',run.text);run.text=re.sub(r'(?<=\d)(?=[A-Za-z])',' ',run.text)
  run.text=re.sub(r'(?<=[A-Za-z])&(?=[A-Za-z])',' & ',run.text)
  run.text=re.sub(r'([,:;])(?=[A-Za-z])',r'\1 ',run.text)
  run.text=re.sub(r'(?<=[A-Za-z)])\.(?=[A-Za-z])','. ',run.text)

for table in d.tables:
 table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in table.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
for idx in [2,12,15,19,21]:p[idx].paragraph_format.page_break_before=True
for idx in [17,19]:p[idx].paragraph_format.keep_with_next=True
for t,r,original in [(0,9,manual[0]),(1,6,manual[1]),(1,7,manual[2])]:d.tables[t].rows[r]._tr.getparent().replace(d.tables[t].rows[r]._tr,original)
temp=Q/'generated.docx';d.save(temp);changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with zipfile.ZipFile(REF)as ref,zipfile.ZipFile(temp)as gen,zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED)as final:
 for item in ref.infolist():final.writestr(item,gen.read(item.filename)if item.filename in changed else ref.read(item.filename))
 for item in gen.infolist():
  if item.filename not in ref.namelist():final.writestr(item,gen.read(item.filename))
check=Document(OUT);assert[etree.tostring(check.tables[t].rows[r]._tr)for t,r in [(0,9),(1,6),(1,7)]]==[etree.tostring(x)for x in manual];assert etree.tostring(check.sections[0]._sectPr)==etree.tostring(Document(REF).sections[0]._sectPr)
with zipfile.ZipFile(REF)as ref,zipfile.ZipFile(OUT)as final:
 preserved=[n for n in ref.namelist()if n not in changed];assert all(ref.read(n)==final.read(n)for n in preserved)
(E/'document_fidelity.json').write_text(json.dumps(dict(manual_rows_identical=True,section_geometry_identical=True,preserved_reference_package_parts=preserved,template_sha256=sha(REF)),indent=2),encoding='utf-8')
print('DOCX',OUT)
