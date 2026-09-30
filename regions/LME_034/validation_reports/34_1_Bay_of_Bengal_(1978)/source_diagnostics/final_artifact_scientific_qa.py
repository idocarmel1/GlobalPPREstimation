from pathlib import Path
import json, hashlib, math, re, sys, zipfile
from urllib.parse import urlsplit, unquote
from collections import Counter, defaultdict
from xml.etree import ElementTree as ET
import numpy as np
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding='utf-8')

out = Path(__file__).resolve().parent
evidence = out.parent
region = evidence.parent.parent
project = region.parent.parent
docpath = region/'Model_validation_34_1_Bay_of_Bengal_(1978).docx'
bookpath = region/'LME034_taxon_mapping_appendix.xlsx'
def read(path): return json.loads(path.read_text(encoding='utf8'))
def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def close(a,b,atol=1e-8,rtol=1e-12): return math.isclose(float(a),float(b),abs_tol=atol,rel_tol=rtol)
extracted = read(out/'final_artifact_extracted_content.json')
audit = read(evidence/'adopted_taxon_audit.json')
baseline = read(evidence/'baseline_evidence.json')
source = read(out/'source_audit.json')
direct = {r['option']:r for r in read(out/'direct_summary.json')}
context = read(evidence/'geography_sources/context_and_geography.json')
checks=[]
errors=[]
notes=[]
def check(name, ok, details):
    checks.append({'check':name,'status':'PASS' if ok else 'FAIL','details':details})
    if not ok: errors.append({'check':name,'details':details})
check('Exact inspected artifact identities',
    extracted['document_sha256']==sha(docpath) and extracted['appendix_sha256']==sha(bookpath),
    {'report_sha256':sha(docpath),'appendix_sha256':sha(bookpath)})
check('Canonical model remains identical to direct evidence',sha(region/'models'/audit['model_id']/'model.json')==source['model_sha256'],
      {'current_canonical_sha256':sha(region/'models'/audit['model_id']/'model.json'),'direct_evidence_sha256':source['model_sha256']})
fields={}
tables=[]
for b in extracted['document_blocks']:
    if b['type']=='table':
        tables.append(b['rows'])
        for r in b['rows']:
            if len(r)==2: fields[r[0]]=r[1]
text='\n'.join([b['text'] if b['type']=='paragraph' else '\n'.join(' | '.join(r) for r in b['rows']) for b in extracted['document_blocks']])
check('Automatic field inventory', all(k in fields for k in ['Region','Catch source','Selected article','Other known articles','Selected model','Other models from the same article','Selection rationale','Model extraction','GE diagnostics','TE diagnostics','Geographic fit','Temporal fit','Other']),
      {'fields_checked':['Region','Catch source','Selected article','Other known articles','Selected model','Other models from the same article','Selection rationale','Model extraction','GE diagnostics','TE diagnostics','Geographic fit','Temporal fit','Other'],
       'mapping_section_checked_separately':True})
check('Model, source period and inventory', 'LME_034' in fields['Region'] and audit['model_id'] in fields['Selected model'] and 'Baseline 1978' in fields['Selected model'] and '49 source ecological groups (+1 computational' in fields['Selected model'] and '6,205,051' in fields['Selected model'],
      {'model_id':audit['model_id'],'source_ecological_groups':49,'loaded_computational_groups':50,'source_area_km2':6205051,'source_identity':'December2013 local62-page report;2014 distribution distinguished'})
extraction=fields['Model extraction']
check('Published, canonical and derived reconstruction distinguished', all(s in extraction for s in ['Balanced Table 16','habitat fraction × habitat biomass','Published diet columns contain omissions','JSON normalizes nonzero','all-zero 3 Meiobenthos','77.3%','constructor derives biomass accumulation','sets detritus EE to 1','Derived BA is not an observation','fidelity remains unverified']),
      {'location':'Word Model extraction','source_closure_maximum':'canonical diet residual at3Meiobenthos is77.3012757% of production',
       'published_inputs':'Table16;Table17/A3.2 omissions','canonical_transforms':'nonzero-column normalization and all-zero Meiobenthos detritus diet','derived_loaded_state':'closure BA and forced detritus EE1',
       'full_default_ledger':'GE/transformation_ledger.json; GE/runtime_identity.json'})
notes.append({'type':'document_scope','finding':'The concise report does not enumerate source-consistent Table15 GS, default zero migration or default closed one-pool detritus routing. It makes no claim that these defaults were published; complete default/derived identities remain in the linked reconstruction evidence and direct ledger.'})
matrixchecks=[]
for method,folder in [('GE','GE'),('TE','TE'),('With Egestion','With_Egestion')]:
    m=np.load(out/folder/'SPPR.npy',allow_pickle=False)
    axes=read(out/folder/'SPPR_axes.json')
    negatives=int(np.count_nonzero(m<0))
    finite=int(np.count_nonzero(np.isfinite(m)))
    ok=m.shape==(50,6) and negatives==0 and finite==300 and axes['column_ids']==[50,49,45,44,43,42] and axes['row_ids'].count(49)==1
    item={'method':method,'shape':list(m.shape),'entries':int(m.size),'strict_negative_entries':negatives,'nonfinite_entries':int(m.size-finite),'row_ids':'all50exactrows','basal_column_ids':axes['column_ids'],'basal_column_names':axes['column_names']}
    matrixchecks.append(item)
    check(method+' complete matrix strict sign finding',ok,item)
for method in ['GE','TE']:
    f=fields[method+' diagnostics']
    expected=direct[method]
    rho=float(re.search(r'ρ_living\s*=\s*([0-9]+(?:\.[0-9]+)?)',f).group(1))
    det=float(re.search(r'Detritus \(49\) SPPR\s*=\s*([0-9]+(?:\.[0-9]+)?)',f).group(1))
    ok='No negative SPPR entries across the basal-source columns' in f and close(rho,expected['rho_living'],5e-8,0) and close(det,expected['detritus']['sppr'],5e-8,0)
    details={'location':'Word '+method+' diagnostics','displayed_rho':rho,'exact_rho':expected['rho_living'],'displayed_named_detritus_sppr':det,'exact_named_detritus_sppr':expected['detritus']['sppr'],'display_rounding_tolerance':5e-8,'returned_status':expected['status'],'strict_model_input_balance':expected['model_input_strict_balanced'],'strict_loaded_balance':expected['balance_strict_balanced']}
    if method=='GE':
        b=float(re.search(r'; b\s*=\s*([0-9]+(?:\.[0-9]+)?)',f).group(1));ok &= close(b,expected['b'],5e-8,0)
        details.update({'displayed_b':b,'exact_b':expected['b']})
    else: details.update({'TE_b_omitted':True,'no_claim_of_absent_physical_recycling':True})
    check(method+' diagnostic scalars and scientific interpretation',ok,details)
check('Provisional source-fidelity and transfer caveats', 'Model-derived estimates remain provisional' in fields['Other'] and 'Numerical reproducibility does not establish author-native balance' in fields['Other'] and 'pooled 1978 total-catch weights' in fields['Temporal fit'] and 'catch bases' in fields['Temporal fit'] and '1950–2019' in fields['Catch source'] and '2019 landings, excluding discards' in fields['Catch source'],
      {'Word_fields':['Other','Temporal fit','Catch source'],'current_reconstruction_source_faithful':False,'native_author_model_recovered':False,'1978_total_catch_to2019_landings':'assumed transfer;Medium allocation'})
check('Manual research fields and no claimed approval', '[MANUAL' in fields['SPPR calculation'] and '[Researcher:' in fields['Open issues and next action'] and '[name]' in fields['Review and reproducibility'] and '[date]' in fields['Review and reproducibility'] and not re.search(r'(researcher approved|approved by|approval granted|production.valid|validated native)',text,re.I),
      {'SPPR_calculation':'unfilled manual','Open_issues':'unfilled manual','Review':'unfilled researcher name/date','researcher_approval_claimed':False})
check('Other automatic context fields', 'Article: ?  Model: ?' in fields['Selection rationale'] and 'Comparative reasons were not recorded' in fields['Selection rationale'] and 'same 1978 baseline' in fields['Other models from the same article'] and '1978–2010 Ecosim' in fields['Other models from the same article'] and 'A ≈ 100%' in fields['Geographic fit'] and 'B ≈ 59%' in fields['Geographic fit'],
      {'selection_rationale':'unknown comparative reasons retained','other_static_models':'none identified;source states/scenarios distinguished','geography':'agrees with separate geography review;approximate containment/area ratio;no exactauthorpolygon claim','other_articles':'2014guide and unverified2019inventory distinguished from verified Dutta2023/Karim2018'})

book=load_workbook(bookpath,data_only=False)
rows=[list(r) for r in book['Taxon mapping'].iter_rows(min_row=7,values_only=True) if r[0] is not None]
records={r['taxon']:r for r in audit['records']}
base={r['taxon']:r for r in baseline['taxa']}
rowerrors=[]
maxcatch=maxppr=0.0
for row in rows:
    name,tl,catch,ppr,mapped,confidence,reason=row
    r=records.get(name)
    if r is None:rowerrors.append({'taxon':name,'error':'not in adopted audit'});continue
    if not close(catch,r['catch_tonnes']):rowerrors.append({'taxon':name,'error':'catch differs'})
    if not close(ppr,r['simple_chain_ppr_tC']):rowerrors.append({'taxon':name,'error':'PPR differs'})
    if confidence!=r['overall_confidence']:rowerrors.append({'taxon':name,'error':'confidence differs'})
    if reason!=r['readable_reason']:rowerrors.append({'taxon':name,'error':'readable reason differs'})
    expectedtl='?' if r['tl'] is None else r['tl']
    if tl!=expectedtl:rowerrors.append({'taxon':name,'error':'TL differs','value':tl,'expected':expectedtl})
    expectedmapped='\n'.join(g['group']+' — '+format(g['weight']*100,'.4f')+'%' for g in r['proposed_groups'])
    if mapped!=expectedmapped:rowerrors.append({'taxon':name,'error':'candidate names/displayed weights differ'})
    b=base[name]
    expected=0 if catch==0 else catch*float(b['classic_sppr'])/9
    if not close(ppr,expected):rowerrors.append({'taxon':name,'error':'independent simple-chain arithmetic differs'})
    maxcatch=max(maxcatch,abs(catch-r['catch_tonnes']))
    maxppr=max(maxppr,abs(ppr-expected))
sortkey=lambda r:(-float(r[3]),r[0])
check('All appendix taxon rows and independent simple-chain arithmetic', len(rows)==315 and len({r[0] for r in rows})==315 and set(r[0] for r in rows)==set(records)==set(base) and not rowerrors and rows==sorted(rows,key=sortkey),
      {'Taxon_mapping_range':'A7:G321','rows':len(rows),'exact_universe':315,'sorted_by_unrounded_numeric_PPR_then_taxon':rows==sorted(rows,key=sortkey),'max_catch_delta_t':maxcatch,'max_independent_PPR_delta_tC':maxppr,'row_errors':rowerrors})
zero=sum(r[2]==0 for r in rows)
miss=sum(r[1]=='?' for r in rows)
positive_unknown=[r[0] for r in rows if r[2]>0 and (r[3]=='?' or r[3] is None)]
totcatch=sum(r[2] for r in rows);totppr=sum(r[3] for r in rows)
check('Zero-catch, missing input and total disclosures',zero==28 and miss==27 and all(r[2]==0 and r[3]==0 for r in rows if r[1]=='?') and not positive_unknown and close(totcatch,audit['summary']['catch_t']) and close(totppr,audit['summary']['ppr_tC']) and f'{totppr:,.3f}' in text,
      {'zero_landings_labels':zero,'missing_TL_coefficient_labels':miss,'positive_catch_unknown_PPR':positive_unknown,'independent_total_landings_t':totcatch,'independent_total_simple_chain_PPR_tC':totppr,'Word_display':f'{totppr:,.3f}'})
coverage=book['Coverage']
categoryerrors=[]
summarydoc=next(t for t in tables if t[0]==['Overall confidence','Taxa (n)','Catch (%)','Simple-chain PPR (%)'])
for k,rowno in zip(['High','Medium','Low','Very low','Unresolved'],range(8,13)):
    subset=[r for r in rows if r[5]==k]
    calc=(len(subset),sum(r[2] for r in subset)/totcatch,sum(r[3] for r in subset)/totppr)
    vals=[coverage.cell(rowno,j).value for j in range(1,5)]
    d=next(r for r in summarydoc[1:] if r[0]==k)
    if vals[0]!=k or vals[1]!=calc[0] or not close(vals[2],calc[1],1e-12) or not close(vals[3],calc[2],1e-12) or int(d[1])!=calc[0] or not close(float(d[2]),calc[1]*100,5e-5,0) or not close(float(d[3]),calc[2]*100,5e-5,0):categoryerrors.append(k)
check('Word/appendix confidence counts and shares',not categoryerrors,{'categories_compared':5,'independent_weakest_component_summary':True,'mismatching_categories':categoryerrors,'total_categories_taxa':sum(coverage.cell(i,2).value for i in range(8,13))})
ruleerrors=[]
for title,key,start,stop in [('Group assignment rules','membership_rules',16,22),('Allocation weight rules','weight_rules',26,27)]:
    expected=audit['summary'][key]
    stored=[[coverage.cell(i,j).value for j in range(1,4)] for i in range(start,stop+1)]
    doctable=next(t for t in tables if t[0]==['Plain-language rule','Confidence','PPR percentage'] and len(t)-1==len(expected))
    for a,s,d in zip(expected,stored,doctable[1:]):
        if s[0]!=a['Plain-language rule'] or s[1]!=a['Confidence'] or not close(s[2]*100,a['PPR percentage']) or d[:2]!=s[:2] or not close(float(d[2]),a['PPR percentage'],5e-5,0):ruleerrors.append(a['Plain-language rule'])
    if not close(sum(s[2] for s in stored),1,1e-10):ruleerrors.append(title+' denominator partition')
check('Independent rule table denominators and confidence',not ruleerrors,{'component_tables':2,'each_sums_to_one':True,'arithmetic_matches_adopted_rule_summaries':not ruleerrors,'errors':ruleerrors})

sourcecatch={r['seq']:r for r in source['catch_rows']}
allocrows=[list(r) for r in book['Allocation evidence'].iter_rows(min_row=7,values_only=True) if r[0] is not None]
expectedalloc={(r['taxon'],g['seq']):(r,g) for r in audit['records'] for g in r['proposed_groups']}
allocerrors=[]
maxdelta=0.0
bytaxon=defaultdict(list)
for row in allocrows:
    name,seq,group,printed,canonical,loaded,total,weight,mc,wc,formula,reason,period,basis,exclusions=row
    key=(name,seq)
    if key not in expectedalloc:allocerrors.append({'taxon':name,'seq':seq,'error':'unknown candidate'});continue
    r,g=expectedalloc[key]
    sc=sourcecatch[seq]
    if group!=sc['group'] or group!=g['group'] or not close(printed,sc['source_displayed_total'],1e-12) or not close(canonical,sc['canonical_catch'],1e-12) or not close(loaded,g['loaded_model_catch_density'],1e-12) or not close(weight,g['weight'],1e-12) or not close(total,r['allocation_total'],1e-12) or mc!=r['membership_confidence'] or wc!=r['allocation_confidence']:
        allocerrors.append({'taxon':name,'seq':seq,'error':'raw catch/identity/total/weight/confidence differs'})
    bytaxon[name].append(row)
    maxdelta=max(maxdelta,abs(loaded-canonical))
for name,candidates in bytaxon.items():
    total=sum(r[5] for r in candidates)
    if len(candidates)>1:
        if total<=0 or any(not close(r[7],r[5]/total,1e-12) for r in candidates):allocerrors.append({'taxon':name,'error':'catch proportion mismatch'})
    elif not close(candidates[0][7],1,1e-12):allocerrors.append({'taxon':name,'error':'single-group weight mismatch'})
    if not close(sum(r[7] for r in candidates),1,1e-12):allocerrors.append({'taxon':name,'error':'weights do not sum to1'})
check('Every candidate density and allocation weight',len(allocrows)==len(expectedalloc)==672 and len(set((r[0],r[1]) for r in allocrows))==672 and not allocerrors,
      {'Allocation_evidence_range':'A7:O678','candidate_rows':len(allocrows),'taxa':len(bytaxon),'max_loaded_vs_canonical_catch_delta':maxdelta,'common_denominator_km2':6205051,'source_basis':'1978 total landings+discards','W4_scope':'assumed pooled composition transfer;Medium','errors':allocerrors})

sr=[list(r) for r in book['Sources'].iter_rows(min_row=7,values_only=True) if r[0] is not None]
sourceurls={r[1] for r in sr}
manifest=read(evidence/'mapping_review/sources_read.json')
requiredurls={r['url'] for r in manifest if 'url' in r}
missingurls=sorted(requiredurls-sourceurls)
wormfiles=['taxonomy_authority_records.json','taxonomy_exception_records.json','taxonomy_extension_records.json','taxonomy_final_exception_records.json']
missingwormurls=[];recordcounts={}
for file in wormfiles:
    recordsfile=read(evidence/'mapping_review'/file)
    recordcounts[file]=len(recordsfile)
    for r in recordsfile:
        if r.get('status')=='read' and r.get('url') not in sourceurls:missingwormurls.append({'file':file,'name':r.get('name'),'url':r.get('url')})
check('Descriptive primary Sources coverage',not missingurls and not missingwormurls and all(len(r)>3 and r[2] and r[3] for r in sr) and any(r[1].endswith('sources_read.json') for r in sr),
      {'Sources_range':'A7:D120','source_entries':len(sr),'primary_manifest_web_entries':len(requiredurls),'missing_primary_urls':missingurls,'WoRMS_file_record_counts':recordcounts,'missing_read_WoRMS_record_urls':missingwormurls,'four_local_WoRMS_files_covered_via_manifest':True,'direct_local_WoRMS_file_links':[f for f in wormfiles if any(f in str(r[1]) for r in sr)],'supports_and_limitations_present_for_every_entry':True})
notes.append({'type':'source_navigation','finding':'All four retained WoRMS record files are named by the linked sources_read.json manifest, and all successfully read record URLs have individual Sources entries. Only taxonomy_final_exception_records.json has a direct local-file Sources hyperlink. This is complete scientific source coverage via the manifest, not four direct local-file entries.'})
notes.append({'type':'precision_limit','finding':'Allocation evidence calls six printed Total/component-sum discrepancies rounding. They are immaterial display-scale differences (five1e-6 and one3e-6), consistent with rounding/internal printed precision. Exact unrounded author-native totals are unavailable; this wording does not establish that native precision was recovered.'})
localissues=[];localcount=0;externalcount=0
links=[]
with zipfile.ZipFile(docpath) as z:
    for relname in z.namelist():
        if relname.endswith('.rels'):
            for r in ET.fromstring(z.read(relname)):
                if r.attrib.get('Type','').endswith('/hyperlink'):
                    links.append(('Word',r.attrib.get('Target','')))
for sheet in book:
    for row in sheet:
        for cell in row:
            if cell.hyperlink and cell.hyperlink.target:links.append((sheet.title+'!'+cell.coordinate,cell.hyperlink.target))
for where,target in links:
    if target.startswith(('https://','http://','mailto:')):externalcount+=1;continue
    if target.startswith('#'):continue
    localcount+=1
    if re.match(r'^[A-Za-z]:',target) or target.startswith(('file:','\\\\','/')):localissues.append({'location':where,'target':target,'error':'absolute local target'});continue
    local=unquote(target.split('#')[0])
    if not (region/local).resolve().exists():localissues.append({'location':where,'target':target,'error':'missing local target'})
check('Source link targets',not localissues,{'relative_local_links_checked':localcount,'external_links_preserved':externalcount,'issues':localissues,'network_refetch_performed':False})
notes.append({'type':'resolved_template_residue','finding':'The earlier report SHA9300e50c15b9aefda8adb823484ee17ad2c13f7fb727055fb24ad453a9502972 contained unused rIdValidation1 pointing to [Excel_appendix.xlsx]. Root removed the orphan template relationship and rebuilt the report. The current package was reread and all stored hyperlink targets rechecked; no residual missing local target remains. This reviewer did not edit either deliverable.'})
tablefilters=[]
for s in book:
    tablefilters.extend({'sheet':s.title,'table':t.name,'ref':t.ref,'autoFilter':None if t.autoFilter is None else t.autoFilter.ref} for t in s.tables.values())
check('Numeric record types and native row filters',all(isinstance(r[2],(int,float)) and isinstance(r[3],(int,float)) for r in rows) and any(t['sheet']=='Taxon mapping' and t['autoFilter'] for t in tablefilters),
      {'catch_and_PPR_numeric':True,'native_table_filters':tablefilters,'layout_rendering':'owned by root;not repeated in this read-only scientific QA'})
final={'schema_version':1,'review_date':'2026-09-30','scope':'read-only final scientific consistency QA;no deliverable rewrite or scientific solver execution',
 'deliverables':{'report':{'path':'../../../Model_validation_34_1_Bay_of_Bengal_(1978).docx','sha256':sha(docpath)},'appendix':{'path':'../../../LME034_taxon_mapping_appendix.xlsx','sha256':sha(bookpath)}},
 'evidence_basis':{'canonical_sha256':source['model_sha256'],'source_pdf_sha256':source['source_sha256'],'source_audit':'source_audit.json','direct_summary':'direct_summary.json','complete_matrix_sources':['GE/SPPR.npy','TE/SPPR.npy','With_Egestion/SPPR.npy'],'adopted_taxon_audit':'../adopted_taxon_audit.json','mapping_source_manifest':'../mapping_review/sources_read.json'},
 'status':'PASS' if not errors else 'FAIL','checks':checks,'errors':errors,'notes':notes,
 'scientific_eligibility':{'numerical_scenario_reproducible':True,'source_fidelity_verified':False,'author_native_model_recovered':False,'researcher_approval_claimed':False,'approval_or_production_validation_conferred_by_this_QA':False},
 'deliverables_unchanged':{'report':sha(docpath)==extracted['document_sha256'],'appendix':sha(bookpath)==extracted['appendix_sha256']}}
(out/'final_artifact_scientific_qa.json').write_text(json.dumps(final,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'status':final['status'],'checks':len(checks),'errors':errors,'notes':notes,'report_sha256':sha(docpath),'appendix_sha256':sha(bookpath)},ensure_ascii=False,indent=2))
