"""Retry only the two timed-out 75-draw methods; preserve completed outputs."""
import json,shutil,sys,time,warnings
from pathlib import Path
import openpyxl
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[4]
sys.path.insert(0,str(ROOT/'tools'));sys.path.insert(0,str(ROOT/'tools/scientific_code/PPREstimation'))
from workbooks import sha
import create_PPRS_excel as cpe
MC=['MC_new_GE','MC_new_TE_EEfix']
MODEL=ROOT/'regions/LME_034/models/34_1_Bay_of_Bengal_(1978)/model.json'
SPPR=MODEL.parent/'sppr_source.xlsx'

class ProgressRunner(cpe._TaskRunner):
    def run(self,kind,payload,mc_samples):
        name=payload if isinstance(payload,str) else str(payload)
        print('START',kind,name,'samples',mc_samples,flush=True)
        t=time.perf_counter();result=super().run(kind,payload,mc_samples)
        print('FINISH',kind,name,'status',result[0],'seconds',time.perf_counter()-t,flush=True)
        return result

def main():
    config=json.loads((OUT/'refresh_configuration.json').read_text(encoding='utf-8'))
    assert config['canonical_model_sha256']==sha(MODEL)
    assert all(config['method_status'][m]=='timeout' for m in MC)
    assert all(config['method_status'][m]=='ok' for m in config['methods'] if m not in MC)
    assert config['fresh_sppr_sha256']==sha(SPPR)
    prior=OUT/'export/before_mc_retry.xlsx';shutil.copy2(SPPR,prior)
    (OUT/'initial_export_configuration.json').write_text(json.dumps(config,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always');calc,label=cpe.load_model(str(MODEL))
    runner=ProgressRunner(600);runner.set_model(calc)
    t=time.perf_counter()
    try:tables=cpe.build_model_tables(calc,method_keys=MC,mc_samples=75,model_label=label,source_file=str(MODEL.relative_to(ROOT)),method_timeout=600,runner=runner)
    finally:runner.close()
    assert all(tables.method_status[m]=='ok' for m in MC),tables.method_status
    retry=cpe.write_tables_excel(tables,str(MODEL),str(OUT/'mc_retry'))
    a=openpyxl.load_workbook(SPPR);b=openpyxl.load_workbook(retry,data_only=True)
    protected={}
    for sheet in ['sppr_all','sppr_inner','sppr_PP']:
        old=list(a[sheet].values);new=list(b[sheet].values);h=old[0];nh=new[0]
        protected[sheet]=[[v for i,v in enumerate(row) if h[i] not in MC] for row in old]
        by_seq={row[0]:row for row in new[1:]}
        for ri,row in enumerate(old[1:],2):
            for m in MC:a[sheet].cell(ri,h.index(m)+1).value=by_seq[row[0]][nh.index(m)]
    for sheet in ['footprint','run_notes']:
        old=list(a[sheet].values);new=list(b[sheet].values);assert old[0]==new[0],sheet
        lookup={row[0]:row for row in new[1:] if row[0] in MC}
        for ri,row in enumerate(old[1:],2):
            if row[0] in MC:
                for ci,v in enumerate(lookup[row[0]],1):a[sheet].cell(ri,ci).value=v
    # This sheet contains only these two MC methods and was empty after timeouts.
    s=a['mc_diagnostics'];s.delete_rows(1,s.max_row)
    for row in b['mc_diagnostics'].values:s.append(row)
    b.close();a.save(SPPR);a.close()
    check=openpyxl.load_workbook(SPPR,read_only=True,data_only=True)
    for sheet,old in protected.items():
        new=list(check[sheet].values);h=new[0]
        assert [[v for i,v in enumerate(row) if h[i] not in MC] for row in new]==old,'Completed coefficient outputs changed: '+sheet
    sample_rows=list(check['mc_diagnostics'].values);check.close()
    assert all(int(r[1])==75 for r in sample_rows[1:])
    config['method_status'].update(tables.method_status)
    config['initial_method_timeout_seconds']=180
    config['method_timeout_overrides_seconds']={m:600 for m in MC}
    config['initial_timed_out_methods']=MC
    config['mc_retry_seconds']=time.perf_counter()-t
    config['issues']=[x for x in config['issues'] if x.get('method') not in MC]+tables.issues
    config['fresh_sppr_sha256']=sha(SPPR)
    config['mc_retry_preserved_twenty_completed_methods_exactly']=True
    config['final_mc_diagnostics']=[dict(zip(sample_rows[0],r)) for r in sample_rows[1:]]
    (OUT/'refresh_configuration.json').write_text(json.dumps(config,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    shutil.copy2(SPPR,OUT/'export/model.xlsx')
    print(json.dumps({'all22_status':config['method_status'],'retry_seconds':config['mc_retry_seconds'],'mc_diagnostics':config['final_mc_diagnostics'],'other20_coefficients_exactly_preserved':True}),flush=True)

if __name__=='__main__':main()
