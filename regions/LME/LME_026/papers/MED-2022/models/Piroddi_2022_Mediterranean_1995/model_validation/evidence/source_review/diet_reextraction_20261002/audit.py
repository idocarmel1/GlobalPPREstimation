"""Read-only source normalization audit; writes only this evidence directory."""
from pathlib import Path
from zipfile import ZipFile
from decimal import Decimal
import csv, hashlib, json, shutil, xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
REGION = ROOT / 'regions/LME_026'
OUT = Path(__file__).parent
MID = 'Piroddi_2022_Mediterranean_1995'
MD = REGION / 'models' / MID
ED = MD / 'extracted_tables'
REVIEW = REGION / 'extraction_review_20260928'
SELECTED = MD / 'model.json'
SOURCE = REGION / 'papers/MED-2022/41598_2022_18017_MOESM1_ESM-75e4d6d3.xlsx'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p): return p.relative_to(ROOT).as_posix()
def save(name, obj): (OUT/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf8')

protected = {rel(p):sha(p) for p in REGION.rglob('*') if p.is_file() and OUT not in p.parents}
before = sha(SELECTED)
overview = {r[0]:r[1] for r in load_workbook(REGION/'LME_026.xlsx', read_only=True, data_only=True)['Overview'].values if len(r)>=2}
assert overview['selected_model_id'] == MID
assert overview['model_path'] == 'models/'+MID+'/model.json'
model = json.loads(SELECTED.read_text(encoding='utf8'))
copies = [SELECTED, ED/'26_2602022_Mediterranean_Piroddi_(1995).json', MD/'diagnostics/26_2602022_Mediterranean_Piroddi_Source_(1995).json', MD/'diagnostics/26_2602022_Mediterranean_Piroddi_(1995).json']
archive = OUT/'originals_by_sha256'; archive.mkdir(exist_ok=True)
archived = []
for p in copies:
 h=sha(p); target=archive/(h+'.json')
 if not target.exists(): shutil.copyfile(p,target)
 assert sha(target)==h
 archived.append({'path':rel(p), 'sha256':h, 'archive':rel(target)})

wb = load_workbook(SOURCE, read_only=False, data_only=False)
ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with ZipFile(SOURCE) as z:
 raw = {c.attrib['r']:c.find('s:v',ns).text for c in ET.fromstring(z.read('xl/worksheets/sheet2.xml')).findall('.//s:c',ns) if c.find('s:v',ns) is not None and c.attrib.get('t') not in ['s','str','inlineStr']}
assert wb.sheetnames[1]=='Diets'
rows = list(csv.reader((ED/'Diet_composition.csv').open(encoding='utf8',newline='')))
imports = next(row for row in rows if row[1]=='Import')
retained_evidence = {(int(e['group_seq']), str(e.get('prey_seq','import'))):e for e in json.loads((REVIEW/'cell_evidence.json').read_text(encoding='utf8')) if e['field'] in ('diet','diet_import')}
loader_rows = list(csv.reader((MD/'diagnostics/loader_diet.csv').open(encoding='utf8',newline='')))
loaded = {int(row[0]):dict(zip(map(int,loader_rows[0][1:]), row[1:])) for row in loader_rows[1:]}
ledger=[]; sums=[]; mismatches=[]; runtime_deltas=[]
variants={rel(p):json.loads(p.read_text(encoding='utf8')) for p in copies}
for consumer in range(1,66):
 g=model['group'][consumer-1]
 assert int(g['group_seq'])==consumer and int(wb['Diets'].cell(3,consumer+2).value)==consumer
 entries=g['diet_descr']['diet']; prey_map={int(d['prey_seq']):(k,d) for k,d in enumerate(entries)}
 total=Decimal(0)
 for prey in list(range(1,72))+['import']:
  source_row=75 if prey=='import' else prey+3
  coord=wb['Diets'].cell(source_row,consumer+2).coordinate
  literal=raw[coord]; total+=Decimal(literal)
  if prey=='import':
   adopted=g['diet_imp']; import_literal=imports[consumer+1]; pointer=f'/group/{consumer-1}/diet_imp'; runtime_preylevel=72
  else:
   k,d=prey_map[prey]; adopted=d['proportion']; import_literal=rows[prey][consumer+1]; pointer=f'/group/{consumer-1}/diet_descr/diet/{k}/proportion'; runtime_preylevel=prey
  evidence=retained_evidence[(consumer,str(prey))]
  variant_literals=[]
  for path,v in variants.items():
   vg=v['group'][consumer-1]
   vl=vg['diet_imp'] if prey=='import' else next(d['proportion'] for d in vg['diet_descr']['diet'] if int(d['prey_seq'])==prey)
   variant_literals.append({'path':path,'literal':vl})
  equal=(literal==adopted==evidence['value'] and Decimal(literal)==Decimal(import_literal) and all(literal==v['literal'] for v in variant_literals))
  runtime_literal=loaded[consumer][runtime_preylevel]
  delta=abs(float(adopted)-float(runtime_literal)); runtime_deltas.append(delta)
  item={'consumer_seq':consumer,'consumer_name':g['group_name'],'prey_seq':prey,'source_file':rel(SOURCE),'source_sha256':sha(SOURCE),'source_sheet':'Diets','source_cell':coord,'stored_xml_literal':literal,'source_number_format':wb['Diets'].cell(source_row,consumer+2).number_format,'canonical_pointer':pointer,'canonical_literal':adopted,'extracted_import_literal':import_literal,'retained_extraction_literal':evidence['value'],'variant_literals':variant_literals,'source_canonical_literals_equal':literal==adopted==evidence['value'],'extracted_import_decimal_equal':Decimal(literal)==Decimal(import_literal),'extracted_import_lexical_equal':literal==import_literal,'all_source_values_equal':equal,'accepted_correction':None,'review_state':'researcher_review_not_established_by_this_audit','retained_loader_value':runtime_literal,'source_runtime_abs_difference':delta}
  ledger.append(item)
  if not equal:mismatches.append(item)
 sums.append({'consumer_seq':consumer,'consumer_name':g['group_name'],'prey_plus_import_decimal_sum':str(total),'import_literal':g['diet_imp'],'missing_cells':0,'runtime_normalization_setting':False,'retained_loader_sum':sum(float(v) for v in loaded[consumer].values() if v),'normalization_factor_in_retained_run':1,'review_state':'researcher_review_not_established_by_this_audit'})
if mismatches: print(json.dumps(mismatches[:2],ensure_ascii=True,indent=2))
assert not mismatches, 'Unexpected source mismatch: do not mutate canonical'
save('cell_ledger.json',ledger); save('consumer_sums.json',sums)

# Researcher DOCX inspection is read-only. Record paragraphs relevant to normalization/corrections.
w_ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
docx_records=[]
for p in REGION.glob('*.docx'):
 with ZipFile(p) as z: tree=ET.fromstring(z.read('word/document.xml'))
 paragraphs=[''.join(t.text or '' for t in x.findall('.//w:t',w_ns)) for x in tree.findall('.//w:p',w_ns)]
 relevant=[t for t in paragraphs if any(k in t.lower() for k in ['diet','normaliz','correct','תזונ','נרמ','תיקו'])]
 docx_records.append({'path':rel(p),'sha256':sha(p),'relevant_paragraphs':relevant,'accepted_diet_override_identified':False})
save('researcher_edit_inspection.json',docx_records)

history_script=REVIEW/'build_extraction.py'
history={'script':rel(history_script),'sha256':sha(history_script),'restoration_ledger':rel(REVIEW/'converter_source_restoration.json'),'restoration_ledger_sha256':sha(REVIEW/'converter_source_restoration.json'),'earlier_VERIFICATION':json.loads((REVIEW/'VERIFICATION.json').read_text(encoding='utf8')),'finding':'Retained build_extraction.py explicitly replaced converter diet_descr and diet_imp with original XLSX XML literals. Current selected and extracted JSON retain those literals exactly. No normalization inferred from sums.'}
save('converter_history.json',history)
settings=json.loads((MD/'diagnostics/staging_transformations.json').read_text(encoding='utf8'))
runtime={'tolerance':{'atol':1e-12,'rtol':0},'canonical_changed':False,'all_complete_input_bytes_identical_before_after':True,'actual_retained_settings':settings['requested_runtime_defaults'],'canonical_sha256_before':before,'canonical_sha256_after':sha(SELECTED),'retained_artifact_hashes':{rel(p):sha(p) for p in (MD/'diagnostics').iterdir() if p.is_file()},'source_vs_saved_loader_diet_cells':len(ledger),'max_source_saved_loader_abs_difference':max(runtime_deltas),'equivalent_within_tolerance':max(runtime_deltas)<=1e-12,'comparison_basis':'No canonical/staged input, settings, group parameters, routing or complete loaded state changed. Source-vs-saved ModelData diet compared cellwise; no new PPRCalculator run. Retained calculator construction is EXCEPTION, so no completed normalized calculator state or SPPR coefficients exists. Runtime normalization remains permitted on a separate copy.','retained_construction':json.loads((MD/'diagnostics/loader_audit.json').read_text(encoding='utf8')),'fresh_scientific_runs_performed':False,'needed_refresh_scope':[]}
save('runtime_comparison.json',runtime)
unchanged={p:(ROOT/p).is_file() and sha(ROOT/p)==h for p,h in protected.items()}
save('protected_file_verification.json',{'files':[{'path':p,'before_sha256':protected[p],'after_sha256':sha(ROOT/p) if (ROOT/p).is_file() else None,'unchanged':v} for p,v in unchanged.items()],'all_unchanged':all(unchanged.values()),'scope':'All preexisting region files; no shared file mutations performed.'})
assert all(unchanged.values()),'Concurrent regional changes: receipt must not claim unchanged'
save('source_archive_manifest.json',archived)
paths=[rel(p) for p in OUT.iterdir() if p.is_file() and p.name!='root_receipt.json']
blockers=['Existing scientific limitation: missing detritus routing between Discards70 and Detritus71 blocks calculator construction; unchanged, no repair authorized.','Diet review trust remains not established by this audit; exact source fidelity is not researcher signoff.']
save('root_receipt.json',{'region':'LME_026','selected_model_id':MID,'selected_path':rel(SELECTED),'outcome':'verified_unnormalized_unchanged','normalization_status':'Source-faithful unnormalized canonical verified from all 4680 exact original XLSX diet/import literals and restoration history.','reextraction_performed':False,'canonical_before_sha256':before,'canonical_after_sha256':sha(SELECTED),'changed_diet_cells':[],'verified_diet_import_cells':len(ledger),'evidence_paths':paths,'precise_blockers':blockers,'protected_files_unchanged':all(unchanged.values()),'runtime_equivalent_at_abs1e_12':True,'needed_refresh_scope':[],'source_fidelity_limits':'Primary parameter-bearing XLSX stored XML literal fidelity proven. Author upstream/native revisions and researcher review trust are not inferred.','review_approval':False})
print(json.dumps({'cells':len(ledger),'mismatches':len(mismatches),'sha256':before,'runtime_max_delta':max(runtime_deltas),'protected_files':len(unchanged),'all_unchanged':all(unchanged.values()),'sum_range':[str(min(Decimal(s['prey_plus_import_decimal_sum']) for s in sums)),str(max(Decimal(s['prey_plus_import_decimal_sum']) for s in sums))],'receipt':rel(OUT/'root_receipt.json')},indent=2))
