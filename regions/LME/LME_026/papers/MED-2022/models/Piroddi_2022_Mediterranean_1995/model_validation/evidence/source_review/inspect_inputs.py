from pathlib import Path
import sys,json,hashlib
ROOT=Path(__file__).resolve().parents[4]
sys.path.insert(0,str(ROOT/'tools'))
from workbooks import read_book,records,overview,digest_tables
from docx import Document
import openpyxl
OUT=Path(__file__).parent
REGION=ROOT/'regions/LME_026'
b=read_book(REGION/'LME_026.xlsx')
snap={s:{t:{'header':h,'records':[dict(zip(h,r)) for r in rows]} for t,(h,rows) in block.items()} for s,block in b.items()}
(OUT/'workbook_input_snapshot.json').write_text(json.dumps(snap,indent=2,ensure_ascii=False),encoding='utf-8')
model=json.loads((REGION/'models/Piroddi_2022_Mediterranean_1995/model.json').read_text(encoding='utf-8'))
groups=model.get('group',model.get('groups'))
meta={}
pw=openpyxl.load_workbook(ROOT/'Project.xlsx',read_only=True)
for sn in ['Papers','Models & coverage','Regions & status']:
 if sn not in pw: continue
 found=[];head=None
 for row in pw[sn].values:
  if not row or not any(v is not None for v in row):continue
  if row[0]=='@table': head=None;continue
  if head is None: head=list(row);continue
  r=dict(zip(head,row))
  if r.get('unit_id')=='LME_026':found.append(r)
 meta[sn]=found
pw.close()
(OUT/'central_metadata_snapshot.json').write_text(json.dumps(meta,indent=2,ensure_ascii=False,default=str),encoding='utf-8')
template=Document(ROOT/'tools/templates/Model_validation_template.docx')
t={'paragraphs':[p.text for p in template.paragraphs],'tables':[[[c.text for c in r.cells] for r in tb.rows] for tb in template.tables]}
(OUT/'template_content.json').write_text(json.dumps(t,indent=2,ensure_ascii=False),encoding='utf-8')
baseline=ROOT/'original_research_archive/research/selected_regions_validation_20260930/baseline/LME_026'
before=read_book(baseline/'LME_026.xlsx')
protected={}
for sn in ['Catch','Classic PPR','Selected model groups','NPP']:
 for tn,(h,rows) in before[sn].items():
  if tn not in ['Catch','Taxa','Groups','Group SPPR','NPP','Provenance']:continue
  cur=b[sn][tn]
  protected[sn+'/'+tn]={'baseline_digest':digest_tables([(h,rows)]),'current_digest':digest_tables([cur]),'exact_equal':(h,rows)==cur}
out={'overview':overview(b),'blocks':{s:{t:{'headers':h,'count':len(r)} for t,(h,r) in block.items()} for s,block in b.items()},'model_top_keys':list(model),'group_count':len(groups) if groups else None,'first_groups':groups[:3] if groups else None,'protected':protected,'template':t,'central_metadata':meta}
print(json.dumps(out,ensure_ascii=True,indent=2))
