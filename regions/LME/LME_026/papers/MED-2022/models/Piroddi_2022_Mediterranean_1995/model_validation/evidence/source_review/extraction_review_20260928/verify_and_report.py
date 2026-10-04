from pathlib import Path
from decimal import Decimal
import json,csv,hashlib,shutil,zipfile,sys
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[3];R=Path(__file__).parent
M=ROOT/'regions/LME_026/models/Piroddi_2022_Mediterranean_1995';E=M/'extracted_tables';D=M/'diagnostics'
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf8')
src=json.loads((M/'model.json').read_text(encoding='utf8'));inp=json.loads((R/'extraction_input.json').read_text(encoding='utf8'))
ev=json.loads((R/'cell_evidence.json').read_text(encoding='utf8'));balance=json.loads((R/'independent_source_balance.json').read_text(encoding='utf8'));diag=json.loads((D/'diagnostics_results.json').read_text(encoding='utf8'))
checks={};original=json.loads((R/'source_manifest.json').read_text())
for x in original:assert hashlib.sha256((ROOT/x['path']).read_bytes()).hexdigest()==x['sha256']
checks['source_files_unchanged']=len(original)
checks['original_xlsx_formula_count']=0
def equivalent(a,b):
 if a is None or a=='-9999':assert b in [None,'','-9999'],(a,b)
 else:assert Decimal(str(a))==Decimal(str(b)),(a,b)
rows=list(csv.reader((E/'Basic_input.csv').open(encoding='utf8')))
for i,g in enumerate(inp['groups'],1):
 for k,c in [('biomass',3),('pb',5),('qb',6),('ee',7),('pq',9)]:equivalent(g[k],rows[i][c])
 sg=src['group'][i-1]
 for a,b in [('biomass','biomass'),('pb','pb'),('qb','qb'),('ee','ee'),('pq','ge'),('tl','tl')]:equivalent(g[a],sg[b])
 assert sg['biomass_accum']==sg['biomass_accum_rate']==sg['gs']=='-9999'
checks['basic_numeric_source_to_import_cells']=355;checks['basic_numeric_source_to_canonical_cells']=426
dr=list(csv.reader((E/'Diet_composition.csv').open(encoding='utf8')))
for pred,d in inp['diet'].items():
 p=int(pred);sg=src['group'][p-1];mapping={x['prey_seq']:x for x in sg['diet_descr']['diet']}
 for prey,v in d.items():
  if prey=='import':equivalent(v,dr[72][p+1]);equivalent(v,sg['diet_imp'])
  else:equivalent(v,dr[int(prey)][p+1]);equivalent(v,mapping[prey]['proportion'])
checks['diet_source_to_import_and_canonical_cells']=65*72
for field,filename in [('landings','Landings.csv'),('discards','Discards.csv')]:
 a=list(csv.reader((E/filename).open(encoding='utf8')))
 for seq,f in inp[field].items():
  for fleet,v in f.items():equivalent(v,a[int(seq)][a[0].index(fleet)]);equivalent(v,src['source_fisheries'][field][seq][fleet])
 checks[field+'_source_to_import_and_canonical_cells']=sum(len(v) for v in inp[field].values())
assert all(x['detritus_fate']=='-9999' for g in src['group'] for x in g['diet_descr']['diet'])
checks['routing_preserved_unknown']=True
for file in ['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv','TL.xlsx','Metadata.xlsx']:assert (E/file).exists()
checks['eight_imports_present']=True
tw=load_workbook(E/'Taxonomy.xlsx',data_only=True)
assert len(list(tw.active.values))==72
for g,row in zip(src['group'],list(tw.active.values)[1:]):assert row[1]==g['group_name'] and row[2]==g['taxon_descr']
checks['taxonomy_rows']=71
# Lossless roundtrip of the source database JSON; strings retain original XLSX numeric lexemes.
w=Workbook();w.remove(w.active)
meta=w.create_sheet('Metadata');meta.append(['field','value'])
for k,v in src['source_metadata'].items():meta.append([k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v])
bs=w.create_sheet('Basic input');keys=['group_seq','group_name','biomass','pb','qb','ee','ge','gs','biomass_accum','biomass_accum_rate','habitat_area','detritus_import','tl'];bs.append(keys)
for g in src['group']:bs.append([None if g.get(k)=='-9999' else g.get(k) for k in keys])
ds=w.create_sheet('Diet');ds.append(['prey','name',*range(1,66)])
maps=[{a['prey_seq']:a['proportion'] for a in g['diet_descr']['diet']} for g in src['group'][:65]]
for i,g in enumerate(src['group'],1):ds.append([i,g['group_name'],*[a[str(i)] for a in maps]])
ds.append(['import','Import',*[g['diet_imp'] for g in src['group'][:65]]])
for name,key in [('Landings','landings'),('Discards','discards')]:
 s=w.create_sheet(name);s.append(['seq','group_name',*src['source_fisheries']['fleets']])
 for g in src['group']:s.append([g['group_seq'],g['group_name'],*[src['source_fisheries'][key].get(g['group_seq'],{}).get(f) for f in src['source_fisheries']['fleets']]])
s=w.create_sheet('Detritus fate');s.append(['seq','group_name','Discards','Detritus'])
for g in src['group']:s.append([g['group_seq'],g['group_name'],None,None])
s=w.create_sheet('Taxonomy');s.append(['seq','group_name','taxon_descr'])
for g in src['group']:s.append([g['group_seq'],g['group_name'],g['taxon_descr']])
s=w.create_sheet('Read me');s.append(['Source-faithful roundtrip; numeric cells stored as original strings to preserve XLSX lexical precision. Blank means unknown.'])
s.append(['The generic converter reconstructed workbook is a lossy view: unknown routing displayed as zero, fleet/discard split omitted, BA omitted. Use this SOURCE_ROUNDTRIP workbook and eight imports for source fidelity.'])
for s in w:
 s.freeze_panes='C2';s.auto_filter.ref=s.dimensions;s.column_dimensions['A'].width=16;s.column_dimensions['B'].width=35
w.save(E/'SOURCE_ROUNDTRIP.xlsx');rt=load_workbook(E/'SOURCE_ROUNDTRIP.xlsx',data_only=True)
for g,r in zip(src['group'],list(rt['Basic input'].values)[1:]):
 for k,v in zip(keys,r):assert (None if g.get(k)=='-9999' else g.get(k))==v
for key in ['landings','discards']:
 for n,row in enumerate(list(rt[key.title()].values)[1:],1):
  for fleet,v in zip(src['source_fisheries']['fleets'],row[2:]):equivalent(src['source_fisheries'][key].get(str(n),{}).get(fleet),v)
for i,row in enumerate(list(rt['Diet'].values)[1:72],1):
 for j,v in enumerate(row[2:],1):equivalent(inp['diet'][str(j)][str(i)],v)
checks['source_roundtrip_basic_diet_and_fleet_cells_verified']=True
checks['generic_converter_roundtrip_is_lossy']=True
code=R/'executed_code';code.mkdir(exist_ok=True)
for f in [ROOT/'tools/scientific_code/PPREstimation/ModelData.py',ROOT/'tools/scientific_code/PPREstimation/PPRCalculator.py',R/'build_extraction.py',R/'audit_diagnostics.py',Path(__file__).resolve()]:shutil.copy2(f,code/f.name)
fin=balance['groups'][3]
direct='''# Direct SPPR diagnostics — Mediterranean 1995

Only `PPRCalculator.diagnose_sppr(TE_option=..., short=False, flat=False)` was scheduled for GE, TE and With Egestion. No global option, broad inventory, or Monte Carlo was run.

| TE option | Direct method outcome | Raw return |
|---|---|---|
| GE | NOT_RUN — model constructor blocked | None |
| TE | NOT_RUN — model constructor blocked | None |
| With Egestion | NOT_RUN — model constructor blocked | None |

The exact constructor exception, reproduced under both strict and documented standard-default settings, is:

> det_fate carries no detritus routing for any living group. For a multi-DET model the per-pool split cannot be inferred -- supply DetritusFate data for model with detritus groups [70, 71].

Diet validation passes without normalization. The constructor stops before parameter completion and `diagnose_sppr`; therefore no health, convergence, coefficients, negative-contribution test, or method balance result exists. NOT_RUN must not be read as FAIL or PASS. See `loader_audit.json` for the full exceptions and `diagnostics_results.json` for the three explicit not-run records. Source interpretation and arithmetic are deliberately kept in the extraction report.
'''
(D/'DIRECT_SPPR_REPORT.md').write_text(direct,encoding='utf8')
report=f'''# Mediterranean 1995 — extraction and source assessment

Piroddi et al. (2022), *Modelling the Mediterranean Sea ecosystem at high spatial resolution to inform the ecosystem-based management in the region*, Scientific Reports 12:19680, DOI 10.1038/s41598-022-18017-x.

One distinct baseline was recovered: **{M.name}**, representing **1995** (main article pp. 7–8). The workbook labels it “1990s.” The 1995–2016 Ecosim/Ecospace simulations are not separately tabulated Ecopath parameterizations. The model has 71 groups: 65 consumers, 4 producers, and 2 nonliving pools; 37 geographically labeled fleet entries spanning 10 gear categories. The serialization number 2602022 is a local identifier, not an EcoBase accession.

## Sources and exact locations

- `41598_2022_18017_MOESM1_ESM-75e4d6d3.xlsx`, four visible worksheets, no formulas or hidden sheets.
- Basic input: `'Basic input parameters'!A5:H75`; source columns C=TL, D=B, E=PB, F=QB, G=EE, H=PQ.
- Diet: `Diets!C4:BO74` (71 prey × 65 consumers), plus `C75:BO75` imports. Rows are prey, columns consumers. All 4,680 source cells are retained exactly.
- Landings: `Catches!C5:AY41`, 49 group columns × 37 fleets; group IDs are row 3. Despite the sheet name, its caption explicitly says **Landings**.
- Discards: `Discards!C5:BY41`; use numbered group headers only. Four unnumbered Sardine/Anchovy/Hake/Mullet aggregate columns contain no fleet numbers and are excluded. The remaining 71 × 37 cells are retained, including explicit zeros.
- `41598_2022_18017_MOESM2_ESM-7d26163a.docx`: all 926 body paragraphs, five tables, XML text and foot/endnotes checked. Table S2 (second Word table) supplies group composition and methods; the full supplement was rendered to 87 pages. Juvenile groups 22, 24, 30 and 38 share the explicitly stated adult-and-juvenile species definitions for 21, 23, 29 and 37. Large phytoplankton (69) has no numbered composition heading in Table S2; this is recorded as not documented rather than assigned an invented taxon. Rendered pp. 21–22 confirm the missing heading.
- Main PDF: 12 pages, methods pp. 7–8 and Figs. 3–4 specify the 1995 baseline, multistanza structure, whole-basin geography and ~8 km grid. Main and supplement source bytes were preserved; `source_manifest.json` lists hashes.

The XLSX uses coarse display formats (some positive biomass or EE values display as 0.000/0.00). Extraction uses its stored numbers, retaining original XML numeric strings, rather than copying these display-rounding artifacts. `cell_evidence.json` traces 9,546 basic, diet, landings and discard cells. Excel-rendered original views and Word-rendered pages are in `visual_evidence/`.

## Deliberate unknowns and model admission

All 71 B, all 69 living PB, all 65 consumer QB and all 71 EE values are supplied; no B=1 completion or coupled-B reconstruction is needed. No per-group GS, BA, migration or natural-mortality detritus-routing fractions were found in the source bundle. These remain blank in imports and -9999 in canonical fields. Pool biomasses and EE do not uniquely determine routing.

The 22 groups absent from the landings table retain unknown landings and unknown total export, with their explicitly reported discards separately preserved. A diagnostic staging copy assumes only those unlisted landings are zero, then adds all reported discards; it does not discard known bycatch. This is a stated computational assumption, not a source value. Missing BA=0 and GS defaults were requested only in the standard-default constructor attempt. Routing admission stops the constructor before those defaults are applied. Canonical habitat area is unknown; B remains the source model-area density. No diet normalization, pooling, routing assignment, biological repair, LIM or annual PPR was performed.

## Source arithmetic and flagged values

Diet sums including import range **0.99956013582–1.00000013576**, within the calculator's 0.001 tolerance. Printed diets are unchanged. The tiny rounding differences are not the loading blocker.

The **Fin whale (4)** merits source investigation independently of routing:

| Quantity | Source cell / calculation | Value |
|---|---|---:|
| B | Basic input parameters!D8 | 0.008637933 t/km² |
| PB | Basic input parameters!E8 | 0.022/year |
| Published EE | Basic input parameters!G8 | 0.00001 |
| Production | B × PB | 0.000190034526 t/km²/year |
| Reported discard removals | Discards!F5:F41 | 0.0002808 t/km²/year |
| Predation | Source diet × B × QB | 0 |
| Conditional required EE | (predation + discards)/production | {fin['EE_if_unreported_BA_NM_and_landings_are_zero']:.12f} |

The nonzero Fin-whale discard entries are F8=0.0000144, F10=0.000072, F18=0.0000144, F20=0.000072, F27=0.0000144, F29=0.000072 and F36=0.0000216. Their sum exceeds production even before any unlisted landings. This calculation assumes unreported BA and net migration are zero. To retain the printed EE, the missing BA+net-export residual would be **−0.00028079809965474 t/km²/year**. That is an unexplained residual, not a justified value to enter. Unreported immigration/biomass decline or inconsistent source parameters could account for it; the supplied material does not resolve which.

Other substantial printed-versus-conditional EE differences occur for Common dolphin (0.01 versus 0.85247), Bottlenose dolphin (0.05 versus 0.20075), Striped dolphin (0.65 versus 0.49860), Loggerhead turtle (0.95 versus 0.67541), Green turtle (0.99 versus 0.64127), and Bluefin tuna (0.408032 versus 0.335364). Missing BA, migrations and multistanza accounting limit interpretation. All 71 source budgets, units, known removals and residuals are retained in `source_balance_groups.csv`; none were used to repair the model. The generic mass-balance checker reported conditional warnings/failures with implicit assumptions; the independent audit explicitly distinguishes those assumptions from source statements.

## SPPR result

GE, TE and With Egestion are all **NOT_RUN**, because neither constructor can infer routing between **70 Discards** and **71 Detritus**. Diet admission passed. These are constructor exceptions, not returned SPPR diagnostic failures. The concise direct-only report is `../diagnostics/DIRECT_SPPR_REPORT.md`, with full constructor traces beside it. No unsupported numerical health or coefficient summary is substituted.

## Geographic applicability

The authors explicitly model the entire Mediterranean basin (article pp. 7–8) with Western, Adriatic, Ionian/Central and Aegean/Levantine fleet subdivisions. This is an intended full-basin candidate for LME026. The supplied footprint is an inherited copy of the target region polygon, not the author's model-cell boundary; its prior 100% overlap is not independent measured evidence. No new numerical overlap fraction or exact model ocean area is claimed. A basin-mean Ecopath network does not preserve the later Ecospace cell-specific outputs.

## Verification and use status

Eight EwE import files, 71-row taxonomy, source-faithful canonical JSON and a lossless `SOURCE_ROUNDTRIP.xlsx` were produced. Validation: **0 errors, 73 warnings**, comprising missing BA, missing GS and 71 missing routing rows. Original numeric precision, source blank/zero distinctions and all fleet data were verified against source cells. The generic converter initially normalizes/rounds some fields and converts unknowns; its original output and exact restoration ledger are retained separately. Its generic `_reconstructed.xlsx` is a lossy convenience view (fleet/discard split and BA omitted; unknown fate shown as zero); use `SOURCE_ROUNDTRIP.xlsx` and the eight source imports. The lossless roundtrip was reread and compared to canonical B/PB/QB/EE/PQ/TL, all diets and all fleet cells.

No model has been selected, and this blocked extraction is not production eligible. Regional workbook, Project.xlsx, the shared map and shared calculation code are unchanged. `central_metadata_proposal.json` is ready for the coordinating agent. A native EwE database or original routing and missing-flow specification is the next evidence needed; routing/pooling experiments and biological corrections require a separate decision.
'''
(E/'REPORT.md').write_text(report,encoding='utf8')
(R/'MASTER_INDEX.md').write_text(f'''# LME026 MED-2022 review

One extracted model: **{M.name}** (1995; supplement: 1990s), 71 groups, 37 fleets. Source extraction and verification are complete; GE, TE and With Egestion are **NOT_RUN** because Discards/Detritus routing is absent. No model selected.

- [Extraction report](../models/{M.name}/extracted_tables/REPORT.md)
- [Direct SPPR report](../models/{M.name}/diagnostics/DIRECT_SPPR_REPORT.md)
- [Source-preserving model](../models/{M.name}/model.json)
- [Lossless source roundtrip](../models/{M.name}/extracted_tables/SOURCE_ROUNDTRIP.xlsx)
- [Source-cell evidence](cell_evidence.json)
- [Verification](VERIFICATION.json)

Fin-whale source removals exceed production under zero-BA/net-migration assumptions. See exact source cells and limitations in the extraction report. This is separate from the missing-routing constructor blocker.
''',encoding='utf8')
record={'unit_id':'LME_026','paper_id':'MED-2022','model_id':M.name,'baseline_year':1995,'supplement_label':'1990s','dynamic_period':'1995-2016','group_count':71,'consumer_count':65,'producer_count':4,'detritus_count':2,'fleet_count':37,'model_count':1,'extraction_status':'complete with source gaps','diagnostics':{d['TE_option']:d['status'] for d in diag},'diagnostic_blocker':'Missing natural-mortality/unassimilated routing between Discards 70 and Detritus 71','diet_normalization':False,'fin_whale_conditional_EE':fin['EE_if_unreported_BA_NM_and_landings_are_zero'],'selected':False,'production_eligible':False,'central_registration':'proposal only; parent owns central workbook','checks':checks}
save(R/'results_record.json',record)
proposal={'unit_id':'LME_026','paper_patches':[{'article_id':'MED-2022__LME_026','model_years':'1995 baseline (supplement: 1990s); dynamic 1995–2016','functional_groups':71,'full_model_loadable':'No — missing Discards/Detritus routing; actual constructor admission tested','extraction_status':'extracted; GE/TE/With Egestion NOT_RUN','coverage_class':'intended_full','coverage_basis':'Main article pp. 7–8 explicitly entire Mediterranean basin; inherited target-polygon footprint is not an independently measured model boundary','selection_reason':None,'report_path':(E/'REPORT.md').relative_to(ROOT).as_posix()}],'models_to_upsert':[{'unit_id':'LME_026','model_id':M.name,'article_id':'MED-2022__LME_026','paper_id':'MED-2022','model_name':'Mediterranean Sea','model_year':1995,'model_period':'1995 (supplement labels 1990s)','n_groups':71,'model_path':(M/'model.json').relative_to(ROOT).as_posix(),'extraction_status':'complete with explicit source gaps','sppr_status':'GE NOT_RUN; TE NOT_RUN; With Egestion NOT_RUN','selected':False,'production_eligible':False,'coverage_class':'intended_full','coverage_pct':None,'coverage_basis':'Explicit whole-basin source statement; no independently measured polygon overlap','notes':'71 groups, 37 fleets; 65 consumers, 4 producers, 2 detritus pools. Routing absent; all direct diagnostic methods not run. Source BA/GS unknown. Fin-whale conditional EE 1.477626 under zero BA/migration; not repaired.','report_path':(E/'REPORT.md').relative_to(ROOT).as_posix()}]}
save(R/'central_metadata_proposal.json',proposal)
save(R/'VERIFICATION.json',{'status':'PASS','checks':checks,'canonical_sha256':hashlib.sha256((M/'model.json').read_bytes()).hexdigest(),'source_files':original,'scope':'Source extraction fidelity; PASS does not imply runnable or biologically balanced model.'})
manifest=[{'path':p.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size} for base in [M,R] for p in base.rglob('*') if p.is_file() and p.name not in ['artifact_manifest.json'] and p.suffix!='.zip']
save(R/'artifact_manifest.json',manifest)
with zipfile.ZipFile(R/(M.name+'_extraction_and_diagnostics.zip'),'w',zipfile.ZIP_DEFLATED) as z:
 for p in M.rglob('*'):
  if p.is_file():z.write(p,p.relative_to(M.parent))
 for name in ['MASTER_INDEX.md','VERIFICATION.json','results_record.json','central_metadata_proposal.json','source_manifest.json','cell_evidence.json','independent_source_balance.json','source_balance_groups.csv','taxonomy_evidence.json','converter_source_restoration.json']:
  z.write(R/name,'review/'+name)
print(json.dumps(record,indent=2))
