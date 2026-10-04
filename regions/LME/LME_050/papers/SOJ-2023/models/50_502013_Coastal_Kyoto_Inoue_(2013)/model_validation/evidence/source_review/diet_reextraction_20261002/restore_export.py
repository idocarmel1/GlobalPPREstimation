"""Prepare a bounded source-faithful export; apply only after coordinator release."""
from pathlib import Path
from decimal import Decimal
import sys, json, hashlib, shutil
import openpyxl

OUT=Path(__file__).parent
ROOT=next(p for p in OUT.parents if (p/'Project.xlsx').exists())
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p): return json.loads(p.read_text(encoding='utf-8'))
def dump(p,x): p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def snapshot(p):
    w=openpyxl.load_workbook(p,data_only=False)
    cells={s.title+'!'+c.coordinate:c.value for s in w for row in s for c in row if c.value is not None}
    w.close();return cells

plan_path=OUT/'export_restoration_plan.json'
candidate=OUT/'source_faithful_reconstructed_candidate.xlsx'
if len(sys.argv)>1 and sys.argv[1]=='apply':
    plan=read(plan_path); target=ROOT/plan['target_path']
    assert sha(target)==plan['original_sha256'],'Concurrent export change; stop'
    selected=ROOT/plan['selected_path']
    assert sha(selected)==plan['selected_sha256'],'Concurrent selected model change; stop'
    assert sha(candidate)==plan['candidate_sha256']
    archive=OUT/'input_archive'/sha(target)/target.name
    assert archive.exists() and sha(archive)==sha(target)
    protected=read(OUT/'protected_before.json')
    changes_before={p:sha(ROOT/p) for p in protected if p!=plan['target_path']}
    shutil.copyfile(candidate,target)
    assert sha(target)==plan['candidate_sha256']
    assert changes_before=={p:sha(ROOT/p) for p in changes_before}
    plan.update(applied=True,after_sha256=sha(target),protected_other_files_unchanged=True)
    dump(OUT/'export_restoration_receipt.json',plan)
    receipt=read(OUT/'root_receipt.json')
    receipt['outcome']='verified_unnormalized_canonical_unchanged_normalized_export_restored'
    receipt['export_restoration']=dict(path=plan['target_path'],before_sha256=plan['original_sha256'],after_sha256=sha(target),changed_diet_import_cells=plan['diet_import_cells'],changed_sum_cells=plan['sum_cells'])
    receipt['evidence_paths'] += [str(p.relative_to(ROOT)).replace('\\','/') for p in [OUT/'export_restoration_receipt.json',plan_path,OUT/'restore_export.py']]
    receipt['precise_blockers']=[b for b in receipt['precise_blockers'] if not b.startswith('Historical reconstructed converter export')]
    receipt['protected_files_unchanged']=True
    receipt['protected_export_exception']=plan['target_path']
    dump(OUT/'root_receipt.json',receipt)
    print('Applied',len(plan['diet_import_cells']),'diet/import cells and',len(plan['sum_cells']),'summary sum cells')
else:
    diff=read(OUT/'historical_export_differences.json'); target=ROOT/diff['path']
    selected=ROOT/read(OUT/'root_receipt.json')['selected_path']
    original_hash=sha(target); original=snapshot(target)
    w=openpyxl.load_workbook(target);s=w['Diet composition']
    diet_changes=[];sum_changes=[]
    for d in diff['changed_cells']:
        cell=s.cell(d['row'],d['column']);before=cell.value;cell.value=float(Decimal(d['source_literal']))
        diet_changes.append(dict(cell=cell.coordinate,before=before,after=cell.value,consumer_id=d['consumer_id'],prey_id=d['prey_id'],source_literal=d['source_literal']))
    sums=read(OUT/'consumer_sums.json');headers={str(c.value):c.column for c in s[1]}
    sumrow=next(r[0].row for r in s if str(r[0].value).lower()=='sum')
    for d in sums:
        cell=s.cell(sumrow,headers[d['consumer_id']]);after=float(Decimal(d['raw_prey_plus_import']))
        if abs(cell.value-after)>1e-12:
            sum_changes.append(dict(cell=cell.coordinate,before=cell.value,after=after,consumer_id=d['consumer_id']))
            cell.value=after
    w.save(candidate);w.close()
    revised=snapshot(candidate)
    actual={k:(original.get(k),revised.get(k)) for k in set(original)|set(revised) if original.get(k)!=revised.get(k)}
    expected={'Diet composition!'+d['cell']:(d['before'],d['after']) for d in diet_changes+sum_changes}
    assert actual==expected,(actual,expected)
    assert sha(target)==original_hash
    plan=dict(target_path=diff['path'],original_sha256=original_hash,candidate_sha256=sha(candidate),selected_path=str(selected.relative_to(ROOT)).replace('\\','/'),selected_sha256=sha(selected),
        diet_import_cells=diet_changes,sum_cells=sum_changes,total_changed_cells=len(actual),all_other_workbook_values_unchanged=True,
        detritus_fate_unchanged=True,applied=False,canonical_reextraction=False,
        limitation='Reconstructed Excel remains a lossy view: missing diet/fate cells are represented as zero; exact source JSON/CSV and cell ledger remain authoritative.')
    dump(plan_path,plan)
    print('Prepared',len(diet_changes),'diet/import cells and',len(sum_changes),'sum cells; original unchanged')
