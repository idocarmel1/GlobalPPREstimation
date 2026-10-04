from pathlib import Path
import hashlib,json,zipfile
from docx import Document
HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[3]
source=ROOT/'tools/templates/Model_validation_template.docx';doc=Document(source)
out=HERE/'qa';out.mkdir(exist_ok=True)
data={'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sections':[{'width_inches':s.page_width.inches,'height_inches':s.page_height.inches,'margins_inches':{x:getattr(s,x).inches for x in ['top_margin','bottom_margin','left_margin','right_margin']}} for s in doc.sections],
 'paragraphs':[{'text':p.text,'style':p.style.name} for p in doc.paragraphs],
 'tables':[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]}
with zipfile.ZipFile(source) as z:data['package_hashes']={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
(out/'template_audit.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
(out/'artifact.md').write_text('# Template contract\n\nReference: tools/templates/Model_validation_template.docx; SHA256 '+data['sha256']+'\n\nOne A4 portrait section, 0.65-inch margins; Calibri11 body, title23, heading1 17, heading2 12. Standard field/entry table10.5-point Calibri, columns1.48/5.49inch. Preserve styles, section geometry and package chrome. Edit the source copy table cells by field label; preserve manual SPPR calculation, Open issues and Review cells. Replace automatic slots; clone coverage table styles for repeated rows. Coverage outside main grid: five confidence rows, separate membership and allocation rule summaries, exact Very low taxa. Replace geographic placeholders with labeled context figures; links relative blue underlined. Do not insert signature or researcher verdict. Template audit JSON retains each edit slot, package inventory and untouched structures. Final rendering may increase pagination naturally; keep widths/fonts and repeated headers. Source remains unchanged.\n',encoding='utf-8')
print(json.dumps({'tables':data['tables'],'paragraphs':data['paragraphs']},ensure_ascii=False))
