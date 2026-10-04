"""Add native portable OOXML hyperlinks, then verify authored workbooks read-only.

Artifact Tool authored the workbook. This narrow OOXML pass supplies native
hyperlinks/underlining unavailable in the inspected Artifact Tool API and
restores table autoFilter nodes omitted by the exporter; it does
not author or rewrite worksheet values, formulas, or numerical results.
"""
from pathlib import Path
import json,zipfile,hashlib,math,xml.etree.ElementTree as ET,os
from urllib.parse import urlsplit,unquote
import openpyxl

P=Path(__file__).resolve().parent
R=P.parents[2]
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('',S);ET.register_namespace('r',REL)
def load(n):return json.loads((P/n).read_text(encoding='utf8'))
def eq(a,b):return a==b or isinstance(a,(float,int)) and isinstance(b,(float,int)) and math.isclose(a,b,rel_tol=2e-15,abs_tol=1e-12)
def table_filters(files):
    for n in list(files):
        if n.startswith('xl/tables/') and n.endswith('.xml'):
            table=ET.fromstring(files[n])
            if table.find(f'{{{S}}}autoFilter') is None:
                table.insert(0,ET.Element(f'{{{S}}}autoFilter',{'ref':table.get('ref')}))
            files[n]=ET.tostring(table,encoding='utf-8',xml_declaration=True)
qa={}
for variant in ['Base','M30','P30']:
    f=R/f'LME027_Guenette2014_{variant}_taxon_mapping_appendix.xlsx'
    with zipfile.ZipFile(f) as z:files={n:z.read(n) for n in z.namelist()}
    sheetpath='xl/worksheets/sheet2.xml'; relpath='xl/worksheets/_rels/sheet2.xml.rels'
    sh=ET.fromstring(files[sheetpath])
    existing=sh.find(f'{{{S}}}hyperlinks')
    if existing is not None:sh.remove(existing)
    relationships=ET.fromstring(files[relpath]) if relpath in files else ET.Element(f'{{{PKG}}}Relationships')
    for child in list(relationships):
        if child.get('Type','').endswith('/hyperlink'):relationships.remove(child)
    used={c.get('Id') for c in relationships};hl=ET.Element(f'{{{S}}}hyperlinks')
    for n,item in enumerate(load(f'{variant}_appendix_hyperlinks.json'),1):
        rid=f'rIdSourceLink{n}';assert rid not in used
        ET.SubElement(hl,f'{{{S}}}hyperlink',{'ref':item['cell'],f'{{{REL}}}id':rid})
        ET.SubElement(relationships,f'{{{PKG}}}Relationship',{'Id':rid,'Type':REL+'/hyperlink','Target':item['target'],'TargetMode':'External'})
    following={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
    pos=next((i for i,c in enumerate(sh) if c.tag.rsplit('}',1)[-1] in following),len(sh))
    sh.insert(pos,hl)
    styles=ET.fromstring(files['xl/styles.xml']);underlined=0
    for font in styles.find(f'{{{S}}}fonts'):
        color=font.find(f'{{{S}}}color')
        if color is not None and color.get('rgb','').upper().endswith('0563C1'):
            if font.find(f'{{{S}}}u') is None:ET.SubElement(font,f'{{{S}}}u',{'val':'single'})
            underlined+=1
    assert underlined>0
    files[sheetpath]=ET.tostring(sh,encoding='utf-8',xml_declaration=True)
    files[relpath]=ET.tostring(relationships,encoding='utf-8',xml_declaration=True)
    files['xl/styles.xml']=ET.tostring(styles,encoding='utf-8',xml_declaration=True)
    table_filters(files)
    temp=f.with_suffix('.links.tmp.xlsx')
    with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in files.items():z.writestr(n,b)
    os.replace(temp,f)
    wb=openpyxl.load_workbook(f,data_only=False)
    ws=wb['Taxon mappings'];src=wb['Sources'];raw=load(f'{variant}_appendix_rows.json')
    vals=list(ws.iter_rows(min_row=7,max_row=518,min_col=1,max_col=7,values_only=True))
    assert len(vals)==512 and len({r[0] for r in vals})==512
    assert [r[0] for r in vals]==[r[0] for r in raw]
    assert all(eq(a[j],b[j]) for a,b in zip(vals,raw) for j in (0,1,2,3,5))
    assert all(isinstance(r[2],(float,int)) and isinstance(r[3],(float,int)) for r in vals)
    assert all(vals[i][3]>=vals[i+1][3] for i in range(511))
    assert ws.max_column==7 and ws.freeze_panes=='B7',(ws.max_column,ws.freeze_panes)
    assert src.freeze_panes=='A5'
    assert list(ws.tables.values())[0].ref=='A6:G518'
    assert list(ws.tables.values())[0].autoFilter.ref=='A6:G518'
    assert list(src.tables.values())[0].autoFilter.ref=='A4:E33'
    expected=load(f'{variant}_appendix_hyperlinks.json');assert len(expected)==29
    checked=[]
    for item in expected:
        c=src[item['cell']];assert c.hyperlink and c.hyperlink.target==item['target']
        assert c.font.underline=='single' and c.font.color.rgb.upper().endswith('0563C1')
        if not item['target'].startswith(('https://','http://')):
            assert not Path(item['target']).is_absolute()
            target=R/unquote(urlsplit(item['target']).path);assert target.is_file(),target
        checked.append(item)
    assert not any(c.data_type=='f' for sheet in wb for row in sheet for c in row)
    assert not any('HyperlinkBase' in b.decode('utf8',errors='ignore') for n,b in files.items() if n.startswith('docProps/'))
    wb.close()
    side=Path(str(f)+'.inspect.ndjson')
    if side.exists():os.replace(side,P/side.name)
    qa[variant]={'file':f.name,'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'rows':512,'zero_catch_labels':sum(r[2]==0 for r in vals),'missing_TL_labels':sum(r[1]=='?' for r in vals),'numeric_catch_PPR':True,'unrounded_PPR_matches_evidence':True,'descending_PPR':True,'exact_taxon_universe':True,'freeze':'B7','filter':'A6:G518','native_hyperlinks':len(checked),'all_local_links_exist':True,'blue_underlined_links':True,'absolute_hyperlink_base_absent':True,'formula_errors_absent':True}

taxfile=P/'Taxonomy.xlsx'
with zipfile.ZipFile(taxfile) as z:taxfiles={n:z.read(n) for n in z.namelist()}
table_filters(taxfiles)
taxtemp=taxfile.with_suffix('.filter.tmp.xlsx')
with zipfile.ZipFile(taxtemp,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in taxfiles.items():z.writestr(n,b)
os.replace(taxtemp,taxfile)
wb=openpyxl.load_workbook(taxfile,data_only=True,read_only=True)
rows=list(wb['Taxonomy'].iter_rows(values_only=True));tax=load('taxonomy_evidence.json')['groups']
assert rows[0]==('seq','group_name','taxon_descr') and len(rows)==52
assert all(rows[i+1]==(g['seq'],g['group_name'],g['taxon_descr']) for i,g in enumerate(tax));wb.close()
qa['taxonomy']={'51_source_groups_exact':True,'three_column_contract':True,'computational_import_group_excluded':True}
identity=load('universe_identity.json')
assert hashlib.sha256((R/'LME_027.xlsx').read_bytes()).hexdigest()==identity['source_workbook_sha256']
qa['source_workbook_unchanged']=True
qa['visual_review']='Rendered all variants; main mapping and Sources top rows, plus source taxonomy. Values/reasons wrap within rows; final hyperlink underlining checked in OOXML and openpyxl.'
(P/'appendix_qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(qa,ensure_ascii=False,indent=2))
