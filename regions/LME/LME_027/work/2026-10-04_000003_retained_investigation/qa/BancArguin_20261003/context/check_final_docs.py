from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import unquote
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import json,zipfile,shutil,hashlib
context=Path(__file__).resolve().parent
run=context.parent
root=context.parents[4]
region=root/'regions/LME_027'
mapping=json.loads((run/'mapping/report_mapping_text.json').read_text(encoding='utf8'))
expected=Counter(t for r in mapping['very_low_decisions'] for t in r['taxa'])
template=Document(root/'tools/templates/Model_validation_template.docx')
manual_names=['SPPR calculation','Open issues and next action','Review and reproducibility']
template_manual={r.cells[0].text:etree.tostring(r.cells[1]._tc) for r in template.tables[0].rows[1:] if r.cells[0].text in manual_names}
qa=[]
for variant in ['Base','M30','P30']:
    path=region/f'Model_validation_Guenette2014_BancArguin_{variant}_1991.docx'
    doc=Document(path)
    decision_tables=[t for t in doc.tables if t.rows[0].cells[0].text=='Affected taxa']
    actual=Counter(t for table in decision_tables for r in table.rows[1:] for t in r.cells[0].text.split('; '))
    assert actual==expected,(actual-expected,expected-actual)
    assert sum(actual.values())==68
    manual={r.cells[0].text:etree.tostring(r.cells[1]._tc) for r in doc.tables[0].rows[1:] if r.cells[0].text in manual_names}
    assert manual==template_manual
    with zipfile.ZipFile(path) as archive:
        rels=etree.fromstring(archive.read('word/_rels/document.xml.rels'))
        links=[e.get('Target') for e in rels if e.get('Type','').endswith('/hyperlink')]
        app=etree.fromstring(archive.read('docProps/app.xml'))
        hyperlink_base=[e.text for e in app.iter() if etree.QName(e).localname=='HyperlinkBase']
    assert not any(hyperlink_base)
    local=[unquote(x.split('#')[0]) for x in links if not x.startswith(('http://','https://','#'))]
    assert all(not Path(x).is_absolute() and ':' not in x for x in local)
    with TemporaryDirectory(prefix='docx_relocation_',dir=run/'qa') as tmp:
        moved_root=Path(tmp)/'relocated_project'
        moved_doc=moved_root/path.relative_to(root)
        moved_doc.parent.mkdir(parents=True)
        shutil.copy2(path,moved_doc)
        for target in local:
            source=(path.parent/target).resolve()
            assert source.is_relative_to(root)
            moved_target=moved_root/source.relative_to(root)
            if source.is_dir():moved_target.mkdir(parents=True,exist_ok=True)
            else:
                moved_target.parent.mkdir(parents=True,exist_ok=True)
                shutil.copy2(source,moved_target)
        assert all((moved_doc.parent/target).exists() for target in local)
    qa.append({'variant':variant,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'very_low_taxa_exactly_once':68,'very_low_grouped_rows':sum(len(t.rows)-1 for t in decision_tables),'manual_cells_byte_identical':True,'hyperlink_base_empty':True,'relative_local_links':len(local),'relocation_copy_check':'PASS','remote_destination_count':len(links)-len(local),'remote_destinations_retrieved_for_evidence':'Publication DOI previously retrieved; formatting verification does not imply reopening every remote destination.'})
(run/'qa/final_docx_acceptance.json').write_text(json.dumps(qa,indent=2),encoding='utf8')
print(json.dumps(qa,indent=2))
