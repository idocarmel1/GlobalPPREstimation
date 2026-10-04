from pathlib import Path
from copy import deepcopy
import json,zipfile,io,hashlib,sys,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

out=Path(__file__).resolve().parent
run=out.parent;root=out.parents[4];region=root/'regions/LME_027'
sys.path.insert(0,str(root/'tools'))
from validation_percentage_format import format_percent,verify_report
template=root/'tools/templates/Model_validation_template.docx'
assert hashlib.sha256(template.read_bytes()).hexdigest()==json.loads((run/'qa/template_audit.json').read_text())['sha256']
summary=json.loads((run/'mapping/reference_summaries.json').read_text(encoding='utf8'))
mapping_text=json.loads((run/'mapping/report_mapping_text.json').read_text(encoding='utf8'))
exp=json.loads((run/'runtime/Base_runtime_PQ_approximation/summary.json').read_text(encoding='utf8'))
baseprefix='validation_reports/BancArguin_20261003/'
manual=['SPPR calculation','Open issues and next action','Review and reproducibility']
qa=[]

def text(p,txt,bold=False,size=None):
 r=p.add_run(txt);r.bold=bold
 if size:r.font.size=Pt(size)
 return r
def hyperlink(p,label,target,size=10.5):
 rid=p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 for tag,atts in [('rStyle',{'val':'Hyperlink'}),('color',{'val':'0563C1'}),('u',{'val':'single'}),('sz',{'val':str(round(size*2))})]:
  e=OxmlElement('w:'+tag)
  for k,v in atts.items():e.set(qn('w:'+k),v)
  pr.append(e)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def clear(p):
 for c in list(p._p):
  if c.tag!=qn('w:pPr'):p._p.remove(c)
def para(p,txt='',size=None):
 clear(p);text(p,txt,size=size)
def cell(c,lines,links=None):
 for e in list(c._tc):
  if e.tag!=qn('w:tcPr'):c._tc.remove(e)
 for line in lines:
  p=c.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=False;text(p,line,size=10.5)
 if links:
  p=c.add_paragraph();p.paragraph_format.space_after=Pt(3)
  for i,(label,target) in enumerate(links):
   if i:text(p,' | ',size=10.5)
   hyperlink(p,label,target)
def table_rows(table,rows):
 proto=deepcopy(table.rows[1]._tr)
 for r in list(table._tbl.findall(qn('w:tr')))[1:]:table._tbl.remove(r)
 for vals in rows:
  tr=deepcopy(proto);table._tbl.append(tr);row=table.rows[-1]
  for c,txt in zip(row.cells,vals):cell(c,[str(txt)])
  pr=tr.get_or_add_trPr();e=OxmlElement('w:cantSplit');pr.append(e)
 for r in table.rows:
  for c in r.cells:
   for p in c.paragraphs:p.paragraph_format.keep_with_next=False
def remove_paragraph(p):p._p.getparent().remove(p._p)

for variant in ['Base','M30','P30']:
 d=Document(template);ps=list(d.paragraphs);ts=list(d.tables)
 cells={r.cells[0].text:r.cells[1] for r in ts[0].rows[1:]}
 manuals={k:etree.tostring(cells[k]._tc) for k in manual}
 p=ps[0];para(p,f'Canary Current {variant} model validation');p.style='Title'
 for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
 # The current template intentionally has no introductory prose.
 if variant=='Base':ps[2].paragraph_format.page_break_before=True
 model=f'models/Guenette2014_BancArguin_{variant}_1991/'
 cell(cells['Region'],['Canary Current (LME 027)'],[('Regional workbook','LME_027.xlsx')])
 cell(cells['Catch source'],['Sea Around Us retained regional dataset, 1950–2019; landings reference year 2019. Provider revision unspecified.'],[('Catch archive','raw/LME_027-catch.zip'),('Catch and coefficient context',baseprefix+'mapping/reference_summaries.json')])
 cell(cells['Selected article'],['Candidate article: Guénette S, Meissa B, Gascuel D (2014). Assessing the Contribution of Marine Protected Areas to the Trophic Functioning of Ecosystems: A Model for the Banc d’Arguin and the Mauritanian Shelf. PLoS ONE 9(4), e94742.'],[('DOI','https://doi.org/10.1371/journal.pone.0094742'),('Article PDF','papers/CAN-2014/file-e30dfe50.pdf'),('Supplement','papers/CAN-2014/pone.0094742.s001-eb18ca66.docx')])
 cell(cells['Other known articles'],['Morissette, Melgo, Kaschner, Gerber & Bamy (2009): Northwest African food-web model; the currently selected EcoBase 118 version represents 1987.','Villanueva, Tito-de-Morais, Weigel & Moreau (2004): An Ecopath model of the Sine-Saloum Delta biosphere reserve (Senegal).'],[('Central article records','../../Project.xlsx'),('Retained source folders','papers')])
 cell(cells['Selected model'],[f'Candidate: Guenette2014_BancArguin_{variant}_1991. 51 source groups, 47 consumers and 3 reported fleets; Mauritanian shelf to 200 m, approximately 33,224 km². Baseline 1991; Ecosim calibration 1991–2006.'],[('Canonical model JSON',model+'model.json'),('Source extraction',model+'extracted_tables/REPORT.md')])
 other={'Base':'M30 lowers assumed subtidal benthos density by 30%; P30 raises it by 30%. Both change Banc feeding proportions and selected balanced biomasses.','M30':'Base assumes equal subtidal and intertidal benthos density; P30 raises assumed subtidal density by 30%. The alternatives alter Banc feeding proportions and selected balanced biomasses.','P30':'Base assumes equal subtidal and intertidal benthos density; M30 lowers assumed subtidal density by 30%. The alternatives alter Banc feeding proportions and selected balanced biomasses.'}[variant]
 cell(cells['Other models from the same article'],[other,'All three share the 1991 shelf domain. EcoBase 689 is a separately deposited numerical version, with material differences from the printed tables.'])
 cell(cells['Selection rationale'],['Article – ? Candidate review requested; no recorded decision selects this article.','Model – ? No decision selects this variant. Overview retains Northwest Africa 1987 after the earlier comparison recorded failures for the Banc d’Arguin configurations.'],[('Current selection record','LME_027.xlsx'),('Exact recorded rationale',baseprefix+'context/workbook_context.json')])
 extraction=['Table 1 and S8 Base conflict; for BA phytoplankton, EE is 0.260 versus 0.40. Both values and their table contexts remain documented.','S2 prints Groupers ad → BA L crustaceans as 1.34 (diet sum 2.2006); Coastal birds sum to 0.980331. Four blank cells remain unknown. The source matrix was not repaired.','Multi-stanza Z is distinct from P/B; missing P/B and growth/transition details are retained separately. No discard observations were reported; unknowns remain distinct from zero.']
 if variant!='Base':extraction.insert(1,f'{variant}: S8 supplies scenario biomass, EE and aggregate Banc-feeding proportions, but leaves 114 diet cells across 14 consumers unpublished. Complete scenario diet/import matrices were not recovered.')
 cell(cells['Model extraction'],extraction,[('Source conflicts',model+'extracted_tables/SOURCE_CONFLICTS.json'),('Diet cell evidence',model+'extracted_tables/DIET_CELL_LEDGER.json'),('Online recovery',baseprefix+'context/context_summary.json')])
 for method in ['GE','TE']:
  if variant=='Base':
   m=exp[method];dv=m['divergence']
   lines=['Source input: NOT_RUN. Five published diet rows fail the admission check.','Separate runtime experiment: FAIL. No negative SPPR entries across the basal-source columns; all 52 × 5 entries are finite.',f"rho_living = {dv['rho_living']:.4f}" ]
   if method=='GE':lines.append(f"b = {dv['b']:.4f}")
   lines += [f"Detritus (51): SPPR = {dv['sppr_det']['51']:.4f}.",'Strict model-input balance fails; numerical convergence does not validate the source.']
   links=[('Source admission',baseprefix+'runtime/Base_source_admission/summary.json'),('Runtime experiment',baseprefix+'runtime/Base_runtime_PQ_approximation/summary.json')]
  else:
   lines=[f'NOT_RUN. The {variant} scenario lacks 114 diet cells across 14 consumers. No scenario-specific coefficient or per-source SPPR matrix is available.','Negative SPPR, rho_living'+(', b' if method=='GE' else '')+' and Detritus (51) SPPR: not determined.']
   links=[('Source completeness',model+'extracted_tables/EXTRACTION_STATUS.json')]
  cell(cells[method+' diagnostics'],lines,links)
 cell(cells['Geographic fit'],['A. Region covered by study: approximately 3.0% (about 2.5–3.5%).','B. Study area covered by region: approximately 95–100%.'])
 cell(cells['Temporal fit'],['The baseline represents 1991 and calibration spans 1991–2006; the mapping reference is 2019. Transferring fixed coefficients and source catch/biomass proportions across years and across the much larger LME is an explicit approximation.'])
 other=['The currently selected Northwest Africa 1987 model was disqualified by Ido Carmel on 02/10/2026: “Reason – too strong living compartments recycling due to near-zero EE values (rho_living = 0.99).” This is a separate model’s review.']
 if variant=='Base':
  other.insert(0,'The experiment normalizes a runtime diet copy and approximates missing P/B as Q/B × P/Q; printed P/B = 0 remains unchanged. It is not the published Base parameterization. With Egestion also FAILS. EcoBase 689 also fails the flat-model checks and does not establish source validation.')
  other.insert(1,'Diagnostic-FAIL experiment, 2019 landings, all sources: GE PPR 105.84 million t C, TE 1,362.76 million t C and With Egestion 55.02 million t C; respectively 22%, 290% and 12% of ensemble NPP (472.96 million t C). No validated source-model estimate is available.')
 else:other.insert(0,'Base or EcoBase 689 coefficients cannot stand in for this scenario. The independent simple-chain PPR and catch-mapping assessment below do not require scenario SPPR.')
 cell(cells['Other'],other,[('Regional assessment',baseprefix+'regional_comparison.json'),('Independent native-deposit audit',baseprefix+'runtime/EcoBase689_native_companion/summary.json')])
 s=summary[variant];vl=next(r for r in s['confidence'] if r['confidence']=='Very low')
 para(ps[3],'Reference: 2019 landings; all retained regional catch labels, including zero landings.')
 para(ps[4],'51 source groups; 512 catch taxa. Computational import groups are separate.')
 para(ps[5],f"Simple-chain PPR totals {s['simple_chain_ppr_tC']:,.0f} t C in 2019.")
 para(ps[6],'100 labels lack a saved TL/coefficient. All have zero recorded 2019 landings and contribute zero annual PPR; the annual total is complete.')
 para(ps[7],'Method: independent simple trophic chain, using retained taxon coefficients.')
 para(ps[8],'');hyperlink(ps[8],'Excel taxon appendix and descriptive Sources sheet',f'LME027_Guenette2014_{variant}_taxon_mapping_appendix.xlsx',11)
 table_rows(ts[1],[[r['confidence'],r['taxa_n'],format_percent(r['catch_pct']),format_percent(r['ppr_pct'])] for r in s['confidence']])
 table_rows(ts[2],[[mapping_text['rule_names'][r['rule']],r['confidence'],format_percent(r['ppr_pct'])] for r in s['membership_rules']])
 table_rows(ts[3],[[mapping_text['rule_names'][r['rule']].replace('Source-model1991','Source model 1991'),r['confidence'],format_percent(r['ppr_pct'])] for r in s['allocation_rules']])
 para(ps[11],'Membership evidence: source Tables S1, S4 and S5, documented guild definitions and regional taxonomic evidence. ');hyperlink(ps[11],'Source supplement','papers/CAN-2014/pone.0094742.s001-eb18ca66.docx',11)
 para(ps[13],'Weights and assumptions: source 1991 catch proportions; complete model biomass where candidate catch totals are zero. The spatial, temporal and life-stage transfer is assumed. Juvenile source zeros do not show that later catches or discards contain no juveniles. ');hyperlink(ps[13],'Allocation evidence',baseprefix+'mapping/reference_summaries.json',11)
 para(ps[15],f"No unresolved mappings remain. The 68 Very low assignments account for {format_percent(vl['catch_pct'])} of catch and {format_percent(vl['ppr_pct'])} of simple-chain PPR.")
 lowrows=[]
 for r in mapping_text['very_low_decisions']:
  reason=r['reason'].replace('so those are removed for that label below.','which are excluded from its candidates.')
  lowrows.append(['; '.join(r['taxa']),reason])
 table_rows(ts[4],lowrows)
 # Keep the long decision list readable in Word PDF export. Explicit table
 # continuations retain the same header where Word omits its repeated header.
 remaining=[deepcopy(r._tr) for r in ts[4].rows[7:]]
 for r in list(ts[4]._tbl.findall(qn('w:tr')))[7:]:ts[4]._tbl.remove(r)
 preceding=ts[4]._tbl
 for chunk in [remaining[:9],remaining[9:]]:
  continuation=deepcopy(ts[4]._tbl)
  for r in list(continuation.findall(qn('w:tr')))[1:]:continuation.remove(r)
  for r in chunk:continuation.append(r)
  boundary=OxmlElement('w:p');props=OxmlElement('w:pPr')
  e=OxmlElement('w:pageBreakBefore');props.append(e)
  e=OxmlElement('w:spacing');e.set(qn('w:before'),'0');e.set(qn('w:after'),'0');e.set(qn('w:line'),'1');e.set(qn('w:lineRule'),'exact');props.append(e)
  boundary.append(props);preceding.addnext(boundary);boundary.addnext(continuation);preceding=continuation
 para(ps[18],'Geographic evidence');ps[18].paragraph_format.page_break_before=True
 para(ps[19],'The study covers a narrow Mauritanian shelf portion of the Canary Current LME. Approximate coverage accounts for land exclusion and the shelf’s 200 m offshore boundary. ');hyperlink(ps[19],'Boundary and area evidence',baseprefix+'context/geographic_assessment.json',11)
 para(ps[21],'Canary Current LME 27 in blue; approximate study shelf in orange. Project LME boundary and Natural Earth 1:50m land mask. ');hyperlink(ps[21],'Regional comparison',baseprefix+'context/geographic_comparison.png',11)
 cell(ts[5].rows[1].cells[0],['']);p=ts[5].rows[1].cells[0].paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(out/'target_region.png'),width=Inches(5.9))
 ps[22].paragraph_format.page_break_before=True
 para(ps[23],'Guénette, Meissa & Gascuel (2014), Figure 1, PDF page 3. The study area is the shelf shallower than 200 m, approximately 33,224 km² including the 6,450 km² Banc. The solid black offshore line marks the larger EEZ. ');hyperlink(ps[23],'Original figure','papers/CAN-2014/Figure_1-78156383.tif',11)
 cell(ts[6].rows[1].cells[0],['']);p=ts[6].rows[1].cells[0].paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(out/'source_figure1.png'),width=Inches(5.6))
 remove_paragraph(ps[24]);remove_paragraph(ps[17])
 # Replace only edited OOXML and new image parts, preserving styles and all package chrome byte-for-byte.
 for k in manual:assert etree.tostring(cells[k]._tc)==manuals[k],k
 b=io.BytesIO();d.save(b)
 with zipfile.ZipFile(template) as z:original={n:z.read(n) for n in z.namelist()}
 with zipfile.ZipFile(io.BytesIO(b.getvalue())) as z:generated={n:z.read(n) for n in z.namelist()}
 changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
 dest=region/f'Model_validation_Guenette2014_BancArguin_{variant}_1991.docx'
 with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
  for name,data in original.items():z.writestr(name,generated[name] if name in changed else data)
  for name,data in generated.items():
   if name not in original:z.writestr(name,data)
 verify_report(dest)
 with zipfile.ZipFile(dest) as z:
  preserved=all(z.read(n)==v for n,v in original.items() if n not in changed)
  rels=etree.fromstring(z.read('word/_rels/document.xml.rels'));links=[]
  for e in rels:
   if e.get('Type','').endswith('/hyperlink'):
    target=e.get('Target');local=not target.startswith(('https://','http://','#'));exists=(dest.parent/target.split('#')[0]).exists() if local else None
    links.append({'target':target,'local':local,'exists':exists})
  doc=etree.fromstring(z.read('word/document.xml'));ns={'w':qn('w:x').split('}')[0][1:]};bad=[]
  for h in doc.findall('.//w:hyperlink',ns):
   for r in h.findall('w:r',ns):
    if r.find('w:rPr/w:color',ns).get(qn('w:val'))!='0563C1' or r.find('w:rPr/w:u',ns).get(qn('w:val'))!='single':bad.append(etree.tostring(r).decode())
 qa.append({'variant':variant,'path':dest.relative_to(root).as_posix(),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'preserve_only_parts_unchanged':preserved,'manual_cells_xml_preserved':True,'percentage_format_valid':True,'links':links,'blue_underlined_links':not bad,'very_low_taxa_n':sum(len(x['taxa']) for x in mapping_text['very_low_decisions'])})
(run/'qa/docx_structure_verification.json').write_text(json.dumps(qa,indent=2),encoding='utf8')
print(json.dumps([{'variant':x['variant'],'preserve_only_parts':x['preserve_only_parts_unchanged'],'missing_links':[y for y in x['links'] if y['local'] and not y['exists']]} for x in qa],indent=2))
