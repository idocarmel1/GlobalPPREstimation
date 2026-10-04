from pathlib import Path
import json,hashlib,xml.etree.ElementTree as E
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[5]
out=Path(__file__).resolve().parent
rec=[]
for p in [root/'regions/LME_027/LME_027.xlsx',root/'Project.xlsx']:
 w=load_workbook(p,read_only=True,data_only=False)
 result={'file':p.relative_to(root).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':w.sheetnames,'rows':[]}
 for s in w:
  if p.name=='LME_027.xlsx' and s.title!='Overview':continue
  for row in s.iter_rows(max_row=min(s.max_row or 5000,5000),max_col=min(s.max_column or 80,80)):
   cells=[{'cell':c.coordinate,'value':str(c.value) if c.value is not None else None} for c in row if c.value is not None]
   txt=' '.join(str(x['value']) for x in cells)
   if (p.name=='LME_027.xlsx' and s.title=='Overview') or (p.name=='Project.xlsx' and any(t in txt for t in ['LME_027','GuenetteMeissa','CAN-2014','27_118_Northwest_Africa','27_Morissette2009'])):
    result['rows'].append({'sheet':s.title,'cells':cells})
 rec.append(result);w.close()
(out/'workbook_context.json').write_text(json.dumps(rec,ensure_ascii=False,indent=2),encoding='utf8')
native=E.parse(root/'regions/LME_027/papers/CAN-2014/recovery_20261003/ecobase689_input.xml').getroot()
output=E.parse(root/'regions/LME_027/papers/CAN-2014/recovery_20261003/ecobase689_output.xml').getroot()
meta={e.tag:e.text for e in native.find('model_descr')}
stanza=E.tostring(native.find('stanza_descr'),encoding='unicode')
(out/'native_stanzas.xml').write_text(stanza,encoding='utf8')
groups={int(g.findtext('group_seq')):g for g in native.findall('./group_descr/group')}
base=json.loads((root/'regions/LME_027/papers/CAN-2014/extracted/work/base_input.json').read_text(encoding='utf8'))
rows=[];diet=[]
for g in base['groups']:
 n=g['n'];ng=groups[n]
 for field in ['biomass','pb','qb','ee']:
  rows.append({'group_id':n,'paper_name':g['name'],'native_name':ng.findtext('group_name'),'field':field,'Table1_retained_transcription':g.get(field),'Table1_Z':g.get('z') if field=='pb' else None,'native_input':ng.findtext(field)})
 nd={int(e.findtext('prey_seq')):e.findtext('proportion') for e in ng.findall('./diet_descr/diet')}
 for prey,pval in base['diet'].get(str(n),{}).items():
  nv=ng.findtext('diet_imp') if prey=='import' else nd.get(int(prey),'0')
  if float(nv or 0)!=float(pval or 0):diet.append({'consumer_id':n,'consumer_name':g['name'],'prey_id':prey,'S2_retained_transcription':pval,'native':nv,'difference':float(nv or 0)-float(pval or 0)})
fleets=[]
for f in native.findall('./fleet_descr/fleet'):
 attrs={e.tag:e.text for e in f if len(e)==0}; cs=[{e.tag:e.text for e in c} for c in f.findall('./catch_descr/catch')];attrs['catch_types']=sorted(set(c['catch_type'] for c in cs));attrs['nonzero_catches']=[c for c in cs if c['catch_type'] not in ['market','prop mort'] and float(c['catch_value'])!=0];fleets.append(attrs)
(out/'native689_signature_comparison.json').write_text(json.dumps({'metadata':meta,'native_groups':len(groups),'scalar_comparison':rows,'diet_differences':diet,'fleets':fleets,'provisional_caution':'Compare retained Table1/S2 transcription to originals; native is distinct until adjudicated. Native pb fields for stanza groups correspond printed Z, not proof flat PB. Native calculated/solved fields retain flags.'},indent=2),encoding='utf8')
print(json.dumps({'workbook_rows':[len(x['rows']) for x in rec],'metadata':meta,'diet_different_cells':len(diet),'diet_material_differences':[x for x in diet if abs(x['difference'])>.0001],'fleets':fleets},ensure_ascii=False,indent=2))
