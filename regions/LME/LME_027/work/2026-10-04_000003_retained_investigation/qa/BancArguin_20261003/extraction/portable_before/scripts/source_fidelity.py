"""Lossless source-table and companion retention for the EwE converter.

The eight standard files stay importer-compatible. Companion CSVs are evidence,
never passed off as EwE imports. Unknown cells remain unknown; no normalization.
"""
from pathlib import Path
from decimal import Decimal
import csv,json,hashlib,re
import openpyxl

STANDARD=['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv',
          'Detritus_fate.csv','Biomass_accumulation.csv','TL.xlsx','Metadata.xlsx']
def _nd(value):return '-9999' if value is None or str(value).strip()=='' else str(value)
def _rows(path):
    if path.suffix=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.reader(f))
    wb=openpyxl.load_workbook(path,data_only=False)
    return [list(row) for row in wb.active.iter_rows(values_only=True)]
def _sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def write_companions(model,outdir):
    """Write optional arbitrary source tables separately from EwE imports."""
    outdir=Path(outdir);manifest=[];seen=set()
    old_manifest=outdir/'companions/manifest.json'
    previous=json.loads(old_manifest.read_text(encoding='utf-8')).get('files',[]) if old_manifest.exists() else []
    for filename,spec in model.get('companions',{}).items():
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]*\.csv',filename) or filename.casefold() in {x.casefold() for x in STANDARD} or filename.casefold() in seen:
            raise ValueError('Companion names must be safe CSV basenames distinct from standard imports')
        seen.add(filename.casefold())
        rows=spec.get('rows');description=spec.get('description','').strip()
        if not isinstance(rows,list) or not rows or not description:
            raise ValueError('Each companion requires nonempty rows and a description')
        layout=spec.get('layout','table')
        if layout not in ('table','source_grid'):raise ValueError('Unknown companion layout')
        if any(not isinstance(r,list) for r in rows):raise ValueError('Companion rows must be lists')
        width=max(map(len,rows))
        if layout=='table':
            header=[str(x).strip() for x in rows[0]]
            if not header or any(not x for x in header) or len(set(header))!=len(header) or any(len(r)!=len(header) for r in rows):raise ValueError('Table companion requires rectangular rows and unique nonempty headers')
        else:rows=[r+[None]*(width-len(r)) for r in rows]
        path=outdir/'companions'/filename;path.parent.mkdir(exist_ok=True)
        with path.open('w',encoding='utf-8',newline='') as f:
            csv.writer(f,lineterminator='\r\n').writerows(rows)
        manifest.append({'filename':filename,'description':description,'layout':layout,'role':'source_evidence_not_EwE_import','sha256':_sha(path)})
    for old in previous:
        name=old['filename']
        if name.casefold() in seen:continue
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]*\.csv',name):raise ValueError('Unsafe prior companion manifest')
        stale=outdir/'companions'/name
        if stale.exists():
            if _sha(stale)!=old['sha256']:raise ValueError('Refusing to remove manually changed omitted companion: '+name)
            stale.unlink()
    if manifest or old_manifest.exists():
        (outdir/'companions/manifest.json').write_text(json.dumps({'schema_version':1,'files':manifest},ensure_ascii=False,indent=2),encoding='utf-8')

def preserve_source(input_dir,data):
    """Restore canonical values from exact imports and retain all companion cells."""
    p=Path(input_dir);tables={name:_rows(p/name) for name in STANDARD if (p/name).exists()}
    source_model={}
    for filename in ['extraction.json','model.json']:
        if (p/filename).exists():
            candidate=json.loads((p/filename).read_text(encoding='utf-8'))
            if isinstance(candidate.get('groups'),list):source_model=candidate;break
    basic=tables.get('Basic_input.csv',[])
    if not basic:return data
    data['extraction_source_tables']=tables
    metadata=source_model.get('metadata',{})
    if metadata:data['metadata']=metadata
    dt=tables.get('Diet_composition.csv',[]);dr={r[0]:r for r in dt[1:] if r and r[0]}
    imp=next((r for r in dt if len(r)>1 and r[1]=='Import'),[])
    fate=tables.get('Detritus_fate.csv',[]);fd={r[0]:r for r in fate[1:] if r and r[0]}
    ba={str(r[0]):r for r in tables.get('Biomass_accumulation.csv',[])[1:]}
    tl={str(r[0]):r[2] for r in tables.get('TL.xlsx',[])[1:] if len(r)>2}
    known={str(g['group_seq']):g for g in data.get('group',[])};groups=[];changes=[]
    names={str(r[1]):str(r[0]) for r in basic[1:] if len(r)>1 and r[0]}
    dets={names[x] for x in (fate[0][2:-2] if fate else []) if x in names}
    for row in basic[1:]:
        if len(row)<2 or not row[0]:continue
        seq=str(row[0]);g=known.get(seq,{'group_seq':seq,'group_name':row[1]});old=json.loads(json.dumps(g))
        for field,col in [('habitat_area',2),('biomass_habitat_area',3),('z',4),('pb',5),('qb',6),('ee',7),('other_mort',8),('ge',9),('gs',10),('detritus_import',11)]:g[field]=_nd(row[col] if len(row)>col else None)
        source_group=next((x for x in source_model.get('groups',[]) if str(x['n'])==seq),{})
        roles=source_group.get('parameter_roles',{})
        role_fields={'biomass_habitat_area':'biomass','pb':'pb','qb':'qb','ee':'ee','ge':'pq'}
        g['parameter_source_roles']=roles
        for flag,field in [('b_hab_area_input','biomass_habitat_area'),('pb_input','pb'),('qb_input','qb'),('ee_input','ee'),('ge_input','ge')]:
            role=roles.get(role_fields[field],'reported_unspecified')
            g[flag]='false' if g[field]=='-9999' or role in ['model_estimated','derived'] else 'true'
        g['input_flag_scope']='known values supplied to importer; author-estimated/derived roles explicitly false; unspecified reported provenance remains separate'
        if g['biomass_habitat_area']=='-9999':g['biomass']='-9999'
        elif metadata.get('biomass_basis')=='whole_model_area':g['biomass']=g['biomass_habitat_area']
        elif g['habitat_area']!='-9999':g['biomass']=str(Decimal(g['biomass_habitat_area'])*Decimal(g['habitat_area']))
        else:g['biomass']='-9999'
        source_abs=_nd(ba.get(seq,[None]*4)[2]);g['biomass_accum']=source_abs;g['biomass_accum_rate']=_nd(ba.get(seq,[None]*4)[3]);g['tl']=_nd(tl.get(seq))
        g['biomass_accum_source_literal']=source_abs
        if source_abs!='-9999' and metadata.get('ba_basis',metadata.get('biomass_basis'))!='whole_model_area':
            g['biomass_accum']=str(Decimal(source_abs)*Decimal(g['habitat_area'])) if g['habitat_area']!='-9999' else '-9999'
        # A zero growth constant/price supplied by the old converter is not source evidence.
        g['vbk']='-9999';g['shadow_price']='-9999'
        for key,fname in [('landings','Landings.csv'),('discards','Discards.csv')]:
            tab=tables.get(fname,[]);r=next((x for x in tab[1:] if str(x[0])==seq),[])
            fleets=tab[0][2:-1] if tab else []
            g[key+'_by_fleet']={f:_nd(r[2+i] if len(r)>2+i else None) for i,f in enumerate(fleets)}
            # A partial numeric row total does not establish complete fleet coverage.
            complete=bool(fleets) and all(v!='-9999' for v in g[key+'_by_fleet'].values())
            g[key+'_total']=_nd(r[-1]) if r and complete else '-9999'
            g[key+'_known_subtotal']=_nd(r[-1]) if r else '-9999'
        g['total_removals']=str(Decimal(g['landings_total'])+Decimal(g['discards_total'])) if g['landings_total']!='-9999' and g['discards_total']!='-9999' else '-9999'
        g['export']=g['landings_total'] if metadata.get('export_basis')=='reported_landings' else g['total_removals']
        g['export_scope']='explicit reported_landings convention; discards unknown or separately retained' if metadata.get('export_basis')=='reported_landings' else 'landings plus discards; unknown when either component incomplete'
        fr=fd.get(seq,[]);g['detritus_fate_by_pool']={names[name]:_nd(fr[i] if len(fr)>i else None) for i,name in enumerate(fate[0] if fate else []) if i>=2 and name in names};g['detritus_export']=_nd(fr[-2]) if fr else '-9999'
        g['diet_imp']='-9999';g['diet_descr']=None
        if dt and seq in dt[0][2:]:
            col=dt[0].index(seq);g['diet_imp']=_nd(imp[col] if len(imp)>col else None)
            g['diet_descr']={'diet':[{'prey_seq':prey,'proportion':_nd(r[col] if len(r)>col else None),'detritus_fate':g['detritus_fate_by_pool'].get(prey,'-9999')} for prey,r in dr.items()]}
        g['pp']='2' if seq in dets else ('0' if dt and seq in dt[0][2:] else g.get('pp','1'))
        for key in set(old)|set(g):
            if old.get(key)!=g.get(key):changes.append({'group_seq':seq,'field':key,'converted':old.get(key),'source_preserved':g.get(key)})
        groups.append(g)
    data['group']=groups;data['source_fidelity']={'version':1,'normalization_applied':False,'biomass_basis':metadata.get('biomass_basis','unspecified'),'unknown_sentinel':'-9999','standard_table_hashes':{f:_sha(p/f) for f in tables},'restorations':changes}
    source_groups={str(g['n']):g for g in source_model.get('groups',[])}
    for g in groups:
        if source_groups.get(g['group_seq'],{}).get('source_status'):g['source_status']=source_groups[g['group_seq']]['source_status']
        literal=source_groups.get(g['group_seq'],{}).get('tl')
        if literal is not None and g['tl']!='-9999' and Decimal(str(literal))==Decimal(g['tl']):g['tl']=str(literal)
    diet_sums=[]
    for g in groups:
        if not g.get('diet_descr'):continue
        diet=g['diet_descr']['diet'];missing=[x['prey_seq'] for x in diet if x['proportion']=='-9999']
        prey=sum((Decimal(x['proportion']) for x in diet if x['proportion']!='-9999'),Decimal(0))
        imported=None if g['diet_imp']=='-9999' else Decimal(g['diet_imp'])
        diet_sums.append({'consumer_seq':g['group_seq'],'consumer_name':g['group_name'],'prey_only_known_sum':str(prey),'diet_plus_import_known_sum':str(prey+(imported or 0)),'diet_plus_import_sum':None if missing or imported is None else str(prey+imported),'unknown_prey_ids':missing,'import':g['diet_imp'],'import_unknown':imported is None,'extraction_normalized':False,'review_status':'not_established_by_conversion'})
    data['source_fidelity']['diet_sums']=diet_sums
    for name in ['stanzas','extraction_notes']:
        if name in source_model:data[name]=source_model[name]
    manifest=p/'companions/manifest.json';companions={}
    if manifest.exists():
        for item in json.loads(manifest.read_text(encoding='utf-8'))['files']:
            name=item['filename']
            if Path(name).name!=name:raise ValueError('Unsafe companion path')
            cp=manifest.parent/name
            if _sha(cp)!=item['sha256']:raise ValueError('Companion changed after manifest: '+name)
            companions[name]={'description':item['description'],'layout':item.get('layout','table'),'role':item['role'],'rows':_rows(cp),'sha256':item['sha256']}
    data['published_companions']=companions
    return data

def reconstruct_source_workbook(data,output):
    """Numeric standard sheets, literal companion sheets, explicit missing masks."""
    wb=openpyxl.Workbook();wb.remove(wb.active);manifest=[]
    for filename,rows in data['extraction_source_tables'].items():
        ws=wb.create_sheet(Path(filename).stem)
        for ri,row in enumerate(rows,1):
            for ci,value in enumerate(row,1):
                if value is None or value=='':continue
                numeric=filename.endswith('.csv') and ri>1 and (ci==1 or ci>=3)
                if numeric:
                    try:
                        number=Decimal(str(value));cell=ws.cell(ri,ci,float(number))
                        if 'e' in str(value).lower():cell.number_format='0.'+'0'*max(0,-number.as_tuple().exponent)+'E+00'
                        else:cell.number_format='0'+('.'+'0'*max(0,-number.as_tuple().exponent) if number.as_tuple().exponent<0 else '')
                        continue
                    except (ValueError,ArithmeticError):pass
                cell=ws.cell(ri,ci,value)
                if isinstance(value,str):cell.data_type='s'
        ws.freeze_panes='C2';ws.column_dimensions['B'].width=32
        manifest.append([ws.title,filename,'standard_EwE_import'])
    for i,(filename,spec) in enumerate(data.get('published_companions',{}).items(),1):
        title=('Source_'+str(i)+'_'+Path(filename).stem)[:31];ws=wb.create_sheet(title)
        for row in spec['rows']:
            ws.append([None if x=='' else x for x in row])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='A2';manifest.append([ws.title,'companions/'+filename,spec['description']])
    ws=wb.create_sheet('Source_manifest');ws.append(['Sheet','Source file','Role / description'])
    for row in manifest:ws.append(row)
    wb.save(output);reopen=openpyxl.load_workbook(output,data_only=False);errors=[];count=0;maxdiff=Decimal(0)
    for sheet,filename,_ in manifest:
        rows=data['extraction_source_tables'][filename] if filename in data['extraction_source_tables'] else data['published_companions'][filename.split('/',1)[1]]['rows']
        got=[list(r) for r in reopen[sheet].iter_rows(values_only=True)];width=max(map(len,rows))
        expected=[[None if x=='' else x for x in r]+[None]*(width-len(r)) for r in rows]
        count+=sum(len(r) for r in expected)
        if len(got)!=len(expected):errors.append(filename);continue
        for ri,(er,ar) in enumerate(zip(expected,got),1):
            for ci,(e,a) in enumerate(zip(er,ar),1):
                if isinstance(a,(int,float)) and e is not None:
                    try:
                        delta=abs(Decimal(str(e))-Decimal(str(a)));maxdiff=max(maxdiff,delta)
                        if delta>max(Decimal('1e-12'),abs(Decimal(str(e)))*Decimal('1e-14')):errors.append([filename,ri,ci])
                    except:errors.append([filename,ri,ci])
                elif a!=e:errors.append([filename,ri,ci])
    result={'schema_version':1,'standard_tables':len(data['extraction_source_tables']),'companion_tables':len(data.get('published_companions',{})),'cells_checked':count,'exact_value_and_missing_mask_match':not errors and maxdiff==0,'numeric_value_and_missing_mask_match':not errors,'max_numeric_difference':str(maxdiff),'absolute_tolerance':'1e-12','relative_tolerance':'1e-14','errors':errors}
    Path(output).with_name('SOURCE_FIDELITY_CHECK.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    if errors:raise ValueError('Source workbook mismatch: '+str(errors))
    return str(output)
