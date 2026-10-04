from pathlib import Path
from docx import Document
import zipfile,json,hashlib
C=Path(__file__).resolve().parents[1];ROOT=C.parents[3]
p=ROOT/'tools/templates/Model_validation_template.docx';d=Document(p)
info={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sections':[{'width':s.page_width/914400,'height':s.page_height/914400,'left':s.left_margin/914400,'right':s.right_margin/914400,'top':s.top_margin/914400,'bottom':s.bottom_margin/914400} for s in d.sections],'paragraphs':[(i,x.text,x.style.name) for i,x in enumerate(d.paragraphs)],'tables':[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]}
(C/'qa/template_inventory.json').write_text(json.dumps(info,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps(info,ensure_ascii=False,indent=2))
with zipfile.ZipFile(p) as z:
 parts={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
(C/'qa/template_parts.json').write_text(json.dumps(parts,indent=2),encoding='utf8')
