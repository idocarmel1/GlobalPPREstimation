from pathlib import Path,PureWindowsPath
import json,hashlib,zipfile,math,sys
from decimal import Decimal
from collections import Counter
from lxml import etree as E
from openpyxl import load_workbook
from docx import Document
C=Path(__file__).resolve().parents[1];ROOT=C.parents[3]
sys.path.insert(0,str(ROOT/'tools'))
from validation_percentage_format import verify_report
read=lambda f:json.loads((C/f).read_text(encoding='utf8'))
checks={};errors=[]
def check(name,cond):
 checks[name]=bool(cond)
 if not cond:errors.append(name)
def equal(a,b):return a is None and b is None or a is not None and b is not None and Decimal(str(a))==Decimal(str(b))
model=read('model.json');literal=read('extracted_tables/thesis/literal_source_extraction.json');partial=read('extracted_tables/thesis/partial_import_source.json')
check('canonical source data identical to final partial extraction',all(model[k]==v for k,v in partial.items() if k!='metadata'))
check('37 original groups and exact order',[g['n'] for g in model['groups']]==list(range(1,38)))
param_map={'B':'biomass','PB':'pb','QB':'qb','EE':'ee','PQ':'pq','TL':'tl'}
check('222 parameter values and missing masks',all(equal(g[k],model['groups'][i][field]) for i,g in enumerate(literal['groups']) for k,field in param_map.items()))
check('unreported inputs preserved',all(g[k] is None for g in model['groups'] for k in ['hab_area','unassim','ba','ba_rate']))
check('1258 source diet cells',len(literal['diet_cells'])==37*34)
check('source diet numeric and blank counts',Counter(x['value'] is None for x in literal['diet_cells'])=={True:940,False:318})
diet_errors=[]
for x in literal['diet_cells']:
 p=x['prey_id'];consumer=x['consumer_id_source_number'];expected=None if p==20 or consumer==20 else x['value']
 actual=model['diet'].get(str(consumer),{}).get(str(p))
 if not equal(actual,expected):diet_errors.append([p,consumer,expected,actual])
check('canonical diet retains exact nonconflicting values and identity mask',not diet_errors)
check('unmodified source diet sums',all(Decimal(x['sum_source_literals'])==(Decimal('.999') if x['consumer_id'] in [2,5,10] else Decimal('1')) for x in literal['diet_column_sums']))
check('candidate IDs and year',model['metadata']['model_id']==C.name and model['metadata']['variant_id']=='thesis_1991-1992_Table6_5_AnnexIIA' and model['metadata']['model_year']=='1991-1992')
source=C/model['metadata']['source'];check('registered thesis hash',hashlib.sha256(source.read_bytes()).hexdigest()==literal['sha256'])
diag=read('diagnostics/direct_reports.json');check('three diagnostics genuinely NOT_RUN',len(diag)==3 and all(x['status']=='NOT_RUN' and all(x[k] is None for k in ['rho_living','b','detritus_sppr','matrix','coefficients']) for x in diag.values()))
rows=read('mapping/appendix_rows.json');cov=read('mapping/coverage_summary.json');weak=read('mapping/very_low_decisions.json');audit=read('mapping/taxon_mapping_audit.json')
headers=['Taxon name','TL','Catch (t)','Simple-chain PPR (t C)','Mapped group names and weights','Confidence level','Reason']
check('512 unique taxa and required seven columns',len(rows)==512 and len({r['Taxon name'] for r in rows})==512 and all(list(r)==headers for r in rows))
check('descending unrounded PPR with name ties',rows==sorted(rows,key=lambda r:(r[headers[3]] is None,-(r[headers[3]] or 0),r[headers[0]])))
check('reference landings total',math.isclose(math.fsum(r['Catch (t)'] for r in rows),cov['total_catch_tonnes'],rel_tol=0,abs_tol=1e-8))
check('reference simple PPR total',math.isclose(math.fsum(r[headers[3]] for r in rows),cov['total_simple_chain_ppr_tC'],rel_tol=0,abs_tol=1e-6))
check('confidence counts reconcile',Counter(r['Confidence level'] for r in rows)=={v['label']:v['taxa'] for v in cov['confidence_summary'] if v['taxa']})
check('very low decisions each once',Counter(t for w in weak for t in w['taxa'])==Counter(r['Taxon name'] for r in rows if r['Confidence level']=='Very low'))
check('positive catch coefficient missingness explicit',all(a['classic_sppr_wet_weight'] is not None or a['catch_t']==0 for a in audit))
annual=read('calculations/annual_independent.json');check('historical incomplete totals remain null',len(annual)==210 and sum(r['simple_ppr_total_tC'] is None for r in annual)==149)
xlsx=C/'LME027_Villanueva2004_thesis_taxon_mapping_appendix.xlsx';w=load_workbook(xlsx,data_only=False);s=w['Taxon mapping'];src=w['Sources']
def office_equal(a,b):return isinstance(a,(int,float)) and isinstance(b,(int,float)) and math.isclose(a,b,rel_tol=2e-15,abs_tol=1e-12) or a==b
check('Excel typed appendix cells within XLSX numeric precision',list(next(s.iter_rows(min_row=7,max_row=7,values_only=True)))==headers and all(office_equal(s.cell(i,j).value,r[h] if r[h] is not None else '?') for i,r in enumerate(rows,8) for j,h in enumerate(headers,1)))
check('freeze headers and native filters',s.freeze_panes=='A8' and src.freeze_panes=='A5' and all(t.autoFilter and t.autoFilter.ref==t.ref for sh in w for t in sh.tables.values()))
check('Excel no formulas or error cells',all(c.data_type not in ['e','f'] for sh in w for row in sh for c in row))
check('16 descriptive blue underlined native Sources links',sum(bool(c.hyperlink) for row in src for c in row)==16 and all(src.cell(i,3).font.color.rgb=='000563C1' and src.cell(i,3).font.underline=='single' for i in range(5,21)))
doc=C/'Model_validation_27_Villanueva2004_SineSaloum_Thesis_1991-1992.docx';d=Document(doc);verify_report(doc)
check('unsigned researcher slots preserved',d.tables[0].cell(15,1).text=='' and '____________________' in d.tables[0].cell(16,1).text and '[Researcher calculation choices and notes]'==d.tables[0].cell(11,1).text)
wordweak=[t.strip() for row in d.tables[4].rows[1:] for t in row.cells[0].text.split(';')];check('Word all very low taxa each once',Counter(wordweak)==Counter(r['Taxon name'] for r in rows if r['Confidence level']=='Very low'))
check('template opaque parts preserved',read('qa/template_fidelity.json')['preserve_only_parts_equal'])
links=[];broken=[]
for file in [doc,xlsx]:
 with zipfile.ZipFile(file) as z:
  for n in z.namelist():
   if not n.endswith('.rels'):continue
   for r in E.fromstring(z.read(n)):
    if not r.get('Type','').endswith('/hyperlink'):continue
    target=r.get('Target');links.append({'file':file.name,'target':target})
    if target.startswith(('https://','http://')):continue
    base=target.split('#')[0]
    if PureWindowsPath(base).drive or base.startswith('/') or not (C/base).exists():broken.append(target)
    else:
     rel=(C/base).resolve().relative_to(ROOT)
     fake=Path('Z:/relocated_project')/C.relative_to(ROOT)
     check('relocated link '+str(len(links)),(fake/base).resolve().as_posix().endswith(rel.as_posix()))
check('all native Office local links portable and resolve',not broken)
baseline=read('qa/protected_baseline.json');changes=[f for f,h in baseline.items() if not (ROOT/f).exists() or hashlib.sha256((ROOT/f).read_bytes()).hexdigest()!=h]
shared={'Project.xlsx','interactive_map/index.html','interactive_map/trends.html'}
check('374 local protected files unchanged',len(baseline)==377 and not (set(changes)-shared))
pw=load_workbook(ROOT/'Project.xlsx',read_only=True,data_only=False);current={}
for sheet in ['Papers','Models & coverage']:
 blocks={};name=None;header=None;rr=[]
 for raw in pw[sheet].values:
  vals=list(raw)
  while vals and vals[-1] is None:vals.pop()
  if not vals:continue
  if vals[0]=='@table':
   if name is not None:blocks[name]=rr
   name=vals[1];header=None;rr=[]
  elif name is not None:
   if header is None:header=vals
   elif 'LME_027' in vals:rr.append(dict(zip(header,(vals+[None]*len(header))[:len(header)])))
 if name is not None:blocks[name]=rr
 current[sheet]=blocks
pw.close();previous=read('work/project_context.json');project_diff=[]
for sh,ts in previous.items():
 for table,old in ts.items():
  new=current[sh][table]
  for a,b in zip(old,new):
   dd={k:[a.get(k),b.get(k)] for k in set(a)|set(b) if a.get(k)!=b.get(k)}
   if dd:project_diff.append({'sheet':sh,'table':table,'identity':a.get('model_id',a.get('article_id')),'fields':dd})
  if len(new)!=len(old):project_diff.append({'table':table,'row_counts':[len(old),len(new)]})
(C/'qa/concurrent_project_differences.json').write_text(json.dumps(project_diff,ensure_ascii=False,indent=2),encoding='utf8')
check('LME_027 Project paper and model rows unchanged',not project_diff)
from original_atlas_data import embedded
network=embedded(ROOT/'interactive_map/index.html','DB')[0]['network']['units']['LME_027']
selected=network['models'][network['default_model']]['id']
check('map retains LME_027 selected model and excludes candidate',selected=='27_118_Northwest_Africa_(1987)' and C.name not in network['models'])
result={'status':'PASS' if not errors else 'FAIL','checks':checks,'errors':errors,'protected_files_checked':len(baseline),'protected_changes':changes,'concurrent_change_note':'Shared Project/map files changed while another user-authorized chat (Evaluate Northern Humboldt auto mode, 01a102bb-2c92-7d93-a1e7-00c9615d4757) integrated LME_013. Read-only thread evidence confirmed its Project and map commands. No shared files were authored by this candidate task. Original baseline hashes retained; LME_027 registry rows and selected map model independently checked.','diet_errors':diet_errors,'office_hyperlinks':links,'broken_links':broken,'scientific_pipeline':'BLOCKED_SOURCE_IDENTITY','review_package':'READY_FOR_RESEARCHER_REVIEW' if not errors else 'QA_FAILED'}
(C/'qa/package_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({k:result[k] for k in ['status','errors','protected_files_checked','protected_changes','diet_errors','broken_links']},ensure_ascii=False));sys.exit(bool(errors))
