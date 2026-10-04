from pathlib import Path
import json,sys,copy,zipfile,hashlib
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.paragraph import Paragraph
C=Path(__file__).resolve().parents[1];ROOT=C.parents[3]
sys.path.insert(0,str(ROOT/'tools'))
from validation_percentage_format import format_percent,verify_report
TEMPLATE=ROOT/'tools/templates/Model_validation_template.docx'
DEST=C/'Model_validation_27_Villanueva2004_SineSaloum_Thesis_1991-1992.docx'
APPENDIX='LME027_Villanueva2004_thesis_taxon_mapping_appendix.xlsx'
def read(f):return json.loads((C/f).read_text(encoding='utf8'))
cov=read('mapping/coverage_summary.json');weak=read('mapping/very_low_decisions.json');geo=read('evidence/geography/geographic_assessment.json')
d=Document(TEMPLATE)
def put(p,text):
 pr=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(str(text))
 if pr is not None:r._r.insert(0,pr)
 return p
def cell(c,text):
 prototype=copy.deepcopy(c.paragraphs[0]._p)
 for p in list(c.paragraphs)[1:]:p._p.getparent().remove(p._p)
 parts=str(text).split('\n');put(c.paragraphs[0],parts[0])
 for t in parts[1:]:
  el=copy.deepcopy(prototype);c._tc.append(el);put(Paragraph(el,c),t)
def link(p,label,target):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 for tag,val in [('w:color','0563C1'),('w:u','single'),('w:sz','21')]:
  e=OxmlElement(tag);e.set(qn('w:val'),val);pr.append(e)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def links(c,items):
 p=c.add_paragraph()
 for i,(label,target) in enumerate(items):
  if i:p.add_run(' | ')
  link(p,label,target)
def fill(t,rows):
 proto=copy.deepcopy(t.rows[1]._tr)
 for r in list(t.rows)[1:]:t._tbl.remove(r._tr)
 for values in rows:
  t._tbl.append(copy.deepcopy(proto));r=t.rows[-1]
  for c,txt in zip(r.cells,values):cell(c,str(txt))
def find(prefix):return next(p for p in d.paragraphs if p.text.startswith(prefix))
put(d.paragraphs[0],'Canary Current Sine Saloum thesis candidate review')
e=OxmlElement('w:p');d.paragraphs[0]._p.addnext(e);p=Paragraph(e,d)
p.add_run('The preferred thesis describes a 37-group estuary model for 1991–1992. Its basic table and diet annex disagree about group 20, preventing a defensible SPPR calculation. Its reported area can cover at most 0.049% of the Canary Current LME. This is an unsigned candidate review draft; the active selection remains EcoBase 118.').font.size=Pt(11)
main=d.tables[0]
entries={
1:'Canary Current | LME_027\nCandidate: Sine-Saloum estuary, Senegal.',
2:'Sea Around Us; retained regional catch records 1950–2019. Reference: 2019 landings, excluding discards.',
3:'Villanueva, Maria Concepcion S. (2004). Biodiversité et relations trophiques dans quelques milieux estuariens et lagunaires de l’Afrique de l’Ouest : adaptations aux pressions environnementales. Doctoral thesis, Institut National Polytechnique de Toulouse / ENSAT. Source title page: defended 29 October 2004; a harvested catalogue date differs. No DOI located.',
4:'Villanueva, Tito-de-Morais, Weigel and Moreau, An Ecopath model of the Sine-Saloum Delta Biosphere Reserve (Senegal), proceedings report cover 2004 / legal imprint 2005: a different 38-group preliminary version.\nVillanueva, Diouf, Albaret and Moreau (2005), Complexity in trophic structure and stability in a modelled West African estuary, in Coastal Ecosystems of West-Africa, edited by J.-J. Symoens, pp. 15–26: a related 37-group release.\nMorissette et al. (2009), Northwest Africa food-web model, FCRR 17(2): same model lineage as selected EcoBase 118. Guénette et al. (2014), Banc d’Arguin and Mauritanian Shelf, PLOS ONE e94742: distinct shelf models.',
5:'27_Villanueva2004_SineSaloum\nPreferred variant: thesis Table 6.5 / Annex II.A, Sine-Saloum, 1991–1992, 37 source groups. Related EcoBase accession 612 identifies this model lineage, but its official numerical input is unavailable. The local canonical representation is explicitly incomplete.',
6:'The thesis also models Gambia (41 groups, 2000–2002; Table 6.6), Ébrié Lagoon (42 groups, 1980–1981; Table 6.7), and Nokoué Lagoon (31 groups, 2000–2001; Table 6.8). These are separate ecosystem models and were not evaluated as Sine-Saloum replacements in this review.',
7:'Article – The researcher authorized the downloaded and documented thesis as the preferred source after the supplied chapter lacked a quantitative diet matrix.\nModel – Review the thesis Sine-Saloum model as a regional alternative; its stated period and quantitative annex improve evidence. Geographic proximity does not establish suitability.',
8:'Table 6.5 and the complete 37-prey × 34-consumer Annex II.A were extracted separately. The annex uses fractions despite its percent caption; consumer sums are 0.999–1.000 and are preserved.\nGroup 27 labels refer to the same documented mullet pool. Group 20 is unresolved: the basic table names Epinephelus aeneus, but the annex names Hemichromis fasciatus; Table 3.3 separates those compartments. Canonical cells requiring that join remain unknown. GS, BA and routing are unreported; no defaults, normalization or numerical repair were applied.',
9:'NOT_RUN — unresolved group 20 diet identity blocks source admission.\nNegative SPPR entries, rho_living, b and Detritus (37) SPPR cannot be determined. No GE matrix or coefficients are available.',
10:'NOT_RUN — the same source identity conflict blocks source admission.\nNegative SPPR entries, rho_living and Detritus (37) SPPR cannot be determined. No TE matrix or coefficients are available.',
11:'[Researcher calculation choices and notes]',
12:'A. Region covered by study: at most approximately 0.049% (reported-area upper bound).\nB. Study area covered by region: not determined; the model water boundary is ambiguous in the schematic and source bounds. Main channels lie within the target geometry.',
13:'The modeled period is 1991–1992; thesis field sampling also includes 1992–1993. The reference landings year 2019 is 27–28 years later. Fixed source catch/biomass mixtures across 1950–2019 are assumptions and cannot reconstruct annual estuarine or offshore food-web change.',
14:'With Egestion is also NOT_RUN. No candidate GE/TE/Egestion regional PPR or PPR/NPP is available. Independent simple-chain PPR remains calculable from the regional taxon coefficients.\nThe selected EcoBase 118 remains disqualified on 02/10/2026: “Reason – too strong living compartments recycling due to near-zero EE values (rho_living = 0.99).” The thesis is a distinct local model, but shared source lineage and model distinction do not establish independent observations or whole-LME representativeness.',
15:'',16:'Researcher name: ____________________ | review date: dd/mm/yyyy'}
for i,text in entries.items():cell(main.cell(i,1),text)
cell(main.cell(3,0),'Preferred source');cell(main.cell(5,0),'Candidate model')
links(main.cell(1,1),[('Regional workbook','../../LME_027.xlsx')])
links(main.cell(2,1),[('Catch evidence','../../raw/LME_027.csv.gz')])
links(main.cell(3,1),[('Thesis PDF','../../papers/LME027-Villanueva-2004-Thesis/Villanueva2004_thesis_29305.pdf#page=138'),('Source metadata','../../papers/LME027-Villanueva-2004-Thesis/metadata.json'),('Archimer record','https://archimer.ifremer.fr/doc/00180/29122/')])
links(main.cell(4,1),[('Supplied chapter','../../papers/LME027-Villanueva-2004/010056023-adc18d15.pdf'),('2005 follow-up','evidence/identity/Villanueva2005_2171.pdf')])
links(main.cell(5,1),[('Candidate folder','.'),('Canonical partial model','model.json'),('Identity assessment','evidence/identity/identity_review.json')])
links(main.cell(8,1),[('Full literal extraction','extracted_tables/thesis/literal_source_extraction.json'),('Diet identity crosswalk','extracted_tables/thesis/diet_identity_crosswalk.csv'),('Admission blocker','computational_inputs/ADMISSION.json')])
links(main.cell(14,1),[('Direct diagnostic records','diagnostics/DIRECT_DIAGNOSTICS.md'),('Preserved signed review','../../Model_validation_27_118_Northwest_Africa_(1987).docx')])
put(find('Reference:'),'Reference: 2019 landings, excluding discards; all recorded taxon labels, including zero-landings rows. Candidate mappings below remain unadopted.')
put(find('[n] source groups'),f"37 source groups; {cov['taxa']} catch taxa.")
put(find('Simple-chain PPR totals'),f"Simple-chain PPR totals {cov['total_simple_chain_ppr_tC']:,.0f} t C in 2019.")
missing=cov['missing_classic_coefficients'];missing=missing if isinstance(missing,int) else len(missing)
put(find('[Missing TL'),f"{missing} taxa have no saved classic coefficient; all have zero 2019 landings and contribute zero annual PPR. No positive-catch contribution is unknown.")
put(find('Method:'),'Method: independent simple trophic chain using the saved regional taxon coefficients.')
p=put(find('Excel taxon appendix'),'');link(p,'Excel taxon appendix and descriptive Sources sheet',APPENDIX)
fill(d.tables[1],[[r['label'],r['taxa'],format_percent(r['catch_percentage']),format_percent(r['ppr_percentage'])] for r in cov['confidence_summary']])
fill(d.tables[2],[[r['plain_language_rule'],r['confidence'],format_percent(r['ppr_percentage'])] for r in sorted(cov['membership_summary'],key=lambda r:-r['ppr_percentage'])])
fill(d.tables[3],[[r['plain_language_rule'],r['confidence'],format_percent(r['ppr_percentage'])] for r in sorted(cov['allocation_summary'],key=lambda r:-r['ppr_percentage'])])
p=put(find('Membership evidence:'),'Thesis Tables 3.3 and 6.5, group-definition prose and verified taxonomy. ');link(p,'Full candidate audit','mapping/taxon_mapping_audit.json')
p=put(find('Weights and assumptions:'),'Source catch proportions, then complete biomass where usable, supply assumed mixtures. One-group assignments need no split. Fixed historical mixtures are transferred to 2019 landings. ');link(p,'Candidate allocation evidence','mapping/allocation_evidence.json')
very=next(r for r in cov['confidence_summary'] if r['label']=='Very low');unresolved=next(r for r in cov['confidence_summary'] if r['label']=='Unresolved')
unres='none' if not unresolved['taxa'] else '; '.join(r['Taxon name']+': '+r['Reason'] for r in read('mapping/appendix_rows.json') if r['Confidence level']=='Unresolved')
put(find('Unresolved taxa:'),f"Unresolved taxa: {unres}. {very['taxa']} Very low decisions account for {format_percent(very['catch_percentage'])} of catch and {format_percent(very['ppr_percentage'])} of simple-chain PPR. High assignment coverage does not establish valid SPPR or whole-LME ecological fit.")
decision_rows=[]
for w in weak:
 taxa=w['taxa'];reason=w.get('reason',w.get('why_confidence_is_very_low',w.get('reason_short','')))
 for start in range(0,len(taxa),7):decision_rows.append(['; '.join(taxa[start:start+7]),reason])
fill(d.tables[4],decision_rows)
put(find('Geographic evidence screenshots'),'Geographic evidence')
put(find('Target region R'),'R is the Canary Current LME and S is the thesis Sine-Saloum estuary. The reported 546 km² gives an upper bound on A even under full containment. Figure 2.2 does not delimit the exact model water area, and its eastward channels conflict with the printed longitude strip; B remains uncertain.')
put(find('[Region name/ID'),'Canary Current LME_027, retained project LME geometry. The estuary is located near the southern end of this much larger marine region.')
put(find('[Article; figure/page'),'Villanueva thesis (2004), Figure 2.2, printed p. 18 / PDF p. 42: the original Sine-Saloum channel schematic. The frame is not a model boundary.')
put(find('[If a figure is unavailable'),'')
for ti,filename,width in [(5,'target_region.png',6.1),(6,'source_thesis_Figure2_2.png',6.1)]:
 cell(d.tables[ti].cell(1,0),'');d.tables[ti].cell(1,0).paragraphs[0].add_run().add_picture(str(C/'evidence/geography'/filename),width=Inches(width));d.tables[ti].rows[1].height=None
 p=d.tables[ti].cell(1,0).add_paragraph();link(p,'Geographic assessment','evidence/geography/geographic_assessment.json')
find('Study area in the article').paragraph_format.page_break_before=True
cell(d.tables[6].cell(0,0),'Original thesis study area schematic')
d.save(DEST)
with zipfile.ZipFile(TEMPLATE) as z:original={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(DEST) as z:filled={n:z.read(n) for n in z.namelist()}
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
for n,v in original.items():
 if n not in editable:filled[n]=v
with zipfile.ZipFile(DEST,'w',zipfile.ZIP_DEFLATED) as z:
 for n,v in filled.items():z.writestr(n,v)
verify_report(DEST)
(C/'qa/template_fidelity.json').write_text(json.dumps({'template_sha256':hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),'preserve_only_parts_equal':all(filled[n]==v for n,v in original.items() if n not in editable),'manual_unsigned':True,'candidate_directory_override':'Explicit user instruction to co-locate candidate package','report':DEST.name,'appendix':APPENDIX},indent=2),encoding='utf8')
print(DEST.name)
