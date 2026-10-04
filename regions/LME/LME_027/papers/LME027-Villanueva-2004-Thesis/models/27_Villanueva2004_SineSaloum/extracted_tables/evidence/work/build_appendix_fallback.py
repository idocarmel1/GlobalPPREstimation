from pathlib import Path
import json,math,posixpath
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
C=Path(__file__).resolve().parents[1]
read=lambda f:json.loads((C/f).read_text(encoding='utf8'))
rows=read('mapping/appendix_rows.json');sources=read('mapping/appendix_sources.json');cov=read('mapping/coverage_summary.json')
headers=['Taxon name','TL','Catch (t)','Simple-chain PPR (t C)','Mapped group names and weights','Confidence level','Reason']
w=Workbook();s=w.active;s.title='Taxon mapping';src=w.create_sheet('Sources')
opening=['Sine Saloum thesis candidate taxon mapping','2019 landings excluding discards. Independent simple trophic chain in tonnes C. Descending unrounded PPR then taxon. Unknown last.','37 source groups from thesis Table 6.5, 1991–1992. Candidate assignments remain unadopted. Model SPPR unavailable because group 20 diet identity is unresolved.',f"{cov['taxa']} taxa. Landings {cov['total_catch_tonnes']:,.2f} t. Simple-chain PPR {cov['total_simple_chain_ppr_tC']:,.2f} t C.",'Missing inputs display ?. Zero recorded catch contributes zero PPR even with an unavailable coefficient. Detailed source and allocation evidence are on Sources.']
for i,text in enumerate(opening,1):s.merge_cells(start_row=i,start_column=1,end_row=i,end_column=7);s.cell(i,1,text);s.row_dimensions[i].height=30
for i,h in enumerate(headers,1):s.cell(7,i,h)
for i,v in enumerate(rows,8):
 for j,h in enumerate(headers,1):s.cell(i,j,v[h] if v[h] is not None else '?')
 lines=max(math.ceil(len(str(v['Reason']))/91),math.ceil(len(str(v['Mapped group names and weights']))/55),math.ceil(len(str(v['Taxon name']))/31))
 s.row_dimensions[i].height=max(44,lines*15+16)
for col,width in {'A':31,'B':8,'C':18,'D':23,'E':56,'F':16,'G':87}.items():s.column_dimensions[col].width=width
src.merge_cells('A1:D1');src['A1']='Sources and applicability';src.merge_cells('A2:D2');src['A2']='Thesis source membership and proxy weights are distinct from whole LME applicability and diagnostic availability.'
src.append([]);src.append(['Source','What this evidence supports','Open evidence','Record date'])
for i,v in enumerate(sources,5):
 src.append([v['title'],v['supports'],'Open '+v['id'],v.get('date','2026-10-03')]);src.row_dimensions[i].height=max(62,math.ceil(len(v['supports'])/100)*15+18)
 target=v['target'] if v['target'].startswith(('http://','https://')) else posixpath.normpath('mapping/'+v['target'])
 if not target.startswith(('http://','https://')):assert (C/target.split('#')[0]).exists(),target
 src.cell(i,3).hyperlink=target
for col,width in {'A':44,'B':104,'C':34,'D':21}.items():src.column_dimensions[col].width=width
for sh,head,end,cols,name in [(s,7,7+len(rows),7,'ThesisCandidateTaxonMapping'),(src,4,4+len(sources),4,'ThesisCandidateSources')]:
 sh.sheet_view.showGridLines=False;sh.sheet_properties.tabColor='315D70';sh.freeze_panes=f'A{head+1}'
 for row in sh.iter_rows(min_row=1,max_row=end,max_col=cols):
  for c in row:
   c.font=Font(name='Arial',size=10);c.alignment=Alignment(wrap_text=True,vertical='center')
   if c.row>head and c.row%2==1:c.fill=PatternFill('solid',fgColor='F1F5F7')
 for c in sh[head]:c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='315D70')
 sh.row_dimensions[head].height=34;sh['A1'].font=Font(name='Arial',size=15,bold=True,color='172F3A');sh.row_dimensions[1].height=32
 t=Table(displayName=name,ref=f'A{head}:{chr(64+cols)}{end}');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False);sh.add_table(t)
 sh.sheet_properties.pageSetUpPr.fitToPage=True;sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0;sh.print_title_rows=f'1:{head}'
for i in range(8,8+len(rows)):
 s.cell(i,2).number_format='0.00'
 for j in (3,4):s.cell(i,j).number_format='#,##0.00;[Red]-#,##0.00;0.00'
 if s.cell(i,6).value=='Very low':s.cell(i,6).fill=PatternFill('solid',fgColor='FFF1D6');s.cell(i,6).font=Font(name='Arial',size=10,bold=True,color='784B00')
for i in range(5,5+len(sources)):src.cell(i,3).font=Font(name='Arial',size=10,color='0563C1',underline='single')
src.row_dimensions[2].height=28
w.save(C/'LME027_Villanueva2004_thesis_taxon_mapping_appendix.xlsx')
(C/'qa/spreadsheet_runtime.json').write_text(json.dumps({'preferred_author':'@oai/artifact-tool','availability':'unavailable','evidence':'Module import ERR_MODULE_NOT_FOUND; bundled node_modules has no @oai directory. Loader was called and dependency path verified.','fallback_author':'openpyxl','values_source':'mapping/appendix_rows.json','rows':len(rows),'sources':len(sources)},indent=2),encoding='utf8')
print(f'Workbook created: {len(rows)} taxa; {len(sources)} native hyperlinks.')
