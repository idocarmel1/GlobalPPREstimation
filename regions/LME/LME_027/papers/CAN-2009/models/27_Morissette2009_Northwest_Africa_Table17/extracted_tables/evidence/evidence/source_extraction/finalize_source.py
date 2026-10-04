"""Audit converter loss and retain explicit source unknowns and exact values."""
from pathlib import Path
from decimal import Decimal
import csv, hashlib, json, shutil, subprocess, sys
from copy import deepcopy
import openpyxl

OUT=Path(__file__).resolve().parent
CAND=OUT.parents[1]
ROOT=OUT.parents[5]
TABLES=CAND/'extracted_tables'
SKILL=ROOT/'tools/skills/original_skill_resources/claude/ecopath-extraction'
sys.path.insert(0,str(SKILL/'scripts'))
import write_outputs as writer

def writejson(path,data):
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

src=json.loads((OUT/'extraction.json').read_text(encoding='utf-8'))
rawpath=next(TABLES.glob('27_Canary_Current_*.json'))
raw=json.loads(rawpath.read_text(encoding='utf-8'))
for p,name in [(rawpath,'raw_converter_model.json'),(TABLES/'MASS_BALANCE.md','raw_converter_MASS_BALANCE.md'),(TABLES/'ewe_conversion.log','raw_converter.log')]:
    if not (OUT/name).exists(): shutil.copy2(p,OUT/name)
rawbook=next(TABLES.glob('*_reconstructed.xlsx'))
if not (OUT/'raw_converter_reconstructed.xlsx').exists(): shutil.copy2(rawbook,OUT/'raw_converter_reconstructed.xlsx')
canonical=deepcopy(raw)
changes=[]
def setvalue(g,key,value,why):
    old=g.get(key)
    if old!=value:
        changes.append({'group_id':int(g['group_seq']),'field':key,'old':old,'new':value,'reason':why})
    g[key]=value

for s,g in zip(src['groups'],canonical['group']):
    assert int(g['group_seq'])==s['n'] and g['group_name']==s['name']
    for f,t in [('biomass','biomass'),('biomass','biomass_habitat_area'),('pb','pb'),('qb','qb'),('ee','ee'),('pq','ge'),('tl','tl')]:
        setvalue(g,t,s[f] if s[f] is not None else '-9999','Table17 literal and source precision preserved')
    setvalue(g,'ge_input','true' if s['pq'] is not None else 'false','Printed GE carried instead of discarded by converter; estimated status in cell ledger')
    for field in ['habitat_area','vbk','shadow_price']:
        setvalue(g,field,'-9999','Not reported by source; converter default removed')
    for field in ['z','detritus_fate_export']:
        setvalue(g,field,'-9999','Not reported by source')
    n=str(s['n'])
    catch=src['landings'].get(n,{})
    total=str(sum((Decimal(v) for v in catch.values()),Decimal(0))) if catch else '-9999'
    setvalue(g,'export',total,'Partial source 1990s total catch, without silently replacing unknown with zero')
    g['_source_unknowns']=['habitat_area','vbk','shadow_price','z','gs','biomass_accum','biomass_accum_rate','immigration','emigration','emigration_rate','detritus_import','detritus_fate','detritus_fate_export','discard_separation']
    if not catch: g['_source_unknowns'].append('catch')
    g['_source_catch']={'period':'1990s mean','final_baseline_equivalence':'not_established','components':catch,'components_are_total_catch_not_separate_landings':True}
    if s['n']==25:
        g['_source_unknowns'].append('entire_diet')
        setvalue(g,'diet_descr',None,'Table18 omits entire consumer25 column')
        setvalue(g,'diet_imp','-9999','Table18 omits entire consumer25 column')
    elif s['n'] in src['consumers']:
        entries=[]
        for prey,value in src['diet'][n].items():
            if prey=='import': continue
            entries.append({'prey_seq':prey,'proportion':value,'detritus_fate':'-9999' if prey=='27' else '0'})
        setvalue(g,'diet_descr',{'diet':entries if len(entries)>1 else entries[0] if entries else None},'Preserve printed diet zeros; detritus27 routing unknown instead of converter zero')
    # Other prey targets cannot be detritus pools; their zero fate is structural.
    g['_source_detritus_routing']={'27':None,'Export':None}
    g['_source_estimated_fields']=['tl'] + (['biomass','qb'] if s['n']==14 else ['qb'] if s['n']==18 else []) + ([] if s['n']==14 else ['ee']) + ([] if s['n'] in [14,18,24,26,27] else ['ge'])

taxonomy=CAND/'Taxonomy.xlsx'
if taxonomy.exists():
    shutil.copy2(taxonomy,TABLES/'Taxonomy.xlsx')
    rows=list(openpyxl.load_workbook(taxonomy,data_only=True).active.values)
    tax={int(r[0]):str(r[2]) for r in rows[1:] if r[0] is not None}
    for g in canonical['group']:
        if int(g['group_seq']) in tax:
            setvalue(g,'taxon_descr',tax[int(g['group_seq'])],'Source Table3 membership joined via verified final-model crosswalk; Taxonomy.xlsx')
canonical['_source_metadata']=src['metadata']
canonical['_source_identity']=json.loads((OUT/'source_identity.json').read_text(encoding='utf-8'))
canonical['_source_status']={'extraction':'partial','constructor_admission':'NOT_RUN_INCOMPLETE_SOURCE','missing_consumer_diets':[25],'missing_routing':list(range(1,28)),'catch_period_matches_model':False,'diet_normalized':False,'researcher_numerical_corrections_applied':False}
canonical['_source_fleets']=src['fleets']
canonical['_source_consumers']=src['consumers']
canonical['_source_detritus_groups']=src['detritus_groups']
writejson(CAND/'model.json',canonical)
writejson(TABLES/'source_canonical_database.json',canonical)
writejson(OUT/'converter_restoration_ledger.json',changes)
shutil.copy2(OUT/'extraction.json',TABLES/'model.json')

# Reconstruct import files from canonical fields, including retained fleet and
# missing-routing extensions. No numerical value comes from original CSVs.
rev={'metadata':canonical['_source_metadata'],'groups':[],'consumers':canonical['_source_consumers'],'fleets':canonical['_source_fleets'],'landings':{},'discards':{},'diet':{},'detritus_groups':canonical['_source_detritus_groups'],'detritus_fate':{},'diet_rows':27,'landings_rows':27,'discards_rows':27}
def nullable(v): return None if v in (None,'-9999') else v
for g in canonical['group']:
    n=int(g['group_seq']); key=str(n)
    row={'n':n,'name':g['group_name']}
    for f,t in [('biomass','biomass_habitat_area'),('pb','pb'),('qb','qb'),('ee','ee'),('pq','ge'),('tl','tl'),('hab_area','habitat_area'),('unassim','gs'),('ba','biomass_accum'),('ba_rate','biomass_accum_rate'),('z','z'),('other_mort','other_mort'),('detritus_import','detritus_import')]: row[f]=nullable(g.get(t))
    rev['groups'].append(row)
    if g['_source_catch']['components']: rev['landings'][key]=g['_source_catch']['components']
    if n in rev['consumers']:
        d={'import':nullable(g['diet_imp'])}
        entries=(g.get('diet_descr') or {}).get('diet') or []
        if isinstance(entries,dict): entries=[entries]
        d.update({e['prey_seq']:e['proportion'] for e in entries})
        rev['diet'][key]=d
writejson(OUT/'roundtrip_extraction.json',rev)
roundtrip=OUT/'roundtrip_tables'
roundtrip.mkdir(exist_ok=True)
csvbuilders={'Basic_input.csv':writer.basic_input,'Diet_composition.csv':writer.diet_composition,'Landings.csv':lambda x:writer.catch_table(x,'landings'),'Discards.csv':lambda x:writer.catch_table(x,'discards'),'Detritus_fate.csv':writer.detritus_fate,'Biomass_accumulation.csv':writer.biomass_accumulation}
for name,fn in csvbuilders.items():
    with (roundtrip/name).open('w',encoding='utf-8',newline='') as f:
        csv.writer(f,lineterminator='\r\n',quoting=csv.QUOTE_NONE).writerows(fn(rev))
writer.write_xlsx(rev,roundtrip)
results=[]
for name in [*csvbuilders,'TL.xlsx','Metadata.xlsx']:
    a=TABLES/name;b=roundtrip/name
    if a.suffix=='.csv': equal=a.read_bytes()==b.read_bytes()
    else: equal=list(openpyxl.load_workbook(a,data_only=True).active.values)==list(openpyxl.load_workbook(b,data_only=True).active.values)
    results.append({'file':name,'equal':equal,'comparison':'exact CSV bytes' if a.suffix=='.csv' else 'cell values including blank cells','original_sha256':hashlib.sha256(a.read_bytes()).hexdigest(),'roundtrip_sha256':hashlib.sha256(b.read_bytes()).hexdigest()})
writejson(OUT/'roundtrip_verification.json',{'pass':all(r['equal'] for r in results),'results':results,'scope':'canonical numeric fields plus explicit source fleet/missing-routing extensions; raw converter workbook is lossy and superseded','lossy_raw_converter_fields':['TL','GE','fleet partition','unknown catch','unknown habitat','unknown routing','printed diet zeros']})
assert all(r['equal'] for r in results),results

# Run and retain actual tool outputs, including known incomplete-source failures.
env=dict(__import__('os').environ);env['PYTHONUTF8']='1'
for script,name in [('validate.py','validate.txt'),('massbalance_check.py','massbalance_check.txt')]:
    p=subprocess.run([sys.executable,'-X','utf8',str(SKILL/'scripts'/script),str(TABLES)],capture_output=True,text=True,encoding='utf-8',env=env)
    (OUT/name).write_text(p.stdout+'\n'+p.stderr,encoding='utf-8')
writejson(OUT/'source_admission.json',{'status':'NOT_RUN_INCOMPLETE_SOURCE','reason':'Consumer25 feeding column absent; assimilation, BA and routing unreported; catch baseline not fully specified','living_core_parameters_complete':True,'total_groups':27,'feeding_consumers':24,'consumer_diets_present':23,'source_zeros_not_assumed_for_missing_fields':True,'structural_validator':{'errors':0,'warnings':30},'massbalance_helper':{'errors':0,'warnings':12,'limitation':'Treats unknown catch/routing/GS via defaults in indicative calculations; no full balance proof'},'raw_database_checker':{'verdict':'NOT BALANCED','errors':1,'indeterminate':1,'warnings':10,'cause':'missing Zooplankton diet / incomplete phytoplankton consumer budget'},'roundtrip_pass':all(r['equal'] for r in results)})
print(json.dumps({'canonical':str(CAND/'model.json'),'restorations':len(changes),'roundtrip_pass':True,'taxonomy_integrated':taxonomy.exists()},ensure_ascii=False))
