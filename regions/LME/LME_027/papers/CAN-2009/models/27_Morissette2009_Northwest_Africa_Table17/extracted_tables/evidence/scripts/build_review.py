"""Fill the maintained validation template for this unadopted candidate review."""
from pathlib import Path
import json,sys,copy,zipfile,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

C=Path(__file__).resolve().parents[1];ROOT=C.parents[3]
sys.path.insert(0,str(ROOT/'tools'))
from validation_percentage_format import format_percent,verify_report
TEMPLATE=ROOT/'tools/templates/Model_validation_template.docx'
DEST=C/'Model_validation_27_Morissette2009_Northwest_Africa_Table17.docx'
APPENDIX='LME027_candidate_taxon_mapping_appendix.xlsx'
def read(p):return json.loads((C/p).read_text(encoding='utf-8'))
coverage=read('mapping/coverage_summary.json');weak=read('mapping/very_low_decisions.json')
summary=read('diagnostics/summary.json');report=read('diagnostics/direct_reports.json')
d=Document(TEMPLATE)

def put(p,text):
    rp=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(text)
    if rp is not None:r._r.insert(0,rp)
    return p
def cell(c,text):
    prototype=copy.deepcopy(c.paragraphs[0]._p)
    for p in list(c.paragraphs)[1:]:p._p.getparent().remove(p._p)
    parts=text.split('\n');put(c.paragraphs[0],parts[0])
    for text in parts[1:]:
        el=copy.deepcopy(prototype);c._tc.append(el)
        from docx.text.paragraph import Paragraph
        put(Paragraph(el,c),text)
def link(p,label,target):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    for tag,val in [('w:color','0563C1'),('w:u','single'),('w:sz','21')]:
        e=OxmlElement(tag);e.set(qn('w:val'),val);pr.append(e)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def cell_links(c,items):
    p=c.add_paragraph()
    for i,(label,target) in enumerate(items):
        if i:p.add_run(' | ')
        link(p,label,target)
def fill_table(t,rows):
    proto=copy.deepcopy(t.rows[1]._tr)
    for r in list(t.rows)[1:]:t._tbl.remove(r._tr)
    for values in rows:
        t._tbl.append(copy.deepcopy(proto));r=t.rows[-1]
        for c,text in zip(r.cells,values):cell(c,str(text))
def find(prefix):return next(p for p in d.paragraphs if p.text.startswith(prefix))

put(d.paragraphs[0],'Canary Current model candidate review')
intro=d.paragraphs[0].insert_paragraph_before('')
intro._p.getparent().remove(intro._p)
from docx.text.paragraph import Paragraph
e=OxmlElement('w:p');d.paragraphs[0]._p.addnext(e)
p=Paragraph(e,d);p.add_run('The supplied final report describes the same Northwest Africa model as EcoBase 118. It improves source verification but does not establish an independent replacement or remove the TE recycling concern.').font.size=Pt(11)
main=d.tables[0]
entries={
1:'Canary Current | LME_027\nCandidate review; regional selection remains EcoBase 118.',
2:'Sea Around Us; catch records 1950–2019. Reference: 2019 landings, excluding discards. Existing regional catch and classic-PPR inputs retained.',
3:'Morissette, L., Melgo, J.L., Kaschner, K., Gerber, L. and Bamy, I.L. (2009). Food web model and data for studying the interactions between marine mammals and fisheries in the Northwest African ecosystem. Fisheries Centre Research Reports 17(2), Northwest Africa chapter, printed pp. 6–52.',
4:'Guénette et al. (2014), Mauritanian Shelf and Banc d’Arguin, baseline 1991. Villanueva (2004), Sine-Saloum estuary. These are distinct, spatially narrower alternatives; neither is the model extracted here.',
5:'27_Morissette2009_Northwest_Africa_Table17\nPublished balanced Table 17, Northwest Africa, late 1980s; native baseline 1987. 27 source groups. The official EcoBase 118 companion adds one computational Import group. Same published model lineage as the current selection; this directory is a source-review candidate, not a new independent model.',
6:'Northwest Africa Table 16 is an unbalanced parameterization, not an independent ecosystem. The report also contains a 29-group Caribbean model (printed pp. 53–116), outside LME_027. Whale-removal and uncertainty simulations are scenarios.',
7:'Article – The user supplied the previously unavailable final report to assess whether it resolves source gaps or provides another suitable model.\nModel – Review the final balanced Northwest Africa Table 17 and its Table 18 diet. It is the geographically relevant chapter; no replacement has been selected.',
8:'All 126 living B, P/B, Q/B, EE and GE values agree with EcoBase 118 at printed precision. Detritus EE is 0.3695 in Table 17 versus 0 in the native record. Table 3 and Tables 16–18 reorder whale identities; names and parameter signatures resolve the crosswalk.\nTable 18 omits the zooplankton predator column. Printed diet sums range 0.999–1.001 and remain unchanged. The paper alone lacks a complete catch, GS, BA and detritus-routing specification. Native supplementation is retained separately; runtime diet normalization and derived detritus throughput/BA are documented, with no new living-BA adjustment.',
11:'[Researcher calculation choices and notes]',
12:'A. Region covered by study: approximately 95–100%.\nB. Study area covered by region: approximately 20–35%.',
13:'The paper describes the late 1980s; native metadata and catch-table comparison support 1987. The review uses 2019 landings, 32 years later. Fixed coefficients and model-catch mixture proxies across 1950–2019 do not reconstruct annual food-web changes.',
14:'The existing EcoBase 118 model was disqualified on 02/10/2026. Its recorded reason remains: “Reason – too strong living compartments recycling due to near-zero EE values (rho_living = 0.99).”\nNative catch data differ from two 1987 printed cells: local mesopelagic predators are about half the printed value; foreign benthos is absent in the native fleet although the paper reports 426 t. Native catch proportions are therefore explicit model proxies, not exact reproduced observations.\nThe paper’s western bound (30° W) differs from its map/native extent (about 25° W). With Egestion also returns WARN (rho_living 0.1329). No result here establishes suitability for adoption.',
15:'',16:'Researcher name: ____________________ | review date: dd/mm/yyyy'}
for method,idx in [('GE',9),('TE',10)]:
    s=summary[method];dv=s['divergence'];bal=report[method]['balance']
    verdict='WARN — zero-EE source groups and strict balance checks remain unresolved.' if method=='GE' else 'WARN — strong living-compartment recycling remains close to divergence; Coastal tunas have near-zero TE.'
    text=verdict+'\nNo negative SPPR entries across the basal-source columns; all 112 entries are finite.'
    text+=f"\nrho_living = {dv['rho_living']:.4f}"
    if method=='GE':text+=f"; b = {dv['b']:.4f}"
    text+=f"\nDetritus (27): SPPR = {dv['sppr_det']['27']:.4f}."
    text+=f"\nStrict source and PP-balance checks are false; PP relative gap {bal['rel_gap']:.2e}."
    text+=' Diagnostics use the separately retained official native companion; the paper-only reconstruction is NOT_RUN.'
    entries[idx]=text
for i,text in entries.items():cell(main.cell(i,1),text)
cell(main.cell(3,0),'Source article');cell(main.cell(5,0),'Candidate model')
cell_links(main.cell(1,1),[('Regional workbook','../../LME_027.xlsx')])
cell_links(main.cell(2,1),[('Catch evidence','../../raw/LME_027.csv.gz')])
cell_links(main.cell(3,1),[('Final report PDF','../../papers/FCRR_2009_17-2.pdf.pdf#page=10')])
cell_links(main.cell(5,1),[('Paper extraction','model.json'),('Official native input','computational_inputs/official_EcoBase118.json'),('Identity assessment','evidence/identity/identity_review.json')])
cell_links(main.cell(8,1),[('Extraction evidence','extracted_tables/REPORT.md'),('Native runtime changes','diagnostics/transformation_ledger.json')])
cell_links(main.cell(14,1),[('Preserved signed report','../../Model_validation_27_118_Northwest_Africa_(1987).docx'),('Full direct diagnostics','diagnostics/DIRECT_DIAGNOSTICS.md')])

put(find('Reference:'),'Reference: 2019 landings, excluding discards; all recorded catch taxa, including zero-landings labels. Candidate mappings remain unadopted.')
put(find('[n] source groups'),f"27 source groups (+1 computational Import in the native companion); {coverage['taxa']} catch taxa.")
put(find('Simple-chain PPR totals'),f"Simple-chain PPR totals {coverage['total_simple_chain_ppr_tC']:,.0f} t C in 2019.")
put(find('[Missing TL'),f"{len(coverage['missing_classic_coefficients'])} taxa have no saved classic coefficient; all have zero 2019 landings and contribute zero annual PPR. No positive-catch contribution is unknown.")
put(find('Method:'),'Method: independent simple trophic chain using the saved taxon coefficients.')
p=put(find('Excel taxon appendix'),'');link(p,'Excel taxon appendix and descriptive Sources sheet',APPENDIX)
fill_table(d.tables[1],[[r['label'],r['taxa'],format_percent(r['catch_percentage']),format_percent(r['ppr_percentage'])] for r in coverage['confidence_summary']])
fill_table(d.tables[2],[[r['plain_language_rule'],r['confidence'],format_percent(r['ppr_percentage'])] for r in sorted(coverage['membership_summary'],key=lambda r:-r['ppr_percentage'])])
fill_table(d.tables[3],[[r['plain_language_rule'],r['confidence'],format_percent(r['ppr_percentage'])] for r in sorted(coverage['allocation_summary'],key=lambda r:-r['ppr_percentage'])])
p=put(find('Membership evidence:'),'Final Table 3, verified taxonomic identities and reporting-scope evidence. ');link(p,'Complete taxon decisions','mapping/taxon_audit.json')
p=put(find('Weights and assumptions:'),'Native model-catch proportions supply assumed mixtures where a split is needed, transferred from 1987 to 2019 landings. ');link(p,'Allocation quantities and alternatives','mapping/allocation_evidence.json')
v=next(r for r in coverage['confidence_summary'] if r['label']=='Very low')
method_exposure=read('calculations/candidate_calculation_summary.json')['methods']
te_exposure=next(x for x in method_exposure if x['method']=='TE')['very_low_ppr_pct']
put(find('Unresolved taxa:'),f"Unresolved taxa: none. {v['taxa']} Very low decisions account for {format_percent(v['catch_percentage'])} of catch and {format_percent(v['ppr_percentage'])} of independent simple-chain PPR, but {format_percent(te_exposure)} of the native-companion TE calculation. Complete assignment coverage is not evidence that these ecological analogues or mixture assumptions are correct.")
short_reasons={
'Actinopterygii':'Broad bony-fish reporting spans seven named and residual fish pools. Its actual composition is unknown; native model-catch proportions impose a historical mixture, including mixed-guild limitations.',
'Alepisaurus ferox':'Deep-pelagic source members support Mesopelagic predators as the nearest habitat analogue. Alepisauridae is absent, and the pool does not establish this large lancetfish’s size, depth or diet composition.',
'Anguilla anguilla':'Coastal demersal contains marine eels and other bottom fishes. The unlisted European eel is diadromous; that marine pool does not represent its freshwater-to-sea life cycle.',
'Branchiostegus':'A demersal tilefish genus is placed with source bottom fishes in Coastal demersal. Malacanthidae, burrow-associated ecology and the genus’s depth composition are unrepresented; a deeper pool remains plausible.',
'Echelus myrus':'Source conger and moray groups provide an eel and bottom-habitat connection to Coastal demersal. Ophichthidae is absent; ecological equivalence is unverified. The deeper eel-bearing pool fits the reported habitat less well.',
'Echinoidea':'Source Benthos supplies a bottom-invertebrate analogue through bivalves, gastropods and anthozoans. No echinoderm member is printed; extending this pooled shellfish/benthos coefficient is a weak ecological assumption.',
'Epigonus telescopus':'Source deep-bottom fishes support Bathydeersal predators as a depth analogue. Epigonidae is absent, so transferring its mixed composition to black cardinalfish remains uncertain.',
'Heteropriacanthus cruentatus':'Source reef/coastal fishes support Coastal demersal for these priacanthids. Priacanthidae is absent, and equivalent bigeye feeding and depth composition are unverified.',
'Kyphosus sectatrix':'Sarpa salpa, Acanthuridae and other source reef/coastal fishes provide an analogue. Kyphosidae is absent; pooled membership and equivalent feeding composition remain uncertain.',
'Lepas':'Crustaceans provides a taxonomic connection to source decapods. This pelagic attached barnacle has a body-form and habitat mismatch; no cirripede compartment exists.',
'Marine groundfishes not identified':'The reporting label spans Coastal demersal, Bathydeersal and source Mesopelagic predators, which includes demersal Mora moro. Unknown composition and mixed cartilage/pelagic source members weaken the historical native-catch mixture.',
'Marine pelagic fishes not identified':'Compatible offshore, mesopelagic, tuna and coastal pools are retained, including explicitly pelagic members of Coastal demersal. Unobserved composition is approximated by native model-catch proportions.',
'Megalops atlanticus':'Large pelagics supplies a mobile large-fish analogue. Megalopidae is absent and the tarpon’s coastal/estuarine ecology differs from this offshore pool; the coastal Elops-containing pool is a meaningful alternative.',
'Miscellaneous aquatic invertebrates':'Unknown catch may span cephalopods, crustaceans, benthos and zooplankton. Mixed membership and composition remain uncertain. Native catch weighting retains Zooplankton at a recorded zero weight, not proven absence.',
'Perciformes':'Historical percomorph reporting spans several modern orders and five compatible source fish pools. The actual catch mixture is unknown and approximated by native model-catch proportions.',
'Peristediidae':'The reported bathydemersal category supports a deep-bottom analogue, but the source lacks Peristediidae. Coastal demersal contains competing true-gurnard relatives; depth motivates the weaker selected placement.',
'Polymixia nobilis':'Beryx and other deep-bottom source fishes support Bathydeersal predators. Polymixiidae is absent; this habitat analogue does not establish actual source membership.',
'Sarotherodon melanotheron':'Coastal demersal contains estuarine/coastal fishes, providing the nearest habitat connection. The unlisted brackish-water cichlid lacks a dedicated estuarine/freshwater compartment; transfer of the marine coefficient is weak.',
'Scombroidei':'Historical reporting breadth spans the source large-pelagic, coastal-tuna and other coastal-pelagic pools. Native catch proportions approximate an unknown mixture; a single printed family label does not define the entire category.',
'Uranoscopus':'Coastal demersal contains benthic predatory fish and flatfish. Uranoscopidae is absent, and the stargazers’ benthic ambush ecology is not isolated; deeper predators are a weaker habitat alternative.'}
assert {r['taxa'][0] for r in weak}==set(short_reasons)
fill_table(d.tables[4],[[r['affected_taxa'],short_reasons[r['taxa'][0]]] for r in weak])
put(find('Geographic evidence screenshots'),'Geographic evidence')
put(find('Target region R'),'Target region R is the Canary Current LME; S is the much larger Northwest Africa study domain. The approximate overlap range brackets the final figure/native extent and the conflicting printed bounds.')
put(find('[Region name/ID'),'Canary Current LME_027; retained regional boundary. Comparison uses compatible geodesic areas and excludes land.')
put(find('[Article; figure/page'),'Morissette et al. (2009), Figure 1, printed p. 8 / PDF p. 12. Original article figure; the offshore western edge differs from the longitude given in the prose.')
put(find('[If a figure is unavailable'),'')
find('Study area in the article').paragraph_format.page_break_before=True
for ti,filename,width in [(5,'target_region.png',6.1),(6,'final_figure1_with_caption.png',6.1)]:
    cell(d.tables[ti].cell(1,0),'')
    d.tables[ti].cell(1,0).paragraphs[0].add_run().add_picture(str(C/'evidence/geography'/filename),width=Inches(width))
    d.tables[ti].rows[1].height=None
    p=d.tables[ti].cell(1,0).add_paragraph();link(p,'Boundary and overlap evidence','evidence/geography/geographic_assessment.json')
cell(d.tables[6].cell(0,0),'Original study figure for the reviewed model')

for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for h in p._p.findall(qn('w:hyperlink')):
                    assert h.find('.//'+qn('w:u')) is not None
d.save(DEST)
# Keep opaque template parts byte-identical. python-docx is used only for editable content/layout slots.
with zipfile.ZipFile(TEMPLATE) as z:original={x:z.read(x) for x in z.namelist()}
with zipfile.ZipFile(DEST) as z:filled={x:z.read(x) for x in z.namelist()}
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
for name,data in original.items():
    if name not in editable:filled[name]=data
with zipfile.ZipFile(DEST,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in filled.items():z.writestr(name,data)
verify_report(DEST)
proof={'template_sha256':hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
       'preserve_only_parts_equal':all(filled[n]==v for n,v in original.items() if n not in editable),
       'manual_unsigned':True,'candidate_directory_override':'User explicitly requested co-located candidate outputs.',
       'docx':DEST.name,'appendix':APPENDIX}
(C/'qa/template_fidelity.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
print(DEST.name)
