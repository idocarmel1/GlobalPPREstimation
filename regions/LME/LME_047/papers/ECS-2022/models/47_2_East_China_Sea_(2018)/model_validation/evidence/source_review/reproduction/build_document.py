from pathlib import Path
import sys,json,copy,zipfile,urllib.parse,shutil
from lxml import etree
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path.cwd();Q=Path(__file__).parent;REG=ROOT/'regions/LME_047';MID='47_2_East_China_Sea_(2018)';OUT=REG/'validation_reports'/MID
sys.path.insert(0,str(ROOT/'tools'));from workbooks import sha
REF=ROOT/'tools/templates/Model_validation_template.docx';OUTPUT=REG/f'Model_validation_{MID}.docx';TEMP=Q/'generated.docx';shutil.copy2(REF,Q/'template_reference.docx')
def j(p):return json.loads(Path(p).read_text(encoding='utf-8'))
a=j(OUT/'coverage_summary.json');di=j(OUT/'matrix_inspection.json');geo=j(OUT/'geography/geographic_assessment.json');ex=j(OUT/'method_contribution_exposure.json')['methods']
d=Document(REF);p=d.paragraphs;manual=[copy.deepcopy(d.tables[t].rows[r]._tr) for t,r in [(0,9),(1,6),(1,7)]]
def clear(par):
 for n in list(par._p):
  if n.tag!=qn('w:pPr'):par._p.remove(n)
 return par
def text(par,value):clear(par);par.add_run(value);return par
def cell(c,value):
 first=c.paragraphs[0]
 for pp in list(c._tc):
  if pp.tag==qn('w:p') and pp is not first._p:c._tc.remove(pp)
 return text(first,value)
def link(par,label,target):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),par.part.relate_to(urllib.parse.quote(target,safe='/:?=&%#'),RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'0563C1');rp.append(co);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def addlinks(par,links):
 for i,(label,target) in enumerate(links):par.add_run(' | ' if i else ' ');link(par,label,target)
def put(t,r,value,links=[]):par=cell(d.tables[t].cell(r,1),value);addlinks(par,links)
def pct(x):return f'{x:.4f}%' if x>=.00005 or x==0 else '<0.0001%'
O=OUT
ev='validation_reports/'+MID+'/';main='papers/ECS-2022/pdf-e60d0358.pdf';sup='papers/ECS-2022/DataSheet_1_EstimatingtheImpactofaSeasonal-7491f315.docx';legacy='papers/LME047-Li-2012/10152_2011_278_MOESM1_ESM-d161693b.doc'
text(p[0],'East China Sea model validation')
put(0,1,'East China Sea, LME_047; Sea Around Us target boundary.',[('Regional workbook','LME_047.xlsx')])
put(0,2,'Sea Around Us v50.1, 1950–2019. Reference: 2019 landings, including reported and unreported retained catch; discards separate.',[('Raw catch','raw/LME_047-catch.zip'),('Taxon/TL source','raw/LME_047-exploited.json')])
put(0,3,'Xu, Song, Wang, Xie, Huang, Li, Zheng & Lin (2022), Estimating the Impact of a Seasonal Fishing Moratorium on the East China Sea Ecosystem From 1997 to 2018. Frontiers in Marine Science 9, 865645; DOI 10.3389/fmars.2022.865645.',[('Main paper',main),('Publisher supplement',sup),('Paper folder','papers/ECS-2022/')])
put(0,4,'Li & Zhang (2012), A trophic model of the East China Sea and its possible application. Helgoland Marine Research; DOI 10.1007/s10152-011-0278-7. Publisher supplement retained; focal main PDF unavailable. Cheng et al. (2009), Changes in the ecosystem structure of the East China Sea 1969–2000; known inventory source unavailable.',[('Li supplement',legacy),('Other source inventory',ev+'source_identity_verification.json')])
put(0,5,MID+'; static M2018, autumn 2018/spring 2019 survey data. 24 ecological groups plus synthetic diet_import (25). Figure 1 shows a localized shelf study area; exact modeled boundary/area is unstated.',[('Selected JSON','models/'+MID+'/model.json'),('Saved computational workbook','models/'+MID+'/sppr_source.xlsx')])
put(0,6,'47_1_East_China_Sea_(1997) remains a separate, unselected counterpart, using the 1997–2000 resource survey and different parameter/diet state. Ecosim M2018 SFM/no-SFM scenarios are simulations, not alternative static M2018 source exports.',[('1997 JSON','models/47_1_East_China_Sea_(1997)/model.json')])
put(0,7,'Article: ? (separate article-choice reason unrecorded). Model: “2018 preferred; 1997 retained separately. Model coefficients held fixed across catch years.” This recorded recency preference does not establish whole-LME spatial suitability.',[('Recorded rationale','LME_047.xlsx#Overview!A1')])
put(0,8,'All populated M2018 Table 1 B/PB/QB/EE cells match accepted inputs. Canonical diets match normalized S3 columns; 91 raw decimal cells differ across columns summing 0.999–1.001. Canonical export 0 becomes catch; those source catch zeros are unverified placeholders. GS 0.2, derived accumulation, migration defaults, single-detritus routing and effective detritus EE 1 versus published 0.241 remain accepted accounting assumptions.',[('Source review',ev+'source_review.md'),('Diet provenance',ev+'diet_normalization_provenance.json')])
for row,label in [(1,'GE'),(2,'TE')]:
 dd=next(x['diagnostic']['divergence']for x in di['methods']if x['method']==label)
 value=f"No negative SPPR entries across the basal-source columns, including unfished groups. rho_living={dd['rho_living']:.8f}; "
 if label=='GE':value+=f"b={dd['b']:.8f}; "
 value+=f"Detritus (24) SPPR={dd['sppr_det']['24']:.8f}.";put(1,row,value)
put(1,3,'Approximate study-footprint A: 10–15% of the target region. B: 90–100% of the displayed study area inside the target. Exact model boundary is unstated.')
put(1,4,'Static M2018 coefficients and catch/biomass proxy splits are transferred across 1950–2019. Reference 2019 overlaps the spring survey, but fixed guild composition does not represent observed change across years, fleets or catch bases.')
put(1,5,'Direct GE, TE and With Egestion remain WARN; strict input and budget balance pass. EE 0 in Sharks (22) and Marine mammals (23) yields TE EEfix SPPR 0. Diagnostic footprint 0 uses unverified model catch 0; positive SAU catch yields positive regional PPR. Source defaults and spatial transfer remain limits; existing production eligibility is preserved.',[('Full matrices',ev+'matrix_inspection.json'),('Method exposure',ev+'method_contribution_exposure.json')])
text(p[3],'Reference: 2019 landings; all exact catch labels, including unidentified and zero-catch records; no taxon or group filter.')
text(p[4],'24 ecological source groups (+1 computational import); 276 catch labels.')
text(p[5],f"Simple-chain PPR totals {a['total_simple_chain_ppr_tC']:,.2f} t C in 2019.")
text(p[6],'22 labels lack a saved classic coefficient/TL; all have zero 2019 landings and zero annual PPR. Unavailable TL remains ?, not zero.')
text(p[7],'Method: independent classic/simple trophic chain using the saved taxon coefficient and a single carbon conversion.')
clear(p[8]);link(p[8],'Full taxon mapping appendix','LME047_taxon_mapping_appendix.xlsx');p[8].add_run(' and ');link(p[8],'descriptive Sources sheet','LME047_taxon_mapping_appendix.xlsx#Sources!A1')
for row,ss in zip(d.tables[2].rows[1:],a['confidence_summary']):
 for c,v in zip(row.cells,[ss['label'],str(ss['taxa']),pct(ss['catch_percentage']),pct(ss['ppr_percentage'])]):cell(c,v)
desc={'M1':'Exact source member in one group or a documented author union.','M2':'Verified synonym connects the catch label to a source member.','M3':'Verified taxon fits the explicit Bivalves or Cnidaria pool.','M4':'Unlisted taxon extends documented source representatives.','M5':'Habitat and feeding support an inferred guild.','M6':'Partial ecological fit retains size, habitat or competing-pool uncertainty.','M9':'Source members support an inferred family/order group set.','M10':'Broad reporting label uses a plausible pooled mixture of unknown composition.','M11':'Closest represented analogue retains material habitat, feeding or taxonomic mismatch.','M12':'Older ECS lobster evidence transfers to focal Crabs.','W1':'One reviewed source group; no numerical split.','W4':'Complete positive 2018 source catch proportions proxy the taxon mixture across years and bases.','W9':'Complete M2018 group biomass proportions proxy composition after source catch is incomplete or censored.'}
for ti,field in [(3,'membership_summary'),(4,'allocation_summary')]:
 table=d.tables[ti];proto=copy.deepcopy(table.rows[1]._tr)
 for row in list(table.rows)[1:]:table._tbl.remove(row._tr)
 for ss in a[field]:
  table._tbl.append(copy.deepcopy(proto))
  for c,v in zip(table.rows[-1].cells,[desc[ss['rule']],ss['confidence'],pct(ss['ppr_percentage'])]):cell(c,v)
clear(p[11]);p[11].add_run('Membership evidence: ');link(p[11],'focal Table S1',sup);p[11].add_run(' and ');link(p[11],'all taxon decisions',ev+'taxon_audit.json');p[11].add_run('. Representative lists are not exhaustive family definitions. Bony residuals exclude dedicated cartilage but retain eligible mixed guilds and named fish groups.')
text(p[13],'Overall confidence is the weaker membership/allocation component. The rule percentages use the complete known classic-PPR universe. 12 split labels use complete positive 2018 S4 source catch; 15 use complete M2018 biomass after incomplete/censored catch. Both allocations are Medium proxies. S4 is M1997 calibration data, not a native static M2018 catch export. Whole-guild proportions are not measured taxon/stage shares; fixed year, area, fleet and basis transfer remains an assumption. ');link(p[13],'Allocation evidence',ev+'allocation_evidence.json');p[13].add_run(' | ');link(p[13],'Historical donor reconstruction',ev+'identified_catch_proxy_review.json')
vl=next(r for r in a['confidence_summary']if r['label']=='Very low');allmethods=[x for x in ex if x['scope']=='all'];shares={x['method']:next(z['ppr_percentage']for z in x['confidence_contribution']if z['label']=='Very low')for x in allmethods}
text(p[14],f"No unresolved labels remain. Very low: {vl['taxa']} labels, {vl['catch_percentage']:.4f}% of landings and {vl['ppr_percentage']:.4f}% of classic PPR. Broad residuals, northern-water/grazing fish, gastropod and whale-shark analogues retain explicit limits. Mola uses a weak zooplankton-feeding fish analogue; Scaridae a weak coastal mixed-feeding analogue. Their Very low cohort contributes {shares['new_GE']:.2f}% of GE, {shares['new_TE_EEfix']:.2f}% of TE and {shares['new_WithEgestion']:.2f}% of With Egestion PPR under the same reference controls. ");link(p[14],'Method-specific contributions',ev+'method_contribution_exposure.json')
text(p[16],'R is the full East China Sea LME target. S reconstructs the blue hatch in Figure 1, whose caption calls it the study area. Methods identify the M2018 joint autumn 2018/spring 2019 survey but supply no separate explicit model boundary or area. Coverage therefore describes this study footprint, with model-domain representativeness unresolved. Native graticule calibration preserves the offshore gap, western bends and southern/northeastern boundary; eight water controls pass.')
for ti,title,img,width in [(5,'East China Sea target region',O/'geography/target_region.png',5.0),(6,'Original East China Sea study area — Figure 1',O/'geography/original_figure1.png',4.65)]:
 cell(d.tables[ti].cell(0,0),title);cell(d.tables[ti].cell(1,0),'').add_run().add_picture(str(img),width=Inches(width))
text(p[18],f"Sea Around Us LME 47 target R: {geo['target_area_km2']:,.0f} km². ");addlinks(p[18],[('Target boundary',ev+'geography/target_region.geojson'),('Map image',ev+'geography/target_region.png')])
text(p[20],'Xu et al. 2022 Figure 1, PDF 3: “The study area in the East China Sea.” §2.2, PDF 4: M2018 uses the joint ECS survey in autumn 2018 and spring 2019. The generic regional description supplies context, not a whole-LME native model boundary. ');addlinks(p[20],[('Original figure',main+'#page=3'),('Methods',main+'#page=4'),('Geographic assessment',ev+'geography/geographic_assessment.json')])
text(p[21],f"Approximate marine study area {geo['reconstructed_study_marine_area_km2']:,.0f} km²; A≈{geo['A_percent']:.2f}%, B≈{geo['B_percent']:.2f}%. Descriptive ranges 10–15% and 90–100% reflect trace, registration and coastline uncertainty. Natural Earth ne_50m means 1:50 million scale, not 50 metres. Coverage is not a PPR multiplier. ");link(p[21],'Reconstructed footprint',ev+'geography/study_area_reconstructed.geojson');p[21].add_run().add_break();p[21].add_run().add_picture(str(O/'geography/comparison.png'),width=Inches(5.8))
for table in d.tables:
 table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in table.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
for idx in [2,15,19,21]:p[idx].paragraph_format.page_break_before=True
for idx in [17,19]:p[idx].paragraph_format.keep_with_next=True
for t,r,original in [(0,9,manual[0]),(1,6,manual[1]),(1,7,manual[2])]:d.tables[t].rows[r]._tr.getparent().replace(d.tables[t].rows[r]._tr,original)
d.save(TEMP);changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with zipfile.ZipFile(REF)as ref,zipfile.ZipFile(TEMP)as gen,zipfile.ZipFile(OUTPUT,'w',zipfile.ZIP_DEFLATED)as dest:
 for item in ref.infolist():dest.writestr(item,gen.read(item.filename)if item.filename in changed else ref.read(item.filename))
 for item in gen.infolist():
  if item.filename not in ref.namelist():dest.writestr(item,gen.read(item.filename))
check=Document(OUTPUT);assert[etree.tostring(check.tables[t].rows[r]._tr)for t,r in [(0,9),(1,6),(1,7)]]==[etree.tostring(x)for x in manual];assert etree.tostring(check.sections[0]._sectPr)==etree.tostring(Document(REF).sections[0]._sectPr)
with zipfile.ZipFile(REF)as ref,zipfile.ZipFile(OUTPUT)as final:
 preserved=[n for n in ref.namelist()if n not in changed];assert all(ref.read(n)==final.read(n)for n in preserved)
(O/'document_fidelity.json').write_text(json.dumps({'manual_rows_identical':True,'section_geometry_identical':True,'preserved_reference_package_parts':preserved,'template_sha256':sha(REF)},indent=2),encoding='utf-8');print('DOCX',OUTPUT)
