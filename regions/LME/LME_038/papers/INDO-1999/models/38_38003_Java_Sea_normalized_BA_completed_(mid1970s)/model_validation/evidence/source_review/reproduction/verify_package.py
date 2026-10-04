from pathlib import Path
import json,zipfile,urllib.parse,shutil,math
from lxml import etree as X
from openpyxl import load_workbook
Q=Path(__file__).parent;ROOT=Q.parents[4];R=ROOT/'regions/LME_038';MID='38_38003_Java_Sea_normalized_BA_completed_(mid1970s)';E=R/'validation_reports'/MID
def dump(n,v):(E/n).write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
a=json.loads((E/'taxon_audit.json').read_text(encoding='utf-8'));by={r['taxon']:r for r in a};s=json.loads((E/'coverage_summary.json').read_text(encoding='utf-8'));ordered=sorted(a,key=lambda r:(-r['simple_chain_ppr_tC'],r['taxon']));xp=R/'LME038_taxon_mapping_appendix.xlsx';dp=R/f'Model_validation_{MID}.docx';w=load_workbook(xp,data_only=False);assert w.sheetnames==['Taxon mapping','Sources'];ws=w['Taxon mapping'];assert ws.freeze_panes=='B8'and ws.tables
headers=['Taxon name','TL','Catch (t)','Simple-chain PPR (t C)','Mapped group names and weights','Confidence level','Reason'];assert[next(ws.iter_rows(min_row=7,max_row=7,values_only=True))[i]for i in range(7)]==headers
rows=list(ws.iter_rows(min_row=8,max_row=188,values_only=True));assert len(rows)==181 and [r[0]for r in rows]==[r['taxon']for r in ordered];assert len({r[0]for r in rows})==181
for rr in rows:
 r=by[rr[0]];assert rr[1]==(r['tl']if r['tl']is not None else'?');assert isinstance(rr[2],(int,float))and isinstance(rr[3],(int,float));assert math.isclose(rr[2],r['catch_tonnes'],abs_tol=1e-9,rel_tol=1e-12);assert math.isclose(rr[3],r['simple_chain_ppr_tC'],abs_tol=1e-7,rel_tol=1e-12);assert rr[4]==r['mapping_display']and rr[5]==r['overall_confidence']
for field in ['membership_summary','allocation_summary']:assert math.isclose(math.fsum(r['ppr_percentage']for r in s[field]),100,abs_tol=1e-8)
assert math.isclose(math.fsum(r['ppr_tC']for r in s['confidence_summary']),s['total_simple_chain_ppr_tC'],abs_tol=1e-7);assert sum(r['taxa']for r in s['confidence_summary'])==181
links=[];formulas=[]
for sheet in w:
 for row in sheet:
  for c in row:
   if c.data_type=='f':formulas.append([sheet.title,c.coordinate,c.value])
   if c.hyperlink:
    assert c.font.color.rgb=='FF0563C1'and c.font.underline=='single';links.append(dict(file=str(xp),label=c.value,target=c.hyperlink.target or'#'+c.hyperlink.location))
w.close();assert not formulas
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
with zipfile.ZipFile(dp)as z:
 x=X.fromstring(z.read('word/document.xml'));rels={r.get('Id'):r.get('Target')for r in X.fromstring(z.read('word/_rels/document.xml.rels'))}
 assert len(x.findall('.//w:drawing',ns))==3
 for h in x.findall('.//w:hyperlink',ns):
  assert all(r.find('w:rPr/w:color',ns).get('{'+ns['w']+'}val')=='0563C1'and r.find('w:rPr/w:u',ns).get('{'+ns['w']+'}val')=='single'for r in h.findall('w:r',ns));links.append(dict(file=str(dp),label=''.join(h.itertext()),target=rels[h.get('{'+ns['r']+'}id')]))
for p in [xp,dp]:
 with zipfile.ZipFile(p)as z:
  for n in ['docProps/custom.xml','docProps/app.xml']:
   if n in z.namelist():
    for el in X.fromstring(z.read(n)).iter():
     if X.QName(el).localname=='HyperlinkBase':assert not(el.text or'').strip()
reloc=Q/'Relocated validation repository';reloc.mkdir(exist_ok=True);local=0
for p in [xp,dp]:dest=reloc/p.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
for l in links:
 target=urllib.parse.unquote(l['target'])
 if target.startswith('http')or target.startswith('#'):continue
 pathname=target.split('#')[0];assert not(':'in pathname or pathname.startswith(('/','\\')));dest=(Path(l['file']).parent/pathname).resolve();assert dest.exists(),str(dest);assert ROOT in dest.parents
 moved=reloc/dest.relative_to(ROOT)
 if dest.is_file():moved.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(dest,moved)
 else:moved.mkdir(parents=True,exist_ok=True)
 anchor=reloc/Path(l['file']).relative_to(ROOT);assert(anchor.parent/pathname).exists();local+=1
source_ids={r['id']for r in json.loads((E/'appendix_sources.json').read_text(encoding='utf-8'))};assert{ss for r in a for ss in r['sources']}<=source_ids
dump('link_verification.json',dict(links=links,local_targets=local,physical_relocation_passed=True,relocation_root_relative_to_scratch=reloc.name,all_visible_links_blue_underlined=True,hyperlink_base_absent=True))
dump('package_verification.json',dict(all181_taxa_exact_once=True,full_precision_numeric_order=True,all_numeric_catch_ppr_types=True,seven_columns_exact=True,both_tabs_verified=True,no_formulas_or_errors=True,weights_display_matches=True,summary_denominators_and_percentages_match=True,frozen_headers_filters=True,source_ids_complete=True,docx_3_figures_present=True,hyperlinks_verified=len(links),relative_targets_relocated=local,all_visible_links_blue_underlined=True,browser_verified=False))
print('Package numeric/schema/link checks passed',len(links),'links;',local,'relocated local targets')
