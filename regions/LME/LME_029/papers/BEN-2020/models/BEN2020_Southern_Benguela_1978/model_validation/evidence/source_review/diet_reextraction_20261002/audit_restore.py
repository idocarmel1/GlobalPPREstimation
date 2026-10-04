"""Bounded LME029 normalization audit; preserves selected and historical inputs."""
from pathlib import Path
from decimal import Decimal
from copy import deepcopy
import json, hashlib, csv, shutil, importlib.util, sys, warnings
from zipfile import ZipFile
from lxml import etree
import pymupdf as fitz
import openpyxl

HERE=Path(__file__).resolve().parent
REGION=HERE.parent
ROOT=REGION.parents[1]
MODEL=REGION/'models/BEN2020_Southern_Benguela_1978'
EX=REGION/'extraction_review_20260928'
SELECTED=MODEL/'model.json'
CONVERTED=MODEL/'extracted_tables/29_Benguela_Current_20201978_Southern_Benguela_(1978).json'
RECON=MODEL/'extracted_tables/29_Benguela_Current_20201978_Southern_Benguela_(1978)_reconstructed.xlsx'
SUPP=REGION/'papers/BEN-2020/Data_Sheet_1_Exploring_Temporal_Variabilit-5c4a190c.pdf'
DOCX=REGION/'Model_validation_BEN2020_Southern_Benguela_1978.docx'
INTEGRATION=MODEL/'integration_20260928'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p):return p.relative_to(ROOT).as_posix()
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(name,x):(HERE/name).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def cells(d):
    out={}
    for gi,g in enumerate(d['group']):
        n=str(g['group_seq']); ds=(g.get('diet_descr') or {}).get('diet',[])
        if isinstance(ds,dict):ds=[ds]
        for di,v in enumerate(ds):out[(n,str(v['prey_seq']))]=(v['proportion'],f'/group/{gi}/diet_descr/diet/{di}/proportion')
        out[(n,'import')]=(g.get('diet_imp','0'),f'/group/{gi}/diet_imp')
    return out
def differences(d,ex):
    cc=cells(d)
    return [(n,p,v,cc.get((n,p),('0','absent_zero'))[0]) for n,col in ex['diet'].items() for p,v in col.items() if Decimal(v)!=Decimal(cc.get((n,p),('0',''))[0])]

assert not (HERE/'root_receipt.json').exists(),'Do not repeat a completed mutation'
paths=[p for p in REGION.rglob('*') if p.is_file() and HERE not in p.parents]
before={rel(p):sha(p) for p in paths}
write('before_manifest.json',before)
archive=HERE/'originals';archive.mkdir(exist_ok=True)
inputs=[SELECTED,CONVERTED,RECON,INTEGRATION/'29_20201978_Southern_Benguela_PQ_completed_(1978).json',MODEL/'diagnostics/29_20201978_Southern_Benguela_PQ_completed_(1978).json']
for p in inputs:
    target=archive/(sha(p)+p.suffix)
    if not target.exists():shutil.copyfile(p,target)
    assert sha(target)==before[rel(p)]

ex=read(EX/'extraction_input.json');selected=read(SELECTED);old=read(CONVERTED)
assert differences(selected,ex)==[],'Selected differs; manual investigation required'
for p in inputs[3:]:assert differences(read(p),ex)==[]
diff=differences(old,ex)
assert len(diff)==31 and {v[0] for v in diff}=={'37','42'}
assert Decimal(sum(Decimal(v) for v in ex['diet']['37'].values()))==Decimal('1.0001')
assert Decimal(sum(Decimal(v) for v in ex['diet']['42'].values()))==Decimal('1.0002')
log=(MODEL/'extracted_tables/database_conversion_execution.txt').read_text(encoding='utf-8')
assert 'Diet Sum is not 1 (1.00010). Normalizing to 1.0' in log
assert 'Diet Sum is not 1 (1.00020). Normalizing to 1.0' in log

# Re-read ONLY the two consumers proven normalized, with continued-page headers.
fresh={};fresh_evidence=[]
with fitz.open(SUPP) as pdf:
    header=pdf[13].find_tables().tables[0].extract()[0]
    predators=[int(v) for v in header[2:]]
    assert '* indicates 0.0001' in ' '.join(pdf[11].get_text().split())
    for pi in [13,14]:
        table=pdf[pi].find_tables().tables[0];rows=table.extract()
        if pi==13:rows=rows[1:]
        for ri,row in enumerate(rows):
            prey=str(int(row[0])) if row[0] else 'import'
            for pred in [37,42]:
                ci=predators.index(pred)+2;literal=row[ci];v='0.0001' if literal=='*' else literal.strip()
                assert ex['diet'][str(pred)][prey]==v
                fresh[(str(pred),prey)]=v
                fresh_evidence.append(dict(consumer=pred,prey=prey,printed_literal=literal,adopted_literal=v,page=pi+1,table='S4',bbox=table.rows[ri+(1 if pi==13 else 0)].cells[ci],source=rel(SUPP),source_sha256=sha(SUPP)))
assert len(fresh)==100
write('confirmed_consumer_reextraction.json',fresh_evidence)

# Read-only researcher-decision inspection. No tracked changes/comments exist.
with ZipFile(DOCX) as z:
    xml=etree.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paras=[' '.join(p.itertext()) for p in xml.findall('.//w:p',ns)]
    doc_review=dict(path=rel(DOCX),sha256=sha(DOCX),tracked_insertions=len(xml.findall('.//w:ins',ns)),tracked_deletions=len(xml.findall('.//w:del',ns)),comments_present='word/comments.xml' in z.namelist(),relevant_paragraphs=[p for p in paras if any(k in p.lower() for k in ['diet','normal','researcher'])],finding='Source sums1.0001/1.0002 explicitly preserved. No accepted diet/import override identified. Researcher name/date remain placeholders; this audit grants no model trust approval.')
write('researcher_decision_inspection.json',doc_review)

ledger=[];cc=cells(selected);ec=cells(old)
previous_evidence={(str(e['predator']),str(e['prey'])):e for e in read(EX/'pdf_cell_evidence.json') if 'predator' in e}
for n,col in ex['diet'].items():
    for prey,literal in col.items():
        ev=previous_evidence[(n,prey)]
        ledger.append(dict(consumer=n,prey=prey,source_literal=ev['printed'],adopted_source_literal=literal,selected_literal=cc.get((n,prey),('0','absent_zero'))[0],selected_pointer=cc.get((n,prey),('0','absent_zero'))[1],converted_before_literal=ec.get((n,prey),('0','absent_zero'))[0],source=rel(SUPP),source_sha256=sha(SUPP),page=ev['page'],table=ev['table'],bbox=ev['bbox'],accepted_correction=None,review_state='not-established: no researcher diet correction/signoff identified'))
write('diet_cell_ledger.json',ledger)

new=deepcopy(old);changed=[]
for gi,g in enumerate(new['group']):
    n=str(g['group_seq'])
    if n not in ['37','42']:continue
    for di,d in enumerate(g['diet_descr']['diet']):
        prey=str(d['prey_seq']);value=fresh[(n,prey)]
        assert Decimal(d['proportion'])!=Decimal(value)
        changed.append(dict(file=rel(CONVERTED),pointer=f'/group/{gi}/diet_descr/diet/{di}/proportion',consumer=n,prey=prey,before=d['proportion'],after=value,provenance=next(e for e in fresh_evidence if str(e['consumer'])==n and e['prey']==prey)))
        d['proportion']=value
assert len(changed)==31 and differences(new,ex)==[]
stripped_before=deepcopy(old);stripped_after=deepcopy(new)
for d in [stripped_before,stripped_after]:
    for g in d['group']:
        for v in (g.get('diet_descr') or {}).get('diet',[]):v['proportion']='DIET_REDACTED'
assert stripped_before==stripped_after,'Non-diet mutation'

# Modify only the corresponding diet cells and their two displayed totals.
wb=openpyxl.load_workbook(RECON);sheet=wb['Diet composition']
cols={str(sheet.cell(1,c).value):c for c in range(3,sheet.max_column+1)}
rows={str(sheet.cell(r,1).value):r for r in range(2,sheet.max_row+1)}
wb_changes=[]
for c in changed:
    cell=sheet.cell(rows[c['prey']],cols[c['consumer']]);assert abs(float(cell.value)-float(c['before']))<=1e-15
    wb_changes.append(dict(sheet=sheet.title,coordinate=cell.coordinate,before=cell.value,after=float(c['after'])))
    cell.value=float(c['after'])
for n in ['37','42']:
    cell=sheet.cell(rows['Sum'],cols[n]);total=sum(Decimal(v) for v in ex['diet'][n].values())
    wb_changes.append(dict(sheet=sheet.title,coordinate=cell.coordinate,before=cell.value,after=float(total),kind='displayed sum'))
    cell.value=float(total)
temp_xlsx=HERE/'restored_reconstructed.xlsx';wb.save(temp_xlsx)
check=openpyxl.load_workbook(temp_xlsx)
original=openpyxl.load_workbook(RECON)
allowed={(c['sheet'],c['coordinate']) for c in wb_changes}
semantic_changes=[]
for name in original.sheetnames:
    a=original[name];b=check[name]
    assert (a.max_row,a.max_column)==(b.max_row,b.max_column)
    for row in a:
        for cell in row:
            other=b[cell.coordinate]
            assert cell.style_id==other.style_id
            if cell.value!=other.value:
                assert (name,cell.coordinate) in allowed,(name,cell.coordinate)
                semantic_changes.append((name,cell.coordinate))
assert len(semantic_changes)==33

# Compare actual retained input/state/settings with a fresh load, without diagnostics.
spec=importlib.util.spec_from_file_location('lme029_retained',INTEGRATION/'reproduce_direct.py')
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
execution=read(INTEGRATION/'execution_evidence.json')
assert execution['constructor']==module.SETTINGS and execution['diagnostic']==module.DIAGNOSTIC
assert sha(module.INPUT)==execution['input_sha256']
md=module.ModelData(str(module.INPUT));md.lme=29
with warnings.catch_warnings(record=True) as notices:
    warnings.simplefilter('always');calc=module.PPRCalculator.from_modeldata(md,**module.SETTINGS)
current=module.state(calc);retained=read(INTEGRATION/'loaded_state.json');numdiff=[];structdiff=[];maxdiff=0.
def compare(a,b,path=''):
    global maxdiff
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):
        delta=abs(a-b);maxdiff=max(maxdiff,delta)
        if delta>1e-12:numdiff.append(dict(path=path,retained=a,reloaded=b,absolute_difference=delta))
    elif isinstance(a,dict) and isinstance(b,dict):
        if set(a)!=set(b):structdiff.append(path+' keys')
        for k in set(a)&set(b):compare(a[k],b[k],path+'/'+k)
    elif isinstance(a,list) and isinstance(b,list):
        if len(a)!=len(b):structdiff.append(path+' length')
        for i,(x,y) in enumerate(zip(a,b)):compare(x,y,path+'/'+str(i))
    elif a!=b:structdiff.append(path)
compare(retained,current)
assert not numdiff and not structdiff,'Actual retained runtime not reproduced; no mutation'
write('actual_loaded_state_reload.json',current)
runtime=dict(tolerance=dict(atol=1e-12,rtol=0),selected_input_sha256=sha(SELECTED),historical_computational_input_sha256=sha(module.INPUT),constructor=module.SETTINGS,diagnostic_settings=module.DIAGNOSTIC,retained_state_file=rel(INTEGRATION/'loaded_state.json'),retained_state_sha256=sha(INTEGRATION/'loaded_state.json'),new_state_file=rel(HERE/'actual_loaded_state_reload.json'),new_state_sha256=sha(HERE/'actual_loaded_state_reload.json'),state_fields=list(current),maximum_absolute_numerical_difference=maxdiff,numerical_differences_exceeding_tolerance=numdiff,structural_differences=structdiff,constructor_warnings=[str(w.message) for w in notices],runtime_equivalent=True,refresh_required=False,reason='Selected canonical and actual historical P/Q computational input are unchanged. Restored conversion artifact is not the input used by retained computations. Actual full loaded state reproduces with identical constructor/diagnostic settings. normalize_DC=False is the retained setting; runtime normalization remains always permitted.',engine_sha256={p.name:sha(p) for p in (INTEGRATION/'engine').glob('*.py')},source_runtime_bridge='Canonical source e11d2bcd…5994 remains distinct from P/Q completed computational f2288b46…b84b. No historical hash is rewritten.')
write('runtime_equivalence.json',runtime)

# Guard exact originals immediately before saving both authorized artifacts.
assert sha(CONVERTED)==before[rel(CONVERTED)] and sha(RECON)==before[rel(RECON)]
assert sha(SELECTED)==before[rel(SELECTED)] and sha(DOCX)==before[rel(DOCX)]
CONVERTED.write_text(json.dumps(new,ensure_ascii=False,indent=2),encoding='utf-8')
assert sha(RECON)==before[rel(RECON)]
shutil.copyfile(temp_xlsx,RECON)
assert differences(read(CONVERTED),ex)==[]
write('changed_cells.json',dict(json_changes=changed,reconstructed_workbook_changes=wb_changes))

sums=[]
for n,col in ex['diet'].items():
    total=sum(Decimal(v) for v in col.values());imported=Decimal(col['import'])
    sums.append(dict(consumer=n,name=selected['group'][int(n)-1]['group_name'],prey_sum=str(total-imported),import_sum=str(imported),raw_source_sum=str(total),canonical_sum=str(total),converted_after_sum=str(total),runtime_sum=float(calc._DC.loc[int(n)].sum()),missing_cells=False,review_state='not-established',normalization_factor_actual_runtime='1 (normalize_DC=False retained)'))
write('consumer_sums.json',sums)
after={rel(p):sha(p) for p in paths}
changed_files=[p for p in before if before[p]!=after[p]]
assert set(changed_files)=={rel(CONVERTED),rel(RECON)},changed_files
historical={p:h for p,h in before.items() if '/diagnostics/' in p or '/integration_20260928/' in p or '/previous_results/' in p}
write('historical_hashes_preserved.json',dict(files=historical,all_unchanged=all(after[p]==h for p,h in historical.items())))
write('protection_verification.json',dict(scope='All pre-existing regional files',changed_files=[dict(file=p,before_sha256=before[p],after_sha256=after[p]) for p in changed_files],protected_files_unchanged=all(before[p]==after[p] for p in before if p not in changed_files),unchanged_file_count=len(before)-len(changed_files),detritus_fate_and_unrelated_json_fields_semantically_unchanged=stripped_before==stripped_after,reconstructed_non_diet_fields_semantically_unchanged=True,canonical_selected_unchanged=before[rel(SELECTED)]==after[rel(SELECTED)],guard='SHA256 checked just before saving; no other current artifacts modified.'))
write('root_receipt.json',dict(region='LME_029',selected_model_id='BEN2020_Southern_Benguela_1978',selected_path=rel(SELECTED),outcome='mixed: verified unnormalized selected model unchanged; confirmed normalized current extracted conversion JSON and reconstructed diet cells restored',normalization_status='selected/canonical model.json and historical computational inputs proven source-faithful; only current extracted conversion artifact had converter normalization of consumers37/42',reextraction_performed=True,reextraction_extent='Only confirmed normalized consumers37/42 freshly read from S4/PDF14–15; no re-extraction of selected unnormalized model',canonical_before_sha256=before[rel(SELECTED)],canonical_after_sha256=after[rel(SELECTED)],converted_before_sha256=before[rel(CONVERTED)],converted_after_sha256=after[rel(CONVERTED)],changed_diet_cells=31,changed_reconstructed_diet_cells=31,changed_reconstructed_sum_cells=2,evidence_paths=[rel(p) for p in HERE.glob('*.json')],precise_blockers=['No researcher diet signoff identified; verification of exact source transcription is not native/model scientific approval. Existing all-method FAIL, missing native stanza/BA interpretation and other source limitations remain unchanged.'],protected_files_unchanged=True,runtime_equivalent=True,runtime_comparison_atol=1e-12,runtime_comparison_rtol=0,refresh_required=False,trust_status='not-established; no new researcher review approval',historical_input_hashes_preserved=True))
print(json.dumps(dict(outcome='completed',selected_sha256=sha(SELECTED),converted_before=before[rel(CONVERTED)],converted_after=sha(CONVERTED),changed_diet_cells=31,runtime_max_abs_diff=maxdiff,refresh_required=False,changed_files=changed_files)))
