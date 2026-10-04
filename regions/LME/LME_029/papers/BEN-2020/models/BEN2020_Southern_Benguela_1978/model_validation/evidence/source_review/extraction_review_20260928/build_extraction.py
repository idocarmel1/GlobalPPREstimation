from pathlib import Path
import pymupdf as fitz, json, re, unicodedata, hashlib, csv, subprocess, sys
from decimal import Decimal
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[3]; REGION=ROOT/'regions/LME_029'; OUT=REGION/'extraction_review_20260928'; PAPER=REGION/'papers/BEN-2020'; MODEL=REGION/'models/BEN2020_Southern_Benguela_1978'; TABLES=MODEL/'extracted_tables'; TABLES.mkdir(parents=True,exist_ok=True)
SKILL=Path('C:/Users/idoca/.agents/skills/ecopath-extraction')
def dump(p,x):p.write_text(json.dumps(x,indent=2,ensure_ascii=False),encoding='utf-8')
def clean(s):return ' '.join(unicodedata.normalize('NFKC',s).split())
def norm(s):return re.sub(r'[^a-z0-9]','',clean(s).lower()).replace('andswordfish','swordfish')
s=fitz.open(next(PAPER.glob('Data*.pdf')));m=fitz.open(PAPER/'pdf-00952390.pdf');evidence=[]
# S4 gridded coordinates, numeric headers anchor predator identities. Different from Table 1 listing order.
diet={}; names={};raw_tables={}
for pi in [11,12,13,14]:
 table=s[pi].find_tables().tables[0];rows=table.extract();raw_tables[str(pi+1)]=rows
 if pi in [11,13]:cons=[int(v) for v in rows.pop(0)[2:]]
 for ri,row in enumerate(rows):
  n=int(row[0]) if row[0] else 'import';name=clean(row[1]);
  if n!='import':
   if n in names:assert norm(names[n])==norm(name),(n,names[n],name)
   names.setdefault(n,name)
  for ci,(pred,v) in enumerate(zip(cons,row[2:])):
   assert v is not None
   value='0.0001' if v=='*' else clean(v);Decimal(value)
   assert str(n) not in diet.setdefault(str(pred),{})
   diet[str(pred)][str(n)]=value
   realri=ri+(1 if pi in [11,13] else 0)
   evidence.append({'source':'supplement','page':pi+1,'table':'S4','prey':n,'predator':pred,'printed':v,'value':value,'bbox':table.rows[realri].cells[ci+2]})
assert len(names)==49 and len(diet)==45 and all(len(v)==50 for v in diet.values())
# Main Table 2: use distinct x anchors for continued page and exact numeric y alignment.
groups=[];basic_evidence=[]
for pi,anchors in [(7,[160,233,317,388,441,488,535]),(8,[147,222,309,381,437,486,535])]:
 words=m[pi].get_text('words'); limit=710 if pi==7 else 160
 starts=[w for w in words if abs((w[0]+w[2])/2-anchors[0])<8 and 95<w[1]<limit and re.fullmatch(r'\d+\.\d+',w[4])]
 for st in starts:
  row=[w for w in words if abs(w[1]-st[1])<1];name=clean(' '.join(w[4] for w in sorted(row) if w[0]<120)).rstrip('*')
  keys=['tl','biomass','z','pb','qb','ee','pq'];g={'name':name}
  for key,x in zip(keys,anchors):
   vals=[w for w in row if abs((w[0]+w[2])/2-x)<15 and re.fullmatch(r'\d+\.\d+',w[4])]
   assert len(vals)<=1,(name,key,vals)
   g[key]=vals[0][4] if vals else None
   if vals:basic_evidence.append({'source':'main','page':pi+1,'table':'2','group_name':name,'field':key,'value':g[key],'bbox':list(vals[0][:4])})
  groups.append(g)
assert len(groups)==49,len(groups)
for i,g in enumerate(groups,1):
 assert norm(g['name'])==norm(names[i]),(i,g['name'],names[i])
 g['n']=i;g['name']=names[i]
# Source S1 explicitly supplies PB for eight stanzas. Deep-water hake S1 PB is blank.
for i,pb in {7:'1.2',8:'1.2',9:'1.4',10:'1.2',13:'1.2',14:'1.0',28:'2',29:'0.8'}.items():
 assert groups[i-1]['pb'] is None;groups[i-1]['pb']=pb
 basic_evidence.append({'source':'supplement','page':1,'table':'S1','group_name':names[i],'field':'pb','value':pb,'note':'Stanza PB explicitly reported in S1; Table 2 separately labels corresponding value total mortality.'})
for i in [9,10]:groups[i-1]['ba_rate']='0.3'
groups[46]['ba']='-0.603'
# Prose GS allocation only for unambiguous zooplanktivorous fish with exclusively plankton diet.
# Other conditional trophic classes are retained as prose evidence because exact scope/threshold is unspecified.
for i in [11,12,13,16,17]:groups[i-1]['unassim']='0.35'
# Catch S3 two 31-fleet column blocks. Resolve cropped rotated-header overflow using Table 1 and visual page 9.
landings={};fleets=[];printed_totals={};fleet_sums={};catch_evidence=[]
for pi in [6,7,8,9]:
 table=s[pi].find_tables().tables[0];rows=table.extract();raw_tables['catch_'+str(pi+1)]=rows
 if pi in [6,8]:
  fs=[clean(x) for x in rows.pop(0)[1:]]
  if pi==8:fs=['Dem Shark Longline','Beach Seine&Gillnet','Large Sparids WC line','Medium Sparids WC line','Sciaenids WC line','Chonds WC line','Snoek WC line','Tuna & billfish WC line','Yellowtail WC line','Large Sparids SC line','Medium Sparids SC line','Sciaenids SC line','Chonds SC line','Other','Total']
  fleets.extend(x for x in fs if x!='Total')
 for ri,row in enumerate(rows):
  nm=clean(row[0]);matches=[n for n,name in names.items() if norm(name)==norm(nm)]
  if nm=='Sum':
   for fl,v in zip(fs,row[1:]):fleet_sums[fl]=v
   continue
  assert len(matches)==1,(nm,matches);n=matches[0]
  for ci,(fl,v) in enumerate(zip(fs,row[1:])):
   Decimal(v)
   if fl=='Total':printed_totals[n]=v
   else:landings.setdefault(str(n),{})[fl]=v
   catch_evidence.append({'source':'supplement','page':pi+1,'table':'S3','group':n,'fleet':fl,'value':v,'bbox':table.rows[ri+(1 if pi in [6,8] else 0)].cells[ci+1]})
assert len(fleets)==31 and len(landings)==48 and all(len(v)==31 for v in landings.values())
catch_checks=[]
for n,v in landings.items():
 total=sum(Decimal(x) for x in v.values());published=Decimal(printed_totals[int(n)])
 catch_checks.append({'seq':int(n),'group':names[int(n)],'sum_fleets':str(total),'printed_total':str(published),'difference':str(total-published)})
# Taxonomic composition follows source definitions; S2 is EXAMPLES, not an exhaustive membership inventory.
tax={i:'not documented — group name only; no exhaustive membership list in main Table 1 or supplement S2' for i in names}
tax.update({1:'Small phytoplankton <10 micrometres (main Table 1).',2:'Larger phytoplankton >=10 micrometres (main Table 1).',3:'Microzooplankton 2–200 micrometres equivalent spherical diameter (main Table 1).',4:'Mesozooplankton 200–2000 micrometres (main Table 1).',5:'Macrozooplankton 2–20 mm (main Table 1).',6:'Gelatinous zooplankton; species not documented.',7:'Engraulis encrasicolus; recruits <1 year (main Table 1).',8:'Engraulis encrasicolus; adults >=1 year (main Table 1).',9:'Sardinops sagax; recruits <1 year (main Table 1).',10:'Sardinops sagax; adults >=1 year (main Table 1).',11:'Etrumeus whiteheadii (main Table 1).',12:'Other small pelagic fish; gobies mentioned in supplement section 3.2; exhaustive membership not documented.',13:'Trachurus trachurus capensis; juveniles <2 years (main Table 1).',14:'Trachurus trachurus capensis; adults >=2 years (main Table 1).',15:'Scomber japonicus (main Table 1).',16:'Lampanyctodes hectoris (main Table 1).',17:'Maurolicus muelleri (main Table 1).',18:'Thyrsites atun (main Table 1).',19:'Thunnus spp.; Sarda sarda (main Table 1). Table 2/S4 call this Tuna and Swordfish but Table 1 also names Atlantic bonito; swordfish identity not explicitly specified.',20:'Examples from supplement S2: Chrysoblephus cristiceps; Petrus rupestris; Polysteganus undulosus.',21:'Examples from supplement S2, South and West Coast subsets combined in this functional group: Argyrozona argyrozona; Chrysoblephus laticeps; Cheimerius nufar; Chrysoblephus gibbiceps; Spondyliosoma emarginatum; Pachymetopon blochii; Rhabdosargus globiceps.',22:'Examples from supplement S2: Atractoscion aequidens; Argyrosomus inodorus.',23:'Seriola lalandii (main Table 1).',24:'Other linefish; exhaustive membership not documented.',25:'Liza richardsonii; haarder (main Table 1).',26:'Loligo reynaudii; Chokka squid (main Table 1).',27:'Other cephalopods; exhaustive membership not documented.',28:'Merluccius capensis; small <3 years (main Table 1).',29:'Merluccius capensis; large >=3 years (main Table 1).',30:'Merluccius paradoxus; small <3 years (main Table 1).',31:'Merluccius paradoxus; large >=3 years (main Table 1).',34:'Austroglossus pectoralis (main Table 1); also appears in S2 benthic-feeding demersals examples, a source overlap requiring priority to this dedicated group.',37:'Examples from supplement S2: Carcharodon carcharias; Hexanchus griseus; Notorhinchus cepedianus.',38:'Arctocephalus pusillus pusillus (main Table 1).',39:'Cetaceans; exhaustive membership not documented.',40:'Spheniscus demersus (main Table 1).',41:'Morus capensis (main Table 1).',42:'Phalacrocorax capensis (main Table 1).',43:'Other seabirds including migrants feeding outside model; 25% imported diet (supplement section 3.8).',44:'Benthic producers; exhaustive membership not documented.',45:'Meiobenthos; exhaustive membership not documented.',46:'Macrobenthos; exhaustive membership not documented.',47:'Jasus lalandii (main Table 1; supplement section 3.9 spells Jasus lalandi).',48:'Source conflict: main Table 1 says Panulirus homarus; supplement section 3.9 says Palinurus gilchristi. Do not automatically treat these distinct taxa as synonyms.',49:'Nonliving detritus pool.'})
# Preserve whole S2 row texts from explicit functional-group headings, including printed spelling.
s2='\n'.join(s[i].get_text() for i in [1,2])
starts={32:'Pelagic-feeding demersal fish',33:'Benthic-feeding demersal fish',35:'Pelagic-feeding chondrichthyans',36:'Benthic-feeding chondrichthyans'}
ends={32:'Benthic-feeding demersal fish',33:'Pelagic-feeding chondrichthyans',35:'Benthic-feeding chondrichthyans',36:'Apex chondrichthyans'}
for i,head in starts.items():tax[i]='Examples of species included (supplement S2; original spellings; not exhaustive): '+clean(s2.split(head,1)[1].split(ends[i],1)[0])
ex={'metadata':{'LME':'29 Benguela Current','model_number':20201978,'model_name':'Southern Benguela','model_year':1978},'groups':groups,'consumers':sorted(map(int,diet)),'fleets':fleets,'landings':landings,'discards':{},'detritus_groups':['Detritus'],'detritus_fate':{},'diet':diet,'diet_rows':49,'landings_rows':48,'discards_rows':48}
dump(OUT/'extraction_input.json',ex);dump(OUT/'pdf_cell_evidence.json',evidence+basic_evidence+catch_evidence);dump(OUT/'raw_coordinate_tables.json',raw_tables);dump(OUT/'catch_total_checks.json',catch_checks)
manifest=[{'file':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'role':'main article' if p.name.startswith('pdf-') else 'supplement' if p.suffix=='.pdf' else 'inherited metadata/geometry'} for p in PAPER.iterdir() if p.is_file()]
dump(OUT/'source_manifest.json',manifest)
subprocess.run([sys.executable,str(SKILL/'scripts/write_outputs.py'),str(OUT/'extraction_input.json'),'--outdir',str(MODEL),'--dir-name','extracted_tables'],check=True)
w=Workbook();ws=w.active;ws.title='Taxonomy';ws.append(['seq','group_name','taxon_descr'])
for i in sorted(names):ws.append([i,names[i],tax[i]])
w.save(TABLES/'Taxonomy.xlsx');dump(OUT/'taxonomy_evidence.json',[{'seq':i,'group_name':names[i],'taxon_descr':tax[i]} for i in names])
for script in ['validate.py','massbalance_check.py']:
 p=subprocess.run([sys.executable,str(SKILL/'scripts'/script),str(TABLES)],capture_output=True,text=True,encoding='utf-8',errors='replace');(TABLES/(script[:-3]+'_execution.txt')).write_text(p.stdout+p.stderr,encoding='utf-8');print(script,p.returncode)
p=subprocess.run([sys.executable,str(SKILL/'scripts/database_json.py'),'-d',str(TABLES),'--update-report'],capture_output=True,text=True,encoding='utf-8',errors='replace');(TABLES/'database_conversion_execution.txt').write_text(p.stdout+p.stderr,encoding='utf-8');print('database',p.returncode)
print('groups',len(groups),'diet cells',sum(len(v) for v in diet.values()),'catch differences',[r for r in catch_checks if Decimal(r['difference'])]);print('diet sums',[(k,str(sum(Decimal(x) for x in v.values()))) for k,v in diet.items()])
