"""Scoped, exact-cell diet audit; never rewrites canonical or frozen history."""
from pathlib import Path
from decimal import Decimal
import csv, hashlib, json, shutil, io, sys
import openpyxl
from docx import Document

ROOT = Path.cwd()
REG = ROOT / 'regions/LME_037'
MID = 'Bacalso2026_Visayan_Sea_1997_baseline'
MODEL = REG / 'models' / MID
EXT = MODEL / 'extracted_tables'
OLD = REG / 'models/extraction_review_20260928'
RUN = REG / 'models/regional_ge_integration_20260928'
OUT = REG / 'diet_reextraction_20261002'
OUT.mkdir(exist_ok=True)
verify_applied = '--verify-applied' in sys.argv

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p): return p.relative_to(ROOT).as_posix()
def load(p): return json.loads(p.read_text(encoding='utf-8'))
def save(name, data): (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
def decimal(v): return Decimal(str(v))

protected = {rel(p): sha(p) for p in REG.rglob('*') if p.is_file() and OUT not in p.parents}
canonical = MODEL / 'model.json'
before = sha(canonical)
groups = load(canonical)['group']
by_id = {g['group_seq']: g for g in groups}
overview_book = openpyxl.load_workbook(REG / 'LME_037.xlsx', read_only=True, data_only=False)
overview = dict((r[0],r[1]) for r in overview_book['Overview'].values if len(r)>1 and r[0])
overview_book.close()
assert overview['selected_model_id'] == MID
assert REG / overview['model_path'] == canonical
assert overview['results_model_sha256'] == before

source_doc = REG / 'papers/LME037-Bacalso-2026/recovered_sources_20260928/Bacalso2023_DataSheet_1.docx'
source_downloads=[]
for entry in load(source_doc.parent/'DOWNLOAD_MANIFEST.json'):
    actual=sha(source_doc.parent/entry['file'])
    assert actual==entry['sha256']
    source_downloads.append({'path':rel(source_doc.parent/entry['file']),'manifest_sha256':entry['sha256'],'actual_sha256':actual,'matches':True})
doc = Document(source_doc)
source_tables = {i:[[cell.text.strip() for cell in row.cells] for row in doc.tables[i-1].rows] for i in (3,4)}
ledger = []
diet_evidence = [x for x in load(OLD / 'cell_evidence.json') if x['table']=='S2']
assert len(diet_evidence) == 990
extracted = load(OLD / 'extraction_input.json')
csv_path = EXT / 'Diet_composition.csv'
csv_rows = list(csv.reader(csv_path.open(encoding='utf-8-sig', newline='')))
csv_col = {cid: i for i,cid in enumerate(csv_rows[0]) if cid in by_id}
csv_prey_row = {row[0]: i for i,row in enumerate(csv_rows) if row[0] in by_id}
xlsx_path = next(EXT.glob('*reconstructed.xlsx'))
xbook = openpyxl.load_workbook(xlsx_path)
sheet = xbook['Diet composition']
xcol = {str(sheet.cell(1,c).value): c for c in range(3,sheet.max_column+1)}
xrow = {str(sheet.cell(r,1).value):r for r in range(2,sheet.max_row+1)}
for item in diet_evidence:
    cid, prey = item['predator'], item['prey']
    text = source_tables[item['docx_table']][item['row']-1][item['column']-1]
    assert text == item['source_text'], item
    adopted = next(d['proportion'] for d in by_id[cid]['diet_descr']['diet'] if d['prey_seq']==prey)
    assert decimal(adopted)==decimal(item['value'])==decimal(extracted['diet'][cid][prey])
    assert decimal(csv_rows[csv_prey_row[prey]][csv_col[cid]])==decimal(adopted)
    assert decimal(sheet.cell(xrow[prey],xcol[cid]).value)==decimal(adopted)
    ledger.append({**item, 'source_path':rel(source_doc),'source_sha256':sha(source_doc),
                   'canonical_pointer':f'/group/{int(cid)-1}/diet_descr/diet/{int(prey)-1}/proportion',
                   'canonical_literal':adopted,'import_csv_cell':[csv_prey_row[prey]+1,csv_col[cid]+1],
                   'export_xlsx_cell':sheet.cell(xrow[prey],xcol[cid]).coordinate,
                   'source_matches':True,'accepted_correction':None,'researcher_review':'not established by this audit'})

restoration = load(OLD / 'converter_source_restoration.json')
imports_history = {str(x['group']):x for x in restoration if x['field']=='diet_imp'}
imports_csv_row = next(i for i,r in enumerate(csv_rows) if len(r)>1 and r[1]=='Import')
changed = []
for cid,g in by_id.items():
    assert g['diet_imp']=='-9999'
    h=imports_history[cid]
    assert h['source_faithful_value']=='-9999' and h['converter_value']=='0'
    ledger.append({'consumer':cid,'kind':'diet_import','canonical_pointer':f'/group/{int(cid)-1}/diet_imp',
                   'canonical_literal':'-9999','source_literal':None,'source_status':'not reported/unknown',
                   'retained_restoration_path':rel(OLD/'converter_source_restoration.json'),
                   'retained_restoration_sha256':sha(OLD/'converter_source_restoration.json'),
                   'retained_restoration_index':restoration.index(h),'accepted_correction':None,
                   'researcher_review':'not established; runtime zero default is documented separately'})
    if cid in csv_col:
        column=csv_col[cid]
        assert csv_rows[imports_csv_row][column]==('' if verify_applied else '0')
        changed.append({'path':rel(csv_path),'row':imports_csv_row+1,'column':column+1,'consumer':cid,'field':'diet_import','before':'0','after':'','reason':'preserve proven unknown canonical import sentinel as empty import cell'})
        csv_rows[imports_csv_row][column]=''
    cell=sheet.cell(xrow['Import'],xcol[cid])
    assert cell.value is None if verify_applied else cell.value==0
    changed.append({'path':rel(xlsx_path),'cell':cell.coordinate,'consumer':cid,'field':'diet_import','before':0,'after':None,'reason':'preserve proven unknown canonical import sentinel as empty export cell'})
    cell.value=None

# Archive exact originals and compare immediately before every save.
archives=[]
if verify_applied:
    archives=load(OUT/'export_repairs_pending.json')['archived_originals']
    for item in archives:
        assert sha(ROOT/item['archive_path'])==item['before_sha256']
for target in ([] if verify_applied else [csv_path,xlsx_path]):
    digest=protected[rel(target)]
    assert sha(target)==digest, 'Concurrent change: '+rel(target)
    archive=OUT/'originals'/digest/target.name
    archive.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(target,archive)
    assert sha(archive)==digest
    archives.append({'path':rel(target),'before_sha256':digest,'archive_path':rel(archive)})
if '--plan-only' in sys.argv:
    save('source_cell_ledger.json',ledger)
    save('export_repairs_pending.json',{'changed_cells':changed,'archived_originals':archives,'applied':False,'source_cells_checked':990})
    print(json.dumps({'outcome':'PLAN_VERIFIED_HELD','canonical_sha256':before,'source_cells_checked':990,'pending_export_import_cells':len(changed)}))
    sys.exit(0)
if not verify_applied:
    assert sha(csv_path)==protected[rel(csv_path)]
    stream=io.StringIO(newline='')
    csv.writer(stream,lineterminator='\n').writerows(csv_rows)
    csv_path.write_text(stream.getvalue(),encoding='utf-8',newline='')
    assert sha(xlsx_path)==protected[rel(xlsx_path)]
    xbook.save(xlsx_path)
xbook.close()
for item in archives: item['after_sha256']=sha(ROOT/item['path'])

# Verify full export-cell scope: exactly imports differ, with no other sheet/field changes.
original_xlsx=ROOT/archives[1]['archive_path']
wb_old=openpyxl.load_workbook(original_xlsx,data_only=False)
wb_new=openpyxl.load_workbook(xlsx_path,data_only=False)
export_differences=[]
assert wb_old.sheetnames==wb_new.sheetnames
for s in wb_old.sheetnames:
    a,b=wb_old[s],wb_new[s]
    assert (a.max_row,a.max_column)==(b.max_row,b.max_column)
    for row in a:
        for oldcell in row:
            newcell=b[oldcell.coordinate]
            assert oldcell.style_id==newcell.style_id
            if oldcell.value!=newcell.value:
                export_differences.append({'sheet':s,'cell':oldcell.coordinate,'old':oldcell.value,'new':newcell.value})
assert len(export_differences)==33
assert all(x['sheet']=='Diet composition' and x['old']==0 and x['new'] is None for x in export_differences)
wb_old.close();wb_new.close()
csv_old=list(csv.reader((ROOT/archives[0]['archive_path']).open(encoding='utf-8-sig',newline='')))
assert len(csv_rows)==len(csv_old)
csv_differences=[(r,c) for r in range(len(csv_old)) for c in range(len(csv_old[r])) if csv_old[r][c]!=csv_rows[r][c]]
assert len(csv_differences)==30 and all(r==imports_csv_row for r,c in csv_differences)

rv=load(RUN/'runtime_verification.json')
engine_trust=[]
for name, expected in rv['engine_sha256'].items():
    actual=sha(RUN/'executed_code'/name)
    assert actual==expected
    engine_trust.append({'path':rel(RUN/'executed_code'/name),'expected_historical_sha256':expected,'actual_sha256':actual,'matches':True})
comp_input=next((RUN/'input').glob('37_*.json'))
comp=load(comp_input)
assert rv['canonical_sha256']==before
assert rv['computational_input_sha256']==sha(comp_input)
assert rv['computational_state_sha256']==sha(RUN/'computational_state.json')
assert overview['computational_state_sha256']==sha(RUN/'computational_state.json')
canonical_vs_comp=[]
for a,b in zip(groups,comp['group']):
    assert a['group_seq']==b['group_seq']
    for field in a:
        if a[field]!=b[field]: canonical_vs_comp.append({'group':a['group_seq'],'field':field,'source':a[field],'computational':b[field]})
assert len(canonical_vs_comp)==33 and all(x['field']=='diet_imp' and x['source']=='-9999' and x['computational']=='0' for x in canonical_vs_comp)
runtime_rows=list(csv.reader((RUN/'runtime_diet.csv').open()))
runtime_cols={c:i for i,c in enumerate(runtime_rows[0]) if i}
runtime={r[0]:r for r in runtime_rows[1:]}
source_runtime_rows=list(csv.reader((RUN/'source_loaded_diet.csv').open()))
assert runtime_rows==source_runtime_rows
sums=[]
for cid,g in by_id.items():
    known_diet=[d for d in g['diet_descr']['diet'] if decimal(d['proportion']) != Decimal(-9999)]
    raw=sum((decimal(d['proportion']) for d in known_diet),Decimal(0))
    runtime_sum=sum((decimal(v) for v in runtime[cid][1:]),Decimal(0))
    diffs=[abs(decimal(runtime[cid][runtime_cols[d['prey_seq']]])-decimal(d['proportion'])) for d in known_diet]
    max_diff=max(diffs,default=Decimal(0))
    assert max_diff<=Decimal('1e-12')
    sums.append({'consumer':cid,'name':g['group_name'],'prey_sum_exact':str(raw),'canonical_import_literal':g['diet_imp'],
                 'diet_plus_import_sum':None,'diet_plus_import_sum_status':'unknown import; prey sum is not full source total',
                 'runtime_import_default':0,'runtime_diet_plus_import_sum':str(runtime_sum),
                 'runtime_normalize_DC':False,'normalization_factor_applied':'1','max_source_runtime_diet_difference':str(max_diff),
                 'unknown_nonfeeding_placeholder_cells':[d['prey_seq'] for d in g['diet_descr']['diet'] if decimal(d['proportion']) == Decimal(-9999)],
                 'review_state':'not established by audit'})
copies=[MODEL/'37_3702026_Visayan_Sea_Bacalso_baseline_(1997).json',EXT/'37_3702026_Visayan_Sea_Bacalso_baseline_(1997).json',EXT/'canonical_source_database.json']
assert all(sha(x)==before for x in copies)
assert sha(canonical)==before
changed_paths={rel(csv_path),rel(xlsx_path)}
protected_checks=[{'path':p,'before_sha256':h,'after_sha256':sha(ROOT/p),'unchanged':sha(ROOT/p)==h} for p,h in protected.items() if p not in changed_paths]
assert all(x['unchanged'] for x in protected_checks)
history_checks={'canonical':before,'computational_input':sha(comp_input),'computational_state':sha(RUN/'computational_state.json'),
                'retained_runtime_settings':rv['constructor'],'retained_input_and_state_hashes_match':True,
                'historical_direct_diagnostics_preserved':True,'saved_runtime_diet_equals_loaded_source_diet':True,
                'historical_engine_trust':engine_trust,'source_download_manifest_verified':source_downloads}
review_doc_path=REG/'Model_validation_Bacalso2026_Visayan_Sea_1997_baseline.docx'
review_doc=Document(review_doc_path)
review_texts=[p.text for p in review_doc.paragraphs]+[c.text for t in review_doc.tables for row in t.rows for c in row.cells]
review_snippets=[t for t in review_texts if any(s in t.lower() for s in ['accepted diet','without normalization','researcher name','manual','other adjustments'])]
save('researcher_edits_readonly.json',{'path':rel(review_doc_path),'sha256':sha(review_doc_path),'relevant_text':review_snippets,
                                     'interpretation':'Existing diet source sums and accounting defaults preserved. Researcher name/date and adjustment sections remain placeholders; no new signoff inferred.'})
save('source_cell_ledger.json',ledger)
save('consumer_sums.json',sums)
save('export_repairs.json',{'changed_cells':changed,'archived_originals':archives,'xlsx_value_differences':export_differences,
                          'prey_cells_checked':990,'canonical_changed':False,'reextraction_performed':False,
                          'sum_rows_retained':'Numerical rows remain exact prey-only sums. Full diet-plus-import totals are unavailable because imports are unknown; this is explicit in consumer_sums.json.'})
save('protected_files_verification.json',protected_checks)
save('historical_hash_trust.json',history_checks)
save('runtime_comparison.json',{'abs_tolerance':1e-12,'relative_tolerance':0,'runtime_differs':False,'full_loaded_state_changed':False,
                              'comparison_ignored_fields':['biomass accumulation','predation'],'ignored_changes':[],
                              'comparison_basis':'Canonical, explicit computational input, saved complete loaded state and actual settings retain identical verified identities. Repairs affect only source export views, not any runtime input. No new scientific run is needed.',
                              'canonical_before_sha256':before,'canonical_after_sha256':sha(canonical),'runtime_identity':history_checks,
                              'source_vs_computational_differences':canonical_vs_comp,'normalized_runtime':False,
                              'scientific_refresh_required':False,'normalization_always_permitted':True})
save('root_receipt.json',{'region':'LME_037','selected_model_id':MID,'selected_path':rel(canonical),
                        'outcome':'verified unnormalized canonical unchanged; corrected 63 exact unknown-import export cells',
                        'normalization_status':'Canonical retains all 990 exact S2 source diet cells; old converter normalization already restored in 20260928 history. Source exports lost 30/33 unknown imports as zeros; now corrected.',
                        'reextraction_performed':False,'canonical_before_sha256':before,'canonical_after_sha256':sha(canonical),
                        'changed_diet_cells':changed,'changed_canonical_diet_cells':[],'runtime_differs':False,'runtime_tolerance_abs':1e-12,
                        'runtime_comparison_ignored_fields':['biomass accumulation','predation'],'ignored_runtime_changes':[],
                        'required_refresh_scope':[],'evidence_paths':[rel(p) for p in sorted(OUT.glob('*.json'))],
                        'precise_blockers':[],'scientific_limits':['Unknown source imports remain unknown; computational zero imports are retained documented defaults.','Researcher signoff is not established by this audit; manual researcher sections remain placeholders.','1997 baseline uses verified 2023 predecessor S2 diets under documented 2026 lineage; unselected 2018 endpoint is outside scope.'],
                        'protected_files_unchanged':True})
(OUT/'AUDIT_REPORT.md').write_text('''# LME037 diet normalization audit — 2026-10-02

The selected Bacalso2026_Visayan_Sea_1997_baseline model is verified unnormalized and unchanged. All 990 diet cells agree exactly with the retained official 2023 supplement Table S2, extraction ledger, active import CSV and reconstructed workbook. Canonical source SHA-256 remains `'''+before+'''`. The three named/source canonical copies retain the same bytes.

September converter output had normalized some columns; the retained restoration ledger already reinstated source literals. That frozen converter output and its historical diagnostic reports remain evidence of their actual earlier inputs, and are not current canonical data. No source re-extraction was justified or performed.

The active import/export views had lost canonical unknown imports: 30 import CSV zeros and 33 reconstructed XLSX zeros now preserve unknowns as empty cells. Exact original files were archived by SHA-256 before guarded writes. A complete workbook comparison proves that only these 33 XLSX values changed; all other cell values, dimensions and styles are preserved. The CSV comparison proves that only its 30 import cells changed. Prey sums are retained as prey-only figures; full source diet-plus-import sums cannot be established while imports are unknown.

The selected canonical, explicit computational input, complete saved loaded state, historical engine snapshots and actual constructor settings retain verified hashes. Saved loaded-source and runtime diet matrices are byte-identical. Runtime uses the established zero-import default, DC tolerance 0.0021 and normalize_DC=False. The audit applies absolute tolerance 1e-12 and relative tolerance zero; runtime is unchanged, so no diagnostic/coefficient/calculation/map/Word refresh is required. Runtime normalization remains permitted if requested separately.

The existing researcher Word report was read only. It describes the accepted source sums and computational defaults; manual researcher name/date and adjustment fields remain placeholders. This audit grants no signoff, source completeness or whole-region validity. Unknown detritus fate, accepted defaults, scientific parameters and all frozen history remain unchanged. The unselected 2018 endpoint remains outside this bounded audit.

`root_receipt.json` is the machine-readable handoff. `export_repairs.json` supersedes the preparatory `export_repairs_pending.json` without changing its historical preparation record.
''',encoding='utf-8')
print(json.dumps({'outcome':'PASS','canonical_sha256':before,'source_cells_checked':990,'export_import_cells_corrected':len(changed),'runtime_changed':False}))
