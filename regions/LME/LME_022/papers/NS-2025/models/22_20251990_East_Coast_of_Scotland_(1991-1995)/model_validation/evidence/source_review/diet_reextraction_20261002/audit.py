"""Bounded read-only selected LME022 diet normalization audit; source artifacts remain untouched."""
from pathlib import Path
from decimal import Decimal
import csv, hashlib, json, shutil, zipfile, xml.etree.ElementTree as ET
import openpyxl

ROOT = next(p for p in Path(__file__).resolve().parents if (p/'Project.xlsx').exists())
REG = ROOT/'regions/LME_022'
OUT = Path(__file__).resolve().parent
MID = '22_20251990_East_Coast_of_Scotland_(1991-1995)'
MOD = REG/'models'/MID
CAN = MOD/'model.json'
SRC = next((REG/'papers/NS-2025').glob('*.xlsx'))
RUN = REG/'models/regional_ge_integration_20260928'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
rel = lambda p: p.relative_to(ROOT).as_posix()
def dump(name, value):
    (OUT/name).write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding='utf-8')

protected = {rel(p):sha(p) for p in REG.rglob('*') if p.is_file() and OUT not in p.parents}
dump('protected_files_before.json', protected)
inputs = [CAN, MOD/'extracted_tables/model.json', MOD/'extracted_tables/Diet_composition.csv',
          MOD/'extracted_tables'/f'{MID}.json', MOD/'converter_normalized_intermediate.json',
          MOD/'source_diet_restoration.json', MOD/'source_provenance.json', SRC,
          MOD.parent/'extraction_review_20260928/validate_and_stage.py',
          MOD.parent/'extraction_review_20260928/extract_candidates.py',
          RUN/'input'/f'{MID}.json', RUN/'computational_state.json', RUN/'runtime_verification.json',
          REG/'LME_022.xlsx', next(REG.glob('Model_validation*.docx'))]
identities=[]
for p in inputs:
    h=sha(p); a=OUT/'archive_by_sha256'/h/p.name
    a.parent.mkdir(parents=True,exist_ok=True)
    if not a.exists(): shutil.copy2(p,a)
    assert sha(a)==h
    identities.append({'path':rel(p),'sha256':h,'archive':rel(a)})
dump('input_identities.json',identities)

b=openpyxl.load_workbook(REG/'LME_022.xlsx',read_only=True,data_only=True)
overview={r[0]:r[1] for r in b['Overview'].iter_rows(values_only=True) if len(r)>1 and isinstance(r[0],str) and r[0] not in ('@table','field')}
b.close()
assert overview['selected_model_id']==MID and REG/overview['model_path']==CAN
dump('overview_read_only.json',overview)
canonical=json.loads(CAN.read_text(encoding='utf-8'))
extracted=json.loads((MOD/'extracted_tables/model.json').read_text(encoding='utf-8'))
intermediate=json.loads((MOD/'converter_normalized_intermediate.json').read_text(encoding='utf-8'))
groups={g['group_seq']:g for g in canonical['group']}
diet=lambda j:{g['group_seq']:{x['prey_seq']:x['proportion'] for x in (g.get('diet_descr') or {}).get('diet',[])} for g in j['group']}
cd=diet(canonical); nd=diet(intermediate)
workbook=openpyxl.load_workbook(SRC,data_only=False)
sheet=workbook['Table S3']
assert [sheet.cell(2,c).value for c in range(3,26)]==list(range(2,25))
assert [sheet.cell(r,1).value for r in range(3,28)]==list(range(1,26))
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(SRC) as z:
    xml=ET.fromstring(z.read('xl/worksheets/sheet3.xml'))
    xml_cells={c.attrib['r']:c for c in xml.findall('.//m:c',ns)}
    stored={k:c.find('m:v',ns).text for k,c in xml_cells.items() if c.find('m:v',ns) is not None}
    doc=ET.fromstring(zipfile.ZipFile(next(REG.glob('Model_validation*.docx'))).read('word/document.xml'))
    wns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs=[''.join(t.text or '' for t in p.findall('.//w:t',wns)) for p in doc.findall('.//w:p',wns)]
dump('researcher_document_read_only_diet_mentions.json',[p for p in paragraphs if any(s in p.lower() for s in ['diet','normaliz','correction'])])

typed=json.loads((RUN/'computational_state.json').read_text(encoding='utf-8'))
attrs=dict(typed['items']); dc=attrs['_DC']
runtime={(str(c),str(p)):v for c,row in zip(dc['index'],dc['values']) for p,v in zip(dc['columns'],row)}
ledger=[]; mismatches=[]; native_count=0; normalized_different=[]
for r in range(3,28):
    prey=str(sheet.cell(r,1).value)
    for c in range(3,26):
        consumer=str(sheet.cell(2,c).value); cell=sheet.cell(r,c); raw=cell.value
        assert cell.data_type!='f', (cell.coordinate,'unexpected formula')
        adopted=cd[consumer].get(prey)
        ext=extracted['diet'][consumer].get(prey)
        known=raw is not None
        if known:
            native_count+=1
            if adopted is None or float(adopted)!=raw or float(ext)!=raw:
                mismatches.append([cell.coordinate,consumer,prey,raw,adopted,ext])
            if float(nd[consumer].get(prey,0))!=raw:
                normalized_different.append({'cell':cell.coordinate,'consumer':consumer,'prey':prey,'source':str(raw),'normalized_intermediate':nd[consumer].get(prey)})
        elif adopted is not None or ext is not None:
            mismatches.append([cell.coordinate,consumer,prey,'blank',adopted,ext])
        ledger.append({'consumer_id':consumer,'consumer_name':groups[consumer]['group_name'],'prey_id':prey,'prey_name':groups[prey]['group_name'],
                       'source_path':rel(SRC),'source_sha256':sha(SRC),'source_table':'Table S3','source_coordinate':cell.coordinate,
                       'source_status':'native_numeric' if known else 'native_blank','source_xml_numeric_literal':stored.get(cell.coordinate) if known else None,
                       'source_numeric_literal':str(raw) if known else None,'source_number_format':cell.number_format,
                       'canonical_literal':adopted,'extracted_literal':str(ext) if ext is not None else None,
                       'canonical_json_pointer':f'/group/{int(consumer)-1}/diet_descr/diet (prey_seq={prey})',
                       'exact_native_binary64_equal':float(adopted)==raw if known and adopted is not None else adopted is None,
                       'runtime_value':runtime[(consumer,prey)],'runtime_minus_source':runtime[(consumer,prey)]-float(raw or 0),
                       'accepted_corrected_literal':None,'correction_evidence':'No diet override identified; unchanged canonical matches every native numeric value.',
                       'review_status':'source equality independently verified; researcher diet review not newly established'})
for consumer,g in groups.items():
    ledger.append({'consumer_id':consumer,'consumer_name':g['group_name'],'prey_id':'import','source_path':rel(SRC),'source_sha256':sha(SRC),
                   'source_status':'not reported','source_coordinate':None,'source_numeric_literal':None,'canonical_literal':g['diet_imp'],
                   'canonical_json_pointer':f'/group/{int(consumer)-1}/diet_imp','runtime_value':runtime[(consumer,'26')],
                   'accepted_corrected_literal':None,'correction_evidence':'Converter default zero; source does not report diet import. No restoration authorized from missing source.',
                   'review_status':'source import absence retained; runtime accepted historical transformation; no new review'})
assert native_count==232 and not mismatches,mismatches
assert json.loads((MOD/'extracted_tables'/f'{MID}.json').read_text(encoding='utf-8'))==canonical
with (MOD/'extracted_tables/Diet_composition.csv').open(encoding='utf-8',newline='') as f:
    rows=list(csv.reader(f)); head=rows[0]
for row in rows[1:26]:
    prey=row[0]
    for c,consumer in enumerate(head[2:],2):
        assert (float(row[c]) if row[c] else None)==(float(cd[consumer][prey]) if prey in cd[consumer] else None)
dump('cell_ledger.json',ledger)
dump('normalization_proof.json',{'selected_is_normalized':False,'nonempty_native_cells_verified':native_count,'blank_cells_verified':25*23-native_count,
      'canonical_source_mismatches':mismatches,'historical_normalized_intermediate_different_cells':normalized_different,
      'evidence':'Current canonical and extraction/import retain exact source numeric values. Retained validate_and_stage.py explicitly restored source cells after the legacy converter; normalized intermediate differs at listed cells.',
      'native_precision_note':'XML stored numeric literals and number formats are retained in ledger. Native Excel binary64 values match exactly; XML binary64 decimal tails are not additional scientific precision. No equality classification relies on sum=1.'})
sums=[]
for consumer,g in groups.items():
    total=sum((Decimal(v) for v in cd[consumer].values()),Decimal(0))+Decimal(g['diet_imp'])
    rs=sum(runtime[(consumer,str(p))] for p in dc['columns'])
    native_cells=[x for x in ledger if x.get('consumer_id')==consumer and x.get('source_status')=='native_numeric']
    divisor=sum(map(float,cd[consumer].values()))+float(g['diet_imp'])
    ratios=[x['runtime_value']/float(x['source_numeric_literal']) for x in native_cells if float(x['source_numeric_literal'])!=0]
    sums.append({'consumer_id':consumer,'consumer_name':g['group_name'],'canonical_diet_plus_import_decimal_sum':str(total),
                 'native_numeric_sum_decimal_from_loaded_literals':str(sum((Decimal(x['source_numeric_literal']) for x in native_cells),Decimal(0))),
                 'source_import':'not reported','source_blank_cell_count':sum(1 for x in ledger if x.get('consumer_id')==consumer and x.get('source_status')=='native_blank'),
                 'retained_runtime_sum':rs,'runtime_divisor_canonical_float_sum':divisor,
                 'nominal_runtime_normalization_factor':1/divisor if divisor>0 else None,
                 'actual_retained_cell_ratio_min':min(ratios) if ratios else None,'actual_retained_cell_ratio_max':max(ratios) if ratios else None,
                 'factor_provenance':'Retained runtime cell ratios and source floating sum; exact cell values in cell_ledger.json remain authoritative. No new runtime construction or canonical write.',
                 'review_status':'independent native comparison completed; no new researcher review/approval'})
dump('raw_runtime_sums.json',sums)
rv=json.loads((RUN/'runtime_verification.json').read_text(encoding='utf-8'))
assert sha(CAN)==rv['canonical_sha256']==sha(RUN/'input'/f'{MID}.json')==overview['results_model_sha256']
assert sha(RUN/'computational_state.json')==rv['computational_state_sha256']
dump('runtime_comparison.json',{'canonical_before_sha256':sha(CAN),'canonical_after_sha256':sha(CAN),
      'retained_computation_input_sha256':sha(RUN/'input'/f'{MID}.json'),'retained_computational_state_sha256':sha(RUN/'computational_state.json'),
      'constructor':rv['constructor'],'diagnostics_config':rv['diagnostics_config'],'equality_atol':1e-12,'equality_rtol':0,
      'canonical_changed':False,'runtime_differs':False,'actual_old_new_load_required':False,
      'comparison_basis':'Selected canonical is byte-identical to retained actual computation input; constructor/configuration and exact persisted runtime are retained. No canonical mutation and no proposed runtime/settings change. Source-versus-retained-runtime cells/sums are saved separately.',
      'scientific_refresh_required':False,'required_refresh_scope':[],'retained_results':'Preserve all saved diagnostics, coefficients and calculations under actual historical identities.'})
after={rel(p):sha(p) for p in REG.rglob('*') if p.is_file() and OUT not in p.parents}
assert after==protected,{'changed':[(p,h,after.get(p)) for p,h in protected.items() if after.get(p)!=h]}
dump('protected_files_verification.json',{'all_existing_regional_files_unchanged':True,'verified_file_count':len(protected),'identities':after})
dump('root_receipt.json',{'region':'LME_022','selected_model_id':MID,'selected_path':rel(CAN),'outcome':'verified unnormalized unchanged',
      'normalization_status':'Current selected canonical/source extraction/import values are unnormalized and exactly equal to 232 native supplementary numeric cells; 343 native blanks preserved. Historical normalized intermediate is separately retained.',
      'reextraction_performed':False,'canonical_before_sha256':sha(CAN),'canonical_after_sha256':sha(CAN),'changed_diet_cells':[],
      'evidence_paths':[rel(p) for p in OUT.glob('*.json')],
      'precise_blockers':[],'source_fidelity_limits':['Native source supplement fidelity established. Native author EwE database is not retained; paper-rendered and earlier upstream fidelity not expanded in this bounded audit.','Diet import absent from source; converter zero convention retained.','Independent equality verification is evidence, not researcher review registration.'],
      'protected_files_unchanged':True,'runtime_differs':False,'required_refresh_scope':[],'review_status':'pending/not newly established'})
print(json.dumps({'outcome':'verified unnormalized unchanged','source_cells':native_count,'historical_normalized_different_cells':len(normalized_different),'canonical_sha256':sha(CAN),'runtime_differs':False,'protected_files':len(protected)},indent=2))
