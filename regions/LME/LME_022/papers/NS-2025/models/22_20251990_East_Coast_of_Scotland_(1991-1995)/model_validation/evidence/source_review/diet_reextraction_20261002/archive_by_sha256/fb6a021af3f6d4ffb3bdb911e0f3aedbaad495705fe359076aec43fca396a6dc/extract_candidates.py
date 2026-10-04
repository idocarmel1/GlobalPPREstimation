"""Source-faithful candidate extraction; does not change regional selection."""
from pathlib import Path
import sys,json,re,subprocess,shutil,unicodedata,hashlib,csv
import openpyxl
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'Project.xlsx').exists())
REG=ROOT/'regions/LME_022'; OUT=Path(__file__).parent
SKILL=Path('C:/Users/idoca/.agents/skills/ecopath-extraction')
sys.path.insert(0,str(SKILL/'scripts'))
from pdfgrid import get_words
def norm(s):return unicodedata.normalize('NFKC',s).strip()
def dump(p,d):p.write_text(json.dumps(d,indent=2,ensure_ascii=False),encoding='utf-8')
def csvout(p,head,rr):
    with p.open('w',encoding='utf-8',newline='') as f:
        w=csv.writer(f);w.writerow(head);w.writerows(rr)
def save_model(mid,model,tax,provenance):
    dest=REG/'models'/mid; dest.mkdir(exist_ok=True)
    et=dest/'extracted_tables';et.mkdir(exist_ok=True)
    dump(et/'model.json',model)
    subprocess.run([sys.executable,'-X','utf8',str(SKILL/'scripts/write_outputs.py'),str(et/'model.json'),'--outdir',str(dest),'--dir-name','extracted_tables'],check=True)
    tw=openpyxl.Workbook();ws=tw.active;ws.title='Taxonomy';ws.append(['seq','group_name','taxon_descr'])
    for g in model['groups']:ws.append([g['n'],g['name'],tax[g['n']]])
    tw.save(et/'Taxonomy.xlsx')
    csvout(et/'taxonomy.csv',['seq','group_name','taxon_descr'],[[g['n'],g['name'],tax[g['n']]] for g in model['groups']])
    dump(dest/'source_provenance.json',provenance)
    return dest
def saygu():
    src=REG/'papers/NS-2025';pdf=src/'pdf-aa2228c7.pdf'
    lines=get_words(str(pdf),6,merge_gap=.4)
    anchors=[c.xc for c in next(l for l in lines if l.text.startswith('Both ')).cells]
    fields=['tl','B_1890','B_1990','PB_1890','PB_1990','QB_1890','QB_1990','EE_1890','EE_1990','PQ_1890','PQ_1990','E_1890','E_1990']
    rr=[]
    for l in lines:
        if not l.cells or not re.fullmatch(r'\d+',l.cells[0].text) or l.cells[0].x0>110:continue
        n=int(l.cells[0].text)
        if not 1<=n<=25:continue
        r={'n':n};nn=[]
        for c in l.cells[1:]:
            if c.xc<anchors[0]-20:nn.append(c.text)
            elif c.xc<anchors[-1]+20 and re.fullmatch(r'\d+(?:\.\d+)?',c.text):
                k=fields[min(range(13),key=lambda i:abs(c.xc-anchors[i]))];assert k not in r;r[k]=c.text
        r['name']=norm(' '.join(nn))
        if nn:rr.append(r)
    rr.sort(key=lambda r:r['n']);assert [r['n'] for r in rr]==list(range(1,26));dump(OUT/'saygu_table1_cells.json',rr)
    wb=openpyxl.load_workbook(next(src.glob('*.xlsx')),data_only=False)
    tax={n:(d.strip() if d and d.strip()!='-' else 'not documented — Table S2 gives no species inventory for this basal group.')+' [Source: supplementary Table S2, Main Species column; selected main species, not an exhaustive taxonomic inventory.]' for n,name,d in list(wb['Table S2'].values)[2:27]}
    consumers=list(range(2,25));diet={str(n):{} for n in consumers};ws=wb['Table S3']
    assert [ws.cell(2,c).value for c in range(3,26)]==consumers
    for row in range(3,28):
        assert ws.cell(row,1).value==row-2
        for c,pred in enumerate(consumers,3):
            v=ws.cell(row,c).value
            if v is not None:assert isinstance(v,(int,float));diet[str(pred)][str(row-2)]=v
    for suffix,period,cols in [('1890','1890-1895',[3,4,5]),('1990','1991-1995',[7,8,9])]:
        groups=[]
        for r in rr:
            g={k:r[k] for k in ['n','name','tl']}
            for field,out in [('B','biomass'),('PB','pb'),('QB','qb'),('EE','ee'),('PQ','pq')]:
                if field+'_'+suffix in r:g[out]=r[field+'_'+suffix]
            groups.append(g)
        fleets=['Herring fisheries','Whitefish fisheries','Creels'];landings={}
        for row in range(4,29):landings[str(row-3)]={f:wb['Table S1'].cell(row,c).value for f,c in zip(fleets,cols) if wb['Table S1'].cell(row,c).value is not None}
        mid=f'22_2025{suffix}_East_Coast_of_Scotland_({period})'
        model={'metadata':{'LME':'22','model_number':int('2025'+suffix),'model_name':'East Coast of Scotland','model_year':period},'groups':groups,'consumers':consumers,'fleets':fleets,'landings':landings,'discards':{},'detritus_groups':['Detritus'],'detritus_fate':{},'diet':diet,'diet_rows':25,'landings_rows':25,'discards_rows':25}
        save_model(mid,model,tax,{'paper_id':'NS-2025','model_id':mid,'period':period,'local_id_note':'Extraction-local identifier; not EcoBase accession','basic_input':'PDF p6 Table 1; paired period columns. E is exploitation, not EE or migration.','diet':'Workbook Table S3 C3:Y27, same published diet for both periods','landings':f'Workbook Table S1 columns {cols}, rows 4:28; fleet values retained, rounded independent total not substituted','source_files':[str(p.relative_to(ROOT)) for p in src.iterdir() if p.suffix in ['.pdf','.xlsx','.docx']],'unreported':['GS','numeric BA','habitat fraction','detritus fate','detritus import','diet import'],'discards':'Explicitly excluded by authors (paper p14); blank source field does not assert observed zero','taxonomy':'Table S2 A3:C27 main-species descriptions; source spelling retained','source_preservation':'No diet normalization, no suggested balancing BA inserted'})
    return rr
def mackinson():
    src=REG/'papers/LME022-Mackinson-2007';pdf=src/'tech142-ead77c0e.pdf'
    all_pages={}
    def page(n):
        if n not in all_pages:
            all_pages[n]=get_words(str(pdf),n,merge_gap=.4)
            for l in all_pages[n]:l.cells=[c for c in l.cells if c.x0>=50]
        return all_pages[n]
    groups=[];raw=[]
    for p in [29,30]:
        lines=page(p);shift=0 if p==29 else 14.18
        anchors=[186.46,225.42,297.42,351.42,396.43,441.43,486.43]
        # The continuation page expands the column widths; anchor directly on group 50.
        if p==30:
            a=next(l for l in lines if l.cells and l.cells[0].text=='50');anchors=[c.x0 for c in a.cells if re.fullmatch(r'\d+(\.\d+)?',c.text)][1:]
        starts=[]
        for l in lines:
            cs=l.cells
            if cs and re.fullmatch(r'\d+',cs[0].text) and cs[0].x0<90 and 1<=int(cs[0].text)<=68 and l.y>120:starts.append((int(cs[0].text),cs[0].yc,l))
        starts.sort(key=lambda v:v[0])
        for ix,(n,y,l) in enumerate(starts):
            nexty=starts[ix+1][1] if ix+1<len(starts) else y+18
            cells=[c for li in lines for c in li.cells if y-2<=c.yc<nexty-2 and c.x0>l.cells[0].x1+1 and c.x0<anchors[0]-2]
            name=norm(' '.join(c.text for c in sorted(cells,key=lambda c:(round(c.yc,1),c.x0))))
            fields=['tl','biomass','pb','qb','ee','pq','unassim'];g={'n':n,'name':name.replace(',',';')};rr={'n':n,'printed_name':name,'pdf_page':p}
            for j,f in enumerate(fields):
                vals=[c.text for c in l.cells if anchors[j]-1<=c.x0<(anchors[j+1]-1 if j<6 else anchors[j]+42)]
                val=' '.join(vals);rr[f]=val
                m=re.match(r'^(\d+(?:\.\d+)?)',val)
                if m:g[f]=m.group(1)
            if n==17:
                # The source text layer glues the TL to the closing parenthesis.
                assert g['name'].endswith(')4.06')
                g['name']=g['name'][:-4];g['tl']='4.06';rr['tl']='4.06';rr['printed_name']=name[:-4]
            g['name']=g['name'].replace('swarm- ing','swarming')
            raw.append(rr);groups.append(g)
    assert [g['n'] for g in groups]==list(range(1,69)),[g['n'] for g in groups]
    dump(OUT/'mackinson_table3_3_cells.json',raw)
    # Matrix has six predator blocks, each continued onto a second prey page.
    diet={str(n):{} for n in range(1,65)};evidence=[]
    for p in range(31,43):
        ls=page(p);head=next(l for l in ls if l.text.startswith('Predator'))
        aa=[c for c in head.cells if re.fullmatch(r'\d+',c.text)];preds=[int(c.text) for c in aa]
        assert preds==list(range(1+(p-31)//2*12,min(69,13+(p-31)//2*12))),(p,preds)
        for l in ls:
            cs=l.cells
            if cs and cs[0].text=='Import':
                for c in cs[1:]:
                    if re.fullmatch(r'\d+\.\d+',c.text):
                        pred=preds[min(range(len(aa)),key=lambda i:abs(c.x0-aa[i].x0))]
                        diet[str(pred)]['import']=c.text;evidence.append([p,'import',pred,c.text,c.x0,c.y0,c.x1,c.y1])
                continue
            if not cs or not re.fullmatch(r'\d+',cs[0].text) or cs[0].x0>100 or l.y<=head.y:continue
            n=int(cs[0].text)
            if not 1<=n<=68:continue
            for c in cs[1:]:
                if re.fullmatch(r'\d+\.\d+',c.text) and c.x0>=aa[0].x0-3:
                    j=min(range(len(aa)),key=lambda i:abs(c.x0-aa[i].x0));pred=preds[j]
                    assert pred<=64,(p,n,pred,c.text)
                    assert str(n) not in diet[str(pred)],(p,n,pred)
                    diet[str(pred)][str(n)]=c.text;evidence.append([p,n,pred,c.text,c.x0,c.y0,c.x1,c.y1])
    dump(OUT/'mackinson_diet_cells.json',evidence)
    print('DIET SUMS',[(n,round(sum(float(v) for v in d.values()),5)) for n,d in diet.items() if abs(sum(float(v) for v in d.values())-1)>.01])
    # Summary Table 3.5 provides named eight-fleet landings and total discards.
    fleets=['Demersal trawl & seine','Beam trawl','Sandeel trawl','Pelagic trawl','Drift and fixed nets','Nephrops trawl','Hooks','Other','Discards total (fleet unspecified in Table 3.5)']
    landing={};discard={};fishraw=[]
    for p in [43,44]:
        ls=page(p);head=next(l for l in ls if l.text.startswith('Functional'))
        row=next(l for l in ls if l.cells and l.cells[0].text==('6' if p==43 else '45'))
        aa=[c.x0 for c in row.cells if re.fullmatch(r'\d[\d,]*|\-',c.text)][1:]
        assert len(aa)==10,(p,aa)
        starts=[(int(l.cells[0].text),l.cells[0].yc,l) for l in ls if l.cells and re.fullmatch(r'\d+',l.cells[0].text) and l.cells[0].x0<100 and l.y>head.y and 1<=int(l.cells[0].text)<=61]
        for i,(n,y,l) in enumerate(starts):
            ny=starts[i+1][1] if i+1<len(starts) else y+30
            vals={}
            for li in ls:
                if 'TOTAL' in li.text:continue
                for c in li.cells:
                    if y-2<=c.yc<ny-2 and c.x0>=aa[0]-3 and re.fullmatch(r'\d[\d,]*|\-',c.text):
                        j=min(range(10),key=lambda k:abs(c.x0-aa[k]));assert j not in vals,(p,n,j);vals[j]=c.text
            fishraw.append({'pdf_page':p,'n':n,'values':vals})
            landing[str(n)]={fleets[j]:str(float(v.replace(',',''))/570000) for j,v in vals.items() if j<8 and v!='-'}
            if vals.get(9,'-')!='-':discard[str(n)]={fleets[8]:str(float(vals[9].replace(',',''))/570000)}
    dump(OUT/'mackinson_table3_5_tonnes.json',fishraw)
    names={g['n']:g['name'] for g in groups}
    fate={'65':{names[66]:'.3',names[67]:'.7'},'66':{names[67]:'1'},'67':{'Export':'1'}}
    tax={g['n']:'Source group definition: '+g['name']+'. Detailed composition review retained in taxonomy_evidence; do not treat group name as an exhaustive species list.' for g in groups}
    taxpath=OUT/'mackinson_taxonomy.json'
    if taxpath.exists():tax={int(k):v for k,v in json.loads(taxpath.read_text(encoding='utf-8')).items()}
    model={'metadata':{'LME':'22','model_number':20071991,'model_name':'North Sea report Table 3.3','model_year':1991},'groups':groups,'consumers':list(range(1,65)),'fleets':fleets,'landings':landing,'discards':discard,'detritus_groups':[names[i] for i in [66,67,68]],'detritus_fate':fate,'diet':diet,'diet_rows':68,'landings_rows':68,'discards_rows':68}
    dest=save_model('22_20071991_North_Sea_report_Table_3.3_(1991)',model,tax,{'paper_id':'LME022-Mackinson-2007','basic_input':'Table 3.3 printed pp27–28 / PDF29–30; first number of each cell is balanced parameter; bracketed initial estimates preserved separately','diet':'Table 3.4 printed pp29–40 / PDF31–42; coordinates anchored on predator IDs; no normalization','fishery':'Table 3.5 printed pp41–42 / PDF43–44; eight fleet landings and unallocated total discards, source tonnes divided by reported 570000 km2; dashes remain unknown','fate':'Section 8 printed p98 / PDF100: phytoplankton 30% DOM 70% POM, unutilised DOM to POM then export. Other shares unknown. Obsolete references to sections19/15 do not provide a routing table in actual report.','model_period':'Nominal 1991; report additionally discusses a 1973 Ecosim fitted version without complete static parameter table','source_files':[str(pdf.relative_to(ROOT))],'group_names':'Printed punctuation preserved in audit; commas replaced by semicolons solely to satisfy unquoted EwE CSV format','unreported':['numeric BA','most detritus fate fractions','habitat fraction','diet import'],'known_conflicts':['Gelatinous zooplankton QB .18 and PB 2.858 imply PQ 15.8778 as actually printed; retained, no unapproved repair','Summary Table 3.5 and detailed fishery Table14.7/14.8 differ; preserved separately and not silently mixed'],'source_preference':'User preferred this paper for perceived better spatial fit and detailed documentation; no exact model selected'})
    dump(OUT/'mackinson_pdf_geometry.json',{str(p):[[{'text':c.text,'x0':c.x0,'x1':c.x1,'y0':c.y0,'y1':c.y1} for c in l.cells] for l in ll] for p,ll in all_pages.items()})
    return dest
if __name__=='__main__':
    if 'saygu' in sys.argv:saygu()
    if 'mackinson' in sys.argv:mackinson()

