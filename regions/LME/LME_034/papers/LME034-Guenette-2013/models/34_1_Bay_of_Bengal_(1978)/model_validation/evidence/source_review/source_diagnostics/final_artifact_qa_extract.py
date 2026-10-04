from pathlib import Path
import json, hashlib, sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

project = Path(__file__).resolve().parents[5]
out = Path(__file__).resolve().parent
region = project / 'regions/LME_034'
docpath = region / 'Model_validation_34_1_Bay_of_Bengal_(1978).docx'
bookpath = region / 'LME034_taxon_mapping_appendix.xlsx'
doc = Document(docpath)
blocks = []
for node in doc.element.body:
    if node.tag.endswith('}p'):
        p = Paragraph(node, doc)
        blocks.append({'type':'paragraph', 'text':p.text})
    elif node.tag.endswith('}tbl'):
        t = Table(node,doc)
        blocks.append({'type':'table','rows':[[c.text for c in r.cells] for r in t.rows]})
book = load_workbook(bookpath, data_only=False)
sheets = {}
for s in book:
    sheets[s.title] = {'max_row':s.max_row,'max_column':s.max_column,
        'values':[[c.value for c in r] for r in s],
        'hyperlinks':[{'cell':c.coordinate,'target':c.hyperlink.target,'location':c.hyperlink.location} for r in s for c in r if c.hyperlink],
        'freeze_panes':s.freeze_panes,'auto_filter':s.auto_filter.ref,
        'formula_cells':[{'cell':c.coordinate,'formula':c.value} for r in s for c in r if c.data_type=='f']}
data = {'document_sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),
    'appendix_sha256':hashlib.sha256(bookpath.read_bytes()).hexdigest(),
    'document_blocks':blocks,'sheets':sheets}
(out/'final_artifact_extracted_content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf8')
for b in blocks:
    if b['type']=='paragraph':print(b['text'])
    else:
        for r in b['rows']:print(' | '.join(r))
print('\nWORKBOOK SHEETS')
for name, s in sheets.items():
    print(name,s['max_row'],s['max_column'],s['freeze_panes'],s['auto_filter'])
    if s['max_row']<100:
        for i,r in enumerate(s['values'],1): print(i,r)
    else:
        for i,r in enumerate(s['values'][:8],1):print(i,r)
print('HASHES',data['document_sha256'],data['appendix_sha256'])
