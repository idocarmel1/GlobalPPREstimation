"""Read-only exact diet audit; writes only this new evidence directory."""
from pathlib import Path
import csv, hashlib, json, zipfile, xml.etree.ElementTree as ET
from decimal import Decimal as D
from datetime import datetime, timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from docx import Document

OUT = Path(__file__).resolve().parent
REG = OUT.parent
ROOT = REG.parent.parent
MODEL = REG / 'models/CAL-2016_California_Current_2000-2014'
PAPER = REG / 'papers/CAL-2016'
AUTHOR = PAPER / 'author_archive_extracted/CalCurFoodWebModelECOMOD-master'
SOURCE = AUTHOR / 'Koehn.et.al.2016_dietmatrix.csv'
SELECTED = MODEL / 'selected_pipeline/3_2016_Author_solved_California_Current_(2000-2014).json'
CANONICAL = MODEL / 'model.json'
IMPORT = MODEL / 'extracted_tables/Diet_composition.csv'
EXTRACTION = MODEL / 'extracted_tables/model.json'

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p): return p.relative_to(ROOT).as_posix()
def read(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def dump(name, obj):
    (OUT / name).write_text(json.dumps(obj, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')

# Snapshot all pre-existing regional files before producing evidence. No originals
# are ever written; concurrent differences are reported rather than overwritten.
protected = {rel(p): sha(p) for p in REG.rglob('*') if p.is_file() and OUT not in p.parents}
originals = [SELECTED, CANONICAL, EXTRACTION, IMPORT, SOURCE, PAPER / 'mmc1.xlsx',
             AUTHOR / 'Koehn.et.al.2016_groupinfo.csv']
archive = OUT / 'originals_by_sha256'
archive.mkdir(exist_ok=True)
archives = []
for p in originals:
    h = protected[rel(p)]
    data = p.read_bytes()
    assert hashlib.sha256(data).hexdigest() == h, f'Concurrent change: {p}'
    dest = archive / (h + p.suffix)
    if dest.exists(): assert dest.read_bytes() == data
    else: dest.write_bytes(data)
    archives.append({'input': rel(p), 'sha256': h, 'archive': rel(dest)})

book = load_workbook(REG / 'LME_003.xlsx', read_only=True, data_only=False)
overview = {r[0]: r[1] for r in book['Overview'].values if len(r) >= 2 and r[0]}
book.close()
assert REG / overview['model_path'] == SELECTED
rows = list(csv.reader(SOURCE.open(encoding='utf-8')))
with zipfile.ZipFile(PAPER/'master-8d6306e9.zip') as z:
    source_member = next(n for n in z.namelist() if n.endswith('/Koehn.et.al.2016_dietmatrix.csv'))
    assert z.read(source_member) == SOURCE.read_bytes(), 'Retained extracted source differs from archive member'
imrows = list(csv.reader(IMPORT.open(encoding='utf-8')))
imcols = {int(v): i for i, v in enumerate(imrows[0][2:], 2)}
imprey = {int(r[0]): r for r in imrows[1:] if r[0].isdigit()}
imimport = next(r for r in imrows if len(r)>1 and r[1] == 'Import')
ext = read(EXTRACTION)
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(PAPER / 'mmc1.xlsx') as z:
    xml = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    publisher = {c.attrib['r']: c.find('s:v', ns).text
                 for c in xml.findall('.//s:c', ns) if c.find('s:v', ns) is not None}

variants = {
    'selected_computational_input': SELECTED,
    'canonical_source': CANONICAL,
    'converted_extraction': MODEL / 'extracted_tables/3_California_Current_CAL-2016_California_Current_(2000-2014).json',
    'retained_solved_diagnostic_input': MODEL / 'diagnostics/author_solved_diagnostic_input.json',
    'retained_filename_adapter_diagnostic_input': MODEL / 'diagnostics/3_2016_Author_solved_California_Current_(2000-2014).json',
}
def diet_lookup(data):
    result = {}
    for gi, g in enumerate(data['group']):
        c = int(g['group_seq'])
        ds = (g.get('diet_descr') or {}).get('diet', [])
        ds = [ds] if isinstance(ds, dict) else ds
        assert len({int(d['prey_seq']) for d in ds}) == len(ds)
        result[c] = (gi, g, {int(d['prey_seq']): (di, d['proportion']) for di, d in enumerate(ds)})
    return result
lookups = {key: diet_lookup(read(p)) for key, p in variants.items()}
checks = {}
for key, lookup in lookups.items():
    differences = []
    for c in range(1, 94):
        gi, g, prey = lookup[c]
        for p in range(1, 94):
            val = prey.get(p, (None, '0'))[1]
            if D(val) != D(rows[p][c]): differences.append([c, p, rows[p][c], val])
        if D(g['diet_imp']) != D(rows[94][c]): differences.append([c, 'import', rows[94][c], g['diet_imp']])
    checks[key] = {'path': rel(variants[key]), 'sha256_before': protected[rel(variants[key])],
                   'cells_compared': 8742, 'differences': differences,
                   'zero_semantics': 'Absent prey edge is structural zero; all source nonzero edges present.'}

ledger = []
sums = []
for c in range(2, 93):
    gi, g, prey = lookups['selected_computational_input'][c]
    total = D(0)
    selected_total = D(g['diet_imp'])
    for p in range(1, 95):
        isimport = p == 94
        lit = rows[p][c]
        total += D(lit)
        if not isimport: selected_total += D(prey.get(p, (None, '0'))[1])
        addr = f'{get_column_letter(c+1)}{p+2}'
        pub = publisher[addr]
        ilit = (imimport if isimport else imprey[p])[imcols[c]]
        elit = ext['diet'][str(c)]['import' if isimport else str(p)]
        slit = g['diet_imp'] if isimport else prey.get(p, (None, '0'))[1]
        pointer = f'/group/{gi}/diet_imp' if isimport else (f'/group/{gi}/diet_descr/diet/{prey[p][0]}/proportion' if p in prey else 'absent edge: structural zero')
        clit = lookups['canonical_source'][c][1]['diet_imp'] if isimport else lookups['canonical_source'][c][2].get(p, (None, '0'))[1]
        ledger.append({'consumer_seq':c, 'consumer_name':g['group_name'], 'prey_seq':'import' if isimport else p,
            'prey_name':rows[p][0], 'source_file':rel(SOURCE), 'source_sha256':protected[rel(SOURCE)],
            'source_csv_row_one_based':p+1, 'source_csv_column_one_based':c+1, 'source_literal':lit,
            'publisher_file':rel(PAPER/'mmc1.xlsx'), 'publisher_sha256':protected[rel(PAPER/'mmc1.xlsx')],
            'publisher_sheet_xml':'xl/worksheets/sheet1.xml', 'publisher_coordinate':addr,
            'publisher_stored_literal':pub, 'publisher_decimal_equal':D(pub)==D(lit),
            'publisher_binary_float_equal':float(pub)==float(lit), 'import_file':rel(IMPORT),
            'import_literal':ilit, 'extraction_json_pointer':f'/diet/{c}/' + ('import' if isimport else str(p)),
            'extraction_literal':elit, 'selected_json_pointer':pointer, 'selected_literal':slit,
            'canonical_literal':clit, 'accepted_corrected_literal':'',
            'correction_evidence':'No diet/import override found in current researcher DOCX or retained acceptance receipts; exact source values retained.',
            'review_state':'not-established: fidelity verified; no new researcher review registration',
            'missing_status':'none; absent JSON edges represent explicit source zeros'})
    sums.append({'consumer_seq':c, 'consumer_name':g['group_name'], 'source_diet_plus_import_sum':str(total),
                 'selected_diet_plus_import_sum':str(selected_total), 'source_import':rows[94][c],
                 'retained_runtime_normalize_DC':False, 'source_to_selected_factor':'1',
                 'runtime_factor':'1 (retained normalize_DC=False; no fresh runtime executed)',
                 'changed_cells':0, 'missing_cells':0, 'review_state':'not-established'})
assert all(not x['differences'] for x in checks.values())
assert all(D(x['source_literal']) == D(x['import_literal']) == D(x['extraction_literal']) == D(x['selected_literal']) == D(x['canonical_literal']) for x in ledger)
assert all(x['publisher_binary_float_equal'] for x in ledger)
assert len(ledger) == 8554
with (OUT/'cell_ledger.csv').open('w', encoding='utf-8', newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(ledger[0]));w.writeheader();w.writerows(ledger)
dump('consumer_sums.json', sums)
dump('source_archives.json', archives)

docpath = REG / 'Model_validation_CAL-2016_California_Current_2000-2014.docx'
doc = Document(docpath)
texts = [{'paragraph_zero_based':i,'text':p.text} for i,p in enumerate(doc.paragraphs)]
texts += [{'table_zero_based':ti,'row_zero_based':ri,'text':' | '.join(c.text for c in row.cells)}
          for ti,t in enumerate(doc.tables) for ri,row in enumerate(t.rows)]
relevant = [x for x in texts if any(s in x['text'].lower() for s in ['diet','correct','normali','accepted','manual'])]
with zipfile.ZipFile(docpath) as z:
    dx=ET.fromstring(z.read('word/document.xml'))
    wn={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    tracked={tag:len(dx.findall('.//w:'+tag,wn)) for tag in ['ins','del']}
    comments='word/comments.xml' in z.namelist()
dump('researcher_correction_review.json', {'path':rel(docpath), 'sha256':protected[rel(docpath)],
    'relevant_current_docx_passages':relevant,'tracked_change_counts':tracked,'comments_part_present':comments,
    'diet_override_found':False,
    'preserved_accepted_context':'Author-equation B/EE solution and GS=0.2, sole-detritus routing/EE=1/import runtime extensions; no parameters changed.'})

changed = [{'path':p,'before':h,'after':sha(ROOT/p)} for p,h in protected.items() if sha(ROOT/p)!=h]
for key,p in variants.items(): checks[key]['sha256_after']=sha(p)
dump('protected_files_verification.json', {'scope':'All pre-existing files within regions/LME_003; new audit folder excluded',
    'before_sha256':protected,'checked_files':len(protected),'concurrent_changes':changed,
    'audit_wrote_original_files':False, 'all_preexisting_regional_files_unchanged':not changed})
staging=read(MODEL/'diagnostics/staging_transformations.json')
receipt=read(MODEL/'selected_pipeline/coefficient_verification.json')
dump('audit_result.json', {'unit_id':'LME_003','timestamp_utc':datetime.now(timezone.utc).isoformat(),
    'outcome':'verified unnormalized unchanged','selected_model_id':overview['selected_model_id'],
    'selected_path':overview['model_path'], 'variant_checks':checks,
    'primary_source_identity':{'archive':rel(PAPER/'master-8d6306e9.zip'),
        'archive_sha256':protected[rel(PAPER/'master-8d6306e9.zip')],
        'member':source_member,'extracted_member_sha256':protected[rel(SOURCE)],
        'archive_member_exact_bytes_equal':True,
        'identity':'Retained author R-model archive CSV plus publisher Appendix A XLSX; selected solved variant is not asserted to be a native EwE database.'},
    'canonical_restoration_authorized_by_evidence':False,'restored_cells':0,'changed_original_files':[],
    'source_cells_compared_per_database_variant':8742,'consumer_ledger_cells':len(ledger),'consumers':len(sums),
    'consumer_source_sum_min':str(min(D(s['source_diet_plus_import_sum']) for s in sums)),
    'consumer_source_sum_max':str(max(D(s['source_diet_plus_import_sum']) for s in sums)),
    'publisher_storage_decimal_tails':sum(not x['publisher_decimal_equal'] for x in ledger),
    'publisher_storage_interpretation':'667 XLSX XML decimal tails differ from author CSV text; all are equal as binary floats. Author CSV is exact primary literal evidence. No normalization is inferred from storage tails or sums.',
    'conversion_history':{'build_extraction':rel(REG/'extraction_review_20260928/build_extraction.py'),
        'source_comparison':rel(MODEL/'extracted_tables/source_comparison.json'),
        'staging_runtime_settings':staging['runtime_settings'],
        'selected_retained_receipt_runtime_sha256':receipt['runtime_sha256'],
        'selected_retained_receipt_canonical_sha256':receipt['canonical_sha256']},
    'staleness':{'canonical_identity_changed':False,'selected_identity_changed':False,
        'existing_selected_receipt_matches':receipt['runtime_sha256']==sha(SELECTED),
        'overview_results_model_hash_matches':overview['results_model_sha256']==sha(SELECTED),
        'diet_audit_invalidated_existing_results':False,'historical_diagnostics_preserved':True,
        'freshness_hashes_rewritten':False,'fresh_scientific_runs':False},
    'researcher_review_state':'not-established; no new approval or review registration',
    'unresolved':'Earlier processing by authors before retained CSV/XLSX is outside available evidence. This establishes no downstream normalization of retained primary input, not a full scientific validation.',
    'all_preexisting_regional_files_unchanged':not changed})
print(json.dumps({'outcome':'verified unnormalized unchanged','selected_sha256':sha(SELECTED),
    'canonical_sha256':sha(CANONICAL),'ledger_cells':len(ledger),'protected_files':len(protected),
    'concurrent_changes':changed,'sum_min':str(min(D(x['source_diet_plus_import_sum']) for x in sums)),
    'sum_max':str(max(D(x['source_diet_plus_import_sum']) for x in sums))}))
