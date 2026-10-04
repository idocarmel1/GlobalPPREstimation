"""Read-only reconciliation of delivered Word tables to frozen source ledgers."""
from pathlib import Path
import hashlib
import json
import math
from docx import Document

E = Path(__file__).resolve().parent
R = E.parents[3]
REGION = E.parents[1]
A = REGION / "models/52_GM2019_Fig9_Pelagic_(2000-2014)/assumption_variants/adopted_balanced_20261003"
I = A.parents[1] / "research_20261003/integration"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def hash_file(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def equivalent(display, source, rel_tol=6e-8):
    if source is None:
        return display == "absent"
    if isinstance(source, bool) or str(source) in ("true", "false"):
        return display == str(source).lower()
    if str(source) == "-9999":
        return display == "unknown"
    try:
        a, b = float(display), float(source)
        return math.isfinite(a) and math.isclose(a, b, rel_tol=rel_tol, abs_tol=1e-12)
    except (ValueError, TypeError):
        return display == str(source)


def main():
    spec = read(E / "report_inputs.json")
    ledger = read(A / "source_to_final_ledger.json")
    model_hash = hash_file(A / "model.json")
    assert model_hash == spec["provenance"]["model_sha256"] == ledger["exact_final_sha256"]
    validation = Document(REGION / spec["validation_filename"])
    departures = Document(REGION / spec["departures_filename"])
    tables = {tuple(c.text for c in t.rows[0].cells):t for t in departures.tables}
    scalar = tables[("Group ID and name", "Field", "Baseline value", "Final input")]
    diet = tables[("Native row group", "Field", "Prey or origin ID", "Baseline", "Final")]
    actual = {}
    for table, is_diet in ((scalar,False),(diet,True)):
        for row in table.rows[1:]:
            cells = [c.text for c in row.cells]
            gid = int(cells[0].split()[0])
            key = (gid,cells[1],int(cells[2]) if is_diet else None)
            assert key not in actual
            actual[key] = cells[-2:]
    assert len(actual) == len(ledger["input_cell_changes"]) == 448
    for change in ledger["input_cell_changes"]:
        key = (change["group"], change["field"], change.get("prey"))
        assert equivalent(actual[key][0],change["before"]), (key,"before")
        assert equivalent(actual[key][1],change["after"]), (key,"after")
    food = tables[("Consumer","Prey or external category","Wet food","Wet per C","Carbon food","Source or assumption")]
    components = read(A / "food_component_ledger.json")
    assert len(food.rows)-1 == len(components) == 74
    for row, component in zip(food.rows[1:],components):
        cells = [c.text for c in row.cells]
        assert int(cells[0].split()[0]) == component["consumer"]
        assert int(cells[1].split()[0]) == component["prey"]
        for value,key in zip(cells[2:5],("wet_million_t","wet_per_carbon","carbon_million_t")):
            assert equivalent(value,component[key],6e-6)
    very_low = next(t for t in validation.tables if tuple(c.text for c in t.rows[0].cells) == ("Affected taxa","Why confidence is very low"))
    delivered_taxa = [taxon.strip() for row in very_low.rows[1:] for taxon in row.cells[0].text.split(";")]
    audit = read(I / "taxon_audit.json")
    expected_taxa = [r["taxon"] for r in audit if r["overall_confidence"] == "Very low"]
    assert len(delivered_taxa) == len(set(delivered_taxa)) == len(expected_taxa) == 115
    assert set(delivered_taxa) == set(expected_taxa)
    main_fields = {r.cells[0].text:r.cells[1].text for r in validation.tables[0].rows[1:]}
    assert main_fields["Open issues and next action"] == ""
    assert main_fields["Review and reproducibility"] == "Researcher name: [name] | review date: [dd/mm/yyyy]"
    text = "\n".join(p.text for p in validation.paragraphs)
    assert ("Pending alignment draft." in text) == spec["provenance"]["pending_map_alignment"]
    for method in ("GE","TE","With_Egestion"):
        matrix = read(A / "diagnostics" / (method + "_full_return.json"))
        assert matrix["model_sha256"] == model_hash
        assert len(matrix["SPPR"]["index"]) == 23
    checks = {
        "exact_model_sha256":model_hash,
        "all_448_native_input_changes_present_and_reconciled":True,
        "all_74_food_components_present_and_reconciled":True,
        "all_115_very_low_taxa_once":True,
        "researcher_fields_undecided":True,
        "same_input_all_three_retained_matrices":True,
        "pending_map_alignment":spec["provenance"]["pending_map_alignment"],
        "numeric_display_tolerance":"8 significant digits in changed-cell tables; 6 in compact food and budget tables. Exact values in linked JSON.",
        "report_hashes":{spec[k]:hash_file(REGION/spec[k]) for k in ("validation_filename","departures_filename")}
    }
    (E / "report_content_verification.json").write_text(json.dumps(checks,indent=2),encoding="utf-8")
    print(json.dumps(checks))


if __name__ == "__main__":
    main()
