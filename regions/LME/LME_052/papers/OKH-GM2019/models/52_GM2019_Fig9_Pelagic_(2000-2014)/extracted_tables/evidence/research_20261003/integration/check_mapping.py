"""Independent checks of saved mapping rows and appendix bytes."""
import hashlib,json,math,zipfile
from collections import defaultdict
from pathlib import Path
from urllib.parse import unquote,urlsplit
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[5]
REGION=ROOT/"regions/LME_052"
MODEL_ID="52_GM2019_Fig9_Pelagic_balanced_(2000-2014)"
audit=json.loads((HERE/"taxon_audit.json").read_text(encoding="utf-8"))
matching=json.loads((HERE/"matching_rows.json").read_text(encoding="utf-8"))
summary=json.loads((HERE/"mapping_summary.json").read_text(encoding="utf-8"))
baseline=json.loads((HERE/"baseline_inventory.json").read_text(encoding="utf-8"))
path=REGION/f"{MODEL_ID}_taxon_mapping_appendix.xlsx"
wb=openpyxl.load_workbook(path,data_only=False)
sheet=wb["Taxon mappings"]
rows=[list(r)for r in sheet.iter_rows(min_row=3,values_only=True)]
by_taxon={r[0]:r for r in rows}
catch={r["taxon"]:r for r in baseline["catch"]if r["catch_basis"]=="landings"}
groups={g["group_name"]for g in baseline["candidate_groups"]}
mapping=defaultdict(list)
for r in matching:
    if r["group"]:mapping[r["taxon"]].append(r)
checks={
    "all151unique_catch_labels":len(rows)==len(by_taxon)==len(audit)==151 and set(by_taxon)==set(catch),
    "allmatching_model_ids":all(r["model_id"]==MODEL_ID for r in matching),
    "weights_positive_sum_one":all(all(isinstance(r["weight"],(int,float))and r["weight"]>0 for r in rr)and math.isclose(math.fsum(r["weight"]for r in rr),1,abs_tol=1e-12)for rr in mapping.values()),
    "exact_group_names":all(r["group"]in groups for rr in mapping.values()for r in rr),
    "no_microbe_producer_detritus_catch":all(r["group"]not in {"Phytoplankton","Bacteria","Protozoa","Detritus"}for rr in mapping.values()for r in rr),
    "one_confidence_category_per_taxon":math.fsum(r["taxa"]for r in summary["confidence"])==151,
    "catch_total_reconciles":math.isclose(math.fsum(float(r[2])for r in rows),summary["total_catch_tonnes"],rel_tol=1e-14),
    "simple_ppr_total_reconciles":math.isclose(math.fsum(float(r[3])for r in rows if isinstance(r[3],(float,int))),summary["total_known_simple_ppr_tC"],rel_tol=1e-14),
    "category_shares100":math.isclose(math.fsum(r["catch_percent"]for r in summary["confidence"]),100,abs_tol=1e-10)and math.isclose(math.fsum(r["simple_ppr_percent"]for r in summary["confidence"]),100,abs_tol=1e-10),
    "rule_shares100_sorted":math.isclose(math.fsum(r["simple_ppr_percent"]for r in summary["rule_shares"]),100,abs_tol=1e-10)and all(a["simple_ppr_percent"]>=b["simple_ppr_percent"]for a,b in zip(summary["rule_shares"],summary["rule_shares"][1:])),
    "numeric_unrounded_descending_sort":all(float(a[3])>=float(b[3])for a,b in zip(rows,rows[1:])if isinstance(a[3],(float,int))and isinstance(b[3],(float,int))),
    "numeric_catch_ppr":all(isinstance(r[2],(float,int))and(isinstance(r[3],(float,int))or r[3]=="?")for r in rows),
    "zero_catch_yields_genuine_zero_ppr":all(r[3]==0 for r in rows if r[2]==0),
    "frozen_header_and_filter":sheet.freeze_panes=="C3"and sheet.auto_filter.ref=="A2:G153",
    "wrapped_readable_reason_rows":sheet.column_dimensions["G"].width>=100 and all(sheet.row_dimensions[i].height>=72 and sheet.cell(i,7).alignment.wrap_text for i in range(3,154)),
}
links=[]
for row in wb["Sources"]:
    for c in row:
        if c.hyperlink:
            target=c.hyperlink.target
            parsed=urlsplit(target)
            external=parsed.scheme in {"https","http"}
            resolved=(path.parent/unquote(parsed.path)).resolve()if not external else None
            links.append(dict(target=target,external=external,present=True if external else resolved.is_file(),portable=external or(not parsed.scheme and not target.startswith(("/","\\"))),blue_underlined=c.font.color.type=="rgb"and c.font.color.rgb=="000563C1"and c.font.underline=="single"))
checks["source_links_present_portable_blue_underlined"]=all(r["present"]and r["portable"]and r["blue_underlined"]for r in links)
with zipfile.ZipFile(path)as z:
    checks["no_machine_hyperlink_base"]=not any(b"HyperlinkBase"in z.read(n)for n in z.namelist()if n.endswith(".xml")and n.startswith("docProps/"))
wb.close()
result=dict(schema_version=1,model_id=MODEL_ID,appendix=path.relative_to(ROOT).as_posix(),appendix_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),checks=checks,links=links,passed=all(checks.values()))
(HERE/"mapping_verification.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(result,ensure_ascii=False,indent=2))
if not result["passed"]:raise SystemExit(1)
