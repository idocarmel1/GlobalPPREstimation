"""Adopt exact native model and direct-only returns into LME_052.

No solver call, central workbook edit, or model JSON mutation occurs here.
The existing project selection/calculation/writer functions own the workflow.
"""
import csv,hashlib,json,math,shutil,sys
from pathlib import Path

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[5]
sys.path.insert(0,str(ROOT/"tools"))
from workbooks import ANNUAL_HEADER,digest_tables,finite,input_hash,overview,records,rows,sha,table_dict,validate_region,write_book
from regional import recalculate,result_hash,set_setting
from run_region import prepare_selection
from regional_book_io import read_fast

REGION=ROOT/"regions/LME_052"
WORKBOOK=REGION/"LME_052.xlsx"
MODEL_ID="52_GM2019_Fig9_Pelagic_balanced_(2000-2014)"
MODEL_PATH="models/52_GM2019_Fig9_Pelagic_(2000-2014)/assumption_variants/adopted_balanced_20261003/model.json"
MODEL=REGION/MODEL_PATH
VARIANT=MODEL.parent
FINAL_HASH="9beff5a4e525dcf4253bf83472253c7a4a9f7e7db10379ecef64eb33cd32a656"
METHODS={"GE":"new_GE","TE":"new_TE_EEfix","With_Egestion":"new_WithEgestion"}
SCOPES={"all":{23,22,1},"inner":{22,1},"PP":{1}}

def read(path):return json.loads(path.read_text(encoding="utf-8"))
def save(path,data):path.write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")
def flatten(report):
    # Exact static flatten contract of PPRCalculator._flatten_diagnostics.
    out={}
    for section,body in report.items():
        if not isinstance(body,dict):out[f"n_{section}"if isinstance(body,(list,tuple))else section]=len(body)if isinstance(body,(list,tuple))else body;continue
        for key,value in body.items():
            name=f"{section}_{key}"
            if isinstance(value,dict):out.update({f"{name}_{k}":v for k,v in value.items()})
            elif isinstance(value,(list,tuple)):out[f"{section}_n_{key}"]=len(value)
            else:out[name]=value
    return out

assert sha(MODEL)==FINAL_HASH,"Final native bytes differ from frozen diagnostic identity"
native=read(VARIANT/"native_reload_result.json")
ledger=read(VARIANT/"source_to_final_ledger.json")
assert native["main_is_model_balanced"]is True and native["exact_input_sha256"]==FINAL_HASH
baseline=read(HERE/"baseline_inventory.json")
before_hash=sha(WORKBOOK)
assert before_hash==baseline["workbook_sha256"],"Regional workbook changed since reviewed baseline; reconcile before applying"
book=read_fast(WORKBOOK)
assert input_hash(book)==overview(book)["calculation_input_sha256"],"Baseline inputs are stale"
backup=HERE/"LME_052_before_GM2019_adoption.xlsx"
if backup.exists():assert sha(backup)==before_hash
else:shutil.copy2(WORKBOOK,backup)
protected={(s,n):digest_tables([h,r])for s,n in [("Catch","Catch"),("Classic PPR","Taxa"),("NPP","NPP"),("NPP","Provenance")]for h,r in [book[s][n]]}
old_classic={tuple(r[k]for k in ["scope","method","catch_basis","unidentified","metric"]):[r[y]for y in range(1950,2020)]for r in records(book,"Classic PPR","Annual")if r["metric"]in {"ppr","catch","covered_catch"}}
original_model_hash=sha(REGION/str(overview(book)["model_path"]))
taxonomy={int(r["seq"]):r["taxon_descr"]for r in csv.DictReader((HERE/"taxonomy.csv").open(encoding="utf-8"))}
factors={int(g["group_id"]):g["body_wet_per_carbon"]for g in ledger["group_parameters"]}
factors[23]=1.0
fg=native["final_groups"]
assert len(set(fg["index"]))==len(fg["index"])==23
group_by_id={int(i):dict(zip(fg["columns"],r))for i,r in zip(fg["index"],fg["data"])}
group_header=["seq","group_name","group_type",*[c for c in fg["columns"]if c!="group_name"],"biomass_units","flow_units","body_wet_per_carbon"]
group_rows=[]
for i in fg["index"]:
    d=group_by_id[i].copy()
    d["taxon_descr"]=taxonomy.get(i,"Synthetic external food import group; carbon currency; no catch taxon.")
    typ={"PP":"Primary producer","Detritus":"Detritus","Import":"Import"}.get(d["trophic_info"],"Regular")
    group_rows.append([i,d["group_name"],typ,*[d[c]for c in fg["columns"]if c!="group_name"],"tC/km2","tC/km2/year",factors.get(i)])

native_coeff={}
group_sppr=[];bridge=[];health=[];run_notes=[]
for option,method in METHODS.items():
    data=read(VARIANT/f"diagnostics/{option}_full_return.json")
    assert data["model_sha256"]==FINAL_HASH
    report=data["report"]
    assert report["status"]in {"OK","WARN"},"FAIL needs a separately authorized interpretation before provisional publication"
    assert report["divergence"]["status"]==report["balance"]["status"]=="OK"
    matrix=data["SPPR"]
    assert len(set(matrix["index"]))==len(matrix["index"])==23 and set(matrix["columns"])=={23,22,1}
    coeff={int(i):dict(zip(matrix["columns"],r))for i,r in zip(matrix["index"],matrix["data"])}
    assert all(finite(v)and v>=0 for r in coeff.values()for v in r.values())
    native_coeff[method]={}
    for scope,cols in SCOPES.items():
        native_coeff[method][scope]={}
        for i in fg["index"]:
            raw=math.fsum(coeff[i][col]for col in cols)
            factor=factors.get(i)
            # Unknown detritus wet/C stays unavailable for wet-catch coefficients.
            stored=raw*9/factor if factor is not None else None
            native_coeff[method][scope][i]=raw
            group_sppr.append([MODEL_ID,group_by_id[i]["group_name"],scope,method,stored])
            bridge.append(dict(group_id=i,group=group_by_id[i]["group_name"],scope=scope,method=method,native_sppr_C_per_C_catch=raw,source_body_wet_per_carbon=factor,regional_sppr_wet_equivalent_per_wet_catch=stored,basis="Native carbon coefficient×9/source body wet:C; map/PPR–NPP divides wet-equivalent output by9 once"if factor is not None else"Detritus wet:C unknown; no wet-catch assignment; native coefficient retained separately"))
        status=f"provisional: native-balanced assumption model; direct {report['config']['TE_option']} {report['status']} (zero model catch; terminal EE=0); source/mapping/carbon-transfer assumptions; scientific review pending"
        run_notes.append(dict(topic=method+"/"+scope,note=status,status=report["status"]))
    health.append({"TE_option":report["config"]["TE_option"],**flatten(report)})
    for warning in report.get("warnings",[]):run_notes.append(dict(topic=report["config"]["TE_option"],note=str(warning),status="WARN"))

set_setting(book,"selected_model_id",MODEL_ID)
set_setting(book,"model_path",MODEL_PATH)
set_setting(book,"selection_rationale","User-authorized adoption of 2000–2014 whole-Okhotsk epipelagic Figure9 plus source-text reconstruction. Exact persisted native carbon model passes real PPRCalculator main/physical checks with all living BA0, positive detritus BA, zero catch/migration and source-supported external diets. Direct GE/TE/WithEgestion return WARN (all-zero model catch and four terminal EE0 groups); source repairs, historical catch transfers and weak analogues retain provisional research display. Mathematical balance is distinct from scientific validation.")
set_setting(book,"selected_paper_ids","OKH-GM2019__LME_052")
prepare_selection(book,WORKBOOK)
book["Selected model groups"]["Groups"]=(group_header,group_rows)
book["Selected model groups"]["Group SPPR"]=(['model_id','group','scope','method','sppr'],group_sppr)
book["Selected model groups"]["Carbon unit bridge"]=table_dict(bridge)
book["Diagnostics"]["model_health"]=table_dict(health)
book["Diagnostics"]["mc_diagnostics"]=(book["Diagnostics"].get("mc_diagnostics",(["method","status"],[]))[0],[])
run_notes.extend([
    dict(topic="Monte Carlo/global",note="NOT_RUN: excluded from the authorized direct-only diagnostic scope; no inherited old-model coefficients/grades",status="NOT_RUN"),
    dict(topic="Native vs regional catch",note="Native model catch is0 by explicit user assumption; diagnostic model footprint0 is not the LME real-catch footprint. Regional catch is unchanged SeaAroundUs wet mass.",status="INFO"),
    dict(topic="Carbon bridge",note="Native source/matrix/Groups flows remain carbon. Stored GroupSPPR = native coefficient×9/bodywet:C. Annual PPR is a wet-equivalent accounting output divided by9 exactly once for tC. Body-carbon transfer to analogue catch taxa is a disclosed weak assumption.",status="INFO"),
    dict(topic="TE zero terminal coefficients",note="Native TE returns finite0 for source groups17(Predatory salmon),18(Baleenwhales),20(Predatoryfish),21(Predatorymammals) withEE0. These outputs and warnings are retained, not replaced with invented values.",status="WARN"),
])
book["Diagnostics"]["run_notes"]=table_dict(run_notes)
matching=read(HERE/"matching_rows.json")
audits=read(HERE/"taxon_audit.json")
book["PPR"]["Matching"]=table_dict(matching)
book["PPR"]["Allocation assumptions"]=table_dict(read(HERE/"allocation_ledger.json"))
book["PPR"]["Mapping review"]=table_dict([{k:r[k]for k in ["taxon","membership_rule","membership_confidence","allocation_rule","allocation_confidence","overall_confidence","assumed","reason"]}for r in audits])
annual=[]
for scope in SCOPES:
    for option,method in METHODS.items():
        status=next(r["note"]for r in run_notes if r["topic"]==method+"/"+scope)
        annual.append([MODEL_ID,scope,method,"landings","method","ppr",status,*[None]*70])
book["PPR"]["Annual"]=(ANNUAL_HEADER,annual)
set_setting(book,"results_model_id",MODEL_ID)
set_setting(book,"results_model_sha256",FINAL_HASH)
set_setting(book,"production_eligible",False)
set_setting(book,"source_note","151 catch labels reviewed for this exact22-group epipelagic reconstruction. 2019landings mapping99.7604625%; Verylow23.8779694%; eight labels unresolved. Source wet-biomass proxy allocates pollock stages and squid/salmon guilds; no observed caught-mass split is claimed. Native carbon coefficient→wet-equivalent bridge is explicit. DirectGE/TE/WithEgestion WARN and zeroTE terminal coefficients remain; no MC/global calculations. Evidence: models/52_GM2019_Fig9_Pelagic_(2000-2014)/research_20261003/integration/.")
set_setting(book,"mapping_limitation","Whole-sea horizontal coverage with epipelagic0–200m trophic scope. Catch1950–2019 extends source2000–2014. Verylow analogue/broad placements include benthic fish/crustaceans/filterfeeders and body-carbon transfer; freshwater/sediment/macroalgal taxa unresolved. Higher mapping coverage and runtime balance do not establish scientific validation.")
set_setting(book,"diet_source_state","Frozen source baselines plus explicitly authorized adopted repairs; source-to-final ledger and failed cases retained")
set_setting(book,"runtime_revalidation_evidence",MODEL_PATH.rsplit('/',1)[0]+"/native_reload_result.json")
set_setting(book,"regional_sppr_units","wet-equivalentPPR/wet catch; native carbon coefficients×9/sourcebodywet:C; output/9once yieldsPPRtC")
recalculate(book,WORKBOOK)
set_setting(book,"calculation_status","provisional: exact native-balanced assumption reconstruction adopted; direct GE/TE/WithEgestion WARN and terminalEE0 retained; mapping/PPR/NPP refreshed; researcher scientific validation pending")
for (s,n),h in protected.items():assert digest_tables(list(book[s][n]))==h,(s,n,"protected inputs changed")
write_book(WORKBOOK,book)
saved=read_fast(WORKBOOK)
validate_region(saved,WORKBOOK)
assert overview(saved)["calculation_result_sha256"]==result_hash(saved)
assert sha(MODEL)==FINAL_HASH
assert sha(REGION/"models/52_1_Sea_of_Okhotsk_NE_(1980)/model.json")==original_model_hash

# Standard upstream source workbook: native Groups, native diagnostic grades;
# the atlas-consumed SPPR sheets explicitly use the regional unit bridge.
import openpyxl
from openpyxl.styles import Alignment,Font,PatternFill
sw=openpyxl.Workbook();sw.remove(sw.active)
def add_sheet(name,header,data):
    s=sw.create_sheet(name);s.append(header)
    for row in data:s.append(row)
    s.freeze_panes="C2";s.auto_filter.ref=s.dimensions;s.sheet_view.showGridLines=False
    for c in s[1]:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="174C58")
    for row in s:
        for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)
    s.column_dimensions["A"].width=15;s.column_dimensions["B"].width=34
    return s
add_sheet("groups_df",group_header,group_rows)
for scope in SCOPES:
    values={(g,s,m):v for _,g,s,m,v in group_sppr}
    add_sheet("sppr_"+scope,["seq","group_name",*METHODS.values()],[[i,group_by_id[i]["group_name"],*[values[group_by_id[i]["group_name"],scope,m]for m in METHODS.values()]]for i in fg["index"]])
    add_sheet("native_C_"+scope,["seq","group_name",*METHODS.values()],[[i,group_by_id[i]["group_name"],*[native_coeff[m][scope][i]for m in METHODS.values()]]for i in fg["index"]])
hh,hr=table_dict(health);add_sheet("model_health",hh,hr)
add_sheet("mc_diagnostics",["method","status"],[["Monte Carlo","NOT_RUN (direct-only scope)"]])
nh,nr=table_dict(run_notes);add_sheet("run_notes",nh,nr)
bh,br=table_dict(bridge);add_sheet("Carbon bridge",bh,br)
add_sheet("Provenance",["field","value"],[["model_id",MODEL_ID],["model_sha256",FINAL_HASH],["native_model_currency","tC/km2, peryear forflows"],["atlas_scope_sheet_currency","wet-equivalentPPR/wet catch"],["unit_conversion","nativeC/C×9/bodywet:C; displayoutput/9once"],["source_evidence","source_to_final_ledger.json"],["matrices","diagnostics/{GE,TE,With_Egestion}_full_return.json; rawcarbon matrices retained"],["production_eligible",False]])
sw.save(VARIANT/"sppr_source.xlsx")
save(HERE/"carbon_unit_bridge.json",dict(schema_version=1,model_id=MODEL_ID,exact_model_sha256=FINAL_HASH,formula="group_SPPR_wet_equivalent = group_SPPR_carbon * 9 / group_body_wet_per_carbon; PPR_tC = wet_catch * weighted_SPPR_wet_equivalent / 9",source_factors="../../assumption_variants/adopted_balanced_20261003/source_to_final_ledger.json",rows=bridge))

checks={"saved_input_hash_fresh":overview(saved)["calculation_input_sha256"]==input_hash(saved),"saved_result_hash_fresh":overview(saved)["calculation_result_sha256"]==result_hash(saved),"selected_and_results_identity":overview(saved)["selected_model_id"]==overview(saved)["results_model_id"]==MODEL_ID,"exact_model_hash_retained":sha(MODEL)==FINAL_HASH,"old_model_bytes_unchanged":sha(REGION/"models/52_1_Sea_of_Okhotsk_NE_(1980)/model.json")==original_model_hash,"protected_catch_classic_taxa_npp_inputs":all(digest_tables(list(saved[s][n]))==h for (s,n),h in protected.items()),"all23_native_groups_retained":len(records(saved,"Selected model groups","Groups"))==23,"direct_only9_scope_methods":{r["method"]for r in records(saved,"Selected model groups","Group SPPR")}==set(METHODS.values()),"diagnostic_warns_unmodified":all(r["status"]=="WARN"for r in records(saved,"Diagnostics","model_health")),"production_eligible_false":overview(saved)["production_eligible"]is False}
new_classic={tuple(r[k]for k in ["scope","method","catch_basis","unidentified","metric"]):[r[y]for y in range(1950,2020)]for r in records(saved,"Classic PPR","Annual")}
checks["independent_classic_annual_values_unchanged"]=all(all(a==b or(finite(a)and finite(b)and math.isclose(a,b,rel_tol=1e-14))for a,b in zip(vals,new_classic.get(key,[])))and key in new_classic for key,vals in old_classic.items())
catch_rows=records(saved,"Catch","Catch")
catch_by={(r["taxon"],r["catch_basis"]):r for r in catch_rows}
taxon_sppr={(r["taxon"],r["scope"],r["method"]):r["sppr"]for r in records(saved,"PPR","Taxon SPPR")}
annual2019=[]
for r in records(saved,"PPR","Annual"):
    if r["metric"]!="ppr"or r["unidentified"]!="method":continue
    basis=r["catch_basis"];scope=r["scope"];method=r["method"]
    pairs=[(c[2019],taxon_sppr.get((t,scope,method)))for(t,b),c in catch_by.items()if b==basis]
    independent=math.fsum(c*v for c,v in pairs if finite(c)and finite(v))
    covered=math.fsum(c for c,v in pairs if finite(c)and finite(v))
    allmass=math.fsum(c for c,v in pairs if finite(c))
    annual2019.append(dict(scope=scope,method=method,catch_basis=basis,status=r["status"],regional_ppr_wet_equivalent=r[2019],regional_ppr_tC=r[2019]/9 if finite(r[2019])else None,independent_arithmetic_wet=independent,covered_catch_tonnes=covered,total_catch_tonnes=allmass,coverage_percent=100*covered/allmass if allmass else None,arithmetic_reconciles=finite(r[2019])and math.isclose(r[2019],independent,rel_tol=1e-14)))
checks["annual2019_arithmetic_reconciles"]=all(r["arithmetic_reconciles"]for r in annual2019)
evidence=dict(schema_version=1,model_id=MODEL_ID,model_path=MODEL_PATH,final_model_sha256=FINAL_HASH,before_workbook_sha256=before_hash,after_workbook_sha256=sha(WORKBOOK),baseline_workbook="LME_052_before_GM2019_adoption.xlsx",sppr_source_sha256=sha(VARIANT/"sppr_source.xlsx"),overview=overview(saved),checks=checks,annual2019=annual2019,passed=all(checks.values()),central_integration="Root coordinator performsProject.xlsx/map refresh after this regional handoff")
save(HERE/"regional_adoption_verification.json",evidence)
print(json.dumps({k:v for k,v in evidence.items()if k!="overview"},ensure_ascii=False,indent=2))
if not evidence["passed"]:raise SystemExit(1)
