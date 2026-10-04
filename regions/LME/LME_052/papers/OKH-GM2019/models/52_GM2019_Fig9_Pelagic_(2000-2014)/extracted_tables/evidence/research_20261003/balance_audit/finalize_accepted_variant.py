"""Verify saved import cells, reject lossy conversion and package this revision."""
from pathlib import Path
from decimal import Decimal
from datetime import datetime, timezone
import csv, hashlib, importlib.util, json, os
from openpyxl import load_workbook
from pypdf import PdfReader

WORK=Path(__file__).resolve().parent
HERE=WORK.parents[1]
ROOT=HERE.parents[3]
OUT=HERE/'assumption_variants/researcher_readings_20261003'
TABLES=OUT/'import_tables'
PAPERS=HERE.parents[1]/'papers/OKH-GM2019'
D=Decimal
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rows(p):
    with p.open(encoding='utf-8',newline='') as f:return list(csv.reader(f))

ext=read(OUT/'extraction.json');model=read(OUT/'model.json');ready=read(OUT/'READINESS.json')
budgets=read(OUT/'flow_budget_checks.json');factors=read(OUT/'conversion_factors.json')
original_checks=read(OUT/'verification.json')
assert all(sha(Path(p))==h for p,h in original_checks['protected_original_hashes_unchanged'].items())
names={g['n']:g['name'] for g in ext['groups']}
for filename in ['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv']:
    raw=(TABLES/filename).read_bytes()
    assert raw.endswith(b'\r\n') and not raw.startswith(b'\xef\xbb\xbf')
    assert b'\n' not in raw.replace(b'\r\n',b'') and b'"' not in raw
basic=rows(TABLES/'Basic_input.csv')
assert len(basic)==23
for row,g in zip(basic[1:],ext['groups']):
    assert int(row[0])==g['n'] and row[1]==g['name']
    for column,key in [(3,'biomass'),(5,'pb'),(6,'qb'),(10,'unassim'),(11,'detritus_import')]:
        assert row[column]==('' if g.get(key) is None else str(g[key]))
diet=rows(TABLES/'Diet_composition.csv')
assert diet[0][2:]==[str(n) for n in range(2,22)]
for i,row in enumerate(diet[1:23],1):
    assert row[:2]==[str(i),names[i]]
    for j,value in zip(range(2,22),row[2:]):
        expected=ext['diet'][str(j)][str(i)]
        assert value==('' if expected is None else str(expected))
assert sum(v=='' for row in diet[1:23] for v in row[2:])==6
fate=rows(TABLES/'Detritus_fate.csv')
assert fate[0][2]=='Detritus' and all(row[2]=='1' for row in fate[1:22])
assert all(g['biomass_accum']=='-9999' for g in model['group'])
assert all(g['gs']=='0.2' for g in model['group'][1:21])
assert all(next(c for c in g['diet_descr']['diet'] if c['prey_seq']=='22')['detritus_fate']=='1' for g in model['group'][:21])
for filename in ['TL.xlsx','Metadata.xlsx']:
    book=load_workbook(TABLES/filename,read_only=True,data_only=True)
    if filename=='TL.xlsx':assert book.active.max_row==23 and book.active['C1'].value=='TL'
    else:assert book.active['A1'].value=='LME' and book.active['B2'].value=='GM2019-Fig9'
    book.close()

native_path=next(TABLES.glob('52_*.json'))
native=read(native_path)
native_groups={int(g['group_seq']):g for g in native['group']}
omitted=[n for n in range(1,23) if n not in native_groups]
misclassified=[n for n in native_groups if native_groups[n]['pp']!=model['group'][n-1]['pp']]
assert omitted==[2,3,22] and misclassified==[4,5]
assert len(native_groups)==19
conversion=dict(verdict='REJECTED_LOSSY_CONVERSION',raw_converter_file=native_path.name,
    raw_group_count=19,faithful_group_count=22,omitted_group_ids=omitted,
    wrongly_classified_consumer_ids=misclassified,sole_detritus_group_omitted=True,
    user_detritus_fate_dropped_by_converter=True,final_model='../model.json',
    native_EwE_load_tested=False,ready=False)
save(OUT/'converter_preservation_audit.json',conversion)
(TABLES/'REJECTED_CONVERTER_OUTPUT.txt').write_text(
    'The 52_*.json and its reconstructed workbook are rejected audit outputs. '
    'The converter omits groups2,3,22, misclassifies consumers4,5 as producers, '
    'and loses the sole-detritus routing. Use ../model.json for the preserved22-group '
    'candidate. Its scientific readiness remains BLOCKED. See the parent REPORT.md.\n',encoding='utf-8')

factor_table='\n'.join(f"| {r['group_id']} | {r['group']} | {r['wet_per_carbon'] or 'Unknown'} | {r['basis']}{'; tentative crosswalk' if r['mapping_tentative'] else ''} |" for r in factors)
energy_table='\n'.join(f"| {r['group']} | {r['figure_P_carbon']} | {r['Q_carbon']} | {r['respiration_carbon_GS_0p2']} |" for r in budgets if r['energy_failure_carbon'])
pred_table='\n'.join(f"| {r['group']} | {r['figure_P_carbon']} | {r['predation_carbon']} | {float(r['predation_over_figure_P']):.6f} | {'Yes' if r['clear_predation_already_exceeds_figure_P'] else 'No'} |" for r in budgets if r['predation_failure_carbon'])
report=f'''# Sea of Okhotsk Figure 9 researcher readings revision (2000–2014)

Status: **BLOCKED — not yet loadable and balanced for researcher validation.** This is a derivative of the preserved Figure-only reconstruction. The original extraction, existing selected NE model, validation DOCX and map remain untouched. Artifact arithmetic and import formatting were verified; biological balance and native EwE load are not established.

**Source:** Gorbatenko and Melnikov (2019), DOI10.26428/1606-9919-2019-198-143-163, Table3 printedp154/PDFp12, Figure9 printedp157/PDFp15; explicit bacteria conversion from Gorbatenko2018 dissertationp60. LME52, local model IDGM2019-Fig9,22groups(21living,1detritus),20consumers, fisheries unknown. Revision date2026-10-03.

## Accepted readings and source corrections

- F62 preferred0.023; alternative0.025 retained in F62_0p025_sensitivity.json, as requested by the researcher. F67=0.05 human confirmed. Source-readable fields and the original ledger still retain the overprinted uncertainty.
- F77 source changes from large pollock19 to medium pollock15, destination predatory fish20, flow0.01. This follows the actual green shaft and arrowhead; no balance criterion was used as tracing evidence.
- F13 and F22 endpoints are unchanged, with improved visual routing confidence. F78 remains15→21 tentative, literal0.05. Other dense red routes and three duplicated adopted pairs remain unresolved.
- Bacteria wet/carbon9.4 is explicitly supported by dissertationp60. Protozoa10 is only a conflicting derivation, not adopted. Detritus factor remains unknown.

With these decisions the Figure-only variant has20/20 numerical carbon diet columns,18/20 wet columns and32 tentative routes. These counts establish arithmetic completeness under adopted routes, not source approval.

## All conversion factors

Factor means tonnes wet mass per tonne carbon. Multiply carbon flow by the **prey's** factor before calculating wet diet fractions; the consumer's body factor cannot convert its total food consumption.

| ID | Group | Wet/carbon | Evidence |
|---|---|---:|---|
{factor_table}

The2019 pooled microheterotroph row has9.4; it does not independently establish both individual microbial factors. Dissertationp60 protozoa27gC/m² and270g wet/m² imply10, but the corresponding whole-Sea total conflicts. Table3's predatory-fish Other**** mapping remains tentative.

## User assumptions

GS=0.2 for20consumers. All other mortality and unassimilated consumption route100% to detritus22. Detritus import/export=0. Detritus accumulation is the uncomputed residual BA_D=Σ(GS×Q+other mortality+actual discarded biomass entering detritus)−Q_D. It is not assumed zero. Living catches, BA and net migration remain unknown.

## Remaining basic inputs and diets

Bacteria and protozoa still lack separate annual whole-Sea B and P/B. Their Q/B cannot be identified without stocks; bacteria's wet Q also needs the detritus conversion. Copepod and euphausiid wet diets each have three unknown positive fractions because the protozoa prey factor is unknown. Their other prey fractions cannot be normalized until the denominator exists. Their Q/B is independently recoverable from prose:2945/108.3=27.192982456/year and1445.6/58.7=24.626916525/year. These prose alternatives are not silently substituted for Figure-derived Q/B.

Baleen whale Q is now0.21415milliontC/year or2.25652millionwet t/year; using Table3B0.721 gives Q/B3.129708738/year. F62=.025 instead adds0.002milliontC and0.0144millionwet t/year to whale Q. It cannot fix the current failing groups.

## Existing balance failures

The three carbon energy failures below use Figure production and adopted incoming arrows, with GS=.2. Flows are milliontC/year; respiration R=.8Q−P must be positive.

| Consumer | P | Q | R |
|---|---:|---:|---:|
{energy_table}

All three also haveP>Q, so setting GS=0 would not cure them. Chaetognaths instead pass carbon R=+3.7528 but fail the wet B×PB/Q diagnostic, R=−26.28032millionwet t/year. Different prey/body water content makes wet mass unsuitable as a direct carbon-energy identity here. Source-stock conversions and printed carbon stocks also conflict for some groups; a complete carbon model must document a consistent input variant.

Predation exceeds Figure production for six groups under the adopted full route interpretation. This ratio becomes an implied EE only in the diagnostic with zero living catch, BA and net migration; those zero values are **not** adopted inputs.

| Prey | Figure P | Predation | Predation/P | Clear-route lower bound already fails? |
|---|---:|---:|---:|---|
{pred_table}

Even clear routes give smelt0.125>0.06, capelin0.115>0.03, squidIV0.158>0.1(or Table3P0.105). Minimum carbon production deficits are0.065,0.085,0.058milliontC/year. Catch/positive BA/outward migration worsen a deficit; net inward migration or negative BA could offset it, but no annual measured values are identified. These stock-balance terms cannot fix negative consumer respiration at fixedP,Q,GS.

## Article results available for testing closure hypotheses

Table3 stocks, turnover and production; Figure9 carbon flows; prose annual Q by species/groups; Table2 percentages of zooplankton production eaten; tier production totals and the food split of the nekton community. Independent wet Q targets(milliont/year): copepods2945,euphausiids1445.6,hyperiids152.5,chaetognaths320,herring36.9,smelt10.29,capelin14.74,small pollock16.5,medium pollock80.8,pooled squid32.4. Total nekton195 is better supported by most prose/ratio checks and the2016 target194.9, while the abstract159 conflicts and its English carbon unit is erroneous. These are competing recorded source targets.

The production equation P=predation+catch+BA+net outward migration+other mortality constrains a combined residual; the paper does not independently identify catch, BA and migration. Seasonal herring survey changes are explained by coverage and cannot be treated as measured annual BA/import. Methods include growth and mortality losses, so there is no evidence for treating Table3P/B simply as surviving-stock growth or a source mortality Z.

Additional dissertation targets are saved in ../../research_20261003/article_constraints/FOLLOWUP_RECOVERIES.md: salmonIII Q1.79694millionwet t/year for pink/chum/sockeye,2000–2014; derived Q/B2.246175 with2019B0.8. Jellyfish Q3.1771millionwet t/year covers2006–2014 and different stock5.0975milliont. Its independently converted copepod/euphausiid/chaetognath food gives carbon Q≥0.209584623, enough for the illustrative carbon energy minimum, but its feeding links are absent from the Figure-only network. These are alternate supporting-source evidence, not a validated mixed-source model.

## Biomass accumulation

Methods/prose and supporting source surveys supply no matching group-specific annual living BA/net migration/catch values or explicit closed steady-state assertion. All corresponding unknowns are preserved. Detritus BA is an allowed residual, calculable after complete living budgets; stock depletion feasibility would require an initial detritus stock.

## Import validation and converter audit

Eight mandatory imports are in import_tables/, writer override --dir-name import_tables documented there. Import validation:0errors,9warnings. Wet mass-balance checker:10errors,9warnings,2notes(4energy and6predation flags). Its inferred detritus count and under-counted microbial feeding are artifacts of missing basics, not evidence of actual extra detritus pools. Dedicated22-group carbon/flow audits govern scientific interpretation.

The raw database conversion is retained but rejected:19groups, omitted2/3/22, wrongly classified4/5 as producers, sole-detritus routing lost. Its native check reports4errors/6indeterminate/7warnings, and does not establish native loadability. The faithful parent model.json preserves all22groups and440diet cells. No normalization or fitted source values was used to force balance.

## Text supplementation

The user's2026-10-03 instruction to inspect whether Figure9 omits links is handled in ../../research_20261003/text_feeding_audit/. The strict Figure-only revision is preserved. Any extra prose-supported links or fractions are separate evidence, with unallocated remainder and qualitative links retained as unknowns. See that audit's FINDINGS.md for the exact completeness wording and extraction; do not conflate a generalized scheme with an explicit claim about all omitted arrows.

## English article

The complete26-page unofficial English PDF is in the papers/OKH-GM2019 source bundle. All311 table numbers and original Figure9 pixels were checked; source inconsistencies were preserved. Original Russian sources remain authoritative.
'''
for old,new in {
    'milliontC/year':'million tC/year', 'millionwet t/year':'million wet tonnes/year',
    'milliont/year':'million tonnes/year', 'millionwet t':'million wet tonnes',
    'milliontC':'million tC', 'milliont.':'million tonnes.',
    'Table3':'Table 3', 'Table2':'Table 2', 'Figure9':'Figure 9',
    'dissertationp60':'dissertation p.60', 'dissertationp':'dissertation p.',
    'printedp':'printed p.', 'PDFp':'PDF p.', 'LME52':'LME 52',
    'IDGM2019':'ID GM2019', '22groups':'22 groups', '21living':'21 living',
    '1detritus':'1 detritus', '20consumers':'20 consumers', '22-group':'22-group',
    'wet/carbon9.4':'wet/carbon 9.4', 'protozoa27':'protozoa 27',
    '270g wet':'270 g wet', '27gC':'27 gC', '195 is':'195 is',
    'dissertation2018':'dissertation 2018', 'Gorbatenko2018':'Gorbatenko (2018)',
    '26-page':'26-page', '311 table':'311 table', 'sole22':'sole 22',
    'GS=0.2':'GS = 0.2', 'GS=.2':'GS = 0.2', 'GS=0':'GS = 0',
    '0errors':'0 errors', '9warnings':'9 warnings', '10errors':'10 errors',
    '2notes':'2 notes', '4energy':'4 energy', '6predation':'6 predation',
    '19groups':'19 groups', 'all22groups':'all 22 groups', '440diet':'440 diet',
    '4errors':'4 errors', '6indeterminate':'6 indeterminate', '7warnings':'7 warnings',
    'all311':'all 311', 'complete26':'complete 26', 'salmonIII':'salmon III',
    'wet t/year':'wet tonnes/year', 'cop+':'cop +', '20/20carbon':'20/20 carbon',
    '18/20wet':'18/20 wet', '32tentative':'32 tentative',
    'copepods2945':'copepods 2945', 'euphausiids1445.6':'euphausiids 1445.6',
    'hyperiids152.5':'hyperiids 152.5', 'chaetognaths320':'chaetognaths 320',
    'herring36.9':'herring 36.9', 'smelt10.29':'smelt 10.29', 'capelin14.74':'capelin 14.74',
    'pollock16.5':'pollock 16.5', 'pollock80.8':'pollock 80.8', 'squid32.4':'squid 32.4',
}.items():report=report.replace(old,new)
(OUT/'REPORT.md').write_text(report,encoding='utf-8')
save(OUT/'import_check_summary.json',dict(import_validation={'errors':0,'warnings':9},
    wet_import_massbalance={'errors':10,'warnings':9,'notes':2},
    rejected_native_converter={'errors':4,'indeterminate':6,'warnings':7},
    full_22_group_budget_file='flow_budget_checks.json',native_EwE_load_tested=False,balanced=False))
original_checks['saved_import_cells_and_types_reconciled']=True
original_checks['six_wet_diet_unknowns_preserved']=True
original_checks['eight_import_format_and_workbook_checks_passed']=True
original_checks['lossy_raw_converter_rejected']=conversion
save(OUT/'verification.json',original_checks)
save(HERE/'CURRENT_VARIANT.json',dict(current_figure_variant='assumption_variants/researcher_readings_20261003',
    source_baseline='model.json',source_baseline_preserved=True,
    text_supplement_evidence='research_20261003/text_feeding_audit',
    combined_diet_case='assumption_variants/text_plus_figure_20261003/diet_constraints.json',
    required_balance_cases_file='BALANCE_CHECK_INPUTS.json',
    selected=False,status='BLOCKED',balanced=False,ready_for_researcher_validation=False))
index=HERE.parent/'MASTER_INDEX.md'
index.write_text('''# LME052 regional reconstruction extraction set

This index covers the new extraction attempt. Existing NE and SD models remain unchanged; it does not alter selection or map status.

| ID | Model | Period | LME | Groups | Fleets | Status | Notes |
|---|---|---|---|---|---|---|---|
| GM2019-Fig9 | [Sea of Okhotsk Figure9 researcher readings revision](52_GM2019_Fig9_Pelagic_(2000-2014)/assumption_variants/researcher_readings_20261003/REPORT.md) |2000–2014|52|22|unknown|partial|20/20carbon,18/20wet diets under adopted routes;32tentative routes; F62=.023(.025alternative),F67=.05; bacteria9.4; energy/balance failures. Original snapshot preserved. Text supplementation audit separate. Not selected or ready for SPPR/PPR.|
''',encoding='utf-8')

english=PAPERS/'gorbatenko_melnikov_2019_English_translation.pdf'
assert len(PdfReader(english).pages)==26
assert sha(english)=='89b68a296e6be280f42028d536f786fab9a23776fbe8593dcc7e98c11f1054e4'
artifacts=[]
def add(role,p):
    artifacts.append(dict(role=role,availability='present',path=Path(os.path.relpath(p,OUT)).as_posix(),sha256=sha(p)))
for s in read(PAPERS/'source_manifest.json')['sources']:add('original_source_pdf',PAPERS/s['file'])
add('english_translation',english)
for p in [OUT/'REPORT.md',OUT/'READINESS.json',OUT/'accepted_decisions.json',OUT/'model.json',OUT/'extraction.json',
    OUT/'flow_readings.json',OUT/'diet_cell_ledger.json',OUT/'flow_budget_checks.json',OUT/'carbon_reconstruction.json',
    OUT/'F62_0p025_sensitivity.json',OUT/'conversion_factors.json',OUT/'assumptions.json',OUT/'verification.json',
    OUT/'import_check_summary.json',OUT/'converter_preservation_audit.json']:
    add('revision_evidence',p)
for p in sorted(TABLES.iterdir()):
    if p.is_file():add('imports_and_rejected_conversion_audit',p)
for p in [HERE/'research_20261003/arrow_audit/proposed_route_changes.json',
    HERE/'research_20261003/arrow_audit/production_node_audit.json',
    HERE/'research_20261003/article_constraints/FINDINGS.md',
    HERE/'research_20261003/article_constraints/FOLLOWUP_RECOVERIES.md',
    HERE/'research_20261003/article_constraints/followup_salmon_recovery.json',
    HERE/'research_20261003/article_constraints/followup_jellyfish_recovery.json',
    WORK/'create_accepted_variant.py',WORK/'check_both_diet_sources.py',
    WORK/'implied_protozoa_conversion_checks.json',
    HERE/'BALANCE_CHECK_INPUTS.json',HERE/'CURRENT_VARIANT.json',
    HERE/'research_20261003/arrow_audit/copepod_smelt_finding.json',
    Path(__file__).resolve()]:add('supporting_audit',p)
for p in sorted((HERE/'assumption_variants/text_plus_figure_20261003').iterdir()):
    if p.is_file():add('mandatory_combined_balance_case',p)
text_audit=HERE/'research_20261003/text_feeding_audit'
if text_audit.exists():
    for p in sorted(text_audit.iterdir()):
        if p.is_file():add('primary_article_text_supplement',p)
save(OUT/'evidence_index.json',dict(schema_version=1,run_id='LME052-GM2019-FIG9-20261003-RESEARCHER-READINGS',
    region_id='LME_052',model_id=HERE.name,variant_id='researcher_readings_20261003',
    source_identity=dict(article_doi='10.26428/1606-9919-2019-198-143-163',figure=9,
        source_pdf_sha256=sha(PAPERS/'gorbatenko_melnikov_2019.pdf')),
    computational_input_identity=dict(accepted_decisions_sha256=sha(OUT/'accepted_decisions.json'),
        source_flow_ledger_sha256=sha(HERE/'audit/flow_readings.json'),
        table3_source_ledger_sha256=sha(HERE/'audit/table3_source_readings.json')),
    methods=['Source figure transcription with retained route uncertainty',
        'Accepted human numeric readings and independent source-path correction',
        'Carbon-flow normalization and prey-specific wet conversion',
        'Independent carbon and wet energy/production diagnostics',
        'Eight import file and faithful 22-group JSON reconciliation',
        'Lossy raw converter audit and explicit rejection'],
    required_roles=sorted({r['role'] for r in artifacts}),artifacts=artifacts,
    scientific_status='BLOCKED; artifact verification does not establish balance or author approval'))
spec=importlib.util.spec_from_file_location('checker',ROOT/'tools/skills/original_skill_resources/combined-src/scripts/check_evidence.py')
checker=importlib.util.module_from_spec(spec);spec.loader.exec_module(checker)
integrity=checker.check(OUT/'evidence_index.json');save(OUT/'evidence_integrity.json',integrity)
assert integrity['complete'],integrity
assert all(sha(Path(p))==h for p,h in original_checks['protected_original_hashes_unchanged'].items())
print(json.dumps(dict(artifact_verification_passed=True,evidence_files=len(integrity['verified_artifacts']),
    groups=22,wet_complete=18,carbon_complete=20,all_original_hashes_unchanged=True,
    english_PDF_pages=26,scientifically_ready=False),ensure_ascii=False,indent=2))
