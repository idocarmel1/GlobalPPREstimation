"""Preserve all figure groups, verify derived artifacts, and inventory evidence.

This is a research extraction, not a balanced or selected Ecopath model. Native
converter outputs are retained as rejected audit evidence, never used to repair
source values. Run after reconstruct.py and audit/build_review.mjs.
"""
from pathlib import Path
from datetime import datetime, timezone
from decimal import Decimal
import csv, hashlib, importlib.util, json, subprocess, sys
from openpyxl import load_workbook
from pypdf import PdfReader

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[3]
AUDIT=HERE/'audit'
TABLES=HERE/'extracted_tables'
PAPERS=HERE.parents[1]/'papers/OKH-GM2019'
SKILL=Path('C:/Users/idoca/.agents/skills/ecopath-extraction')
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
data=read(AUDIT/'workbook_data.json')
source=read(PAPERS/'source_manifest.json')
ext=read(TABLES/'extraction.json')
raw_path=next(TABLES.glob('52_*.json'))
raw=read(raw_path)
raw_groups={int(g['group_seq']):g for g in raw['group']}

# Preserve a faithful database-shaped review candidate, including unknown diets.
# Unknown values remain sentinel strings, not importer-generated defaults.
groups=[]
for b,g in zip(ext['groups'],data['groups']):
 n=g['n']
 x={k:'-9999' for k in ['biomass_habitat_area','biomass','vbk','pb','ee','biomass_accum','biomass_accum_rate','qb','detritus_import','respiration','immigration','emigration','emigration_rate','other_mort','export','gs','shadow_price','ge']}
 x.update(group_name=g['name'],group_seq=str(n),habitat_area='1',pp='1' if n==1 else '2' if n==22 else '0',diet_imp='0',b_hab_area_input='false',pb_input='false',qb_input='false',ee_input='false',ge_input='false',taxon_descr=None,pedigree_assignment_descr=None,diet_descr=None)
 for key,inputkey in [('biomass','b_hab_area_input'),('pb','pb_input'),('qb','qb_input')]:
  if b[key] is not None:x[key]=str(b[key]);x[inputkey]='true'
 x['biomass_habitat_area']=x['biomass']
 if n in ext['consumers']:
  x['diet_descr']={'diet':[{'prey_seq':str(i),'proportion':ext['diet'][str(n)][str(i)] if ext['diet'][str(n)][str(i)] is not None else '-9999','detritus_fate':'-9999'} for i in range(1,23)]}
 groups.append(x)
model={'group':groups,'_reconstruction':dict(status='BLOCKED_RESEARCH_CANDIDATE_NOT_FOR_SPPR',selected=False,model_id=HERE.name,article_doi='10.26428/1606-9919-2019-198-143-163',period='2000-2014',scope='epipelagic 0-200 m',currency='wet-weight biomass t/km2; P/B and Q/B /year; diet wet weight',area_km2='1544000',area_basis='same implied whole-Sea reporting denominator as Table 3; not a geographic coverage measurement',unknown_sentinel='-9999',zero_diet_import_basis='research convention: only displayed feeding arrows represented; not an author import estimate',missing_displayed_arrows_zero=True,tentative_routes=33,unresolved_label_readings=['F62','F67'],unknown_wet_conversion_prey_ids=[2,3,22],ee_gs_catches_ba_migration_and_routing_not_estimated=True,source_pdf_sha256=source['sources'][0]['sha256'],raw_converter_output_rejected=str(raw_path.relative_to(HERE)).replace('\\','/'),limitations='19/20 carbon and 16/20 wet diet columns numerically identifiable under tentative routes; energy failures and missing inputs prevent a validated Ecopath model')}
write(HERE/'model.json',model)

omitted=[g['n'] for g in data['groups'] if g['n'] not in raw_groups]
wrong=[g['n'] for g in data['groups'] if g['n'] in raw_groups and raw_groups[g['n']]['pp']!=groups[g['n']-1]['pp']]
conversion=dict(raw_converter_file=str(raw_path.relative_to(HERE)).replace('\\','/'),verdict='REJECTED_LOSSY_CONVERSION',raw_group_count=len(raw_groups),faithful_group_count=22,omitted_group_ids=omitted,misclassified_group_ids=wrong,corrected_candidate='model.json',corrections=['Restore all original figure groups without inventing missing stock parameters','Use explicit source group type, not Q/B availability, for pp','Preserve all 440 prey/consumer cells, including unknown fractions as -9999','Keep GS, catches, BA, migration, export and detritus routing unknown','Retain original native-converter JSON and round-trip workbook unchanged as rejected evidence'],source_values_changed=False,unknowns_restored=True,balanced=False)
write(AUDIT/'converter_preservation_audit.json',conversion)
readiness=dict(status='BLOCKED',user_readiness_gate='loadable and balanced before researcher validation',json_structure_preserved=True,native_Ecopath_load_tested=False,balanced=False,ready_for_researcher_validation=False,missing_living_biomass_group_ids=[2,3],missing_living_PB_group_ids=[2,3],missing_consumer_QB_group_ids=[2,3,4,5,18],incomplete_wet_diet_group_ids=[3,4,5,18],unresolved_label_ids=['F62','F67'],tentative_routes=33,energy_failures_carbon=data['summary']['energy_failure_groups'],source_values_adjusted_to_force_balance=False,possible_next_evidence=['Resolve or recover original vector/native network for uncertain endpoints','Recover separate microbial stock/conversion and complete consumer ration inputs','Establish whether Figure 9 omits feeding links and reconcile independent reported consumption'],limits='No native-model load or balance claim is supported. A normalized diet or JSON file alone does not meet the readiness gate.')
write(HERE/'READINESS.json',readiness)
(TABLES/'REJECTED_CONVERTER_OUTPUT.txt').write_text('The 52_*.json and *_reconstructed.xlsx files are original native-converter evidence, not final models. The converter dropped groups 2, 3 and 22; misclassified 4, 5 and 18 as producers; and substituted defaults for unknowns. Use ../model.json and the source ledgers for the preserved review candidate. It remains scientifically BLOCKED. See ../REPORT.md and ../audit/converter_preservation_audit.json.\n',encoding='utf-8')

# Capture the actual import-validation and mass-balance outcomes separately.
checks={}
for script,label in [('validate.py','import_validation'),('massbalance_check.py','import_massbalance')]:
 r=subprocess.run([sys.executable,str(SKILL/'scripts'/script),str(TABLES)],capture_output=True,text=True,encoding='utf-8')
 (AUDIT/f'{label}.txt').write_text(r.stdout+r.stderr,encoding='utf-8')
 checks[label]={'exit_code':r.returncode,'log':f'{label}.txt','summary_lines':[line for line in r.stdout.splitlines() if 'error(s)' in line or 'warning(s)' in line]}

# Independent reconciliation: exact-string source numbers, Decimal calculations,
# saved formula cache and CSV matrices must describe the same extraction.
results={}
for s in source['sources']:
 p=PAPERS/s['file'];assert sha(p)==s['sha256'];assert len(PdfReader(p).pages)==s['pages']
results['six_original_pdf_hashes_and_page_counts']=True
ledger=read(AUDIT/'diet_cell_ledger.json')
assert len(ledger)==440 and len(data['flows'])==78 and len(groups)==22
assert sum(not r['flow_ids'] for r in ledger)==365
assert all(r['carbon_flow']=='0' and r['wet_flow']=='0' and r['carbon_dc']=='0' and r['wet_dc']=='0' for r in ledger if not r['flow_ids'])
results['all_440_cells_and_365_absent_arrow_zeros']=True
assert omitted==[2,3,22] and wrong==[4,5,18]
assert [g['group_seq'] for g in model['group']]==[str(n) for n in range(1,23)]
assert sum(g['pp']=='0' for g in groups)==20 and sum(g['pp']=='1' for g in groups)==1 and groups[-1]['pp']=='2'
for g in groups:
 assert g['gs']=='-9999' and g['ee']=='-9999' and g['biomass_accum']=='-9999' and g['export']=='-9999'
 if g['diet_descr']:
  for x in g['diet_descr']['diet']:
   v=ext['diet'][g['group_seq']][x['prey_seq']]
   assert x['proportion']==(v if v is not None else '-9999') and x['detritus_fate']=='-9999'
results['all_groups_correct_types_and_unknowns_preserved_in_model_json']=True
assert all(f['readable_carbon_flow'] is None for f in data['flows'] if f['flow_id'] in ['F62','F67'])
results['overprinted_readings_unknown_in_main_variant']=True
wb=load_workbook(HERE/'DC_reconstruction_review.xlsx',read_only=True,data_only=True)
for sheet,key,csvname,expected in [('DC carbon','carbon_dc','Diet_composition_carbon.csv',19),('DC wet','wet_dc','Diet_composition_wet_review.csv',16),('Label hypothesis','carbon_hypothesis_dc','Diet_composition_carbon_label_hypothesis.csv',20),('Carbon flows','carbon_flow','Food_flows_carbon.csv',None)]:
 with (TABLES/csvname).open(encoding='utf-8',newline='') as f:rows=list(csv.reader(f))[1:]
 assert len(rows)==22
 sh=wb[sheet]
 for i in range(22):
  for j in range(20):
   val=data[key][i][j];cached=sh.cell(i+7,j+3).value
   assert rows[i][j+2]==('' if val is None else val)
   if val is None:assert cached=='n.a.',(sheet,i,j,cached)
   else:assert isinstance(cached,(int,float)) and abs(Decimal(str(cached))-Decimal(val))<Decimal('1e-12'),(sheet,i,j,val,cached)
 if expected is not None:
  complete=0
  for j in range(20):
   vals=[data[key][i][j] for i in range(22)]
   total=sh.cell(30,j+3).value
   if all(v is not None for v in vals):
    complete+=1;assert abs(sum(Decimal(v) for v in vals)-1)<Decimal('1e-12');assert abs(total-1)<1e-12
   else:assert total=='n.a.'
  assert complete==expected
wb.close()
results['csv_json_and_saved_workbook_matrix_reconciliation']=True
results['complete_diet_columns_sum_to_one_without_rounding']=True
assert len(data['table3'])==23
assert len(read(AUDIT/'table3_source_readings.json'))==23
results['all_23_table3_rows_preserved']=True
selected=ROOT/'regions/LME_052/models/52_1_Sea_of_Okhotsk_NE_(1980)/model.json'
assert sha(selected)=='f85c62fa8e4c57dcca72130aa649482d5716a87ac1ef3184324c2ca378ac2268'
results['existing_selected_model_hash_unchanged']=True
verification=dict(timestamp_utc=datetime.now(timezone.utc).isoformat(),artifact_integrity_passed=True,scientific_validity=False,source_interpretation_approved=False,native_model_ready=False,checks=checks,reconciliation=results,limits='Reconciliation verifies preservation and arithmetic, not the 33 tentative endpoint readings or biological consistency. No selected-model, map, DOCX or regional-workbook write operation is performed.')
write(AUDIT/'verification.json',verification)

# Portable evidence index. Do not recurse into the temporary node_modules junction.
items=[]
def add(role,p):items.append(dict(role=role,availability='present',path=Path(__import__('os').path.relpath(p,AUDIT)).as_posix(),sha256=sha(p)))
for s in source['sources']:add('source_pdf',PAPERS/s['file'])
for role,p in [('source_manifest',PAPERS/'source_manifest.json'),('source_use',PAPERS/'source_use.json'),('report',HERE/'REPORT.md'),('readiness_gate',HERE/'READINESS.json'),('review_workbook',HERE/'DC_reconstruction_review.xlsx'),('preserved_model',HERE/'model.json'),('direct_carbon_extraction',HERE/'carbon_reconstruction.json'),('illustrative_variant',HERE/'carbon_label_hypothesis.json'),('source_figure',AUDIT/'figure9_original.tif'),('source_arrow_ledger',AUDIT/'flow_readings.json'),('cell_ledger',AUDIT/'diet_cell_ledger.json'),('table3_ledger',AUDIT/'table3_source_readings.json'),('verification',AUDIT/'verification.json'),('conversion_audit',AUDIT/'converter_preservation_audit.json'),('import_validation',AUDIT/'import_validation.txt'),('massbalance_checks',AUDIT/'import_massbalance.txt'),('independent_consumption',AUDIT/'independent_consumption_comparison.json'),('budget_checks',AUDIT/'flow_budget_checks.json'),('reconstruction_script',HERE/'reconstruct.py'),('source_readings_script',HERE/'figure9_data.py'),('workbook_builder',AUDIT/'build_review.mjs'),('package_verifier',HERE/'finalize_package.py')]:add(role,p)
for filename in ['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv','TL.xlsx','Metadata.xlsx']:add('eight_import_files',TABLES/filename)
for p in sorted(AUDIT.glob('review_*.png')):add('workbook_visual_qa',p)
add('rejected_raw_converter',raw_path)
add('rejected_raw_roundtrip',next(TABLES.glob('*_reconstructed.xlsx')))
index=dict(schema_version=1,run_id='LME052-GM2019-FIG9-20261002',region_id='LME_052',model_id=HERE.name,variant_id='main_unresolved_labels_preserved',source_identity=dict(doi='10.26428/1606-9919-2019-198-143-163',figure=9,printed_page=157,pdf_page=15,sha256=source['sources'][0]['sha256']),computational_input_identity=dict(arrow_readings_sha256=sha(AUDIT/'flow_readings.json'),table3_sha256=sha(AUDIT/'table3_source_readings.json')),methods=['manual figure transcription with tentative routes retained','absent displayed feeding arrows set to zero by user instruction','consumer-column carbon-flow normalization','prey-specific wet/carbon conversion before wet normalization','independent prose consumption and budget comparison','lossy converter audit and faithful unknown preservation','CSV/JSON/saved-XLSX reconciliation'],required_roles=sorted({x['role'] for x in items}),artifacts=items,reconciliation=results,scientific_status='PROVISIONAL; not selected, balanced, or approved',out_of_scope=['taxonomy','SPPR','PPR','regional integration','map and validation DOCX update'])
write(AUDIT/'evidence_index.json',index)
spec=importlib.util.spec_from_file_location('check_evidence',ROOT/'tools/skills/original_skill_resources/combined-src/scripts/check_evidence.py')
checker=importlib.util.module_from_spec(spec);spec.loader.exec_module(checker)
integrity=checker.check(AUDIT/'evidence_index.json');write(AUDIT/'evidence_integrity.json',integrity)
assert integrity['complete'],integrity
print(json.dumps(dict(artifact_integrity_passed=True,evidence_files=len(integrity['verified_artifacts']),source_pdfs=6,groups=22,arrow_labels=78,absent_arrow_zero_cells=365,carbon_complete=19,wet_complete=16,illustrative_complete=20,import_checks=checks,native_model_ready=False,selected_model_unchanged=True),ensure_ascii=False,indent=2))
