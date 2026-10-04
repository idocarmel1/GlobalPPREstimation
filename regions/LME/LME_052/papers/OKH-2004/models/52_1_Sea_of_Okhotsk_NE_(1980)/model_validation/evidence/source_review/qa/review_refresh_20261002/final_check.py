from pathlib import Path
from collections import Counter
from urllib.parse import unquote,urlsplit
import hashlib,json,sys,zipfile,subprocess,tempfile,shutil,re
from lxml import etree as E
from docx import Document
from openpyxl import load_workbook
import pypdfium2 as pdfium

QA=Path(__file__).resolve().parent
ROOT=QA.parents[5]
REGION=ROOT/'regions/LME_052'
REPORT=REGION/'Model_validation_52_1_Sea_of_Okhotsk_NE_(1980).docx'
BASELINE=QA/('baseline_'+REPORT.name)
sys.path.insert(0,str(ROOT/'tools'))
from validation_percentage_format import verify_report
verify_report(REPORT)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
check=json.loads((QA/'verification.json').read_text(encoding='utf-8'))
for path,digest in check['protected_hashes'].items():assert sha(ROOT/path)==digest,path
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
N={'w':W}
with zipfile.ZipFile(BASELINE) as z:before={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(REPORT) as z:after={n:z.read(n) for n in z.namelist()}
assert before.keys()==after.keys()
assert [n for n in before if before[n]!=after[n]]==['word/document.xml']
b=E.fromstring(before['word/document.xml']);a=E.fromstring(after['word/document.xml'])
def txt(el):return ''.join(n.text or '' for n in el.iter('{'+W+'}t'))
oldrows=b.find('w:body/w:tbl',N).findall('w:tr',N)
newrows=a.find('w:body/w:tbl',N).findall('w:tr',N)
changed={'Catch source','Model extraction','GE diagnostics','TE diagnostics','Other'}
for x,y in zip(oldrows,newrows):
    label=txt(x.find('w:tc',N))
    if label not in changed:assert E.tostring(x)==E.tostring(y),label
for x,y in zip(b.findall('w:body/w:tbl',N)[1:],a.findall('w:body/w:tbl',N)[1:]):assert E.tostring(x)==E.tostring(y)
for x,y in zip(b.findall('w:body/w:p',N),a.findall('w:body/w:p',N)):assert E.tostring(x)==E.tostring(y)
styles=E.fromstring(after['word/styles.xml'])
style_link=styles.find("w:style[@w:styleId='Hyperlink']/w:rPr",N)
def visible_link(run):
    pr=run.find('w:rPr',N)
    color=pr.find('w:color',N) if pr is not None else None
    under=pr.find('w:u',N) if pr is not None else None
    if color is None and style_link is not None:color=style_link.find('w:color',N)
    if under is None and style_link is not None:under=style_link.find('w:u',N)
    assert color is not None and color.get('{'+W+'}val').lower() in ['0563c1','0000ff','0070c0'],txt(run)
    assert under is not None and under.get('{'+W+'}val') not in ['none','0'],txt(run)
for link in a.findall('.//w:hyperlink',N):
    for run in link.findall('w:r',N):
        if txt(run):visible_link(run)
props=E.fromstring(after['docProps/app.xml'])
assert all(not (n.text or '').strip() for n in props.iter() if E.QName(n).localname=='HyperlinkBase')

# Portable download verification: use a temporary repository root with identical relative files.
git='C:/Users/idoca/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/git/cmd/git.exe'
tracked=set(subprocess.run([git,'-C',str(ROOT),'ls-files','-z'],capture_output=True,check=True).stdout.decode('utf-8').split('\0'))
local_paths={x['relative_path'].replace('\\','/') for x in check['relative_links']}
new_evidence=local_paths-tracked
for path in new_evidence:
    # New retained evidence can live in a tracked directory before the user's next commit.
    assert any(t.startswith(str(Path(path).parents[1]).replace('\\','/')+'/') for t in tracked),path
    ignored=subprocess.run([git,'-C',str(ROOT),'check-ignore','--',path],capture_output=True)
    assert ignored.returncode==1,path
with tempfile.TemporaryDirectory(prefix='LME052_report_links_') as tmp:
    clone=Path(tmp)
    for path in local_paths|{str(REPORT.relative_to(ROOT)).replace('\\','/')}:
        dest=clone/path;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(ROOT/path,dest)
    cloned_report=clone/REPORT.relative_to(ROOT)
    d=Document(cloned_report)
    for r in d.part.rels.values():
        if 'hyperlink' not in r.reltype or urlsplit(r.target_ref).scheme in ['http','https']:continue
        rel=unquote(r.target_ref.split('#')[0]);p=cloned_report.parent/rel
        assert p.exists() and sha(p)==sha(REGION/rel),r.target_ref
    check['physical_relocation']={'copied_repository_files':len(local_paths)+1,'all_local_word_links_resolve_with_identical_bytes':True,'tracked_target_files':len(local_paths&tracked),'new_retained_evidence_in_tracked_directories':sorted(new_evidence),'new_evidence_ignored':False}

wb=load_workbook(REGION/'LME052_taxon_mapping_appendix.xlsx')
excel_link_count=0
for sheet in wb:
    for row in sheet:
        for c in row:
            if not c.hyperlink:continue
            excel_link_count+=1
            assert c.font.underline in ['single','double']
            assert c.font.color and c.font.color.type=='rgb' and c.font.color.rgb[-6:].lower() in ['0563c1','0000ff','0070c0']
            target=c.hyperlink.target
            if not target or target.startswith('#') or urlsplit(target).scheme in ['http','https']:continue
            p=REGION/unquote(target.split('#')[0]);assert p.exists(),target
            assert not re.match(r'(?i)^(?:[a-z]:|file:|\\\\|/)',target),target
check['hyperlink_appearance']={'word_all_visible_links_blue_and_underlined':True,'excel_links_checked':excel_link_count}
doc=pdfium.PdfDocument(QA/'render/validation.pdf')
changed_pages=[]
for i in range(len(doc)):
    p=QA/'render'/f'page-{i+1}.png';previous=sha(p) if p.exists() else None
    doc[i].render(scale=2).to_pil().save(p)
    if previous!=sha(p):changed_pages.append(i+1)
assert len(doc)==8
check['report_sha256']=sha(REPORT)
check['visual_review']={'status':'awaiting_reinspection_of_changed_pages','page_count':8,'changed_pages_after_link_spacing_fix':changed_pages,'renderer':'Independent hidden Microsoft Word read-only export; bundled PDFium rasterization','packaged_renderer_limitation':'Bundled LibreOffice unavailable'}
(QA/'verification.json').write_text(json.dumps(check,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'page_count':len(doc),'changed_pages':changed_pages,'protected_files_unchanged':len(check['protected_hashes']),'manual_and_other_unedited_content_preserved':True,'portable_local_targets':len(local_paths),'excel_links_checked':excel_link_count,'percentage_check':'passed','report_sha256':check['report_sha256']},indent=2))
