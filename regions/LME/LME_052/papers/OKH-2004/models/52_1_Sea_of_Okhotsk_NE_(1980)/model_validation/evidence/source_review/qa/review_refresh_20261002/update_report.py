from pathlib import Path
from copy import deepcopy
from collections import Counter, defaultdict
from urllib.parse import unquote, urlsplit
from xml.parsers import expat
from xml.sax.saxutils import escape
import sys, json, hashlib, math, re, shutil, zipfile
from lxml import etree as E
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[6]
sys.path.insert(0, str(ROOT / 'tools'))
from workbooks import read_book, records, overview
from validation_percentage_format import format_percent, verify_report

QA = Path(__file__).resolve().parent
REGION = ROOT / 'regions/LME_052'
MODEL = '52_1_Sea_of_Okhotsk_NE_(1980)'
BUNDLE = REGION / 'validation_reports' / MODEL
REPORT = REGION / f'Model_validation_{MODEL}.docx'
BASELINE = QA / ('baseline_' + REPORT.name)
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p): return json.loads(p.read_text(encoding='utf-8'))
protected = [REGION/'LME_052.xlsx', REGION/'LME052_taxon_mapping_appendix.xlsx',
             REGION/'models'/MODEL/'model.json', ROOT/'Project.xlsx',
             ROOT/'tools/templates/Model_validation_template.docx',
             ROOT/'tools/knowledge_graph/graph.json',
             *[ROOT/'interactive_map'/p for p in ['index.html','time_series.html','source_archive.html'] if (ROOT/'interactive_map'/p).exists()]]
before = {str(p.relative_to(ROOT)): sha(p) for p in protected}
if BASELINE.exists():
    assert sha(BASELINE)==sha(REPORT), 'The current document differs from this refresh baseline.'
else:
    shutil.copy2(REPORT, BASELINE)
book = read_book(REGION/'LME_052.xlsx')
assert overview(book)['selected_model_id'] == MODEL
audit = read(BUNDLE/'taxon_audit.json')
summary = read(BUNDLE/'coverage_summary.json')
assert len(audit) == 151 == len({x['taxon'] for x in audit})
catch_rows = [x for x in records(book,'Catch','Catch') if x['catch_basis']=='landings']
catch = {x['taxon']: x for x in catch_rows}
assert len(catch_rows)==len(catch), 'Duplicate landings taxon labels.'
classic = {x['taxon']: x for x in records(book,'Classic PPR','Taxa')}
assert set(catch)=={x['taxon'] for x in audit}
matching = defaultdict(list)
for x in records(book,'PPR','Matching'):
    if x['model_id']==MODEL: matching[x['taxon']].append(x)
for x in audit:
    t=x['taxon']; c=catch[t][2019]; coef=classic.get(t,{}).get('sppr')
    v=0 if c==0 else c*coef/9 if coef is not None else None
    assert math.isclose(c,x['catch_tonnes'],rel_tol=1e-12,abs_tol=1e-8)
    assert v is not None and math.isclose(v,x['simple_chain_ppr_tC'],rel_tol=1e-12,abs_tol=1e-7)
    expected=dict(zip(x['group_names'],x['weights']))
    actual={m['group']:m['weight'] for m in matching[t] if m['group']}
    assert actual.keys()==expected.keys(), t
    for k in actual: assert math.isclose(actual[k],expected[k],rel_tol=1e-12,abs_tol=1e-14)
    assert all(m['confidence'].lower().replace('_',' ')==x['overall_confidence'].lower() for m in matching[t]), t
totals = {'catch':sum(x['catch_tonnes'] for x in audit),'ppr':sum(x['simple_chain_ppr_tC'] for x in audit)}
assert math.isclose(totals['ppr'],summary['total_simple_chain_ppr_tC'],rel_tol=1e-12)
assert math.isclose(totals['catch'],summary['total_catch_tonnes'],rel_tol=1e-12)
for categories,key in [('confidence_summary','overall_confidence'),('membership_summary','membership_rule'),('allocation_summary','allocation_rule')]:
    for s in summary[categories]:
        xs=[x for x in audit if x[key]==s.get('label',s.get('rule')) and (categories=='confidence_summary' or x['membership_confidence' if categories=='membership_summary' else 'allocation_confidence']==s['confidence'])]
        assert math.isclose(sum(x['simple_chain_ppr_tC'] for x in xs),s['ppr_tC'],rel_tol=1e-12,abs_tol=1e-7)
    assert abs(sum(x['ppr_percentage'] for x in summary[categories])-100)<1e-8

runtime = read(BUNDLE/'diet_reextraction_20261002/scientific_revalidation.json')
assert runtime['canonical_sha256']==sha(REGION/'models'/MODEL/'model.json')
diagnostics = read(BUNDLE/'diet_reextraction_20261002/direct_reports.json')
solutions = read(BUNDLE/'diet_reextraction_20261002/direct_solutions.json')
negatives = {}
for method in ['GE','TE']:
    m=solutions[method]['SPPR']; assert len(m['index'])==len(set(m['index']))==30
    assert len(m['columns'])==len(set(m['columns']))==4
    assert len(m['data'])==30 and all(len(r)==4 for r in m['data'])
    pairs=[(sid,gid,value) for gid,row in zip(m['index'],m['data']) for sid,value in zip(m['columns'],row) if math.isfinite(value) and value<0]
    assert all(math.isfinite(v) for row in m['data'] for v in row)
    assert len({x[0] for x in pairs})==diagnostics[method]['divergence']['n_negative_sources']
    negatives[method]=pairs
assert negatives['GE']==[]
assert {(s,g) for s,g,v in negatives['TE']}=={(s,g) for s in [29,28,27] for g in [14,12,4,2]}

appendix = load_workbook(REGION/'LME052_taxon_mapping_appendix.xlsx',data_only=False)
sheet=appendix['Taxon mapping']; rows=list(sheet.values)[7:]
assert len(rows)==151
assert list(sheet.values)[6]==('Taxon name','TL','Catch (t)','Simple-chain PPR (t C)','Mapped group names and weights','Confidence level','Reason')
assert sheet.freeze_panes and any(t.autoFilter and t.autoFilter.ref for t in sheet.tables.values())
assert 'Sources' in appendix.sheetnames
ordered=sorted(audit,key=lambda x:(-x['simple_chain_ppr_tC'],x['taxon']))
assert [r[0] for r in rows]==[x['taxon'] for x in ordered]
for r,x in zip(rows,ordered):
    assert r[5]==x['overall_confidence'] and r[4]==x['mapping_display'] and x['reason'] in r[6]
    assert math.isclose(r[2],x['catch_tonnes'],rel_tol=1e-12,abs_tol=1e-9)
    assert math.isclose(r[3],x['simple_chain_ppr_tC'],rel_tol=1e-12,abs_tol=1e-7)

with zipfile.ZipFile(BASELINE) as z: parts=[(i,z.read(i.filename)) for i in z.infolist()]
xml=dict((i.filename,v) for i,v in parts)['word/document.xml']
tree=E.fromstring(xml)
main=tree.find('w:body/w:tbl',NS)
manual_labels=['SPPR calculation','Open issues and next action','Review and reproducibility','Selection rationale']
def text(el): return ''.join(n.text or '' for n in el.iter('{'+W+'}t'))
def cell_label(row): return text(row.find('w:tc',NS))
manual_before={cell_label(r):E.tostring(r) for r in main.findall('w:tr',NS) if cell_label(r) in manual_labels}

# Bounded row replacement: unchanged XML bytes, Office parts, styles and figures are retained.
changes={
 'Catch source':[
   'Sea Around Us, 1950–2019 reconstructed catch; reference: 2019 landings.',
 ],
 'Model extraction':[
   'Related-chapter Tables 4a/b supply 174 printed diet values and 580 blanks. The canonical diet retains the printed proportions; eight consumer sums are 0.998–1.001. Their cause remains undetermined. Normalization is applied only in the calculation runtime.',
   'Supplied Table 1 B/PB/QB/EE values agree with the canonical inputs. Native group catches and the pollock cutoff are absent; export zeros, detritus B = 1 and inherited defaults/balance remain assumptions.',
 ],
 'GE diagnostics':[
   'WARN — living network near divergence; EE = 0 in Seabirds (5) and Sperm whales (3).',
   'No negative SPPR entries across the basal-source columns.',
   'rho_living = 0.8672', 'b = 0.1520', 'Detritus (29): SPPR = 1.251',
   'Model and system balance checks pass.',
 ],
 'TE diagnostics':[
   'FAIL — living network diverges (rho_living > 1).',
   'Negative SPPR: Detritus (29) → Cottidae (14), Other Gadidae (12), Pinnipeds (4), Toothed whales (2).',
   'Phytoplankton (28) → Cottidae (14), Other Gadidae (12), Pinnipeds (4), Toothed whales (2).',
   'Phytobenthos (27) → Cottidae (14), Other Gadidae (12), Pinnipeds (4), Toothed whales (2).',
   'rho_living = 1.0251', 'Detritus (29): SPPR = 0.9486',
   'b = 0 is the TE convention.',
   'Model and system balance checks pass; living-network divergence remains.',
 ],
 'Other':[
   'Diagnostics use the restored printed diet with runtime normalization; native group catch is unavailable. Whether the diet discrepancies reflect rounding or other errors remains for researcher review. Runtime normalization is always allowed.',
   'With Egestion remains WARN. Near-singular Monte Carlo TE and negative/unsupported methods remain historical; failed TE annual PPR is unavailable.',
 ],
}
parser=expat.ParserCreate(namespace_separator='}'); spans=[];active=[]
def start(name,attrs):
    if name==W+'}tr':active.append(parser.CurrentByteIndex)
def end(name):
    if name==W+'}tr':
        lo=active.pop();hi=xml.index(b'>',parser.CurrentByteIndex)+1;spans.append((lo,hi))
parser.StartElementHandler=start;parser.EndElementHandler=end;parser.Parse(xml,True)
row_elems=tree.findall('.//w:tr',NS); spans.sort()
assert len(row_elems)==len(spans)
patches=[]
for row,(lo,hi) in zip(row_elems,spans):
    label=cell_label(row)
    if row.getparent()!=main or label not in changes: continue
    revised=deepcopy(row);cell=revised.findall('w:tc',NS)[1]; ps=cell.findall('w:p',NS)
    # Preserve all hyperlink-bearing paragraphs and their complete anchors.
    linked=[deepcopy(p) for p in ps if p.find('w:hyperlink',NS) is not None]
    for p in linked:
        for child in list(p):
            if child.tag not in ['{'+W+'}pPr','{'+W+'}hyperlink']:p.remove(child)
        hyperlinks=p.findall('w:hyperlink',NS)
        for h in hyperlinks[1:]:
            sep=E.Element('{'+W+'}r');st=E.SubElement(sep,'{'+W+'}t');st.text=' | '
            st.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
            p.insert(p.index(h),sep)
    seed=deepcopy(ps[0])
    for p in ps:cell.remove(p)
    for value in changes[label]:
        p=deepcopy(seed)
        for item in list(p):
            if item.tag!='{'+W+'}pPr':p.remove(item)
        rr=E.SubElement(p,'{'+W+'}r')
        oldrun=seed.find('w:r',NS)
        if oldrun is not None and oldrun.find('w:rPr',NS) is not None:rr.append(deepcopy(oldrun.find('w:rPr',NS)))
        tt=E.SubElement(rr,'{'+W+'}t');tt.text=value
        cell.append(p)
    for p in linked:cell.append(p)
    # Serialise only the changed row; root namespaces keep unrelated bytes intact.
    patches.append((lo,hi,E.tostring(revised,encoding='utf-8',with_tail=False),label))
for lo,hi,new,label in sorted(patches,reverse=True):xml=xml[:lo]+new+xml[hi:]
tmp=QA/'updated.docx'
with zipfile.ZipFile(tmp,'w') as z:
    for i,v in parts:z.writestr(i,xml if i.filename=='word/document.xml' else v)
verify_report(tmp)
newdoc=Document(tmp);newmain=newdoc.tables[0]._element
manual_after={cell_label(r):E.tostring(r) for r in newmain.findall('w:tr',NS) if cell_label(r) in manual_labels}
assert manual_before==manual_after, 'Manual cells or selection rationale changed.'
with zipfile.ZipFile(tmp) as z:
    changed=[i.filename for i,v in parts if z.read(i.filename)!=v]
assert changed==['word/document.xml']
with zipfile.ZipFile(tmp) as z:
    finalxml=z.read('word/document.xml')
for label in manual_labels:
    oldraw=[xml0 for (lo,hi),r in zip(spans,row_elems) if cell_label(r)==label for xml0 in [dict((i.filename,v) for i,v in parts)['word/document.xml'][lo:hi]]][0]
    assert oldraw in finalxml, label

# Verify coverage tables against exact unrounded evidence and complete Very low enumeration.
for row,s in zip(newdoc.tables[1].rows[1:],summary['confidence_summary']):
    assert [c.text for c in row.cells]==[s['label'],str(s['taxa']),format_percent(s['catch_percentage']),format_percent(s['ppr_percentage'])]
for ti,key in [(2,'membership_summary'),(3,'allocation_summary')]:
    for row,s in zip(newdoc.tables[ti].rows[1:],summary[key]):assert row.cells[2].text==format_percent(s['ppr_percentage'])
low=[t.strip() for row in newdoc.tables[4].rows[1:] for t in row.cells[0].text.split(';')]
assert len(low)==len(set(low))==46 and set(low)=={x['taxon'] for x in audit if x['overall_confidence']=='Very low'}
assert [s['ppr_percentage'] for s in summary['membership_summary']]==sorted([s['ppr_percentage'] for s in summary['membership_summary']],reverse=True)

links=[]
for rel in newdoc.part.rels.values():
    if 'hyperlink' not in rel.reltype:continue
    target=rel.target_ref
    if urlsplit(target).scheme in ['http','https']:continue
    assert not re.match(r'(?i)^(?:[a-z]:|file:|\\\\|/)',target),target
    path=(REGION/unquote(target.split('#')[0])).resolve();assert path.exists(),target
    assert path.is_relative_to(ROOT),target
    links.append({'target':target,'relative_path':str(path.relative_to(ROOT)),'sha256':sha(path) if path.is_file() else None})
assert all(sha(p)==before[str(p.relative_to(ROOT))] for p in protected)
shutil.copy2(tmp,REPORT)
assert all(sha(p)==before[str(p.relative_to(ROOT))] for p in protected)
evidence={
 'region': 'LME_052','model_id':MODEL,'reference_year':2019,'basis':'landings',
 'baseline_sha256':sha(BASELINE),'report_sha256':sha(REPORT),
 'changed_zip_parts':changed,'updated_rows':[p[3] for p in patches],
 'manual_rows_and_selection_rationale_byte_preserved':True,
 'template_layout_styles_and_figures_preserved':True,
 'counts_and_total':totals,'taxa':151,'very_low_taxa_once':46,
 'confidence_and_rule_shares_reconciled':True,'word_percentage_check':True,
 'appendix_all_rows_and_order_verified':True,'appendix_unchanged':True,
 'diagnostic_identity_verified':True,'negative_source_group_pairs':negatives,
 'relative_links':links,'protected_hashes':before,
 'scientific_runs_performed':False,'map_updated':False,
 'visual_review':{'status':'pending'}
}
(QA/'verification.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in evidence.items() if k not in ['relative_links','protected_hashes','negative_source_group_pairs']},ensure_ascii=False,indent=2))
