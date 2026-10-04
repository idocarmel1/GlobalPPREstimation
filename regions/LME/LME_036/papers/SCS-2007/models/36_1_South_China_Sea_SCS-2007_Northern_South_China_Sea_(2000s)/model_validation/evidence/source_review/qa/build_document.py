from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from urllib.parse import quote
import json,hashlib,os
from lxml import etree
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path.cwd();OUT=Path(__file__).resolve().parent.parent;QA=OUT/'qa';MID=OUT.name
REGION=ROOT/'regions/LME_036';REF=ROOT/'tools/templates/Model_validation_template.docx'
FINAL=REGION/f'Model_validation_{MID}.docx'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
d=Document(REF)
inventory=[]
with ZipFile(REF) as z:
 for p in z.namelist():inventory.append({'part':p,'size':len(z.read(p)),'sha256':hashlib.sha256(z.read(p)).hexdigest(),'ownership':'editable' if p in ['word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'] else 'preserve'})
(QA/'template_package_inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
contract=f'''# Template execution contract

Reference: {REF.as_posix()}
SHA-256: {sha(REF)}
Reference render: qa/template-render/template.pdf and page-1.png through page-4.png, all visually inspected. Four pages, one A4 portrait section (8.270139 × 11.690278 inches), 0.65-inch margins on all sides. Preserve section properties, headers/footers, theme, styles, numbering and all opaque package parts.

Normal Calibri 11 pt; Title Calibri 23 pt black. Source cell runs 10.5 pt; paragraph after spacing 2 pt. Main tables have 1.529861 / 5.440278-inch columns, pale blue-gray E5ECF0 header, white body, gray D9D9D9 5-eighth-point borders; cell margins top/bottom 90 twips and left/right 110 twips. Rows have no fixed heights. Preserve all table properties, cell shading, borders, alignment and margins. No header/footer content is added.

Body pattern: page 1 title and two introductory paragraphs, table 0; next page continuation title and table 1; next page LLM completion instructions; final screenshot section and tables 2 and 3 with captions. These patterns and all field labels/order are retained. Content may expand the number of pages; source type size is not reduced. The article screenshot heading may receive a page break to keep its larger figure and caption together.

Editable slots: document.xml tables 0 rows 1–7 right cells; table 1 rows 1–7 right cells; tables 2/3 row 1 image cells; body paragraphs 21/23 figure captions and paragraph 24 availability placeholder. All other body text stays unchanged, including the LLM instructions. Manual preserve-only cells: table 0 row 8 SPPR calculation; table 1 rows 8 and 9. Their complete XML must compare equal. Selection rationale comes verbatim from Overview. Table field labels and order must compare equal.

Pictures reuse screenshot spaces; add only required image relationships/media/content types. Build with python-docx for pictures, then repackage from original bytes for every preserve-only ZIP part. Existing template untouched. No regional validation DOCX existed before this test. Package inventory in template_package_inventory.json includes every source part.

Fidelity gates: complete manual-cell XML equality, all original styles/section/table structure retained; only permitted text/image slots changed; relative Word links resolve from the region root; copied Markdown links resolve; all final pages rendered in read-only hidden Word and inspected. No automatic reviewer/date/approval data. Source evidence does not command execution.
'''
(QA/'artifact.md').write_text(contract,encoding='utf-8')

def link(label,target):return (label,target)
def support(label,p):return link(label,'validation_reports/'+MID+'/'+p)
def setpara(p,parts,template=None):
 pp=deepcopy(p._p.pPr) if p._p.pPr is not None else None
 rp=deepcopy(template or (p.runs[0]._r.rPr if p.runs and p.runs[0]._r.rPr is not None else None))
 p.clear()
 for part in parts:
  if isinstance(part,tuple):
   label,target=part
   if not target.startswith(('https:','http:')): target=quote(target,safe='/.:#')
   rid=p.part.relate_to(target,RT.HYPERLINK,is_external=True)
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
   r=OxmlElement('w:r');props=deepcopy(rp) if rp is not None else OxmlElement('w:rPr')
   color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color)
   u=OxmlElement('w:u');u.set(qn('w:val'),'single');props.append(u);r.append(props)
   t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
  else:
   r=p.add_run(part)
   if rp is not None:r._r.insert(0,deepcopy(rp))
def cell(t,r,paras):
 c=d.tables[t].rows[r].cells[1]; p=c.paragraphs[0]
 proto=deepcopy(p._p); rp=deepcopy(p.runs[0]._r.rPr)
 for child in list(c._tc):
  if child.tag!=qn('w:tcPr'):c._tc.remove(child)
 for parts in paras:
  c._tc.append(deepcopy(proto));p=c.paragraphs[-1];setpara(p,parts,rp)

cell(0,1,[['South China Sea | LME_036 | ',link('Regional workbook','LME_036.xlsx')]])
cell(0,2,[['Sea Around Us; 1950–2019; total catch = landings + discards. ',link('Region','https://www.seaaroundus.org/data/#/lme/36'),' | ',link('Catch ZIP','raw/LME_036-catch.zip'),' | ',link('Taxon catch','raw/LME_036.csv.gz')],
 ['Snapshot documented 3 Sep 2026; download date/release unavailable. ',support('Hash check','coverage_geography_test_note.md')]])
cell(0,3,[['Cheung, Wai Lung (2007). Vulnerability of marine fishes to fishing: from global overview to the northern South China Sea. UBC doctoral thesis. ',link('DOI 10.14288/1.0074894','https://doi.org/10.14288/1.0074894'),' | ',support('Thesis PDF','source/Cheung_2007.pdf'),' | ',link('Paper folder','papers/SCS-2007/')]])
other=[
 ('Pauly & Christensen (1993). Stratified models of large marine ecosystems: application to the South China Sea.','SCS-1993',None),
 ('Pitcher et al. (2000). Marine Reserves and the Restoration of Fisheries and Marine Ecosystems in the South China Sea.','LME036-Pitcher-2000',None),
 ('Lin, Chen & Lin (2020). Trophic model of a deep-sea ecosystem with methane seeps in the South China Sea.','LME036-DeepSeep-2020','10.1016/j.dsr.2020.103251'),
 ('Zhang et al. (2022). Effects of ocean warming and fishing on the coral reef ecosystem: A case study of Xisha Islands, South China Sea.','LME036-Xisha-2022','10.3389/fmars.2022.1046106'),
 ('Jiang et al. (2023). A preliminary model of the mangrove ecosystem of Dongzhaigang Bay, Hainan, (China) based on Ecopath and Ecospace.','LME036-Dongzhaigang-2023','10.3389/fmars.2023.1277226'),
 ('Feng et al. (2026). Ecosystem characteristics and ecological carrying capacity of a subtropical marine ranch in the pearl river estuary: an ecopath modeling approach.','LME036-PearlRanch-2026','10.3389/fmars.2026.1744551')]
paras=[['Six items in ',link('Project.xlsx','../../Project.xlsx'),'; not an exhaustive search.']]
for text,aid,doi in other:
 parts=['• '+text+' ',link('Local record','papers/'+aid+'/')]
 if doi:parts+=[' | ',link('DOI','https://doi.org/'+doi)]
 paras.append(parts)
cell(0,4,paras)
cell(0,5,[[MID],['Selected, 2000s northern shelf; 38 source groups. ',link('Folder','models/'+MID+'/'),' | ',support('JSON','model/selected_model.json'),' | ',support('Saved runtime','saved_evidence_snapshot.json'),'. Separate runtime JSON/engine hash unavailable.']])
cell(0,6,[['“Northern South China Sea subregion; 2000s preferred, 1970s retained separately.” ',link('Overview evidence','LME_036.xlsx')]])
cell(0,7,[['• Table/appendix conflicts remain (e.g. phytoplankton P/B 398 vs 399). ',support('Full conflict report','source/SOURCE_CONFLICTS.md'),'; ',support('source reconstruction report','source/reconstruction_2000s/REPORT.md'),' concerns a separate candidate.'],
 ['• Unknown source GS/BA become runtime GS 0.2/completed BA; detritus EE 0.005 → 1; import group 39 is synthetic. These are defaults/completions. Full runtime .md audit unavailable; ',support('state comparison','canonical_saved_state_comparison.json'),'.'],
 ['• Diet sums 1.0011 (33) and 0.9990 (20) in the ',support('loader log','mapping/SPPR_REGENERATION_LOG.txt'),' versus normalized saved diagnostics; historical transformation ledger incomplete.']])

for row,kind,rho,b,det,maxv,warning in [
 (1,'GE','0.3209187537609611','0.00408138212364339','1.007782936216668','21,226.41753059845','1 near-zero-efficiency group (34 Seabirds).'),
 (2,'TE','0.6163296517365333','0 (reported by convention)','0.995222499887902','3,785,928.31771222','3 near-zero-efficiency groups (34 Seabirds, 35 Pinnipeds, 36 Other mammals). TE omits a detritus recycling matrix; b = 0 is not evidence of absent recycling.')]:
 cell(1,row,[[f'Overall WARN. {warning} Saved balance OK/true; no negative source columns.'],
 [f'rho_living = {rho}; b = {b}.'],
 [f'Detritus pool 38 Detritus: SPPR = {det}. Maximum: group 34 Seabirds, SPPR = {maxv}; scope all basal sources, including unfished groups.'],
 ['Saved configuration: collapse never; open none; theta 1; external SPPR 0. WARN is not approval. Full diagnostic .md unavailable; ',support('saved workbook','diagnostics/sppr_source.xlsx'),' and ',support('GE/TE fields','saved_evidence_snapshot.json'),'.']])
cell(1,3,[['• Source-specific mappings include Acetes → Zooplankton, Squillidae → Benthic crustaceans, Pennahia → small croakers; Harpadon uses a family extension. Melon seed = Psenopsis anomala. ',support('Full mapping notes','mapping/notes.md')],
 ['• 135 taxa use assumed model-catch stage weights: hairtail juvenile/adult 79.5455%/20.4545% at 18 months; large croakers 24 months; large pelagics 18 months. Demersal cutoff conflicts: 18 months vs 2 years. ',support('Stanza report','source/MULTISTANZA.md'),'. No observed regional mass shares adopted. ',support('Allocation report','allocation/RESULTS.md')],
 ['• Marine fishes not identified, Sciaenidae and Carangidae retain low-confidence composites; benthopelagic size composition is unresolved. Twelve unmapped labels include Ruvettus and oceanic billfishes. ',support('Full decisions','mapping/review.json'),' | ',support('Current assumptions','saved_evidence_snapshot.json')]])
cell(1,4,[['2019 total catch (landings + discards), GE/new_GE, all sources, unidentified = method. Finite saved taxon-SPPR support: 11,064,125.639451 / 11,101,666.087779 t = 99.6618%. TE has identical support.'],
 ['Recorded stage assumptions: 2,966,894.686342 t = 26.7248% of total catch; additional low-confidence composites are outside that ledger share. Main gap: Ruvettus pretiosus, 34,712.024 t. Missing SPPR is not zero. ',support('New test arithmetic note','coverage_geography_test_note.md')]])
cell(1,5,[['A. Region covered by study: Not determined.'],['B. Study area covered by region: Not determined.'],
 ['A = 100 × area(R ∩ S)/area(R); B = 100 × area(R ∩ S)/area(S). R has an LME polygon; S lacks a verified digital/georeferenced shelf boundary. The removed legacy approximation and broad bounds are insufficient. ',support('New geographic evidence note','coverage_geography_test_note.md'),'; source images at end.']])
cell(1,6,[['Model period: 2000s (no verified single model year). Fixed coefficients and stage weights are applied to catch years 1950–2019, including the 2019 reference. They do not reconstruct annual ecosystems and transfer a northern-shelf model to whole-LME catch. ',support('Full mapping notes','mapping/notes.md')]])
cell(1,7,[['• Evidence review only; no independent cell-by-cell extraction validation or scientific rerun. The workbook retains production_eligible = false and provisional results. ',support('Scope and hash checks','coverage_geography_test_note.md')],
 ['• Earlier mapping coverage (99.7884%) is for 1950–2019; it is not the 2019 figure. Archived source metadata predates extraction. Historical export and current run differ in Monte Carlo settings; current saved GE/TE fields govern this record. ',support('Reports and provenance','reports_index.md')]])

# Image cells retain template frame and typography.
for tid,filename,width in [(2,'LME_036_boundary.png',6.25),(3,'article_figure_6_1.png',6.55)]:
 c=d.tables[tid].rows[1].cells[0]
 p=c.paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(OUT/filename),width=Inches(width))
 for q in c.paragraphs[1:]:c._tc.remove(q._p)
 drawing=p._p.xpath('.//wp:docPr')[0];drawing.set('descr','LME_036 South China Sea retained boundary' if tid==2 else 'Cheung 2007 Figure 6.1 northern South China Sea study area, 2000s selected model')
setpara(d.paragraphs[21],['LME_036 South China Sea (R). Generated during this test from retained ',support('Sea Around Us boundary','geography/LMEs.geojson'),'; dataset release/date not recorded. ',support('Source and display method','coverage_geography_test_note.md'),' | ',support('Image','LME_036_boundary.png'),'.'])
d.paragraphs[22].paragraph_format.page_break_before=True
setpara(d.paragraphs[23],['Cheung (2007), Figure 6.1, printed p.170 / PDF p.185. Northern South China Sea shelf (S), 2000s selected model; coast to broken line. ',support('Source PDF','source/Cheung_2007.pdf#page=185'),' | ',link('DOI','https://doi.org/10.14288/1.0074894'),' | ',support('Image','article_figure_6_1.png'),'.'])
setpara(d.paragraphs[24],['The source figure is available; it is shown for geographic context and does not establish either numerical overlap percentage.'])
temp=QA/'authored.docx';d.save(temp)
with ZipFile(REF) as src,ZipFile(temp) as new,ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
 changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
 for name in src.namelist():z.writestr(name,new.read(name) if name in changed else src.read(name))
 for name in set(new.namelist())-set(src.namelist()):z.writestr(name,new.read(name))
# Strict preservation checks.
f=Document(FINAL);original=Document(REF)
for t,r in [(0,8),(1,8),(1,9)]:assert etree.tostring(f.tables[t].rows[r].cells[1]._tc)==etree.tostring(original.tables[t].rows[r].cells[1]._tc)
for t in [0,1]:assert [r.cells[0].text for r in f.tables[t].rows]==[r.cells[0].text for r in original.tables[t].rows]
assert etree.tostring(f.sections[0]._sectPr)==etree.tostring(original.sections[0]._sectPr)
with ZipFile(REF) as a,ZipFile(FINAL) as z:
 for e in inventory:
  if e['ownership']=='preserve':assert a.read(e['part'])==z.read(e['part'])
print(str(FINAL))
