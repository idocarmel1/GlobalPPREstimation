from pathlib import Path
import sys,json,math,hashlib,zipfile,collections
from urllib.parse import unquote,urlsplit
from lxml import etree as E
import numpy as np
import openpyxl
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'Project.xlsx').exists())
R=Path(__file__).resolve().parent.parent; region=R.parent.parent
sys.path.insert(0,str(ROOT/'tools'))
from workbooks import read_book,records,overview,sha
from regional import result_hash
load=lambda p:json.loads(p.read_text('utf-8'))
audit=load(R/'mapping/adopted_taxon_audit.json');summary=load(R/'mapping/coverage_summary.json');old=load(R/'baseline/snapshot.json')
checks={};detail={}
def check(name,value):
 checks[name]=bool(value)
 if not value:raise AssertionError(name)
def equivalent(a,b):return json.loads(json.dumps(a,ensure_ascii=False))==json.loads(json.dumps(b,ensure_ascii=False))
b=read_book(region/'LME_032.xlsx');o=overview(b)
check('selected model identity',o['selected_model_id']=='32_1_Arabian_Sea_off_Karnataka_(2000)')
check('regional result freshness',o['calculation_result_sha256']==result_hash(b))
catch=records(b,'Catch','Catch');classic=records(b,'Classic PPR','Taxa')
check('original catch unchanged',equivalent(catch,old['catch']))
check('classic taxon coefficients unchanged',equivalent(classic,old['classic']))
key=lambda r:tuple(r.get(k) or '' for k in ['model_id','scope','method','catch_basis','unidentified','metric'])
previous={key(r):r for r in old['classic_annual']};current={key(r):r for r in records(b,'Classic PPR','Annual')}
removed=set(previous)-set(current)
check('historical sensitivity bounds explicitly invalidated',all(k[-1] in ['min_tC','max_tC'] for k in removed) and not set(current)-set(previous))
max_abs=0.;max_rel=0.
for k,r in current.items():
 for year in range(1950,2020):
  a=r.get(year);v=previous[k].get(str(year))
  if a is None or v is None:check('classic missing '+str(k)+' '+str(year),a==v);continue
  max_abs=max(max_abs,abs(a-v));max_rel=max(max_rel,abs(a-v)/max(abs(a),abs(v),1.))
  check('classic series '+str(k)+' '+str(year),math.isclose(a,v,rel_tol=1e-12,abs_tol=1e-9))
detail['classic_comparison']={'maximum_absolute_wet_tonne_difference':max_abs,'maximum_relative_difference':max_rel,'removed_historical_bounds':list(removed),'reason':'Current regional recalculate explicitly invalidates historical sensitivity bounds; numeric differences are floating-point summation/write precision, not changed source coefficients or catches.'}
matching=records(b,'PPR','Matching');bytax=collections.defaultdict(list)
for row in matching:bytax[row['taxon']].append(row)
check('all 430 labels represented once in review',len(audit)==len({x['taxon'] for x in audit})==430)
check('adopted mapping taxon set',set(bytax)=={x['taxon'] for x in audit})
for x in audit:
 actual=bytax[x['taxon']]
 check('adopted mapping '+x['taxon'],{r['group'] for r in actual}=={g['group'] for g in x['groups']} and all(r['confidence']==x['overall_confidence'].lower() for r in actual) and all(math.isclose(r['weight'],next(g['weight'] for g in x['groups'] if g['group']==r['group']),abs_tol=1e-12) for r in actual))
 check('weights '+x['taxon'],math.isclose(sum(r['weight'] for r in actual),1.,abs_tol=1e-9))
rank={'High':4,'Medium':3,'Low':2,'Very low':1,'Unresolved':0}
check('weakest necessary component',all(rank[x['overall_confidence']]==min(rank[x['membership_confidence']],rank[x['allocation_confidence']]) for x in audit))
check('independent assumption fields',all(isinstance(x.get('assumed'),bool) and isinstance(x.get('assumption_type'),list) for x in audit))
check('complete zero-catch missing policy',sum(x['missing_coefficient'] for x in audit)==37 and all(x['catch_t']==x['simple_ppr_tC']==0 for x in audit if x['missing_coefficient']))
check('simple PPR one carbon conversion',all(x['missing_coefficient'] or math.isclose(x['simple_ppr_tC'],x['catch_t']*x['classic_sppr']/9,rel_tol=1e-13) for x in audit))
for key in ['catch_pct','ppr_pct']:check('coverage partition '+key,math.isclose(sum(x[key] for x in summary['coverage']),100,abs_tol=1e-8))
for key in ['membership_rules','allocation_rules']:check('rule partition '+key,math.isclose(sum(x['ppr_pct'] for x in summary[key]),100,abs_tol=1e-8))
check('membership rules descending',all(a['ppr_pct']>=c['ppr_pct'] for a,c in zip(summary['membership_rules'],summary['membership_rules'][1:])))
check('catch total',math.isclose(sum(x['catch_t'] for x in audit),summary['catch_t'],rel_tol=1e-13))
check('PPR total',math.isclose(sum(x['simple_ppr_tC'] for x in audit),summary['ppr_tC'],rel_tol=1e-13))
rows=records(b,'Selected model groups','Group SPPR')
detail['group_sppr_headers']=list(rows[0]);detail['mapping_rows']=len(matching);detail['allocation_candidates']=len(records(b,'PPR','Allocation assumptions'))
matrix_detail={}
for method in ['GE','TE','With Egestion']:
 d=load(R/f'diagnostics/adopted_{method.replace(" ","_")}.json');m=d['SPPR'];values=np.array(m['values'],float)
 check('matrix dimensions '+method,values.shape==(25,3) and len(set(m['index']))==25 and len(set(m['columns']))==3)
 check('strict signs/nonfinite '+method,np.isfinite(values).all() and not (values<0).any() and not d['negative_entries'])
 for mask in m['masks'].values():check('nonfinite masks '+method,not np.array(mask,bool).any())
 cols={seq:i for i,seq in enumerate(m['columns'])};ids={seq:i for i,seq in enumerate(m['index'])}
 errors=[]
 for r in rows:
  if r['method']!={'GE':'new_GE','TE':'new_TE_EEfix','With Egestion':'new_WithEgestion'}[method]:continue
  gid=next(int(k) for k,v in d['names'].items() if v==r['group']);scope=r['scope']
  chosen=list(cols) if scope=='all' else [c for c in cols if c!=25] if scope=='inner' else [c for c in cols if d['names'][str(c)]=='Phytoplankton']
  total=float(values[ids[gid],[cols[c] for c in chosen]].sum());errors.append(abs(total-r['sppr']))
  check('matrix scope '+method+' '+str(gid)+' '+scope,math.isclose(total,r['sppr'],rel_tol=1e-10,abs_tol=1e-9))
 matrix_detail[method]={'shape':list(values.shape),'strict_negative_count':int((values<0).sum()),'nonfinite_count':int((~np.isfinite(values)).sum()),'coefficient_max_abs_difference':max(errors),'status':d['report']['status']}
check('adopted input hash matches diagnostics',sha(region/'models'/o['selected_model_id']/'model.json')==load(R/'diagnostics/adopted_provenance.json')['model_sha256'])
wb=openpyxl.load_workbook(region/'LME032_taxon_mapping_appendix.xlsx',data_only=True)
ws=wb['Taxon appendix'];append=list(ws.iter_rows(min_row=6,max_col=7,values_only=True))
check('seven exact appendix headers',[ws.cell(5,i).value for i in range(1,8)]==['Taxon name','TL','Catch (t)','Simple-chain PPR (t C)','Mapped group names and weights','Confidence level','Reason'])
check('appendix all rows',len(append)==430 and {row[0] for row in append}==set(bytax))
check('appendix numeric descending',all((-a[3],a[0])<=(-c[3],c[0]) for a,c in zip(append,append[1:])))
check('appendix native types',all(isinstance(row[2],(float,int)) and isinstance(row[3],(float,int)) for row in append))
for row,x in zip(append,audit):check('appendix agreement '+x['taxon'],row[0]==x['taxon'] and row[5]==x['overall_confidence'] and math.isclose(row[2],x['catch_t'],rel_tol=1e-13) and math.isclose(row[3],x['simple_ppr_tC'],rel_tol=1e-13))
check('appendix frozen filtered',ws.freeze_panes=='B6' and len(ws.tables)==1 and next(iter(ws.tables.values())).autoFilter is not None)
for i,x in enumerate(summary['coverage'],6):
 check('appendix coverage '+x['confidence'],wb['Coverage'].cell(i,2).value==x['taxa'] and math.isclose(wb['Coverage'].cell(i,4).value,x['catch_pct'],abs_tol=1e-10) and math.isclose(wb['Coverage'].cell(i,5).value,x['ppr_pct'],abs_tol=1e-10))
check('appendix cached formulas',math.isclose(wb['Coverage']['B13'].value,summary['ppr_tC'],rel_tol=1e-13))
check('appendix native links',len([c for row in wb['Sources'] for c in row if c.hyperlink])==22)
for filename in ['LME032_taxon_mapping_appendix.xlsx',f'Model_validation_{o["selected_model_id"]}.docx']:
 targets=[]
 with zipfile.ZipFile(region/filename) as z:
  for name in z.namelist():
   if name.endswith('.rels'):
    for rel in E.fromstring(z.read(name)):
     if rel.get('Type','').endswith('/hyperlink'):targets.append(rel.get('Target'))
   if name=='docProps/custom.xml':check('no hyperlink base '+filename,'HyperlinkBase' not in z.read(name).decode())
 local=[]
 for target in targets:
  if target.startswith(('http://','https://')):continue
  clean=unquote(target.split('#',1)[0]);check('relative hyperlink '+clean,not Path(clean).is_absolute() and not __import__('ntpath').splitdrive(clean)[0] and '://' not in clean)
  dest=(region/clean).resolve();check('hyperlink exists '+clean,dest.exists() and dest.is_relative_to(ROOT));local.append(clean)
 detail[filename]={'hyperlinks':len(targets),'local':local}
baseline_w=openpyxl.load_workbook(R/'baseline/LME_032.xlsx',read_only=True,data_only=True)
current_w=openpyxl.load_workbook(region/'LME_032.xlsx',read_only=True,data_only=True)
check('NPP unchanged',list(baseline_w['NPP'].values)==list(current_w['NPP'].values))
baseline_w.close();current_w.close();wb.close()
detail.update({'overview':o,'matrix_checks':matrix_detail,'reference':{'year':2019,'basis':'landings','scope':'all independent taxa','catch_t':summary['catch_t'],'simple_chain_ppr_tC':summary['ppr_tC']},'confidence_changes':sum(x['previous_confidence'].lower()!=x['overall_confidence'].lower() for x in audit),'assignment_changes':sum(x['previous_groups']!=[{'group_id':g['group_id'],'weight':g['weight']} for g in x['groups']] for x in audit)})
out={'all_passed':all(checks.values()),'checks':checks,'details':detail}
(R/'qa/final_scientific_checks.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print('Final regional/artifact checks passed:',len(checks))
