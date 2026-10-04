from pathlib import Path
import sys,json,hashlib,shutil
ROOT=Path(__file__).resolve().parents[5]
sys.path.insert(0,str(ROOT/'tools'))
from workbooks import read_book,overview,records
from docx import Document
OUT=Path(__file__).parent
REG=ROOT/'regions/LME_013'
files=[REG/'LME_013.xlsx',REG/'Model_validation_13_1_Chilean_Patagonia_(1980).docx',REG/'LME013_taxon_mapping_appendix.xlsx']
baseline=OUT/'baseline';baseline.mkdir(exist_ok=True)
manifest={}
for p in files:
    q=baseline/p.name
    if not q.exists(): shutil.copy2(p,q)
    manifest[p.relative_to(ROOT).as_posix()]={'sha256':hashlib.sha256(q.read_bytes()).hexdigest(),'backup':q.relative_to(ROOT).as_posix()}
(OUT/'baseline_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
b=read_book(files[0])
(OUT/'baseline_tables.json').write_text(json.dumps(b,ensure_ascii=False,indent=2),encoding='utf-8')
doc=Document(files[1]);snap={'paragraphs':[p.text for p in doc.paragraphs],'tables':[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]}
(OUT/'baseline_document.json').write_text(json.dumps(snap,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(overview(b),ensure_ascii=False,indent=2))
print('TABLES')
for s,t in b.items():
    for n,(h,r) in t.items(): print(s,n,len(r),h[:18])
print('MATCHING')
for r in records(b,'PPR','Matching'):print(json.dumps(r,ensure_ascii=False))
