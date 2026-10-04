from pathlib import Path
import sys, json
from docx import Document
import openpyxl
ROOT = Path(__file__).resolve().parents[6]
sys.path.insert(0, str(ROOT / 'tools'))
from workbooks import read_book, overview, records
region = ROOT / 'regions/LME_013'
evidence = region / 'validation_reports/13_1_Chilean_Patagonia_(1980)'
book = read_book(region / 'LME_013.xlsx')
print('OVERVIEW', json.dumps(overview(book), ensure_ascii=False))
for label, path in [('REPORT', next(region.glob('Model_validation*.docx'))), ('TEMPLATE', ROOT/'tools/templates/Model_validation_template.docx')]:
    doc = Document(path)
    print(label, 'paragraphs', [(i, p.text) for i,p in enumerate(doc.paragraphs)])
    for i,t in enumerate(doc.tables):
        print(label, 'TABLE', i, json.dumps([[c.text for c in r.cells] for r in t.rows], ensure_ascii=False))
for name in ['coverage_summary.json', 'matrix_inspection.json', 'taxon_audit.json', 'evidence_manifest.json']:
    data = json.loads((evidence/name).read_text(encoding='utf-8-sig'))
    print(name, 'TYPE', type(data).__name__, 'KEYS', list(data)[:30] if isinstance(data,dict) else len(data))
    if name=='taxon_audit.json':
        rows=data if isinstance(data,list) else data.get('taxa',data.get('rows',[]))
        print('AUDIT SAMPLE', json.dumps(rows[:1],ensure_ascii=False))
        print('VERY LOW', json.dumps([r for r in rows if 'very low' in str(r.get('overall_confidence',r.get('confidence',''))).lower()],ensure_ascii=False))
    else: print(json.dumps(data, ensure_ascii=False)[:14000])
w=openpyxl.load_workbook(region/'LME013_taxon_mapping_appendix.xlsx', read_only=True, data_only=False)
for s in w:
    print('APPENDIX', s.title, s.max_row,s.max_column,'HEAD',list(s.values)[:7])
w.close()
g=json.loads((ROOT/'tools/knowledge_graph/graph.json').read_text(encoding='utf-8'))
nodes=g.get('nodes',[])
found=[n for n in nodes if 'LME_013' in json.dumps(n,ensure_ascii=False) or 'Chilean Patagonia' in json.dumps(n,ensure_ascii=False)]
print('GRAPH MATCHES',json.dumps(found,ensure_ascii=False)[:13000])
print('MATCHING',len(records(book,'PPR','Matching')))
