"""Integrate authorized corrected-model validation; never build the map."""
from pathlib import Path
from datetime import datetime, timezone
from collections import defaultdict
import copy, csv, hashlib, json, math, os, sys, shutil

RUN=Path(__file__).resolve().parents[1]
MODEL=Path(__file__).resolve().parents[4]
ROOT=next(p for p in MODEL.parents if (p/'Project.xlsx').exists())
REGION=ROOT/'regions/LME/LME_013'
OUT=RUN/'outputs'
sys.path.insert(0,str(ROOT))
from tools.project_core.workbooks.workbooks import *
from tools.project_core.calculations.regional import recalculate,set_setting,set_result_hash
from tools.project_core.registry.update_project import update
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def load(p):return json.loads(p.read_text(encoding='utf-8'))
def dump(p,d):p.write_text(json.dumps(d,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def csvrows(p):
    with p.open(encoding='utf-8',newline='') as f:return list(csv.DictReader(f))
def decode(x):
    t=x['type']
    if t=='dict':return {decode(k):decode(v) for k,v in x['items']}
    if t in ['list','tuple']:return [decode(v) for v in x['items']]
    if t=='null':return None
    if t=='ndarray':return decode(x['data'])
    return x.get('value',x)
def flatten(x,prefix='',out=None):
    out={} if out is None else out
    for k,v in x.items():
        key=prefix+k
        if isinstance(v,dict):flatten(v,key+'_',out)
        else:out[key]=clean(v)
    return out

METHOD={'GE':'new_GE','TE':'new_TE_EEfix','With Egestion':'new_WithEgestion'}
reports={}
summaries=[]
negatives=[]
for mode in ['unpooled','pooled']:
    folder=OUT/mode
    manifest=load(folder/'run_manifest.json')
    assert manifest['input_sha256']==sha(MODEL/'model.json')
    for method in METHOD:
        key=method.replace(' ','_')
        report=load(folder/key/'diagnostic_return.json')
        assert report['model_input']['is_model_balanced']
        reports[mode,method]=report
        info=decode(load(folder/key/'detritus_resolution_info_lossless.json')) if method!='TE' else None
        summaries.append(dict(configuration=mode,method=method,overall_status=report['status'],
            model_balanced=report['model_input']['is_model_balanced'],input_grade=report['model_input']['status'],
            rho_living=report['divergence']['rho_living'],unpooled_b=report['divergence']['b'],
            pooled_b=info.get('collapse_b') if info else None,
            pooled_SPPR=info.get('collapse_scalar') if info else None,
            actual_formulation=report['config']['method'],
            negative_source_columns=report['divergence']['n_negative_sources'],
            SPPR_balanced=report['balance']['is_balanced'],PP_relative_gap=report['balance']['rel_gap']))
    for row in csvrows(folder/'all_negative_matrix_entries.csv'):
        if row['matrix']=='SPPR':negatives.append({'configuration':mode,**row})
dump(OUT/'validation_summary.json',dict(scope='Fresh GE/TE/With Egestion diagnostics and pooled comparisons; no Monte Carlo or map refresh.',
    model_sha256=sha(MODEL/'model.json'),methods=summaries,
    full_diagnostics={mode:{m:reports[mode,m] for m in METHOD} for mode in ['unpooled','pooled']}))

path=REGION/'LME_013.xlsx';baseline_hash=sha(path)
b=read_book(path);original=copy.deepcopy(b);o=overview(b)
assert o['selected_model_id']==MODEL.name
mapping=load(MODEL/'extracted_tables/evidence/mapping/mapping_review.json')
names={int(g['group_seq']):g['group_name'] for g in load(MODEL/'model.json')['group']}
classic={r['taxon']:r for r in records(b,'Classic PPR','Taxa')}
catch={r['taxon']:r for r in records(b,'Catch','Catch') if r['catch_basis']==o['catch_basis']}
assert len(catch)==len(mapping['taxa'])==218
assert set(catch)=={r['taxon'] for r in mapping['taxa']}
adopted=defaultdict(list)
for r in records(b,'PPR','Matching'):
    assert r['model_id']==MODEL.name
    adopted[r['taxon']].append((r['group'],r['weight']))
category=defaultdict(lambda:dict(taxa=0,catch_t=0.,simple_chain_ppr_tC=0.))
membership=defaultdict(float);allocation=defaultdict(float);new_taxa=[]
for r in mapping['taxa']:
    t=r['taxon'];cr=catch[t];c=cr.get(int(o['taxon_detail_year']))
    assert finite(c) and c>=0
    coefficient=classic.get(t,{}).get('sppr')
    ppr=0. if c==0 else c*coefficient/9 if finite(coefficient) else None
    assert ppr is not None
    assert math.isclose(c,r['catch_t'],rel_tol=1e-12,abs_tol=1e-8)
    assert math.isclose(ppr,r['simple_chain_ppr_tC'],rel_tol=1e-12,abs_tol=1e-7)
    candidates=r['candidates']
    expected=sorted((v['group_name'],v['weight']) for v in candidates)
    actual=sorted(adopted[t])
    assert len(actual)==len(expected)
    assert all(g==h and math.isclose(w,v,rel_tol=1e-14,abs_tol=1e-16)
               for (g,w),(h,v) in zip(actual,expected)), t
    assert all(names[v['group_id']]==v['group_name'] for v in candidates)
    assert math.isclose(math.fsum(v['weight'] for v in candidates),1.,abs_tol=1e-10)
    rank={'High':4,'Medium':3,'Low':2,'Very low':1,'Unresolved':0}
    assert rank[r['overall_confidence']]==min(rank[r['membership_confidence']],rank[r['allocation_confidence']])
    entry=category[r['overall_confidence']];entry['taxa']+=1;entry['catch_t']+=c;entry['simple_chain_ppr_tC']+=ppr
    membership[r['membership_rule']]+=ppr;allocation[r['allocation_rule']]+=ppr
    new_taxa.append(dict(taxon=t,catch_t=c,tl=classic.get(t,{}).get('tl'),classic_sppr=coefficient,
        simple_chain_ppr_tC=ppr,membership_confidence=r['membership_confidence'],
        allocation_confidence=r['allocation_confidence'],overall_confidence=r['overall_confidence'],
        assumed_membership=r['assumed_membership'],assumed_allocation=r['assumed_allocation'],
        groups=[dict(seq=v['group_id'],name=v['group_name'],weight=v['weight']) for v in candidates]))
total_ppr=math.fsum(r['simple_chain_ppr_tC'] for r in new_taxa)
total_catch=math.fsum(r['catch_t'] for r in new_taxa)
dump(OUT/'mapping_validation.json',dict(status='PASS',reference_year=o['taxon_detail_year'],basis=o['catch_basis'],
    mapping_decisions_preserved=True,source_membership_and_confidence_evidence_reused=True,
    source_evidence_sha256=sha(MODEL/'extracted_tables/evidence/mapping/mapping_review.json'),
    taxa=new_taxa,category_totals=dict(category),total_catch_t=total_catch,simple_chain_ppr_tC=total_ppr,
    membership_rule_ppr=dict(membership),allocation_rule_ppr=dict(allocation),
    source_proxy_policy='Retain explicitly published native source-model landings proxies. No numerical allocation adoption requested; accepted sardine correction changes removals/EE, not original source proxy evidence.',
    missing_coefficients=sum(not finite(r['classic_sppr']) for r in new_taxa),
    geography_evidence_sha256=sha(MODEL/'extracted_tables/evidence/geography/geographic_assessment.json')))

# Retain the previously authorized pooled research-preview configuration; exact
# FAIL grades remain and no model method receives ordinary production eligibility.
runtime=csvrows(OUT/'pooled/runtime_groups.csv')
def value(v):
    if v=='':return None
    try:x=float(v);return x if math.isfinite(x) else None
    except ValueError:return v
groups=[dict(seq=int(r['group_seq']),**{k:value(v) for k,v in r.items() if k!='group_seq'}) for r in runtime]
b['Selected model groups']['Groups']=table_dict(groups)
coeff=csvrows(OUT/'pooled/group_scope_coefficients.csv')
b['Selected model groups']['Group SPPR']=(['model_id','group','scope','method','sppr'],
    [[MODEL.name,r['group_name'],r['scope'],METHOD[r['method']],float(r['sppr_wet'])] for r in coeff])
statuses={METHOD[m]:f"provisional: DIRECT {m} overall {reports['pooled',m]['status']}; balanced model; det_collapse_mode=always; pending researcher review; production-ineligible" for m in METHOD}
for row in b['PPR']['Annual'][1]:
    if row[2] in statuses and not(row[1]!='all' and row[4]=='simple'):row[6]=statuses[row[2]]
set_setting(b,'results_model_sha256',sha(MODEL/'model.json'))
recalculate(b,path)

# The general calculation rounds taxon coefficients. Retain full precision for
# this reviewed direct source-matrix run, including genuine zero contributions.
group_coeff={(r['group_name'],r['scope'],METHOD[r['method']]):float(r['sppr_wet']) for r in coeff}
taxon_coeff={}
for t,assign in adopted.items():
    for scope in ['all','inner','PP']:
        for method in METHOD.values():
            vals=[group_coeff.get((g,scope,method)) for g,w in assign]
            taxon_coeff[t,scope,method]=math.fsum(w*v for (g,w),v in zip(assign,vals)) if all(finite(v) for v in vals) else None
b['PPR']['Taxon SPPR']=(['model_id','taxon','scope','method','sppr'],
    [[MODEL.name,t,s,m,v] for (t,s,m),v in sorted(taxon_coeff.items())])
catch_all={(r['taxon'],r['catch_basis']):r for r in records(b,'Catch','Catch')}
simple={t:r.get('sppr') for t,r in classic.items()}
unknown_annual=[]
for row in b['PPR']['Annual'][1]:
    model,scope,method,basis,treatment,metric,status=row[:7]
    unsupported=scope!='all' and treatment=='simple'
    for i,y in enumerate(YEARS):
        contributions=[];quantities=[];covered=[];unknown=False;unknown_catch=False
        for t in adopted:
            source=catch_all.get((t,basis),{});quantity=source.get(y)
            v=taxon_coeff[t,scope,method]
            if source.get('unidentified') and treatment=='zero':v=0.
            if source.get('unidentified') and treatment=='simple':v=simple.get(t)
            if finite(quantity):
                quantities.append(quantity)
                if quantity==0:contributions.append(0.);covered.append(0.)
                elif finite(v):contributions.append(quantity*v);covered.append(quantity)
                else:unknown=True
            else:unknown=True;unknown_catch=True
        row[7+i]=(math.fsum(quantities) if not unknown_catch else None) if metric=='catch' else (
            None if unsupported or unknown else math.fsum(contributions if metric=='ppr' else covered))
        if unknown and metric=='ppr':unknown_annual.append([scope,method,basis,treatment,y])
den={r['method']:r for r in records(b,'NPP','NPP')}
classic_ratios=[r for r in original['PPR–NPP']['Ratios'][1] if not r[0]]
ratios=[]
for row in b['PPR']['Annual'][1]:
    if row[5]!='ppr':continue
    for name,npp in den.items():ratios.append([*row[:5],name,*[
        100*row[7+i]/9/npp[y] if finite(row[7+i]) and finite(npp.get(y)) and npp[y]>0 else None
        for i,y in enumerate(YEARS)]])
b['PPR–NPP']['Ratios']=(original['PPR–NPP']['Ratios'][0],classic_ratios+ratios)
detail=[]
refyear=int(o['taxon_detail_year']);basis=o['catch_basis']
for (t,s,m),v in sorted(taxon_coeff.items()):
    c=catch_all[t,basis].get(refyear)
    contribution=0. if c==0 else c*v if finite(c) and finite(v) else None
    detail.append([MODEL.name,t,s,m,refyear,basis,c,v,contribution])
b['PPR']['Taxon PPR inspected year']=(original['PPR']['Taxon PPR inspected year'][0],detail)
b['Classic PPR']=original['Classic PPR']
b['Diagnostics']['model_health']=table_dict([flatten(reports['pooled',m])|{'production_eligible':False} for m in METHOD])
b['Diagnostics']['Validation configurations']=table_dict(summaries)
notes=records(b,'Diagnostics','run_notes')
notes=[r for r in notes if r.get('topic') not in ['Corrected-model validation','Monte Carlo','Map and researcher decision']]
notes.extend([dict(topic='Corrected-model validation',note=(OUT/'validation_summary.json').relative_to(REGION).as_posix(),status='fresh unpooled and pooled; all overall FAIL'),
    dict(topic='Monte Carlo',note='Not rerun; prior sensitivity belongs to the old input. No current uncertainty bounds claimed.',status='NOT_RUN'),
    dict(topic='Map and researcher decision',note='Map refresh deferred by user until human validation. No researcher verdict registered.',status='pending researcher review')])
b['Diagnostics']['run_notes']=table_dict(notes)
manifest=load(OUT/'pooled/run_manifest.json')
for k,v in dict(production_eligible=False,computational_input_sha256=sha(MODEL/'model.json'),
    results_model_id=MODEL.name,computational_runtime_evidence=(OUT/'pooled/run_manifest.json').relative_to(REGION).as_posix(),
    computational_state_sha256=manifest['runtime_before_sha256'],diagnostic_manifest=(OUT/'validation_summary.json').relative_to(REGION).as_posix(),
    sppr_constructor=json.dumps(manifest['constructor']),sppr_configuration=json.dumps(manifest['method_call']),
    sppr_engine_hashes=json.dumps(manifest['engine']),det_collapse_mode='always',
    validation_document=(MODEL/'model_validation/validation.docx').relative_to(REGION).as_posix(),
    source_note_reference=(MODEL/'model_notes.md').relative_to(REGION).as_posix(),
    canonical_source_path=(MODEL.parents[0]/(MODEL.name+'__source')/'model.json').relative_to(REGION).as_posix(),
    source_note='Corrected Northern Humboldt model passes mass balance; GE/TE/With Egestion retain overall FAIL. Pooled results retained provisionally for human review; map update deferred.',
    calculation_status='corrected model balanced; full unpooled and pooled validation complete; all SPPR overall FAIL; provisional research estimates; pending human review').items():set_setting(b,k,v)
set_setting(b,'calculation_input_sha256',input_hash(b));set_result_hash(b)
assert b['Catch']==original['Catch'] and b['Classic PPR']==original['Classic PPR'] and b['NPP']==original['NPP']
assert b['PPR']['Matching']==original['PPR']['Matching']
validate_region(b,path)
assert sha(path)==baseline_hash
update_book(path,b)
reopened=read_book(path);validate_region(reopened,path)
assert reopened['PPR']['Matching']==original['PPR']['Matching']
assert digest_tables([reopened['Catch'],reopened['Classic PPR'],reopened['NPP']])==digest_tables([original['Catch'],original['Classic PPR'],original['NPP']])
print('Regional workbook updated and reopened; catch/classic/NPP/mapping preserved.',flush=True)

# Coefficient workbook includes both complete diagnostic configurations and all
# negative source-to-recipient pairs as a referenced report attachment.
w=openpyxl.Workbook();w.remove(w.active)
def sheet(title,header,rows):
    s=w.create_sheet(title);s.append(header)
    for row in rows:s.append([clean(v) for v in row])
    s.freeze_panes='C2';s.auto_filter.ref=s.dimensions
    for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='174C58')
    for row in s.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    for j in range(1,len(header)+1):s.column_dimensions[get_column_letter(j)].width=28 if j<5 else 23
    return s
sheet('Validation summary',list(summaries[0]),[list(r.values()) for r in summaries])
sheet('groups_df',*b['Selected model groups']['Groups'])
sheet('model_health',*b['Diagnostics']['model_health'])
for scope in ['all','inner','PP']:
    sheet('sppr_'+scope,['seq','group_name',*METHOD.values()],[[g['seq'],g['group_name'],*[group_coeff.get((g['group_name'],scope,m)) for m in METHOD.values()]] for g in groups])
if negatives:sheet('Negative SPPR',list(negatives[0]),[list(r.values()) for r in negatives])
for mode in ['unpooled','pooled']:
    for m in METHOD:
        folder=OUT/mode/m.replace(' ','_')
        frame=csvrows(folder/'SPPR.csv')
        keys=list(frame[0]);sheet(mode[:1]+'_'+m.replace(' ','_')+'_matrix',keys,
            [[value(r[k]) for k in keys] for r in frame])
sheet('Sources',['Field','Value'],[
    ['model_id',MODEL.name],['model_sha256',sha(MODEL/'model.json')],
    ['Input','Corrected canonical model; accepted primary recommendations of 8 October 2026.'],
    ['Configuration','Unpooled never and pooled always. TE has no detritus scaling; its always call is unchanged, not a pooled TE formulation.'],
    ['Restrictions','All returned overall FAIL grades retained. Pooled nonnegative SPPR does not establish convergent unpooled recycling or human approval.']])
w.save(MODEL/'sppr_source.xlsx');w.close()

appendix=MODEL/'model_validation/taxon_mapping.xlsx'
w=openpyxl.load_workbook(appendix);s=w['Taxon mappings']
old_rows=[[c.value for c in row] for row in s.iter_rows(min_row=8,max_col=7)]
s['A1']='Northern Humboldt corrected-model taxon mappings'
s['A2']='2019 landings; all 218 labels including zero catch; corrected Northern Humboldt 1995–1998 model'
# The native table already provides filters; a second worksheet filter overlaps it.
s.auto_filter.ref=None
src=w['Sources'];src['A1']='Sources for model validation, membership and allocation'
src['A2']='Mapping evidence retained from 3 October 2026; corrected-model diagnostics and source corrections verified 8 October 2026.'
for row in src.iter_rows(min_row=5):
    if row[0].value=='Native':row[3].value='Original Table A supplies published landings/discards and allocation proxies. These source proxies are retained; accepted sardine removals/EE corrections are separate.'
nr=next((row[0].row for row in src.iter_rows(min_row=5) if row[0].value=='AcceptedCorrections'),src.max_row+1)
for j,v in enumerate(['AcceptedCorrections','Source-backed jellyfish reconstruction and article-priority sardine correction, 8 October 2026','Model notes','Adopted diet 68% meso / 17% macro / 15% eggs; sardine landings 1.4, source discards retained, EE recalculated. Original author-source contradictions remain documented.'],1):src.cell(nr,j,v)
src.cell(nr,3).hyperlink='../model_notes.md';src.cell(nr,3).font=Font(color='0563C1',underline='single')
src.row_dimensions[nr].height=85
for cell in src[nr]:cell.alignment=Alignment(wrap_text=True,vertical='top')
for t in src.tables.values():
    t.ref=f'A4:D{nr}'
    if t.autoFilter is not None:t.autoFilter.ref=t.ref
w.save(appendix);w.close()
w=openpyxl.load_workbook(appendix)
assert digest_tables([old_rows])==digest_tables([[[c.value for c in row] for row in w['Taxon mappings'].iter_rows(min_row=8,max_col=7)]])
w.close()
print('Coefficient workbook and linked mapping appendix updated.',flush=True)

# This API writes Project.xlsx only and serializes central writes; no map call.
update(ROOT,[path])
shutil.copy2(path,MODEL/'results/regional_snapshot.xlsx')
dump(MODEL/'results/result_manifest.json',dict(schema_version=1,model_id=MODEL.name,
    input_sha256=sha(MODEL/'model.json'),snapshot_sha256=sha(MODEL/'results/regional_snapshot.xlsx'),
    snapshot_selection=MODEL.name,timestamp_utc=datetime.now(timezone.utc).isoformat(),
    diagnostics=manifest['method_call'],constructor=manifest['constructor'],engine_hashes=manifest['engine'],
    evidence=(OUT/'validation_summary.json').relative_to(MODEL).as_posix(),production_eligible=False,
    researcher_review='pending',map_refresh='deferred by user',
    independent_inputs_and_mapping_preserved=True))
dump(OUT/'workbook_integration.json',dict(status='PASS',regional_sha256=sha(path),project_sha256=sha(ROOT/'Project.xlsx'),
    full_precision_coefficients=True,catch_classic_NPP_mapping_preserved=True,unknown_annual_keys=unknown_annual,
    annual_rows=len(b['PPR']['Annual'][1]),taxon_coefficients=len(taxon_coeff),model_sha256=sha(MODEL/'model.json'),
    map_refreshed=False,researcher_verdict_registered=False))
print('Project.xlsx and complete regional snapshot updated; map not refreshed.',flush=True)
