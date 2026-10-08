"""Create a source-only scientific rationale beside the validation report."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

RUN=Path(__file__).resolve().parents[1]
MODEL=Path(__file__).resolve().parents[4]
OUT=MODEL/'model_validation/source_value_corrections.docx'
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
    s=d.styles[name]; s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.06
d.styles['Title'].font.size=Pt(23)
d.styles['Title'].paragraph_format.space_after=Pt(9)
d.styles['Heading 1'].font.size=Pt(15)
d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
d.styles['Heading 1'].paragraph_format.space_after=Pt(7)
def p(t):return d.add_paragraph(t)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt,w in zip(t.rows[0].cells,headers,widths):c.text=txt;c.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for rr in rows:
        for c,txt,w in zip(t.add_row().cells,rr,widths):c.text=txt;c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if i==0 else ('F2F5F8' if i%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),'95');x.set(qn('w:type'),'dxa');margins.append(x)
            c._tc.get_or_add_tcPr().append(margins)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.02
                if j:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in par.runs:
                    r.font.size=Pt(10.5)
                    if i==0:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
def link(par,text,target):
    rid=par.part.relate_to(target,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');pr.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline)
    r.append(pr);tx=OxmlElement('w:t');tx.text=text;r.append(tx);h.append(r);par._p.append(h)

d.add_paragraph('Reasons for correcting the Northern Humboldt source values','Title')
p('Detailed model for 1995–1998 from Chiaverano et al 2018\nScientific rationale prepared 8 October 2026')
p('We chose a coherent reconstruction of the author evidence: large jellyfish consume 68% mesozooplankton, 17% macrozooplankton and 15% anchovy eggs; sardine landings are 1.4 t km⁻² year⁻¹; sardine ecotrophic efficiency is recalculated as 0.68720654. The diet is inferred from the authors’ production matrix. The landings follow their explicit balancing statement. The efficiency follows their mass balance equation. These choices resolve contradictions between published source components; they are not newly measured parameters.')
d.add_paragraph('Reconstructing the large jellyfish diet','Heading 1')
p('Supplementary Table B, column H, gives the following Chrysaora plocamia diet. Its fractions sum to 1.045, so they cannot all be a complete diet as printed. Simply dividing by that sum preserves the disputed links and does not explain the independently published production matrix.')
table(['Prey','Table B fraction','Chosen fraction'],[
 ['Diatoms','0.00985222','0'],['Mesozooplankton','0.73118227','0.68'],
 ['Macrozooplankton','0.10470443','0.17'],['Small gelatinous zooplankton','0.04926108','0'],
 ['Anchovy eggs','0.15','0.15'],['Total','1.045','1.00']], [3.4,1.6,1.9])
p('Table E identifies mesozooplankton, macrozooplankton, small jellyfish and large jellyfish as separate singleton groups in the aggregated system. Table H row 14 therefore supplies independent direct connections for the same large jellyfish consumer. Cells F14 and K14 are zero for diatoms and small jellyfish. Cells I14 and J14 are 0.1070741622 and 0.0462539820 for mesozooplankton and macrozooplankton.')
p('Using Table A consumption and the physiological partition of the living connections in Table H, each nonzero coefficient is multiplied by prey consumption and divided by large jellyfish consumption. The mesozooplankton calculation is 0.1070741622 × 3010.2500916 ÷ 474.0000076 = 0.6800000031; the macrozooplankton calculation is 0.0462539820 × 1742.1203679 ÷ 474.0000076 = 0.1700000061. We retain Table B H39’s egg fraction of 0.15. Together these give a complete diet without adding unsupported prey.')
p('This interpretation also recovers anchovy egg EE independently: reconstructing source predation and detritus routing gives 0.88433221, matching Table A H41’s 0.88433224 to source precision. This agreement strengthens the reconstruction, but cannot establish which source version was the authors’ final intended model. [1, 2]')
d.add_page_break()
d.add_paragraph('Following the stated sardine landings correction','Heading 1')
p('The article Methods on printed page 30 explicitly describe reducing sardine landings from 5.65 to 1.4 t km⁻² year⁻¹ to achieve balance. Supplementary Table A J14 still reports 5.6513425. We prioritize the specific correction described in the article, as requested, while keeping the independently reported biological inputs unchanged. [1, 2]')
table(['Sardine quantity','Source evidence','Chosen value'],[
 ['Landings','Article page 30','1.4'],['Discards','Table A L14','0.17462648325'],
 ['Total fishery removals','Landings plus discards','1.57462648325'],
 ['Ecotrophic efficiency','Recalculated from Eq 1','0.68720654021']], [2.45,2.25,2.2])
p('We retain reported discards because no explicit revised quantity was recovered. The Methods describe using gear specific discard rates, so scaling with landings is a plausible alternative: discards 0.04326 and EE 0.67846812. That would impose a rate preservation assumption rather than retain the reported quantity. [1, 2]')
d.add_paragraph('Why sardine efficiency must be recalculated','Heading 1')
p('Article Eq 1 balances production utilized within the system against predation, fishery catch, biomass accumulation and net migration. Table A H14’s sardine EE of 0.97000349 is bold, marking it as an estimated parameter. Once landings decrease, retaining that efficiency while holding biomass and production fixed would leave utilization inconsistent with the losses. We therefore recompute the estimated efficiency rather than changing biomass, production or consumption to absorb the discrepancy. [1, 2]')
p('With the reconstructed diets and the zero accumulation and zero net migration used for the steady state calculation, sardine production is 10.7379999161 × 1.3999999762 = 15.0331996265 t km⁻² year⁻¹. Predation is approximately 8.75627983 in the same units. Adding 1.4 landings and 0.17462648325 discards, then dividing by production, gives EE = 0.68720654021. This is a derived value. It is not printed in the source, and zero accumulation and migration are calculation conventions rather than independently reported measurements.')
d.add_paragraph('Remaining source uncertainty','Heading 1')
p('Table 1 on page 31 gives forage fish landings of 28.13, and Scenario III describes harvest near 29 including discards. These support the higher catch branch. We choose the explicit balancing statement while acknowledging this unresolved source conflict. Living group balance in our reconstruction does not verify the authors’ native Ecopath or ECOTRAN implementation, fleet ancestry, detritus routing or downstream SPPR suitability. [1]')
d.add_paragraph('Source references','Heading 1')
par=p('[1] Chiaverano LM et al 2018. Evaluating the role of large jellyfish and forage fishes as energy pathways, and their interplay with fisheries, in the Northern Humboldt Current System. Progress in Oceanography 164, 28–36. Relevant locations: page 30, Eq 1 and Methods; page 31, Table 1; Scenario III, pages 30–31. ')
link(par,'Article PDF','../../../sources/1-s2.0-S0079661117303312-main.pdf')
par=p('[2] Author supplementary workbook, Supplementary material revised and final.xls. Table A C14–H14, J14, L14 and H41; Table B H4, H7–H9 and H39; Table C detritus fate; Table E group identities; Table H F14, I14, J14 and K14. ')
link(par,'Author supplementary workbook','../../../sources/Supplementary%20material%20revised%20and%20final.xls')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Northern Humboldt source correction rationale  |  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
for r in footer.runs:r.font.size=Pt(9)
d.core_properties.title='Reasons for correcting the Northern Humboldt source values'
d.core_properties.subject='Source evidence and interpretation for the 1995–1998 detailed model'
d.save(OUT)
print(OUT)
