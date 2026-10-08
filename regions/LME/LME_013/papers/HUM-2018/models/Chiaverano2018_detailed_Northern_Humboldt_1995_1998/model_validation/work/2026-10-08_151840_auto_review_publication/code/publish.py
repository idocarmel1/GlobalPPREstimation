"""Publish the explicitly authorized conditional researcher decision and auto evidence."""
from pathlib import Path
import copy, json, shutil, sys, zipfile
from datetime import datetime, timezone

RUN=Path(__file__).resolve().parents[1]
MODEL=Path(__file__).resolve().parents[4]
ROOT=next(p for p in MODEL.parents if (p/'Project.xlsx').exists())
sys.path.insert(0,str(ROOT))
from tools.project_core.workbooks.workbooks import read_book,update_book,overview,records,sha,input_hash,clean,table_dict
from tools.project_core.calculations.regional import set_setting,set_result_hash
from tools.project_core.registry.writes import central_lock
from tools.project_core.registry.update_project import update
from tools.project_core.maps.build_html import build
from tools.project_core.maps.original_atlas_data import embedded
from tools.project_core.validation import researcher_review as review
from tools.project_core.validation.validation_percentage_format import verify_report
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font

def load(p):return json.loads(p.read_text(encoding='utf-8'))
def dump(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def flatten(x,prefix='',out=None):
    out={} if out is None else out
    for k,v in x.items():
        if isinstance(v,dict):flatten(v,prefix+k+'_',out)
        else:out[prefix+k]=clean(v)
    return out

REGION=ROOT/'regions/LME/LME_013/LME_013.xlsx'
AUTO=RUN/'outputs/auto'
manifest=load(AUTO/'run_manifest.json')
assert manifest['method_call']['det_collapse_mode']=='auto'
assert manifest['input_sha256']==sha(MODEL/'model.json')
assert all(v['SPPR']['exact_equal'] and v['A']['exact_equal'] and v['L']['exact_equal'] for v in load(RUN/'qa/auto_matches_previous_pooled.json').values())
report=MODEL/'model_validation/validation.docx';verify_report(report)
name,date,summary=review.read_report(ROOT,report,MODEL.name)
assert name=='Ido Carmel' and date=='2026-10-08'
assert "det_collapse_mode='auto'" in summary['sections'][2]['rows'][0]['text']
excluded_names=['Cetaceans','Fishery offal','Pinnipeds','Seabirds','Chrysaora plocamia']
book=read_book(REGION);original=copy.deepcopy(book);settings=overview(book)
assert settings['results_model_id']==MODEL.name
group_ids={g['group_name']:int(g['seq']) for g in records(book,'Selected model groups','Groups')}
excluded_seq=[group_ids[n] for n in excluded_names]
for n in excluded_names:assert n in summary['sections'][2]['rows'][0]['text']
assert input_hash(book)==settings['calculation_input_sha256']
relative=lambda p:p.relative_to(REGION.parent).as_posix()
set_setting(book,'det_collapse_mode','auto')
set_setting(book,'sppr_configuration',json.dumps(manifest['method_call']))
set_setting(book,'computational_runtime_evidence',relative(AUTO/'run_manifest.json'))
set_setting(book,'computational_state_sha256',manifest['runtime_before_sha256'])
set_setting(book,'diagnostic_manifest',relative(AUTO/'methods_summary.json'))
set_setting(book,'source_note','Researcher validated on 8 October 2026 only with det_collapse_mode=auto for every SPPR_new-related calculation. Exact auto outputs equal prior pooled values; engine overall FAIL restrictions retained.')
set_setting(book,'calculation_status','researcher validated conditionally: det_collapse_mode=auto; engine GE/TE/With Egestion overall FAIL; numerical estimates retain diagnostic restrictions')
set_setting(book,'selection_rationale','Corrected Northern Humboldt 1995–1998 detailed model; researcher validated on 8 October 2026 only with det_collapse_mode=auto. Engine overall FAIL grades and production-ineligible numerical restrictions remain.')
set_setting(book,'researcher_validation_condition',"All SPPR_new-related calculations require det_collapse_mode='auto'.")
methods=['GE','TE','With Egestion']
diagnostics={m:load(AUTO/m.replace(' ','_')/'diagnostic_return.json') for m in methods}
book['Diagnostics']['model_health']=table_dict([{**flatten(diagnostics[m]),'production_eligible':False} for m in methods])
book['Diagnostics']['run_notes'][1].append(['conditional researcher validation',"Signed 2026-10-08; all SPPR_new-related calculations require det_collapse_mode=auto. Fresh exact matrix equivalence verified; researcher display exclusions do not change saved scientific inputs or coefficients.",'validated by researcher; engine FAIL retained'])
for s in ['Catch','Classic PPR','Selected model groups','PPR','NPP','PPR–NPP']:
    assert book[s]==original[s],s
assert input_hash(book)==settings['calculation_input_sha256']
update_book(REGION,book)
fresh=read_book(REGION)
assert input_hash(fresh)==settings['calculation_input_sha256']
for s in ['Catch','Classic PPR','Selected model groups','PPR','NPP','PPR–NPP']:assert fresh[s]==original[s],s
shutil.copy2(REGION,MODEL/'results/regional_snapshot.xlsx')

# This coefficient workbook retains original never/always audit comparisons;
# auto has independently reproduced every adopted value exactly.
p=MODEL/'sppr_source.xlsx';w=load_workbook(p)
baseline_values={s.title:list(s.values) for s in w}
src=w['Sources']
src['B5']='Active validated configuration: det_collapse_mode=auto for all SPPR_new-related calculations. Fresh auto GE/TE/With Egestion matrices exactly match the retained pooled always comparison. Historical never/always audit rows remain.'
src['B6']='Researcher validated conditionally on 8 October 2026. All engine overall FAIL grades remain; TE PP budget failure remains. Display-only exclusions preserve original scientific groups and coefficients.'
src.cell(7,1,'Active auto evidence');src.cell(7,2,(AUTO/'run_manifest.json').relative_to(p.parent).as_posix())
src['B7'].hyperlink=(AUTO/'run_manifest.json').relative_to(p.parent).as_posix();src['B7'].font=Font(color='0563C1',underline='single')
for i in [5,6,7]:src.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top');src.row_dimensions[i].height=60
health=w['model_health'];heads=[c.value for c in health[1]]
if 'config_det_collapse_mode' in heads:
    ci=heads.index('config_det_collapse_mode')+1
    for row in range(2,health.max_row+1):health.cell(row,ci,'auto')
for s in w:
    if s.title not in ['Sources','model_health']:assert list(s.values)==baseline_values[s.title]
w.save(p);w.close()

result=load(MODEL/'results/result_manifest.json')
result.update(snapshot_sha256=sha(MODEL/'results/regional_snapshot.xlsx'),timestamp_utc=datetime.now(timezone.utc).isoformat(),diagnostics=manifest['method_call'],evidence=relative(AUTO/'run_manifest.json'),researcher_review='Validated by researcher',researcher_validation_condition="Only det_collapse_mode='auto' for all SPPR_new-related calculations",map_refresh='authorized; publication in progress',production_eligible=False)
dump(MODEL/'results/result_manifest.json',result)

with central_lock(ROOT):
    # Read the latest shared state only after obtaining the guard. Never use
    # a Project snapshot from before the other session's publication.
    project=ROOT/'Project.xlsx';before=RUN/'inputs/Project_before_publication.xlsx'
    shutil.copy2(project,before)
    old_central=read_book(project)
    old_pages={name:embedded(ROOT/'interactive_map'/name,var)[0] for name,var in [('index.html','DB'),('trends.html','SERIES_DB')]}
    hs_before={s:{t:[r for r in records(old_central,s,t) if r.get('unit_id')=='HS_077'] for t in old_central[s]} for s in old_central}
    update(ROOT,[REGION],lock_held=True)
    # The signed source uses the plural 'calculations'. Resolve that wording
    # in the computed display note without altering the researcher's text or
    # the shared parser. Its five exact exclusions were independently joined.
    original_note=review.exclusion_note
    review.exclusion_note=lambda text:original_note(text.replace('Removed groups from calculations:','Removed groups from calculation:'))
    try:review._register_review(project,report,'LME_013',MODEL.name,excluded_seq)
    finally:review.exclusion_note=original_note
    build(project,only_units={'LME_013'})
    central=read_book(project)
    hs_after={s:{t:[r for r in records(central,s,t) if r.get('unit_id')=='HS_077'] for t in central[s]} for s in central}
    assert hs_before==hs_after,'HS_077 central state changed'
    md=next(r for r in records(central,'Models & coverage','Models') if r['unit_id']=='LME_013' and r['model_id']==MODEL.name)
    approved=review.approved_review(ROOT,md,fresh)
    assert approved and set(approved['excluded_group_ids'])==set(excluded_names)
    for page,var in [('index.html','DB'),('trends.html','SERIES_DB')]:
        payload,layout=embedded(ROOT/'interactive_map'/page,var)
        units=payload['network']['units'] if var=='DB' else payload['units']
        old_units=old_pages[page]['network']['units'] if var=='DB' else old_pages[page]['units']
        assert units['HS_077']==old_units['HS_077'],'HS_077 page changed'
        target=next(m for m in units['LME_013']['models'] if m['id']==MODEL.name)
        assert target['researcher_review']==approved
        assert 'recorded_review_pending' not in target
        assert set(target['display_ppr_excluded_group_ids'])==set(excluded_names)
        assert target['workbook_sha256']==sha(REGION)
        assert sha(project) in layout
    dump(RUN/'qa/publication_checks.json',{'status':'PASS','reviewer':name,'review_date':date,'condition':"det_collapse_mode='auto' for all SPPR_new-related calculations",'excluded_seq':excluded_seq,'excluded_names':excluded_names,'auto_matrices_exactly_match_previous_pooled':True,'scientific_tables_unchanged':True,'HS_077_central_and_model_payload_unchanged':True,'Project_sha256':sha(project),'report_sha256':sha(report),'regional_sha256':sha(REGION),'approved_review':approved})
result['map_refresh']='published and verified';dump(MODEL/'results/result_manifest.json',result)
print('Conditional auto review published; HS_077 preserved.',flush=True)
