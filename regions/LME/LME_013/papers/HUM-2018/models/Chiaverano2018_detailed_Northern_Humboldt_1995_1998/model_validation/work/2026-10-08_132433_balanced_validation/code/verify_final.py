"""Fresh checks of scientific, workbook, attachment and protected-map claims."""
from pathlib import Path
import csv, json, math, hashlib, sys, zipfile, warnings
from collections import defaultdict
from urllib.parse import unquote
import numpy as np
from lxml import etree
import openpyxl
RUN=Path(__file__).resolve().parents[1];MODEL=Path(__file__).resolve().parents[4]
ROOT=next(p for p in MODEL.parents if (p/'Project.xlsx').exists())
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tools/scientific_code/PPREstimation'))
from tools.project_core.workbooks.workbooks import *
from tools.project_core.calculations.regional import result_hash
from ModelData import ModelData
from PPRCalculator import PPRCalculator
def load(p):return json.loads(p.read_text(encoding='utf-8'))
def dump(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def csvrows(p):
    with p.open(encoding='utf-8',newline='') as f:return list(csv.DictReader(f))
def close(a,b):
    if a is None or b is None:return a is b
    return math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-8)
diff=[]
def compare(a,b,key=''):
    if isinstance(a,dict):
        assert a.keys()==b.keys()
        for k in a:compare(a[k],b[k],key+'/'+str(k))
    elif isinstance(a,list):
        assert len(a)==len(b)
        for i,(x,y) in enumerate(zip(a,b)):compare(x,y,key+'/'+str(i))
    elif a!=b:diff.append([key,a,b])
compare(load(RUN/'inputs/model.json'),load(MODEL/'model.json'))
assert len(diff)==6,diff
adopted=load(RUN/'outputs/adopted_corrections.json')
assert sha(MODEL/'model.json')==load(RUN/'outputs/validation_summary.json')['model_sha256']
settings=dict(underdetermined=True,zero_catch=True,zero_biomass_accum=False,default_gs=True,
 weight_flow=1.,weight_guess=1.,normalize_DC=True,DC_tol=.001)
with warnings.catch_warnings(record=True) as ww:
    calc=PPRCalculator.from_modeldata(ModelData(str(MODEL/'model.json'),model_name='Humboldt Current',model_year='1995'),**settings)
balanced,p,q=calc.is_model_balanced();assert balanced
dump(RUN/'qa/notebook_reproduction.json',dict(balanced=bool(balanced),settings=settings,
 model_sha256=sha(MODEL/'model.json'),max_relative_production_residual=float(((p-calc.p)/calc.p).abs().max()),
 max_relative_consumption_residual=float(((q-calc.q)/calc.q).abs().max()),warnings=[str(w.message) for w in ww]))
print('Canonical model has exactly six adopted changes; exact notebook constructor is balanced.',flush=True)
negatives=[];matrices={};coeff={}
for mode in ['unpooled','pooled']:
    manifest=load(RUN/f'outputs/{mode}/run_manifest.json')
    assert manifest['input_sha256']==sha(MODEL/'model.json')
    for filename,digest in manifest['engine'].items():assert sha(ROOT/'tools/scientific_code/PPREstimation'/filename)==digest
    for m in ['GE','TE','With Egestion']:
        f=RUN/f'outputs/{mode}/{m.replace(" ","_")}'
        axes=load(f/'SPPR_axes.json');arr=np.load(f/'SPPR.npy');matrices[mode,m]=arr
        assert arr.shape==(len(axes['rows']),len(axes['columns']))
        assert load(f/'diagnostic_return.json')['status']=='FAIL'
        masks=np.load(f/'SPPR_masks.npz')
        assert all(x.shape==arr.shape for x in masks.values())
        for r,c in zip(*np.where(arr<0)):
            negatives.append((mode,m,str(axes['columns'][c]['id']),str(axes['rows'][r]['id']),float(arr[r,c])))
    coeff[mode]=csvrows(RUN/f'outputs/{mode}/group_scope_coefficients.csv')
    for rr in coeff[mode]:
        axes=load(RUN/f'outputs/{mode}/{rr["method"].replace(" ","_")}/SPPR_axes.json')
        row=next(i for i,a in enumerate(axes['rows']) if a['id']==int(rr['group_id']))
        ids=[int(x) for x in rr['source_ids'].split('|') if x]
        cols=[i for i,c in enumerate(axes['columns']) if c['id'] in ids]
        assert close(float(rr['sppr_wet']),float(matrices[mode,rr['method']][row,cols].sum()))
assert np.array_equal(matrices['unpooled','TE'],matrices['pooled','TE'],equal_nan=True)
w=openpyxl.load_workbook(MODEL/'sppr_source.xlsx',data_only=False)
s=w['Negative SPPR'];h=[c.value for c in s[1]]
record=[dict(zip(h,r)) for r in s.iter_rows(min_row=2,values_only=True)]
expected=sorted(negatives,key=lambda x:x[:4]);actual=sorted([(r['configuration'],r['method'],str(r['source_column_id']),str(r['recipient_group_id']),float(r['value'])) for r in record],key=lambda x:x[:4])
assert len(actual)==len(expected)
assert all(a[:4]==b[:4] and close(a[4],b[4]) for a,b in zip(actual,expected))
assert len(w['Validation summary']['A'])==7
assert all(s.freeze_panes and s.auto_filter.ref for s in w)
w.close()
print(f'All six matrices, source axes, masks, scoped coefficients and {len(actual)} negative pairings verified.',flush=True)
b=read_book(ROOT/'regions/LME/LME_013/LME_013.xlsx');o=validate_region(b,ROOT/'regions/LME/LME_013/LME_013.xlsx')
old=read_book(RUN/'inputs/LME_013.xlsx')
assert digest_tables([b['Catch'],b['Classic PPR'],b['NPP'],b['PPR']['Matching']])==digest_tables([old['Catch'],old['Classic PPR'],old['NPP'],old['PPR']['Matching']])
assert o['calculation_input_sha256']==input_hash(b)
assert o['calculation_result_sha256']==result_hash(b)
assert o['production_eligible'] is False
mapping=defaultdict(list)
for r in records(b,'PPR','Matching'):mapping[r['taxon']].append((r['group'],r['weight']))
method={'GE':'new_GE','TE':'new_TE_EEfix','With Egestion':'new_WithEgestion'}
gc={(r['group_name'],r['scope'],method[r['method']]):float(r['sppr_wet']) for r in coeff['pooled']}
tc={}
for r in records(b,'PPR','Taxon SPPR'):
    assign=mapping[r['taxon']];vals=[gc.get((g,r['scope'],r['method'])) for g,v in assign]
    ex=math.fsum(weight*v for (g,weight),v in zip(assign,vals)) if all(finite(v) for v in vals) else None
    assert close(r['sppr'],ex);tc[r['taxon'],r['scope'],r['method']]=ex
catch={(r['taxon'],r['catch_basis']):r for r in records(b,'Catch','Catch')}
simple={r['taxon']:r.get('sppr') for r in records(b,'Classic PPR','Taxa')}
checked=0
for r in records(b,'PPR','Annual'):
    if r['metric']!='ppr':continue
    for year in YEARS:
        if r['scope']!='all' and r['unidentified']=='simple':assert r[year] is None;continue
        terms=[];unknown=False
        for t in mapping:
            cr=catch.get((t,r['catch_basis']),{});amount=cr.get(year);v=tc[t,r['scope'],r['method']]
            if cr.get('unidentified') and r['unidentified']=='zero':v=0.
            if cr.get('unidentified') and r['unidentified']=='simple':v=simple.get(t)
            if amount==0:terms.append(0.)
            elif finite(amount) and finite(v):terms.append(amount*v)
            else:unknown=True
        ex=None if unknown else math.fsum(terms);assert close(r[year],ex),(r['scope'],r['method'],year,r[year],ex)
        checked+=1
npp={r['method']:r for r in records(b,'NPP','NPP')}
annual={(r['model_id'],r['scope'],r['method'],r['catch_basis'],r['unidentified']):r for r in records(b,'PPR','Annual') if r['metric']=='ppr'}
for r in records(b,'PPR–NPP','Ratios'):
    if not r['model_id']:continue
    rr=annual[r['model_id'],r['scope'],r['method'],r['catch_basis'],r['unidentified']]
    den=npp[r['npp_method']]
    for year in YEARS:
        ex=100*rr[year]/9/den[year] if finite(rr[year]) and finite(den.get(year)) and den[year]>0 else None
        assert close(r[year],ex)
assert sha(MODEL/'results/regional_snapshot.xlsx')==sha(ROOT/'regions/LME/LME_013/LME_013.xlsx')
print(f'Regional values reopened and independently reconciled for {checked} annual PPR cells, taxon coefficients and NPP ratios.',flush=True)
# Reopen bounded central integration and unchanged human/source metadata.
sheets=['Regions & status','Models & coverage','Papers','Diagnostics & sensitivity']
project=read_book(ROOT/'Project.xlsx',sheets=sheets)
prev=read_book(RUN/'inputs/Project.xlsx',sheets=['Models & coverage','Papers'])
for sheet in ['Models & coverage','Papers']:
    for name in prev[sheet]:assert digest_tables([project[sheet][name]])==digest_tables([prev[sheet][name]])
region=next(r for r in records(project,'Regions & status','Regions') if r['unit_id']=='LME_013')
assert region['sha256']==sha(ROOT/'regions/LME/LME_013/LME_013.xlsx')
assert region['production_eligible'] is False
appendix=MODEL/'model_validation/taxon_mapping.xlsx'
wa=openpyxl.load_workbook(appendix);wb=openpyxl.load_workbook(RUN/'inputs/taxon_mapping.xlsx')
av=[list(r) for r in wa['Taxon mappings'].iter_rows(min_row=8,max_col=7,values_only=True)]
bv=[list(r) for r in wb['Taxon mappings'].iter_rows(min_row=8,max_col=7,values_only=True)]
assert len(av)==218 and digest_tables([av])==digest_tables([bv])
assert wa['Taxon mappings'].tables['CandidateMappings'].autoFilter.ref=='A7:G225'
wa.close();wb.close()
links=[]
for filename in ['validation.docx','source_value_corrections.docx']:
    p=MODEL/'model_validation'/filename
    with zipfile.ZipFile(p) as z:
        rel=etree.fromstring(z.read('word/_rels/document.xml.rels'))
        targets=[x.get('Target') for x in rel if x.get('Type').endswith('/hyperlink')]
        for target in targets:
            if target.startswith(('http:','https:','mailto:')):continue
            local=(p.parent/unquote(target.split('#')[0])).resolve()
            assert local.exists(),(filename,target)
            links.append([filename,target])
        if filename=='source_value_corrections.docx':
            text=' '.join(etree.fromstring(z.read('word/document.xml')).itertext())
            assert 'json' not in text.lower()
            assert all('/sources/' in t for t in targets) and len(targets)==2
baseline=load(RUN/'inputs/baseline_hashes.json')['map_files']
now={p.relative_to(ROOT).as_posix():sha(p) for p in (ROOT/'interactive_map').rglob('*') if p.is_file()}
assert baseline.keys()==now.keys()
map_changed=[k for k in baseline if baseline[k]!=now[k]]
assert all(k in ['interactive_map/index.html','interactive_map/trends.html'] for k in map_changed)
map_text=(ROOT/'interactive_map/index.html').read_text(encoding='utf-8')
assert sha(MODEL/'model.json') not in map_text
assert load(RUN/'inputs/baseline_hashes.json')[str((MODEL/'model.json').relative_to(ROOT).as_posix())] in map_text
dump(RUN/'qa/map_deferred.json',dict(map_builder_called=False,corrected_model_not_published=True,
 baseline_and_current_data_files_identical=True,HTML_files_changed_independently_during_run=map_changed,
 baseline_sha256=baseline,current_sha256=now))
source=load(MODEL/'model_validation/work/2026-10-08_123302_source_corrections/outputs/correction_scenarios.json')
fresh_source_hash=next(v for k,v in source['protected_input_hashes'].items() if k.endswith('Supplementary material revised and final.xls'))
assert fresh_source_hash==sha(MODEL.parents[1]/'sources/Supplementary material revised and final.xls')
dump(RUN/'qa/final_verification.json',dict(status='PASS',model_diff=diff,exact_notebook_balanced=True,
 six_full_matrices_verified=True,negative_pair_count=len(negatives),all_scoped_coefficients_verified=True,
 annual_PPR_cells_checked=checked,regional_taxon_coefficients_NPP_ratios_verified=True,
 independent_inputs_and_218_mapping_rows_preserved=True,central_source_and_human_metadata_preserved=True,
 hyperlink_targets=links,corrections_document_references_only_sources=True,map_data_files_unchanged=True,
 corrected_model_not_published_to_map=True,independently_changed_map_HTML=map_changed,
 all_SPPR_overall_FAIL_retained=True,pooled_results_reported=True,Monte_Carlo='NOT_RUN',researcher_review='pending'))
print('PASS: central metadata, 218 mapping rows, local report references and source supplement; corrected model not published to map.',flush=True)
