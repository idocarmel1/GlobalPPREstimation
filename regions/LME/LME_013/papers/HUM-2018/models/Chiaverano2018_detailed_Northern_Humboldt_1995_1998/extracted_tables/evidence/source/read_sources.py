from pathlib import Path
import sys,json,hashlib,zipfile
from collections import Counter
ROOT=Path(__file__).resolve().parent
REPO=Path.cwd()
PAPERS=REPO/'regions/LME_013/papers/HUM-2018'
OUT=ROOT/'evidence'
OUT.mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/'vendor'))
def dump(name,data):
    (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
inventory=[]
for name in ['1-s2.0-S0079661117303312-main.pdf','Supplementary material revised and final.xls','Peru paper figures and tables II.docx']:
    p=PAPERS/name
    inventory.append({'filename':name,'relative_path':'../../../papers/HUM-2018/'+name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})
dump('original_sources.json',inventory)
import pdfplumber
with pdfplumber.open(PAPERS/'1-s2.0-S0079661117303312-main.pdf') as pdf:
    pages=[]
    for i,p in enumerate(pdf.pages,1):
        txt=p.extract_text() or ''
        (OUT/f'pdf-page-{i:02}.txt').write_text(txt,encoding='utf-8')
        pages.append({'pdf_page':i,'printed_page':i+27,'width':p.width,'height':p.height,'text':txt,'words':p.extract_words()})
        if i in [2,3,4]:
            p.to_image(resolution=200).save(str(OUT/f'pdf-page-{i:02}.png'))
    dump('pdf_pages_words.json',pages)
import docx
d=docx.Document(PAPERS/'Peru paper figures and tables II.docx')
dump('docx_content.json',{'paragraphs':[p.text for p in d.paragraphs], 'tables':[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]})
with zipfile.ZipFile(PAPERS/'Peru paper figures and tables II.docx') as z:
    media=[]
    for n in z.namelist():
        if n.startswith('word/media/'):
            p=OUT/Path(n).name
            p.write_bytes(z.read(n))
            media.append({'source_zip_member':n,'local_file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    dump('docx_media.json',media)
try:
    import xlrd
except ImportError:
    print('PDF and DOCX extracted; xlrd absent.')
    sys.exit(2)
b=xlrd.open_workbook(PAPERS/'Supplementary material revised and final.xls',formatting_info=True)
sheets=[]
from openpyxl.utils import get_column_letter
for s in b.sheets():
    cells=[]
    rows=[]
    for r in range(s.nrows):
        row=[]
        for c in range(s.ncols):
            cell=s.cell(r,c)
            xf=b.xf_list[cell.xf_index]
            fmt=b.format_map[xf.format_key].format_str
            font=b.font_list[xf.font_index]
            val=None if cell.ctype in [xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK] else cell.value
            row.append(val)
            cells.append({'row':r+1,'column':c+1,'address':f'{get_column_letter(c+1)}{r+1}','value':val,'cell_type':cell.ctype,'number_format':fmt,'bold':bool(font.bold),'italic':bool(font.italic)})
        rows.append(row)
    sheets.append({'name':s.name,'rows':s.nrows,'columns':s.ncols,'cells':cells,'values':rows})
    print(s.name,s.nrows,s.ncols)
dump('supplement_cells.json',{'source':'../../../papers/HUM-2018/Supplementary material revised and final.xls','source_sha256':inventory[1]['sha256'],'extraction_method':'xlrd stored BIFF8 cell values and source number formats/styles; no recalculation or source edits','sheets':sheets})
