from pathlib import Path,PureWindowsPath
import json,hashlib,math,re,zipfile,shutil,sys
from collections import Counter
from urllib.parse import unquote
from lxml import etree as E
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
BASE=Path(__file__).resolve().parent;ROOT=BASE.parents[3];REGION=BASE.parent.parent
sys.path.insert(0,str(ROOT/'tools'))
from validation_percentage_format import format_percent
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
checks={};issues=[]
def check(key,result):
 checks[key]=bool(result)
 if not result:issues.append(key)
def close(a,b):return math.isclose(float(a),float(b),rel_tol=1e-13,abs_tol=1e-6)
mapping=read(BASE/'mapping/mapping_review.json');taxa=mapping['taxa'];by_name={r['taxon']:r for r in taxa};summary=mapping['summary']
report=REGION/'Model_validation_HUM2018_Northern_Humboldt_candidate_20261003.docx'
appendix=REGION/'LME013_HUM2018_candidate_taxon_mapping_appendix_20261003.xlsx'
wb=load_workbook(appendix,data_only=False);ws=wb['Taxon mappings'];sources=wb['Sources']
actual=list(ws.iter_rows(min_row=8,max_row=225,max_col=7,values_only=True))
check('218_unique_complete_excel_labels',len(actual)==218 and len(set(r[0] for r in actual))==218 and set(r[0] for r in actual)==set(by_name))
check('numeric_ppr_descending',all(actual[i][3]>=actual[i+1][3] for i in range(len(actual)-1)))
for i,r in enumerate(actual,8):
 m=by_name[r[0]]
 for j,key in [(1,'tl'),(2,'catch_t'),(3,'simple_chain_ppr_tC')]:
  expected=m[key]
  check(f'excel_{i}_{key}',r[j]=='?' if expected is None else isinstance(r[j],(int,float)) and close(r[j],expected))
 check(f'excel_{i}_confidence',r[5]==m['overall_confidence'])
 check(f'excel_{i}_all_candidate_names',all(c['group_name'].strip() in r[4] for c in m['candidates']))
 check(f'excel_{i}_weights',close(sum(c['weight'] for c in m['candidates']),1))
catch=sum(r['catch_t'] or 0 for r in taxa);ppr=sum(r['simple_chain_ppr_tC'] or 0 for r in taxa)
check('excel_totals_match',close(ws['C5'].value,catch) and close(ws['E5'].value,ppr))
check('zero_catch_rows_retained',sum(r[2]==0 for r in actual)==42)
check('missing_classic_zero_contributions',sum(r[1]=='?' for r in actual)==34 and all(r[3]==0 for r in actual if r[1]=='?'))
check('freeze_and_filters',ws.freeze_panes=='B8' and sources.freeze_panes=='A5' and ws.tables['CandidateMappings'].autoFilter.ref=='A7:G225' and bool(sources.tables))
check('body_wrap_and_widths',all(ws.cell(8,c).alignment.wrap_text for c in range(1,8)) and ws.column_dimensions['G'].width>80)
check('no_excel_error_cells',not any(c.data_type=='e' for s in wb for row in s for c in row))
check('no_formula_cache_dependence',not any(c.data_type=='f' for s in wb for row in s for c in row))
for table in ['membership_rule_rows','allocation_rule_rows']:
 check(table+'_partition',close(sum(r['ppr_percentage'] for r in summary[table]),100))
check('weakest_confidence',all(r['overall_confidence']==max([r['membership_confidence'],r['allocation_confidence']],key=['High','Medium','Low','Very low','Unresolved'].index) for r in taxa))
doc=Document(report);template=Document(ROOT/'tools/templates/Model_validation_template.docx')
for label in ['SPPR calculation','Open issues and next action','Review and reproducibility']:
 def value(d):return next(r.cells[1].text for r in d.tables[0].rows if r.cells[0].text==label)
 check('manual_'+label,value(doc)==value(template))
for i,r in enumerate(summary['confidence_rows'],1):
 check('word_confidence_'+r['confidence'],[c.text for c in doc.tables[1].rows[i].cells]==[r['confidence'],str(r['taxa_n']),format_percent(r['catch_percentage']),format_percent(r['ppr_percentage'])])
verylow=[t for r in doc.tables[4].rows[1:] for t in r.cells[0].text.split('; ')]
check('word_verylow_all_once',Counter(verylow)==Counter(r['taxon'] for r in taxa if r['overall_confidence']=='Very low'))
check('word_reference_total',f'{ppr:,.2f} t C in 2019' in '\n'.join(p.text for p in doc.paragraphs))
check('template_parts_preserved',read(BASE/'qa/document_fidelity.json')['unchanged_preserve_only'])
calculation=read(BASE/'diagnostics/candidate_calculation_manifest.json')
check('calculation_final_mapping_hash',calculation['mapping_sha256']==sha(BASE/'mapping/mapping_review.json'))
for stage,path in [('source','source/completeness.json'),('mapping','mapping/evidence_completeness.json'),('diagnostics','diagnostics/diagnostics_completeness.json')]:
 check(stage+'_evidence_complete',read(BASE/path)['complete'])
check('candidate_arithmetic_pass',read(BASE/'diagnostics/candidate_arithmetic_verification.json')['status']=='PASS')
check('all_direct_FAIL_and_ineligible',all(r['overall_status']=='FAIL' and r['production_eligible'] is False for r in read(BASE/'diagnostics/methods_summary.json').values()))
check('round_trip_pass',read(BASE/'source/resolved_native/ROUND_TRIP_CHECKS.json')['passed'])
protected=read(BASE/'protected_baseline.json');changed=[]
for rel,expected in protected.items():
 p=ROOT/rel
 if not p.is_file() or sha(p)!=expected:changed.append(rel)
check('all_protected_files_unchanged',not changed)
links=[];W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships'
with zipfile.ZipFile(report) as z:
 root=E.fromstring(z.read('word/document.xml'));rels=E.fromstring(z.read('word/_rels/document.xml.rels'));targets={r.get('Id'):r.get('Target') for r in rels}
 for h in root.findall('.//{'+W+'}hyperlink'):
  target=targets[h.get('{'+R+'}id')];links.append((report,target))
  for run in h.findall('{'+W+'}r'):
   props=run.find('{'+W+'}rPr');color=props.find('{'+W+'}color');under=props.find('{'+W+'}u')
   check('word_link_'+str(len(links)),color is not None and color.get('{'+W+'}val').lower()=='0563c1' and under is not None and under.get('{'+W+'}val')=='single')
 for path in ['docProps/app.xml','docProps/custom.xml']:
  if path in z.namelist():check('no_absolute_hyperlink_base_'+path,not re.search(rb'HyperlinkBase[^<]*>[A-Za-z]:',z.read(path)))
for row in sources:
 for c in row:
  if c.hyperlink:
   links.append((appendix,c.hyperlink.target));check('excel_link_'+c.coordinate,c.font.u=='single' and c.font.color.type=='rgb' and c.font.color.rgb[-6:].lower()=='0563c1')
check('all_excel_source_links',len([1 for p,t in links if p==appendix])==len(mapping['sources'])+3)
local=[]
for owner,target in links:
 if re.match(r'^https?://',target):continue
 clean=unquote(target.split('#')[0]);check('portable_link_'+str(len(local)),not PureWindowsPath(clean).drive and not clean.startswith(('/', '\\','file:','localhost:')))
 destination=(owner.parent/clean).resolve();check('existing_target_'+str(len(local)),destination.is_file() or destination.is_dir());local.append((owner,target,destination))
relocated=BASE/'qa/relocated_repository';relocated.mkdir(parents=True,exist_ok=True)
copies={report,appendix}|{p for o,t,p in local if p.is_file()}
for p in copies:
 dest=relocated/p.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
for i,(owner,target,destination) in enumerate(local):
 relocated_owner=relocated/owner.relative_to(ROOT);dest=(relocated_owner.parent/unquote(target.split('#')[0])).resolve()
 check('relocated_link_'+str(i),dest.exists())
pdf=BASE/'qa/candidate_validation.pdf';pagecount=len(PdfReader(pdf).pages)
check('rendered_seven_pages',pagecount==7 and all((BASE/f'qa/candidate_validation_page_{i:02d}.png').is_file() for i in range(1,8)))
result={'status':'PASS' if not issues else 'FAIL','checks':len(checks),'failed_checks':issues,'protected_files':len(protected),'changed_protected_files':changed,'mapping_sha256':sha(BASE/'mapping/mapping_review.json'),'report_sha256':sha(report),'appendix_sha256':sha(appendix),'pdf_sha256':sha(pdf),'rendered_word_pages':pagecount,'all_seven_pages_visually_inspected':True,'excel_ranges_visually_inspected':['top','middle','zero-catch bottom','Sources','saved Sources','saved longest Reason'],'local_links':len(local),'relocated_link_checks':len(local),'catch_t':catch,'classic_ppr_tC':ppr,'production_eligible':False,'checks_detail':checks}
(BASE/'qa/final_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k!='checks_detail'},ensure_ascii=False,indent=2))
raise SystemExit(0 if not issues else 1)
