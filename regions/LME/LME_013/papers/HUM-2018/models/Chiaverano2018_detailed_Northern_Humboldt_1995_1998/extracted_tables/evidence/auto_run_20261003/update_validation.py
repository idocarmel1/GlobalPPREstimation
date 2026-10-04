from pathlib import Path
import json,copy,hashlib,zipfile,os,tempfile
from docx import Document
from docx.oxml.ns import qn
OUT=Path(__file__).resolve().parent;BASE=OUT.parent;ROOT=BASE.parents[3];REGION=ROOT/'regions/LME_013'
MODEL=BASE.name;report=REGION/'Model_validation_HUM2018_Northern_Humboldt_candidate_20261003.docx'
d=Document(report);table=d.tables[0]
protected_labels=['SPPR calculation','Open issues and next action','Review and reproducibility']
manual={r.cells[0].text:r._tr.xml for r in table.rows if r.cells[0].text in protected_labels}
independent_tables=[t._tbl.xml for t in d.tables[1:5]]
diag={m:json.loads((OUT/m.replace(' ','_')/'diagnostic_return.json').read_text(encoding='utf-8')) for m in ['GE','TE','With Egestion']}
pool=json.loads((OUT/'never_vs_auto_comparison.json').read_text(encoding='utf-8'))
def replace_text(p,text):
    props=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();run=p.add_run(text)
    if props is not None:run._r.insert(0,props)
def metrics(m):
    r=diag[m];v=r['divergence']
    text=[f"FAIL overall: input {r['model_input']['status']}; convergence {v['status']}; SPPR balance {r['balance']['status']}.",
      'No negative SPPR entries across any basal-source column or recipient, including unfished groups.',f"rho_living = {v['rho_living']:.4f}"]
    if m=='GE':
        p=pool[m]['detritus_resolution_info'];text.append(f"Unpooled recycling b = {v['b']:.4f}; it does not converge. Auto pools four detritus columns; pooled b = {p['collapse_b']:.4f}.")
    text.append('Detritus SPPR: '+ '; '.join(f"{next(g for g in pool['GE']['detritus_resolution_info']['det_names'] if pool['GE']['detritus_resolution_info']['det_names'].index(g)==[36,37,38,39].index(int(k)))} ({k}) = {x:.4f}" for k,x in v['sppr_det'].items())+'.')
    text.append(f"Strict model balance and SPPR balance both fail. PP gap {r['balance']['rel_gap']*100:.1f}%.")
    if m=='TE':text.append('The engine cannot recredit EE=0 losses with four operational detritus pools; auto does not change this TE formulation.')
    return '\n'.join(text)
for row in table.rows:
    label=row.cells[0].text;c=row.cells[1]
    if label=='Candidate article':replace_text(row.cells[0].paragraphs[0],'Selected article')
    if label=='Candidate model':
        replace_text(row.cells[0].paragraphs[0],'Selected model')
        replace_text(c.paragraphs[0],MODEL+'; selected audited detailed Northern Humboldt baseline, 1995–1998. '+c.paragraphs[0].text.split('1995–1998. ',1)[1])
    elif label=='Region':replace_text(c.paragraphs[0],c.paragraphs[0].text.replace('Candidate evaluation','Selected-model evaluation'))
    elif label=='Other known articles':replace_text(c.paragraphs[0],c.paragraphs[0].text.replace('Its static 1980 model remains selected.','Its static 1980 model is retained as previous-selection evidence.'))
    elif label=='Selection rationale':replace_text(c.paragraphs[0],c.paragraphs[0].text.replace('Requested candidate review','Requested source review').replace('The active selection remains Chilean Patagonia 1980.','The user selected this audited model after auto returned finite, nonnegative SPPR in all three methods. Selection does not establish scientific validity or researcher validation.'))
    elif label=='GE diagnostics':replace_text(c.paragraphs[0],metrics('GE'))
    elif label=='TE diagnostics':replace_text(c.paragraphs[0],metrics('TE'))
    elif label=='Other':
        r=diag['With Egestion'];p=pool['With Egestion']['detritus_resolution_info']
        replace_text(c.paragraphs[0],f"All direct calls explicitly use det_collapse_mode='auto' with the audited constructor, normalization, EE=0 fix and closed-detritus settings unchanged. With Egestion remains overall FAIL: input FAIL, convergence FAIL, SPPR balance WARN; rho_living = {r['divergence']['rho_living']:.4f}; unpooled b = {r['divergence']['b']:.4f}; pooled b = {p['collapse_b']:.4f}; pooled detritus SPPR = {p['collapse_scalar']:.4f}; PP gap {r['balance']['rel_gap']*100:.1f}%. All its SPPR entries are nonnegative, including unfished groups. Auto eliminates the previous negative coefficients but does not resolve source production conflicts, native-flow translation limits or strict balance failures. Selected-model regional arithmetic is displayed provisionally with production eligibility false, pending researcher review.")
replace_text(d.paragraphs[0],'Northern Humboldt model validation')
for rel in d.part.rels.values():
    if rel.reltype.endswith('/hyperlink') and rel.is_external:
        target=rel.target_ref.replace('candidate_studies/HUM2018_20261003/',f'models/{MODEL}/')
        for suffix in ['GE/diagnostic_return.json','GE/SPPR.csv','TE/diagnostic_return.json','TE/SPPR.csv','full_direct_report.json','full_direct_report.md','candidate_calculation_manifest.json']:
            target=target.replace('/diagnostics/'+suffix,'/auto_run_20261003/'+suffix)
        rel._target=target
assert manual=={r.cells[0].text:r._tr.xml for r in table.rows if r.cells[0].text in protected_labels}
assert independent_tables==[t._tbl.xml for t in d.tables[1:5]]
d.save(report)
reopened=Document(report)
assert manual=={r.cells[0].text:r._tr.xml for r in reopened.tables[0].rows if r.cells[0].text in protected_labels}
assert independent_tables==[t._tbl.xml for t in reopened.tables[1:5]]
# Rebase only Office link relationships; preserve all seven appendix columns,
# full-precision ordering, numbers, source labels, styles and confidence tables.
appendix=REGION/'LME013_HUM2018_candidate_taxon_mapping_appendix_20261003.xlsx'
with zipfile.ZipFile(appendix) as z:
    parts=[(i,z.read(i.filename)) for i in z.infolist()]
fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=REGION);os.close(fd)
with zipfile.ZipFile(tmp,'w') as z:
    for info,data in parts:
        if info.filename.endswith('.rels'):
            data=data.replace(b'candidate_studies/HUM2018_20261003/',('models/'+MODEL+'/').encode())
        z.writestr(info,data)
os.replace(tmp,appendix)
(OUT/'document_preservation.json').write_text(json.dumps({'manual_rows_preserved_exact_OOXML':protected_labels,'confidence_and_rule_tables_OOXML_unchanged':True,'appendix_changes':'hyperlink targets only; worksheet and style parts byte-identical','no_researcher_verdict_registered':True,'report_sha256':hashlib.sha256(report.read_bytes()).hexdigest()},indent=2),encoding='utf-8')
print('Validation report updated locally; researcher fields and independent confidence tables preserved exactly.',flush=True)
