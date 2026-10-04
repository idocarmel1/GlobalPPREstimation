from pathlib import Path
import json,csv,hashlib,math
from decimal import Decimal
import openpyxl,numpy as np
R=Path(__file__).resolve().parent;N=R/'resolved_native';E=R/'evidence';C=R/'computational'
def jw(p,d):p.write_text(json.dumps(d,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
x=json.loads((N/'extraction.json').read_text(encoding='utf-8'));m=json.loads((N/'model.json').read_text(encoding='utf-8'));s=json.loads((E/'supplement_cells.json').read_text(encoding='utf-8'));sheets={a['name']:a for a in s['sheets']}
checks={};failures=[]
def check(label,ok):checks[label]=bool(ok);failures.extend([] if ok else [label])
def eq(a,b):return (a is None and b is None) or a==b or (a is not None and b is not None and float(a)==float(b))
for filename,cols in [('Basic_input.csv',{'biomass':3,'pb':5,'qb':6,'ee':7,'pq':9,'unassim':10}),('Biomass_accumulation.csv',{'ba':2,'ba_rate':3})]:
    rows=list(csv.reader((N/filename).open(encoding='utf-8',newline='')))
    for i,g in enumerate(x['groups'],1):
        check(f'{filename}.group_name.{i}',rows[i][1]==g['name'])
        for field,c in cols.items():
            a=rows[i][c];b=g[field];check(f'{filename}.{field}.{i}',a=='' if b is None else Decimal(a)==Decimal(str(b)))
diet=list(csv.reader((N/'Diet_composition.csv').open(encoding='utf-8',newline='')))
for c,n in enumerate(x['consumers'],2):
    for p in range(1,40):check(f'diet.{n}.{p}',diet[p][c]==x['diet'][str(n)][str(p)])
    check(f'diet.import.{n}',diet[40][c]==x['diet'][str(n)]['import'])
for n,g in enumerate(m['group'],1):
    a=x['groups'][n-1];check(f'canonical.name.{n}',a['name']==g['group_name'])
    for field,mfield in [('biomass','biomass'),('pb','pb'),('qb','qb'),('ee','ee'),('unassim','gs')]:
        check(f'canonical.{field}.{n}',g[mfield]=='-9999' if a[field] is None else Decimal(g[mfield])==Decimal(str(a[field])))
    for p in range(1,40):
        d=g['diet_descr']['diet'][p-1];check(f'canonical.diet.{n}.{p}',d['proportion']==x['diet'][str(n)][str(p)])
        f=x['detritus_fate'][str(n)].get(x['groups'][p-1]['name'],'0');check(f'canonical.fate.{n}.{p}',d['detritus_fate']==f)
for name in ['Taxonomy.xlsx','TL.xlsx','Metadata.xlsx','reconstructed.xlsx']:
    wb=openpyxl.load_workbook(N/name,data_only=False)
    if name=='Taxonomy.xlsx':
        ws=wb.active;check('Taxonomy.three_columns',ws.max_column==3);check('Taxonomy.39_rows',ws.max_row==40)
        for i,g in enumerate(x['groups'],2):check(f'Taxonomy.identity.{i-1}',[ws.cell(i,j).value for j in [1,2,3]]==[g['n'],g['name'],g['taxon_descr']])
    if name=='TL.xlsx':check('TL.all_source_TL_unknown',all(wb.active.cell(i,3).value is None for i in range(2,41)))
    if name=='reconstructed.xlsx':
        ws=wb['Groups']
        for i,g in enumerate(m['group'],2):
            check(f'reconstructed.name.{i-1}',ws.cell(i,2).value==g['group_name'])
            for col,f in enumerate(['biomass','pb','qb','ee','ge','gs','pp','diet_imp','export','habitat_area','biomass_accum','biomass_accum_rate'],3):
                v=ws.cell(i,col).value; src=g[f];check(f'reconstructed.{f}.{i-1}',v is None if src=='-9999' else v==float(src))
        ws=wb['Diet']
        for col,n in enumerate(x['consumers'],3):
            for p in range(1,40):check(f'reconstructed.diet.{n}.{p}',ws.cell(p+1,col).value==float(x['diet'][str(n)][str(p)]))
check('all_original_source_hashes_unchanged',all(sha(Path.cwd()/'regions/LME_013/papers/HUM-2018'/a['filename'])==a['sha256'] for a in json.loads((E/'original_sources.json').read_text(encoding='utf-8'))))
jw(N/'ROUND_TRIP_CHECKS.json',{'passed':not failures,'checks':len(checks),'failures':failures,'exact_checks':checks,'source_layers':['original stored cell values','six EwE CSV tables plus two XLSX imports','canonical source model JSON','reconstructed workbook'],'computed_note':'Native IEEE binary values retained via decimal repr; numeric Excel re-export checked at float equality. Format evidence separate.'})
# Independent arithmetic, not a source repair and not a calculator diagnosis.
cons=np.array([float(g['biomass'])*float(g['qb']) if g['qb'] is not None else 0 for g in x['groups']]); dc=np.array([[float(x['diet'][str(c+1)][str(p+1)]) for p in range(39)] for c in range(39)]); pred=cons@dc
balances=[]
for g in x['groups'][:35]:
    n=g['n'];nn=str(n);prod=float(g['biomass'])*float(g['pb']); catch=float(sum(Decimal(x['landings'][nn][f])+Decimal(x['discards'][nn][f]) for f in x['fleets'])); used=prod*float(g['ee']);needed=catch+pred[n-1]
    balances.append({'seq':n,'group_name':g['name'],'source_production_t_km2_y':prod,'source_consumption_t_km2_y':cons[n-1],'source_EE':float(g['ee']),'raw_source_predation_t_km2_y':pred[n-1],'source_total_catch_t_km2_y':catch,'EE_implied_without_BA_or_migration':needed/prod,'EE_difference':needed/prod-float(g['ee']),'used_production_minus_catch_and_predation':used-needed,'source_BA':'unknown','arithmetic_BA_assumption':0,'native_diet_normalized':False})
jw(N/'SOURCE_BUDGET_ARITHMETIC.json',{'convention':'Source raw diet and stored B/PB/QB/EE; indicative steady-state identity BA=netmigration=0 only for arithmetic; no field changed','groups':balances})
# Geometric reconstruction of printed aggregate Table1; exact two-decimal text, blanks retained.
pages=json.loads((E/'pdf_pages_words.json').read_text(encoding='utf-8'));words=pages[3]['words']; anchors=[43.6,99.4,209.9,266.3,309.0,355.2,390.7,426.2,461.7,518.1]
names=['Large phytoplankton','Small phytoplankton','Microzooplankton','Mesozooplakton','Macrozooplankton','Small jellyfish','Large jellyfish','Macrobenthos','Forage fish','Mesopelagics','Cephalopods','Pelagic planktivorous fish','Pelagic piscivorous fish','Demersal piscivorous fish','Demersal planktivorous fish','Demersal benthivorous fish','Apex predatory fish','Seabirds','Marine mammals','Sea turtles','Fish eggs','Detritus offal','Pelagic detritus','Benthic detritus','Fisheries']
rowanchors=[w for w in words if 108<w['top']<320 and abs(w['x0']-43.6)<.6 and w['text'].isdigit()]
agg=[]
fields=['source_code','group_name','biomass','pb','qb','pq','ae','ee','landings','discards']
for j,a in enumerate(rowanchors):
    values=[None]*10
    for w in words:
        if abs(w['top']-a['top'])<2:
            k=min(range(10),key=lambda z:abs(w['x0']-anchors[z]))
            if abs(w['x0']-anchors[k])<3:values[k]=w['text']
    values[1]=names[j]
    agg.append(dict(zip(fields,values),pdf_page=4,printed_page=31,table='Table 1',row_y=a['top']))
jw(R/'aggregated_Table1_printed.json',agg)
with (R/'aggregated_Table1_printed.csv').open('w',encoding='utf-8',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(agg[0]));w.writeheader();w.writerows(agg)
# Compare source names independent of conflicting numeric aggregation codes.
crosswalk=[]
for rowidx,row in enumerate(sheets['Table E-trophic aggregations']['values'][7:],8):
    matching=[g for g in x['groups'] if g['name']==row[0]]
    crosswalk.append({'resolved_source_id':matching[0]['n'] if matching else 40 if row[0]=='Artisanal fisheries' else 41 if row[0]=='Commercial fisheries' else None,'resolved_source_name':row[0],'aggregated_source_name':row[1],'TableE_source_code':row[2],'TableE_cell':f'A{rowidx}:C{rowidx}'})
jw(R/'AGGREGATION_CROSSWALK.json',crosswalk)
sumdiscard=sum(Decimal(v) for d in x['discards'].values() for v in d.values());sumland=sum(Decimal(v) for d in x['landings'].values() for v in d.values())
jw(C/'DISCARD_FLOW_TRANSLATION_LIMIT.json',{'fleet_nodes_excluded':[40,41],'source_detritus_destination':37,'source_destination_name':'Fishery offal','fleet_return_proportion':1,'total_source_discards_t_km2_y':str(sumdiscard),'source_total_landings_t_km2_y':str(sumland),'computational_removals_include_discards':True,'discard_return_represented':False,'external_detritus_import_invented':False,'limitation':'The simplified stock JSON preserves donor removals but cannot retain source internal fleet discard-return ancestry. Source fleet fate evidence retained. Scientific production use ineligible until representation resolved.'})
identity={'schema_version':1,'paper_id':'HUM-2018','title':'Evaluating the role of large jellyfish and forage fishes as energy pathways, and their interplay with fisheries, in the Northern Humboldt Current System','historical_short_title_alias':'Large jellyfish and forage fishes as energy pathways in the Northern Humboldt Current System','doi':'10.1016/j.pocean.2018.04.009','publication_year':2018,'static_base_period':'1995–1998 average','dynamic_predecessor_period_1995_2004':'Not the focal static baseline; earlier cached metadata conflates the cited dynamic predecessor','domain':{'latitude_south':[4,16],'offshore_km':111,'area_km2':165000,'source':'PDFp2 Figure1; PDFp3/printedp30 Methods'},'source_table_identity':{'A':'Fully resolved balanced ECOPATH parameters','B':'resolved diet','C':'resolved Ecopath fate','D':'resolved uncertainty CV','E':'aggregation crosswalk','F_sheet_caption_G':'aggregated diets','G_sheet_caption_H':'ECOTRAN separately divided feces/senescence/excretion fate','H_sheet_caption_I':'ECOTRAN production matrix including nutrients','I_sheet_caption_J':'aggregated uncertainty CV'},'source_bundle_hashes':json.loads((E/'original_sources.json').read_text(encoding='utf-8')),'predecessor':{'title':'Trophic modeling of the Northern Humboldt Current Ecosystem, Part I: Comparing trophic linkages under La Niña and El Niño conditions','doi':'10.1016/j.pocean.2008.10.007','url':'https://epic.awi.de/id/eprint/22464/1/Tam2008a.pdf','file':'evidence/Tam2008a.pdf','sha256':sha(E/'Tam2008a.pdf'),'verified_lineage':'Focal Methods explicitly averaged Tam et al. 2008 groups/parameters for 1995–1998; membership inherited with crosswalk, numerical parameters not copied.'}}
jw(R/'SOURCE_IDENTITY.json',identity)
largest=sorted(balances,key=lambda a:abs(a['EE_difference']),reverse=True)[:8]
print('ROUND TRIP',not failures,len(checks),'failures',failures[:10]);print('Largest EE disagreements',[(a['seq'],a['EE_implied_without_BA_or_migration'],a['source_EE']) for a in largest]);print('Total discards',sumdiscard,'landings',sumland,'aggregate rows',len(agg))
