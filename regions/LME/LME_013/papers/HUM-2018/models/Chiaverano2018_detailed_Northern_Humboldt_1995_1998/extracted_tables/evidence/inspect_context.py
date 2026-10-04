from pathlib import Path
import sys,json,hashlib,zipfile
from docx import Document
import openpyxl
ROOT=Path(__file__).resolve().parents[4]
OUT=Path(__file__).resolve().parent
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
protected=[ROOT/'Project.xlsx',ROOT/'regions/LME_013/LME_013.xlsx',ROOT/'regions/LME_013/LME013_taxon_mapping_appendix.xlsx',ROOT/'regions/LME_013/Model_validation_13_1_Chilean_Patagonia_(1980).docx',ROOT/'tools/templates/Model_validation_template.docx',ROOT/'tools/knowledge_graph/graph.json',ROOT/'interactive_map/index.html',ROOT/'interactive_map/trends.html',ROOT/'interactive_map/archive/index.html']
protected+=list((ROOT/'regions/LME_013/models').rglob('*'))
protected=[p for p in protected if p.is_file()]
baseline=OUT/'protected_baseline.json'
if not baseline.exists(): baseline.write_text(json.dumps({str(p.relative_to(ROOT)):sha(p) for p in protected},indent=2),encoding='utf-8')
wb=openpyxl.load_workbook(ROOT/'regions/LME_013/LME_013.xlsx',read_only=True,data_only=True)
context={'sheets':wb.sheetnames,'overview':list(wb['Overview'].values)}
for name in ['NPP','Diagnostics']:
 context[name]=list(wb[name].values)
wb.close()
central=openpyxl.load_workbook(ROOT/'Project.xlsx',read_only=True,data_only=True)
context['central_rows']={s:[list(r) for r in central[s].values if any('HUM-2018' in str(v) or 'Chiaverano' in str(v) or '13_2_Northern' in str(v) for v in r)] for s in ['Papers','Models & coverage']}
central.close()
selected=Document(ROOT/'regions/LME_013/Model_validation_13_1_Chilean_Patagonia_(1980).docx')
context['selected_report']={'paragraphs':[p.text for p in selected.paragraphs], 'tables':[[[c.text for c in r.cells] for r in t.rows] for t in selected.tables]}
template=ROOT/'tools/templates/Model_validation_template.docx'
d=Document(template)
context['template']={'paragraphs':[(i,p.style.name,p.text) for i,p in enumerate(d.paragraphs)],'tables':[[[c.text for c in r.cells] for r in t.rows] for t in d.tables],'sections':[{'width':s.page_width.inches,'height':s.page_height.inches,'left':s.left_margin.inches,'right':s.right_margin.inches,'top':s.top_margin.inches,'bottom':s.bottom_margin.inches} for s in d.sections]}
(OUT/'context.json').write_text(json.dumps(context,indent=2,ensure_ascii=False,default=str),encoding='utf-8')
print(json.dumps({'overview':context['overview'],'sheets':context['sheets'],'central_rows':context['central_rows'],'template':context['template'],'selected_core':context['selected_report']['tables'][0]},ensure_ascii=False,default=str))
