"""Source-preserving extraction and independently checked author linear solve."""
from pathlib import Path
import csv, json, hashlib, sys, re
from decimal import Decimal
import numpy as np
import pymupdf
from docx import Document
from openpyxl import load_workbook, Workbook

ROOT=Path(__file__).resolve().parents[3]
REVIEW=Path(__file__).parent
PAPER=ROOT/'regions/LME_003/papers/CAL-2016'
SRC=PAPER/'author_archive_extracted/CalCurFoodWebModelECOMOD-master'
MODEL=ROOT/'regions/LME_003/models/CAL-2016_California_Current_2000-2014'
OUT=MODEL/'extracted_tables'
OUT.mkdir(parents=True,exist_ok=True)
def dump(path,x):path.write_text(json.dumps(x,indent=2,ensure_ascii=False),encoding='utf-8')
def readcsv(name):return list(csv.DictReader((SRC/name).open(encoding='utf-8')))
groups=readcsv('Koehn.et.al.2016_groupinfo.csv')
params=readcsv('Koehn.et.al.2016_parameters.csv')
dietrows=list(csv.reader((SRC/'Koehn.et.al.2016_dietmatrix.csv').open()))
assert len(groups)==len(params)==93
name_aliases=[{'seq':i+1,'groupinfo_name':g['GroupName'],'parameters_name':p['Groups'],'resolution':'Retain groupinfo identity/order; PB QB B EE cross-check against manuscript Table 1. parameters.csv swaps Shelf/Slope labels at rows 39/55.' if i in [37,53] else 'Author abbreviation/spelling alias; aligned by verified parameter signature and author R row indexing.'} for i,(g,p) in enumerate(zip(groups,params)) if g['GroupName']!=p['Groups']]
dump(OUT/'source_name_aliases.json',name_aliases)
for g,p in zip(groups,params):
 for a,b in [('Biomass','density'),('EE','EE'),('PB','PB'),('QB','QB')]:assert Decimal(g[a])==Decimal(p[b])
DC=np.array([row[1:94] for row in dietrows[1:]],float)
assert DC.shape==(94,93)
assert all(float(g['Import'])==DC[-1,j] for j,g in enumerate(groups))
sheet=load_workbook(PAPER/'mmc1.xlsx',data_only=True).active
supp=np.array([[sheet.cell(i,j).value for j in range(2,95)] for i in range(3,97)],float)
assert np.array_equal(DC,supp),'Publisher and author diets differ'
dump(OUT/'source_comparison.json',{'parameter_duplicate_cells_equal':372,'diet_publisher_vs_zip_cells_equal':8742,'diet_total_min':float(DC[:,1:92].sum(0).min()),'diet_total_max':float(DC[:,1:92].sum(0).max()),'diet_import_max_discrepancy':0})

# Preserve source row identities; manuscript displays seabirds before mammals.
starts=[32,6,9,12,35,38,41,44,47,50,15,18,21,28,60,63,66,54,70,75,82,99,106,111,116,121,126,132,135,138,141,144,150,155,158,163,168,171,178,183,188,195,200,205,210,215,220,225,228,233,238,247,252,257,262,267,270,275,280,294,301,308,315,326,333,426,433,444,451,348,464,469,474,479,484,493,498,503,508,519,528,533,546,551,558,359,370,377,386,397,408,415,None]
doc=Document(PAPER/'mmc2.docx')
headings=sorted(x for x in starts if x is not None)+[565]
taxonomy=[]
for i,(g,start) in enumerate(zip(groups,starts),1):
 if start is None:desc='Detritus; nonLiving pool in author groupinfo.csv row 94. No taxonomic membership applies.'; paras=[];heading='Detritus'
 else:
  end=next(x for x in headings if x>start)
  paras=[{'paragraph_zero_based':j,'text':doc.paragraphs[j].text} for j in range(start+1,end) if doc.paragraphs[j].text.strip()]
  heading=doc.paragraphs[start].text.strip()
  desc=f'Appendix B: {heading}. '+paras[0]['text']
 taxonomy.append({'seq':i,'group_name':g['GroupName'],'taxon_descr':desc,'appendix_heading':heading,'heading_paragraph_zero_based':start,'source_paragraphs':paras})
dump(OUT/'taxonomy_source_evidence.json',taxonomy)
wb=Workbook();ws=wb.active;ws.title='Taxonomy';ws.append(['seq','group_name','taxon_descr'])
for t in taxonomy:ws.append([t['seq'],t['group_name'],t['taxon_descr']])
wb.save(OUT/'Taxonomy.xlsx')
with (OUT/'taxonomy.csv').open('w',encoding='utf-8',newline='') as f:
 w=csv.writer(f);w.writerow(['seq','group_name','taxon_descr']);w.writerows((t['seq'],t['group_name'],t['taxon_descr']) for t in taxonomy)

model={'metadata':{'LME':'3 California Current','model_number':'CAL-2016','model_name':'California Current','model_year':'2000-2014'},'groups':[], 'consumers':list(range(2,93)), 'fleets':['Total catch including discards'], 'landings':{},'discards':{},'detritus_groups':['Detritus'],'detritus_fate':{},'diet':{},'diet_rows':93,'landings_rows':93,'discards_rows':93}
provenance=[]
for i,(g,p) in enumerate(zip(groups,params),1):
 def known(k):return None if Decimal(g[k])<0 else g[k]
 model['groups'].append({'n':i,'name':g['GroupName'],'biomass':known('Biomass'),'pb':known('PB'),'qb':known('QB'),'ee':known('EE'),'ba':g['BA']})
 model['landings'][str(i)]={'Total catch including discards':p['Yield']}
 # Author runecopath lines 98-104 route all unused living production to the sole pool.
 # Egestion is not represented by author code; its runtime extension is documented separately.
 if i<93:model['detritus_fate'][str(i)]={'Detritus':'1'}
 if 2<=i<=92:
  model['diet'][str(i)]={str(j):dietrows[j][i] for j in range(1,94)}
  model['diet'][str(i)]['import']=dietrows[94][i]
 provenance.append({'seq':i,'group_name':g['GroupName'],'source_csv_row':i+1,'diet_predator_header':dietrows[0][i],'diet_prey_label':dietrows[i][0],'publisher_diet_column':i+1,'publisher_diet_row':i+2,'B_source_flag':'solve_unknown' if float(g['Biomass'])<0 else 'input','EE_source_flag':'solve_unknown' if float(g['EE'])<0 else 'input','GCE_source_zero':'placeholder; author recomputes PB/QB, not extracted as GE=0','GS_source_flag':'not reported','BA_source_flag':'explicit 0','catch_basis':'Yield total catch including discards; no split supplied'})
dump(OUT/'model.json',model);dump(OUT/'cell_provenance.json',provenance)

# Reconstruct the author linear system without executing archived R code.
B=np.array([float(g['Biomass']) for g in groups]);EE=np.array([float(g['EE']) for g in groups]);PB=np.array([float(g['PB']) for g in groups]);QB=np.array([float(g['QB']) for g in groups]);Y=np.array([float(p['Yield']) for p in params]);BA=np.zeros(93)
unknown=np.where(B[:92]<0)[0];known=np.where(B>0)[0]
A=np.zeros((92,92));rhs=Y[:92]+DC[:92,known]@(B[known]*QB[known])
for i in range(92):
 for j in range(92):
  if i==j:A[i,j]=EE[i]*PB[i]-QB[i]*DC[i,i] if j in unknown else B[i]*PB[i]
  elif j in unknown:A[i,j]=-DC[i,j]*QB[j]
x=np.linalg.solve(A,rhs);Bout=B.copy();EEout=EE.copy()
for i in range(92):
 if i in unknown:Bout[i]=x[i]
 else:EEout[i]=x[i]
# Independent smaller system directly eliminating known-B equations.
A2=np.diag(EE[unknown]*PB[unknown])-DC[np.ix_(unknown,unknown)]*QB[unknown][None,:]
x2=np.linalg.solve(A2,rhs[unknown]);assert np.allclose(Bout[unknown],x2,rtol=1e-12)
pred=DC[:93,:]@(Bout*QB)
res=Bout[:92]*PB[:92]*EEout[:92]-Y[:92]-pred[:92]
det_in=float(np.sum(Bout[:92]*PB[:92]*(1-EEout[:92])))
EEout[92]=pred[92]/det_in
result={'method':'Independent NumPy transcription of author runecopath linear system; independently checked by reduced unknown-B system and production identities','unknown_B_count':len(unknown),'unknown_EE_living_count':92-len(unknown),'max_abs_linear_residual':float(max(abs(A@x-rhs))),'max_abs_production_residual':float(max(abs(res))),'max_abs_independent_B_discrepancy':float(max(abs(x2-Bout[unknown]))),'min_B':float(min(Bout)),'max_living_EE':float(max(EEout[:92])),'living_EE_gt_1':[groups[i]['GroupName'] for i in range(92) if EEout[i]>1],'author_detritus_inflow_unused_production_only':det_in,'author_detritus_consumption':float(pred[92]),'author_detritus_EE':float(EEout[92]),'groups':[{'seq':i+1,'group_name':g['GroupName'],'biomass':float(Bout[i]),'ee':float(EEout[i]),'biomass_origin':'derived_author_equations' if i in unknown else 'source_input','ee_origin':'source_input' if i in unknown else 'derived_author_equations'} for i,g in enumerate(groups)]}
dump(OUT/'author_equation_solution.json',result)

# Coordinate-based paper table extraction, with explicit page-specific column boundaries.
pdf=pymupdf.open(next(PAPER.glob('*.pdf')));paperrows=[]
for pageidx, ymin,ymax,bounds in [(2,107,730,[130,180,227,276,327,380,435,484,529]),(3,81,270,[119,164,213,262,313,367,423,472,519])]:
 words=[w for w in pdf[pageidx].get_text('words') if ymin<w[1]<ymax]
 ys=[]
 for w in sorted(words,key=lambda w:w[1]):
  hit=next((r for r in ys if abs(r[0]-w[1])<1),None)
  if hit is None:ys.append([w[1],[w]])
  else:hit[1].append(w)
 for y,ws2 in ys:
  cells=['']*10
  for w in sorted(ws2,key=lambda w:w[0]):
   col=sum(w[0]>=b for b in bounds);cells[col]+=(' ' if cells[col] else '')+w[4]
  if all(cells[1:]):paperrows.append({'pdf_page':pageidx+1,'printed_page':pageidx+87,'y':y,'cells':cells})
dump(OUT/'paper_table1_coordinate_rows.json',paperrows)
manifest=[]
for f in sorted(PAPER.rglob('*')):
 if f.is_file():manifest.append({'path':f.relative_to(ROOT).as_posix(),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
dump(REVIEW/'source_manifest.json',manifest)
print(json.dumps({k:v for k,v in result.items() if k!='groups'},indent=2));print('paper_rows',len(paperrows),'model',MODEL)
