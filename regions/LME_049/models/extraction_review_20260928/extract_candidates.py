"""Source-coordinate extraction, preserving strings, censoring, and evidence."""
from pathlib import Path
import json,csv,re,hashlib,sys,subprocess,shutil
import fitz
from docx import Document
import openpyxl
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'Project.xlsx').exists())
REG=ROOT/'regions/LME_049'; REVIEW=Path(__file__).parent
SKILL=Path('C:/Users/idoca/.agents/skills/ecopath-extraction/scripts')
def spans(p):
    return [s for b in p.get_text('dict')['blocks'] if 'lines' in b for l in b['lines'] for s in l['spans']]
def clean(s):return ' '.join(s.replace('\u200b','').split())
def num(s):return s if re.fullmatch(r'\d+(?:\.\d+)?',s) else None
def dump(p,obj):p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def extract(year,mid,number,name,period):
    src=REG/'papers'/f'KUR-{year}'; pdf=next(src.glob('*.pdf')); doc=fitz.open(pdf)
    dest=REG/'models'/mid;et=dest/'extracted_tables'; et.mkdir(parents=True,exist_ok=True)
    page=doc[5 if year==2019 else 4];ss=spans(page)
    if year==2019:
        fields=['tl','hab_area','biomass','biomass_model_area','pb','qb','ee','catch']
        centers=[219,255,300,345,390,444,496,543];n=41;lo=155;hi=560
    else:
        fields=['tl','biomass','pb','qb','ee','catch','omnivory_index'];centers=[161,208,254,330,413,477,518];n=25;lo=95;hi=323
    anchors=sorted([(int(clean(s['text'])),s['bbox'][1]) for s in ss if s['bbox'][0]<55 and lo<s['bbox'][1]<hi and clean(s['text']).isdigit()])
    assert [i for i,y in anchors]==list(range(1,n+1)),anchors
    groups=[];evidence=[];bounds=[];landings={}
    for k,(seq,y) in enumerate(anchors):
        yend=anchors[k+1][1]-.5 if k+1<n else y+8
        row=[s for s in ss if y-.2<=s['bbox'][1]<yend]
        nameparts=[clean(s['text']) for s in row if 55<s['bbox'][0]<(205 if year==2019 else 151)]
        gn=' '.join(nameparts);original=gn;gn=gn.replace(',',';')
        g={'n':seq,'name':gn};vals={}
        for s in row:
            x0,y0,x1,y1=s['bbox'];t=clean(s['text'])
            if x0<(205 if year==2019 else 151):continue
            fld=fields[min(range(len(centers)),key=lambda j:abs((x0+x1)/2-centers[j]))]
            if not t:continue
            assert fld not in vals,(seq,fld,t,vals)
            vals[fld]=t
            evidence.append({'group_seq':seq,'group_name':original,'field':fld,'source_text':t,'source':str(pdf.relative_to(ROOT)),'pdf_page':6 if year==2019 else 5,'printed_page':299 if year==2019 else 5,'table':'2','bbox':list(s['bbox']),'font':s['font'],'bold':bool(s['flags']&16)})
            if t.startswith('<'):bounds.append({'group_seq':seq,'field':fld,'source_text':t,'numeric_value':None})
        for f in fields:
            if f not in ['catch','omnivory_index','biomass_model_area']:g[f]=num(vals.get(f,''))
        if year==2019:g['z']=g['pb'] # Section 2.1 explicitly equates P/B and Z.
        if num(vals.get('catch','')) is not None:landings[str(seq)]={'Total fishery':vals['catch']}
        groups.append(g)
    model={'metadata':{'LME':'49 Kuroshio Current','model_number':number,'model_name':name,'model_year':period},'groups':groups,'consumers':list(range(1,36 if year==2019 else 24)),'fleets':['Total fishery'],'landings':landings,'discards':{},'detritus_groups':[g['name'] for g in groups if g['name'].startswith('Detritus')],'detritus_fate':{},'diet':{},'diet_rows':n}
    taxonomy={};diet_evidence=[]
    if year==2019:
        words=doc[7].get_text('words');header=[w for w in words if abs(w[0]-100.279)<.1 and w[4].isdigit()]
        colanchors={int(w[4]):w[1] for w in header};assert len(colanchors)==35
        rowanchors=sorted([w for w in words if w[4].isdigit() and 709<w[1]<718 and 109<w[0]<537],key=lambda w:w[0]);assert len(rowanchors)==41,len(rowanchors)
        for rw in rowanchors:
            prey=int(rw[4]);x=rw[0]
            for w in words:
                if abs(w[0]-x)>1 or w[1]>610 or w[1]<80:continue
                predator=min(colanchors,key=lambda k:abs(w[1]-colanchors[k]))
                assert abs(w[1]-colanchors[predator])<2,(w,predator)
                assert num(w[4]) is not None,w
                model['diet'].setdefault(str(predator),{})[str(prey)]=w[4]
                diet_evidence.append({'prey_seq':prey,'predator_seq':predator,'source_text':w[4],'pdf_page':8,'printed_page':301,'table':'3','bbox':list(w[:4])})
        assert len(diet_evidence)==41*35,len(diet_evidence)
        species={6:'Katsuwonus pelamis',8:'Seriola quinqueradiata',9:'Sardinops melanostictus',10:'Engraulis japonicus',11:'Cololabis saira',12:'Scomber japonicus',13:'Scomber australasicus',14:'Etrumeus teres',15:'Trachurus japonicus',17:'Gadus chalcogrammus (also known as Alaska pollock)',18:'Gadus macrocephalus'}
        for g in groups:
            seq=g['n'];taxonomy[seq]=(species.get(seq,'not documented in locally available main article; group definition: '+g['name'])+'; spatial distribution '+g['name'][g['name'].find('('):]+'; source: main article Table 1 and Section 2.3 (printed pp.297–298). Detailed species inventory is in publisher supplement; exact bytes unavailable due HTTP401 bot check.')
    else:
        supp=next(src.glob('*.docx'));dt=Document(supp).tables[1]
        for r in dt.rows[1:]:
            cells=[clean(c.text) for c in r.cells];prey=int(cells[1]);assert cells[0].replace(',',';')==groups[prey-1]['name'],(prey,cells[0],groups[prey-1]['name'])
            for predator,t in enumerate(cells[2:],1):
                diet_evidence.append({'prey_seq':prey,'predator_seq':predator,'source_text':t,'table':'S2','docx_table_index':2,'docx_row':prey+1,'docx_column':predator+2})
                if t=='+':bounds.append({'prey_seq':prey,'predator_seq':predator,'field':'diet','source_text':'+','bound':'percentage <0.01 (verbatim caption; scale ambiguous)','numeric_value':None})
                elif t:
                    assert num(t) is not None,t
                    model['diet'].setdefault(str(predator),{})[str(prey)]=t
        tp=doc[2];ts=spans(tp)
        left=[s for s in ts if 311<s['bbox'][0]<315 and 98<s['bbox'][1]<446]
        names=[g['name'].replace(';',',') for g in groups]
        starts=[]
        for gn in names:
            match=[s for s in left if gn.startswith(clean(s['text']))]
            # Starts match complete first line; exact y-order is verified against 25 group sequence.
            match=[s for s in match if clean(s['text']) not in ['and driftfish']]
            assert len(match)==1,(gn,match)
            starts.append(match[0]['bbox'][1])
        assert starts==sorted(starts)
        for i,g in enumerate(groups):
            tx=' '.join(clean(s['text']) for s in ts if s['bbox'][0]>390 and starts[i]-.1<=s['bbox'][1]<(starts[i+1]-.1 if i+1<n else 453))
            taxonomy[g['n']]=tx+'; scope: main species composition (not claimed exhaustive), main article Table 1, PDF/printed p.3; source spelling preserved.'
        dump(et/'supplement_tables.json',[[[c.text for c in r.cells] for r in t.rows] for t in Document(supp).tables])
    for c in model['consumers']:model['diet'].setdefault(str(c),{})
    dump(et/'model.json',model);dump(et/'source_cells.json',evidence);dump(et/'diet_source_cells.json',diet_evidence);dump(et/'censored_values.json',bounds)
    subprocess.run([sys.executable,'-X','utf8',str(SKILL/'write_outputs.py'),str(et/'model.json'),'--outdir',str(dest),'--dir-name','extracted_tables'],check=True)
    w=openpyxl.Workbook();ws=w.active;ws.title='Taxonomy';ws.append(['seq','group_name','taxon_descr'])
    for g in groups:ws.append([g['n'],g['name'],taxonomy[g['n']]])
    w.save(et/'Taxonomy.xlsx')
    with (et/'taxonomy.csv').open('w',newline='',encoding='utf-8') as f:
        cw=csv.writer(f);cw.writerow(['seq','group_name','taxon_descr']);cw.writerows([[g['n'],g['name'],taxonomy[g['n']]] for g in groups])
    manifest=[{'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in src.iterdir() if p.suffix in ['.pdf','.docx']]
    dump(dest/'source_manifest.json',manifest)
    print(mid,'groups',n,'diet_cells',len(diet_evidence),'censored',len(bounds),'diet_sums',[(i,sum(float(v) for v in d.values())) for i,d in model['diet'].items()])
    return dest
if __name__=='__main__':
    extract(2019,'49_20192013_Western_North_Pacific_Watari_(2013)',20192013,'Western North Pacific Watari',2013)
    extract(2025,'49_20252023_Kuroshio_Oyashio_Extension_Chen_(2023)',20252023,'Kuroshio Oyashio Extension Chen',2023)
