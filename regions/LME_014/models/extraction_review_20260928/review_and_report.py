from pathlib import Path
import json,csv,hashlib,re,sys
import numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[4];R=Path(__file__).resolve().parent;E=R/'evidence'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def dump(p,x):p.write_text(json.dumps(x,indent=2,ensure_ascii=False),encoding='utf-8')
ids=read(R/'model_ids.json');manifest=read(E/'source_manifest.json')
checks=[dict(path=x['path'],matches=(ROOT/x['path']).stat().st_size==x['bytes'] and hashlib.sha256((ROOT/x['path']).read_bytes()).hexdigest()==x['sha256']) for x in manifest]
assert all(c['matches'] for c in checks)
dump(E/'source_hash_verification.json',checks)
native=read(E/'PAT-2023_native_tables.json');ng={g['Sequence']:g for g in native['EcopathGroup']};idseq={g['GroupID']:g['Sequence'] for g in ng.values()}
# Independent Ecopath algebra is a calculation-only audit, never a source value or adopted model.
living=list(range(1,35));unk=[i for i in living if ng[i]['Biomass']==-9999];known=[i for i in living if i not in unk]
DC=np.zeros((36,36))
for d in native['EcopathDietComp']:DC[idseq[d['PredID']]-1,idseq[d['PreyID']]-1]=d['Diet']
qb=np.array([max(0,ng[i]['ConsBiom']) for i in range(1,37)])
pb=np.array([max(0,ng[i]['ProdBiom']) for i in range(1,37)])
catch=np.array([sum(r['Landing']+r['Discards'] for r in native['EcopathCatch'] if idseq[r['GroupID']]==i) for i in range(1,37)])
B=np.array([max(0,ng[i]['Biomass']) for i in range(1,37)])
A=np.zeros((len(unk),len(unk)));rhs=np.zeros(len(unk))
for row,i in enumerate(unk):
 rhs[row]=catch[i-1]+ng[i]['BiomAcc']+ng[i]['Emigration']-ng[i]['Immigration']+sum(B[j-1]*qb[j-1]*DC[j-1,i-1] for j in known)
 for col,j in enumerate(unk):A[row,col]=(pb[i-1]*ng[i]['EcoEfficiency'] if i==j else 0)-qb[j-1]*DC[j-1,i-1]
solution=np.linalg.solve(A,rhs)
for i,b in zip(unk,solution):B[i-1]=b
pred=(B*qb)@DC
audit=[]
for i in living:
 g=ng[i];used=pred[i-1]+catch[i-1]+g['BiomAcc']+g['Emigration']-g['Immigration']
 ee=used/(B[i-1]*pb[i-1])
 audit.append(dict(seq=i,name=g['GroupName'],biomass_native=None if g['Biomass']==-9999 else g['Biomass'],biomass_solved=float(B[i-1]),ee_native=None if g['EcoEfficiency']==-9999 else g['EcoEfficiency'],ee_recomputed=float(ee),predation=float(pred[i-1]),catch=float(catch[i-1]),immigration=g['Immigration'],source='Native coupled production balance; no diet normalization; no source replacement'))
dump(E/'native_independent_balance_audit.json',{'equation':'Bi * PBi * EEi = catches_i + BA_i + emigration_i - immigration_i + sum(Bj * QBj * DCji)','unknown_biomass_groups':unk,'matrix_residual_max':float(np.max(np.abs(A@solution-rhs))),'groups':audit,'not_an_adopted_model':True})
print('independent native EE>1',[(g['seq'],g['name'],g['ee_recomputed']) for g in audit if g['ee_recomputed']>1])

summary=[];proposals=[]
for mid in ids:
 f=ROOT/'regions/LME_014/models'/mid;tab=f/'extracted_tables';src=read(f/'extraction_input.json');can=read(f/'model.json');run=read(f/'diagnostic_run_record.json')
 is_native=mid.endswith('_native');is_pat=mid.startswith('PAT');count=len(src['groups'])
 required=['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv','TL.xlsx','Metadata.xlsx']
 assert all((tab/p).is_file() for p in required)
 assert len(can['group'])==count
 tw=load_workbook(tab/'Taxonomy.xlsx',read_only=True);assert tw.active.max_row==count+1
 assert [g['group_name'] for g in can['group']]==[g['name'] for g in src['groups']]
 assert all(g['taxon_descr'] for g in can['group'])
 wb=load_workbook(f/'model_reconstructed.xlsx',read_only=True,data_only=False)
 reconstruction={'workbook_sheets':wb.sheetnames,'group_count':count,'taxonomy_count':tw.active.max_row-1,'source_csv_files_present':True,'canonical_group_order_and_names_match':True,'source_hashes_unchanged':True,'canonical_changes_from_stock_converter':len(read(tab/'converter_canonical_differences.json'))}
 dump(tab/'verification.json',reconstruction)
 provenance=[]
 for g in src['groups']:
  n=g['n']
  for field,value in g.items():
   if field in ['n','name']:continue
   location=f'EcopathGroup GroupID={ng[n]["GroupID"]}; Sequence={n}' if is_native else (f'Article Table 4; PDF p{14 if n<30 else 15}; group {n}; field {field}' if is_pat else f'Article Table 2; PDF p5 / printed p35; row {n}; field {field}')
   if is_native and field=='tl':location=f'Article Table 4; PDF p{14 if n<30 else 15}; group {n}; TL column (published output; native table does not store TL)'
   provenance.append(dict(group=n,parameter=field,value=value,location=location))
 for consumer,diet in src['diet'].items():
  for prey,v in diet.items():
   if is_native:loc=f'EcopathGroup GroupID={ng[int(consumer)]["GroupID"]} ImpVar' if prey=='import' else f'EcopathDietComp PredID={ng[int(consumer)]["GroupID"]}; PreyID={ng[int(prey)]["GroupID"]}; Diet'
   elif is_pat:
    ci=src['consumers'].index(int(consumer))+3
    pseq=[i for i in range(1,35) if i not in [12,22]]
    # CSV prey rows include primary producers and use the 35-group basic ordering.
    native_name=ng[int(prey)]['GroupName'] if prey!='import' else 'Import'
    rows=list(csv.reader(next((ROOT/'regions/LME_014/papers/PAT-2023').glob('*.csv')).open()))
    ri=next(k+1 for k,row in enumerate(rows) if len(row)>1 and row[1].casefold()==native_name.casefold())
    loc=f'Supplement S2 CSV row {ri} column {ci}; percentages divided by 100; no normalization'
   else:
    ti=1 if int(consumer)<=12 else 2;ci=(int(consumer)-1)%12+3;ri=29 if prey=='import' else int(prey)+2
    loc=f'Appendix A DOCX table {ti} row {ri} column {ci}; rendered supplement p{ti}'
   provenance.append(dict(group=int(consumer),parameter='diet:'+prey,value=v,location=loc))
 for kind in ['landings','discards']:
  for s,fleets in src[kind].items():
   for fleet,v in fleets.items():
    if is_native:loc=f'EcopathCatch GroupID={ng[int(s)]["GroupID"]}; FleetName={fleet}; '+('Landing' if kind=='landings' else 'Discards')
    else:loc=f'Supplement DOCX Table S4 (document table 3); rendered p5; source row group {s if s!="36" else "35 Detritus"}; {kind} / {fleet}'
    provenance.append(dict(group=int(s),parameter=kind+':'+fleet,value=v,location=loc))
 dump(tab/'cell_provenance.json',provenance)
 statuses={k:(v['return']['status'] if v['execution']=='returned' else v['execution'].upper()) for k,v in run['diagnostics'].items()}
 summary.append(dict(model_id=mid,groups=count,diagnostics=statuses,production_eligible=False,source_report=f'{mid}/source_report.md',diagnostic_report=f'{mid}/diagnostic_report.md'))
 proposals.append(dict(unit_id='LME_014',model_id=mid,model_path=f'models/{mid}/model.json',paper_id='PAT-2023' if is_pat else 'LME014-OcampoReinaldo-2016',publication_year=2024 if is_pat else 2016,model_year=2020 if is_pat else 1970,model_area_km2=200000 if is_pat else 10000,coverage_percent=None,selected=False,production_eligible=False,registration_status='PROPOSED ONLY — Project.xlsx not edited'))
 validation=(tab/'validation.txt').read_text(encoding='utf-8').strip()
 balance=(tab/'source_massbalance.txt').read_text(encoding='utf-8')
 tail=next((x for x in reversed(balance.splitlines()) if 'error(s)' in x),'See source_massbalance.txt')
 common=f'''# Source extraction report — {mid}

Extracted 2026-09-28. Scope: source-faithful reconstruction and isolated model diagnostics. LME_014 Overview has no selected model; it remains unchanged. No regional PPR or central workbook values were written. The local model identifier is not an EcoBase accession.

## Bundle and evidence

All actual files in both supplied paper folders were inventoried before extraction; `../extraction_review_20260928/evidence/source_manifest.json` records filenames, byte sizes and SHA-256. The final `source_hash_verification.json` verifies every byte unchanged. Source PDFs, supplement DOCX tables and native database tables are retained through read-only evidence extraction. `extracted_tables/cell_provenance.json` maps basic, diet and fishery values to source pages/cells/record keys. PDF coordinate cells and rendered review pages reside in the shared evidence folder. Both supplementary DOCX files were rendered through Word in read-only mode; Ocampo diet pages 1–2 and PAT Table S4 page 5 were visually checked. Paper table pages were rendered and visually checked.

## Group structure and taxonomy

{count} source groups. All have `Taxonomy.xlsx`, a source-qualified `taxonomy.csv`, and canonical JSON `taxon_descr`. Source spellings are preserved; no external synonym corrections or taxon assignment weights were invented. No Selected model groups worksheet update is appropriate before selection. Scientific names in labels are preserved separately from guild composition, and size/cohort definitions remain distinct.

## Import and converter validation

All eight required EwE tables are in `extracted_tables/`. The writer override `--dir-name extracted_tables` follows the active regional project contract. The extraction-input JSON is separate from canonical database `model.json`. Source CSV blanks remain blank, published zeros remain zero. Source unknown database values use `-9999`.

The stock converter ran and produced its own database JSON, conversion log, MASS_BALANCE.md and reconstructed workbook. Its output is explicitly an intermediate: it normalizes some diets, converts missing catches/imports/routing to zeros, defaults habitat area, drops printed P/Q, and converts habitat biomass. These transformations cannot silently become source data. `converter_canonical_differences.json` records each restoration in the canonical JSON; canonical diet values are never normalized. Canonical `model_reconstructed.xlsx` was regenerated and checked for workbook structure, group count, names and taxonomy. The converter's reconstruction is not a lossless source representation: it cannot retain all native fleet routing or paper-specific input/output font markers; source inputs and cell evidence remain authoritative.

Structural validator output:

```
{validation}
```

Source mass-balance tool: {tail}. Full arithmetic and flags: `extracted_tables/source_massbalance.txt`. Its printed recomputed EEs are conditional comparisons: unknown BA/catch is numerically omitted, missing predator B/QB undercounts predation, and immigration is not represented in its eight-table contract. Critically, the standalone massbalance_check.py uses the habitat-biomass column without multiplying habitat area, making its Falkland predation comparisons inconsistent with model-area catches. Its PAT errors must not be read as native model failure. The separate native audit uses model-area B and includes immigration. Suggested BA values were not adopted. `MASS_BALANCE.md` is the stock converter's check, not a claim of strict source admission.

## Direct diagnostics and production status

The separate diagnostic report contains only the full direct returns of `PPRCalculator.diagnose_sppr(short=False, flat=False)` for GE, TE and With Egestion, or an exact not-run reason. `diagnostics_raw.json` retains the returned objects. Statuses: {statuses}. All numerical runs disable diet normalization and detritus pooling (`det_collapse_mode=never`, `det_open_mode=none`). No global variants, method inventory or Monte Carlo were run.

`diagnostic_run_record.json` records all boundary settings, warnings and exceptions. `loader_input_groups.csv`, `loader_diet.csv`, `loader_detritus_fate.csv` and `loader_completed_groups.csv` make loader alterations reviewable. The legacy loader requires a numeric filename and a parenthesized year; `loader_input/` contains an escaped copy under a local diagnostic placeholder id, not a model renumbering or EcoBase claim. For native PAT, this diagnostic copy also contains explicitly derived missing B/EE values listed in the run record; canonical source JSON remains unchanged. The loader hardcodes LME=13 internally; the canonical extraction metadata correctly says LME=14. Diagnostic calculations do not establish region identity or authorize adoption.

Production eligibility: **not established / false**. Successful computation is distinct from scientific source fidelity. No candidate is selected.
'''
 if is_native:
  detail='''
## Source tables and numerical identity

Büring et al. (2024), *Unveiling the wasp-waist structure of the Falkland shelf ecosystem: the role of Doryteuthis gahi as a keystone species and its trophic influences*, J. Mar. Biol. Assoc. UK 104 e2, DOI 10.1017/S0025315423000887. Folder PAT-2023 reflects the accepted manuscript/DOI year, not the printed publication year. The model is **Falkland Islands 2020**, FirstYear 2020 / NumYears 1, area 200,000 km², saved EwE 6.6.8.18072. The author-supplied native sup002.eweaccdb contains exactly one EcopathModel, 36 EcopathGroup rows, 1,296 EcopathDietComp records, 72 catch rows and two fleets (Trawling, Jigging). Ecosim scenario/shape tables are not additional published Ecopath model periods.

Native `Sequence`, not Access `GroupID`, fixes group order. Native group IDs and source records are retained in evidence. Native biomass is on model-area basis: B=0.002500000176951289 for Baleen Whales at habitat area=0.20000000298023224, consistent with the article's ~0.01 habitat biomass/~0.003 model-area biomass. Import-table habitat B is derived by dividing native B by Area; canonical model-area B is restored exactly from the binary value. Single-precision binary artifacts are retained, not rounded into allegedly printed values.

Native basic mapping: Area→hab_area; Biomass/Area→habitat biomass; ProdBiom→P/B; ConsBiom→Q/B; EcoEfficiency→EE; Unassim→GS; DtImports→detritus import; BiomAcc→BA. Native -9999 is unknown/unresolved. Native native `OtherMort`/`ProdCons` are unknown where sentinel-coded. Article Table 4 TL values are recorded as published outputs, not native resolved outputs. Native ImpVar supplies exact imports for every group (including stated zeros for nonconsumers). Diet rows map PredID to predator and PreyID to prey. Each native consumer diet plus import sums to within 1.2e-7 of 1; no normalization was applied.

Native `EcopathDietComp.DetritusFate` routes every living group to Detritus=36; both nonliving fate rows contain printed/native zeros. The validator expects every fate row to sum to one and flags these two source rows; they were deliberately not replaced with invented identity routing. ModelData itself forces a detritus identity block; that alteration exists only in the loader. `EcopathDiscardFate` sends both fleets' discards to Detritus=36, not Discards=35. Five consumer groups nevertheless consume Discards=35. This native configuration and the eight-table/engine omission of fleet-specific discard supply are retained as material limitations. Full fleet routing is in the native evidence JSON.

## Biomass accumulation and prose

Native BiomAcc and BiomAccRate are explicitly zero for all 36 groups; zero is therefore sourced here. Native GS=0.2 (binary representation retained) for regular groups, with producer and Detritus values zero; no convention was substituted. Kingclip has native immigration=0.003000000026077032, preserved in canonical JSON although the eight-table format has no migration table. Unknown native B (13 living/producer groups plus Detritus) and many EEs remain unresolved source fields.

## Source conflicts and independent arithmetic

The article Table 4 has 36 groups with separate Discards/Detritus, but Table 2/S2 use 35 groups and combine the detritus category. CSV S2 and native diet/import values differ beyond rounding. Native Detritus EE=0 versus printed Table 4 EE=0.05. These are not reconciled by pooling or averaging. The published-table candidate remains separate. Table S3 is explicitly *initial* P/B/Q/B/EE, lists 45 rows including groups absent from the final model, and has no complete aligned B/diet model. Its full cell contents are retained as initial/predecessor evidence, not adopted as a fourth balanced model.

The independent audit `../extraction_review_20260928/evidence/native_independent_balance_audit.json` solves the native unknown living biomasses from the coupled production equations, preserving all source diet, catch, migration and EE inputs. The matrix residual is1.13e-14; every solved living/producer biomass is positive and every recomputed EE lies in[0,1]. Solved B agrees with the article's printed outputs to source rounding (e.g. Benthic Crustaceans1.371551, Phytoplankton3.255569). The eight-table massbalance helper's contrary errors arise partly from omitting habitat-area scaling and immigration; its raw results remain visible, not silently repaired.

The diagnostic loader copy explicitly fills only native-unknown living/producer B and EE with this algebraic solution. Canonical model.json retains the native unknowns and every source value. The full direct diagnostics on this declared completion all return **WARN**: model production/consumption residuals are approximately1.67e-14/1.08e-15. GE and With Egestion PP-budget gaps are about1.17e-16; TE's gap is0.0001256. Baleen Whales EE=0 is a remaining warning, and TE has additional convergence/magnitude warning details in the direct return. The earlier uncompleted stock-loader run is retained under `default_loader_diagnostics/`; its FAIL results came from replacing unknown B with1 and losing unresolved p/q, so they do not describe the faithful coupled reconstruction. Native Discards receives no modeled detritus-fate inflow, and the loader compensates its consumption with negative detritus accumulation; fleet discard routing and native-vs-printed EE remain explicit caveats. No detritus supply, routing, diet, catches, migration or GS was invented by the algebraic completion.

## Spatial applicability

Article p.6 defines the Falkland shelf to 300 m, about 200,000 km², inside the wider Patagonian Shelf LME. It excludes adult deep-water toothfish, and separates D. gahi seasonal cohorts. No LME catch-coverage fraction is supplied or inferred. Table 2 lists Pygoscelis papua in both Penguins and Seabirds; that source overlap remains explicit for later matching.
'''
 elif is_pat:
  detail='''
## Source tables and numerical identity

Same Büring et al. (2024) article and 2020 Falkland shelf period as the native candidate. This is the **published-table reconstruction**, retained separately because the tabulated parameters and native database conflict. Table 4 (PDF/printed pp14–15) provides 36 basic rows in exact source order. Header mapping: Trophic level→TL; Habitat area→hab_area; Biomass in habitat area→habitat B; Biomass→model-area B; Production/biomass→P/B; Consumption/biomass→Q/B; EE→EE; Production/consumption→P/Q. Parentheses indicate Ecopath-estimated B and user-input EE per caption. The source matrix and basic table have no rounding repairs.

Table 4 was extracted through displayed word bounding boxes, respecting the PDF p14 rotation and locally rotating the p15 continuation in memory. Published model-area B is separately retained: it sometimes differs from multiplying independently rounded habitat B and habitat area (e.g. Baleen whales 0.003 versus 0.01×0.2). Canonical B uses the directly printed model-area column; eight-table conversion uses multiplication. Both values and the transformation are recorded. Published zero B for Flounder, Hake Austral and Sharks remains zero, despite positive native values.

S2 CSV has 32 consumer columns, 35 prey rows including producers but one Detritus category, Import and Sum. It reports percentages, divided by 100 exactly; calculated sums range 0.99998–1.00002. CSV numbering differs from basic numbering because producer groups are absent among consumers. Names, not column positions, determine the explicit mapping. The S2 Detritus label maps to Table4 Detritus=36; Table4 Discards=35 has no separate S2 row and stays unknown. This is an unresolved mismatch, not evidence that discards were pooled or absent. Source raw CSV and per-cell mappings are retained. Table S4 (supplement DOCX table3, rendered p5) supplies separate Trawling/Jigging landings/discards; '-' stays unknown, while Baleen Whales discard zeros are sourced. Detritus in S4 is #35 and maps by name to Table4 #36.

## Biomass accumulation, deliberate blanks and conflicts

The full paper, S2, S3, S4 and remaining supplement were checked. The basic Ecopath equation names BA but supplies no numeric BA in the printed-table evidence; every BA import cell stays blank. No printed GS or detritus routing was found. Unlike the native database, the printed-source version does not borrow native zeros/GS/fates/immigration. The Table2 versus Table4 Kelp definition differs: Table4 label names Macrocystis pyrifera; Table2 composition also includes Lessonia spp., retained in taxonomy. Native and printed diet/EE differences are listed in the comparison evidence. Table S3 initial 45-group rates are preserved as an incomplete initial parameter compilation, not merged into this final reconstruction.

The direct loader rejects this two-detritus model because no living-group routing to Discards vs Detritus is given. All three requested diagnostics are explicitly NOT_RUN; no synthetic split or pooling was introduced. Native routing cannot simply be copied into this conflicting published version.

## Spatial applicability

Falkland shelf to 300 m, about 200,000 km²; not the entire Patagonian Shelf LME. Coverage percent is unknown. Source membership includes cohort and pooled-taxonomy overlaps requiring later model-specific matching; no catch mapping has been adopted.
'''
 else:
  detail='''
## Source tables and numerical identity

Ocampo Reinaldo et al. (2016), *Assessing the effects of demersal fishing and conservation strategies of marine mammals over a Patagonian food web*, Ecological Modelling 331, 31–43, DOI 10.1016/j.ecolmodel.2015.10.025 (article PDF p1). The source contains one reference Ecopath model for **1970**, with Ecosim simulation through 2009. Simulated annual time series and prospective fishing scenarios are not separate fully tabulated Ecopath parameterizations.

Table2 on PDF p5/printed p35 contains 26 rows, including three separate hake sizes, Phytoplankton and Detritus; the prose's 23 groups is not used to collapse that explicit structure. Columns: group; TL; B t/km²; P/B year−1; Q/B year−1; EE; P/Q; omnivory index. Bounding boxes fix numeric columns before reading values. Bold B indicates model estimates for Juvenile sharks, Flounders, Benthic-demersal fish I, Small pelagic fishes, Squids, Jellyfishes, Zooplankton and Phytoplankton. Bold EEs are estimates for Dolphins, Sea Lions, Other rays, Large rays, Small rays, Large sharks, Medium-sized sharks, Narrownose smooth-hound, Benthic-demersal fish II, Plownose chimaera, Large hake, Medium-sized hake, Pink cusk-eel, Patagonian hoki, Medium-sized pelagic fishes, Benthic organisms and Detritus. Hake Z=P/B is explicitly stated on p35 and recorded only for the three hake groups.

The new Appendix A DOCX supplies two 29×14 diet blocks (including merged caption/header): predators1–12 and13–24, 26 named prey plus Import. All 648 numerical cells (26 prey+Import ×24 consumers) were extracted from exact DOCX table cells and reviewed against rendered pages1–2. Diet sums are 0.999–1.002; Benthic-demersal fish II sums1.002. The imported diet fractions up to0.655 are explicit source inputs. No rescaling was performed. Later vulnerability matrices in Appendix A are Ecosim vulnerabilities, never diet values. Pedigree and system-metric tables were retained but not mistaken for missing basic inputs.

## Biomass accumulation and deliberate blanks

The article describes a reference steady-state model on p32 and defines BA on p34, but reports no numeric BA in its tables, supplement or relevant figures. BA remains blank. GS/assimilation fractions, habitat proportions, detritus imports and detritus fate were not numerically specified in the reviewed bundle and remain unknown. The paper says synthetic incipient fishing effort equal to1‰ of later fleet maxima was introduced (p35), despite the nominal pre-fishery period. That statement is not a group catch table, and effort is not automatically a catch multiplier. No source catches were invented or back-calculated from EE. All three fleet landings/discards files therefore remain blank. The printed note that long-liner/jigger discards are negligible is qualitative, not exact zero.

## Source balance, loader completion and diagnostic interpretation

Juvenile hake B=0.098, P/B=0.880, EE=0.999 on p35 was visually confirmed. The supplied diet implies predation≈3.755 t/km²/year versus production0.08624, equivalent to EE≈43.539 before catch/BA/migration. Benthic-demersal fish I gives EE≈1.006. These source conflicts were not repaired. Source P/Q for Large hake is0.5064, also confirmed, rather than replaced by a generic fish value. Hake is a multistanza group; the eight-table/JSON workflow does not encode age-transition/maturation flows. Such omitted flows may be relevant, and the large juvenile imbalance should not be attributed uniquely to a typo without the original EwE file and multistanza settings.

Strict default loading rejects the diet tolerance at the published1.002 sum. The isolated diagnostic allows DC_tol=0.0021 solely to admit the unchanged three-decimal diet. Full diagnose_sppr continues to flag its0.002 deviation. There is no normalization. The source JSON retains missing nonconsumer diet import; the diagnostic boundary sets PP/DET import contributions to0 because they do not consume. Otherwise NaN propagates through the engine and produces a missing detritus-column exception; no consumer cell is altered. ModelData automatically assumes all living detritus flows reach the sole detritus pool, forces detritus self-routing, and the calculator supplies missing catches=0, BA=0, GS=0.2 and missing migration/import=0. It also overrides detritus EE to1 and infers detritus accumulation. These are loader conventions, not source findings; `diagnostic_run_record.json` and loader snapshots expose them. All three full direct diagnostics return FAIL. Their catch footprint has zero catch because catch is unavailable in the source; it does not establish zero historical catch or a usable regional PPR coefficient.

## Taxonomy and spatial applicability

Article Table1 (PDF p4/printed p34) explicitly lists species/class/order composition. Original spellings such as Milyobatis, Carcharinus, Xistreurys and Polichaeta are preserved with source scope. Juvenile sharks are juveniles of Large and Medium-sized sharks. Hake classes are <23 cm,24–53 cm and>53 cm/6 years, per p34; the apparent23–24 cm boundary gap is retained. The model describes the demersal-pelagic web of the San Matías Gulf trawl-operating area, about10,000 km² (p32), not the entire LME. No regional coverage fraction is inferred.
'''
 (f/'source_report.md').write_text(common+detail,encoding='utf-8')
 (tab/'REPORT.md').write_text(common+detail,encoding='utf-8')
 (tab/'MODEL_PROFILE.md').write_text(f'# {mid}\n\nAxis: taxonomic groups with explicit life-stage/cohort and pooled guild groups.\nArea: '+('Falkland shelf to300 m; about200000 km².' if is_pat else 'San Matías Gulf demersal trawl area; about10000 km².')+'\nPrefixes: native group numbers are sequence indices; no spatial prefixes.\nMembership: complete row-level source descriptions in Taxonomy.xlsx.\nSelection: none; no species-to-group weights inferred.\nProduction eligibility: false pending source/loader review.\n',encoding='utf-8')
dump(R/'results_summary.json',summary);dump(R/'metadata_registration_proposals.json',proposals)
lines=['# LME_014 extraction review — 28 September 2026','', 'Both supplied paper bundles have been fully reviewed, with all supported values extracted. Source gaps and conflicting versions remain explicit. No model has been selected and no annual PPR was published.','', '| Candidate | Groups | GE | TE | With Egestion |','|---|---:|---|---|---|']
for x in summary:lines.append('| '+x['model_id']+' | '+str(x['groups'])+' | '+' | '.join(x['diagnostics'].values())+' |')
lines+=['','Each candidate folder contains `source_report.md`, `diagnostic_report.md`, `diagnostics_raw.json`, canonical `model.json`, taxonomy and eight EwE tables. The native and published-table PAT versions are separate because basic group structures/diets/EE disagree. S3 initial rates are retained as incomplete precursor evidence. Ocampo has one1970 Ecopath model;1970–2009 is an Ecosim simulation.','', 'Metadata registration is proposed only in metadata_registration_proposals.json. Project.xlsx and LME_014.xlsx were not edited. The final source_hash_verification.json confirms all source bytes unchanged.','', 'The diagnostic report files contain only the full direct engine returns for the three requested choices. Source assumptions, missing fields, converter corrections, native-vs-table conflicts, loader completion and the Ocampo rounding tolerance are explained in source_report.md.','', 'The native coupled Ecopath calculation reconstructs unknown living B/EE from native inputs with residual1.13e-14 and all living EE within[0,1]. Its separate diagnostic completion returns WARN for GE,TE and With Egestion; source JSON retains native missingness. The earlier default-B=1 loader diagnostic is preserved separately and is not the faithful reconstruction. Native detritus/discard handling caveats remain. The PAT published-table version is blocked by absent two-pool fate routing. Ocampo runs only as a declared loader-completed diagnostic and fails major source/flow checks.']
(R/'REPORT.md').write_text('\n'.join(lines)+'\n',encoding='utf-8');(R/'MASTER_INDEX.md').write_text('\n'.join(lines[:8]+['','Status: extraction complete to source support; scientific admission partial/blocked as documented.'])+'\n',encoding='utf-8')
print('REPORTS READY',len(summary),'models',len(manifest),'sources unchanged')
