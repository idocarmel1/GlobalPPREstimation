from pathlib import Path
import json,hashlib,csv,fitz
from docx import Document
import pyodbc
ROOT=Path(__file__).resolve().parents[4]
OUT=Path(__file__).resolve().parent/'evidence'
manifest=[]
for paper in ['PAT-2023','LME014-OcampoReinaldo-2016']:
 for p in (ROOT/'regions/LME_014/papers'/paper).iterdir():
  if not p.is_file():continue
  manifest.append(dict(path=p.relative_to(ROOT).as_posix(),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
  if p.suffix=='.pdf':
   doc=fitz.open(p)
   (OUT/(paper+'_article.txt')).write_text('\n'.join(f'\n=== PDF PAGE {i+1} ===\n'+pg.get_text() for i,pg in enumerate(doc)),encoding='utf-8')
   (OUT/(paper+'_words.json')).write_text(json.dumps([dict(page=i+1,words=pg.get_text('words')) for i,pg in enumerate(doc)]),encoding='utf-8')
  if p.suffix=='.docx':
   doc=Document(p)
   data={'file':p.name,'paragraphs':[x.text for x in doc.paragraphs],'tables':[[[c.text for c in row.cells] for row in tab.rows] for tab in doc.tables]}
   (OUT/(paper+'_supplement.json')).write_text(json.dumps(data,indent=2,ensure_ascii=False),encoding='utf-8')
  if p.suffix=='.eweaccdb':
   cn=pyodbc.connect('DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ='+str(p)+';READONLY=1;')
   cur=cn.cursor(); tables=[r.table_name for r in cur.tables(tableType='TABLE') if not r.table_name.startswith('MSys')]
   data={}
   for t in tables:
    q=cn.cursor().execute('SELECT * FROM ['+t+']'); cols=[x[0] for x in q.description]
    data[t]=[dict(zip(cols,row)) for row in q.fetchall()]
   (OUT/'PAT-2023_native_tables.json').write_text(json.dumps(data,indent=2,default=str,ensure_ascii=False),encoding='utf-8')
   cn.close()
   print('NATIVE TABLES',[(k,len(v)) for k,v in data.items()])
(OUT/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print('Inventory complete')
