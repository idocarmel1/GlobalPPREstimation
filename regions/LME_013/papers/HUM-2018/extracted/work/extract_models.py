from pathlib import Path
import sys,json,re,subprocess,csv
from decimal import Decimal
import openpyxl
sys.path.insert(0,r'C:\Users\idoca\.agents\skills\ecopath-extraction\scripts')
from pdfgrid import get_words
folder=Path(__file__).parents[2];work=folder/'extracted'/'work'
book=openpyxl.load_workbook(work/'supplement.xlsx',data_only=False)
def val(x): return None if x is None else str(x)
def norm(s):return re.sub(r'\s+',' ',str(s).strip()).replace('ﬁ','fi').replace('ﬀ','ff')
def write(data,key):
 path=work/(key+'.json');path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
 r=subprocess.run([sys.executable,'-X','utf8',r'C:\Users\idoca\.agents\skills\ecopath-extraction\scripts\write_outputs.py',str(path),'--outdir',str(folder/'extracted')],capture_output=True,text=True,encoding='utf-8');r.check_returncode();print(r.stdout);return Path(r.stdout.strip().splitlines()[-1])
groups=[];land={};disc={};prov=[]
sh=book['Table A-resolved parameters']
for row in list(sh)[5:44]:
 n,name=row[0].value,norm(row[1].value)
 assert n==len(groups)+1
 g={'n':n,'name':name,'source_group_code':n}
 for fld,c in zip(['biomass','pb','qb','pq','ae','ee'],row[2:8]):
  if c.value is None:continue
  assert c.data_type!='f'
  if fld=='ae':g['unassim']=str(Decimal(1)-Decimal(str(c.value)))
  else:g[fld]=val(c.value)
  prov.append([n,name,fld,sh.title,c.coordinate,str(c.value),c.number_format,bool(c.font.bold)])
 groups.append(g)
 land[str(n)]={f:val(row[col].value) for f,col in [('Artisanal',8),('Commercial',9)] if row[col].value is not None}
 disc[str(n)]={f:val(row[col].value) for f,col in [('Artisanal',10),('Commercial',11)] if row[col].value is not None}
names={norm(g['name']):g['n'] for g in groups}
diet={str(i):{} for i in range(3,36)};sh=book['Table B-resolved diet']
for c in range(2,41):
 name=norm(sh.cell(3,c).value)
 if name not in names:continue
 pred=names[name]
 if pred not in range(3,36):continue
 for r in range(4,44):
  pname=norm(sh.cell(r,1).value);prey='import' if pname=='Import' else str(names[pname])
  x=sh.cell(r,c).value
  if x is not None:diet[str(pred)][prey]=val(x)
fates={};sh=book['Table C-resolved detritus fate'];dets=[g['name'] for g in groups[35:]]
for r in range(4,43):
 n=names[norm(sh.cell(r,2).value)];assert n==sh.cell(r,1).value
 fates[str(n)]={norm(sh.cell(3,c).value):val(sh.cell(r,c).value) for c in range(3,8) if sh.cell(r,c).value is not None}
resolved={'metadata':{'LME':'13 Humboldt Current','model_number':'13_2','model_name':'Northern Humboldt Current','model_year':'1995-1998','variant':'fully resolved'},
'groups':groups,'consumers':list(range(3,36)),'fleets':['Artisanal','Commercial'],'landings':land,'discards':disc,'diet':diet,
'detritus_groups':dets,'detritus_fate':fates,'diet_rows':39}
p=write(resolved,'resolved-input')
with (p/'SOURCE_CELLS.csv').open('w',newline='',encoding='utf-8') as f:
 wr=csv.writer(f);wr.writerow(['group','name','parameter','sheet','cell','stored_value','source_number_format','bold_model_estimate']);wr.writerows(prov)
(p/'FLEET_DETRITUS_FATE.json').write_text(json.dumps({'Artisanal':{'Fishery offal':1},'Commercial':{'Fishery offal':1},'source':'Supplement Table C rows44-45; fleets treated separately from biological groups'},indent=2),encoding='utf-8')
# Aggregated source order 4..27 is remapped into contiguous import rows1..24, never add ECOTRAN nutrients1..3.
groups=[];land={};disc={};coords=[209.9,266.3,309.0,355.2,390.7,426.2,461.7,518.1]
for line in get_words(str(next(folder.glob('*.pdf'))),4,merge_gap=2):
 if not 105<line.y<312:continue
 cells=line.cells;src=int(cells[0].text);n=src-3
 name=norm(' '.join(c.text for c in cells if 98<c.x0<200))
 g={'n':n,'name':name,'source_group_code':src}
 for cell in cells:
  if cell.x0<200:continue
  i=min(range(8),key=lambda k:abs(cell.x0-coords[k]));assert abs(cell.x0-coords[i])<2
  fld=['biomass','pb','qb','pq','ae','ee','land','disc'][i]; x=cell.text;Decimal(x)
  if fld=='ae':g['unassim']=str(1-Decimal(x))
  elif fld=='land':land[str(n)]={'Fisheries':x}
  elif fld=='disc':disc[str(n)]={'Fisheries':x}
  else:g[fld]=x
 assert n==len(groups)+1;groups.append(g)
assert len(groups)==24
names={norm(g['name']):g['n'] for g in groups}
aliases={'Sea turtle':'Sea turtles','Eggs':'Fish eggs'}
def agg_n(name):return names[aliases.get(norm(name),norm(name))]
diet={str(i):{} for i in range(3,21)};sh=book['Table F-aggregated diet']
for c in range(2,26):
 name=norm(sh.cell(3,c).value)
 pred=agg_n(name)
 if pred not in range(3,21):continue
 for r in range(4,29):
  name=norm(sh.cell(r,1).value);prey='import' if name=='Import' else str(agg_n(name))
  x=sh.cell(r,c).value
  if x is not None:diet[str(pred)][prey]=val(x)
agg={'metadata':{'LME':'13 Humboldt Current','model_number':'13_3','model_name':'Northern Humboldt Current','model_year':'1995-1998','variant':'aggregated'},
'groups':groups,'consumers':list(range(3,21)),'fleets':['Fisheries'],'landings':land,'discards':disc,'diet':diet,'detritus_groups':[g['name'] for g in groups[20:]],'detritus_fate':{},'diet_rows':24}
p=write(agg,'aggregated-input')
with (p/'GROUP_NUMBER_MAP.csv').open('w',newline='',encoding='utf-8') as f:
 wr=csv.writer(f);wr.writerow(['import_group','paper_Table1_code','group_name']);wr.writerows((g['n'],g['source_group_code'],g['name']) for g in groups)
for data in [resolved,agg]:
 print(data['metadata']['model_number'],[(n,str(sum(Decimal(v) for v in d.values()))) for n,d in data['diet'].items() if abs(sum(Decimal(v) for v in d.values())-1)>Decimal('.01')])

