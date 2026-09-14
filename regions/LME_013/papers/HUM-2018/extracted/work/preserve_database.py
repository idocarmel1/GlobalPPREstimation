"""Local source-preserving adapter. Installed skill is unchanged."""
from pathlib import Path
import csv,json,sys,shutil,re,logging
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
sys.path.insert(0,r'C:\Users\idoca\.agents\skills\ecopath-extraction\scripts')
from database_json import EwEConverter
NODATA='-9999'
def readcsv(p):return list(csv.reader(p.open(encoding='utf-8-sig',newline='')))
def keep(x):return x if x not in [None,''] else NODATA
def adapt(folder):
 folder=Path(folder);model=json.loads((folder/'model.json').read_text(encoding='utf-8'))
 dbfile=next(f for f in folder.glob('*.json') if '_(' in f.name)
 rawdir=folder.parent/'work'/('raw-converter-'+model['metadata']['model_number']);rawdir.mkdir(parents=True,exist_ok=True)
 for p in [dbfile,folder/'MASS_BALANCE.md',folder/'ewe_conversion.log',dbfile.with_name(dbfile.stem+'_reconstructed.xlsx')]:
  if p.exists() and not (rawdir/p.name).exists():shutil.copy2(p,rawdir/p.name)
 data=json.loads(dbfile.read_text(encoding='utf-8'))
 # The writer substitutes zero for absent diet imports. Restore source silence.
 diet_csv=readcsv(folder/'Diet_composition.csv')
 import_row=next(r for r in diet_csv if len(r)>1 and r[1]=='Import')
 for ci,n in enumerate(diet_csv[0][2:],2):
  if model['diet'].get(n,{}).get('import') is None:import_row[ci]=''
 with (folder/'Diet_composition.csv').open('w',encoding='utf-8',newline='') as f:csv.writer(f,lineterminator='\r\n').writerows(diet_csv)
 source={p.name:readcsv(p) for p in folder.glob('*.csv') if p.name in ['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv']}
 for name in ['TL.xlsx','Metadata.xlsx']:
  b=openpyxl.load_workbook(folder/name,data_only=True);source[name]=[list(row) for row in b.active.values]
 basic=source['Basic_input.csv']; diet=source['Diet_composition.csv']; fate=source['Detritus_fate.csv']
 byn={str(g['n']):g for g in model['groups']}; name_to_n={g['name']:str(g['n']) for g in model['groups']}
 catch={n:{str(row[0]):row for row in source[n][1:]} for n in ['Landings.csv','Discards.csv']}
 dby={r[0]:r for r in diet[1:] if r[0]};imp=next(r for r in diet if len(r)>1 and r[1]=='Import')
 for g,row in zip(data['group'],basic[1:]):
  n=row[0]; assert n==g['group_seq']
  for field,c in {'habitat_area':2,'biomass_habitat_area':3,'z':4,'pb':5,'qb':6,'ee':7,'other_mort':8,'ge':9,'gs':10,'detritus_import':11}.items():g[field]=keep(row[c])
  g['biomass']=keep(row[3]);g['biomass_basis']='published model-area density; habitat fraction unspecified' if not row[2] else 'biomass_habitat_area times stated habitat_area'
  if row[2] and row[3]:g['biomass']=str(Decimal(row[2])*Decimal(row[3]))
  g['ge_input']='true' if row[9] else 'false'
  for field in ['vbk','shadow_price']:g[field]=NODATA
  for field,c in [('pb_input',5),('qb_input',6),('ee_input',7),('b_hab_area_input',3)]:g[field]='true' if row[c] else 'false'
  g['pp']='2' if row[1] in model['detritus_groups'] else ('0' if int(n) in model['consumers'] else '1')
  g['source_group_code']=byn[n].get('source_group_code',int(n))
  g['tl']=keep(source['TL.xlsx'][int(n)][2])
  sums=[]
  for fn,key in [('Landings.csv','landings'),('Discards.csv','discards')]:
   r=catch[fn].get(n)
   g[key+'_by_fleet']={fleet:keep(r[i]) for i,fleet in enumerate(source[fn][0][2:-1],2)} if r else {}
   g[key+'_total']=keep(r[-1]) if r else NODATA
   if r and r[-1]!='':sums.append(Decimal(r[-1]))
  g['export']=str(sum(sums)) if sums else NODATA
  g['export_basis']='sum of reported landings and discards only; omitted fields remain unknown separately'
  dc_col=diet[0].index(n) if n in diet[0] else None
  g['diet_imp']=keep(imp[dc_col]) if dc_col is not None else NODATA
  entries={};frow=fate[int(n)]
  for prey in range(1,len(data['group'])+1):
   p=str(prey);value=dby[p][dc_col] if dc_col is not None and p in dby else ''
   fname=byn[p]['name'];fvalue=frow[fate[0].index(fname)] if fname in fate[0] else ''
   if value!='' or fname in model['detritus_groups']:
    entries[p]={'prey_seq':p,'proportion':value if value!='' else '0','detritus_fate':keep(fvalue)}
  g['diet_descr']={'diet':list(entries.values())} if entries else None
  g['detritus_export']=keep(frow[fate[0].index('Export')])
 data['metadata']=model['metadata']
 data['source_tables']=source
 data['conversion_notes']=['Local adapter restores CSV precision and diet totals without normalization.','Unknown source habitat, growth K, prices, imports and detritus fate remain -9999.','ge carries source P/Q; z and tl extensions carry stated Z/TL.','source_tables preserves all eight original tables for lossless reconstruction, including fleet splits.']
 dbfile.write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
 # Rebuild all eight sheets exclusively from saved database JSON.
 saved=json.loads(dbfile.read_text(encoding='utf-8'));out=openpyxl.Workbook();out.remove(out.active)
 for name,rows in saved['source_tables'].items():
  ws=out.create_sheet(Path(name).stem[:31])
  for row in rows:ws.append(row)
  ws.freeze_panes='C2';ws.column_dimensions['A'].width=9;ws.column_dimensions['B'].width=33
  for c in ws[1]:c.font=Font(bold=True);c.fill=PatternFill('solid',fgColor='E9EEF3');c.alignment=Alignment(wrap_text=True,vertical='center')
  ws.row_dimensions[1].height=44
  for c in range(3,ws.max_column+1):ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=20
  if name=='Metadata.xlsx':ws.column_dimensions['A'].width=21;ws.column_dimensions['B'].width=34
 outpath=dbfile.with_name(dbfile.stem+'_reconstructed.xlsx');out.save(outpath)
 check=openpyxl.load_workbook(outpath,data_only=True);count=0
 for name,rows in saved['source_tables'].items():
  ws=check[Path(name).stem[:31]]
  for ri,row in enumerate(rows,1):
   for ci,v in enumerate(row,1):
    assert ws.cell(ri,ci).value==(None if v=='' else v),(name,ri,ci,v,ws.cell(ri,ci).value)
    count+=1
 # Cross-check database core vs input strings, including non-unit diet sums.
 for g,row in zip(saved['group'],basic[1:]):
  assert g['biomass_habitat_area']==keep(row[3]);assert g['pb']==keep(row[5]);assert g['qb']==keep(row[6]);assert g['ge']==keep(row[9])
  n=g['group_seq'];col=diet[0].index(n) if n in diet[0] else None
  if col is not None:
   entries={e['prey_seq']:e for e in g['diet_descr']['diet']}
   for p,r in dby.items():
    if r[col]!='': assert entries[p]['proportion']==r[col];count+=1
 converter=EwEConverter(str(folder/'ewe_conversion.log'))
 disc_total=sum(float(g['discards_total']) for g in saved['group'] if g['discards_total']!=NODATA)
 findings=converter.check_mass_balance(saved['group'],disc_total)
 converter.report_mass_balance(findings,model['metadata']['model_name'],str(folder/'MASS_BALANCE.md'))
 bad=[]
 for n,vals in model['diet'].items():
  total=sum(Decimal(x) for x in vals.values())
  if abs(total-1)>Decimal('.01'):bad.append(f"Group {n} ({byn[n]['name']}): diet+import = {total}")
 extra='\n\n### Source-fidelity and completeness limits\n\nThese findings were recomputed from the final, unnormalized JSON. Known-flow EE checks omit unknown catch/BA/migration. The checker uses GS=0.2 internally for missing GS in respiration/detritus diagnostics only; the JSON retains -9999. Unstated detritus routing prevents pool-specific validation.\n'
 extra+='\nDiet columns outside +/-0.01: '+('; '.join(bad) if bad else 'none')+'.\n'
 extra+='\nThe physiological check derives P/B divided by Q/B; any differently rounded stated P/Q is preserved separately in ge and Basic_input.csv.\n'
 if model['detritus_fate']:
  extra=extra.replace('Unstated detritus routing prevents pool-specific validation.','All group detritus routes are explicitly supplied in Table C. The bundled checker reports a pooled detritus diagnostic; it does not certify each routed pool.')
 with (folder/'MASS_BALANCE.md').open('a',encoding='utf-8') as f:f.write(extra)
 mass=(folder/'MASS_BALANCE.md').read_text(encoding='utf-8')
 mass=mass.replace('which is a lower bound on what the model may actually contain, not a steady-state result.','as a diagnostic with unreported terms omitted. It is not a bound because BA may be positive or negative.')
 (folder/'MASS_BALANCE.md').write_text(mass,encoding='utf-8')
 (folder/'ROUND_TRIP_AUDIT.md').write_text(f'# Source fidelity audit\n\n{count} exact cell comparisons passed across the saved JSON, eight import tables and reconstructed workbook. Core biomass, P/B, Q/B and P/Q strings and every stated diet cell match the CSVs. No diet normalization was applied. Unknown habitat fields remain -9999.\n\nUnmodified converter files are archived under ../work/raw-converter-{model["metadata"]["model_number"]}/. The final files require this local adapter when regenerating.\n',encoding='utf-8')
 with (folder/'ewe_conversion.log').open('a',encoding='utf-8') as f:f.write('\nSOURCE-PRESERVING ADAPTER: undid diet/fate normalization, retained all source strings, unknowns, fleet splits, P/Q, metadata and TL; reconstructed eight tables from JSON; exact comparisons '+str(count)+' passed.\n')
 print(model['metadata']['model_number'],findings['verdict'],count,bad)
if __name__=='__main__':
 for p in sys.argv[1:]:adapt(p)
