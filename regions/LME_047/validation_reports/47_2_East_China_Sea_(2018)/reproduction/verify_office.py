from pathlib import Path,PureWindowsPath
import json,zipfile,math,re,shutil,urllib.parse,sys,hashlib
import openpyxl
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[5];Q=Path(__file__).parent;R=ROOT/'regions/LME_047';MID='47_2_East_China_Sea_(2018)';O=R/'validation_reports'/MID;sys.stdout.reconfigure(encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
audit=json.loads((O/'taxon_audit.json').read_text(encoding='utf-8'));audit.sort(key=lambda r:(-(r['simple_chain_ppr_tC'] if r['simple_chain_ppr_tC'] is not None else -math.inf),r['taxon']))
doc=R/f'Model_validation_{MID}.docx';xlsx=R/'LME047_taxon_mapping_appendix.xlsx';targets=[];word_links=0
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';RR='http://schemas.openxmlformats.org/officeDocument/2006/relationships';NS={'w':W}
with zipfile.ZipFile(doc) as z:
 rel={r.get('Id'):r.get('Target') for r in E.fromstring(z.read('word/_rels/document.xml.rels')) if r.get('Type').endswith('/hyperlink')}
 d=E.fromstring(z.read('word/document.xml'))
 for h in d.findall('.//w:hyperlink',NS):
  if h.get('{'+RR+'}id'):targets.append((doc,rel[h.get('{'+RR+'}id')]))
  for r in h.findall('w:r',NS):
   pr=r.find('w:rPr',NS);assert pr.find('w:color',NS).get('{'+W+'}val')=='0563C1';assert pr.find('w:u',NS).get('{'+W+'}val')=='single'
  word_links+=1
wb=openpyxl.load_workbook(xlsx,data_only=False);s=wb['Taxon mapping'];assert wb.sheetnames==['Taxon mapping','Sources'];assert s.max_column==7 and s.max_row==283;assert s.freeze_panes=='B8';assert wb['Sources'].freeze_panes=='A5';assert len(s.tables)==1 and next(iter(s.tables.values())).autoFilter.ref=='A7:G283'
for i,(cells,a) in enumerate(zip(s.iter_rows(min_row=8,max_row=283,values_only=True),audit),8):
 assert cells[0]==a['taxon'] and cells[4]==a['mapping_display'] and cells[5]==a['overall_confidence'],(i,cells[0],a['taxon'])
 for actual,expected in [(cells[1],a['tl']),(cells[2],a['catch_tonnes']),(cells[3],a['simple_chain_ppr_tC'])]:
  assert actual=='?' if expected is None else isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=2e-14,abs_tol=1e-10)
 assert s.cell(i,7).alignment.wrap_text and s.row_dimensions[i].height>60
 assert 'Membership:' in cells[6] and 'allocation:' in cells[6]
assert len({s.cell(i,1).value for i in range(8,284)})==276
formula_errors=[];excel_links=0
for ss in wb:
 for row in ss:
  for c in row:
   if c.data_type=='e':formula_errors.append([ss.title,c.coordinate,c.value])
   if c.hyperlink:
    assert c.font.color.type=='rgb' and c.font.color.rgb=='FF0563C1' and c.font.underline=='single',(ss.title,c.coordinate)
    target=c.hyperlink.target or '#'+c.hyperlink.location;targets.append((xlsx,target));excel_links+=1
assert not formula_errors;wb.close()
support=O/'source_review.md'
for target in re.findall(r'\]\(((?:[^()\n]|\([^()\n]*\))*)\)',support.read_text(encoding='utf-8')):
 targets.append((support,target))
local=[];external=[];internal=[]
for source,t in targets:
 t=urllib.parse.unquote(t)
 if t.startswith('#'):internal.append(t);continue
 if t.startswith(('https://','http://')):external.append(t);continue
 assert not PureWindowsPath(t).drive and not t.startswith(('/','\\','file:','localhost:')),(source,t)
 target=(source.parent/t.split('#')[0]).resolve();assert target.exists(),(source,t);target.relative_to(ROOT);local.append((source,t,target))
for artifact in [doc,xlsx]:
 with zipfile.ZipFile(artifact) as z:
  for f in ['docProps/app.xml','docProps/custom.xml']:
   if f in z.namelist():assert all(not x.text for x in E.fromstring(z.read(f)).iter() if E.QName(x).localname=='HyperlinkBase')
rel=Q/'physical_relocation/repository';rel.mkdir(parents=True,exist_ok=True)
# Copy final artifacts and each directly referenced local evidence target under a different root.
copies={doc,xlsx,support,*[p for a,t,p in local]}
for p in sorted(copies,key=lambda p:len(str(p))):
 dest=rel/p.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True)
 if p.is_dir():shutil.copytree(p,dest,dirs_exist_ok=True,ignore=shutil.ignore_patterns('notes for AI.txt'))
 else:shutil.copy2(p,dest)
verified=[]
for source,t,p in local:
 relocated_source=rel/source.relative_to(ROOT);rp=(relocated_source.parent/t.split('#')[0]).resolve();assert rp.exists() and rp==rel/p.relative_to(ROOT)
 if p.is_file():assert sha(p)==sha(rp)
 verified.append({'from':source.relative_to(ROOT).as_posix(),'target':t,'relocated_exists':True})
v={'word_links_blue_underlined':word_links,'excel_links_blue_underlined':excel_links,'local_targets_physically_relocated':verified,'external_destinations_preserved':external,'internal_sheet_anchors':internal,'all_276_rows_numeric_sorted_exact_labels':True,'seven_columns_two_sheets':True,'filters_and_frozen_headers':True,'wrapped_readable_rows':True,'formula_errors':formula_errors,'all_blue_underlined_effective_styles':True,'no_absolute_HyperlinkBase':True,'appendix_sha256':sha(xlsx),'docx_sha256':sha(doc),'relocation_root':'Disjoint ignored scratch/physical_relocation/repository, preserving repository-relative hierarchy','limits':'Source destinations are preserved. Fresh remote retrieval is separate from this link-format/relative-target check. Versioned inclusion is coordinator staging responsibility; all local targets are inside regional deliverable/source trees.'}
(O/'office_and_portability_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8');print('OFFICE_VERIFIED',word_links,excel_links,len(local),'276rows')
