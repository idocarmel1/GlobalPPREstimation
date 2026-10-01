from pathlib import Path,PureWindowsPath
from collections import Counter,defaultdict
from urllib.parse import unquote,urlsplit
import json,sys,math,zipfile,shutil
from lxml import etree as E
import openpyxl
OUT=Path(__file__).parent;ROOT=OUT.parents[3];REGION=OUT.parents[1];sys.path.insert(0,str(ROOT/'tools'))
from workbooks import *
from regional import result_hash
def save(name,obj):(OUT/name).write_text(json.dumps(obj,indent=2,ensure_ascii=True,allow_nan=False),encoding='utf-8')
BASE=ROOT/'original_research_archive/research/selected_regions_validation_20260930/baseline/LME_026'
b=read_book(REGION/'LME_026.xlsx');old=read_book(BASE/'LME_026.xlsx');audit=json.loads((OUT/'adopted_taxon_audit.json').read_text(encoding='utf-8'));summary=json.loads((OUT/'coverage_summary.json').read_text(encoding='utf-8'));allocation=json.loads((OUT/'allocation_evidence.json').read_text(encoding='utf-8'))['records']
modelpath=REGION/'models/Piroddi_2022_Mediterranean_1995/model.json';model=json.loads(modelpath.read_text(encoding='utf-8'));groups={int(g['group_seq']):g for g in model['group']}
protected={s:digest_tables(old[s])==digest_tables(b[s]) for s in ['Catch','NPP']};protected['classic_TL_coefficients']=digest_tables(old['Classic PPR']['Taxa'])==digest_tables(b['Classic PPR']['Taxa']);protected['canonical_model_bytes']=modelpath.read_bytes()==(BASE/'accepted_model.json').read_bytes()
for key in ['selected_model_id','model_path','selection_rationale','catch_basis','TE','taxon_detail_year','production_eligible']:protected['Overview/'+key]=overview(old).get(key)==overview(b).get(key)
for n in old['Diagnostics']:protected['historical_diagnostic/'+n]=digest_tables(old['Diagnostics'][n])==digest_tables(b['Diagnostics'][n])
assert all(protected.values()),protected
inv=records(b,'Selected model groups','Groups');assert len(inv)==71
for r in inv:
    g=groups[r['seq']]
    for k,v in g.items():assert (json.loads(r[k]) if isinstance(v,(dict,list)) else r[k])==v,(r['seq'],k)
validate_region(b,REGION/'LME_026.xlsx');assert overview(b)['calculation_result_sha256']==result_hash(b)
assert not records(b,'Selected model groups','Group SPPR') and not records(b,'PPR','Taxon SPPR') and not records(b,'PPR','Annual')
taxa={r['taxon'] for r in records(b,'Catch','Catch') if r['catch_basis']=='landings'};assert len(audit)==522 and {r['taxon'] for r in audit}==taxa
rank={'High':0,'Medium':1,'Low':2,'Very low':3,'Unresolved':4};matches=defaultdict(list)
for r in records(b,'PPR','Matching'):matches[r['taxon']].append(r)
for r in audit:
    assert r['overall_confidence']==max([r['membership_confidence'],r['allocation_confidence']],key=rank.get)
    assert math.isclose(math.fsum(g['weight'] for g in r['groups']),1,abs_tol=1e-12,rel_tol=0)
    assert {g['group_name'] for g in r['groups'] if g['weight']>0}=={g['group'] for g in matches[r['taxon']]}
    for g in r['groups']:
        assert groups[g['seq']]['group_name']==g['group_name'] and finite(g['weight']) and g['weight']>=0
        if g['weight']>0:assert math.isclose(next(m['weight'] for m in matches[r['taxon']] if m['group']==g['group_name']),g['weight'],rel_tol=1e-14,abs_tol=1e-14)
    expected=0. if r['catch_tonnes']==0 else r['catch_tonnes']*r['classic_coefficient_wet']/9 if finite(r['catch_tonnes']) and finite(r['classic_coefficient_wet']) else None
    assert expected==r['simple_chain_ppr_tC']
assert len(summary['missing_classic_coefficients'])==40 and not summary['positive_catch_unknown_ppr']
pp=math.fsum(r['simple_chain_ppr_tC'] for r in audit);tc=math.fsum(r['catch_tonnes'] for r in audit)
annual=next(r for r in records(b,'Classic PPR','Annual') if r['scope']=='all' and r['catch_basis']=='landings' and r['unidentified']=='method' and r['metric']=='ppr')
assert math.isclose(annual[2019]/9,pp,rel_tol=2e-15)
for key in ['membership_rules','allocation_rules']:assert abs(math.fsum(r['ppr_percent'] for r in summary[key])-100)<1e-8
assert sum(v['taxa_count'] for v in summary['categories'].values())==522
assert math.isclose(math.fsum(v['ppr_tC'] for v in summary['categories'].values()),pp,rel_tol=2e-15)
app=REGION/'LME026_taxon_mapping_appendix.xlsx';w=openpyxl.load_workbook(app,data_only=False);s=w['Taxon mappings'];ss=w['Sources']
assert w.sheetnames==['Taxon mappings','Sources'] and s.max_row==528 and s.max_column==7
rows=list(s.iter_rows(min_row=7,values_only=True));assert len(rows)==522 and {r[0] for r in rows}==taxa
last=float('inf');am={r['taxon']:r for r in audit}
for r in rows:
    assert r[3]<=last;last=r[3];a=am[r[0]]
    assert math.isclose(r[2],a['catch_tonnes'],rel_tol=1e-14,abs_tol=1e-14) and math.isclose(r[3],a['simple_chain_ppr_tC'],rel_tol=1e-14,abs_tol=1e-14) and r[5]==a['overall_confidence']
assert len(s.tables)==1 and s.freeze_panes and len(ss.tables)==1 and ss.freeze_panes
links=[]
for cellrow in ss.iter_rows(min_row=5):
    c=cellrow[4];assert c.value=='Open source' and c.hyperlink and c.font.color.rgb=='FF0563C1' and c.font.underline=='single';links.append({'container':app.name,'target':c.hyperlink.target,'label':c.value,'cell':c.coordinate,'blue_underlined':True})
assert not any(c.data_type=='e' or c.data_type=='f' for ws in w for row in ws for c in row)
w.close()
report=REGION/'Model_validation_Piroddi_2022_Mediterranean_1995.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
with zipfile.ZipFile(report) as z:
    doc=E.fromstring(z.read('word/document.xml'));rels=E.fromstring(z.read('word/_rels/document.xml.rels'));rm={r.get('Id'):r.get('Target') for r in rels}
    for h in doc.findall('.//w:hyperlink',ns):
        if h.get('{'+ns['r']+'}id'):
            runs=h.findall('w:r',ns);assert runs
            for r in runs:
                assert r.find('w:rPr/w:color',ns).get('{'+ns['w']+'}val')=='0563C1' and r.find('w:rPr/w:u',ns).get('{'+ns['w']+'}val')=='single'
            links.append({'container':report.name,'target':rm[h.get('{'+ns['r']+'}id')],'label':''.join(h.itertext()),'blue_underlined':True})
    assert not doc.findall('.//w:instrText',ns),'Unexpected field hyperlink requires audit'
    if 'docProps/app.xml' in z.namelist():assert b'HyperlinkBase' not in z.read('docProps/app.xml') or b'<HyperlinkBase/>' in z.read('docProps/app.xml')
# Sparse relocated repository: copy every actual local file destination plus both deliverables.
relocated=OUT/'qa/relocated_repository';relocated.mkdir(exist_ok=True)
copy_targets=set([report,app]);errors=[]
for l in links:
    url=l['target'];parts=urlsplit(url)
    if parts.scheme in ['https','http','mailto']:l['public_url']=True;continue
    ptext=unquote(url.split('#',1)[0]);assert not Path(ptext).is_absolute() and not PureWindowsPath(ptext).drive
    p=(REGION/ptext).resolve();l['resolved_repository_path']=p.relative_to(ROOT).as_posix();l['anchor']=url.split('#',1)[1] if '#' in url else None
    if not p.exists():errors.append(l);continue
    copy_targets.add(p)
    if p.suffix=='.xlsx' and l['anchor']:
        wsname=l['anchor'].split('!',1)[0];cw=openpyxl.load_workbook(p,read_only=True);assert wsname in cw.sheetnames;cw.close()
assert not errors,errors
for p in copy_targets:
    target=relocated/p.relative_to(ROOT)
    if p.is_dir():target.mkdir(parents=True,exist_ok=True)
    else:target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,target)
for l in links:
    if l.get('public_url'):continue
    container=relocated/REGION.relative_to(ROOT)/l['container'];dest=(container.parent/unquote(l['target'].split('#',1)[0])).resolve();assert dest.exists();l['relocated_resolution']=True
save('qa/hyperlink_audit.json',{'links':links,'all_relative_repository_targets_resolve':True,'relocated_repository_check':True,'all_effective_link_text_blue_underlined':True,'formula_hyperlinks':0,'hyperlink_base_empty':True})
save('qa/protected_input_checks.json',{'checks':protected,'all_passed':True,'source_inventory_exact_accepted_fields':True,'canonical_model_sha256':sha(modelpath),'baseline_workbook_sha256':sha(BASE/'LME_026.xlsx'),'final_workbook_sha256':sha(REGION/'LME_026.xlsx')})
save('qa/reconciliation.json',{'taxa':522,'matching_rows':sum(map(len,matches.values())),'source_groups':71,'usable_model_coefficients':0,'model_annual_rows':0,'protected_inputs_preserved':True,'weights_valid':True,'separate_confidence_weakest_component':True,'total_catch_tonnes':tc,'total_simple_chain_ppr_tC':pp,'classic_annual_matches':True,'missing_coeff_zero_catch_labels':40,'positive_catch_unknown_ppr':0,'appendix_rows_match':True,'appendix_numeric_sort':True,'rule_tables_each_sum_100':True,'all_links_resolve':True,'link_count':len(links),'word_page_count':6,'model_methods':'GE/TE/With Egestion NOT_RUN','scientific_approval':False})
print(json.dumps({'passed':True,'links':len(links),'protected_checks':len(protected),'total_simple_chain_ppr_tC':pp},ensure_ascii=True))
