from pathlib import Path
from copy import deepcopy
from io import BytesIO
import json,zipfile,hashlib,sys
from lxml import etree
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
OUT=Path(__file__).parent;ROOT=OUT.parents[3];REGION=OUT.parents[1]
template=ROOT/'tools/templates/Model_validation_template.docx'
audit=json.loads((OUT/'adopted_taxon_audit.json').read_text(encoding='utf-8'))
summary=json.loads((OUT/'coverage_summary.json').read_text(encoding='utf-8'))
def percent(v):return '<0.0001%' if 0<v<0.0001 else f'{v:.4f}%'
doc=Document(template)
doc.paragraphs[0].text='Mediterranean Sea model validation'+('' if '--aligned' in sys.argv else ' draft')
doc.paragraphs[0].style='Title'
def link(p,label,target):
 h=OxmlElement('w:hyperlink');rid=p.part.relate_to(target,RT.HYPERLINK,is_external=True);h.set(qn('r:id'),rid)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def set_cell(cell,text,links=()):
 p=cell.paragraphs[0];p.clear();p.add_run(text)
 for p2 in list(cell.paragraphs)[1:]:p2._element.getparent().remove(p2._element)
 if links:
  p=cell.add_paragraph()
  for i,(label,target) in enumerate(links):
   if i:p.add_run(' | ')
   link(p,label,target)
base='validation_reports/Piroddi_2022_Mediterranean_1995/'
rows={r.cells[0].text:r for tb in doc.tables[:2] for r in tb.rows[1:]}
texts={
 'Region':('Mediterranean Sea, LME_026',[('Regional workbook','LME_026.xlsx')]),
 'Catch source':('Sea Around Us reconstructed catches, 1950–2019. Reference: 2019 landings, excluding discards. Frozen regional extract.',[('Sea Around Us region','https://www.seaaroundus.org/data/#/lme/26'),('Catch files','raw/LME_026.csv.gz')]),
 'Selected article':('Piroddi, Coll, Macias, Steenbeek, Garcia-Gorriz, Mannini, Vilas and Christensen (2022). Modelling the Mediterranean Sea ecosystem at high spatial resolution to inform the ecosystem-based management in the region. Scientific Reports 12, 19680.',[('DOI','https://doi.org/10.1038/s41598-022-18017-x'),('Paper PDF','papers/MED-2022/s41598-022-18017-x-75a92a71.pdf'),('Paper folder','papers/MED-2022')]),
 'Other known articles':('Piroddi et al. (2017). Historical changes of the Mediterranean Sea ecosystem: modelling the role and impact of primary productivity and fisheries changes over time. Scientific Reports 7, 44491. 1950s baseline with 1950–2011 simulations and four sub-basin networks.',[('2017 article','https://doi.org/10.1038/srep44491'),('Local PDF','papers/LME026-Piroddi-2017/srep44491-e23185b8.pdf')]),
 'Selected model':('Piroddi_2022_Mediterranean_1995. Basin-wide 1995 baseline (supplement: 1990s); 71 source groups: 65 consumers, 4 producers and 2 nonliving pools; 37 fleet entries.',[('Model folder','models/Piroddi_2022_Mediterranean_1995'),('Selected JSON','models/Piroddi_2022_Mediterranean_1995/model.json')]),
 'Other models from the same article':('One distinct Ecopath baseline is supplied. The 1995–2016 Ecosim/Ecospace simulations and spatial outputs are dynamic applications of that baseline, with no separately tabulated alternative Ecopath parameterization.',[]),
 'Selection rationale':('Article — The recorded recommendation prefers the recent whole-basin source with supplied parameter-bearing XLSX and DOCX supplements.\nModel — Explicit user selection of the extracted 1995 model despite blocked construction; missing routing and the conditional fin-whale conflict remain.',[]),
 'Model extraction':('Source B/PB/QB/EE, diets and fleet values are preserved. Diet sums including imports span 0.99956013582–1.00000013576; no normalization. GS, BA, migration and routing remain unknown. Diagnostic staging assumes zero unlisted landings but preserves reported discards; construction stops before requested GS/BA defaults. Fin-whale discards exceed production under the stated zero-BA/net-migration arithmetic; the residual is unexplained and unapplied.',[('Source audit','models/Piroddi_2022_Mediterranean_1995/extracted_tables/REPORT.md'),('Lossless roundtrip','models/Piroddi_2022_Mediterranean_1995/extracted_tables/SOURCE_ROUNDTRIP.xlsx')]),
 'GE diagnostics':('NOT_RUN: missing routing between Discards (70) and Detritus (71) blocks construction. Negative SPPR entries cannot be assessed. rho_living and b are unavailable; SPPR for Discards (70) and Detritus (71) is unavailable.',[]),
 'TE diagnostics':('NOT_RUN: the same constructor restriction prevents calculation. Negative SPPR entries and rho_living cannot be assessed; SPPR for Discards (70) and Detritus (71) is unavailable.',[]),
 'Geographic fit':('A. Region covered by study: approximately 95–100%.\nB. Study area covered by region: approximately 95–100%.',[]),
 'Temporal fit':('The 1995 baseline precedes reference catch year 2019 by 24 years. Fixed source composition and stage proxies are transferred across 1950–2019 catches and bases. These do not reconstruct annual ecosystem structure; model PPR remains unavailable.',[]),
 'Other':('With Egestion also remains NOT_RUN. Unknown routing cannot be replaced by a source balance check or a taxonomic mapping.',[('Constructor and method evidence','models/Piroddi_2022_Mediterranean_1995/diagnostics/DIRECT_SPPR_REPORT.md')])
}
for label,(text,links) in texts.items():set_cell(rows[label].cells[1],text,links)
# All three manually maintained rows stay byte-equivalent under exclusive C14N.
manual_labels=['SPPR calculation','Open issues and next action','Review and reproducibility']
manual_before={k:etree.tostring(rows[k]._tr,method='c14n',exclusive=True) for k in manual_labels}
for p in list(doc.paragraphs):
 text=p.text
 replacements={
 'Reference: [year and catch basis]; [source scope and taxon/group filters].':'Reference: 2019 landings. Independent all-taxon simple-chain estimate; all source support, all recorded labels and no group filter.',
 '[n] source groups [+ synthetic imports separately]; [n] catch taxa.':'71 source groups; 522 catch taxa. No synthetic computational import group.',
 'Simple-chain PPR totals [value] t C in [year].':f"Simple-chain PPR totals {summary['total_simple_chain_ppr_tC']:,.3f} t C in 2019.",
 '[If applicable: n taxa lack assigned TL/coefficient. Zero recorded catch gives zero PPR; positive catch with no coefficient has unknown PPR.]':'40 taxa lack a classic TL/coefficient; all have zero 2019 landings and contribute zero annual PPR. No positive-catch contribution is unknown.',
 'Simple-chain PPR is calculated with [method name].':'Simple-chain PPR uses the fixed taxon coefficients of the simple trophic chain method.',
 'Each percentage is the simple-chain PPR associated with taxa using that rule. [Use known PPR only when unknown annual contributions remain.]':'Each rule percentage is the share of independent simple-chain PPR associated with taxa using that rule.',
 'Target region R and selected model study area S. [Overlap method and approximation, if applicable.]':'Target region R and selected-model study area S. The basin-wide geographic estimate is approximate; coastal grid cells and strait boundaries remain uncertain.',
 }
 if text in replacements:p.text=replacements[text]
 if text=='[Linked Excel taxon appendix and descriptive Sources sheet]':
  p.clear();link(p,'All-taxon Excel appendix','LME026_taxon_mapping_appendix.xlsx');p.add_run(' | ');link(p,'Descriptive Sources sheet','LME026_taxon_mapping_appendix.xlsx#Sources!A1')
 if text.startswith('[Add a Sum row') or text.startswith('[If a figure is unavailable'):p._element.getparent().remove(p._element)
 if text.startswith('[Unresolved taxa:'):
  unresolved=[r for r in audit if r['overall_confidence']=='Unresolved']
  vl=summary['categories']['Very low']
  p.text=('Unresolved taxa: '+('; '.join(r['taxon']+': '+r['membership_reason'] for r in unresolved) if unresolved else 'none.')+f" Very low approximations account for {vl['catch_percent']:.4f}% of landings and {vl['ppr_percent']:.4f}% of simple-chain PPR. They concern broad mixed categories and absent-compartment analogues; exact taxa, assumptions and weights are flagged in the appendix.")
 if text=='[Region name/ID; boundary source/version | Source | Image]':
  p.clear();p.add_run('Mediterranean Sea LME_026. Current project marine polygon. ');link(p,'Boundary source','../../common_reference_data/geography/LMEs.geojson');p.add_run(' | ');link(p,'Geographic assessment',base+'geography/geographic_assessment.json')
 if text=='[Article; figure/page; represented model area | Source PDF | Image]':
  p.clear();p.add_run('Piroddi et al. (2022), Figure 4h, p. 9. Whole Mediterranean ocean domain and four sub-basins; Figure 4a provides the matching marine mask. ');link(p,'Source PDF','papers/MED-2022/s41598-022-18017-x-75a92a71.pdf#page=9')
cat=doc.tables[2]
for c,v in zip(cat.rows[0].cells,['Overall confidence','Taxa (n)','Catch (%)','Simple-chain PPR (%)']):set_cell(c,v)
for row,level in zip(cat.rows[1:],['High','Medium','Low','Very low','Unresolved']):
 x=summary['categories'][level]
 for cell,value in zip(row.cells,[level,str(x['taxa_count']),f"{x['catch_percent']:.4f}%",f"{x['ppr_percent']:.4f}%"]):set_cell(cell,value)
rule_words={
 'M1':'Explicitly listed source member of the selected group or stage pair',
 'M2':'Verified synonym of an explicitly listed source member',
 'M3':'Unambiguous fit to a documented taxonomic group with no competing pool',
 'M4':'Taxonomic extension from source-listed representatives',
 'M5':'Supported ecological extension of the source-defined group',
 'M6':'Partial or conflicting ecological fit among plausible source pools',
 'M7':'Broad regional reporting category with assumed composition',
 'M8':'No meaningful source-group assignment',
 'M9':'Assumed eligible group set supported by source members and taxonomy',
 'M10':'Broad reporting category approximated by named and residual source pools',
 'M11':'Closest represented taxonomic or ecological analogue for an absent compartment',
 'M12':'Evidence-supported transfer from a compatible nearby model',
 'W1':'One group receives the complete taxon catch; no split uncertainty',
 'W4':'Selected-model caught-mass proportions used as a fixed composition proxy',
 'W9':'Complete accepted model-biomass proportions when source landings are unusable',
 'W11':'Explicit last-resort numerical allocation among meaningful candidates',
 'W8':'No usable numerical allocation'
}
for tb,key in [(doc.tables[3],'membership_rules'),(doc.tables[4],'allocation_rules')]:
 prototype=deepcopy(tb.rows[1]._tr)
 for row in list(tb.rows)[1:]:tb._tbl.remove(row._tr)
 for rec in sorted(summary[key],key=lambda r:-r['ppr_percent']):
  tb._tbl.append(deepcopy(prototype));row=tb.rows[-1]
  for cell,value in zip(row.cells,[rule_words[rec['rule']],rec['confidence'],percent(rec['ppr_percent'])]):set_cell(cell,value)
for p in doc.paragraphs:
 if p.text=='Study area in the article':p.paragraph_format.page_break_before=True
 if p.text=='Group assignment rules':
  np=OxmlElement('w:p');p._p.addnext(np)
  from docx.text.paragraph import Paragraph
  q=Paragraph(np,p._parent);q.add_run('Membership evidence: ');link(q,'Source group definitions and authoritative taxonomic audit',base+'taxonomy_audit/taxonomy_mapping_proposals.json')
 if p.text=='Allocation weight rules':
  p.paragraph_format.page_break_before=True
  np=OxmlElement('w:p');p._p.addnext(np)
  from docx.text.paragraph import Paragraph
  q=Paragraph(np,p._parent);q.add_run('Weights and assumptions: ');link(q,'All candidates and source quantities',base+'allocation_evidence.json');q.add_run(' | ');link(q,'Primary-source composition search',base+'allocation_search.json')
for tb,img in [(doc.tables[5],'target_region.png'),(doc.tables[6],'study_area_figure_4h.png')]:
 cell=tb.rows[1].cells[0];p=cell.paragraphs[0];p.clear();p.add_run().add_picture(str(OUT/'geography'/img),width=Inches(6.6))
 # Remove placeholder fixed image height, allowing the retained cell to fit pixels.
 for trh in tb.rows[1]._tr.xpath('./w:trPr/w:trHeight'):trh.getparent().remove(trh)
for tb in doc.tables:
 pr=tb.rows[0]._tr.get_or_add_trPr()
 if not pr.xpath('./w:tblHeader'):pr.append(OxmlElement('w:tblHeader'))
dest=REGION/'Model_validation_Piroddi_2022_Mediterranean_1995.docx'
buf=BytesIO();doc.save(buf)
# Preserve every original package part outside the authored document/media slots.
with zipfile.ZipFile(template) as z:original={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(BytesIO(buf.getvalue())) as z:final={n:z.read(n) for n in z.namelist()}
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
for n,data in original.items():
 if n not in editable:final[n]=data
with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in final.items():z.writestr(n,data)
check=Document(dest);rows2={r.cells[0].text:r for tb in check.tables[:2] for r in tb.rows[1:]}
preserved={k:manual_before[k]==etree.tostring(rows2[k]._tr,method='c14n',exclusive=True) for k in manual_labels}
assert all(preserved.values()),preserved
changed=[n for n in original if original[n]!=final[n]]
assert set(changed)<=editable,changed
(OUT/'qa/docx_integrity.json').write_text(json.dumps({'template_sha256':hashlib.sha256(template.read_bytes()).hexdigest(),'report_sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'manual_rows':preserved,'changed_original_parts':changed,'section_xml_preserved':etree.tostring(Document(template).sections[0]._sectPr)==etree.tostring(check.sections[0]._sectPr)},indent=2),encoding='utf-8')
print(json.dumps({'report':str(dest),'manual_rows_preserved':preserved},ensure_ascii=True))
