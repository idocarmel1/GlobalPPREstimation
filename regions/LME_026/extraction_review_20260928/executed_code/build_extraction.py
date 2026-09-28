"""Reproduce the 1995 Mediterranean baseline without filling unpublished inputs."""
from pathlib import Path
from decimal import Decimal
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import json, csv, hashlib, re, sys, subprocess, copy
from openpyxl import load_workbook, Workbook
from docx import Document

ROOT=Path(__file__).resolve().parents[3]
REVIEW=Path(__file__).parent
SRC=ROOT/'regions/LME_026/papers/MED-2022'
MID='Piroddi_2022_Mediterranean_1995'
MD=ROOT/'regions/LME_026/models'/MID
MD.mkdir(parents=True,exist_ok=True)
ED=MD/'extracted_tables'
SKILL=Path('C:/Users/idoca/.agents/skills/ecopath-extraction/scripts')
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf8')
xlsx=next(SRC.glob('*.xlsx')); wb=load_workbook(xlsx,data_only=False)
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
raw={}
with ZipFile(xlsx) as z:
 for i,s in enumerate(wb,1):
  raw[s.title]={c.attrib['r']:c.find('s:v',ns).text for c in ET.fromstring(z.read(f'xl/worksheets/sheet{i}.xml')).findall('.//s:c',ns) if c.find('s:v',ns) is not None and c.attrib.get('t') not in ['s','str','inlineStr']}
evidence=[]
def value(s,r,c,field,seq=None,**kw):
 cell=wb[s].cell(r,c);v=raw[s].get(cell.coordinate) if isinstance(cell.value,(int,float)) else cell.value
 evidence.append({'file':xlsx.name,'sheet':s,'cell':cell.coordinate,'field':field,'group_seq':seq,'value':v,'number_format':cell.number_format,**kw})
 return v
groups=[]
for r in range(5,76):
 n=wb['Basic input parameters'].cell(r,1).value
 assert n==r-4
 g={'n':n,'name':wb['Basic input parameters'].cell(r,2).value.strip()}
 for c,k in [(3,'tl'),(4,'biomass'),(5,'pb'),(6,'qb'),(7,'ee'),(8,'pq')]:g[k]=value('Basic input parameters',r,c,k,n)
 groups.append(g)
diet={}
for c in range(3,68):
 pred=wb['Diets'].cell(3,c).value;assert pred==c-2
 diet[str(pred)]={str(n):value('Diets',n+3,c,'diet',pred,prey_seq=n) for n in range(1,72)}
 diet[str(pred)]['import']=value('Diets',75,c,'diet_import',pred)
fleets=[wb['Catches'].cell(r,2).value.strip() for r in range(5,42)]
assert len(fleets)==len(set(fleets))==37
landings={};discards={};crosswalk=[];excluded=[]
for s,dest in [('Catches',landings),('Discards',discards)]:
 for c in range(3,wb[s].max_column+1):
  n=wb[s].cell(3,c).value
  if n is None:
   excluded.append({'sheet':s,'column':c,'name':wb[s].cell(4,c).value,'reason':'Unnumbered aggregate multistanza heading; no numeric fleet data','nonempty_values':[wb[s].cell(r,c).value for r in range(5,42) if wb[s].cell(r,c).value is not None]})
   continue
  n=int(n);dest[str(n)]={}
  crosswalk.append({'sheet':s,'column':c,'group_seq':n,'source_name':wb[s].cell(4,c).value,'canonical_name':groups[n-1]['name']})
  for r in range(5,42):
   fleet=wb[s].cell(r,2).value.strip();assert fleet==fleets[r-5]
   dest[str(n)][fleet]=value(s,r,c,s.lower(),n,fleet=fleet)
assert len(landings)==49 and len(discards)==71
assert not any(e['nonempty_values'] for e in excluded)
model={'metadata':{'LME':'26','model_number':2602022,'model_name':'Mediterranean Piroddi','model_year':1995},'groups':groups,'consumers':list(range(1,66)),'fleets':fleets,'landings':landings,'discards':discards,'detritus_groups':['Discards','Detritus'],'detritus_fate':{},'diet':diet,'diet_rows':71,'landings_rows':71,'discards_rows':71}
save(REVIEW/'extraction_input.json',model);save(REVIEW/'cell_evidence.json',evidence)
save(REVIEW/'group_name_crosswalk.json',crosswalk);save(REVIEW/'excluded_aggregate_columns.json',excluded)
p=subprocess.run([sys.executable,str(SKILL/'write_outputs.py'),str(REVIEW/'extraction_input.json'),'--outdir',str(MD),'--dir-name','extracted_tables'],capture_output=True,text=True,encoding='utf8',errors='replace');(REVIEW/'writer.log').write_text(p.stdout+p.stderr,encoding='utf8');assert p.returncode==0,p.stderr

doc=Document(next(SRC.glob('*.docx')));table=doc.tables[1];heads={}
for i,row in enumerate(table.rows):
 text=row.cells[0].text.strip();m=re.match(r'^(\d+)\.\s*(.*)',text,re.S)
 if m:heads[int(m[1])]={'table':'S2','docx_table_one_based':2,'row_one_based':i+1,'header':text,'description':m[2]}
tax=[];taxevidence=[]
for g in groups:
 n=g['n'];parent={22:21,24:23,30:29,38:37}.get(n,n)
 if parent in heads:
  h=heads[parent];desc=h['description']
  if n in [21,22,23,24,29,30,37,38]:desc+='; source basic-input life stage: '+g['name']
  taxevidence.append({'group_seq':n,'group_name':g['name'],**h})
 else:
  desc=('Nonliving pool: '+g['name']+'; no taxonomic membership applies.' if n in [70,71] else 'not documented — Table S2 has no numbered composition heading for this group; XLSX group name: '+g['name'])
  taxevidence.append({'group_seq':n,'group_name':g['name'],'source':'Basic input parameters row '+str(n+4),'description':desc})
 tax.append([n,g['name'],desc])
tw=Workbook();ts=tw.active;ts.title='Taxonomy';ts.append(['seq','group_name','taxon_descr'])
for row in tax:ts.append(row)
tw.save(ED/'Taxonomy.xlsx')
with (ED/'taxonomy.csv').open('w',encoding='utf8',newline='') as f:
 w=csv.writer(f);w.writerow(['seq','group_name','taxon_descr']);w.writerows(tax)
save(REVIEW/'taxonomy_evidence.json',taxevidence)
for script,args,log in [('validate.py',[str(ED)],'validate'),('massbalance_check.py',[str(ED)],'source_massbalance'),('database_json.py',['-d',str(ED)],'converter')]:
 p=subprocess.run([sys.executable,str(SKILL/script),*args],capture_output=True,text=True,encoding='utf8',errors='replace');(REVIEW/f'{log}.log').write_text(p.stdout+p.stderr,encoding='utf8');print(log,p.returncode)
converted=next(ED.glob('26_2602022*.json'));source=json.loads(converted.read_text());save(REVIEW/'converter_raw.json',source);changes=[]
for sg,g,tr in zip(source['group'],groups,tax):
 old=copy.deepcopy(sg);n=g['n'];sg['habitat_area']='-9999';sg['tl']=g['tl'];sg['ge']=g['pq'] or '-9999';sg['ge_input']='true' if g['pq'] is not None else 'false';sg['taxon_descr']=tr[2]
 for sf,tf in [('biomass','biomass'),('biomass','biomass_habitat_area'),('pb','pb'),('qb','qb'),('ee','ee')]:sg[tf]=g[sf] if g[sf] is not None else '-9999'
 sg['diet_imp']=diet.get(str(n),{}).get('import','-9999');sg['detritus_import']='-9999'
 # Every source discard is retained independently even where landings are not tabulated.
 sg['source_landings']=str(sum((Decimal(v) for v in landings[str(n)].values()),Decimal(0))) if str(n) in landings else '-9999'
 sg['source_discards']=str(sum((Decimal(v) for v in discards[str(n)].values()),Decimal(0)))
 sg['export']=str(Decimal(sg['source_landings'])+Decimal(sg['source_discards'])) if str(n) in landings else '-9999'
 entries=[]
 for prey in range(1,72):
  v=diet.get(str(n),{}).get(str(prey))
  if v is not None or prey in [70,71]:entries.append({'prey_seq':str(prey),'proportion':v if v is not None else '-9999','detritus_fate':'-9999'})
 sg['diet_descr']={'diet':entries}
 for k,v in sg.items():
  if v!=old.get(k):changes.append({'group_seq':n,'field':k,'converter_value':old.get(k),'source_value':v})
source['source_metadata']={'model_id':MID,'paper_id':'MED-2022','unit_id':'LME_026','baseline_year':1995,'supplement_period_label':'1990s','dynamic_simulation_period':'1995-2016','serialization_number_is_ecobase_accession':False,'missing_landing_groups':[g['n'] for g in groups if str(g['n']) not in landings],'routing':'not reported','ba':'not reported','gs':'not reported','selected':False,'production_eligible':False}
source['source_fisheries']={'fleets':fleets,'landings':landings,'discards':discards,'units':'t/km2 per baseline annual time step; workbook caption shortens to t/km2'}
save(MD/'model.json',source);save(converted,source);save(REVIEW/'converter_source_restoration.json',changes)
# Build round-trip workbook from the source-restored canonical database JSON.
p=subprocess.run([sys.executable,str(SKILL/'database_json.py'),'-j',str(converted)],capture_output=True,text=True,encoding='utf8',errors='replace');(REVIEW/'roundtrip.log').write_text(p.stdout+p.stderr,encoding='utf8');assert p.returncode==0
save(REVIEW/'source_manifest.json',[{'path':f.relative_to(ROOT).as_posix(),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sorted(SRC.iterdir()) if f.is_file()])
print('Source cells',len(evidence),'groups',len(groups),'model',MD)
