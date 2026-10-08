"""Retain native LME047 evidence and bounded applicability to released reviews."""
from pathlib import Path
import hashlib
import json
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
BASE = Path(__file__).resolve().parent


def identity(path):
    return {"path": path.relative_to(ROOT).as_posix(),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def main():
    source = ROOT / "regions/LME_047/papers/ECS-2022/DataSheet_1_EstimatingtheImpactofaSeasonal-7491f315.docx"
    document = Document(source)
    rows = [[c.text.strip() for c in r.cells] for r in document.tables[6].rows]
    caption = next(p.text for p in document.paragraphs if "Supplementary Table 4." in p.text)
    assert "absolute catch (6)" in caption and "M1997" in caption
    row2018 = next(r for r in rows if r[0] == "2018")
    catches, biomasses = {}, {}
    for i, kind in enumerate(rows[2][1:], 1):
        if kind in {"1", "6"}:
            group = int(rows[1][i])
            (biomasses if kind == "1" else catches)[group] = float(row2018[i])
    expected = {8:.22, 9:.86, 10:.28, 11:.16, 13:.39, 14:.62,
                16:.24, 17:.39, 18:.31, 19:.62, 20:0., 21:.13}
    assert catches == expected
    assert biomasses[20] == .01
    membership = [[c.text.strip() for c in r.cells] for r in document.tables[3].rows]
    assert "Planktivores/Piscivores" in membership[16][0]
    assert "Planktivores/Benthivores" in membership[17][0]
    selected = ROOT / "regions/LME_047/validation_reports/47_2_East_China_Sea_(2018)/source_parameter_cell_review.json"
    cells = json.loads(selected.read_text(encoding="utf-8"))["cells"]
    b = {int(c["seq"]): c for c in cells if c["field"] == "biomass"}
    assert b[20]["published"] == .0339
    assert "benthivores" in b[16]["group"].lower()
    assert "piscivores" in b[17]["group"].lower()
    dispositions = {
        "EEZ_598": "Final31-group source Table5 catches already reviewed; added-group catches unavailable rather than measured zero. Complete biomass retained where catch is incomplete. No additional applicable absolute-catch vector demonstrated in retained source inventory.",
        "EEZ_941": "Same independently reviewed final31-group source; final-stage/forage catch omissions and initial/final distinction documented. No replacement catch vector demonstrated.",
        "HS_071": "Same source reviewed locally; Table5 initial-active rows do not establish catches for all final compartments. Incomplete vectors already use complete biomass.",
        "HS_077": "Final ETP7 table catches and source stage pairs already reviewed; initial/intermediate versions distinguished. No overlooked compatible vector demonstrated.",
        "LME_003": "Complete author Yield fields and source-zero assumptions already reviewed; current40-bony candidate release uses them. No missing static-catch condition requiring a new proxy.",
        "LME_013": "Main static source and1980–2020 dynamic distinction reviewed. Referenced S1–S7 supplements remain unavailable; this is an explicit evidence gap, not evidence of absent series.",
        "LME_014": "Complete native72fleet/group catches and discards already verified; source Sequence/GroupID join is explicit. No new proxy justified.",
        "LME_022": "Main source, four-table supplement and methods already inventoried. Different regional/native variants distinguished; no demonstrated overlooked compatible catch vector.",
        "LME_024": "Supplement B3 supplies complete54-group1985 landings/discards; all113 multi-group allocations use complete vectors. Dynamic snapshots are distinguished; no missing vector to replace.",
        "LME_026": "Native supplement Catches/Discards sheets and DOCX time-series discussion inspected again.1990s gear/area catches already support current allocation evidence; no additional explicit catch time-series table identified. Missing candidates retain documented biomass fallback.",
        "LME_027": "Accepted EcoBase catch literals already supply complete proxies; final2009 native article/catch fidelity unavailable. Context2008 material is not silently substituted.",
        "LME_028": "Selected44-group1998 source and revised35-group later report distinguished; no separate selected-source supplement identified. Current catch and actual W5 donor provenance reviewed.",
        "LME_029": "Source S3 biological catch vector complete including printed zeros;1978 baseline and fitted1978–2015 series differentiated. No new vector needed.",
        "LME_032": "Published basic/catch/diet reconciliation already distinguishes blanks, explicit zeros and conventions. Seven W5 donor routes checked. No overlooked compatible series demonstrated.",
        "LME_034": "Source1978 guild catch vectors already used for current complete candidates; fixed year/basis transfer explicit. No overlooked vector demonstrated.",
        "LME_035": "Accepted complete export literals used; focal publication lacks full numeric catch/parameter tables and no applicable supplement located. Broader source fidelity gap remains.",
        "LME_036": "Selected2000s source reconstruction and all-six-fleet catch totals already retained; current complete-bony allocations use them. Conflicting stage cutoff remains explicit.",
        "LME_037": "1997 baseline uses2023 TableS3 catch and2026 group composition crosswalk;2018 endpoint is a separate unselected/NOT_RUN identity. No endpoint catch substitution justified."
    }
    progress = json.loads((BASE / "verification/review_package_progress.json").read_text(encoding="utf-8"))
    reviews = []
    for row in progress["regions"]:
        if row["unit_id"] not in dispositions:
            continue
        folder = ROOT / Path(row["handoff"]["path"]).parent
        source_index = folder / "source_review.md"
        assert source_index.exists()
        reviews.append({"unit_id": row["unit_id"], "source_inventory": identity(source_index),
                        "workbook": identity(ROOT / row["workbook"]["path"]),
                        "finding": dispositions[row["unit_id"]],
                        "decision": "No new mapping/parameter change supported by this bounded applicability check."})
    assert len(reviews) == 18
    result = {
        "status": "PASS native-table extraction and18-package bounded source-inventory applicability",
        "source": identity(source), "main_pdf": identity(ROOT / "regions/LME_047/papers/ECS-2022/pdf-e60d0358.pdf"),
        "source_caption": caption,
        "locator": "S4 seventh native DOCX table (python index6); S1 fourth (index3). Main PDF4 sections2.2/2.3 and PDF5 Table1 independently read.",
        "source_2018_catch_t_per_km2": catches,
        "source_2018_biomass_t_per_km2": biomasses,
        "important_exception": "S4group20 B0.01 differs from M2018 Table1 B0.0339; other11 match rounded M2018 biomass. S4printed catch0.00 is not demonstrated exact zero.",
        "join_check": "S1 row-order16=P/P and17=P/B, but numbered main/diet/canonical16=P/B and17=P/P. Join by verified group names/IDs.",
        "parameter_review_binding": identity(selected),
        "interpretation": "2018 S4 absolute catches are source-group composition proxies with common CFSY lineage, not recovered native M2018 catch inputs or observed regional catch composition. Main methods disclose no discard data. Candidate completeness and local uncertainty require separate lead/QA review; this lesson alone adopts no individual allocation.",
        "rule": "Inspect relevant time-series supplements before rejecting source catch; verify own caption/type/year/units/basis and name/ID crosswalk, completeness and rounded/censored zeros. Keep proxy/source-export identities separate and accepted scientific inputs unchanged.",
        "completed_applicability_scope": "Read current18 released source inventories and relevant allocation locators; re-inspected MED supplement sheets and time-series prose. Bounded review, not exhaustive new source discovery or proof that inaccessible material has no catch evidence.",
        "completed_reviews": reviews,
        "active_and_pending": ["LME_038", "LME_047", "LME_049", "LME_050", "LME_052"]
    }
    output = BASE / "verification/supplementary_catch_timeseries_lesson.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": output.relative_to(ROOT).as_posix(), "status": result["status"]}))


if __name__ == "__main__":
    main()
