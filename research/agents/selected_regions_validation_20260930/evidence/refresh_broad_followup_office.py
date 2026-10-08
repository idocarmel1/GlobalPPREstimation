"""Small OOXML edits to existing regional artifacts; preserve researcher rows."""
import argparse
import copy
import json
import os
import re
import shutil
import sys
import urllib.parse
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree as E

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
import workbooks as W
from bounded_workbook_update import _cell_value, _move_row, q, REL

PR = 'http://schemas.openxmlformats.org/package/2006/relationships'
MANUAL = ['SPPR calculation', 'Open issues and next action', 'Review and reproducibility', 'Selection rationale']


def readj(p):
    return json.loads(p.read_text(encoding='utf-8'))


def writej(p, value):
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def load_parts(p):
    with ZipFile(p) as z:
        infos = z.infolist()
        return infos, {i.filename: z.read(i.filename) for i in infos}


def save_parts(p, infos, parts):
    temp = p.with_suffix('.office.tmp')
    with ZipFile(temp, 'w') as z:
        for info in infos:
            z.writestr(info, parts[info.filename])
        for name in set(parts) - {i.filename for i in infos}:
            z.writestr(name, parts[name])
    os.replace(temp, p)


def paths(parts):
    rels = E.fromstring(parts['xl/_rels/workbook.xml.rels'])
    targets = {r.get('Id'): r.get('Target') for r in rels}
    wb = E.fromstring(parts['xl/workbook.xml'])
    return {s.get('name'): 'xl/' + targets[s.get('{' + REL + '}id')].lstrip('/').removeprefix('xl/') for s in wb.find(q('sheets'))}


def set_cell(root, address, value):
    row_number = int(re.sub('[A-Z]', '', address))
    data = root.find(q('sheetData'))
    row = data.find(q('row') + f"[@r='{row_number}']")
    assert row is not None
    cell = row.find(q('c') + f"[@r='{address}']")
    if cell is None:
        cell = E.SubElement(row, q('c'), r=address)
    _cell_value(cell, value)


def make_row(template, number, values):
    row = copy.deepcopy(template)
    old = {re.sub(r'\d+$', '', c.get('r')): c for c in row}
    for child in list(row):
        row.remove(child)
    for i, value in enumerate(values, 1):
        col = openpyxl.utils.get_column_letter(i)
        cell = copy.deepcopy(old[col]) if col in old else E.Element(q('c'))
        cell.set('r', col + str(number))
        _cell_value(cell, value)
        row.append(cell)
    _move_row(row, number)
    return row


def manual_xml(doc):
    return {row.cells[0].text: E.tostring(row._tr, method='c14n', exclusive=True)
            for table in doc.tables for row in table.rows if row.cells[0].text in MANUAL}


def add_link(par, label, target):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), par.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    prop = OxmlElement('w:rPr')
    col = OxmlElement('w:color'); col.set(qn('w:val'), '0563C1'); prop.append(col)
    under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); prop.append(under)
    r.append(prop)
    t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); par._p.append(h)


def run(unit):
    folder = ROOT / 'regions' / unit
    doc = next(folder.glob('Model_validation_*.docx'))
    app = next(folder.glob('*taxon_mapping_appendix.xlsx'))
    mid = doc.stem.removeprefix('Model_validation_')
    ev = folder / 'validation_reports' / mid
    receipt_path = ev / 'broad_candidate_followup.json'
    receipt = readj(receipt_path)
    assert receipt['after_workbook_sha256'] == W.sha(folder / (unit + '.xlsx'))
    audit = readj(ev / ('taxon_audit.json' if unit == 'LME_003' else 'adopted_taxon_audit.json'))
    by = {r['taxon']: r for r in audit['rows' if unit == 'LME_003' else 'records']}
    targets = [c['taxon'] for c in receipt['changes']]
    if unit == 'LME_003':
        targets.append('Marine pelagic fishes not identified')
    scratch = HERE / 'work/coordinator_broad_followup' / unit
    for source in [app, doc]:
        destination = scratch / source.name
        if destination.exists():
            assert W.sha(destination) == W.sha(source), 'Keep original Office baseline; current artifact has already changed'
        else:
            shutil.copy2(source, destination)
    initial = {p.name: W.sha(p) for p in [app, doc]}
    wb = openpyxl.load_workbook(app)
    infos, parts = load_parts(app)
    sheet_paths = paths(parts)
    roots = {name: E.fromstring(parts[path]) for name, path in sheet_paths.items()}
    mapping = wb['Taxon mapping']
    changed_rows = []
    for row in mapping:
        taxon = row[0].value
        if taxon not in targets:
            continue
        a = by[taxon]
        display = a['mapping_display'] if unit == 'LME_003' else '; '.join(c['group'] + f" ({c['weight'] * 100:.4f}%)" for c in a['proposed_groups'])
        reason = a['reason'].replace('Allocation W4 Medium', 'Allocation confidence Medium').replace('as W4 Medium pooled catch weights', 'as pooled catch weights with Medium allocation confidence')
        reason += ' Source definitions and complete candidates: Sources sheet, Complete broad fish candidates.'
        set_cell(roots['Taxon mapping'], 'E' + str(row[0].row), display)
        set_cell(roots['Taxon mapping'], 'G' + str(row[0].row), reason)
        node = roots['Taxon mapping'].find(q('sheetData')).find(q('row') + f"[@r='{row[0].row}']")
        node.set('ht', '360' if unit == 'LME_003' and taxon != 'Marine pelagic fishes not identified' else '285')
        node.set('customHeight', '1')
        changed_rows.append({'taxon': taxon, 'row': row[0].row, 'display': display, 'reason': reason})
    assert len(changed_rows) == len(targets)
    # Add a descriptive, portable source link while retaining every existing one.
    source_sheet = wb['Sources']
    last = max(c.row for row in source_sheet for c in row if c.value is not None)
    root = roots['Sources']; data = root.find(q('sheetData'))
    template = data.find(q('row') + f"[@r='{last}']")
    target = 'validation_reports/' + mid + '/broad_candidate_followup.json'
    description = ('Source definitions, complete eligible candidates, exact source catch denominators and genuine zeros; '
                   'assumed fixed composition across all retained years and catch bases. Scientific inputs are preserved.')
    values = ['Complete broad fish candidates', description, 'Open reviewed evidence', '30 September 2026'] if unit == 'LME_003' else ['BROAD_FISH_COMPLETE', 'Open reviewed evidence', description, '30 September 2026']
    row = make_row(template, last + 1, values)
    row.set('ht', '95'); row.set('customHeight', '1'); data.append(row)
    link_col = 'C' if unit == 'LME_003' else 'B'
    link_style = next(c.style_id for row in source_sheet for c in row if getattr(c, 'column_letter', '') == link_col and c.hyperlink)
    row.find(q('c') + f"[@r='{link_col}{last + 1}']").set('s', str(link_style))
    member = sheet_paths['Sources']
    relpath = str(Path(member).parent / '_rels' / (Path(member).name + '.rels')).replace('\\', '/')
    rels = E.fromstring(parts[relpath])
    rid = 'rIdCompleteBroadCandidates'
    assert not any(r.get('Id') == rid for r in rels)
    E.SubElement(rels, '{' + PR + '}Relationship', Id=rid, Type=REL + '/hyperlink', Target=urllib.parse.quote(target, safe='/:?=&%#'), TargetMode='External')
    E.SubElement(root.find(q('hyperlinks')), q('hyperlink'), ref=link_col + str(last + 1), attrib={'{' + REL + '}id': rid})
    parts[relpath] = E.tostring(rels, encoding='utf-8', xml_declaration=True)
    if root.find(q('dimension')) is not None:
        root.find(q('dimension')).set('ref', 'A1:D' + str(last + 1))
    # The Bay appendix has a native, value-backed complete-candidate table.
    if unit == 'LME_034':
        sheet = wb['Allocation evidence']
        root = roots['Allocation evidence']; data = root.find(q('sheetData'))
        assert not root.findall('.//' + q('f'))
        assert root.find(q('hyperlinks')) is None
        old_rows = {int(r.get('r')): r for r in data}
        out, done = [], set()
        for row in list(sheet.values)[6:]:
            taxon = row[0]
            if taxon not in targets:
                out.append(list(row)); continue
            if taxon in done:
                continue
            done.add(taxon); a = by[taxon]
            for c in a['proposed_groups']:
                out.append([taxon, c['seq'], c['group'], c['source_printed_catch_density'], c['native_model_catch_density'],
                            c['loaded_model_catch_density'], a['allocation_total'], c['weight'], a['membership_confidence'],
                            a['allocation_confidence'], a['allocation_formula'], a['candidate_inclusion_rationale'],
                            a['allocation_period_transfer'], 'Guénette2013 A2.1p46; total1978 landings+discards; t wet/km²/year; common whole-study-area denominator.',
                            'Excluded IDs: ' + ', '.join(str(c['seq']) for c in a['candidate_exclusions']) + '. Actual source and geographic reasons in Complete broad fish candidates (Sources).'])
        assert done == set(targets)
        for row in list(data):
            if int(row.get('r')) >= 7:
                data.remove(row)
        for i, values in enumerate(out, 7):
            data.append(make_row(old_rows[min(i, max(old_rows))], i, values))
        allocation_last = len(out) + 6
        if root.find(q('dimension')) is not None:
            root.find(q('dimension')).set('ref', 'A1:O' + str(allocation_last))
    for name, root in roots.items():
        if name in ['Taxon mapping', 'Sources'] or (unit == 'LME_034' and name == 'Allocation evidence'):
            parts[sheet_paths[name]] = E.tostring(root, encoding='utf-8', xml_declaration=True)
    for name in list(parts):
        if name.startswith('xl/tables/') and name.endswith('.xml'):
            root = E.fromstring(parts[name]); end = None
            if root.get('name') == 'EvidenceSources':
                end = 'D' + str(last + 1)
            elif root.get('name') == 'CandidateAllocation' and unit == 'LME_034':
                end = 'O' + str(allocation_last)
            if end:
                root.set('ref', root.get('ref').split(':')[0] + ':' + end)
                if root.find(q('autoFilter')) is not None:
                    root.find(q('autoFilter')).set('ref', root.get('ref'))
                parts[name] = E.tostring(root, encoding='utf-8', xml_declaration=True)
    wb.close()
    save_parts(app, infos, parts)
    # Modify only the automatic mapping summary, keeping all manual rows exact.
    d = Document(doc); protected = manual_xml(d)
    par = next(p for p in d.paragraphs if p.text.startswith('Unresolved taxa: none.') or p.text.startswith('No unresolved mappings remain.'))
    note = (' The four broad bony-fish labels now include all 40 eligible source groups, including named pelagics and genuine zero-catch stages. '
            'Juvenile sablefish/lingcod31 remains pelagic; adult Sablefish45 does not. Complete author Yield supplies assumed catch proportions.' if unit == 'LME_003' else
            ' Broad marine fish now includes all 23 eligible bony-fish pools; the pelagic subset has 21. Named tuna/billfish categories do not prove absence from unidentified catches. '
            'Source-zero small bathypelagics remain eligible; Maldives shelf pools retain the documented geographic exclusion.')
    par.add_run(note + ' ')
    add_link(par, 'Complete candidate evidence', target)
    stage = scratch / 'report_stage.docx'; d.save(stage)
    assert manual_xml(Document(stage)) == protected
    infos, parts = load_parts(doc)
    with ZipFile(stage) as z:
        for name in ['word/document.xml', 'word/_rels/document.xml.rels']:
            parts[name] = z.read(name)
    save_parts(doc, infos, parts)
    assert manual_xml(Document(doc)) == protected
    check = openpyxl.load_workbook(app, data_only=False)
    baseline = openpyxl.load_workbook(scratch / app.name, data_only=False)
    for sn in baseline.sheetnames:
        if sn == 'Allocation evidence' and unit == 'LME_034':
            continue
        for row in baseline[sn]:
            for cell in row:
                allowed = sn == 'Taxon mapping' and cell.row in [x['row'] for x in changed_rows] and cell.column in [5, 7]
                if not allowed:
                    assert check[sn][cell.coordinate].value == cell.value, (sn, cell.coordinate)
                if cell.hyperlink:
                    current = check[sn][cell.coordinate].hyperlink
                    assert current and current.target == cell.hyperlink.target and current.location == cell.hyperlink.location
    check.close(); baseline.close()
    receipt['office_refresh'] = {'before_hashes': initial, 'after_hashes': {p.name: W.sha(p) for p in [app, doc]},
                                 'mapping_rows': changed_rows, 'source_row': last + 1, 'manual_rows_and_selection_exact': True,
                                 'existing_links_preserved': True, 'render_and_portability_pending': True}
    receipt['status'] = 'Mapping, arithmetic and Office values updated; independent QA/render/handoff pending'
    writej(receipt_path, receipt)
    print(json.dumps({'unit': unit, 'changed_rows': [x['row'] for x in changed_rows], 'manual_exact': True}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('unit', choices=['LME_003', 'LME_034'])
    run(parser.parse_args().unit)
