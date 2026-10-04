"""Change reviewed authoritative block values without rewriting protected XLSX parts.

Only ordinary, formula-free regional @table sheets are supported. Other package
parts retain their exact bytes; changed sheets retain existing row/cell styles.
Numerical values use Python round-trip serialization, including accepted precision.
"""
import copy
import os
import re
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as E
import openpyxl
import workbooks as W

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
q = lambda s: '{' + NS + '}' + s


def _cell_value(cell, value):
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop('t', None)
    value = W.clean(value)
    if value is None:
        return
    if isinstance(value, bool):
        cell.set('t', 'b')
        E.SubElement(cell, q('v')).text = '1' if value else '0'
    elif isinstance(value, (int, float)):
        E.SubElement(cell, q('v')).text = repr(value)
    else:
        assert len(str(value)) <= 32767
        cell.set('t', 'inlineStr')
        text = E.SubElement(E.SubElement(cell, q('is')), q('t'))
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        text.text = str(value)


def _move_row(row, number):
    row.set('r', str(number))
    for cell in row:
        if cell.tag == q('c'):
            cell.set('r', re.sub(r'\d+$', str(number), cell.get('r')))


def update_blocks(path, before, after, *, expected_sha256):
    path = Path(path)
    assert W.sha(path) == expected_sha256, 'Workbook changed before bounded update'
    assert before.keys() == after.keys(), 'Sheet additions are not supported'
    observed = W.read_book(path)
    assert observed == before, 'Expected-old tables differ'
    with ZipFile(path) as z:
        infos = z.infolist()
        parts = {i.filename: z.read(i.filename) for i in infos}
    rels = E.fromstring(parts['xl/_rels/workbook.xml.rels'])
    targets = {n.get('Id'): n.get('Target') for n in rels}
    wb = E.fromstring(parts['xl/workbook.xml'])
    paths = {s.get('name'): 'xl/' + targets[s.get('{' + REL + '}id')].lstrip('/').removeprefix('xl/')
             for s in wb.find(q('sheets'))}
    reader = openpyxl.load_workbook(path, read_only=True, data_only=False)
    changed_parts = []
    try:
        for sn, old_blocks in before.items():
            new_blocks = after[sn]
            assert list(old_blocks) == list(new_blocks), 'Block additions/reordering are not supported'
            if old_blocks == new_blocks:
                continue
            root = E.fromstring(parts[paths[sn]])
            assert not root.findall('.//' + q('f')), 'Formula-bearing sheet is unsupported'
            for feature in ['tableParts', 'hyperlinks', 'mergeCells', 'autoFilter',
                            'conditionalFormatting', 'dataValidations', 'drawing', 'legacyDrawing']:
                assert root.find(q(feature)) is None, (sn, feature, 'coordinate-bearing feature unsupported')
            data = root.find(q('sheetData'))
            raw_rows = {int(r.get('r')): r for r in data}
            positions = {}
            active = None
            for number, row in enumerate(reader[sn].values, 1):
                vals = list(row)
                if vals and vals[0] == '@table':
                    active = vals[1]
                    positions[active] = {'marker': number, 'header': None, 'rows': []}
                elif active and any(v is not None for v in vals):
                    if positions[active]['header'] is None:
                        positions[active]['header'] = number
                    else:
                        positions[active]['rows'].append(number)
            replacement = {}
            for tn, (header, rows) in old_blocks.items():
                new_header, new_rows = new_blocks[tn]
                assert header == new_header, (sn, tn, 'header change unsupported')
                assert len(rows) == len(positions[tn]['rows'])
                if rows == new_rows:
                    continue
                loc = positions[tn]
                start = loc['header'] + 1
                stop = max(loc['rows'], default=loc['header']) + 1
                assert loc['rows'] == list(range(start, stop)), 'Sparse block unsupported'
                templates = [raw_rows[i] for i in loc['rows']]
                made = []
                for i, values in enumerate(new_rows):
                    row = copy.deepcopy(templates[min(i, len(templates) - 1)]) if templates else E.Element(q('row'))
                    prior_cells = {re.sub(r'\d+$', '', c.get('r')): c for c in row if c.tag == q('c')}
                    for child in list(row):
                        row.remove(child)
                    for ci, value in enumerate(values, 1):
                        col = openpyxl.utils.get_column_letter(ci)
                        cell = copy.deepcopy(prior_cells[col]) if col in prior_cells else E.Element(q('c'))
                        cell.set('r', col + str(start + i))
                        _cell_value(cell, value)
                        if value is not None or cell.get('s'):
                            row.append(cell)
                    made.append(row)
                replacement[start] = (stop, made)
            new_data = []
            maximum = max(raw_rows, default=0)
            old_number = 1
            shift = 0
            while old_number <= maximum:
                if old_number in replacement:
                    stop, made = replacement[old_number]
                    for i, row in enumerate(made):
                        _move_row(row, old_number + shift + i)
                        new_data.append(row)
                    shift += len(made) - (stop - old_number)
                    old_number = stop
                else:
                    if old_number in raw_rows:
                        row = copy.deepcopy(raw_rows[old_number])
                        _move_row(row, old_number + shift)
                        new_data.append(row)
                    old_number += 1
            for row in list(data):
                data.remove(row)
            data.extend(new_data)
            dimension = root.find(q('dimension'))
            if dimension is not None:
                max_col = max(len(h) for h, _ in new_blocks.values())
                dimension.set('ref', 'A1:' + openpyxl.utils.get_column_letter(max_col) + str(maximum + shift))
            parts[paths[sn]] = E.tostring(root, encoding='utf-8', xml_declaration=True)
            changed_parts.append(paths[sn])
    finally:
        reader.close()
    temp = path.with_suffix('.bounded.tmp.xlsx')
    try:
        with ZipFile(temp, 'w') as z:
            for info in infos:
                z.writestr(info, parts[info.filename])
        assert W.read_book(temp) == after, 'Temporary workbook failed exact table verification'
        assert W.sha(path) == expected_sha256, 'Workbook changed during bounded update'
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()
    return changed_parts
