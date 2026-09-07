from pathlib import Path
import csv,json,math,hashlib,shutil
import openpyxl
ROOT=next(p for p in Path(__file__).resolve().parents if (p/'tools/run_sppr.py').is_file() and (p/'PPRAtlas/archive').is_dir());BASE=ROOT/'data/LME_022/validation/NS-2025_ECS-1990s'
MODEL=BASE/'extraction/20251990_East_Coast_of_Scotland_1991-1995';STEM='22_20251990_East_Coast_of_Scotland_(1991-1995)'
SRC=ROOT/'PPRAtlas/archive/regions/LME_022/NS-2025';source=openpyxl.load_workbook(next(SRC.glob('*.xlsx')),data_only=False)
raw=json.loads((MODEL/'model.json').read_text());db=json.loads((MODEL/f'{STEM}.json').read_text())['group']
numeric=lambda x:None if x in [None,'','-9999'] else float(x)
assert len(raw['groups'])==len(db)==25
required=['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv','TL.xlsx','Metadata.xlsx']
assert all((MODEL/x).exists() for x in required)
assert all(b'\r\n' in (MODEL/x).read_bytes() for x in required if x.endswith('.csv'))
assert all('"' not in (MODEL/x).read_text() for x in required if x.endswith('.csv'))
checks=0
for n in range(1,26):
    for c in range(3,26):
        v=source['Table S3'].cell(n+2,c).value;pred=str(c-1)
        assert numeric(raw['diet'][pred].get(str(n)))==numeric(v),(n,c,v)
        checks+=1
for n in range(1,26):
    for f,c in zip(raw['fleets'],[7,8,9]):
        assert numeric(raw['landings'][str(n)].get(f))==numeric(source['Table S1'].cell(n+3,c).value)
        checks+=1
tax=list(csv.DictReader((MODEL/'taxonomy.csv').open(encoding='utf-8')))
for g,t in zip(db,tax):
    assert g['group_name']==t['group_name'] and g['taxon_descr']==t['taxon_descr']
    assert g['biomass_accum']=='-9999' and g['gs']=='-9999'
normalized=[]
for g in db:
    n=str(g['group_seq'])
    if n not in raw['diet']:continue
    src=raw['diet'][n];total=sum(src.values());dd=g['diet_descr']['diet'];dd=dd if isinstance(dd,list) else [dd]
    for v in dd:
        exp=src.get(str(v['prey_seq']),0)/total
        assert math.isclose(float(v['proportion']),exp,abs_tol=1e-12)
    if not math.isclose(total,1,abs_tol=1e-6):normalized.append({'group':g['group_name'],'source_sum':total,'converter_factor':1/total})
sppr=openpyxl.load_workbook(BASE/'sppr'/f'{STEM}.xlsx',data_only=True)
gh=list(sppr['groups_df'].values);gd={str(r[1]):dict(zip(gh[0],r)) for r in gh[1:]}
for g in db:assert gd[g['group_name']]['taxon_descr']==g['taxon_descr']
hh=list(sppr['model_health'].values);health=[dict(zip(hh[0],r)) for r in hh[1:]]
transforms=[]
for s in raw['groups']:
    g=gd[s['name']]
    transforms.append({'group':s['name'],'source_BA':None,'loaded_BA':g['biomass_accum'],'source_GS':None,'loaded_GS':g['gs'],'loaded_flow_to_detritus':g['flow_to_det'],'source_EE':numeric(s.get('ee')),'loaded_EE':g['ee']})
(BASE/'LOADER_TRANSFORMATIONS.csv').write_text('',encoding='utf-8')
with (BASE/'LOADER_TRANSFORMATIONS.csv').open('w',encoding='utf-8',newline='') as f:
    w=csv.DictWriter(f,list(transforms[0]));w.writeheader();w.writerows(transforms)
rounding=[]
for n in [9,11]:
    victim=db[n-1];pred=0
    for p in db:
        if numeric(p['qb']) is None:continue
        diet=p.get('diet_descr',{});diet=diet.get('diet',[]) if diet else [];diet=diet if isinstance(diet,list) else [diet]
        pred+=sum(float(v['proportion'])*float(p['biomass'])*float(p['qb']) for v in diet if str(v['prey_seq'])==str(n))
    implied=(pred+float(victim['export']))/(float(victim['pb'])*float(victim['ee']))
    observed=float(victim['biomass']);assert observed-.005<=implied<observed+.005
    rounding.append({'group':victim['group_name'],'printed_B':observed,'B_that_would_reconcile_no_BA':implied,'within_printed_rounding_interval':True,'applied_to_extraction':False})
book=BASE/'evaluation/data/LME_022/models'/f'{STEM}.xlsx'
ppr=openpyxl.load_workbook(book,data_only=True);pr=list(ppr['PPR by method'].values);header=next(r for r in pr if 2019 in r);ycol=header.index(2019)
method_values={r[0]:{'status':r[1],'PPR_2019_t_wet':r[ycol]} for r in pr if r[0] in ['new_GE','new_WithEgestion','new_TE_EEfix','SPPR_1995_TE0.1','simple trophic chain, per taxon']}
totals=json.loads((BASE/'mapping_summary.json').read_text())
result={'source_cells_crosschecked':checks,'original_groups':25,'algorithm_added_import_groups':1,'all_taxonomy_propagated':True,'unknown_source_BA_and_GS_preserved_in_database':True,'eight_import_files_present':True,'converter_normalized_diet_columns':normalized,'rounding_diagnostics_not_corrections':rounding,'health':[{k:r[k] for k in ['TE_option','status','model_input_is_model_balanced','divergence_rho_living','divergence_n_negative_sources']} for r in health],'mapping':totals,'selected_2019_PPR':method_values,'source_files':[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in SRC.iterdir() if p.suffix in ['.pdf','.xlsx','.docx']]}
(BASE/'VALIDATION_RESULT.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
