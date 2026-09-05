import openpyxl,sys,pathlib,json
p=pathlib.Path(sys.argv[1])/'extracted/work'
for f in p.glob('Table*.xlsx'):
 w=openpyxl.load_workbook(f,data_only=False);out=[]
 for s in w:
  print(f.name,s.title,s.sheet_state,s.max_row,s.max_column)
  rows=[]
  for r in s:
   cells=[{'cell':c.coordinate,'value':c.value,'format':c.number_format,'bold':c.font.bold,'italic':c.font.italic,'type':c.data_type} for c in r if c.value is not None];rows.append(cells)
   print(' | '.join(str(c['cell'])+'='+str(c['value']) for c in cells))
  out.append({'sheet':s.title,'state':s.sheet_state,'rows':rows})
 (p/(f.stem+'_cells.json')).write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
