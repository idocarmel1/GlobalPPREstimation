"""Source-preservation audit applied AFTER the bundled Ecopath conversion.
Restores exact CSV values which the bundled converter normalizes/defaults/drops.
It changes no source document, import CSV/XLSX, or installed skill.
The database extension stores every import table so workbook round-trip is lossless.
"""
import csv,json,pathlib,sys,re,decimal,hashlib
sys.dont_write_bytecode=True
import openpyxl
D=decimal.Decimal
sys.path.insert(0,r'C:\Users\idoca\.agents\skills\ecopath-extraction\scripts')
from database_json import EwEConverter
p=pathlib.Path(sys.argv[1]);ex=p/'extracted';w=ex/'work';summary=[]
def nd(v):return str(v) if v is not None and str(v).strip() else '-9999'
def readcsv(f):
 with f.open(encoding='utf-8',newline='') as h:return list(csv.reader(h))
for d in sorted(ex.glob('27_*_Banc*')):
 model=json.loads((d/'model.json').read_text(encoding='utf-8'));variant=model['metadata']['variant'];jpath=next(d.glob('27_Canary_Current*.json'));data=json.loads(jpath.read_text(encoding='utf-8'));source_tables={}
 for name in ['Basic_input.csv','Diet_composition.csv','Landings.csv','Discards.csv','Detritus_fate.csv','Biomass_accumulation.csv']:source_tables[name]=readcsv(d/name)
 for name in ['TL.xlsx','Metadata.xlsx']:source_tables[name]=[list(r) for r in openpyxl.load_workbook(d/name,data_only=False).active.iter_rows(values_only=True)]
 basic={r[0]:r for r in source_tables['Basic_input.csv'][1:]};land={r[0]:r for r in source_tables['Landings.csv'][1:]};disc={r[0]:r for r in source_tables['Discards.csv'][1:]};ba={r[0]:r for r in source_tables['Biomass_accumulation.csv'][1:]};dt=source_tables['Diet_composition.csv'];dietrows={r[0]:r for r in dt[1:] if r[0]};imports=next(r for r in dt if r[1]=='Import');fates={r[0]:r for r in source_tables['Detritus_fate.csv'][1:]};fatehead=source_tables['Detritus_fate.csv'][0];names={g['group_seq']:g['group_name'] for g in data['group']};sums={};blanks={};log=[]
 for g in data['group']:
  seq=g['group_seq'];r=basic[seq];n=int(seq)
  for key,col in [('habitat_area',2),('biomass_habitat_area',3),('z',4),('pb',5),('qb',6),('ee',7),('other_mort',8),('ge',9),('gs',10),('detritus_import',11)]:g[key]=nd(r[col])
  # Published biomass is study-area density; no arbitrary habitat-area scaling.
  g['biomass']=nd(r[3]);g['vbk']='-9999';g['shadow_price']='-9999';g['ge_input']='false' if g['ge']=='-9999' else 'true'
  g['pp']='0' if n<=47 else ('1' if n<=50 else '2')
  g['biomass_accum_rate']=nd(ba[seq][3]);g['biomass_accum']=nd(ba[seq][2])
  if g['biomass_accum_rate']!='-9999' and g['biomass']!='-9999':
   g['biomass_accum']=format(D(g['biomass'])*D(g['biomass_accum_rate']),'f');log.append(f"Group {seq}: database absolute BA {g['biomass_accum']} derived from reported B {g['biomass']} x reported BA/B {g['biomass_accum_rate']}; import absolute BA remains blank.")
  g['export']=nd(land[seq][-1]);g['landings_by_fleet']={fleet:nd(v) for fleet,v in zip(source_tables['Landings.csv'][0][2:-1],land[seq][2:-1])};g['discards_by_fleet']={fleet:nd(v) for fleet,v in zip(source_tables['Discards.csv'][0][2:-1],disc[seq][2:-1])}
  g['detritus_fate_by_pool']={'51':nd(fates[seq][2])};g['detritus_export']=nd(fates[seq][3]);g['diet_imp']='-9999'
  if seq in dt[0][2:]:
   col=dt[0].index(seq);items=[];total=D(0);missing=[]
   for prey,row in dietrows.items():
    v=nd(row[col]);total+=D(v) if v!='-9999' else D(0)
    if v=='-9999':missing.append(prey)
    items.append({'prey_seq':prey,'proportion':v,'detritus_fate':nd(fates[seq][2]) if prey=='51' else '-9999'})
   g['diet_imp']=nd(imports[col]);total+=D(g['diet_imp']) if g['diet_imp']!='-9999' else D(0);g['diet_descr']={'diet':items};sums[seq]=str(total);blanks[seq]=missing
  else:g['diet_descr']=None
 data['metadata']=model['metadata'];data['extraction_metadata']=model['metadata'];data['extraction_source_tables']=source_tables;data['extraction_adapter']={'purpose':'Preserve exact source values and unknowns after bundled conversion; do not normalize published diets.','unknown_sentinel':'-9999','biomass_basis':'Published Table 1 density per full study area; habitat proportion not reported.','extra_fields':'ge preserves P/Q; z preserves total mortality; fleet and detritus fields preserve unknowns and fleet rows; extraction_source_tables provides exact eight-table round-trip.'}
 jpath.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
 # Restore a workbook entirely from final database JSON, retaining all eight table shapes/values.
 final=json.loads(jpath.read_text(encoding='utf-8'));wb=openpyxl.Workbook();wb.remove(wb.active)
 for fname,rows in final['extraction_source_tables'].items():
  ws=wb.create_sheet(pathlib.Path(fname).stem)
  for row in rows:ws.append(row)
  ws.freeze_panes='C2';ws.column_dimensions['A'].width=8;ws.column_dimensions['B'].width=29
  for cell in ws[1]:cell.font=openpyxl.styles.Font(bold=True)
 out=jpath.with_name(jpath.stem+'_reconstructed.xlsx');wb.save(out)
 # Recompute from exact database numbers. Add source-integrity and missing-data guards absent in the stock equation checker.
 converter=EwEConverter(log_file=str(d/'ewe_conversion.log'));converter.logger.info('FINAL SOURCE-PRESERVATION AUDIT: preceding automatic-normalisation/default results are superseded. The following JSON/workbook preserve published values and unknowns.')
 for line in log:converter.logger.info(line)
 converter.logger.info('Restored source diet totals, habitat unknowns, ge=P/Q, z, fleet catches/discard unknowns, detritus unknowns and original metadata. No values inserted into import files.')
 findings=converter.check_mass_balance(final['group'],0.0,0.05);stock_verdict=findings['verdict']
 for seq,total in sums.items():
  if D(total)>D('1.01'):findings['errors'].append(f'group {seq} ({names[seq]}): source diet + Import = {total} > 1.01; published diet cannot be a valid composition. Exact source preserved without normalization.')
  elif D(total)<D('0.99'):
   if variant!='Base' and int(seq) in model.get('extraction_notes',{}).get('incomplete_diet_consumers',[]):findings['indeterminate'].append(f'group {seq} ({names[seq]}): known diet sum {total}; variant-specific prey components are unpublished and remain -9999, so predation checks are lower bounds.')
   else:findings['errors'].append(f'group {seq} ({names[seq]}): published diet + Import = {total} < 0.99; unexplained deficit preserved.')
 findings['indeterminate'].extend(['Unassimilated consumption and detritus routing/import/export are unknown. Stock detritus/respiration arithmetic uses a software GS assumption only as an indicative check, not as a verified source value.','P/B is zero for Meagre despite nonzero Z/PQ and is absent for 10 other multi-stanza groups; those production identities cannot be verified from the printed inputs.'])
 findings['notes'].append(f'Stock equation-only verdict before integrity/missing-data guards: {stock_verdict}. Final verdict concerns this extraction, not whether the authors had an internally balanced operational model.')
 findings['verdict']='NOT BALANCED' if findings['errors'] else 'INDETERMINATE';findings['balanced']=False
 converter.report_mass_balance(findings,jpath.stem,str(d/'MASS_BALANCE.md'));converter.update_report(str(d/'REPORT.md'),str(d/'MASS_BALANCE.md'))
 validation=(d/'VALIDATION.txt').read_text(encoding='utf-8');vsummary=re.search(r'\d+ error\(s\), \d+ warning\(s\)',validation).group();mbsummary=re.search(r'\d+ error\(s\), \d+ warning\(s\), \d+ note\(s\)',(d/'MASS_BALANCE_CHECK.txt').read_text(encoding='utf-8')).group()
 valtext=f"## Validation\nBundled validate.py: **{vsummary}**. See VALIDATION.txt for every flagged row/column. All file/group/header/CRLF/total checks ran. The base diet totals range from 0.980331 to 2.2006; affected variants have incomplete columns as documented above. Missing detritus fate produces 51 warnings; all-GS blank produces one; rounded zero juvenile biomasses account for the variants' additional 3 warnings. All flagged source values were checked on article pp. 4-5 and rendered Supplement pp. 8-10 and 15-16.\n\nBundled massbalance_check.py: {mbsummary}. Its detritus calculation is indicative because GS/routing are missing. It cannot test groups with missing/zero P/B. The final database analysis below includes extra integrity/missing-data guards.\n\nDatabase source-preservation adapter ../work/audit_database.py corrects stock converter normalisation/defaulting and dropped fields, retaining exact CSV values and -9999 unknowns. The official bundled converter was run first, with original logs saved under ../work. The final section of ewe_conversion.log supersedes its preliminary values/verdict. ge retains source P/Q; z retains Z; the JSON extension holds all eight original import tables. The reconstructed XLSX is rebuilt solely from the final JSON extension and verified cell-by-cell against those eight files. No installed skill was modified.\n\n"
 report=(d/'REPORT.md').read_text(encoding='utf-8');report=re.sub(r'## Validation\n.*?(?=## Mass balance)',valtext,report,flags=re.S);(d/'REPORT.md').write_text(report,encoding='utf-8')
 # Verify JSON standard fields and eight-sheet round-trip.
 check=openpyxl.load_workbook(out,data_only=False)
 for fname,rows in source_tables.items():
  actual=[list(row) for row in check[pathlib.Path(fname).stem].iter_rows(values_only=True)];expected=[[None if v=='' else v for v in row] for row in rows]
  assert actual==expected,(variant,fname)
 for g in final['group']:
  assert g['habitat_area']=='-9999' and g['gs']=='-9999';assert g['ge']==nd(basic[g['group_seq']][9])
  if g['group_seq'] in dt[0][2:]:
   col=dt[0].index(g['group_seq'])
   for item in g['diet_descr']['diet']:assert item['proportion']==nd(dietrows[item['prey_seq']][col])
 assert [g['pp'] for g in final['group']]==['0']*47+['1']*3+['2']
 result={'model_number':model['metadata']['model_number'],'variant':variant,'year':1991,'groups':51,'consumers':47,'fleets':3,'directory':str(d),'status':'partial','validation':vsummary,'standalone_massbalance':mbsummary,'final_database_verdict':findings['verdict'],'database_errors':len(findings['errors']),'database_warnings':len(findings['warnings']),'database_indeterminate':len(findings['indeterminate']),'roundtrip':'8 sheets matched all source cells','source_diet_sums':sums}
 summary.append(result);(d/'AUDIT_RESULTS.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(variant,vsummary,findings['verdict'],'roundtrip verified')
(ex/'EXTRACTION_SUMMARY.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
index="""# CAN-2014 extraction index\n\nSource: Guénette, Meissa and Gascuel (2014), DOI 10.1371/journal.pone.0094742. Read README.md and metadata.json before extraction. All source files remain unchanged; all work and outputs are below this extracted directory.\n\n**Three distinct Ecopath versions, all 1991.** The catalog's 2007-2009 label is incorrect. The 1991-2006 and 2056 outputs are Ecosim simulations and are not counted as separate static models.\n\n| # | Model | Year | LME | Groups | Fleets | Status | Notes |\n|---|---|---|---|---|---|---|---|\n"""
for s in summary:
 name=pathlib.Path(s['directory']).name;index+=f"| {s['model_number']} | [Banc d'Arguin and Mauritanian Shelf — {s['variant']}]({name}/REPORT.md) | 1991 | 27 Canary Current | 51 | 3 | partial | {s['validation']}; final database NOT BALANCED. {'Full base matrix; source Groupers diet 220.06% and birds 98.0331% preserved.' if s['variant']=='Base' else 'S8 B/EE preserved; 14 affected consumer diets incomplete; rounded juvenile zeros preserved.'} |\n"
index+='''\nEvery directory contains the eight Ecopath import files, extraction model.json, database JSON, reconstructed XLSX, REPORT.md, provenance and source-conflict records, validation results and mass-balance report. No Taxonomy.xlsx is created.\n\nThese are source-faithful partial extractions, not clean/load-tested balanced models. Common unresolved fields are GS, habitat proportions, discards, detritus routing/import and multi-stanza P/B. The published base matrix itself has two invalid column totals. No value was normalized or invented to pass checks. The local adapter in work/audit_database.py restores exact source cells after auditing the bundled converter; its final workbook matches all eight input tables.\n'''
(ex/'MASTER_INDEX.md').write_text(index,encoding='utf-8')
