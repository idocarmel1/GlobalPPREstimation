"""Apply reviewed LME037 metadata without rewriting generated numerical sheets."""
from pathlib import Path
import copy, hashlib, json, os, posixpath, shutil, sys, tempfile, zipfile
import xml.etree.ElementTree as ET
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries
ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'regions/LME_037/models/extraction_review_20260928'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('',NS);ET.register_namespace('r',R)
def tag(s):return '{'+NS+'}'+s
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def val(c):
    t=c.get('t');v=c.find(tag('v'))
    if t=='inlineStr':return ''.join(c.itertext())
    if v is None:return None
    if t=='s':return shared[int(v.text)]
    if t=='b':return v.text=='1'
    if t in ('str','e'):return v.text
    f=float(v.text);return int(f) if f.is_integer() else f
def grid(root):return {c.get('r'):val(c) for c in root.iter(tag('c'))}
def row_values(row):return {column_index_from_string(''.join(filter(str.isalpha,c.get('r')))):val(c) for c in row}
def put(row,col,v):
    ref=f'{get_column_letter(col)}{row.get("r")}'
    c=next((c for c in row if c.get('r')==ref),None)
    if c is None:c=ET.SubElement(row,tag('c'),{'r':ref})
    for ch in list(c):c.remove(ch)
    c.attrib.pop('t',None)
    if v is None:return
    if isinstance(v,str):
        assert len(v)<=32767
        c.set('t','inlineStr');ET.SubElement(ET.SubElement(c,tag('is')),tag('t')).text=v
    elif isinstance(v,bool):c.set('t','b');ET.SubElement(c,tag('v')).text=str(int(v))
    else:ET.SubElement(c,tag('v')).text=str(v)
    row[:]=sorted(row,key=lambda c:column_index_from_string(''.join(filter(str.isalpha,c.get('r')))))
def resolve(base,target):return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(base,target))
proposal=json.loads((OUT/'central_metadata_patch.json').read_text(encoding='utf8'))
p=ROOT/'Project.xlsx';initial=sha(p)
with zipfile.ZipFile(p) as z:infos=z.infolist();raw={i.filename:z.read(i) for i in infos}
shared=[''.join(n.itertext()) for n in ET.fromstring(raw['xl/sharedStrings.xml'])] if 'xl/sharedStrings.xml' in raw else []
rels={r.get('Id'):r.get('Target') for r in ET.fromstring(raw['xl/_rels/workbook.xml.rels'])}
paths={s.get('name'):resolve('xl',rels[s.get('{'+R+'}id')]) for s in ET.fromstring(raw['xl/workbook.xml']).find(tag('sheets'))}
changed={};allowed=set();summary=[]
# Current region rank is reused, never inferred from the source paper rank.
reg=ET.fromstring(raw[paths['Regions & status']]);rr=list(reg.find(tag('sheetData')))
h=list(row_values(rr[1]).values())
ranks=[dict(zip(h,[row_values(row).get(i) for i in range(1,len(h)+1)])) for row in rr[2:]]
rank=next(r['atlas_region_rank'] for r in ranks if r.get('unit_id')=='LME_037')
for sheet,items in [('Papers',proposal['paper_updates']),('Models & coverage',proposal['model_rows_to_register'])]:
    part=paths[sheet];node=ET.fromstring(raw[part]);sd=node.find(tag('sheetData'))
    header=list(row_values(list(sd)[1]).values())
    for item in items:
        patch=item['changes'] if sheet=='Papers' else copy.deepcopy(item)
        keys={'article_id':item['article_id'],'unit_id':'LME_037'} if sheet=='Papers' else {'model_id':item['model_id'],'unit_id':'LME_037'}
        matches=[]
        for row in list(sd)[2:]:
            d={k:row_values(row).get(i) for i,k in enumerate(header,1)}
            if all(d.get(k)==v for k,v in keys.items()):matches.append(row)
        assert len(matches)<=1
        if sheet=='Papers':assert len(matches)==1,'Expected existing source-paper row'
        if matches:row=matches[0]
        else:
            n=max(int(row.get('r')) for row in sd)+1
            row=ET.SubElement(sd,tag('row'),{'r':str(n)})
            patch.update(keys,atlas_region_rank=rank)
        assert set(patch)<=set(header),set(patch)-set(header)
        if sheet=='Models & coverage':assert (ROOT/patch['model_path']).is_file()
        for k,v in patch.items():
            col=header.index(k)+1;ref=f'{get_column_letter(col)}{row.get("r")}'
            old=grid(node).get(ref)
            if old==v:continue
            if k in {'notes','coverage_note','loadability_evidence','documentation_evidence'} and old and v and str(v) not in str(old):v=str(old)+'\n\n[2026-09-28 source recovery] '+str(v)
            put(row,col,v);allowed.add((part,ref));summary.append({'sheet':sheet,'key':keys,'field':k,'before':old,'after':v})
    last=max(int(row.get('r')) for row in sd)
    dim=node.find(tag('dimension'))
    if dim is not None:dim.set('ref',f'A1:{get_column_letter(len(header))}{last}')
    relpath=posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')
    sr={r.get('Id'):r.get('Target') for r in ET.fromstring(raw[relpath])}
    tp=list(node.find(tag('tableParts')));assert len(tp)==1
    tablepart=resolve(posixpath.dirname(part),sr[tp[0].get('{'+R+'}id')]);table=ET.fromstring(raw[tablepart])
    _,first,_,_=range_boundaries(table.get('ref'));ref=f'A{first}:{get_column_letter(len(header))}{last}'
    table.set('ref',ref);table.find(tag('autoFilter')).set('ref',ref)
    changed[tablepart]=ET.tostring(table,encoding='utf-8');changed[part]=ET.tostring(node,encoding='utf-8')
    ET.fromstring(changed[part]);ET.fromstring(changed[tablepart])
    oldgrid=grid(ET.fromstring(raw[part]));newgrid=grid(node)
    for ref in set(oldgrid)|set(newgrid):
        if oldgrid.get(ref)!=newgrid.get(ref):assert (part,ref) in allowed
    assert len([r.get('r') for r in sd])==len(set(r.get('r') for r in sd))
backup=OUT/f'Project_before_metadata_{initial[:12]}.xlsx';shutil.copy2(p,backup)
fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=OUT);os.close(fd)
try:
    with zipfile.ZipFile(tmp,'w') as z:
        for i in infos:z.writestr(i,changed.get(i.filename,raw[i.filename]))
    with zipfile.ZipFile(tmp) as z:
        assert set(z.namelist())==set(raw)
        for k,v in raw.items():
            if k not in changed:assert z.read(k)==v
        tables=[n for n in z.namelist() if n.startswith('xl/tables/') and n.endswith('.xml')];assert len(tables)==13
        for n in tables:
            t=ET.fromstring(z.read(n));assert t.find(tag('autoFilter')).get('ref')==t.get('ref')
    assert sha(p)==initial,'Concurrent central mutation'
    os.replace(tmp,p)
finally:
    if os.path.exists(tmp):os.unlink(tmp)
audit={'status':'PASS','before_sha256':initial,'after_sha256':sha(p),'backup':str(backup.relative_to(ROOT)),'modified_parts':list(changed),'all_other_parts_byte_identical':True,'changes':summary,'native_tables_and_filters':13,'selection':False}
(OUT/'CENTRAL_REGISTRATION_VERIFICATION.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'status':'PASS','after_sha256':audit['after_sha256'],'modified_parts':list(changed)},indent=2))
