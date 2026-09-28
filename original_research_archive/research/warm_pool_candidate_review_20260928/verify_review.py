"""Check only the admitted source artifacts; does not claim model validation."""
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[2]
p=json.loads((BASE/'REGISTRATION_PROPOSAL.json').read_text(encoding='utf-8'))
assert p['workbook_writes_performed'] is False and p['selected'] is False
files=[]
for record in p['records']:
    for item in record['material_files']:
        f=ROOT/item['relative_path']; raw=f.read_bytes()
        assert len(raw)==item['size_bytes']
        assert hashlib.sha256(raw).hexdigest()==item['sha256']
        files.append(item['relative_path'])
    review=ROOT/record['review_path']
    admission=json.loads((review/'ADMISSION_STATUS.json').read_text(encoding='utf-8'))
    assert admission['diagnostic_status']=='NOT_RUN' and admission['diagnostic_results'] is None
    assert not (review.parent/'model.json').exists()
review=ROOT/p['records'][0]['review_path']
partial=json.loads((review/'SOURCE_ONLY_PARTIAL.json').read_text(encoding='utf-8'))
evidence=json.loads((review/'TABLE1_CELL_EVIDENCE.json').read_text(encoding='utf-8'))
groups=partial['groups']
assert len(groups)==46 and len(evidence)==276
for e in evidence:
    assert groups[e['group']-1][e['parameter']]==(None if e['printed_value']=='–' else e['printed_value'])
raw=(review/'Basic_input_PARTIAL.csv').read_bytes()
assert raw.count(b'\r\n')==47 and b'"' not in raw
rows=list(csv.reader(raw.decode('utf-8').splitlines()))
for row,g in zip(rows[1:],groups):
    assert row[1]==g['name'] and row[3]==g['biomass']
    assert row[5]==(g['pb'] or '') and row[6]==(g['qb'] or '')
    assert row[7]==g['ee'] and row[9]==(g['pq'] or '')
    assert row[2]==row[4]==row[8]==row[10]==row[11]==''
for filename in ['TL_source.xlsx','Taxonomy_unverified.xlsx']:
    w=load_workbook(review/filename,read_only=True); assert w.active.max_row==47; w.close()
result={'artifact_verification':'PASS','verified_complete_pdf_files':files,'verified_table1_rows':46,
        'verified_table1_cells':276,'numeric_cells':266,'source_dash_cells':10,
        'unknown_fields_preserved':True,'canonical_models_generated':0,'diagnostic_calls':0,
        'scientific_model_validation':'NOT_RUN_MISSING_VERIFIED_INPUTS',
        'note':'Artifact PASS is not Ecopath balance, SPPR health or extraction completion.'}
(BASE/'FINAL_ARTIFACT_VERIFICATION.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,indent=2))
