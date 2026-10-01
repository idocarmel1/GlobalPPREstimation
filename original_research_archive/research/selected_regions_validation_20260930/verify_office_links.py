"""Read-only Office hyperlink audit with a physical, sparse repository relocation.

Checks destinations and explicit effective hyperlink appearance. Scientific
arithmetic and rendered page layout are verified separately by regional reviews.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path, PureWindowsPath
import re
import shutil
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
BLUES = {"0563C1", "0000FF", "1155CC", "0070C0", "1F4E79"}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def blue(value):
    return isinstance(value, str) and value.upper()[-6:] in BLUES


def word_links(path):
    found = []
    with ZipFile(path) as package:
        styles = ET.fromstring(package.read("word/styles.xml"))
        by_style = {s.get(W + "styleId"): s for s in styles.findall(W + "style")}

        def properties(run):
            props = {}
            rp = run.find(W + "rPr")
            style = rp.find(W + "rStyle") if rp is not None else None
            chain = []
            seen = set()
            sid = style.get(W + "val") if style is not None else None
            while sid and sid in by_style and sid not in seen:
                seen.add(sid)
                s = by_style[sid]
                chain.append(s.find(W + "rPr"))
                parent = s.find(W + "basedOn")
                sid = parent.get(W + "val") if parent is not None else None
            for pr in [*reversed(chain), rp]:
                if pr is not None:
                    for item in pr:
                        props[item.tag] = item.get(W + "val")
            return props

        for member in package.namelist():
            if not re.fullmatch(r"word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml", member):
                continue
            doc = ET.fromstring(package.read(member))
            relpath = "word/_rels/" + Path(member).name + ".rels"
            rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(package.read(relpath))} if relpath in package.namelist() else {}
            for h in doc.iter(W + "hyperlink"):
                target = rels.get(h.get(R + "id"))
                anchor = h.get(W + "anchor")
                if not target and anchor:
                    target = "#" + anchor
                if not target:
                    raise AssertionError(f"Hyperlink without target in {path.name}")
                texts = []
                for run in h.iter(W + "r"):
                    label = "".join(t.text or "" for t in run.iter(W + "t"))
                    if not label:
                        continue
                    props = properties(run)
                    assert blue(props.get(W + "color")), (path.name, label, "non-blue link", props)
                    assert props.get(W + "u") in {"single", "double"}, (path.name, label, "missing underline")
                    texts.append(label)
                found.append({"part": member, "target": target, "label": "".join(texts), "blue_underlined": True})
            fields = ["".join(doc.itertext())] if doc.findall(".//" + W + "instrText") else []
            assert not any("HYPERLINK" in f for f in fields), "Field-code hyperlink needs a separate effective-format review"
        if "docProps/app.xml" in package.namelist():
            app = ET.fromstring(package.read("docProps/app.xml"))
            assert not any(e.text for e in app.iter() if ET.QName(e).localname == "HyperlinkBase"), "Nonempty HyperlinkBase"
    return found


def excel_links(path):
    found = []
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    assert cell.data_type != "e", (path.name, sheet.title, cell.coordinate, cell.value)
                    target = None
                    if cell.hyperlink:
                        target = cell.hyperlink.target or ("#" + cell.hyperlink.location if cell.hyperlink.location else None)
                    elif cell.data_type == "f" and re.search(r"\bHYPERLINK\s*\(", cell.value, re.I):
                        match = re.search(r'HYPERLINK\s*\(\s*"((?:[^"]|"")*)"', cell.value, re.I)
                        assert match, "Dynamic hyperlink formula requires separate target verification"
                        target = match.group(1).replace('""', '"')
                    if target:
                        color = cell.font.color
                        assert color is not None and color.type == "rgb" and blue(color.rgb), (path.name, sheet.title, cell.coordinate, "non-blue link")
                        assert cell.font.underline in {"single", "double", "singleAccounting", "doubleAccounting"}, (path.name, sheet.title, cell.coordinate, "missing underline")
                        found.append({"sheet": sheet.title, "cell": cell.coordinate, "label": str(cell.value), "target": target, "blue_underlined": True})
    finally:
        workbook.close()
    return found


def anchor_exists(path, fragment):
    if not fragment:
        return True
    fragment = unquote(fragment)
    if path.suffix.lower() == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True)
        try:
            sheet = fragment.split("!", 1)[0].strip("'").replace("''", "'")
            return sheet in workbook.sheetnames or fragment in workbook.defined_names
        finally:
            workbook.close()
    if path.suffix.lower() == ".docx":
        with ZipFile(path) as package:
            doc = ET.fromstring(package.read("word/document.xml"))
            return fragment in {e.get(W + "name") for e in doc.iter(W + "bookmarkStart")}
    if path.suffix.lower() == ".pdf" and re.fullmatch(r"page=\d+", fragment):
        import pymupdf
        with pymupdf.open(path) as pdf:
            return 1 <= int(fragment[5:]) <= pdf.page_count
    if path.suffix.lower() == ".md":
        anchors, counts = set(), {}
        for heading in re.findall(r"^#{1,6}\s+(.+?)\s*#*\s*$", path.read_text(encoding="utf-8"), re.M):
            heading = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", heading)
            slug = re.sub(r"[^\w -]", "", heading.lower()).replace(" ", "-")
            n = counts.get(slug, 0)
            anchors.add(slug + (f"-{n}" if n else ""))
            counts[slug] = n + 1
        return fragment in anchors
    return None


def verify(units, relocate):
    inventory = json.loads((HERE / "preflight.json").read_text(encoding="utf-8"))["regions"]
    target_root = HERE / "work" / "office_relocation" / "Relocated validation repository"
    results = []
    for item in inventory:
        unit = item["unit_id"]
        if unit not in units:
            continue
        region = ROOT / "regions" / unit
        paths = [region / ("Model_validation_" + item["model_id"] + ".docx"), region / (unit.replace("_", "") + "_taxon_mapping_appendix.xlsx")]
        for path in paths:
            assert path.is_file(), str(path)
            initial_hash = sha(path)
            links = word_links(path) if path.suffix == ".docx" else excel_links(path)
            assert links, f"No links in {path}"
            if relocate:
                relocated = target_root / path.relative_to(ROOT)
                relocated.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, relocated)
                assert sha(relocated) == initial_hash
            for link in links:
                target = link["target"]
                parsed = urlsplit(target)
                if parsed.scheme in {"https", "http", "mailto"}:
                    link["public_destination"] = True
                    continue
                local = unquote(parsed.path)
                assert not parsed.scheme and not Path(local).is_absolute() and not PureWindowsPath(local).drive and not local.startswith("\\"), (path.name, target, "absolute target")
                destination = (path.parent / local).resolve() if local else path
                assert destination.is_relative_to(ROOT) and destination.exists(), (path.name, target, "missing/outside repository")
                link["repository_destination"] = destination.relative_to(ROOT).as_posix()
                link["anchor_verified"] = anchor_exists(destination, parsed.fragment)
                assert link["anchor_verified"] is not False, (path.name, target, "missing anchor")
                if relocate:
                    moved_target = target_root / destination.relative_to(ROOT)
                    if destination.is_dir():
                        moved_target.mkdir(parents=True, exist_ok=True)
                    else:
                        moved_target.parent.mkdir(parents=True, exist_ok=True)
                        shutil.copy2(destination, moved_target)
                        assert sha(destination) == sha(moved_target)
                    resolved = (relocated.parent / local).resolve() if local else relocated
                    assert resolved == moved_target.resolve() and resolved.exists()
                    link["relocated_destination_verified"] = True
            assert sha(path) == initial_hash, "Source Office file changed during audit"
            results.append({"unit_id": unit, "file": path.relative_to(ROOT).as_posix(), "sha256": initial_hash, "link_count": len(links), "links": links})
    destination = HERE / "verification" / ("office_links_" + "_".join(units) + ".json")
    destination.write_text(json.dumps({"files": results, "physical_relocation": relocate, "all_checked_links_passed": True,
                                     "scope": "Effective link style/destination audit; rendered layout and scientific arithmetic are separate checks."}, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    print(json.dumps({"files": len(results), "links": sum(r["link_count"] for r in results), "physical_relocation": relocate, "passed": True, "evidence": destination.relative_to(ROOT).as_posix()}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("units", nargs="+")
    parser.add_argument("--relocate", action="store_true")
    args = parser.parse_args()
    verify(args.units, args.relocate)
