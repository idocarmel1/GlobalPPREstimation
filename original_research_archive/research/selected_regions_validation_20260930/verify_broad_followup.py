"""Independent arithmetic, allowed-diff and appendix checks after broad mapping edits."""
import argparse
import json
import math
import sys
from collections import defaultdict, Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
import workbooks as W
import regional as R


def verify(unit):
    folder = ROOT / 'regions' / unit
    p = folder / (unit + '.xlsx')
    fingerprint = W.sha(p)
    b = W.read_book(p)
    before = W.read_book(HERE / 'work/coordinator_broad_followup' / unit / p.name)
    mid = W.overview(b)['selected_model_id']
    ev = folder / 'validation_reports' / mid
    receipt = json.loads((ev / 'broad_candidate_followup.json').read_text(encoding='utf-8'))
    changed = {c['taxon'] for c in receipt['changes']}
    o = W.overview(b)
    assert o['calculation_input_sha256'] == W.input_hash(b)
    assert o['calculation_result_sha256'] == R.result_hash(b)
    checks = Counter()
    for sheet in ['Catch', 'Classic PPR', 'Selected model groups', 'NPP', 'Diagnostics']:
        assert b[sheet] == before[sheet], ('protected', sheet)
        checks['protected_tables_exact'] += len(b[sheet])
    coeff = {(r['taxon'], r['scope'], r['method']): r['sppr'] for r in W.records(b, 'PPR', 'Taxon SPPR')}
    old_coeff = {(r['taxon'], r['scope'], r['method']): r['sppr'] for r in W.records(before, 'PPR', 'Taxon SPPR')}
    assert all(v == old_coeff[k] for k, v in coeff.items() if k[0] not in changed)
    group = {(r['group'], r['scope'], r['method']): r['sppr'] for r in W.records(b, 'Selected model groups', 'Group SPPR')}
    mapping = defaultdict(list)
    for r in W.records(b, 'PPR', 'Matching'):
        mapping[r['taxon']].append((r['group'], r['weight']))
    for key, value in coeff.items():
        taxon, scope, method = key
        candidates = [(weight, group[name, scope, method]) for name, weight in mapping[taxon]]
        expected = round(math.fsum(weight * value for weight, value in candidates), 6) if candidates and all(W.finite(value) for _, value in candidates) else None
        if taxon in changed:
            assert value == expected, ('changed taxon coefficient', key, value, expected)
        else:
            assert value == old_coeff[key], ('unrelated coefficient must remain exact', key)
            assert (value is None and expected is None) or (W.finite(value) and W.finite(expected) and math.isclose(value, expected, rel_tol=1e-14, abs_tol=1e-6)), ('stored coefficient reproduction', key, value, expected)
        checks['taxon_coefficients'] += 1
    catch = {(r['taxon'], r['catch_basis']): r for r in W.records(b, 'Catch', 'Catch')}
    taxa = sorted(mapping)
    simple = {r['taxon']: r['sppr'] for r in W.records(b, 'Classic PPR', 'Taxa')}
    unidentified = {r['taxon']: bool(r['unidentified']) for r in W.records(b, 'Catch', 'Catch')}
    annual_index = {}
    max_relative = 0.0

    def close(actual, expected, key):
        nonlocal max_relative
        if expected is None:
            assert actual is None, key
        else:
            assert W.finite(actual) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-6), (key, actual, expected)
            max_relative = max(max_relative, abs(actual - expected) / max(abs(expected), 1))

    for r in W.records(b, 'PPR', 'Annual'):
        key = tuple(r[k] for k in ['model_id', 'scope', 'method', 'catch_basis', 'unidentified', 'metric'])
        annual_index[key] = r
        for year in W.YEARS:
            pairs = []
            catches = []
            for taxon in taxa:
                c = catch[taxon, r['catch_basis']][year]
                v = coeff[taxon, r['scope'], r['method']]
                if unidentified[taxon] and r['unidentified'] == 'zero':
                    v = 0.0
                elif unidentified[taxon] and r['unidentified'] == 'simple':
                    v = simple.get(taxon)
                catches.append(c)
                if W.finite(c) and W.finite(v):
                    pairs.append((c, v))
            expected = (math.fsum(catches) if r['metric'] == 'catch' else
                        (math.fsum(c * v for c, v in pairs) if r['metric'] == 'ppr' else math.fsum(c for c, v in pairs))
                        if W.numeric_status(r['status']) and pairs else None)
            close(r[year], expected, (*key, year))
            checks['annual_cells'] += 1
    for r in W.records(b, 'Classic PPR', 'Annual'):
        annual_index[tuple(r[k] for k in ['model_id', 'scope', 'method', 'catch_basis', 'unidentified', 'metric'])] = r
    npp = {r['method']: r for r in W.records(b, 'NPP', 'NPP')}
    for r in W.records(b, 'PPR–NPP', 'Ratios'):
        key = tuple(r[k] for k in ['model_id', 'scope', 'method', 'catch_basis', 'unidentified']) + ('ppr',)
        for year in W.YEARS:
            v, n = annual_index[key][year], npp[r['npp_method']][year]
            close(r[year], 100 * v / 9 / n if W.finite(v) and W.finite(n) and n > 0 else None, (*key, r['npp_method'], year))
            checks['ratio_cells'] += 1
    for r in W.records(b, 'PPR', 'Taxon PPR inspected year'):
        c = catch[r['taxon'], r['catch_basis']][r['year']]
        v = coeff[r['taxon'], r['scope'], r['method']]
        close(r['ppr_wet_tonnes'], c * v if W.finite(c) and W.finite(v) else None, r['taxon'])
        checks['taxon_details'] += 1
    app = next(folder.glob('*taxon_mapping_appendix.xlsx'))
    w = openpyxl.load_workbook(app, data_only=False)
    for r in receipt['office_refresh']['mapping_rows']:
        row = r['row']
        assert w['Taxon mapping'].cell(row, 1).value == r['taxon']
        assert w['Taxon mapping'].cell(row, 5).value == r['display']
        assert w['Taxon mapping'].cell(row, 7).value == r['reason']
    if unit == 'LME_034':
        source = W.records(b, 'PPR', 'Allocation assumptions')
        office = list(w['Allocation evidence'].values)[6:]
        assert len(source) == len(office)
        office_by_key = {(row[0], row[1], row[2]): row for row in office}
        assert len(office_by_key) == len(office), 'Duplicate appendix candidate keys'
        assert set(office_by_key) == {(a['taxon'], a['seq'], a['group']) for a in source}, 'Appendix candidate universe differs'
        for a in source:
            row = office_by_key[a['taxon'], a['seq'], a['group']]
            for key, i in [('source_printed_catch_density', 3), ('source_model_catch_density', 4), ('loaded_catch_density', 5), ('candidate_total', 6), ('weight', 7)]:
                close(row[i], a[key], (a['taxon'], a['seq'], key))
            checks['appendix_allocation_candidates'] += 1
    w.close()
    assert W.sha(p) == fingerprint
    result = {'unit_id': unit, 'workbook_sha256': fingerprint, 'checks': dict(checks),
              'max_relative_arithmetic_difference': max_relative, 'unrelated_taxon_coefficients_exact': True,
              'complete_source_quantities_and_zero_candidates': receipt['changes'], 'passed': True}
    (ev / 'broad_candidate_calculations_verification.json').write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps({'unit_id': unit, 'checks': dict(checks), 'passed': True}), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('units', nargs='+')
    for unit in parser.parse_args().units:
        verify(unit)
