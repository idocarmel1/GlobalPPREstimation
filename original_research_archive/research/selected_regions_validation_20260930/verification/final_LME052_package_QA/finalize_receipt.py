from pathlib import Path
import json,hashlib,datetime
import openpyxl
R=Path.cwd();Q=R/'original_research_archive/research/selected_regions_validation_20260930/work/final_LME052_package_QA';E=R/'regions/LME_052/validation_reports/52_1_Sea_of_Okhotsk_NE_(1980)';sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();obj=lambda p:{'path':p.relative_to(R).as_posix(),'sha256':sha(p),'bytes':p.stat().st_size}
a=json.loads((Q/'automated_checks.json').read_text(encoding='utf-8'));h=json.loads((E/'coordination_handoff.json').read_text(encoding='utf-8-sig'))
assert not a['manifest_hash_failures'] and not a['link_failures']
assert len(a['manual_researcher_rows'])==3 and all(x['exact_xml']for x in a['manual_researcher_rows']) and a['section_geometry_identical']
assert len(a['visual_files'])==13 and all(sha(R/x['path'])==x['sha256'] for x in a['visual_files'])
app=R/'regions/LME_052/LME052_taxon_mapping_appendix.xlsx';wb=openpyxl.load_workbook(app,data_only=False);ws=wb['Taxon mapping'];taxa=[ws.cell(i,1).value for i in range(8,159)];assert len(taxa)==151 and len(set(taxa))==151
src=json.loads((E/'taxonomy_source_entries.json').read_text(encoding='utf-8-sig'));sel=[x for x in src if x['id']=='Taxon125557'];assert len(sel)==1 and 'reporting rank Family' in sel[0]['supports'] and 'Scarinae' in sel[0]['supports'];assert not any(x['id']=='Taxon398089' for x in src)
tlinks=json.loads((E/'taxon_source_links.json').read_text(encoding='utf-8-sig'));assert tlinks['Scaridae']==['Taxon125557']
raw=R/'regions/LME_052/models/52_1_Sea_of_Okhotsk_NE_(1980)/source_evidence/mapping/52_1_Sea_of_Okhotsk_NE_(1980).taxonomy-sources.json';records=json.loads(raw.read_text(encoding='utf-8-sig'))['queries']['Scaridae'];fam=[x for x in records if x['AphiaID']==125557][0];assert fam['rank']=='Family' and fam['valid_AphiaID']==151787 and fam['valid_name']=='Scarinae'
ind=E/'independent_source_review';m=json.loads((ind/'selected_evidence_manifest.json').read_text(encoding='utf-8-sig'))
notes={
 'word_pages':'All six authenticated pages visually inspected: readable tables and figures, repeated headers, no clipped text or stranded headings; blue underlined links visible.',
 'appendix':'All seven authenticated previews inspected: mapping top/middle/bottom and longest row102 (Perciformes); Sources top/middle/bottom. Long explanations and multi-line weight vectors fit; numerical cells show readable values. Both sheets and all151 exact rows read directly.',
 'manual_fields':'The three researcher-owned rows have byte-identical serialized XML against current template, including calculation choices, open issues/actions and name/date placeholders; section geometry matches.',
 'rank':'Scaridae taxon_source_links selects only Taxon125557. Retained primary response identifies Family125557 -> accepted Scarinae151787; genus398089 -> Calotomus is retained raw but not adopted in taxonomy_source_entries or appendix source link. Mapping15=100%, membership Very low and allocation High remain explicit.',
 'limitations':'Word pages1/2/4/5/6 disclose unavailable original thesis, existing normalized diet/default assumptions, missing native catch/cutoff, fixed1980s proportions, whole-sea proxy and unknown author boundary, weak analogues and unresolved Cyprinidae. TE FAIL/negative signed contributions and unavailable ordinary positive confidence percentages are stated. Handoff browser TE explicitly says DIVERGED/Annual unavailable and excludes negative arithmetic; appendix quantities are labeled independent simple trophic-chain PPR.',
 'scope':'Protected scientific input fidelity is covered by coordinator checks, not rerun here. No source, runtime, model, regional workbook, Project or production Office file was edited. No scientific stages ran. This is package QA, not scientific approval.',
 'coordinator_followup':'Pending alignment draft subtitle remains as explicitly instructed; browser validation is false/pending in this regional handoff. Coordinator owns removal after browser verification and consequent DOCX/render/hash refresh.'
}
for x in a['visual_files']:
 x['independently_visually_inspected']=True;x['qa_findings']='No material layout defect observed'
receipt={
 'status':'PASS for bounded read-only final package QA at these exact identities; coordinator browser/title finalization remains',
 'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
 'reviewer':'/root/region_lme_050',
 'model_id':h['model_id'],
 'production_writes':[], 'scientific_stages_run':[], 'science_approval':False,
 'handoff':obj(E/'coordination_handoff.json'),
 'current_artifacts':[obj(R/h['final']['workbook']['path']),obj(R/h['final']['canonical_model']['path'])]+a['artifact_bindings'],
 'evidence_proofs':[obj(E/n) for n in ['retained_evidence_manifest.json','visual_qa_manifest.json','independent_source_review/selected_evidence_manifest.json','independent_source_review/candidate_vector_verification.json','independent_source_review/findings.json','taxon_audit.json','candidate_and_allocation_evidence.json','taxon_source_links.json','taxonomy_source_entries.json']]+[obj(raw)],
 'hash_check_count':len(a['manifest_hash_checks']),'distinct_checked_paths':len(set(c['path'] for c in a['manifest_hash_checks'])),'hash_failures':[],
 'independent_selected_evidence_files':len(m['files']),
 'manual_researcher_rows':a['manual_researcher_rows'],'section_geometry_identical':True,
 'actual_visible_hyperlinks':a['actual_link_count'],'word_hyperlinks':len([l for l in a['links']if 'sheet'not in l]),'appendix_hyperlinks':len([l for l in a['links']if 'sheet'in l]),'local_relative_links_physically_relocated':a['relative_link_count'],'link_failures':[],
 'appendix_sheet_dimensions':[{'sheet':s.title,'rows':s.max_row,'columns':s.max_column} for s in wb],
 'Scaridae_selected_record':{'rank':'Family','AphiaID':125557,'accepted_name':'Scarinae','valid_AphiaID':151787,'Taxon_mapping_row':131,'Sources_row':131,'adopted_wrong_genus':False},
 'current_TE_expected':h['browser_validation']['TE'],
 'visual_files':a['visual_files'],
 'findings':[], 'observations':notes,
 'scratch_automated_proof':obj(Q/'automated_checks.json'),'scratch_verification_script':obj(Q/'verify_package.py')
}
(Q/'final_package_QA_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(Q/'final_package_QA_notes.md').write_text('# LME 052 bounded final package QA\n\nPASS at the artifact identities in final_package_QA_receipt.json. No production files edited and no scientific stages run.\n\n'+'\n\n'.join(k+': '+v for k,v in notes.items())+'\n',encoding='utf-8')
print(json.dumps({'receipt':obj(Q/'final_package_QA_receipt.json'),'notes':obj(Q/'final_package_QA_notes.md'),'handoff':receipt['handoff'],'artifact_hashes':receipt['current_artifacts'],'hash_checks':receipt['hash_check_count'],'visual_files':len(a['visual_files']),'links':receipt['actual_visible_hyperlinks'],'independent_files':len(m['files'])},ensure_ascii=True,indent=2))
