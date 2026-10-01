"""Bounded current-state update; never loads an old book or executes a solver."""
import sys,json,copy,math,shutil,hashlib,re,os,importlib.util
from pathlib import Path
from decimal import Decimal
from zipfile import ZipFile
from lxml import etree as E
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path.cwd();SHARED=ROOT/'original_research_archive/research/selected_regions_validation_20260930';SCR=SHARED/'work/qa_residual_fish_scope'
sys.path.insert(0,str(ROOT/'tools'));import workbooks as W;import regional as R
TARGET={'Marine fishes not identified':100039,'Marine finfishes not identified':100139,'Marine groundfishes not identified':100239,'Marine pelagic fishes not identified':100339}
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';SN=lambda x:'{'+NS+'}'+x
WN='http://schemas.openxmlformats.org/wordprocessingml/2006/main';wq=lambda x:'{'+WN+'}'+x
def writej(p,j):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(j,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
def readj(p):return json.loads(p.read_text(encoding='utf-8'))
def module(name,path):
 s=importlib.util.spec_from_file_location(name,path);m=importlib.util.module_from_spec(s);s.loader.exec_module(m);return m
AC=module('accepted_inputs_readonly',SHARED/'check_accepted_inputs.py')
def backup(p):
 q=SCR/'before'/p.relative_to(ROOT);q.parent.mkdir(parents=True,exist_ok=True)
 if not q.exists():shutil.copy2(p,q)
 return q
def cell_column(address):
 col=0
 for letter in re.match(r'[A-Z]+',address).group():col=col*26+ord(letter)-64
 return col
def patch_package(p,edits,heights=None):
 """Transplant only cell values, keeping all non-worksheet OOXML byte-for-byte."""
 with ZipFile(p) as z:infos=z.infolist();parts={i.filename:z.read(i.filename) for i in infos}
 rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);targets={n.get('Id'):n.get('Target') for n in rels};wb=E.fromstring(parts['xl/workbook.xml']);paths={s.get('name'):'xl/'+targets[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].lstrip('/').removeprefix('xl/') for s in wb.find(SN('sheets'))}
 roots={}
 for s,address,old,new in edits:
  if s not in roots:roots[s]=E.fromstring(parts[paths[s]])
  root=roots[s];data=root.find(SN('sheetData'));ri=int(re.sub('[A-Z]','',address));rn=data.find(SN('row')+f"[@r='{ri}']")
  if rn is None:
   rn=E.Element(SN('row'),r=str(ri));later=next((x for x in data if int(x.get('r'))>ri),None)
   if later is None:data.append(rn)
   else:data.insert(data.index(later),rn)
  c=rn.find(SN('c')+f"[@r='{address}']")
  if c is None:
   c=E.Element(SN('c'),r=address);later=next((node for node in rn if node.tag==SN('c') and cell_column(node.get('r'))>cell_column(address)),None)
   if later is None:rn.append(c)
   else:rn.insert(rn.index(later),c)
  for child in list(c):c.remove(child)
  c.attrib.pop('t',None)
  if new is None:pass
  elif isinstance(new,bool):c.set('t','b');E.SubElement(c,SN('v')).text='1' if new else '0'
  elif isinstance(new,(int,float)):E.SubElement(c,SN('v')).text=repr(new)
  else:
   c.set('t','inlineStr');t=E.SubElement(E.SubElement(c,SN('is')),SN('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=str(new)
 for (s,ri),height in (heights or {}).items():
  if s not in roots:roots[s]=E.fromstring(parts[paths[s]])
  root=roots[s];rn=root.find(SN('sheetData')).find(SN('row')+f"[@r='{ri}']");assert rn is not None;rn.set('ht',str(height));rn.set('customHeight','1')
 for s,root in roots.items():parts[paths[s]]=E.tostring(root,encoding='utf-8',xml_declaration=True)
 tmp=p.with_suffix('.bounded.tmp')
 with ZipFile(tmp,'w') as z:
  for i in infos:z.writestr(i,parts[i.filename])
 os.replace(tmp,p)
 return {s:paths[s] for s in roots}
def block_edits(p,old,new):
 wb=openpyxl.load_workbook(p,read_only=True,data_only=False);changes=[]
 for sn,sheet in old.items():
  positions={};active=None;header=None
  for ri,row in enumerate(wb[sn],1):
   vals=[c.value for c in row]
   if vals and vals[0]=='@table':active=vals[1];header=None;positions[active]=[]
   elif active and any(v is not None for v in vals):
    if header is None:header=vals
    else:positions[active].append(ri)
  for tn,(hd,rr) in sheet.items():
   nh,nr=new[sn][tn];assert hd==nh,(sn,tn,'headers');assert len(nr)<=len(rr),(sn,tn,len(nr),len(rr))
   if W.digest_tables((hd,rr))==W.digest_tables((nh,nr)):continue
   for i,ri in enumerate(positions[tn]):
    before=rr[i];after=nr[i] if i<len(nr) else [None]*len(hd)
    for ci,(a,b) in enumerate(zip(before,after),1):
     if W.digest_tables([a])!=W.digest_tables([b]):changes.append((sn,openpyxl.utils.get_column_letter(ci)+str(ri),a,W.clean(b)))
 wb.close();return changes
def record_groups(ar):return ar.get('groups',ar.get('mapping',ar.get('adopted_groups',[])))
def source_bridge(folder,labels):
 out=[]
 for p in (folder/'raw').glob('*exploited*.json'):
  def visit(x):
   if isinstance(x,dict):
    if x.get('scientific_name') in labels:out.append({'path':str(p.relative_to(ROOT)),'sha256':W.sha(p),'record':x})
    for v in x.values():visit(v)
   elif isinstance(x,list):
    for v in x:visit(v)
  visit(readj(p))
 assert {r['record']['scientific_name'] for r in out}>=set(labels),('missing local retained-key bridge',labels)
 assert all(r['record']['taxon_key']==TARGET[r['record']['scientific_name']] for r in out)
 return out
def rationale(u,t):
 key=TARGET[t];scope='Bony-fish reporting scope is verified for exact SAU key '+str(key)+' (ISSCAAP 39; FAO Osteichthyes), rather than inferred from the English label.'
 if u=='LME_013':
  choices='Named and pooled source bony fish groups 6–11 and 13 provide an assumed broad composition; Skates 12 (Zearaja chilensis) is excluded.' if 'ground' not in t and 'pelagic' not in t else 'All named demersal/benthopelagic bony fish groups 7–11 and 13 are retained; Skates 12 is excluded and explicitly pelagic group 6 remains outside the groundfish candidate set.' if 'ground' in t else 'Source Small pelagic fish 6 is the sole explicitly pelagic bony-fish pool; the represented small/medium fish differ from unrepresented large oceanic predators.'
 elif u=='LME_014':
  choices='All represented named and pooled bony fish groups are retained, including source-zero Myctophidae 18; dedicated Dogfish 4, Sharks 27 and Skates 28 are excluded.' if 'pelagic' not in t else 'Blue Whiting 3, Hoki 9, Myctophidae 18 and Pelagic Fish 20 provide the pelagic/benthopelagic bony-fish approximation; Lamna-only Sharks 27 is excluded. Source-zero Myctophidae remains an eligible candidate.'
 elif u=='LME_026':
  choices='Represented bony fish groups 14–42 form the broad composition proxy; dedicated shark, ray/skate and torpedo pools 13 and 43–47 are excluded.' if 'ground' not in t and 'pelagic' not in t else 'Benthic/benthopelagic/deep bony fish groups 26–42 remain the broad groundfish proxy. Dedicated shark/ray/torpedo pools 13 and 43–47 are excluded; source 27 retains its listed deep benthic-associated bony fishes.' if 'ground' in t else 'Pelagic/benthopelagic bony fish groups 14–27 remain the approximation; dedicated Pelagic shark 13 and Ray&Skate 46 are excluded.'
  if 'ground' not in t:choices+=' Mixed NC large pelagic 14 remains eligible through source-listed Mola mola; Cetorhinus maximus and Mobula mobular are outside target scope, so full pooled source quantities retain an unobserved-composition mismatch.'
 else:
  choices='Named and residual bony fish groups 13–31 are retained; dedicated Demersal sharks and rays 32 and Pelagic sharks and rays 33 are excluded.' if 'ground' not in t and 'pelagic' not in t else 'Named demersal/benthopelagic bony fish groups 13–28 remain the groundfish approximation; Demersal sharks and rays 32 is excluded.' if 'ground' in t else 'Pelagic bony fish groups 29–31 remain the open-water approximation; Pelagic sharks and rays 33 is excluded.'
 return scope+' '+choices+' Residual taxonomic composition, unrepresented taxa and wider-region transfer remain unobserved; membership and overall confidence remain Very low.'
def allocate(u,t,ledger,ids,groups,canonical):
 if u=='LME_013':
  candidates=[copy.deepcopy(c) for c in ledger['candidates'] if int(c['seq']) in ids];catch=[c['source_catch_value'] for c in candidates];bio=[c['source_biomass_value'] for c in candidates]
 elif u=='LME_014':
  candidates=[{'seq':g,'group_name':groups[g]['group_name'],'source_catch_value':ledger['values'][str(g)],'source_biomass_value':groups[g]['biomass']} for g in ids];catch=[c['source_catch_value'] for c in candidates];bio=[c['source_biomass_value'] for c in candidates]
 elif u=='LME_026':
  candidates=[copy.deepcopy(c) for c in ledger['candidates'] if int(c['seq']) in ids];catch=[None if c['source_landings_sum'] is None else float(c['source_landings_sum']) for c in candidates];bio=[float(c['accepted_biomass_raw']) for c in candidates]
 else:
  candidates=[{'seq':g,'group_name':groups[g]['group_name'],'source_catch_value':float(canonical[g]['export']),'source_biomass_value':float(canonical[g]['biomass'])} for g in ids];catch=[c['source_catch_value'] for c in candidates];bio=[c['source_biomass_value'] for c in candidates]
 assert [int(c['seq']) for c in candidates]==ids,(u,t,'complete source candidates',ids)
 catchok=all(W.finite(x) and x>=0 for x in catch) and math.fsum(catch)>0
 values=catch if catchok else bio;assert all(W.finite(x) and x>=0 for x in values) and math.fsum(values)>0
 total=math.fsum(values);weights=[x/total for x in values];assert math.isclose(math.fsum(weights),1,abs_tol=1e-12)
 rule='W4' if catchok else 'W9';assert rule==ledger.get('allocation_rule',ledger.get('rule')),(u,t,'method changed',rule)
 for c,w in zip(candidates,weights):c['weight']=w
 return candidates,total,rule,{'complete_catch_attempt':catch,'complete_catch_usable':catchok,'source_biomass_values':bio,'selected_field':'source catch' if catchok else 'accepted biomass','denominator':total,'candidate_ids':ids,'weights':weights,'valid_zero_candidate_ids':[c['seq'] for c in candidates if c['weight']==0]}
def audit_update(u,a,ids,weights,groups,reason,allocation_reason,candidates,proof):
 old=copy.deepcopy(a);name=lambda g:groups[g]['group_name'];display='; '.join(name(g)+f' ({w*100:.6f}%)' for g,w in zip(ids,weights));full=reason+' '+allocation_reason
 if u in ['LME_013','LME_026']:
  a['groups']=[{'seq':g,'group_name':name(g),'weight':w} for g,w in zip(ids,weights)];a['membership_reason']=reason
  if u=='LME_013':a['mapping_display']=display;a['reason']=full
  else:
   a['membership_assumptions']=[reason];a['membership_evidence']=[c for c in a['membership_evidence'] if c['group_id'] in ids];a['appendix_reason']=full+' Membership Very low (M10); allocation Medium ('+a['allocation_rule']+'). Exact source values and candidates are retained in allocation_evidence.json and residual_fish_scope_followup.json.'
 elif u=='LME_014':
  a['mapping']=[{'seq':g,'group':name(g),'weight':w} for g,w in zip(ids,weights)];a['mapping_display']=display;a['positive_connection']=reason;a['mismatch_and_alternatives']='Unknown residual composition and 2020 Falkland shelf guild transfer to the whole LME and all catch years; source catches are complete model-pool quantities, not observed target composition.';a['reason']=full
 else:
  a['adopted_groups']=[{'group_id':g,'name':name(g),'weight':w} for g,w in zip(ids,weights)];a['display_mapping']=display;a['reason']=full;a['membership_source']+='; verified SAU exact reporting key and FAO Osteichthyes scope (residual_fish_scope_followup.json)'
  a['candidate_selection']=[{'group_id':g,'group':name(g),'eligibility':reason,'definition':groups[g].get('taxon_descr'),'catch':c['source_catch_value'],'biomass':c['source_biomass_value'],'catch_field':'model.json group.export','biomass_field':'model.json group.biomass','source_units':'model areal mass units; ratios are dimensionless'} for g,c in zip(ids,candidates)]
  a['allocation_calculation'].update({'values':proof['complete_catch_attempt'],'total':proof['denominator'],'weights':weights})
  for k in ['membership_review','allocation_review']:
   if k in a:a[k]['current_reporting_scope_followup']='residual_fish_scope_followup.json; former review retained as historical antecedent to this correction.'
  a['adoption_state']='adopted evidence-supported reporting-scope group/weight correction; accepted biological inputs preserved'
 for c in a.get('candidate_decisions',[]):
  gid=int(c['seq']);inc=gid in ids;k='reason' if 'reason' in c else 'rationale';c['decision']=('included' if inc else 'excluded') if u=='LME_013' else ('include' if inc else 'exclude')
  if inc:c[k]=reason
  elif gid in [g.get('seq') for g in record_groups(old)]:c[k]='Excluded by exact ISSCAAP39 / Osteichthyes reporting scope: this source group has no eligible bony member. Local definition retained in residual_fish_scope_followup.json.'
 a['reporting_scope_taxon_key']=TARGET[a['taxon']];a['reporting_scope_isscaap_id']=39;a['reporting_scope_followup']='residual_fish_scope_followup.json';return full,display
def main(u):
 folder=ROOT/'regions'/u;p=folder/(u+'.xlsx');b=W.read_book(p);o=W.overview(b);mid=o['selected_model_id'];ev=folder/'validation_reports'/mid;handp=ev/'coordination_handoff.json';hp=readj(handp);app=next(folder.glob('*taxon_mapping_appendix.xlsx'));doc=folder/('Model_validation_'+mid+'.docx');expected=hp.get('final_workbook_sha256') or hp.get('outputs',{}).get('regional_workbook',{}).get('sha256') or hp['identities']['current_region_sha256'];assert W.sha(p)==expected
 W.validate_region(b,p);assert R.result_hash(b)==o['calculation_result_sha256'];old=copy.deepcopy(b);initial={str(q.relative_to(ROOT)):W.sha(q) for q in [p,app,doc,handp]};[backup(q) for q in [p,app,doc,handp]]
 workev=ev/'selected_regions_review_20260930' if u=='LME_036' else ev;ap=workev/('taxon_audit_adopted.json' if u=='LME_036' else 'adopted_taxon_audit.json' if u=='LME_026' else 'taxon_audit.json');lp=workev/('allocation_audit.json' if u=='LME_036' else 'allocation_evidence.json');audit=readj(ap);ledj=readj(lp);ledrows=ledj['records'] if isinstance(ledj,dict) else ledj;[backup(q) for q in [ap,lp]]
 by={r['taxon']:r for r in audit};ledby={r['taxon']:r for r in ledrows};groups={int(g['seq']):g for g in W.records(b,'Selected model groups','Groups')};model=readj(folder/o['model_path']);canonical={int(g['group_seq']):g for g in model['group']}
 cands={'LME_013':{'all':[6,7,8,9,10,11,13],'ground':[7,8,9,10,11,13],'pelagic':[6]},'LME_014':{'all':[3,5,6,7,8,9,13,14,18,20,23,24,29,33],'ground':[],'pelagic':[3,9,18,20]},'LME_026':{'all':list(range(14,43)),'ground':list(range(26,43)),'pelagic':list(range(14,28))},'LME_036':{'all':list(range(13,32)),'ground':list(range(13,29)),'pelagic':[29,30,31]}}[u]
 labels=[t for t in TARGET if t in by];bridge=source_bridge(folder,labels);manifest=readj(SHARED/'verification/reporting_scope_sources/source_manifest.json');assert all(W.sha(ROOT/s['path'])==s['sha256'] for s in manifest)
 changes=[];targetmap={};arithmetic=[]
 for t in labels:
  a=by[t];pre=copy.deepcopy(a);ledger=ledby.get(t);ids=cands['ground' if 'ground' in t else 'pelagic' if 'pelagic' in t else 'all'];assert ids
  # Single LME013 pelagic requires no allocation proxy; keep its W1 evidence.
  if u=='LME_013' and t.startswith('Marine pelagic'):
   candidates=[{'seq':6,'group_name':groups[6]['group_name'],'weight':1}];total=1;rule='W1';proof={'candidate_ids':[6],'weights':[1],'single_group_weight':1};allocation_reason=a['allocation_reason'];weights=[1]
  else:
   candidates,total,rule,proof=allocate(u,t,ledger,ids,groups,canonical);weights=[c['weight'] for c in candidates];allocation_reason=a.get('allocation_reason',ledger.get('reason',a.get('transfer_assumption','')))
  reason=rationale(u,t);full,display=audit_update(u,a,ids,weights,groups,reason,allocation_reason,candidates,proof)
  rr=[r for r in W.records(old,'PPR','Matching') if r['taxon']==t];conf=rr[0]['confidence'];targetmap[t]=[{'model_id':mid,'taxon':t,'group':groups[g]['group_name'],'weight':w,'confidence':conf,'evidence':rr[0]['evidence']+'; validation_reports/'+mid+'/residual_fish_scope_followup.json','explanation':full} for g,w in zip(ids,weights)]
  if ledger:
   oldledger=copy.deepcopy(ledger)
   if u=='LME_013':ledger['candidate_set']=ids;ledger['candidates']=candidates;ledger['allocation_denominator']=total
   elif u=='LME_014':ledger.update({'candidate_ids':ids,'values':{str(c['seq']):c['source_catch_value'] for c in candidates},'weights':{str(c['seq']):c['weight'] for c in candidates},'total':total,'zero_candidate_ids':proof['valid_zero_candidate_ids']})
   elif u=='LME_026':
    ledger['candidates']=candidates;ledger['allocation_denominator']=str(sum(Decimal(c['source_landings_sum'] if rule=='W4' else c['accepted_biomass_raw']) for c in candidates));ledger['candidate_inclusion_evidence']=[r for r in ledger['candidate_inclusion_evidence'] if r['group_id'] in ids];ledger['candidate_inclusion_rationale']=reason
   else:ledger.update({'candidate_ids':ids,'weights':weights,'source_values':proof['complete_catch_attempt'],'source_total':total})
   ledger['reporting_scope_followup']='../residual_fish_scope_followup.json' if u=='LME_036' else 'residual_fish_scope_followup.json';ledger['reporting_scope_excluded_candidates']=[c for c in oldledger.get('candidates',[]) if int(c['seq']) not in ids]
  changes.append({'key':{'unit_id':u,'model_id':mid,'taxon':t,'taxon_key':TARGET[t]},'before_matching':rr,'after_matching':targetmap[t],'before_audit':pre,'local_source_definition_review':[{k:g.get(k) for k in ['seq','group_name','taxon_descr','catch','biomass']} for g in groups.values()],'complete_eligible_candidate_ids':ids,'allocation':proof,'overall_confidence_preserved':'Very low','membership_rationale':reason})
  arithmetic.append({'taxon':t,'candidate_count':len(ids),'weight_sum':math.fsum(weights),'rule':rule,'denominator':total,'zero_candidates':proof.get('valid_zero_candidate_ids',[])})
 # Replace the exact labels once; preserve every unrelated mapping record/order.
 matching=W.records(b,'PPR','Matching');new=[];done=set()
 for r in matching:
  if r['taxon'] in targetmap:
   if r['taxon'] not in done:new.extend(targetmap[r['taxon']]);done.add(r['taxon'])
  else:new.append(r)
 b['PPR']['Matching']=W.table_dict(new)
 # Refresh current canonical review/ledger blocks from their exact retained schemas.
 for s,tn in [('PPR','Mapping review'),('Diagnostics','Taxon mapping review')]:
  if tn not in b.get(s,{}):continue
  hh,rr=b[s][tn];newrows=[]
  for r in W.records(b,s,tn):
   if r.get('taxon') in targetmap:
    a=by[r['taxon']]
    for k in ['reason','membership_reason','positive_connection','mismatch_and_alternatives','membership_assumptions','membership_evidence']:
     if k in r and k in a:r[k]=a[k]
   newrows.append([W.clean(r.get(k)) for k in hh])
  b[s][tn]=(hh,newrows)
 for s,tn in [('PPR','Allocation assumptions'),('PPR','Validation allocation review'),('Diagnostics','Allocation assumptions')]:
  if tn not in b.get(s,{}):continue
  hh,rr=b[s][tn];rebuilt=[]
  for r in W.records(b,s,tn):
   t=r.get('taxon')
   if t not in targetmap:rebuilt.append(r);continue
   group=r.get('group');ids=next(c['complete_eligible_candidate_ids'] for c in changes if c['key']['taxon']==t);keep={groups[g]['group_name'] for g in ids}
   if group and group not in keep:continue
   if group:
    g=next(g for g in ids if groups[g]['group_name']==group);r['weight']=next(x['weight'] for x in targetmap[t] if x['group']==group)
   for k in ['reason','definition','positive_connection']:
    if k in r:r[k]=by[t].get('reason',by[t].get('appendix_reason'))
   rebuilt.append(r)
  b[s][tn]=(hh,[[W.clean(r.get(k)) for k in hh] for r in rebuilt])
 # True dependent arithmetic runs solely from unchanged accepted Group SPPR.
 R.recalculate(b,p)
 # Mapping cannot change independent classic results or legitimate bounds.
 classic_differences=[]
 old_classic_keyed={tuple([r[0] or '',*r[1:7]]):r for r in old['Classic PPR']['Annual'][1]}
 for z in b['Classic PPR']['Annual'][1]:
  a=old_classic_keyed[tuple([z[0] or '',*z[1:7]])]
  for i,(x,y) in enumerate(zip(a[7:],z[7:]),1950):
   if x!=y:
    assert W.finite(x) and W.finite(y) and math.isclose(x,y,rel_tol=1e-12,abs_tol=1e-6),(u,a[:7],i,x,y)
    classic_differences.append({'key':a[:7],'year':i,'accepted':x,'fresh_sum':y,'difference':y-x})
 b['Classic PPR']=copy.deepcopy(old['Classic PPR'])
 oldclassic=[copy.deepcopy(r) for r in old['PPR–NPP']['Ratios'][1] if not r[0]]
 h,rs=b['PPR–NPP']['Ratios'];b['PPR–NPP']['Ratios']=(h,[*oldclassic,*[r for r in rs if r[0]]]);R.set_result_hash(b)
 R.set_setting(b,'calculation_status',o['calculation_status'])
 protected=['Catch','Classic PPR','Selected model groups','NPP','Diagnostics']
 assert all(b[s]==old[s] for s in ['Catch','Classic PPR','Selected model groups','NPP'])
 assert [r for r in W.records(b,'PPR','Matching') if r['taxon'] not in targetmap]==[r for r in W.records(old,'PPR','Matching') if r['taxon'] not in targetmap]
 updates=block_edits(p,old,b);patch_package(p,updates);post=W.read_book(p);W.validate_region(post,p);assert W.input_hash(post)==W.overview(post)['calculation_input_sha256'];assert R.result_hash(post)==W.overview(post)['calculation_result_sha256']
 for s in ['Catch','Classic PPR','Selected model groups','NPP']:assert W.digest_tables(post[s])==W.digest_tables(old[s]),(u,s)
 writej(ap,audit);writej(lp,ledj)
 # Existing adoption ledgers retain historical expected_old but current adopted rows refresh.
 mp=workev/'mapping_changes.json'
 if mp.exists():
  backup(mp);j=readj(mp)
  for d in j:
   t=d.get('taxon',d.get('key',{}).get('taxon'))
   if t not in targetmap:continue
   for k in ['adopted_new','adopted','new_mapping','after']:
    if k in d:d[k]=targetmap[t] if k=='adopted_new' else copy.deepcopy(record_groups(by[t]))
   if 'reason' in d:d['reason']=by[t].get('reason',by[t].get('appendix_reason'))
   d['current_reporting_scope_followup']='../residual_fish_scope_followup.json' if u=='LME_036' else 'residual_fish_scope_followup.json'
  writej(mp,j)
 # Update only target E/G cells in the seven-column display, with full retained zeros.
 aw=openpyxl.load_workbook(app,data_only=False);aedits=[];heights={};locations=[]
 for s in aw:
  if s.title in ['Allocation evidence','Coverage','Sources']:continue
  for row in s:
   if row[0].value in targetmap and s.max_column>=7:
    t=row[0].value;a=by[t];display=a.get('mapping_display',a.get('display_mapping')) or '; '.join(g.get('group_name',g.get('group',g.get('name')))+f" ({g['weight']*100:.6f}%)" for g in record_groups(a));reason=a.get('appendix_reason',a.get('reason'))
    aedits.extend([(s.title,row[4].coordinate,row[4].value,display),(s.title,row[6].coordinate,row[6].value,reason)]);heights[(s.title,row[0].row)]=max(s.row_dimensions[row[0].row].height or 60,min(360,15*(1+math.ceil(len(reason)/95)),15*(1+math.ceil(len(display)/60))));locations.append({'taxon':t,'sheet':s.title,'row':row[0].row,'range':f'A{row[0].row}:G{row[0].row}'})
 aw.close();assert len(locations)==len(labels),(u,'appendix target rows',locations);patch_package(app,aedits,heights)
 # Refresh affected year details and method exposure using pure read-only helpers.
 ME=module('method_exposure_pure',SHARED/'verify_method_exposure.py');exposure=ME.inspect(u);writej(ev/'residual_fish_method_exposure.json',exposure)
 deltas=[]
 beforeann={tuple(r[:7]):r for r in old['PPR']['Annual'][1]}
 for r in W.records(post,'PPR','Annual'):
  key=tuple(r[k] for k in ['model_id','scope','method','catch_basis','unidentified','metric','status']);prev=beforeann[key];y=int(o['taxon_detail_year']);v=r[y];a=prev[7+y-1950]
  if a!=v:deltas.append({'key':dict(zip(['model_id','scope','method','catch_basis','unidentified','metric','status'],key)),'year':y,'before':a,'after':v,'delta':v-a if W.finite(a) and W.finite(v) else None})
 receipt={'schema_version':1,'unit_id':u,'model_id':mid,'scope':'Four exact residual-fish reporting labels and necessary dependent calculations only','accepted_scientific_parameters':'preserved; no solver, normalization, routing or model alteration','local_retained_reporting_key_bridges':bridge,'primary_reporting_scope_sources':manifest,'local_decisions':changes,'allocation_arithmetic':arithmetic,'before_sha256':initial,'prior_current_workbook_input_sha256':o['calculation_input_sha256'],'prior_current_workbook_result_sha256':o['calculation_result_sha256'],'final_input_sha256':W.overview(post)['calculation_input_sha256'],'final_result_sha256':W.overview(post)['calculation_result_sha256'],'actual_calculation_provenance':{'function':'tools/regional.py recalculate','code_sha256':W.sha(ROOT/'tools/regional.py'),'Group_SPPR_table_sha256_preserved':W.digest_tables(old['Selected model groups'].get('Group SPPR',([],[]))),'accepted_model_sha256_preserved':W.sha(folder/o['model_path']),'full_source_matrices_and_diagnostic_statuses_reused_unchanged':True,'source_and_canonical_bytes_unchanged':'See accepted baseline comparator and follow-up file hashes; current diagnostic returns were not rerun.','Classic_Annual_TL_coefficients_and_independent_bounds_exactly_retained':True,'all_non_target_matching_exactly_preserved':True,'workbook_cells_changed':len(updates),'OOXML_nonworksheet_parts_preserved':True},'protected_baseline_check':AC.compare(u),'affected_reference_year_annual_rows':deltas,'appendix_changed_target_rows':locations,'method_exposure':'residual_fish_method_exposure.json','Office_QA':'pending targeted render/layout/link checks','shared_actions':'pending coordinator shared integration/browser/Git','descendants':[]}
 receipt['accepted_classic_floating_point_reconstruction_differences']=classic_differences
 assert receipt['protected_baseline_check']['all_checked_protected_inputs_preserved'];writej(ev/'residual_fish_scope_followup.json',receipt)
 writej(SCR/(u+'_pending.json'),{'unit':u,'folder':str(folder),'evidence':str(ev),'audit':str(ap),'ledger':str(lp),'report':str(doc),'appendix':str(app),'workbook':str(p),'prior_handoff_sha256':W.sha(backup(handp)),'changed_target_labels':labels})
 print(json.dumps({'unit':u,'mapped_labels_changed':labels,'allocation_arithmetic':arithmetic,'protected_pass':True,'annual_changed_rows':len(deltas),'office_QA':'pending'},ensure_ascii=False),flush=True)
if __name__=='__main__':main(sys.argv[1])
