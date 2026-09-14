"""Final mappings expose the precise scientific weights without losing unresolved taxa."""
import sys
import unittest
import csv
import gzip
import tempfile
from contextlib import closing
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import build_ecosystem_data as ecosystem
import build_model_workbook as model


class FinalMappingSheet(unittest.TestCase):
    def test_build_one_writes_numeric_final_mappings_identical_to_central(self):
        """Exercise the saved model workbook, not only its row/render helpers."""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            unit, stem = 'LME_001', '1_1_Test_model'
            book = root / 'PPREstimation/output/top10' / f'{stem}.xlsx'
            book.parent.mkdir(parents=True)
            source = openpyxl.Workbook()
            source.active.title = 'groups_df'
            source.active.append(['seq', 'group_name', 'biomass', 'catch'])
            source.active.append([1, 'g1', 1, 1])
            source.active.append([2, 'g2', 2, 2])
            for scope in ('all', 'inner', 'PP'):
                ws = source.create_sheet(f'sppr_{scope}')
                ws.append(['seq', 'group_name', 'new_TE_EEfix'])
                ws.append([1, 'g1', 10])
                ws.append([2, 'g2', 20])
            source.save(book)
            source.close()
            catch = root / 'SeaAroundUsExtraction/data/catch_by_taxon_year' / f'{unit}.csv.gz'
            catch.parent.mkdir(parents=True)
            with gzip.open(catch, 'wt', encoding='utf-8', newline='') as stream:
                writer = csv.writer(stream)
                writer.writerow(['taxon', 'year', 'catch_tonnes'])
                writer.writerows([['A', 2019, 1], ['B', 2019, 2],
                                  ['Composite', 2019, 10], ['Unknown', 2019, 5]])
            model.mio.write_mapping_csv(model.mio.mapping_dir(root, unit) / f'{stem}.csv', [
                {'taxon': 'A', 'group': 'g1', 'confidence': 'high', 'evidence': 'explicit_member', 'explanation': 'source table'},
                {'taxon': 'B', 'group': 'g2', 'confidence': 'high'},
                {'taxon': 'Composite', 'group': 'g1 | g2', 'weights': 'catch_composition',
                 'confidence': 'medium', 'evidence': 'composite_split', 'explanation': 'two groups'},
                {'taxon': 'Unknown', 'group': 'Unresolved', 'confidence': 'unresolved',
                 'evidence': 'none', 'explanation': 'insufficient evidence'},
            ])
            with patch.object(model, 'ROOT', root), patch.object(ecosystem, 'ROOT', root):
                result, error = model.build_one(unit, book, {})
                self.assertIsNone(error)
                with closing(openpyxl.load_workbook(root / result['path'])) as built:
                    self.assertIn('Final mappings', built.sheetnames)
                    model_rows = list(built['Final mappings'].values)
                central = openpyxl.Workbook()
                ecosystem.sheet_final_mappings(central, ecosystem.load_final_mappings(unit))
                saved = BytesIO()
                central.save(saved)
                central.close()
                saved.seek(0)
                with closing(openpyxl.load_workbook(saved)) as reloaded:
                    self.assertEqual(model_rows, list(reloaded['Final mappings'].values))
            data = [dict(zip(model_rows[3], r)) for r in model_rows[4:]]
            self.assertEqual({r['model_id'] for r in data}, {stem})
            self.assertEqual({r['taxon'] for r in data}, {'A', 'B', 'Composite', 'Unknown'})
            split = [r['weight'] for r in data if r['taxon'] == 'Composite']
            self.assertTrue(all(isinstance(w, (int, float)) for w in split))
            self.assertAlmostEqual(split[0], 1 / 3, places=15)
            self.assertAlmostEqual(split[1], 2 / 3, places=15)
            self.assertNotEqual(split, [0.333, 0.667])
            self.assertIsNone(next(r['weight'] for r in data if r['taxon'] == 'Unknown'))

    def test_full_precision_numeric_weights_all_taxa_and_distinct_models(self):
        taxa = {t: {'common_name': '', 'functional_group': '', 'commercial_group': ''}
                for t in ['A', 'B', 'Composite', 'Unknown']}
        rows = [
            {'taxon': 'A', 'group': 'g1', 'confidence': 'high', 'evidence': 'explicit_member', 'explanation': 'paper table'},
            {'taxon': 'B', 'group': 'g2', 'confidence': 'high'},
            {'taxon': 'Composite', 'group': 'g1 | g2', 'weights': 'catch_composition',
             'confidence': 'medium', 'evidence': 'composite_split', 'explanation': 'multiple groups'},
            {'taxon': 'Unknown', 'group': 'Unresolved', 'confidence': 'unresolved',
             'evidence': 'none', 'explanation': 'insufficient evidence'},
        ]
        totals = {'A': 1, 'B': 2, 'Composite': 10, 'Unknown': 5}
        resolved = model.build_taxon_sppr(rows, ['method'], {'g1': [10], 'g2': [20]},
                                          {'g1': {}, 'g2': {}}, totals)
        records = model.final_mapping_rows('LME_001', 'model1', taxa, resolved, totals)
        records += model.final_mapping_rows('LME_001', 'model2', taxa, resolved, totals)
        wb = openpyxl.Workbook()
        ecosystem.sheet_final_mappings(wb, records)
        values = list(wb['Final mappings'].values)
        header = next(r for r in values if r[0] == 'unit_id')
        data = [dict(zip(header, r)) for r in values if r[0] == 'LME_001']
        self.assertEqual(len(data), 10)
        for stem in ('model1', 'model2'):
            selected = [r for r in data if r['model_id'] == stem]
            self.assertEqual({r['taxon'] for r in selected}, set(taxa))
            split = [r for r in selected if r['taxon'] == 'Composite']
            self.assertEqual([r['weight'] for r in split], resolved['Composite']['weights'])
            self.assertTrue(all(isinstance(r['weight'], float) for r in split))
            self.assertEqual(sum(r['weight'] for r in split), 1)
            self.assertEqual(split[0]['weight_basis'], 'catch_composition')
            self.assertEqual(split[0]['explanation'], 'multiple groups')
            self.assertEqual(split[0]['evidence'], 'composite_split')
            single = next(r for r in selected if r['taxon'] == 'A')
            self.assertEqual(single['weight'], 1.0)
            unresolved = next(r for r in selected if r['taxon'] == 'Unknown')
            self.assertEqual(unresolved['group'], 'Unresolved')
            self.assertIsNone(unresolved['weight'])

    def test_no_mapping_has_explanatory_sheet_without_invented_rows(self):
        wb = openpyxl.Workbook()
        ecosystem.sheet_final_mappings(wb, [])
        self.assertIn('No final mappings', wb['Final mappings']['A1'].value)

    def test_central_loader_matches_actual_model_resolution_exactly(self):
        unit = 'LME_034'
        books = [p for p in model.mio.model_workbooks(ecosystem.ROOT, unit)
                 if (model.mio.mapping_dir(ecosystem.ROOT, unit) / f'{p.stem}.csv').exists()]
        if not books:
            self.skipTest('repository model inputs unavailable')
        actual = ecosystem.load_final_mappings(unit)
        taxa, _ = model.mio.read_catch(ecosystem.ROOT, unit)
        totals = {t: sum(round(v, model.CATCH_DP) for v in e['by_year'].values())
                  for t, e in taxa.items()}
        expected = []
        for book in books:
            rows = model.mio.read_mapping_csv(model.mio.mapping_dir(ecosystem.ROOT, unit) / f'{book.stem}.csv')
            groups = {g['group_name']: g for g in model.mio.read_groups(book)}
            methods, sppr = model.mio.read_methods(book)
            resolved = model.build_taxon_sppr(rows, methods, sppr, groups, totals)
            expected.extend(model.final_mapping_rows(unit, book.stem, taxa, resolved, totals))
        self.assertEqual(actual, expected)
        self.assertEqual({r['model_id'] for r in actual}, {p.stem for p in books})
        self.assertEqual({r['taxon'] for r in actual}, set(taxa))


if __name__ == '__main__':
    unittest.main()
