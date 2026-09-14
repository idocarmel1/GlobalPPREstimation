"""A failed workbook write must not truncate the last usable model output."""
import errno
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import build_model_workbook as model
import build_ecosystem_data as ecosystem


class AtomicModelWorkbookSave(unittest.TestCase):
    def test_ecosystem_build_avoids_existing_destination_truncation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / 'EEZ_024' / 'EEZ_024.xlsx'
            output.parent.mkdir()
            output.write_bytes(b'previous workbook')
            real_open = io.open

            def reject_destination_truncation(file, mode='r', *args, **kwargs):
                if isinstance(file, (str, os.PathLike)) and Path(file) == output and 'w' in mode:
                    raise OSError(errno.EINVAL, 'Invalid argument', str(output))
                return real_open(file, mode, *args, **kwargs)

            with patch.multiple(ecosystem, OUT=root, ARCHIVE=root / 'archive',
                                SPPR_DIR=root / 'sppr', MODEL_JSON_DIR=root / 'models',
                                load_catch=lambda unit: ([], []),
                                load_trophic_levels=lambda unit: {},
                                load_sppr_methods=lambda unit: [],
                                load_final_mappings=lambda unit: []), \
                    patch('io.open', side_effect=reject_destination_truncation):
                ecosystem.build_unit('EEZ_024', {}, {}, force=True)
            actual = openpyxl.load_workbook(output)
            try:
                self.assertIn('Catch', actual.sheetnames)
                self.assertIn('NPP', actual.sheetnames)
            finally:
                actual.close()

    def test_save_avoids_truncating_the_destination_path(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / '47_2_East_China_Sea_(2018).xlsx'
            output.write_bytes(b'previous workbook')
            workbook = openpyxl.Workbook()
            self.addCleanup(workbook.close)
            workbook.active['A1'] = 'new workbook'
            real_open = io.open

            def reject_destination_truncation(file, mode='r', *args, **kwargs):
                if isinstance(file, (str, os.PathLike)) and Path(file) == output and 'w' in mode:
                    raise OSError(errno.EINVAL, 'Invalid argument', str(output))
                return real_open(file, mode, *args, **kwargs)

            with patch('io.open', side_effect=reject_destination_truncation):
                with self.assertRaises(OSError):
                    workbook.save(output)
                model.save_workbook_atomic(workbook, output)
            actual = openpyxl.load_workbook(output)
            try:
                self.assertEqual(actual.active['A1'].value, 'new workbook')
            finally:
                actual.close()
            self.assertEqual(list(Path(directory).iterdir()), [output])

    def test_partial_serialization_failure_preserves_original_and_removes_temp(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'model.xlsx'
            original = b'last usable workbook bytes'
            output.write_bytes(original)
            workbook = openpyxl.Workbook()
            self.addCleanup(workbook.close)

            def fail_after_partial_write(stream):
                stream.write(b'incomplete new archive')
                raise OSError('simulated serialization failure')

            with patch.object(workbook, 'save', side_effect=fail_after_partial_write):
                with self.assertRaisesRegex(OSError, 'simulated serialization failure'):
                    model.save_workbook_atomic(workbook, output)
            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(list(Path(directory).iterdir()), [output])

    def test_replacement_failure_preserves_original_and_removes_temp(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'model.xlsx'
            original = b'last usable workbook bytes'
            output.write_bytes(original)
            workbook = openpyxl.Workbook()
            self.addCleanup(workbook.close)
            with patch('os.replace', side_effect=PermissionError('destination locked')):
                with self.assertRaisesRegex(PermissionError, 'destination locked'):
                    model.save_workbook_atomic(workbook, output)
            self.assertEqual(output.read_bytes(), original)
            self.assertEqual(list(Path(directory).iterdir()), [output])


if __name__ == '__main__':
    unittest.main()
