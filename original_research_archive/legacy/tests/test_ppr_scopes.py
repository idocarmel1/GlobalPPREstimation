import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
from ppr_scopes import read_scopes, read_health, add_group_scope_sheets, method_health_flags


def test_scopes_preserve_missing_and_group_identity(tmp_path):
    path = tmp_path / 'upstream.xlsx'
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for scope, val in [('all', 12), ('inner', 8), ('PP', None)]:
        s = wb.create_sheet('sppr_' + scope)
        s.append(['seq', 'group_name', 'new_GE'])
        s.append([7, 'Fish', val])
    s = wb.create_sheet('model_health')
    s.append(['TE_option', 'status', 'divergence_b', 'divergence_rho_living'])
    s.append(['GE', 'FAIL', 1.2, .3])
    wb.save(path)
    scopes = read_scopes(path)
    assert scopes['inner'] == (['new_GE'], {'Fish': [8.0]})
    assert scopes['PP'][1]['Fish'] == [None]
    assert read_health(path)['GE']['b'] == 1.2
    assert read_health(path)['GE']['status'] == 'FAIL'
    output = openpyxl.Workbook()
    add_group_scope_sheets(output, [path])
    assert output['sppr_PP'].cell(5, 3).value is None
    assert 'Recycling' in output.sheetnames


def test_absent_scope_is_not_substituted(tmp_path):
    path = tmp_path / 'old.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'sppr_all'
    wb.active.append(['seq', 'group_name', 'method'])
    wb.active.append([1, 'Fish', 99])
    wb.save(path)
    assert read_scopes(path)['PP'] == ([], {})


def test_unfished_negative_groups_and_exact_configuration_health_are_not_hidden():
    flags = method_health_flags(['new_GE', 'new_TE_EEfix', 'MC_new_GE', 'other'],
                                {'Unfished': [1, 2, 3, -4]}, {'TE': {'status': 'FAIL'}})
    assert 'new_TE_EEfix' in flags
    assert 'other' in flags
    assert 'new_GE' not in flags
    assert 'MC_new_GE' not in flags


def test_partial_ecosystem_refresh_preserves_other_index_rows(tmp_path, monkeypatch):
    import csv
    import build_ecosystem_data as builder
    index = tmp_path / 'INDEX.csv'
    columns = ['unit_id','region_name','region_type','catch_from','catch_to','taxa','taxa_with_tl',
               'articles','ecopath_models','models_mapped','model_workbooks','has_npp']
    with index.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerow(dict(zip(columns, ['LME_002','Keep me','LME',1950,2019,5,5,0,0,0,0,0])))
    (tmp_path / 'LME_001.csv.gz').touch()
    monkeypatch.setattr(builder, 'OUT', tmp_path)
    monkeypatch.setattr(builder, 'CATCH_DIR', tmp_path)
    monkeypatch.setattr(builder, 'load_atlas_regions', lambda: {})
    monkeypatch.setattr(builder, 'load_npp', lambda: {})
    monkeypatch.setattr(builder, 'build_unit', lambda *a, **k: {
        'region_name':'Changed','region_type':'LME','article_count':1,
        'coverage':{'catch_years':[1950,2019],'taxa':8,'taxa_with_trophic_level':7,
                    'ecopath_models':1,'models_mapped':1,'model_workbooks':1,'has_npp':True}})
    monkeypatch.setattr(sys, 'argv', ['build', '--units', 'LME_001'])
    assert builder.main() == 0
    with index.open(encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    assert [(r['unit_id'], r['region_name']) for r in rows] == [('LME_001','Changed'),('LME_002','Keep me')]


def test_atlas_selection_does_not_promote_archived_candidates():
    import json
    p = Path(__file__).resolve().parents[1] / 'data/atlas_selection.json'
    units = json.loads(p.read_text(encoding='utf-8'))['units']
    assert set(units) == {'HS_077','LME_013','LME_027','LME_028','LME_032',
                          'LME_034','LME_035','LME_036','LME_047','LME_052'}
    assert 'LME_022' not in units
