#!/usr/bin/env python3
"""Generate the interactive Ecopath coverage map using only the summary Excel."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def rows_as_dicts(ws):
    headers = [c.value for c in ws[4]]
    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not any(v is not None for v in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    return out


def reconstruct_geometry(ws, id_field):
    pieces = defaultdict(list)
    for row in rows_as_dicts(ws):
        pieces[str(row[id_field])].append((int(row["chunk_index"]), str(row["geometry_chunk"])))
    return {key: json.loads("".join(v for _, v in sorted(vals))) for key, vals in pieces.items()}


def read_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=False)
    region_geoms = reconstruct_geometry(wb["Region Geometry"], "unit_id")
    article_geoms = reconstruct_geometry(wb["Article Geometry"], "article_id")
    regions = rows_as_dicts(wb["Regions"])
    articles = rows_as_dicts(wb["Articles"])
    files = rows_as_dicts(wb["Files"])

    for r in regions:
        r["selected_all84"] = str(r["selected_all84"]).lower() == "yes"
        r["selected_top25"] = str(r["selected_top25"]).lower() == "yes"
        r["ppr_rank"] = int(r["ppr_rank"])
        for k in ["ppr_species_2019", "global_ppr_share", "cumulative_ppr_share", "best_quality_score", "marker_lat", "marker_lon"]:
            r[k] = float(r[k] or 0)
        r["article_count"] = int(r["article_count"] or 0)
        r["geometry"] = region_geoms[str(r["unit_id"])]
    for a in articles:
        for k in ["region_rank", "publication_year", "legacy_documentation_score_60", "legacy_applicability_score_40", "model_loadability_score_55", "documentation_score_20", "spatial_fit_score_15", "recency_validation_score_10", "quality_score_cap", "quality_score_100", "downloaded_file_count", "download_attempt_count"]:
            a[k] = int(a[k] or 0)
        a["target_coverage_ratio"] = float(a["target_coverage_ratio"] or 0)
        a["geometry"] = article_geoms[str(a["article_id"])]
    return {"regions": regions, "articles": articles, "files": files,
            "source_workbook": path.name, "generated_on": "2026-09-04"}


HTML = r'''<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>Global Ecopath Coverage · All 84 PPR Units</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@600;700;800&display=swap" rel="stylesheet">
  <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/leaflet@1.9.4/dist/leaflet.css">
  <style>
    :root{--navy:#09273b;--ink:#14354a;--muted:#647b8b;--line:#dce8ee;--panel:#fff;--ocean:#dceff4;--teal:#0b7e88;--aqua:#4dc3c9;--shadow:0 12px 34px rgba(16,45,63,.13)}
    *{box-sizing:border-box} html,body{height:100%;margin:0;font-family:"DM Sans",sans-serif;color:var(--ink);background:#eef5f7} button,input,select{font:inherit}
    .app{height:100%;display:grid;grid-template-rows:auto 1fr;overflow:hidden}
    header{background:linear-gradient(120deg,#082b42,#0c6570);color:#fff;padding:16px 22px 14px;display:flex;align-items:center;gap:22px;box-shadow:0 3px 18px rgba(4,34,52,.22);z-index:1002}
    .brand{min-width:345px}.eyebrow{text-transform:uppercase;letter-spacing:.14em;font-size:10px;font-weight:700;color:#7de0df}.brand h1{font-family:Manrope,sans-serif;font-size:22px;line-height:1.1;margin:4px 0 2px}.brand p{margin:0;color:#c9e8ec;font-size:12px}
    .kpis{display:flex;gap:9px;flex:1;justify-content:flex-end;flex-wrap:wrap}.kpi{background:rgba(255,255,255,.11);border:1px solid rgba(255,255,255,.17);border-radius:10px;padding:8px 12px;min-width:105px}.kpi b{display:block;font:700 17px Manrope}.kpi span{font-size:10px;color:#cbe6e9;text-transform:uppercase;letter-spacing:.07em}
    .workspace{display:grid;grid-template-columns:320px minmax(420px,1fr) 390px;min-height:0}
    .sidebar,.details{background:var(--panel);overflow:auto;z-index:1001}.sidebar{border-right:1px solid var(--line);padding:17px 16px}.details{border-left:1px solid var(--line);padding:18px}
    #map{height:100%;min-height:550px;background:var(--ocean)}
    h2,h3{font-family:Manrope;margin:0}.section-title{font-size:11px;color:var(--muted);font-weight:800;letter-spacing:.1em;text-transform:uppercase;margin:18px 0 8px}.section-title:first-child{margin-top:0}
    .search{position:relative}.search input{width:100%;padding:10px 12px 10px 34px;border:1px solid var(--line);border-radius:9px;background:#f7fbfc;outline:none}.search:before{content:'⌕';position:absolute;left:12px;top:8px;color:#69808f;font-size:20px}.search input:focus{border-color:#52aeb5;box-shadow:0 0 0 3px rgba(62,169,177,.13)}
    .grid2{display:grid;grid-template-columns:1fr 1fr;gap:8px}.field label{font-size:11px;font-weight:700;color:var(--muted);display:block;margin-bottom:4px}.field select{width:100%;border:1px solid var(--line);border-radius:8px;padding:8px;background:#fff}
    .range-row{display:flex;align-items:center;gap:10px}.range-row input{flex:1;accent-color:var(--teal)}.range-val{font:700 12px Manrope;background:#e8f5f6;color:#0b6b73;padding:4px 7px;border-radius:6px;min-width:34px;text-align:center}
    .checks{display:grid;gap:7px}.check{display:flex;align-items:center;gap:8px;font-size:12px}.check input{accent-color:var(--teal)}
    .region-list{display:grid;gap:5px}.region-btn{display:grid;grid-template-columns:27px 1fr auto;align-items:center;gap:7px;border:1px solid transparent;background:#f7fafb;border-radius:8px;padding:7px 8px;text-align:left;cursor:pointer;color:var(--ink)}.region-btn:hover,.region-btn.active{border-color:#72bdc3;background:#e9f5f6}.rank{height:22px;width:22px;display:grid;place-items:center;border-radius:50%;background:#dcecf0;font:700 10px Manrope}.region-btn b{font-size:12px}.region-btn small{font-size:10px;color:var(--muted)}.qdot{width:9px;height:9px;border-radius:50%}
    .legend-card{background:#f7fafb;border:1px solid var(--line);border-radius:10px;padding:10px}.legend-labels{display:flex;justify-content:space-between;font-size:10px;color:var(--muted);margin-top:3px}.grad{height:9px;border-radius:6px}.ppr-grad{background:linear-gradient(90deg,#d73027,#f46d43,#fee08b,#a6d96a,#1a9850)}.q-grad{background:linear-gradient(90deg,#d94b4b,#f08a4b,#f3c64e,#1f9d68)}
    .ppr-marker{background:transparent;border:0}.ppr-marker span{display:grid;place-items:center;width:25px;height:25px;border-radius:50%;border:2px solid rgba(255,255,255,.96);box-shadow:0 2px 8px rgba(15,43,58,.42);font:800 10px Manrope;color:#1e2830;text-align:center}.ppr-marker span.hot{color:white;text-shadow:0 1px 2px rgba(0,0,0,.5)}
    .map-note{position:absolute;z-index:800;bottom:18px;left:50%;transform:translateX(-50%);background:rgba(255,255,255,.94);box-shadow:var(--shadow);border:1px solid #d9e6eb;border-radius:10px;padding:8px 12px;font-size:11px;pointer-events:none;white-space:nowrap}
    .details-empty{height:100%;display:grid;place-items:center;text-align:center;color:var(--muted)}.details-empty .icon{font-size:42px;margin-bottom:8px}.details-head{padding-bottom:14px;border-bottom:1px solid var(--line)}.details-head .code{font:700 11px Manrope;color:var(--teal);letter-spacing:.08em}.details-head h2{font-size:22px;margin:3px 0 8px}.chips{display:flex;gap:6px;flex-wrap:wrap}.chip{font-size:10px;font-weight:700;background:#eef5f7;border-radius:99px;padding:5px 8px;color:#486778}.chip.primary{background:#dff2f2;color:#07636b}
    .metric-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:7px;margin:12px 0}.metric{background:#f7fafb;border-radius:9px;padding:9px}.metric b{display:block;font:700 15px Manrope}.metric span{font-size:9px;color:var(--muted);text-transform:uppercase;letter-spacing:.05em}
    .article-card{border:1px solid var(--line);border-left:5px solid var(--q);border-radius:11px;padding:11px 12px;margin:9px 0;background:#fff;box-shadow:0 4px 13px rgba(17,53,72,.05)}.article-top{display:flex;justify-content:space-between;gap:8px}.article-card h3{font-size:13px;line-height:1.35;margin:2px 0 4px}.article-card .authors{font-size:11px;color:var(--muted)}.score{height:39px;min-width:42px;border-radius:9px;background:var(--q);color:#fff;display:grid;place-items:center;font:800 14px Manrope;text-shadow:0 1px 2px rgba(0,0,0,.2)}.bar{height:6px;background:#edf2f4;border-radius:5px;overflow:hidden;margin:8px 0}.bar i{display:block;height:100%;background:var(--q)}.meta{display:grid;grid-template-columns:1fr 1fr;gap:5px 10px;font-size:10px;margin:7px 0}.meta b{display:block;color:#24495f}.meta span{color:var(--muted)}.rec{font-size:11px;line-height:1.45;background:#f6fafb;padding:8px;border-radius:7px}.actions{display:flex;gap:6px;margin-top:8px}.actions a,.actions button{font-size:10px;font-weight:700;color:#0b6b73;background:#e7f4f5;border:0;border-radius:6px;padding:6px 8px;text-decoration:none;cursor:pointer}.actions button{color:#496a7c;background:#eef3f5}
    .status-warn{background:#fff4d9;color:#7e5a08}.leaflet-popup-content-wrapper{border-radius:10px}.leaflet-popup-content{font-family:"DM Sans";min-width:230px}.leaflet-control-layers{font-family:"DM Sans"}
    .footer-note{font-size:10px;color:var(--muted);line-height:1.5;margin-top:14px;padding-top:12px;border-top:1px solid var(--line)}
    @media(max-width:1100px){.workspace{grid-template-columns:280px 1fr}.details{position:absolute;right:0;top:84px;bottom:0;width:380px;transform:translateX(100%);transition:.2s;box-shadow:var(--shadow)}.details.open{transform:none}.kpi:nth-last-child(-n+2){display:none}}
    @media(max-width:720px){header{padding:12px}.brand{min-width:0}.brand h1{font-size:17px}.brand p,.kpis{display:none}.workspace{grid-template-columns:1fr}.sidebar{position:absolute;left:0;top:67px;bottom:0;width:290px;transform:translateX(-100%);transition:.2s;box-shadow:var(--shadow)}.sidebar.open{transform:none}.details{top:67px;width:100%}.map-note{white-space:normal;width:80%;text-align:center}.leaflet-control-zoom{margin-top:70px!important}}
  </style>
</head>
<body>
<div class="app">
  <header>
    <div class="brand"><div class="eyebrow">Sea Around Us × Ecopath</div><h1>Global ecosystem-model coverage</h1><p>All 84 LME and high-seas production units · curated for extraction readiness</p></div>
    <div class="kpis"><div class="kpi"><b id="kpiPpr">100%</b><span>Global PPR</span></div><div class="kpi"><b id="kpiArticles">71</b><span>Article packages</span></div><div class="kpi"><b id="kpiRegions">60 / 84</b><span>Units with models</span></div><div class="kpi"><b id="kpiFiles">14</b><span>Verified files</span></div></div>
  </header>
  <main class="workspace">
    <aside class="sidebar" id="sidebar">
      <div class="section-title">Find an ecosystem or article</div><div class="search"><input id="search" placeholder="Search names, authors, titles…"></div>
      <div class="section-title">Filters</div>
      <div class="grid2"><div class="field"><label>Region type</label><select id="typeFilter"><option value="all">LME + High Seas</option><option value="LME">LME only</option><option value="High Seas">High Seas only</option></select></div><div class="field"><label>Rank scope</label><select id="rankFilter"><option value="84" selected>All 84</option><option value="25">Top 25</option><option value="10">Pilot top 10</option></select></div></div>
      <div class="section-title">Downloaded source files</div><div class="field"><label>Area file availability</label><select id="downloadFilter"><option value="all">All areas</option><option value="with">Has verified download</option><option value="without">No verified download</option></select></div>
      <div class="section-title">Minimum article quality</div><div class="range-row"><input type="range" id="qualityFilter" min="0" max="100" value="0"><span class="range-val" id="qualityVal">0</span></div>
      <div class="section-title">Map layers</div><div class="checks"><label class="check"><input id="showRegions" type="checkbox" checked> Original polygons + PPR-rank markers</label><label class="check"><input id="colorRegionsByPpr" type="checkbox"> Color region polygons by PPR</label><label class="check"><input id="showArticles" type="checkbox" checked> Recommended article areas</label></div>
      <div class="section-title">Cumulative global PPR — ranked circles</div><div class="legend-card"><div class="grad ppr-grad"></div><div class="legend-labels"><span>High PPR · low cumulative %</span><span>Low PPR · 100%</span></div></div>
      <div class="section-title">Article quality — area fill</div><div class="legend-card"><div class="grad q-grad"></div><div class="legend-labels"><span>0 · weak</span><span>50</span><span>100 · strong</span></div></div>
      <div class="section-title">PPR-ranked units</div><div id="regionList" class="region-list"></div>
    </aside>
    <section style="position:relative;min-width:0"><div id="map"></div><div class="map-note">Click a region to compare its original boundary with recommended article coverage</div></section>
    <aside class="details" id="details"><div class="details-empty"><div><div class="icon">◉</div><h3>Select a marine region</h3><p>Its PPR rank, original boundary, recommended articles, coverage, quality, files and Zotero links will appear here.</p></div></div></aside>
  </main>
</div>
<script src="https://cdn.jsdelivr.net/npm/leaflet@1.9.4/dist/leaflet.js"></script>
<script>
const DB=__DATA__;
const regions=DB.regions, articles=DB.articles, files=DB.files;
const byUnit=Object.groupBy?Object.groupBy(articles,a=>a.unit_id):articles.reduce((o,a)=>((o[a.unit_id]??=[]).push(a),o),{});
const regionById=Object.fromEntries(regions.map(r=>[r.unit_id,r]));
const map=L.map('map',{worldCopyJump:true,zoomControl:true,minZoom:2}).setView([12,10],2);
L.tileLayer('https://{s}.tile.openstreetmap.org/{z}/{x}/{y}.png',{maxZoom:10,opacity:.55,attribution:'© OpenStreetMap · Sea Around Us polygons'}).addTo(map);
const regionGroup=L.layerGroup().addTo(map), markerGroup=L.layerGroup().addTo(map), articleGroup=L.layerGroup().addTo(map);
const regionLayers={}, regionMarkers={}, articleLayers={}; let selectedId=null;

function lerp(a,b,t){return Math.round(a+(b-a)*t)} function hex(c){return [parseInt(c.slice(1,3),16),parseInt(c.slice(3,5),16),parseInt(c.slice(5,7),16)]} function rgb(x){return '#'+x.map(v=>v.toString(16).padStart(2,'0')).join('')}
function mix(c1,c2,t){let a=hex(c1),b=hex(c2);return rgb(a.map((v,i)=>lerp(v,b[i],t)))}
function pprNorm(r){return Math.max(0,Math.min(1,Number(r.cumulative_ppr_share)||0))}
const pprRamp=['#d73027','#f46d43','#fee08b','#a6d96a','#1a9850'];
function pprColor(r){const t=pprNorm(r),x=t*(pprRamp.length-1),i=Math.min(pprRamp.length-2,Math.floor(x));return mix(pprRamp[i],pprRamp[i+1],x-i)}
function qualityColor(q){if(q<40)return mix('#d94b4b','#f08a4b',q/40);if(q<70)return mix('#f08a4b','#f3c64e',(q-40)/30);return mix('#f3c64e','#1f9d68',(q-70)/30)}
const pct=v=>`${(Number(v)*100).toFixed(v<.01?1:0)}%`, num=v=>Number(v).toLocaleString(undefined,{maximumFractionDigits:0});
function regionFeature(r){return {type:'Feature',properties:r,geometry:r.geometry}} function articleFeature(a){return {type:'Feature',properties:a,geometry:a.geometry}}
function baseRegionStyle(r){const active=r.unit_id===selectedId,pprFill=document.querySelector('#colorRegionsByPpr')?.checked;return{color:active?'#f5b942':'#2b86a6',weight:active?3:1.15,fillColor:pprFill?pprColor(r):'#8ecae6',fillOpacity:active?(pprFill ? .46 : .30):(pprFill ? .34 : .18),opacity:.9}}
function addRegion(r){
  const tooltip=`<b>#${r.ppr_rank} ${r.region_name}</b><br>${r.unit_id} · ${pct(r.global_ppr_share)} unit share · ${pct(r.cumulative_ppr_share)} cumulative PPR<br>${r.article_count} recommended article${r.article_count===1?'':'s'}`;
  const lyr=L.geoJSON(regionFeature(r),{style:()=>baseRegionStyle(r),onEachFeature:(_,l)=>{l.on('click',()=>selectRegion(r.unit_id));l.bindTooltip(tooltip,{sticky:true})}});regionLayers[r.unit_id]=lyr;lyr.addTo(regionGroup);
  const n=String(r.ppr_rank), hot=(pprNorm(r)<.24||pprNorm(r)>.76)?' hot':'';
  const marker=L.marker([r.marker_lat,r.marker_lon],{icon:L.divIcon({className:'ppr-marker',html:`<span class="${hot}" style="background:${pprColor(r)}">${n}</span>`,iconSize:[25,25],iconAnchor:[12,12]})});
  marker.on('click',()=>selectRegion(r.unit_id)); marker.bindTooltip(tooltip,{direction:'top',offset:[0,-10]}); regionMarkers[r.unit_id]=marker; marker.addTo(markerGroup)
}
function addArticle(a){const q=a.quality_score_100,c=qualityColor(q);const lyr=L.geoJSON(articleFeature(a),{style:{color:c,weight:2,fillColor:c,fillOpacity:.46,dashArray:a.geometry_confidence==='high'?null:'6 4'},onEachFeature:(_,l)=>{l.on('click',e=>{L.DomEvent.stopPropagation(e);selectRegion(a.unit_id);highlightArticle(a.article_id)});l.bindTooltip(`<b>${a.authors} (${a.publication_year})</b><br>Quality ${q}/100 · est. ${pct(a.target_coverage_ratio)} of target<br>${a.geometry_confidence} geometry confidence`,{sticky:true})}});articleLayers[a.article_id]=lyr;lyr.addTo(articleGroup)}
regions.forEach(addRegion);articles.forEach(addArticle);

function regionHasDownload(r){return (byUnit[r.unit_id]||[]).some(a=>Number(a.downloaded_file_count)>0)}
function passRegion(r){const rank=+document.querySelector('#rankFilter').value,type=document.querySelector('#typeFilter').value,download=document.querySelector('#downloadFilter').value,q=document.querySelector('#search').value.toLowerCase().trim(),hasDownload=regionHasDownload(r);return r.ppr_rank<=rank&&(type==='all'||r.region_type===type)&&(download==='all'||(download==='with'&&hasDownload)||(download==='without'&&!hasDownload))&&(!q||`${r.region_name} ${r.unit_id}`.toLowerCase().includes(q)||(byUnit[r.unit_id]||[]).some(a=>`${a.title} ${a.authors}`.toLowerCase().includes(q)))}
function passArticle(a){const minQ=+document.querySelector('#qualityFilter').value,q=document.querySelector('#search').value.toLowerCase().trim();return a.quality_score_100>=minQ&&(!q||`${a.title} ${a.authors} ${a.region_name}`.toLowerCase().includes(q))}
function refresh(){
  regionGroup.clearLayers(); markerGroup.clearLayers(); articleGroup.clearLayers();
  if(document.querySelector('#showRegions').checked)regions.filter(passRegion).forEach(r=>{regionLayers[r.unit_id].eachLayer(l=>l.setStyle(baseRegionStyle(r)));regionLayers[r.unit_id].addTo(regionGroup);regionMarkers[r.unit_id].addTo(markerGroup)});
  if(document.querySelector('#showArticles').checked)articles.filter(a=>passArticle(a)&&passRegion(regionById[a.unit_id])).forEach(a=>articleLayers[a.article_id].addTo(articleGroup));
  renderList(); if(selectedId&&(!regionById[selectedId]||!passRegion(regionById[selectedId])))selectedId=null;
}
function renderList(){const el=document.querySelector('#regionList');el.innerHTML='';regions.filter(passRegion).forEach(r=>{const t=pprNorm(r),b=document.createElement('button');b.className='region-btn'+(selectedId===r.unit_id?' active':'');b.innerHTML=`<span class="rank" style="background:${pprColor(r)};color:${t<.24||t>.76?'white':'#1e2830'}">${r.ppr_rank}</span><span><b>${r.region_name}</b><br><small>#${r.ppr_rank} · ${pct(r.cumulative_ppr_share)} cumulative PPR${regionHasDownload(r)?' · file ✓':''}</small></span><span class="qdot" style="background:${r.best_quality_score?qualityColor(r.best_quality_score):'#b7c4ca'}"></span>`;b.onclick=()=>selectRegion(r.unit_id);el.appendChild(b)})}
function selectRegion(id){selectedId=id;const r=regionById[id],lyr=regionLayers[id];Object.values(regionLayers).forEach((x)=>x.eachLayer(l=>l.setStyle(baseRegionStyle(l.feature.properties))));if(lyr){lyr.eachLayer(l=>l.setStyle(baseRegionStyle(r)));try{map.fitBounds(lyr.getBounds(),{padding:[35,35],maxZoom:6})}catch(e){}}renderDetails(r);renderList();document.querySelector('#details').classList.add('open')}
function highlightArticle(id){Object.entries(articleLayers).forEach(([k,x])=>x.eachLayer(l=>l.setStyle({weight:k===id?4:2,fillOpacity:k===id?.62:.46})));const card=document.querySelector(`[data-article="${id}"]`);card?.scrollIntoView({behavior:'smooth',block:'center'})}
function renderDetails(r){const list=(byUnit[r.unit_id]||[]).sort((a,b)=>b.quality_score_100-a.quality_score_100);const html=list.length?list.map((a,i)=>{const c=qualityColor(a.quality_score_100);return `<article class="article-card" data-article="${a.article_id}" style="--q:${c}"><div class="article-top"><div><div class="chip primary">${i===0?'Primary recommendation':'Complementary coverage'}</div><h3>${a.title}</h3><div class="authors">${a.authors} · ${a.publication_year}</div></div><div class="score">${a.quality_score_100}</div></div><div class="bar"><i style="width:${a.quality_score_100}%"></i></div><div class="meta"><div><b>${a.model_loadability_score_55}/55</b><span>Model loadability</span></div><div><b>${a.documentation_score_20}/20</b><span>Documentation</span></div><div><b>${a.spatial_fit_score_15}/15</b><span>Spatial fit</span></div><div><b>${a.recency_validation_score_10}/10</b><span>Recency + validation</span></div><div><b>${a.download_status}</b><span>Material status</span></div><div><b>${a.downloaded_file_count}</b><span>Verified files</span></div></div><div class="rec"><b>${a.loadability_class}</b><br>${a.quality_rationale}</div><div class="rec" style="margin-top:6px">${a.recommendation}</div>${a.download_failure_reason?`<div class="rec" style="margin-top:6px;border-left-color:#c84a3a;background:#fff3ef"><b>File flag:</b> ${a.download_failure_reason}</div>`:''}<div class="actions"><a href="${a.landing_page}" target="_blank" rel="noopener">Open source ↗</a>${a.zotero_key?`<button title="Zotero item key">Zotero ${a.zotero_key}</button>`:''}<button onclick="highlightArticle('${a.article_id}')">Show area</button></div></article>`}).join(''):`<div class="article-card" style="--q:#b66b32"><h3>No balanced Ecopath model located</h3><p class="authors">Web and Zotero searches found background ecosystem sources but no defensible extractable model. This is a documented gap, not evidence that no unpublished model exists.</p></div>`;
  document.querySelector('#details').innerHTML=`<div class="details-head"><div class="code">#${r.ppr_rank} · ${r.unit_id}</div><h2>${r.region_name}</h2><div class="chips"><span class="chip primary">${r.region_type}</span><span class="chip">${r.article_count} recommended source${r.article_count===1?'':'s'}</span><span class="chip">${r.search_status}</span></div></div><div class="metric-grid"><div class="metric"><b>${num(r.ppr_species_2019)}</b><span>PPR 2019</span></div><div class="metric"><b>${pct(r.global_ppr_share)}</b><span>Global share</span></div><div class="metric"><b>${r.best_quality_score||'—'}</b><span>Best quality</span></div></div><div class="section-title">Recommended model sources</div>${html}<div class="footer-note"><b>Boundary rule.</b> The original region polygon is always the supplied Sea Around Us geometry. Article areas use reported/inferred bounds or, for explicitly ecosystem-scale models, the target polygon as an intended envelope. Review geometry confidence before spatial analysis.<br><br><b>Data source.</b> ${DB.source_workbook}; generated ${DB.generated_on}.</div>`;
}
['search','typeFilter','rankFilter','downloadFilter','qualityFilter','showRegions','colorRegionsByPpr','showArticles'].forEach(id=>document.querySelector('#'+id).addEventListener(id==='search'?'input':'change',()=>{if(id==='qualityFilter')document.querySelector('#qualityVal').textContent=document.querySelector('#qualityFilter').value;refresh()}));
document.querySelector('#kpiArticles').textContent=articles.length;document.querySelector('#kpiRegions').textContent=`${new Set(articles.map(a=>a.unit_id)).size} / ${regions.length}`;document.querySelector('#kpiFiles').textContent=files.filter(f=>String(f.status).includes('verified')).length;document.querySelector('#kpiPpr').textContent='100%';
refresh();
</script>
</body></html>'''


def main():
    p = argparse.ArgumentParser()
    p.add_argument("workbook", nargs="?", default="Ecopath_all84_summary.xlsx")
    p.add_argument("output", nargs="?", default="ecopath_all84_interactive_map.html")
    args = p.parse_args()
    data = read_workbook(Path(args.workbook))
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    Path(args.output).write_text(HTML.replace("__DATA__", payload), encoding="utf-8")
    print(json.dumps({"output": str(Path(args.output).resolve()), "regions": len(data["regions"]), "articles": len(data["articles"]), "files": len(data["files"])}, indent=2))


if __name__ == "__main__":
    main()
