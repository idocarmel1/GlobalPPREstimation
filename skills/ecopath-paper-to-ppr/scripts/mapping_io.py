"""Shared plumbing for the mapping scripts: repo discovery, catch, models, mapping files.

Kept in one place so `prepare_mapping.py` and `validate_mapping.py` cannot drift apart on
what a mapping file is allowed to contain. The parser here is the specification.
"""
from __future__ import annotations

import csv
import gzip
import os
import re
from collections import defaultdict
from pathlib import Path

CONFIDENCE = ("high", "medium", "low", "unresolved")

EVIDENCE = (
    "explicit_member",        # the paper lists this exact taxon in this group
    "synonym",                # accepted-name or historical combination of a listed member
    "taxonomic_containment",  # the group is defined by a taxon that contains this one
    "habitat_size_guild",     # habitat + size + feeding traits select one group
    "inherited_model",        # classification adopted from a cited predecessor model
    "analogue",               # ecological analogue, no direct evidence
    "composite_split",        # the taxon spans several groups and is apportioned
    "inherited_mapping",      # carried over from an earlier mapping, not re-derived here
    "none",                   # unresolved
)

WEIGHT_BASES = ("catch_composition", "model_catch", "model_biomass", "equal")

FIELDS = ["taxon", "common_name", "functional_group", "commercial_group",
          "group", "weights", "confidence", "evidence", "explanation"]

# Catch labels that are not identified to a level any Ecopath model can name.
_COARSE_RE = re.compile(
    r"(not identified|nei\b|miscellaneous|^marine |^other |unidentified)", re.I)

# Endings that mark a rank above genus. A genus is a single capitalised word without
# one of these; a species carries a lowercase epithet.
_RANK_SUFFIX = {
    "iidae": "family", "idae": "family", "inae": "subfamily",
    "iformes": "order", "oidei": "suborder", "oidea": "superfamily",
    "acea": "superfamily", "aceae": "family", "ini": "tribe",
    "opoda": "class", "ozoa": "phylum", "ophyta": "division",
}


def find_root(start: Path | None = None) -> Path:
    """The GlobalPPREstimation checkout: nearest ancestor holding the data spine."""
    env = os.environ.get("GLOBALPPR_ROOT")
    if env and (Path(env) / "data").is_dir():
        return Path(env).resolve()
    here = (start or Path.cwd()).resolve()
    for cand in [here, *here.parents]:
        if (cand / "data" / "INDEX.csv").exists() or (cand / "SeaAroundUsExtraction").is_dir():
            return cand
    raise SystemExit(
        "cannot locate the GlobalPPREstimation checkout. Run from inside it, or set "
        "GLOBALPPR_ROOT to its path.")


def taxon_rank(name: str) -> str:
    """Rough rank of a Sea Around Us catch label. Advisory, not authoritative."""
    n = (name or "").strip()
    if not n:
        return "unknown"
    if _COARSE_RE.search(n):
        return "category"
    parts = n.split()
    if len(parts) >= 2 and parts[1][:1].islower():
        return "species"
    if len(parts) == 1:
        for suf, rank in sorted(_RANK_SUFFIX.items(), key=lambda kv: -len(kv[0])):
            if n.lower().endswith(suf):
                return rank
        return "genus"
    return "category"


def is_coarse(name: str) -> bool:
    """True when the label sits above genus, so it may span several model groups."""
    return taxon_rank(name) not in ("species", "genus")


def unit_from_model_filename(name: str) -> str | None:
    m = re.match(r"^(\d+)HS_", name)
    if m:
        return f"HS_{int(m.group(1)):03d}"
    m = re.match(r"^(\d+)_", name)
    if m:
        return f"LME_{int(m.group(1)):03d}"
    return None


def model_workbooks(root: Path, unit: str) -> list[Path]:
    d = root / "PPREstimation" / "output" / "top10"
    if not d.exists():
        return []
    return sorted(p for p in d.glob("*.xlsx")
                  if not p.name.startswith("~$") and unit_from_model_filename(p.name) == unit)


def read_groups(path: Path) -> list[dict]:
    """`groups_df` of one Ecopath SPPR workbook. The authoritative group names."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "groups_df" not in wb.sheetnames:
        wb.close()
        return []
    rows = list(wb["groups_df"].iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    hdr = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        if d.get("group_name") is None:
            continue
        d["group_name"] = str(d["group_name"])
        out.append(d)
    return sorted(out, key=lambda d: d.get("seq") or 0)


def read_methods(path: Path):
    """`sppr_all`: (method names, group_name -> values). Empty when the sheet is absent."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "sppr_all" not in wb.sheetnames:
        wb.close()
        return [], {}
    rows = list(wb["sppr_all"].iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], {}
    methods = [str(h) for h in rows[0][2:]]
    by_group = {str(r[1]): list(r[2:]) for r in rows[1:] if len(r) > 1 and r[1] is not None}
    return methods, by_group


def read_catch(root: Path, unit: str):
    """(taxa, years). taxa: taxon -> {common_name, functional_group, commercial_group, by_year}."""
    p = root / "SeaAroundUsExtraction" / "data" / "catch_by_taxon_year" / f"{unit}.csv.gz"
    if not p.exists():
        return {}, []
    taxa: dict[str, dict] = {}
    years: set[int] = set()
    with gzip.open(p, "rt", encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            try:
                y = int(r["year"])
                t = float(r["catch_tonnes"] or 0.0)
            except (ValueError, KeyError):
                continue
            years.add(y)
            e = taxa.setdefault(r["taxon"], {
                "common_name": r.get("common_name", "") or "",
                "functional_group": r.get("functional_group", "") or "",
                "commercial_group": r.get("commercial_group", "") or "",
                "by_year": defaultdict(float),
            })
            e["by_year"][y] += t
    return taxa, sorted(years)


def read_trophic_levels(root: Path, unit: str) -> dict:
    for base in ("global_output", "eez_output"):
        p = root / "SeaAroundUsExtraction" / base / "tables" / "regions" / unit / "species.csv"
        if not p.exists():
            continue
        out = {}
        with p.open(encoding="utf-8-sig") as fh:
            for r in csv.DictReader(fh):
                v = (r.get("tl") or "").strip()
                if v:
                    try:
                        out[r["taxon"]] = float(v)
                    except ValueError:
                        pass
        return out
    return {}


def mapping_dir(root: Path, unit: str) -> Path:
    return root / "data" / unit / "mapping"


# Sidecars that live beside a mapping and must never be read as one: the group
# dictionary, the paper's transcribed member list, and the weights the builder resolved.
SIDECAR_SUFFIXES = (".groups.csv", ".members.csv", ".resolved.csv")


def mapping_files(root: Path, unit: str) -> list[Path]:
    d = mapping_dir(root, unit)
    if not d.exists():
        return []
    return sorted(p for p in d.glob("*.csv")
                  if not p.name.endswith(SIDECAR_SUFFIXES))


def parse_groups_cell(value) -> list:
    """`A | B | C` -> ['A', 'B', 'C']. A single name returns a one-element list."""
    return [s.strip() for s in str(value or "").split("|") if s.strip()]


def parse_weights_cell(value, n: int):
    """Return (explicit weights, named basis, error).

    Blank means `catch_composition`. A name from WEIGHT_BASES selects a rule the merge
    step applies. Otherwise the cell must hold `n` non-negative numbers separated by
    `|`, normalised here to sum to 1.
    """
    s = str(value or "").strip()
    if not s:
        return None, "catch_composition", None
    if s.lower() in WEIGHT_BASES:
        return None, s.lower(), None
    parts = [p.strip() for p in s.split("|") if p.strip()]
    if len(parts) != n:
        return None, None, f"{len(parts)} weights for {n} groups"
    try:
        vals = [float(p) for p in parts]
    except ValueError:
        return None, None, f"weights are not numeric: {s!r}"
    if any(v < 0 for v in vals):
        return None, None, "negative weight"
    total = sum(vals)
    if total <= 0:
        return None, None, "weights sum to zero"
    return [v / total for v in vals], None, None


def read_mapping_csv(path: Path) -> list:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def write_mapping_csv(path: Path, rows: list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
