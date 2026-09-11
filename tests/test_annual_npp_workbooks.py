"""Workbook denominators follow the catch year and preserve unavailable years."""
import sys
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import build_ecosystem_data as ecosystem
import build_model_workbook as model
import verify_model_workbook as verifier


class AnnualNppWorkbooks(unittest.TestCase):
    def setUp(self):
        self.years = [1950, 2018, 2019, 2020]
        self.npp = {'ens_median_tC_yr': 999, 'annual': {
            '1950': {'status': 'unavailable', 'reason': 'before satellite record'},
            '2018': {'ens_median_tC_yr': 100, 'npp_vgpm_tC_yr': 100, 'status': 'available', 'provenance': 'source18',
                     'n_models': 1, 'available_models': 'vgpm', 'ensemble_basis': 'available models only'},
            '2019': {'ens_median_tC_yr': 200, 'npp_vgpm_tC_yr': 200, 'status': 'available'},
        }}
        self.taxa = {'Fish': {'taxon': 'Fish', 'common_name': '', 'functional_group': '',
                             'commercial_group': '', 'by_year': dict.fromkeys(self.years, 9.0)}}

    def year_rows(self, ws):
        return {r[0]: r for r in ws.values if isinstance(r[0], int)}

    def test_ecosystem_ratios_use_matching_year_and_missing_stays_blank(self):
        wb = openpyxl.Workbook()
        ecosystem.sheet_summary(wb, 'LME_001', {}, list(self.taxa.values()),
                                self.years, {'Fish': 2}, self.npp, [])
        rows = self.year_rows(wb['Summary'])
        self.assertEqual([r[2] for r in rows.values()], [90] * 4)
        self.assertEqual([r[3] for r in rows.values()], [None, 10, 5, None])

    def test_coverage_counts_available_annual_values_only(self):
        self.assertEqual(ecosystem.npp_coverage(self.npp), {
            'has_npp': True, 'npp_years_available': 2, 'npp_from': 2018, 'npp_to': 2019})
        self.assertEqual(ecosystem.npp_coverage({'annual': {'1950': {'status': 'unavailable'}}}), {
            'has_npp': False, 'npp_years_available': 0, 'npp_from': None, 'npp_to': None})

    def test_model_ratios_match_year_for_both_numerators(self):
        wb = openpyxl.Workbook()
        model.sheet_summary(wb, 'LME_001', 'model1', {}, ['Fish'], self.taxa,
                            self.years, {'Fish': {'names': ['fish']}}, ['new_TE_EEfix'],
                            {'new_TE_EEfix': [180] * 4}, [90] * 4, {'new_TE_EEfix': 'ok'},
                            self.npp, {'Fish': 36}, 36, '')
        rows = self.year_rows(wb['Summary'])
        self.assertEqual([r[4] for r in rows.values()], [None, 10, 5, None])
        self.assertEqual([r[5] for r in rows.values()], [None, 20, 10, None])

    def test_npp_sheet_has_annual_values_and_explicit_missingness(self):
        for builder in (ecosystem, model):
            with self.subTest(builder=builder.__name__):
                wb = openpyxl.Workbook()
                builder.sheet_npp(wb, self.npp, self.years)
                values = list(wb['NPP'].values)
                hdr = next(r for r in values if r[0] == 'year')
                rows = self.year_rows(wb['NPP'])
                self.assertEqual(rows[2018][hdr.index('ens_median_tC_yr')], 100)
                self.assertIsNone(rows[1950][hdr.index('ens_median_tC_yr')])
                self.assertIsNone(rows[2020][hdr.index('ens_median_tC_yr')])
                self.assertEqual(rows[1950][hdr.index('reason')], 'before satellite record')
                self.assertEqual(rows[2018][hdr.index('provenance')], 'source18')
                self.assertEqual(rows[2018][hdr.index('n_models')], 1)
                self.assertEqual(rows[2018][hdr.index('available_models')], 'vgpm')
                self.assertEqual(rows[2018][hdr.index('ensemble_basis')], 'available models only')

    def test_verifier_detects_wrong_year_denominator_and_changed_npp(self):
        wb = openpyxl.Workbook()
        model.sheet_npp(wb, self.npp, self.years)
        model.sheet_summary(wb, 'LME_001', 'model1', {}, ['Fish'], self.taxa,
                            self.years, {'Fish': {'names': ['fish']}}, ['new_TE_EEfix'],
                            {'new_TE_EEfix': [180] * 4}, [90] * 4, {'new_TE_EEfix': 'ok'},
                            self.npp, {'Fish': 36}, 36, '')
        self.assertEqual(verifier.check_annual_npp(wb, self.npp), [])
        row = next(r[0].row for r in wb['Summary'] if r[0].value == 2018)
        wb['Summary'].cell(row, 5, 5)  # copied the 2019 denominator
        self.assertTrue(any('2018' in e for e in verifier.check_annual_npp(wb, self.npp)))
        wb['NPP']['G5'] = 999
        self.assertTrue(any('NPP sheet' in e for e in verifier.check_annual_npp(wb, self.npp)))


if __name__ == '__main__':
    unittest.main()
