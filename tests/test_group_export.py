"""Embedded group inputs retain authoritative identities and source precision."""
import copy
import hashlib
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))


def helpers():
    import group_filter_data
    return group_filter_data


def test_source_group_precision_missingness_and_fixed_mapping_weights():
    groups = [{'group_name': 'α fish', 'ge': .1234567890123456, 'ee': .8765432109876543, 'tl': 3.123456789012345},
              {'group_name': 'Detritus', 'ge': None, 'ee': 0, 'tl': 1},
              {'group_name': 'No metadata', 'ge': float('nan'), 'ee': 1, 'tl': None}]
    scopes = {'all': (['m1', 'm2'], {'α fish': [123.1234567890123, None], 'Detritus': [0, 9]}),
              'inner': (['m2', 'm1'], {'α fish': [6, 7]}),
              'PP': ([], {})}
    mappings = {'mixed': [('α fish', .3333333333333333), ('Detritus', .6666666666666666)]}
    result = helpers().build_group_data(groups, scopes, ['mixed', 'unresolved'], mappings,
        {'source_sha256': 'sourcehash', 'workbook_sha256': 'mappinghash'})
    assert [g['id'] for g in result['groups']] == ['α fish', 'Detritus', 'No metadata']
    assert result['groups'][0]['te'] == groups[0]['ge'] * groups[0]['ee']
    assert result['groups'][0]['tl'] == groups[0]['tl']
    assert result['groups'][1]['te'] is None
    assert result['groups'][2]['te'] is None
    assert result['methods'] == ['m1', 'm2']
    assert result['scopes'] == {'all': [[123.1234567890123, None], [0, 9], [None, None]],
                               'inner': [[7, 6], [None, None], [None, None]],
                               'PP': [[None, None]] * 3}
    assert result['mappings'] == [[[0, .3333333333333333], [1, .6666666666666666]], []]
    assert result['provenance']['source_sha256'] == 'sourcehash'
    assert result['provenance']['workbook_sha256'] == 'mappinghash'
    assert 'GE' in result['te_definition'] and 'EE' in result['te_definition']


@pytest.mark.parametrize('groups,mappings,error', [
    ([{'group_name': 'same'}, {'group_name': 'same'}], {}, 'Duplicate group'),
    ([{'group_name': 'fish'}], {'taxon': [('missing', 1)]}, 'Unknown mapped group'),
    ([{'group_name': 'fish'}], {'taxon': [('fish', float('nan'))]}, 'Invalid mapping weight'),
])
def test_ambiguous_or_invalid_group_inputs_are_rejected(groups, mappings, error):
    with pytest.raises(ValueError, match=error):
        helpers().build_group_data(groups, {}, ['taxon'], mappings)


def network_fixture():
    return {'years': [2000, 2002], 'taxa': ['fish'], 'catch': [[1.123, 2]],
            'full_precision_catch': [[1.123456789, 2]], 'landings': [[1, None]],
            'discards': [[.123456789, None]], 'simple_sppr': [123.123456789],
            'unidentified': {'taxa': [], 'catch': [0, 0]},
            'models': [{'id': 'model', 'verified': True,
                'group_data': {'groups': [{'id': 'fish', 'name': 'fish', 'te': .1, 'tl': 3}],
                               'methods': ['m'], 'scopes': {'all': [[10]]}, 'mappings': [[[0, 1]]]},
                'mc_diagnostics': {'MC_m': {'n_samples': 100, 'n_accepted': 2}},
                'scopes': {'all': {'methods': ['m'], 'status': {'m': 'ok'}, 'values': [[10]]}}}],
            'default_model': 0}


def test_annual_copies_group_inputs_once_with_source_years_and_preserves_originals():
    import build_time_series as annual
    source = network_fixture()
    before = copy.deepcopy(source)
    inputs = helpers().annual_group_inputs(source)
    assert inputs['years'] == [2000, 2002]
    for key in ('taxa', 'catch', 'full_precision_catch', 'landings', 'discards', 'unidentified', 'simple_sppr'):
        assert inputs[key] == source[key]
    models, _ = annual.export_models(source, [1999, 2000, 2001, 2002])
    assert models[0]['group_data'] == source['models'][0]['group_data']
    assert models[0]['taxon_scopes'] == source['models'][0]['scopes']
    assert models[0]['mc_diagnostics'] == source['models'][0]['mc_diagnostics']
    assert 'group_inputs' not in models[0]
    plain = copy.deepcopy(source)
    plain['models'][0].pop('group_data')
    plain_models, _ = annual.export_models(plain, [1999, 2000, 2001, 2002])
    assert models[0]['scopes'] == plain_models[0]['scopes']
    assert helpers().annual_group_inputs(plain) is None
    assert source == before


def test_annual_group_inputs_reject_misaligned_catch():
    source = network_fixture()
    source['landings'] = [[1]]
    with pytest.raises(ValueError, match='dimensions'):
        helpers().annual_group_inputs(source)


def test_real_source_uses_authoritative_groups_scopes_and_exact_final_mappings():
    import build_model_workbook as bmw
    from discard_data import read_final_mappings
    from ppr_scopes import read_scopes
    upstream = next(p for p in bmw.mio.model_workbooks(ROOT, 'LME_013')
                    if (ROOT / 'data/LME_013/models' / p.name).exists())
    mapped = ROOT / 'data/LME_013/models' / upstream.name
    wb = openpyxl.load_workbook(mapped, data_only=True, read_only=True)
    try:
        taxa = [r[0] for r in list(wb['Catch'].values)[1:]]
        mappings = read_final_mappings(wb, upstream.stem)
    finally:
        wb.close()
    groups = bmw.mio.read_groups(upstream)
    scopes = read_scopes(upstream)
    result = helpers().build_group_data(groups, scopes, taxa, mappings)
    for g, exported in zip(groups, result['groups']):
        assert exported['id'] == g['group_name']
        assert exported['tl'] == g['tl']
        assert exported['te'] == (g['ge'] * g['ee'] if g['ge'] is not None and g['ee'] is not None else None)
    for taxon, pairs in zip(taxa, result['mappings']):
        assert [(result['groups'][index]['id'], weight) for index, weight in pairs] == mappings.get(taxon, [])
    for scope, (methods, by_group) in scopes.items():
        for group, values in by_group.items():
            row = result['scopes'][scope][next(i for i, g in enumerate(result['groups']) if g['id'] == group)]
            assert [row[result['methods'].index(method)] for method in methods] == values


def test_network_exports_mc_counts_and_group_provenance_from_source(tmp_path, monkeypatch):
    import json
    import build_network_atlas as network
    unit = 'LME_001'
    model_id = 'model'
    upstream = tmp_path / 'upstream.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'groups_df'
    wb.active.append(['seq', 'group_name', 'ge', 'ee', 'tl'])
    wb.active.append([1, 'fish', .123456789, .87654321, 3.23456789])
    for scope in network.SCOPES:
        ws = wb.create_sheet('sppr_' + scope)
        ws.append(['seq', 'group_name', 'MC_m'])
        ws.append([1, 'fish', 12.34567890123456])
    ws = wb.create_sheet('mc_diagnostics')
    ws.append(['method', 'n_samples', 'n_accepted'])
    ws.append(['MC_m', 100, 2])
    wb.save(upstream)
    mapped = tmp_path / 'data' / unit / 'models' / upstream.name
    mapped.parent.mkdir(parents=True)
    wb = openpyxl.Workbook()
    wb.active.title = 'Catch'
    wb.active.append(['taxon', 'common_name', 'functional_group', 'commercial_group', 2000])
    wb.active.append(['fish', 'Fish', None, None, 1])
    for scope in network.SCOPES:
        sn, pn = ('SPPR', 'PPR by method') if scope == 'all' else ('sppr_' + scope, 'PPR ' + scope)
        ws = wb.create_sheet(sn)
        for _ in range(3):
            ws.append(['title'])
        ws.append(['taxon', 'group', 'simple', 'other', 'MC_m'])
        ws.append(['fish', 'fish', 100, None, 12.345679])
        wb.create_sheet(pn).append(['MC_m', 'ok'])
    ws = wb.create_sheet('Final mappings')
    ws.append(['unit_id', 'model_id', 'taxon', 'group', 'weight'])
    ws.append([unit, upstream.stem, 'fish', 'fish', 1])
    wb.save(mapped)
    (tmp_path / 'data/atlas_selection.json').write_text(json.dumps({'units': {
        unit: {'note': '', 'default_model': model_id}}}), encoding='utf-8')
    monkeypatch.setattr(network, 'load_npp', lambda root: {})
    monkeypatch.setattr(network, 'source_paths', lambda root: [])
    monkeypatch.setattr(network, 'export_simple_units', lambda root: {})
    monkeypatch.setattr(network.bmw.mio, 'model_workbooks', lambda root, unit: [upstream])
    monkeypatch.setattr(network.bmw.mio, 'read_trophic_levels', lambda root, unit: {'fish': 3.23456789})
    monkeypatch.setattr(network.verify, 'check', lambda path, unit: [])
    monkeypatch.setattr(network, 'read_catch_components', lambda *args: {
        'landings': [[1]], 'discards': [[0]], 'full_precision_catch': [[1]]})
    monkeypatch.setattr(network, 'model_sensitivity', lambda *args: ({}, {}))
    payload = network.export_network(tmp_path)
    exported = payload['units'][unit]
    model = exported['models'][0]
    assert model['mc_diagnostics'] == {'MC_m': {'n_samples': 100, 'n_accepted': 2}}
    assert exported['simple_sppr'] == [10 ** (3.23456789 - 1)]
    assert model['group_data']['mappings'] == [[[0, 1]]]
    assert model['group_data']['scopes']['all'] == [[12.34567890123456]]
    assert model['group_data']['provenance']['source_sha256'] == hashlib.sha256(upstream.read_bytes()).hexdigest()
    assert model['group_data']['provenance']['workbook_sha256'] == hashlib.sha256(mapped.read_bytes()).hexdigest()


def test_unverified_models_cannot_supply_annual_group_inputs():
    import build_time_series as annual
    source = network_fixture()
    source['models'][0]['verified'] = False
    models, _ = annual.export_models(source, [2000])
    assert 'group_data' not in models[0]
    assert 'taxon_scopes' not in models[0]
    assert 'mc_diagnostics' not in models[0]
    assert helpers().annual_group_inputs(source) is None
