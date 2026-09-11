"""The release audit catches stale annual data and copied historical denominators."""
import csv
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
from npp_data import METHODS, ANNUAL_PATH
from verify_annual_npp import verify


@pytest.fixture
def release(tmp_path):
    source = tmp_path / ANNUAL_PATH
    source.parent.mkdir(parents=True)
    with source.open('w', newline='', encoding='utf-8') as stream:
        writer = csv.writer(stream)
        writer.writerow(['unit_id', 'year', 'npp_vgpm_tC_yr', 'ens_median_tC_yr'])
        writer.writerows([['LME_034', 2000, 100, 100], ['LME_034', 2001, '', '']])
    sources = [{'path': ANNUAL_PATH, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()}]
    arrays = {m['id']: ([100, None] if m['id'] in ('npp_vgpm_tC_yr', 'ens_median_tC_yr') else [None, None])
              for m in METHODS}
    relative = 'data/LME_034/LME_034.xlsx'
    book = tmp_path / relative
    book.parent.mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'NPP'
    ws.append(['year'] + [m['id'] for m in METHODS])
    for i, year in enumerate([2000, 2001]):
        ws.append([year] + [arrays[m['id']][i] for m in METHODS])
    ws = wb.create_sheet('Summary')
    ws.append(['year', 'catch_tonnes', 'ppr_tonnes', 'ppr_over_npp_percent'])
    ws.append([2000, 90, 900, 100])
    ws.append([2001, 90, 900, None])
    wb.create_sheet('Final mappings')
    wb.save(book)
    wb.close()
    destination = tmp_path / 'PPRAtlas/data'
    destination.mkdir(parents=True)
    (destination / 'network_ppr.json').write_text(json.dumps({
        'npp_sources': sources, 'npp_years': [2000, 2001], 'npp': {'LME_034': arrays}}))
    (destination / 'time_series.json').write_text(json.dumps({
        'sources': sources, 'years': [2000, 2001],
        'units': {'LME_034': {'sources': {'workbook': relative}, 'models': [], 'npp': arrays}}}))
    return tmp_path


def test_release_accepts_matching_annual_inputs(release):
    result = verify(release)
    assert result['central_workbooks'] == 1
    assert result['summary_ratios_checked'] == 2


def test_release_rejects_copied_ratio_and_missing_npp_year(release):
    path = release / 'data/LME_034/LME_034.xlsx'
    wb = openpyxl.load_workbook(path)
    wb['Summary']['D3'] = 100
    wb.save(path)
    with pytest.raises(ValueError, match='missingness differs'):
        verify(release)
    wb['Summary']['D3'] = None
    wb['NPP'].delete_rows(3)
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='year coverage differs'):
        verify(release)


def test_release_rejects_stale_source_snapshot(release):
    with (release / ANNUAL_PATH).open('a', encoding='utf-8') as stream:
        stream.write('\n')
    with pytest.raises(ValueError, match='source hashes differ'):
        verify(release)


def test_workbook_scope_does_not_get_replaced_by_export_units(release):
    (release / 'data/LME_034/LME_034.xlsx').unlink()
    result = verify(release, units=['LME_013'])
    assert result['central_workbooks'] == 0
    assert result['annual_values_checked'] > 0
    assert result['workbook_units'] == ['LME_013']
